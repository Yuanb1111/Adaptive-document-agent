"""Readable editable tables for model-authored literal source comparisons."""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from .presentation_brief import BODY_PT, _copy_height


def render_comparison(presentation, title, item, *, notes=''):
    """Fit complete comparisons by measured column widths, preserving conditions."""
    from .slide_compositor import _base
    from .pptx_export import _source_footer, _text, FOURIER_DARK, FOURIER_MUTED, FOURIER_PURPLE, FOURIER_BG_CARD
    from .text_capacity import wrap_copy
    remaining = list(item.table.rows)
    slides = []
    width = presentation.slide_width.inches - 1.3
    def row_height(row, widths):
        return max(.45, max(len(wrap_copy(cell, w - .18, 14)) for cell, w in zip(row, widths)) * 17 / 72 + .12)
    fractions = [1 / len(item.table.headers), .4, .45, .5, .55]
    choices = []
    for fraction in fractions:
        widths = [width * fraction] + [width * (1 - fraction) / (len(item.table.headers) - 1)] * (len(item.table.headers) - 1)
        heights = [row_height(row, widths) for row in [item.table.headers, *remaining]]
        choices.append((sum(heights), abs(fraction - 1 / len(widths)), widths, heights[0]))
    _, _, widths, header_height = min(choices)
    full = '\n'.join(' | '.join(row) for row in [item.table.headers, *item.table.rows])
    while remaining:
        slide, top = _base(presentation, title + (' (continued)' if slides else ''), '')
        heading_h = _copy_height(item.title, width, 18) + .04
        condition_h = _copy_height(item.conditions, width, 12) + .04 if item.conditions else 0
        body_copy = '' if slides else item.text
        body_pt = BODY_PT
        # Keep a complete small comparison together at readable sizes when its
        # narrative would otherwise repeat across several single-row pages.
        total_rows = sum(row_height(row, widths) for row in remaining)
        for size in (BODY_PT, 15, 14):
            body_h = _copy_height(body_copy, width, size) + .06 if body_copy else 0
            if top + heading_h + body_h + .08 + header_height + total_rows + condition_h <= presentation.slide_height.inches - 1.1:
                body_pt = size
                break
        body_h = _copy_height(body_copy, width, body_pt) + .06 if body_copy else 0
        _text(slide, item.title, .65, top, width, heading_h, size=18, bold=True, color=FOURIER_PURPLE).name = 'brief:heading'
        if body_copy:
            _text(slide, body_copy, .65, top + heading_h, width, body_h, size=body_pt, color=FOURIER_DARK).name = 'brief:body'
        table_top = top + heading_h + body_h + .08
        capacity = presentation.slide_height.inches - 1.1 - table_top - condition_h
        group, heights = [], []
        for row in remaining:
            height = row_height(row, widths)
            if header_height + sum(heights) + height > capacity:
                break
            group.append(row)
            heights.append(height)
        if not group:
            raise ValueError('Brief comparison row and qualifications exceed readable page capacity.')
        shape = slide.shapes.add_table(len(group) + 1, len(item.table.headers), Inches(.65), Inches(table_top),
                                      Inches(width), Inches(header_height + sum(heights)))
        shape.name = 'brief:comparison_table'
        table = shape.table
        for column, column_width in zip(table.columns, widths):
            column.width = Inches(column_width)
        for offset, row in enumerate([item.table.headers, *group]):
            table.rows[offset].height = Inches(header_height if offset == 0 else heights[offset - 1])
            for column, text in enumerate(row):
                cell = table.cell(offset, column)
                cell.text = text
                cell.margin_left = cell.margin_right = Inches(.09)
                cell.margin_top = cell.margin_bottom = Inches(.06)
                cell.text_frame.word_wrap = True
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(FOURIER_PURPLE if offset == 0 else FOURIER_BG_CARD)
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = Pt(14)
                    paragraph.font.bold = offset == 0
                    paragraph.font.color.rgb = RGBColor.from_string('FFFFFF' if offset == 0 else FOURIER_DARK)
                    paragraph.line_spacing = Pt(17)
        if item.conditions:
            _text(slide, item.conditions, .65, table_top + header_height + sum(heights) + .05,
                  width, condition_h, size=12, color=FOURIER_MUTED).name = 'brief:conditions'
        _text(slide, _source_footer(item.pages), .55, presentation.slide_height.inches - .82,
              width, .20, size=9, color=FOURIER_MUTED)
        slide.notes_slide.notes_text_frame.text = '\n\n'.join([notes, item.title, item.text, item.conditions, full, _source_footer(item.pages)])
        slides.append(slide)
        remaining = remaining[len(group):]
    return slides
