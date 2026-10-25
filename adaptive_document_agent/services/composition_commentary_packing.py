"""Pack complete short continuation copy beside charts of the same planned topic."""

import json


def pack_commentary_pages(presentation, rendered, charts, index):
    """Reclaim measured plot whitespace, never discard copy or reduce text size.

    Only one compositor invocation supplies these sibling pages. A candidate
    must retain native chart minimum dimensions, all fixed labels/qualifications,
    the same visible citation and the ordinary geometry validation. If no such
    placement exists, the original continuation remains.
    """
    from pptx.util import Inches
    from .chart_space import chart_space_requirements, composition_plot_height
    from .composition_pagination import _remove_pages
    from .qa_reporter import CriticalQAError
    from .slide_compositor import (
        COMMENTARY_PT, COMMENTARY_LINE_PT, Rect, _lines, _put_commentary,
        validate_composed_geometry,
    )

    by_id = {f'chart:{chart.id}': chart for chart in charts}
    retained = list(rendered)
    width = presentation.slide_width.inches - 1.1

    def citations(page):
        return [s.text for s in page.shapes if s.has_text_frame and s.text.startswith('Source:')]

    for donor in list(rendered):
        if donor.name != 'composed_continuation':
            continue
        if any(s.has_chart or s.has_table or s.shape_type in (6, 13) for s in donor.shapes):
            continue
        # A subtitle can carry a sign convention or qualification; leave that
        # page intact rather than implicitly treating it as a repeated heading.
        placeholders = sorted((s for s in donor.placeholders if s.has_text_frame and s.text.strip()),
                              key=lambda s: s.top)
        if len(placeholders) > 1:
            continue
        bodies = [s for s in donor.shapes if s.has_text_frame and s.text.strip()
                  and not s.is_placeholder and not s.text.startswith('Source:')]
        if len(bodies) != 1 or not bodies[0].name.startswith('composed:text:'):
            continue
        donor_citations = citations(donor)
        if not donor_citations:
            continue
        text = bodies[0].text.replace('\v', '\n')
        lines = len(_lines(text, width - .18, COMMENTARY_PT))
        if lines > 6:
            continue
        candidates = sorted((page for page in retained if page != donor),
                            key=lambda page: abs(retained.index(page) - retained.index(donor)))
        for target in candidates:
            native = [s for s in target.shapes if s.has_chart]
            if not native or any(s.name not in by_id for s in native) or citations(target) != donor_citations:
                continue
            # Exact already-visible paragraphs need not be printed twice.
            # Do not infer semantic equivalence from similar wording.
            visible = {s.text.replace('\v', '\n').strip().removesuffix(' (continued)')
                       for s in target.shapes if s.has_text_frame}
            copy = '\n'.join(part for part in text.split('\n') if part.strip() not in visible)
            if not copy.strip():
                continue
            height = len(_lines(copy, width - .18, COMMENTARY_PT)) * COMMENTARY_LINE_PT / 72 + .12
            gap = .03
            bottom = presentation.slide_height.inches - 1.02
            context = [s.top.inches - gap for s in target.shapes
                       if s.name == 'source_context:definitions_and_limits']
            if context:
                bottom = min(bottom, *context)
            rect = Rect(.55, bottom - height, width, height)
            if rect.y < 1.45:
                continue
            # Do not cover a table, label, annotation, source qualification,
            # picture or existing body. Template background artwork is ignored.
            fixed = [s for s in target.shapes if not s.has_chart and (
                s.has_table or s.shape_type in (6, 13)
                or s.has_text_frame and s.text.strip())]
            if any(s.top.inches < bottom and s.top.inches + s.height.inches > rect.y - gap
                   and s.left.inches < rect.x + rect.w and s.left.inches + s.width.inches > rect.x
                   for s in fixed):
                continue
            sizes = []
            for shape in native:
                new_height = min(shape.height.inches, rect.y - gap - shape.top.inches)
                chart = by_id[shape.name]
                values = [index.get(oid) for oid in chart.observation_ids]
                totals = [index.get(oid) for oid in chart.total_observation_ids]
                if any(value is None for value in [*values, *totals]):
                    break
                min_width, min_height, categories = chart_space_requirements(chart, values, totals)
                bounds = (shape.left.inches, shape.top.inches, shape.width.inches, new_height)
                plot_height = (composition_plot_height(bounds, categories, compact=shape.width.inches < 4.5)
                               if categories else new_height)
                if shape.width.inches + 1e-6 < min_width or plot_height + 1e-6 < min_height:
                    break
                sizes.append((shape, shape.height, Inches(new_height)))
            if len(sizes) != len(native):
                continue
            original_shapes = {s.shape_id for s in target.shapes}
            try:
                for shape, _, new in sizes:
                    shape.height = new
                if _put_commentary(target, copy, rect):
                    raise ValueError('Complete commentary did not fit')
                validate_composed_geometry(presentation)
            except (CriticalQAError, ValueError):
                for shape, old, _ in sizes:
                    shape.height = old
                for shape in list(target.shapes):
                    if shape.shape_id not in original_shapes:
                        shape.element.getparent().remove(shape.element)
                continue
            raw_notes = target.notes_slide.notes_text_frame.text
            try:
                notes = json.loads(raw_notes)
            except ValueError:
                notes = None
            if not isinstance(notes, dict):
                notes = {'original_notes': raw_notes}
            notes.setdefault('packed_commentary_pages', []).append({
                'text': text, 'source_citations': donor_citations,
                'original_notes': donor.notes_slide.notes_text_frame.text,
                'original_title': [s.text for s in donor.placeholders if s.has_text_frame],
            })
            target.notes_slide.notes_text_frame.text = json.dumps(notes, ensure_ascii=False)
            _remove_pages(presentation, [donor])
            retained.remove(donor)
            break
    return retained
