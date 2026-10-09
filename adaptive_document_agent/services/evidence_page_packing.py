"""Pack evidence tables without joining their periods, values or provenance."""

from copy import deepcopy
import re

from pptx.util import Inches

from .text_capacity import wrap_copy


def _tables(slide):
    return [shape for shape in slide.shapes
            if shape.has_table and shape.name.startswith("evidence:packable")]


def _note_spec(previous, slide, name):
    shapes = [shape for page in (previous, slide) for shape in page.shapes
              if shape.name == name]
    if not shapes:
        return None
    texts = list(dict.fromkeys(paragraph.text for shape in shapes
                              for paragraph in shape.text_frame.paragraphs if paragraph.text))
    if name == 'evidence:footer':
        from .pptx_export import _source_footer
        pattern = r'^Source: Document disclosures \(p\. ([\d, -]+)\)(.*)$'
        matches = [re.fullmatch(pattern, text) for text in texts]
        if matches and all(matches):
            pages = set()
            tails = []
            for match in matches:
                for span in match[1].split(','):
                    bounds = [int(part) for part in span.strip().split('-')]
                    pages.update(range(bounds[0], bounds[-1] + 1))
                tails.extend(part.strip() for part in match[2].split('|') if part.strip())
            texts = [_source_footer(sorted(pages)) + ''.join(' | ' + tail for tail in dict.fromkeys(tails))]
    template = shapes[0]
    size = template.text_frame.paragraphs[0].font.size
    size = size.pt if size is not None else 9.0
    lines = sum(len(wrap_copy(text, template.width.inches - .10, size)) for text in texts)
    height = max(.22, lines * max(.20, size * 1.3 / 72) + .02)
    return template, texts, height


def _compact_row_height(row, columns):
    """Remove only measured vertical whitespace; never reduce the font size."""
    required = .31
    for cell, column in zip(row.cells, columns):
        height = cell.margin_top.inches + cell.margin_bottom.inches + .04
        width = column.width.inches - cell.margin_left.inches - cell.margin_right.inches
        for paragraph in cell.text_frame.paragraphs:
            if not paragraph.text:
                continue
            sizes = [run.font.size or paragraph.font.size for run in paragraph.runs]
            if not sizes or any(size is None for size in sizes):
                # Inherited fonts cannot be measured reliably here.
                return row.height.inches
            size = max(value.pt for value in sizes)
            spacing = paragraph.line_spacing
            line_height = (spacing.pt / 72 if hasattr(spacing, "pt") else
                           size * max(1.25, spacing or 1.0) / 72)
            height += len(wrap_copy(paragraph.text, width, size)) * line_height
            height += sum(value.pt / 72 for value in
                          (paragraph.space_before, paragraph.space_after) if value is not None)
        required = max(required, height)
    return required


def _copy_shape(slide, shape):
    element = deepcopy(shape._element)
    for prop in element.xpath(".//p:cNvPr"):
        prop.set("id", str(slide.shapes._next_shape_id))
    slide.shapes._spTree.insert_element_before(element, "p:extLst")
    return slide.shapes[-1]


def _write_note(slide, name, spec, top):
    if spec is None:
        return
    template, texts, height = spec
    shape = next((item for item in slide.shapes if item.name == name), None)
    if shape is None:
        shape = _copy_shape(slide, template)
    properties = deepcopy(template.text_frame.paragraphs[0]._p.pPr)
    shape.text_frame.clear()
    for index, text in enumerate(texts):
        paragraph = shape.text_frame.paragraphs[0] if index == 0 else shape.text_frame.add_paragraph()
        paragraph.text = text
        if properties is not None:
            old_properties = paragraph._p.pPr
            if old_properties is not None:
                paragraph._p.remove(old_properties)
            paragraph._p.insert(0, deepcopy(properties))
    shape.top, shape.height = Inches(top), Inches(height)


def pack_evidence_pages(presentation, slides):
    """Combine adjacent, explicitly marked tables when their content fits.

    Each table retains its own complete headers and rows, even when periods or
    units differ. Source paragraphs and raw-record speaker notes are retained.
    Compact rows only when needed, using the existing fonts and cell margins.
    """
    retained = []
    for slide in slides:
        tables = _tables(slide)
        previous = retained[-1] if retained else None
        old_tables = _tables(previous) if previous is not None else []
        if not tables or not old_tables:
            retained.append(slide)
            continue
        combined = old_tables + tables
        footer = _note_spec(previous, slide, "evidence:footer")
        convention = _note_spec(previous, slide, "evidence:convention")
        note_bottom = min(6.47, presentation.slide_height.inches - .60)
        footer_top = note_bottom - (footer[2] if footer else 0)
        convention_top = footer_top - (.08 + convention[2] if convention else 0)
        bottom = min(presentation.slide_height.inches - 1.02, convention_top - .12)
        top = min(shape.top.inches for shape in old_tables)
        gap = .18
        heights = [[row.height.inches for row in shape.table.rows] for shape in combined]
        needed = sum(sum(rows) for rows in heights) + gap * (len(combined) - 1)
        if top + needed > bottom:
            heights = [[_compact_row_height(row, shape.table.columns)
                        for row in shape.table.rows] for shape in combined]
            needed = sum(sum(rows) for rows in heights) + gap * (len(combined) - 1)
        if top + needed > bottom:
            # A source page may contain several independent complete tables.
            # Move the largest fitting prefix instead of leaving a sparse page
            # just because the whole next page cannot fit. Never split a table.
            for count in range(len(tables) - 1, 0, -1):
                proposed = old_tables + tables[:count]
                measured = [[_compact_row_height(row, shape.table.columns)
                             for row in shape.table.rows] for shape in proposed]
                if top + sum(map(sum, measured)) + gap * (len(proposed) - 1) <= bottom:
                    combined, heights = proposed, measured
                    break
            else:
                retained.append(slide)
                continue
        moved_count = len(combined) - len(old_tables)
        moved = tables[:moved_count]
        moved_tables = old_tables + [_copy_shape(previous, shape) for shape in moved]
        for shape, rows in zip(moved_tables, heights):
            shape.name = f"evidence:packable:{shape.shape_id}"
            shape.top = Inches(top)
            for row, height in zip(shape.table.rows, rows):
                row.height = Inches(height)
            top += shape.height.inches + gap
        _write_note(previous, "evidence:footer", footer, footer_top)
        _write_note(previous, "evidence:convention", convention, convention_top)
        previous.notes_slide.notes_text_frame.text += "\n\n" + slide.notes_slide.notes_text_frame.text
        if moved_count < len(tables):
            new_top = min(shape.top.inches for shape in tables)
            for shape in moved:
                shape._element.getparent().remove(shape._element)
            for shape in tables[moved_count:]:
                shape.top = Inches(new_top)
                new_top += shape.height.inches + gap
            retained.append(slide)
            continue
        for slide_id in list(presentation.slides._sldIdLst):
            if presentation.part.related_slide(slide_id.rId) is slide:
                presentation.part.drop_rel(slide_id.rId)
                presentation.slides._sldIdLst.remove(slide_id)
                break
    # Removing a middle page leaves a hole; renumber before appending slides so
    # python-pptx cannot reuse an existing slide part's filename.
    presentation.part.rename_slide_parts(s.rId for s in presentation.slides._sldIdLst)
    return retained
