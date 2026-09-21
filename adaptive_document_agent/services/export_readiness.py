"""Session-local, document-free smoke test before expensive analysis starts."""

import io

from .presentation_rendering import RenderingError, configured_renderer


def check_export_readiness(cache: dict | None = None) -> dict:
    from .presentation_visual_qa import POLICY_VERSION, VisualQAError, verify_presentation
    try:
        renderer = configured_renderer()
        key = (POLICY_VERSION, renderer.identity)
        if cache is not None and key in cache:
            return dict(cache[key])
        # Exercise the actual Impress import, native chart rendering and network
        # policy, not just executable existence. Contains no user document data.
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(.5), Inches(.3), Inches(8), Inches(.5)).text = "Renderer readiness check"
        data = CategoryChartData()
        data.categories = ["A", "B"]
        data.add_series("Check", [1, 2])
        chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.3), Inches(7), Inches(4), data)
        for tag in ("axId", "crossAx"):
            for node in chart_shape.chart._chartSpace.xpath(f".//c:{tag}"):
                node.set("val", str(int(node.get("val")) % (2**32)))
        stream = io.BytesIO()
        deck.save(stream)
        verify_presentation(stream.getvalue(), renderer=renderer, max_repairs=0)
        result = {"ready": True, "backend": type(renderer).__name__, "message": "PowerPoint renderer passed its document-free startup check."}
        if cache is not None:
            cache.clear()
            cache[key] = dict(result)
        return result
    except (RenderingError, VisualQAError) as exc:
        return {"ready": False, "backend": "unavailable", "message": str(exc)}
    except (OSError, ValueError, TypeError) as exc:
        return {"ready": False, "backend": "unavailable", "message": f"PowerPoint startup check could not complete ({type(exc).__name__})."}
