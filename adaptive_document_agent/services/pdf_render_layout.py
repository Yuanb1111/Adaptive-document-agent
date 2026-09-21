"""Bridge real PDF pixels/text and native PPT geometry, with explicit provenance.

LibreOffice's PDF export does not expose PPT object IDs. Bounds/IDs below come
from OOXML; visible text is checked against PDF extraction in the same region.
Never claim that source objects alone prove they were rendered.
"""

import io
import re
import unicodedata

import fitz
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .presentation_rendering import RenderedPage, RenderingError


def _normalized(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    return re.sub(r"\s+", "", text).replace("\u00ad", "")


def _shape_elements(shapes, pdf_page, *, sx, sy, inherited=False):
    elements = []
    for shape in shapes:
        if inherited and shape.is_placeholder:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            # Transformed group coordinates require a dedicated adapter. Do not
            # certify a group using incorrect, untransformed child coordinates.
            raise RenderingError("LibreOffice QA does not yet support grouped slide objects.")
        x, y, w, h = float(shape.left)*sx, float(shape.top)*sy, float(shape.width)*sx, float(shape.height)*sy
        text = shape.text if shape.has_text_frame else ""
        kind = "chart" if shape.has_chart else "table" if shape.has_table else "image" if shape.shape_type == MSO_SHAPE_TYPE.PICTURE else "shape"
        if not text.strip() and kind == "shape":
            continue
        element = {"id": str(shape.shape_id), "name": shape.name, "kind": kind,
                   "bbox": [x, y, w, h], "text": text, "boundsSource": "ooxml"}
        if text.strip() and x >= 0 and y >= 0 and x+w <= pdf_page.rect.width*4/3+2 and y+h <= pdf_page.rect.height*4/3+2:
            # PDF coordinates are points; the common QA layout uses 96dpi px.
            region = fitz.Rect(x*.75-2, y*.75-2, (x+w)*.75+2, (y+h)*.75+2)
            rendered_text = pdf_page.get_textbox(region)
            # Check each paragraph separately: PDF reading order can interleave
            # parallel columns even when their individual text is complete.
            normalized = _normalized(rendered_text)
            paragraphs = [_normalized(p.text) for p in shape.text_frame.paragraphs if p.text.strip()]
            element["renderedTextPresent"] = all(p in normalized for p in paragraphs)
        elements.append(element)
    return elements


def pdf_rendered_pages(payload: bytes, pdf_path) -> list[RenderedPage]:
    deck = Presentation(io.BytesIO(payload))
    pages = []
    total_bytes = 0
    try:
        with fitz.open(pdf_path) as pdf:
            if len(pdf) != len(deck.slides) or not 1 <= len(pdf) <= 150:
                raise RenderingError("LibreOffice omitted or added slides during PDF rendering.")
            for slide, page in zip(deck.slides, pdf):
                width, height = page.rect.width*4/3, page.rect.height*4/3
                if min(width, height) < 100 or width*height > 8_000_000:
                    raise RenderingError("LibreOffice page exceeds rendering dimensions limits.")
                if abs(width/height/(deck.slide_width/deck.slide_height)-1) > .015:
                    raise RenderingError("LibreOffice changed the slide aspect ratio.")
                # Source coordinate scale is independent of the rendered frame:
                # preserve OOXML's true 96dpi size for out-of-bounds checks.
                sx = sy = 96/914400
                layout = {"unit": "px", "slide": {"frame": {"width": width, "height": height}},
                    "elements": _shape_elements(slide.shapes, page, sx=sx, sy=sy),
                    "inheritedLayers": [], "geometrySource": "ooxml", "pixelsSource": "libreoffice-pdf"}
                for scope, native in (("layout", slide.slide_layout), ("master", slide.slide_layout.slide_master)):
                    layout["inheritedLayers"].append({"scope": scope,
                        "elements": _shape_elements(native.shapes, page, sx=sx, sy=sy, inherited=True)})
                png = page.get_pixmap(matrix=fitz.Matrix(4/3, 4/3), alpha=False).tobytes("png")
                total_bytes += len(png)
                if total_bytes > 300_000_000:
                    raise RenderingError("Rendered output exceeds the 300 MB validation limit.")
                pages.append(RenderedPage(png, layout))
    except RenderingError:
        raise
    except (RuntimeError, ValueError) as exc:
        raise RenderingError("LibreOffice produced an unreadable PDF.") from exc
    return pages
