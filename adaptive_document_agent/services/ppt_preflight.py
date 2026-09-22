"""Pre-export presentation quality control, collision detection, and preflight sanitization."""

from __future__ import annotations

from collections import Counter
import re
from typing import Any

# Forbidden internal/debug phrases that must never appear in client-facing decks
BANNED_PHRASES = (
    "calculated result",
    "calculated change result",
    "based on extracted observations",
    "this is a calculated result",
    "internal calculation",
    "debug",
    "parser",
    "extracted observations",
    "generated metric",
    "retained fact",
    "ai says",
    "llm",
    "json",
)


class PreflightIssue:
    def __init__(self, slide_idx: int, check: str, message: str, severity: str = "warning"):
        self.slide_idx = slide_idx
        self.check = check
        self.code = check
        self.message = message
        self.severity = severity

    def __repr__(self) -> str:
        return f"Slide {self.slide_idx + 1} [{self.check}] ({self.severity}): {self.message}"


class PresentationPreflight:
    """Automated pre-export validation for PowerPoint presentations."""

    def __init__(self, presentation: Any):
        self.presentation = presentation
        self.issues: list[PreflightIssue] = []

    def validate_and_sanitize(self) -> list[PreflightIssue]:
        self.issues.clear()
        for idx, slide in enumerate(self.presentation.slides):
            self._check_geometry_invariants(idx, slide)
            self._check_chart_gridlines(idx, slide)
            self._check_chart_quality(idx, slide)
            self._check_banned_phrases(idx, slide)
            self._check_raw_unit_tokens(idx, slide)
            self._check_layout_overflow_and_overlap(idx, slide)
            self._check_unit_consistency(idx, slide)
            self._check_impossible_percentages(idx, slide)
            self._check_title_data_alignment(idx, slide)
            self._check_truncated_fields(idx, slide)
            self._check_empty_slides(idx, slide)
            self._check_template_completeness(idx, slide)
            self._check_duplicate_card_titles(idx, slide)
            self._check_slide_title_quality(idx, slide)
        self._check_thank_you_slide()
        return self.issues

    def _check_geometry_invariants(self, idx: int, slide: Any) -> None:
        from pptx.util import Inches
        for shape in slide.shapes:
            try:
                # Title placeholders must have generous horizontal width
                if getattr(shape, "is_placeholder", False) and shape.placeholder_format.idx in (14, 15):
                    if shape.width.inches < 6.0:
                        self.issues.append(
                            PreflightIssue(
                                idx,
                                "title_geometry_collapsed",
                                f"Title placeholder width collapsed to {shape.width.inches:.2f}in; restoring to 9.06in",
                                severity="warning",
                            )
                        )
                        shape.width = Inches(9.06)
                if shape.has_text_frame and shape.text.strip():
                    if shape.width.inches <= 0.05:
                        shape.width = Inches(4.0)
                        self.issues.append(PreflightIssue(idx, "shape_zero_width", f"Shape '{shape.name}' had zero width; restored", severity="warning"))
                    if shape.height.inches <= 0.05:
                        shape.height = Inches(0.4)
                        self.issues.append(PreflightIssue(idx, "shape_zero_height", f"Shape '{shape.name}' had zero height; restored", severity="warning"))
            except Exception:
                pass

    def _check_chart_gridlines(self, idx: int, slide: Any) -> None:
        for shape in slide.shapes:
            if shape.has_chart:
                chart = shape.chart
                # Enforce no horizontal background gridlines
                try:
                    if hasattr(chart, "value_axis") and chart.value_axis is not None:
                        if chart.value_axis.has_major_gridlines:
                            chart.value_axis.has_major_gridlines = False
                            self.issues.append(PreflightIssue(idx, "chart_gridlines", "Disabled horizontal major gridlines on chart"))
                        if chart.value_axis.has_minor_gridlines:
                            chart.value_axis.has_minor_gridlines = False
                except Exception:
                    pass

    def _check_chart_quality(self, idx: int, slide: Any) -> None:
        for shape in slide.shapes:
            if shape.has_chart:
                chart = shape.chart
                # Reject "Observed pairs"
                try:
                    for series in getattr(chart, "series", []):
                        if series.name and "observed pairs" in series.name.casefold():
                            self.issues.append(
                                PreflightIssue(
                                    idx,
                                    "uninterpretable_chart_series",
                                    f"Chart contains uninterpretable series name '{series.name}'",
                                    severity="warning",
                                )
                            )
                            series.name = "Metric Analysis"
                except Exception:
                    pass

                # Check category axis position for negative bars
                try:
                    if hasattr(chart, "category_axis") and chart.category_axis is not None:
                        from pptx.enum.chart import XL_TICK_LABEL_POSITION
                        has_negative = False
                        for series in getattr(chart, "series", []):
                            if any(v is not None and v < 0 for v in getattr(series, "values", [])):
                                has_negative = True
                                break
                        if has_negative and chart.category_axis.tick_label_position != XL_TICK_LABEL_POSITION.LOW:
                            chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
                            self.issues.append(
                                PreflightIssue(
                                    idx,
                                    "negative_bar_axis_repositioned",
                                    "Repositioned category axis to LOW for negative vertical bars",
                                    severity="warning",
                                )
                            )
                except Exception:
                    pass

    def _check_raw_unit_tokens(self, idx: int, slide: Any) -> None:
        raw_token_pattern = re.compile(r"(?i)\b(rmb|cny|hkd|usd)(?:in)?(thousands?|millions?|billions?|'000)\b")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if raw_token_pattern.search(p.text):
                        orig = p.text
                        clean = re.sub(r"(?i)\b(rmb|cny|hkd|usd)(?:in)?(thousands?|'000)\b", r"\1 '000", orig)
                        clean = re.sub(r"(?i)\b(rmb|cny|hkd|usd)(?:in)?(millions?)\b", r"\1 million", clean)
                        clean = re.sub(r"(?i)\b(rmb|cny|hkd|usd)(?:in)?(billions?)\b", r"\1 billion", clean)
                        if clean != orig:
                            self.issues.append(
                                PreflightIssue(
                                    idx,
                                    "raw_unit_token_sanitized",
                                    f"Sanitized raw unit token in text: '{orig}' -> '{clean}'",
                                    severity="warning",
                                )
                            )
                            p.text = clean
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            if raw_token_pattern.search(p.text):
                                orig = p.text
                                clean = re.sub(r"(?i)\b(rmb|cny|hkd|usd)(?:in)?(thousands?|'000)\b", r"\1 '000", orig)
                                clean = re.sub(r"(?i)\b(rmb|cny|hkd|usd)(?:in)?(millions?)\b", r"\1 million", clean)
                                clean = re.sub(r"(?i)\b(rmb|cny|hkd|usd)(?:in)?(billions?)\b", r"\1 billion", clean)
                                if clean != orig:
                                    p.text = clean


    def _check_banned_phrases(self, idx: int, slide: Any) -> None:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text
                text_lower = text.casefold()
                for phrase in BANNED_PHRASES:
                    if phrase in text_lower:
                        self.issues.append(PreflightIssue(idx, "banned_phrase", f"Found banned phrase '{phrase}' in text"))
                        # Sanitize in place if possible
                        self._sanitize_text_frame(shape.text_frame, phrase)

    @staticmethod
    def _sanitize_text_frame(frame: Any, phrase: str) -> None:
        pattern = re.compile(re.escape(phrase), re.IGNORECASE)
        for p in frame.paragraphs:
            if pattern.search(p.text):
                # Clean professional rephrasing
                if phrase in ("calculated result", "calculated change result", "this is a calculated result"):
                    p.text = pattern.sub("Analysis indicates", p.text)
                elif phrase in ("based on extracted observations", "extracted observations"):
                    p.text = pattern.sub("based on reported disclosures", p.text)
                elif phrase in ("retained fact", "retained facts"):
                    p.text = pattern.sub("reported data", p.text)
                else:
                    p.text = pattern.sub("", p.text)
                p.text = re.sub(r"[ \t]{2,}", " ", p.text).strip()

    def _check_collisions(self, idx: int, slide: Any) -> None:
        layout_name = slide.slide_layout.name if hasattr(slide, "slide_layout") else ""
        if "封面" in layout_name or "短文本" in layout_name or idx == 0:
            return

        title_ph = None
        sub_ph = None
        for shape in slide.placeholders:
            try:
                if shape.placeholder_format.idx in (14, 15):
                    title_ph = shape
                elif shape.placeholder_format.idx == 16:
                    sub_ph = shape
            except Exception:
                pass

        # 1. Check and correct title-subtitle overlap
        if title_ph and sub_ph and sub_ph.has_text_frame and sub_ph.text.strip():
            title_bottom = title_ph.top.inches + title_ph.height.inches
            if sub_ph.top.inches < title_bottom - 0.02:
                self.issues.append(
                    PreflightIssue(
                        idx,
                        "title_subtitle_collision",
                        f"Subtitle at top={sub_ph.top.inches:.2f}in collides with title ending at {title_bottom:.2f}in",
                        severity="warning",
                    )
                )
                from pptx.util import Inches
                sub_ph.top = Inches(title_bottom + 0.08)

        # 2. Compute dynamic content boundary based on subtitle/title bottom
        content_min_top = 1.38
        if sub_ph and sub_ph.has_text_frame and sub_ph.text.strip():
            content_min_top = max(content_min_top, sub_ph.top.inches + sub_ph.height.inches + 0.06)
        elif title_ph and title_ph.has_text_frame and title_ph.text.strip():
            content_min_top = max(content_min_top, title_ph.top.inches + title_ph.height.inches + 0.06)

        for shape in slide.shapes:
            if (shape.name == "decoration:title_rule" and shape.height.inches <= .05
                    and not (shape.has_text_frame and shape.text.strip())):
                continue  # Owned underline belongs inside the heading zone.
            if getattr(shape, "is_placeholder", False):
                try:
                    if shape.placeholder_format.idx in (14, 15, 16):
                        continue
                except Exception:
                    pass
            top = shape.top.inches if hasattr(shape, "top") else 0
            if 0.2 < top < content_min_top - 0.05:
                self.issues.append(
                    PreflightIssue(
                        idx,
                        "title_collision",
                        f"Shape '{shape.name}' at top={top:.2f}in encroaches on title/subtitle zone (min={content_min_top:.2f}in)",
                        severity="warning",
                    )
                )

    def _check_semantic_units(self, idx: int, slide: Any) -> None:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text
                # Multiple metrics mistakenly formatted with %
                for match in re.finditer(r"(?i)\b(current ratio|quick ratio)\s*[:=]?\s*([0-9]+(?:\.[0-9]+)?)\s*%", text):
                    self.issues.append(
                        PreflightIssue(
                            idx,
                            "invalid_ratio_unit",
                            f"Ratio metric formatted with percentage: '{match.group(0)}'",
                            severity="warning",
                        )
                    )
                    ratio_val = match.group(2)
                    new_val = f"{match.group(1)}: {ratio_val}x"
                    for p in shape.text_frame.paragraphs:
                        if match.group(0) in p.text:
                            p.text = p.text.replace(match.group(0), new_val)

                # Multiple metrics with pp change (e.g. +0.2 pp)
                for match in re.finditer(r"(?i)\b(current ratio|quick ratio)[^0-9\n]*([+\-][0-9]+(?:\.[0-9]+)?)\s*pp\b", text):
                    self.issues.append(
                        PreflightIssue(
                            idx,
                            "invalid_ratio_change",
                            f"Ratio change formatted with pp: '{match.group(0)}'",
                            severity="warning",
                        )
                    )
                    for p in shape.text_frame.paragraphs:
                        if match.group(0) in p.text:
                            p.text = p.text.replace(f"{match.group(2)} pp", f"{match.group(2)}x")

                # Days metrics mistakenly formatted with pp (e.g. +0.9 pp)
                for match in re.finditer(r"(?i)\b(turnover\s+days|inventory\s+days|receivables\s+days|payables\s+days|cash\s+conversion\s+cycle|days)[^0-9\n]*([+\-]?[0-9]+(?:\.[0-9]+)?)\s*pp\b", text):
                    self.issues.append(
                        PreflightIssue(
                            idx,
                            "invalid_days_change",
                            f"Days metric change formatted with pp: '{match.group(0)}'",
                            severity="warning",
                        )
                    )
                    for p in shape.text_frame.paragraphs:
                        if match.group(0) in p.text:
                            p.text = p.text.replace(f"{match.group(2)} pp", f"{match.group(2)} days")

                # Days metrics mistakenly formatted with % (e.g. 30.5%)
                for match in re.finditer(r"(?i)\b(turnover\s+days|inventory\s+days|receivables\s+days|payables\s+days)\s*[:=]?\s*([0-9]+(?:\.[0-9]+)?)\s*%", text):
                    self.issues.append(
                        PreflightIssue(
                            idx,
                            "invalid_days_unit",
                            f"Days metric formatted with percentage: '{match.group(0)}'",
                            severity="warning",
                        )
                    )
                    new_val = f"{match.group(1)}: {match.group(2)} days"
                    for p in shape.text_frame.paragraphs:
                        if match.group(0) in p.text:
                            p.text = p.text.replace(match.group(0), new_val)

    def _check_unit_consistency(self, idx: int, slide: Any) -> None:
        """Check and repair unit-family mismatches in client-facing text."""
        self._check_semantic_units(idx, slide)

    def _check_impossible_percentages(self, idx: int, slide: Any) -> None:
        monetary_terms = r"(?:revenue|sales|turnover(?! days)|gross profit|net profit|operating profit|ebitda|operating cash flow|total assets|total liabilities|total equity|cash and cash equivalents|cost of sales|capex)"
        pattern = re.compile(rf"(?i)\b({monetary_terms})\s*[:=]?\s*([0-9,]{{3,}}(?:\.[0-9]+)?)\s*%")
        extreme_pattern = re.compile(r"(?<![\w.])([0-9][0-9,]*(?:\.[0-9]+)?)\s*%")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text
                monetary_spans = [match.span() for match in pattern.finditer(text)]
                for match in pattern.finditer(text):
                    orig = match.group(0)
                    fixed = f"{match.group(1)}: {match.group(2)}"
                    self.issues.append(
                        PreflightIssue(
                            idx,
                            "impossible_percentage",
                            f"Monetary metric formatted as percentage: '{orig}'",
                            severity="error",
                        )
                    )
                    for p in shape.text_frame.paragraphs:
                        if orig in p.text:
                            p.text = p.text.replace(orig, fixed)
                for match in extreme_pattern.finditer(text):
                    if any(match.start() < end and match.end() > start for start, end in monetary_spans):
                        continue
                    try:
                        value = float(match.group(1).replace(",", ""))
                    except ValueError:
                        continue
                    if value > 1000:
                        self.issues.append(
                            PreflightIssue(
                                idx,
                                "impossible_percentage",
                                f"Percentage value exceeds 1,000%: '{match.group(0)}'",
                                severity="error",
                            )
                        )
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if pattern.search(cell.text):
                            for match in pattern.finditer(cell.text):
                                orig = match.group(0)
                                fixed = f"{match.group(1)}: {match.group(2)}"
                                self.issues.append(
                                    PreflightIssue(
                                        idx,
                                        "impossible_percentage",
                                        f"Monetary metric in table cell formatted as percentage: '{orig}'",
                                        severity="error",
                                    )
                                )
                                for p in cell.text_frame.paragraphs:
                                    if orig in p.text:
                                        p.text = p.text.replace(orig, fixed)
                        for match in extreme_pattern.finditer(cell.text):
                            try:
                                value = float(match.group(1).replace(",", ""))
                            except ValueError:
                                continue
                            if value > 1000:
                                self.issues.append(
                                    PreflightIssue(
                                        idx,
                                        "impossible_percentage",
                                        f"Percentage value in table exceeds 1,000%: '{match.group(0)}'",
                                        severity="error",
                                    )
                                )

    def _check_truncated_fields(self, idx: int, slide: Any) -> None:
        dangling_pattern = re.compile(r"(?i)\b(?:to|of|and|with|from|in|for|by|as|at|or|including|such\s+as)\s*$")
        broken_prefix_pattern = re.compile(r"^(?:[a-z]|ing|ed|tion|ment|ly|al|ic)\s+[a-z]{3,}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip().rstrip(".,;:-–—")
                    if broken_prefix_pattern.search(t):
                        self.issues.append(
                            PreflightIssue(
                                idx,
                                "truncated_text_fragment",
                                f"Text fragment appears to begin with a broken word: '{p.text}'",
                                severity="error",
                            )
                        )
                    if dangling_pattern.search(t):
                        orig = p.text
                        clean = dangling_pattern.sub("", t).strip(" .,;:-–—")
                        self.issues.append(
                            PreflightIssue(
                                idx,
                                "truncated_text_fragment",
                                f"Text fragment ending in dangling connector: '{orig}'",
                                severity="warning",
                            )
                        )
                        p.text = clean

    def _check_empty_slide(self, idx: int, slide: Any) -> None:
        text_content: list[str] = []
        has_visual = False
        for shape in slide.shapes:
            if getattr(shape, "has_chart", False) or getattr(shape, "has_table", False):
                has_visual = True
            if shape.has_text_frame and shape.text.strip():
                text_content.append(shape.text.strip())
        joined = " ".join(text_content).strip()
        is_cover = idx == 0
        is_closing = "thank you" in joined.casefold()
        if not is_cover and not is_closing and not has_visual and len(joined) < 20:
            self.issues.append(
                PreflightIssue(
                    idx,
                    "empty_slide",
                    f"Slide {idx + 1} has insufficient substantive content ({len(joined)} text characters and no table/chart)",
                    severity="error",
                )
            )

    def _check_empty_slides(self, idx: int, slide: Any) -> None:
        """Named entry point matching the preflight specification."""
        self._check_empty_slide(idx, slide)

    def _check_content_overflow(self, idx: int, slide: Any) -> None:
        slide_h = 7.5
        slide_w = 13.333
        for shape in slide.shapes:
            try:
                top = shape.top.inches if hasattr(shape, "top") else 0
                height = shape.height.inches if hasattr(shape, "height") else 0
                left = shape.left.inches if hasattr(shape, "left") else 0
                width = shape.width.inches if hasattr(shape, "width") else 0
                if top + height > slide_h + 0.15:
                    self.issues.append(
                        PreflightIssue(
                            idx,
                            "vertical_overflow",
                            f"Shape '{shape.name}' overflows bottom edge (bottom={top + height:.2f}in > {slide_h}in)",
                            severity="warning",
                        )
                    )
                if left + width > slide_w + 0.15:
                    self.issues.append(
                        PreflightIssue(
                            idx,
                            "horizontal_overflow",
                            f"Shape '{shape.name}' overflows right edge (right={left + width:.2f}in > {slide_w}in)",
                            severity="warning",
                        )
                    )
            except Exception:
                pass

    def _check_layout_overflow_and_overlap(self, idx: int, slide: Any) -> None:
        """Check page bounds, header/footer encroachment, and visual-block overlap."""
        self._check_collisions(idx, slide)
        self._check_content_overflow(idx, slide)

        visual_shapes = [
            shape
            for shape in slide.shapes
            if getattr(shape, "has_chart", False) or getattr(shape, "has_table", False)
        ]
        for left_index, first in enumerate(visual_shapes):
            for second in visual_shapes[left_index + 1 :]:
                try:
                    left = max(first.left, second.left)
                    top = max(first.top, second.top)
                    right = min(first.left + first.width, second.left + second.width)
                    bottom = min(first.top + first.height, second.top + second.height)
                    if right <= left or bottom <= top:
                        continue
                    overlap_area = (right - left) * (bottom - top)
                    smaller_area = min(first.width * first.height, second.width * second.height)
                    if smaller_area and overlap_area / smaller_area >= 0.03:
                        self.issues.append(
                            PreflightIssue(
                                idx,
                                "visual_block_overlap",
                                f"Visual blocks '{first.name}' and '{second.name}' overlap materially",
                                severity="error",
                            )
                        )
                except Exception:
                    continue

    def _check_title_data_alignment(self, idx: int, slide: Any) -> None:
        """Flag directional titles that contradict a single linked chart series."""
        positive_words = re.compile(r"(?i)\b(?:grew|rose|increased|expanded|improved|rebounded)\b")
        negative_words = re.compile(r"(?i)\b(?:fell|decreased|declined|contracted|dropped|weakened)\b")

        slide_title = ""
        try:
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                slide_title = slide.shapes.title.text.strip()
        except Exception:
            pass

        for shape in slide.shapes:
            if not getattr(shape, "has_chart", False):
                continue
            chart = shape.chart
            chart_title = ""
            try:
                if chart.has_title:
                    chart_title = chart.chart_title.text_frame.text.strip()
            except Exception:
                pass
            title = chart_title or slide_title
            if not title or not (positive_words.search(title) or negative_words.search(title)):
                continue
            try:
                series = list(chart.series)
                if len(series) != 1:
                    continue
                values = [float(value) for value in series[0].values if value is not None]
            except Exception:
                continue
            if len(values) < 2:
                continue
            delta = values[-1] - values[0]
            tolerance = max(abs(values[0]), abs(values[-1]), 1.0) * 1e-9
            contradicts = (positive_words.search(title) and delta < -tolerance) or (negative_words.search(title) and delta > tolerance)
            if contradicts:
                self.issues.append(
                    PreflightIssue(
                        idx,
                        "title_data_misalignment",
                        f"Directional title '{title}' conflicts with chart endpoints ({values[0]:g} to {values[-1]:g})",
                        severity="error",
                    )
                )

    def _check_template_completeness(self, idx: int, slide: Any) -> None:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    orig_text = p.text
                    clean = orig_text
                    # Check and clean unresolved variables like {variable_name}
                    unresolved = re.findall(r"\{[a-zA-Z0-9_]+\}", clean)
                    if unresolved:
                        self.issues.append(
                            PreflightIssue(idx, "unresolved_template_variable", f"Unresolved template variables {unresolved} in text", severity="warning")
                        )
                        for var in unresolved:
                            clean = clean.replace(var, "")
                    # Check and clean literal "None" or "null" appearing as words
                    if re.search(r"\b(?:None|null)\b", clean):
                        self.issues.append(
                            PreflightIssue(idx, "literal_none_null", "Literal 'None' or 'null' detected in text", severity="warning")
                        )
                        clean = re.sub(r"\b(?:None|null)\b", "", clean)
                    # Check and clean double spaces
                    if "  " in clean:
                        clean = re.sub(r"[ \t]{2,}", " ", clean)
                    clean = clean.strip()
                    if clean != orig_text:
                        p.text = clean

    def _check_duplicate_card_titles(self, idx: int, slide: Any) -> None:
        titles: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                top = getattr(shape, "top", None)
                if top and top.inches > 1.3:
                    paras = [p for p in shape.text_frame.paragraphs if p.text.strip()]
                    if paras and paras[0].font.bold and len(paras[0].text) < 50:
                        t = paras[0].text.strip().casefold()
                        titles.append(t)
        counts = Counter(titles)
        for t, count in counts.items():
            if count > 1 and t not in ("key facts", "business focus", "source data", "notes", "financial metric"):
                self.issues.append(
                    PreflightIssue(
                        idx,
                        "duplicate_card_title",
                        f"Duplicate card title '{t}' detected across {count} elements on slide {idx + 1}",
                        severity="warning",
                    )
                )

    def _check_slide_title_quality(self, idx: int, slide: Any) -> None:
        title = ""
        for shape in slide.placeholders:
            try:
                if shape.placeholder_format.idx in (14, 15) and shape.has_text_frame:
                    title = shape.text.strip()
                    break
            except Exception:
                pass
        if not title and hasattr(slide, "shapes"):
            first = next((s for s in slide.shapes if getattr(s, "has_text_frame", False) and s.text.strip()), None)
            if first:
                title = first.text.strip()

        for pattern in (
            r"(?i)\bunaudited\s+analysis\b",
            r"(?i)\btrend\s+and\s+related\s+measures\b",
            r"(?i)\bevidence-?backed\s+comparison\b",
            r"(?i)\bretained\s+reported\s+values\b",
            r"(?i)^(?:add|less|plus|minus|adjustments?|reconciliation|sub-?total|total)\s*[:\-\u2013\u2014]",
        ):
            if re.search(pattern, title):
                self.issues.append(
                    PreflightIssue(
                        idx,
                        "generic_slide_title",
                        f"Slide {idx + 1} has generic or unparsed raw table title: '{title}'",
                        severity="warning",
                    )
                )
                break

    def _check_thank_you_slide(self) -> None:
        if not self.presentation.slides:
            return
        last_slide = self.presentation.slides[-1]
        text = " ".join(s.text for s in last_slide.shapes if s.has_text_frame).casefold()
        layout_name = last_slide.slide_layout.name if hasattr(last_slide, "slide_layout") else ""
        if not ("thank you" in text or "thank" in layout_name.casefold()):
            self.issues.append(
                PreflightIssue(
                    len(self.presentation.slides) - 1,
                    "missing_thank_you_slide",
                    "The final slide is not the official FOURIER Thank You slide",
                    severity="warning",
                )
            )

