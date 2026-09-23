"""Static illustration layouts for local presentation export.

Artwork is contextual, never analytical evidence. It is not added to the
document model, prompts, disk cache or model requests. Only static PNG/JPEG is
accepted and aspect ratio is preserved without cropping away image content.
"""

import io
from PIL import Image, UnidentifiedImageError


def validate_artwork(payload: bytes | None) -> bytes | None:
    if payload is None:
        return None
    if not payload or len(payload) > 8_000_000:
        raise ValueError("Presentation artwork must be a PNG or JPEG smaller than 8 MB.")
    try:
        with Image.open(io.BytesIO(payload)) as image:
            if image.format not in {"PNG", "JPEG"} or getattr(image, "n_frames", 1) != 1:
                raise ValueError("Presentation artwork must be a static PNG or JPEG.")
            if min(image.size) < 100 or image.width * image.height > 16_000_000:
                raise ValueError("Presentation artwork dimensions must be at least 100 pixels and at most 16 megapixels.")
            image.verify()
    except (OSError, UnidentifiedImageError, Image.DecompressionBombError) as exc:
        raise ValueError("The presentation artwork could not be read as a safe static image.") from exc
    return payload


def _picture(slide, artwork: bytes, rect):
    from pptx.util import Inches
    with Image.open(io.BytesIO(artwork)) as image:
        scale = min(rect.w / image.width, rect.h / image.height)
        width, height = image.width * scale, image.height * scale
    picture = slide.shapes.add_picture(io.BytesIO(artwork), Inches(rect.x + (rect.w - width) / 2),
        Inches(rect.y + (rect.h - height) / 2), width=Inches(width), height=Inches(height))
    picture.name = "user_supplied_illustration"
    return picture


def add_picture_cover(presentation, title: str, purpose: str, artwork: bytes):
    from .pptx_export import _base_slide, _rule
    from .slide_compositor import Rect, _lines, _put_text, _put_commentary
    from .presentation_style import DARK, MUTED, PURPLE
    slide = _base_slide(presentation, "", "")
    slide.name = "picture_cover"
    width, height = presentation.slide_width.inches, presentation.slide_height.inches
    left_width = (width - 1.5) * .48
    lines = _lines(title, left_width, 28)
    title_h = len(lines) * .46
    if title_h > height - 3:
        raise ValueError("Cover title is too long for the picture layout; shorten the planned title.")
    top = 1.55
    _rule(slide, .60, top - .25, 1.4, .035, PURPLE)
    _put_text(slide, title, Rect(.60, top, left_width, title_h), size=28, bold=True, color=DARK)
    text_top = top + title_h + .32
    remaining = _put_commentary(slide, purpose, Rect(.60, text_top, left_width, max(.6, height - 1.2 - text_top))) if purpose else ""
    _picture(slide, artwork, Rect(.95 + left_width, 1.35, width - left_width - 1.55, height - 2.55))
    _put_text(slide, "User-selected illustration", Rect(.95 + left_width, height - 1.06, width - left_width - 1.55, .20), size=8, color=MUTED)
    # Long descriptions stay available rather than disappearing at a character cap.
    if remaining:
        slide.notes_slide.notes_text_frame.text = "Document purpose (continued): " + remaining
    return slide


def add_picture_profile(presentation, company, title: str, pages: list[int], artwork: bytes,
                        *, source_page: int | None = None):
    from .pptx_export import _source_footer
    from .slide_compositor import Rect, _base, _put_text, _put_commentary
    from .presentation_style import MUTED
    slide, top = _base(presentation, title, company.document_type)
    slide.name = "picture_profile"
    width, height = presentation.slide_width.inches, presentation.slide_height.inches
    text_width = (width - 1.38) * .58
    if source_page is not None:
        # The source image is supporting context. Keep the complete company
        # record in notes instead of creating a sparse continuation slide.
        parts = [company.name]
        for label, value in (
            ("Business", company.one_line_description or company.industry),
            ("Model", company.business_model),
            ("Products", ", ".join(company.products)),
            ("Applications", ", ".join(company.application_areas)),
            ("Customers", ", ".join(company.customer_types)),
            ("Markets", ", ".join(company.geographies)),
            ("Listing", ", ".join(v for v in (company.listing_market, company.stock_code) if v)),
        ):
            if value:
                parts.append(f"{label}: {value}")
        remaining = _put_commentary(slide, "\n\n".join(parts),
                                    Rect(.55, top, text_width, height - 1.12 - top))
        picture = _picture(slide, artwork,
                           Rect(.83 + text_width, top, width - text_width - 1.38, height - 1.5 - top))
        picture.name = f"source_document_image:p{source_page}"
        _put_text(slide, f"Image from source document, p. {source_page}",
                  Rect(.83 + text_width, height - 1.29, width - text_width - 1.38, .20),
                  size=8, color=MUTED)
        _put_text(slide, _source_footer(pages), Rect(.55, height - .82, width - 1.1, .20),
                  size=9, color=MUTED)
        slide.notes_slide.notes_text_frame.text = (
            f"Context image reproduced from uploaded PDF page {source_page}.\n\n"
            + company.model_dump_json(indent=2)
            + ("\n\nVisible company details continued:\n" + remaining if remaining else "")
        )
        return slide
    groups = [company.name, company.one_line_description]
    for label, values in (
        ("Profile", [f"{label}: {value}" for label, value in (
            ("Industry", company.industry), ("Headquarters", company.headquarters),
            ("Reporting currency", company.reporting_currency), ("Period", company.track_record_period)) if value]),
        ("Business", [company.business_model, *company.customer_types]),
        ("Products and segments", [*company.products, *company.segments]),
        ("Markets and listing", [*company.geographies, company.market_position, company.listing_market,
                                 company.stock_code, company.offering_type, *company.listing_facts]),
        ("Key facts", [f"{f.label}: {f.value}" for f in company.key_facts]),
    ):
        retained = list(dict.fromkeys(v.strip() for v in values if v.strip()))
        if retained:
            groups.append(label + "\n" + "\n".join(retained))
    text = "\n\n".join(dict.fromkeys(t for t in groups if t.strip()))
    remaining = _put_commentary(slide, text, Rect(.55, top, text_width, height - 1.12 - top))
    picture = _picture(slide, artwork, Rect(.83 + text_width, top, width - text_width - 1.38, height - 1.5 - top))
    _put_text(slide, "User-selected illustration", Rect(.83 + text_width, height - 1.29, width - text_width - 1.38, .20), size=8, color=MUTED)
    _put_text(slide, _source_footer(pages), Rect(.55, height - .82, width - 1.1, .20), size=9, color=MUTED)
    while remaining:
        continuation, ctop = _base(presentation, title + " (continued)", "")
        continuation.name = "picture_profile_continued"
        remaining = _put_commentary(continuation, remaining, Rect(.55, ctop, width - 1.1, height - 1.12 - ctop))
        _put_text(continuation, _source_footer(pages), Rect(.55, height - .82, width - 1.1, .20), size=9, color=MUTED)
    return slide
