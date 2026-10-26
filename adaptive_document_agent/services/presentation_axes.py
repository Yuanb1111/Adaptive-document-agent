"""Readable native axes from retained values and explicit source dates."""

from datetime import date

from .presentation_style import readable_axis_bounds


def explicit_date_categories(values, categories):
    """Never infer a calendar date from a fiscal year or an unknown period."""
    from adaptive_document_agent.document_model.period_semantic_validator import format_observation_period

    dates = {}
    for item in values:
        if item.period_type != "balance_sheet_date":
            return None
        try:
            point = date.fromisoformat(item.as_of_date or item.period)
        except (TypeError, ValueError):
            return None
        label = format_observation_period(item)
        if label in dates and dates[label] != point:
            return None
        dates[label] = point
    return [dates[label] for label in categories] if all(label in dates for label in categories) else None


def style_date_axis(chart, dates, width):
    """Space points by calendar time; keep full dates in the chart workbook."""
    from pptx.oxml.xmlchemy import OxmlElement

    axis = chart.category_axis
    axis.tick_labels.number_format = "dd mmm yyyy" if width >= 5 else "mmm yyyy"
    axis.tick_labels.number_format_is_linked = False
    # Anchor regular calendar ticks to a reported date. Arbitrary padding on
    # the left shifted annual December observations to November tick labels.
    # The right display window may expand without changing the tick origin.
    padding = max(1, (max(dates) - min(dates)).days * .06)
    origin = date(1899, 12, 30)
    for tag, value in (("c:min", (min(dates) - origin).days),
                       ("c:max", (max(dates) - origin).days + padding)):
        nodes = axis._element.scaling.xpath(tag)
        element = nodes[0] if nodes else OxmlElement(tag)
        element.set("val", str(value))
        if not nodes:
            axis._element.scaling.append(element)
    months = (max(dates).year - min(dates).year) * 12 + max(dates).month - min(dates).month
    # A full date is much wider than a month/year label. Reserve axis margins
    # and inter-label space rather than treating each inch as one readable tick.
    font_size = axis.tick_labels.font.size
    points = font_size.pt if font_size is not None else 11
    characters = 11 if width >= 5 else 8
    label_points = characters * points * .6 + 2 * points
    target = max(2, int(max(.5, width - .8) * 72 / label_points))
    step = next((n for n in (1, 3, 6, 12, 24, 60) if n >= months / target), 120)
    for tag, value in (("c:baseTimeUnit", "days"), ("c:majorUnit", str(step)), ("c:majorTimeUnit", "months")):
        nodes = axis._element.xpath(tag)
        element = nodes[0] if nodes else OxmlElement(tag)
        element.set("val", value)
        if not nodes:
            axis._element.append(element)


def label_source_dates(chart, dates, *, width=6, precision=None):
    """Display exact point dates as well as the continuous calendar axis."""
    from pptx.util import Pt
    from pptx.enum.chart import XL_DATA_LABEL_POSITION

    if len(chart.series) != 1 or len(dates) > 6:
        return
    series = chart.series[0]
    for i, (point, day, value) in enumerate(zip(series.points, dates, series.values)):
        label = point.data_label
        label.position = (XL_DATA_LABEL_POSITION.RIGHT if width < 5 and i == 0 else
                          XL_DATA_LABEL_POSITION.ABOVE if i % 2 == 0 else XL_DATA_LABEL_POSITION.BELOW)
        # The plotted values are already scaled; labels use that same unit.
        amount = (f'{value:,.{precision}f}' if precision is not None
                  else f'{value:,.2f}'.rstrip('0').rstrip('.'))
        # A single paragraph is interoperable with renderers that ignore extra
        # paragraphs or breaks in native chart labels. Alternate above/below.
        # On a three-panel slide full dates beside every marker collide with
        # neighboring values. Calendar spacing stays intact; exact dates remain
        # in the editable workbook and source-data appendix.
        # Narrow date axes can omit a final interim tick. Preserve that exact
        # observed endpoint visibly, rather than only in the workbook.
        label.text_frame.text = (f'{day:%d %b %Y}: {amount}'
                                 if width >= 5 or i == len(dates) - 1 else amount)
        for paragraph in label.text_frame.paragraphs:
            paragraph.font.size = Pt(9)
            for run in paragraph.runs:
                # Explicit run formatting prevents the template's large chart
                # font from winning over a paragraph default in PowerPoint.
                run.font.name = 'Arial'
                run.font.size = Pt(9)
                run.font.bold = False


def style_value_range(axis, values, *, height=None):
    """Use zero for one-sided data and rounded ticks for every native chart."""
    if not values:
        return
    low, high = min(values), max(values)
    if low >= 0:
        low, high = 0.0, high * 1.20 if high else 1.0
    elif high <= 0:
        low, high = low * 1.30, 0.0
    else:
        low, high = low * 1.45, high * 1.25
    low, high, step = readable_axis_bounds(low, high)
    if height is not None:
        # Chart frames include data labels and category labels. Short panels
        # need fewer ticks, while full-height charts keep the regular scale.
        intervals = max(2, min(5, int(max(.4, height - 1.0) / .28)))
        from math import ceil, floor, log10
        raw = (high - low) / intervals
        order = 10 ** floor(log10(raw))
        step = next(n * order for n in (1, 2, 2.5, 5, 10) if n * order >= raw)
        low, high = floor(low / step) * step, ceil(high / step) * step
    axis.minimum_scale, axis.maximum_scale, axis.major_unit = low, high, step
    # Axis labels describe the rounded scale; point labels retain source precision.
    axis.tick_labels.number_format = "0.##;-0.##;0" if step < 1 else "#,##0;-#,##0;0"
    axis.tick_labels.number_format_is_linked = False
