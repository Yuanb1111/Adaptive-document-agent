"""Check colours of generated slide content without changing source artwork."""

from collections.abc import Iterable
from typing import Any

from pptx.oxml.ns import qn

from .fourier_brand import APPROVED_COLORS


def _check_rgb(root: Any, location: str, *, exclude_pictures: bool = False) -> None:
    for color in root.iter(qn("a:srgbClr")):
        # Raster/source pictures and their original colour effects are artwork,
        # not generated chart, shape or text formatting.
        if exclude_pictures and any(parent.tag == qn("p:pic") for parent in color.iterancestors()):
            continue
        value = (color.get("val") or "").upper()
        if value not in APPROVED_COLORS:
            raise ValueError(f"{location}: unsupported generated brand colour #{value or '(missing)'}.")


def _has_approved_solid_fill(parent: Any) -> bool:
    if parent is None:
        return False
    fill = parent.find(qn("a:solidFill"))
    if fill is None:
        return False
    color = fill.find(qn("a:srgbClr"))
    return color is not None and (color.get("val") or "").upper() in APPROVED_COLORS


def _shapes(shapes: Iterable[Any]) -> Iterable[Any]:
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from _shapes(shape.shapes)


def validate_generated_brand(presentation: Any) -> None:
    """Reject unsupported explicit RGB colours and automatic chart series fills.

    Call on the generated deck after template sample slides have been removed.
    Layouts, masters and picture media are deliberately excluded: their original
    logos, gradients and artwork are retained. This is a read-only colour check,
    not a rendering/contrast check or a validation of inherited theme colours.
    An explicit no-fill outline is allowed; a series fill must be approved RGB.
    """
    for number, slide in enumerate(presentation.slides, start=1):
        location = f"Slide {number}"
        _check_rgb(slide._element, location, exclude_pictures=True)
        for shape in _shapes(slide.shapes):
            if not shape.has_chart:
                continue
            chart_location = f"{location}, chart '{shape.name}'"
            root = shape.chart._chartSpace
            _check_rgb(root, chart_location)
            for index, series in enumerate(root.iter(qn("c:ser")), start=1):
                series_location = f"{chart_location}, series {index}"
                properties = series.find(qn("c:spPr"))
                if not _has_approved_solid_fill(properties):
                    raise ValueError(f"{series_location}: an explicit approved solid series fill is required.")
                line = properties.find(qn("a:ln"))
                if line is None or (line.find(qn("a:noFill")) is None and not _has_approved_solid_fill(line)):
                    raise ValueError(f"{series_location}: an explicit approved series line or no-fill line is required.")
