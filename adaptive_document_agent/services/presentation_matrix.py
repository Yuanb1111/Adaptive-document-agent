"""Exact category-by-measure matrix from model-selected source observations."""

from __future__ import annotations

from dataclasses import dataclass
import json

from adaptive_document_agent.document_model.period_semantic_validator import format_observation_period


@dataclass(frozen=True)
class ComparisonMatrix:
    row_labels: tuple[str, ...]
    column_labels: tuple[str, ...]
    cells: tuple[tuple[object, ...], ...]
    observations: tuple


def comparison_matrix(block, index) -> ComparisonMatrix:
    """Accept a fully populated matrix with one source observation per cell."""
    dimension = block.matrix_dimension.strip()
    if not dimension or len(block.observation_ids) < 4 or len(set(block.observation_ids)) != len(block.observation_ids):
        raise ValueError("Matrix needs a named category dimension and distinct source records.")
    items = [index.get(oid) for oid in block.observation_ids]
    if any(item is None or item.value is None or not item.period or not item.evidence
           or item.validation_status != "valid" or item.anomaly_notes for item in items):
        raise ValueError("Matrix cells require valid reported values, periods and citations.")
    rows, columns, cells, scope = [], [], {}, None
    for item in items:
        dims = {**item.dimensions, **item.category_dimensions}
        if dimension == "source_metric":
            from adaptive_document_agent.document_model.series import display_metric_name
            row = display_metric_name(item)
        else:
            row = str(dims.pop(dimension, "")).strip()
        if not row:
            raise ValueError("Every matrix record must identify the selected category.")
        other_dims = tuple(sorted((key, str(value)) for key, value in dims.items()
                                  if key not in {"table_context", "section", "column_role", "period_basis"}))
        item_scope = (item.entity, other_dims)
        if scope is None:
            scope = item_scope
        elif item_scope != scope:
            raise ValueError("Matrix cells mix distinct entities or nonselected dimensions.")
        column = ("" if dimension == "source_metric" else item.metric_original,
                  format_observation_period(item), item.unit,
                  item.raw_unit, item.currency, item.parent_section)
        if row not in rows:
            rows.append(row)
        if column not in columns:
            columns.append(column)
        if (row, column) in cells:
            raise ValueError("Matrix has conflicting or duplicate values for one category and measure.")
        cells[row, column] = item
    if not 2 <= len(rows) <= 6 or not 2 <= len(columns) <= 5:
        raise ValueError("A readable matrix needs two to six categories and two to five measures.")
    if len(cells) != len(rows) * len(columns):
        raise ValueError("Matrix has missing category/measure cells; missing is not zero.")
    names = []
    for metric, period, unit, raw_unit, currency, parent in columns:
        name = f"{parent}: {metric}" if parent and parent.casefold() not in metric.casefold() else metric
        parts = []
        if currency and currency.casefold() not in (raw_unit or "").casefold():
            parts.append(currency)
        parts.append(raw_unit or ("%" if unit == "percent" else unit))
        suffix = " ".join(part for part in parts if part)
        names.append((f"{name}\n" if name else "") + period
                     + (f" ({suffix})" if suffix else ""))
    return ComparisonMatrix(tuple(rows), tuple(names),
                            tuple(tuple(cells[row, col] for col in columns) for row in rows), tuple(items))


def topic_matrix_dimension(items: list) -> str | None:
    """Find an exact, readable source matrix without inferring missing cells."""
    from adaptive_document_agent.models import PresentationVisualBlock

    index = {item.id: item for item in items}
    dimensions = {key for item in items
                  for key in {**item.dimensions, **item.category_dimensions}
                  if key not in {"table_context", "section", "column_role", "period_basis"}}
    for dimension in [*sorted(dimensions), "source_metric"]:
        try:
            comparison_matrix(PresentationVisualBlock(
                role="matrix", observation_ids=list(index), matrix_dimension=dimension), index)
            return dimension
        except ValueError:
            continue
    return None


def _matrix_value(item) -> str:
    value = item.raw_value.strip() or str(item.value)
    if item.unit == "percent" and item.raw_unit == "%" and "%" not in value:
        value += "%"
    return value


def render_matrix(presentation, slide_plan, block, index):
    """Continue complete matrix rows on readable, independently cited pages."""
    from pptx.util import Inches, Pt
    from .pptx_export import _source_footer, _text, _rgb, FOURIER_DARK, FOURIER_MUTED, FOURIER_PURPLE
    from .slide_compositor import _base, _lines
    from .presentation_header_overflow import visual_header, append_header_commentary
    from .presentation_style import FONT
    from .fourier_brand import WHITE

    matrix = comparison_matrix(block, index)
    width = presentation.slide_width.inches - 1.1
    first_width = width * .22
    data_width = (width - first_width) / len(matrix.column_labels)
    widths = [first_width, *([data_width] * len(matrix.column_labels))]
    headers = ["Metric" if block.matrix_dimension == "source_metric" else block.matrix_dimension,
               *matrix.column_labels]
    if any(len(_lines(label, cell_width - .20, 12)) > 3
           for label, cell_width in zip(headers, widths)):
        raise ValueError("Matrix column heading exceeds readable capacity.")
    rows = [[label, *(_matrix_value(item) for item in records)]
            for label, records in zip(matrix.row_labels, matrix.cells)]
    # Match the explicit cell margins and line spacing below. A long category
    # or raw value gets more height, never a smaller font or truncated copy.
    row_heights = [max(.78, max(len(_lines(value, cell_width - .20, 14))
                               for value, cell_width in zip(row, widths)) * 18 / 72 + .12)
                   for row in rows]
    header_h = .85
    header = visual_header(slide_plan)
    first_slide, top = _base(presentation, header.title, header.subtitle)
    available = presentation.slide_height.inches - 1.02 - top - .08
    if any(header_h + row_h > available for row_h in row_heights):
        raise ValueError("Matrix row exceeds readable slide capacity.")
    ranges = []
    start, height = 0, header_h
    for i, row_h in enumerate(row_heights):
        if height + row_h > available:
            ranges.append((start, i))
            start, height = i, header_h
        height += row_h
    ranges.append((start, len(rows)))

    for page, (start, end) in enumerate(ranges):
        slide = first_slide if page == 0 else _base(presentation, header.title, header.subtitle)[0]
        slide.name = "evidence_comparison_matrix"
        observations = [item for records in matrix.cells[start:end] for item in records]
        heights = [header_h, *row_heights[start:end]]
        shape = slide.shapes.add_table(end - start + 1, len(headers),
                                       Inches(.55), Inches(top + .08), Inches(width), Inches(sum(heights)))
        shape.name = "table:comparison_matrix:" + ",".join(item.id for item in observations)
        table = shape.table
        for column, cell_width in zip(table.columns, widths):
            column.width = Inches(cell_width)
        for i, (values, row_h) in enumerate(zip([headers, *rows[start:end]], heights)):
            table.rows[i].height = Inches(row_h)
            for j, value in enumerate(values):
                cell = table.cell(i, j)
                cell.text = value
                cell.margin_left = cell.margin_right = Inches(.10)
                cell.margin_top = cell.margin_bottom = Inches(.06)
                cell.text_frame.word_wrap = True
                if i == 0:
                    # The inherited table theme may also use purple; set both
                    # colours so the repeated headings stay visibly readable.
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb(FOURIER_PURPLE)
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.name = FONT
                    paragraph.font.size = Pt(12 if i == 0 else 14)
                    paragraph.font.bold = i == 0
                    paragraph.font.color.rgb = _rgb(WHITE if i == 0 else FOURIER_DARK)
                    paragraph.line_spacing = Pt(15 if i == 0 else 18)
                    paragraph.space_before = paragraph.space_after = Pt(0)
        pages = sorted({source.page for item in observations for source in item.evidence})
        footer = _source_footer(pages)
        if len(ranges) > 1:
            footer += f" | Matrix {page + 1} of {len(ranges)}"
        _text(slide, footer, .55, presentation.slide_height.inches - .82,
              width, .20, size=9, color=FOURIER_MUTED)
        slide.notes_slide.notes_text_frame.text = json.dumps({
            "matrix_dimension": block.matrix_dimension,
            "source_observations": [item.model_dump(mode="json") for item in observations],
        }, ensure_ascii=False)
    append_header_commentary(presentation, slide_plan, header,
        sorted({source.page for item in matrix.observations for source in item.evidence}), {
            "matrix_dimension": block.matrix_dimension,
            "source_observations": [item.model_dump(mode="json") for item in matrix.observations],
        })
    # Existing callers that inspect a single-page matrix still receive its slide.
    return first_slide
