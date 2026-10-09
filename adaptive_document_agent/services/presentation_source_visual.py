"""Reuse first-page artwork locally as cited context.

An embedded source image avoids copying separate cover titles into the picture.
Documents without suitable embedded artwork retain a first-page preview.
"""
from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from typing import Literal

from adaptive_document_agent.models import PipelineResult
from .presentation_artwork import validate_artwork


@dataclass(frozen=True)
class SourceVisual:
    payload: bytes
    page: int
    kind: Literal["embedded_image", "page_snapshot"] = "page_snapshot"


def select_company_source_visual(pdf_bytes: bytes, result: PipelineResult) -> SourceVisual:
    """Prefer a substantial native first-page image, with no model call.

    Identity resolution and document type do not gate a faithful source preview.
    Hash validation prevents a different upload from being attached to results.
    """
    if sha256(pdf_bytes).hexdigest() != result.document.sha256:
        raise ValueError("The uploaded PDF does not match the analysed document; source image was not embedded.")
    import pymupdf

    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as document:
        if not document.page_count or document.needs_pass:
            raise ValueError("The uploaded PDF has no readable first page for the source preview.")
        page = document[0]
        candidates = []
        for info in page.get_images(full=True)[:80]:
            xref, width, height = info[0], info[2], info[3]
            if min(width, height) < 450 or width * height > 16_000_000:
                continue
            areas = [(rect & page.rect).get_area() / max(page.rect.get_area(), 1)
                     for rect in page.get_image_rects(xref)]
            area = max(areas, default=0)
            if area >= .12 and .3 <= width / height <= 3.2:
                candidates.append((area, width * height, xref))
        for _, _, xref in sorted(set(candidates), reverse=True):
            image = document.extract_image(xref)
            if image.get("ext") not in {"png", "jpeg", "jpg"}:
                continue
            try:
                payload = validate_artwork(image["image"])
            except ValueError:
                continue
            if payload is not None:
                return SourceVisual(payload=payload, page=1, kind="embedded_image")
        # Bound even unusually large page boxes before allocating the raster.
        scale = 1800 / max(page.rect.width, page.rect.height)
        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale),
                                colorspace=pymupdf.csRGB, alpha=False)
        payload = validate_artwork(pixmap.tobytes("png"))
        if payload is None:  # pragma: no cover - a pixmap always returns bytes
            raise ValueError("The uploaded PDF first-page preview could not be rendered.")
    return SourceVisual(payload=payload, page=1)


@dataclass(frozen=True)
class _ProfileLayout:
    columns: int
    text_width: float
    image_width: float
    column_width: float
    heights: list[tuple[float, float]]
    row_heights: list[float]
    gap: float


def _profile_layout(items, width: float, height: float) -> _ProfileLayout | None:
    """Measure complete authored sections before choosing text/image columns.

    A vertical reading order is preferable to narrow text cards. Try it with
    progressively more text space before considering a compact two-column grid.
    The picture keeps a substantial share of the page in every candidate.
    """
    from .presentation_brief import _copy_height, _heading_height

    for columns in (1, 2):
        if columns == 2 and not 3 <= len(items) <= 4:
            continue
        for image_fraction in (.36, .32, .28):
            image_width = width * image_fraction
            text_width = width - image_width - .40
            column_width = (text_width - .30 * (columns - 1)) / columns
            heights = [(_heading_height(item.title, column_width),
                        _copy_height(item.text, column_width, 16) + .06) for item in items]
            row_heights = [max(h + b + .06 for h, b in heights[i:i + columns])
                           for i in range(0, len(heights), columns)]
            needed = sum(row_heights) + max(0, len(row_heights) - 1) * .22
            if items and needed <= height:
                gap = min(.34, .22 + (height - needed) / max(len(row_heights) - 1, 1))
                return _ProfileLayout(columns, text_width, image_width, column_width,
                                      heights, row_heights, gap)
    return None


def render_profile_with_source(presentation, title, items, visual: SourceVisual, *, notes=""):
    """Keep complete introductory copy beside uncropped source artwork.

    A long introduction retains its existing pagination. The source image is
    omitted when it would require an otherwise empty preview page.
    """
    from .presentation_brief import render_profile
    from .presentation_artwork import _picture
    from .pptx_export import _source_footer, _text
    from .presentation_style import DARK, MUTED, PURPLE
    from .slide_compositor import Rect, _base

    slide, top = _base(presentation, title, "")
    slide.name = "source_document_profile"
    width, height = presentation.slide_width.inches, presentation.slide_height.inches
    bottom = height - 1.24
    retained = [item for item in items if item.text.strip()]
    layout = _profile_layout(retained, width - 1.3, bottom - top)
    fits = layout is not None
    if not retained:
        image_rect = Rect(.65, top, width - 1.3, bottom - top - .30)
    elif not fits:
        slide_id = presentation.slides._sldIdLst[-1]
        presentation.part.drop_rel(slide_id.rId)
        presentation.slides._sldIdLst.remove(slide_id)
        pages = render_profile(presentation, title, items, notes=notes)
        first = pages[0]
        copy_bottom = max((shape.top.inches + shape.height.inches for shape in first.shapes
                           if shape.name == 'brief:body'), default=bottom)
        available = bottom - copy_bottom - .18
        if available >= 1.45:
            picture = _picture(first, visual.payload,
                               Rect(.65, copy_bottom + .18, width - 1.3, available - .24))
            picture.name = f'source_document_image:p{visual.page}'
            _text(first, f'Source image: document p. {visual.page}', .65, bottom - .20,
                  width - 1.3, .20, size=9, color=MUTED).name = 'profile:image_source'
            cited_pages = sorted({visual.page} | {p for item in items for p in item.pages})
            for shape in first.shapes:
                if shape.has_text_frame and shape.text.startswith('Source: Document disclosures'):
                    shape.text_frame.paragraphs[0].runs[0].text = _source_footer(cited_pages)
                    break
            first.notes_slide.notes_text_frame.text += (
                f'\n\nOriginal image from uploaded PDF page {visual.page}; context artwork only.')
        return pages
    else:
        from pptx.util import Pt

        for index, (item, (heading_h, body_h)) in enumerate(zip(retained, layout.heights)):
            column, row = index % layout.columns, index // layout.columns
            x = .65 + column * (layout.column_width + .30)
            y = top + sum(layout.row_heights[:row]) + row * layout.gap
            if item.title:
                shape = _text(slide, item.title, x, y, layout.column_width, heading_h,
                              size=18, bold=True, color=PURPLE)
                shape.name = "brief:heading"
                for paragraph in shape.text_frame.paragraphs:
                    paragraph.line_spacing = Pt(22.5)
            shape = _text(slide, item.text, x, y + heading_h + .06, layout.column_width,
                          body_h, size=16, color=DARK)
            shape.name = "brief:body"
            for paragraph in shape.text_frame.paragraphs:
                paragraph.line_spacing = Pt(20)
        image_rect = Rect(.65 + layout.text_width + .40, top,
                          layout.image_width, bottom - top)
    picture = _picture(slide, visual.payload,
                       Rect(image_rect.x, image_rect.y, image_rect.w, image_rect.h - .24))
    picture.name = f"source_document_image:p{visual.page}"
    _text(slide, f'Source image: document p. {visual.page}', image_rect.x,
          image_rect.y + image_rect.h - .20, image_rect.w, .20,
          size=9, color=MUTED).name = 'profile:image_source'
    shown_items = retained if fits else []
    pages = sorted({visual.page} | {p for item in shown_items for p in item.pages})
    _text(slide, _source_footer(pages), .55, height - .82, width - 1.1, .20,
          size=9, color=MUTED)
    full_copy = "\n\n".join(f"{item.title}\n{item.text}\n{_source_footer(item.pages)}" for item in shown_items)
    origin = (f"Original embedded image from uploaded PDF page {visual.page}. "
              if visual.kind == "embedded_image" else
              f"Complete first page reproduced from uploaded PDF page {visual.page}. ")
    slide.notes_slide.notes_text_frame.text = (
        origin + "Context artwork only; analytical claims use their separately cited evidence.\n\n"
        + notes + "\n\n" + full_copy
    )
    return [slide]
