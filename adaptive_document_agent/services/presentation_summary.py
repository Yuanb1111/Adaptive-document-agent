"""Paginate selected summary findings without changing their meaning or priority."""

from dataclasses import dataclass
import re
from typing import Any

from pptx.util import Pt

from .presentation_brief import BODY_PT, HEADING_PT, LINE_HEIGHT_FACTOR, BriefItem, _copy_height, _heading_height


@dataclass(frozen=True)
class _Cell:
    item: BriefItem
    left: float
    width: float
    heading_height: float
    body_height: float
    padding: float = .06

    @property
    def height(self) -> float:
        return self.heading_height + self.body_height + self.padding


def _rows(items: list[BriefItem], width: float, columns: int, *, body_pt: int = BODY_PT,
          compact: bool = False, column_fraction: float = .5) -> list[list[_Cell]]:
    """Pair findings in reading order; let an odd final finding use the full width."""
    rows = []
    for offset in range(0, len(items), columns):
        group = items[offset:offset + columns]
        widths = ([width] if len(group) == 1 else
                  [(width - .30) * column_fraction, (width - .30) * (1 - column_fraction)])
        rows.append([
            _Cell(item, .65 + sum(widths[:column]) + column * .30, cell_width,
                  _heading_height(item.title, cell_width),
                  _copy_height(item.text, cell_width, body_pt) + (.03 if compact else .08),
                  .04 if compact else .06)
            for column, (item, cell_width) in enumerate(zip(group, widths))
        ])
    return rows


def _height(rows: list[list[_Cell]], *, compact: bool = False) -> float:
    return sum(max(cell.height for cell in row) for row in rows) + (.14 if compact else .16) * max(0, len(rows) - 1)


def _split_item(item: BriefItem, width: float, capacity: float) -> tuple[BriefItem, BriefItem]:
    """Continue long copy at a measured text boundary, retaining every character."""
    minimum = BriefItem(item.title, item.text[:1], item.pages, item.short_title)
    if item.title and _height(_rows([minimum], width, 1)) > capacity:
        # A title can itself be a long, qualified claim. Retain it verbatim in
        # the body, where ordinary continuation can handle its length. A short
        # model-provided heading is optional and must independently fit.
        short = item.short_title
        if not short.strip() or _height(_rows([BriefItem(short, item.title[:1], item.pages)], width, 1)) > capacity:
            short = ""
        item = BriefItem(short, item.title + "\n" + item.text, item.pages, "")
    low, high = 0, len(item.text)
    while low < high:
        middle = (low + high + 1) // 2
        part = BriefItem(item.title, item.text[:middle], item.pages, item.short_title)
        if _height(_rows([part], width, 1)) <= capacity:
            low = middle
        else:
            high = middle - 1
    if low == 0:
        raise ValueError("Summary page has no readable body capacity.")
    # Prefer a whole clause/word near the capacity boundary. Unspaced CJK and
    # long identifiers still make progress without discarding any characters.
    boundary = max((match.end() for match in re.finditer(r"[.!?。！？]\s*|\s+", item.text[:low])
                    if match.end() >= low * .6), default=low)
    return (BriefItem(item.title, item.text[:boundary], item.pages, item.short_title),
            BriefItem(item.title, item.text[boundary:], item.pages, item.short_title))


def render_complete_summary(presentation: Any, title: str, items: list[BriefItem], *,
                            notes: str = "", single_column: bool = False) -> list[Any]:
    """Keep every selected finding visible at the template's readable text sizes.

    Layout depends only on copy length and available space. Five concise findings
    can share two paired rows plus a full-width final row. Longer summaries use
    continuation pages, retaining the model's order, qualifiers and source pages.
    """
    from .slide_compositor import _base
    from .pptx_export import _source_footer, _text, FOURIER_DARK, FOURIER_MUTED, FOURIER_PURPLE

    remaining = [item for item in items if item.title.strip() or item.text.strip()]
    full_copy = "\n\n".join(f"{item.title}\n{item.text}\n{_source_footer(item.pages)}" for item in items)
    full_notes = "\n\n".join(value for value in (notes, full_copy) if value)
    slides = []
    width = presentation.slide_width.inches - 1.3
    while remaining or not slides:
        slide, top = _base(presentation, title + (" (continued)" if slides else ""), "")
        bottom = presentation.slide_height.inches - 1.02
        capacity = bottom - top
        rows = []
        selected_count = 0
        body_pt = BODY_PT
        # Balance complete findings before adding a sparse continuation. Search
        # readable sizes from largest to smallest and use one shared column
        # split across rows; never shorten evidence or its qualifications.
        if single_column and 3 <= len(remaining) <= 4:
            for size in (BODY_PT, 15, 14):
                candidates = [(_rows(remaining, width, 1, body_pt=size, compact=True), 0)]
                candidates.extend((_rows(remaining, width, 2, body_pt=size, compact=True,
                                         column_fraction=fraction), abs(fraction - .5))
                                  for fraction in (.4, .45, .5, .55, .6))
                fitting = [(candidate, balance) for candidate, balance in candidates
                           if _height(candidate, compact=True) <= capacity]
                if fitting:
                    rows = min(fitting, key=lambda pair: (_height(pair[0], compact=True), pair[1]))[0]
                    selected_count, body_pt = len(remaining), size
                    break
        # Bound the geometric search, not the content: remaining findings always
        # continue on another page. Eight short findings can use four paired rows.
        if not rows:
            for count in range(min(8, len(remaining)), 0, -1):
                for columns in ((1,) if single_column else ((2, 1) if count >= 4 else (1, 2))):
                    candidate = _rows(remaining[:count], width, columns)
                    if _height(candidate) <= capacity:
                        rows, selected_count = candidate, count
                        break
                if rows:
                    break
        if remaining and not rows:
            first, rest = _split_item(remaining[0], width, capacity)
            rows = _rows([first], width, 1)
            remaining = ([rest] if rest.text else []) + remaining[1:]
        else:
            remaining = remaining[selected_count:]

        used = sum(max(cell.height for cell in row) for row in rows)
        # The minimum gap was already included in the selected layout's fit.
        # Do not add a larger default gap back after choosing compact geometry.
        gap = min(.34, max(0, (capacity - used) / max(1, len(rows) - 1)))
        y = top
        for row in rows:
            for cell in row:
                if cell.item.title:
                    heading = _text(slide, cell.item.title, cell.left, y, cell.width, cell.heading_height,
                                    size=HEADING_PT, bold=True, color=FOURIER_PURPLE)
                    heading.name = "brief:heading"
                    heading.text_frame.paragraphs[0].line_spacing = Pt(HEADING_PT * LINE_HEIGHT_FACTOR)
                body = _text(slide, cell.item.text, cell.left, y + cell.heading_height + .04,
                             cell.width, cell.body_height, size=body_pt, color=FOURIER_DARK)
                body.name = "brief:body"
                # Match _copy_height's point-based budget. A proportional value
                # multiplies the renderer's font line metrics, not the font size,
                # and can place complete lines outside the reserved text box.
                body.text_frame.paragraphs[0].line_spacing = Pt(body_pt * LINE_HEIGHT_FACTOR)
            y += max(cell.height for cell in row) + gap
        if not rows:
            _text(slide, "See the evidence pages for supported findings and scope.",
                  .65, top, width, 1, size=BODY_PT, color=FOURIER_MUTED)
        pages = sorted({page for row in rows for cell in row for page in cell.item.pages})
        _text(slide, _source_footer(pages), .55, presentation.slide_height.inches - .82,
              presentation.slide_width.inches - 1.1, .20, size=9, color=FOURIER_MUTED)
        slide.notes_slide.notes_text_frame.text = full_notes
        slides.append(slide)
    return slides
