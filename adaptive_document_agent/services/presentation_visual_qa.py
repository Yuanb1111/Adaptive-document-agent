"""Bounded render/check/repair gate. Layout repair cannot rewrite evidence.

This is automated visual diagnostics, not an aesthetic score or a claim of
PowerPoint parity. Artifact Tool exposes boxes and line counts, not glyph boxes.
"""

from dataclasses import asdict, dataclass, field
import hashlib
import io
import math
from pathlib import Path
import tempfile
import subprocess
import zipfile

from lxml import etree
import numpy as np
from PIL import Image
from pptx import Presentation

from .presentation_rendering import RenderingError, check_render_input, configured_renderer
from .qa_reporter import CriticalQAError

POLICY_VERSION = "visual-qa-v1"
NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


@dataclass(frozen=True)
class VisualIssue:
    slide: int
    code: str
    severity: str
    shape_ids: tuple[str, ...] = ()
    message: str = ""


@dataclass
class VisualQAReport:
    status: str = "failed"
    renderer: str = "unavailable"
    policy: str = POLICY_VERSION
    attempts: int = 0
    slide_count: int = 0
    input_sha256: str = ""
    output_sha256: str = ""
    facts_preserved: bool = False
    cache_hit: bool = False
    issues: list[VisualIssue] = field(default_factory=list)
    repairs: list[dict] = field(default_factory=list)
    history: list[list[VisualIssue]] = field(default_factory=list)
    coverage: tuple[str, ...] = (
        "Gate requires local rendering of every slide, PNG integrity, page count and aspect ratio",
        "Gate checks rendered object bounds, blank pages, chart/table presence and source-footer collisions",
        "Text fit is estimated from renderer line counts, not measured glyphs; warnings require review",
        "No PowerPoint/Google Slides parity or aesthetic-quality certification",
    )

    def to_dict(self) -> dict:
        return asdict(self)


class VisualQAError(CriticalQAError):
    def __init__(self, report: VisualQAReport):
        self.report = report
        detail = "; ".join(f"{i.code} (slide {i.slide}): {i.message}" for i in report.issues[:8])
        super().__init__("PowerPoint visual export gate blocked: " + detail)


@dataclass
class VerifiedPresentation:
    payload: bytes
    report: VisualQAReport


def package_digest(payload: bytes, *, exclude_positions=False) -> str:
    """ZIP timestamps do not affect identity. Only slide positions may change.

    All other parts (including native charts, workbooks, notes, evidence text,
    relationships, images and order) participate in the preservation guard.
    """
    digest = hashlib.sha256()
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        for name in sorted(archive.namelist()):
            data = archive.read(name)
            if exclude_positions and name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                root = etree.fromstring(data, parser=etree.XMLParser(resolve_entities=False, no_network=True))
                for off in root.xpath(".//a:xfrm/a:off | .//p:xfrm/a:off", namespaces=NS):
                    off.set("x", "0")
                    off.set("y", "0")
                data = etree.tostring(root, method="c14n")
            digest.update(name.encode())
            digest.update(b"\0")
            digest.update(hashlib.sha256(data).digest())
    return digest.hexdigest()


def _intersection(a, b):
    return max(0, min(a[0]+a[2], b[0]+b[2])-max(a[0], b[0])) * max(0, min(a[1]+a[3], b[1]+b[3])-max(a[1], b[1]))


def _content(element):
    return bool(element.get("text", "").strip()) or element.get("kind") in {"chart", "table", "image"}


def _elements(page, slide):
    elements = [dict(e, scope="slide") for e in page.layout["elements"] if _content(e)]
    # Importers may renumber graphicFrame IDs for charts/tables. Match their
    # native, unique name and object kind; do not require importer IDs to survive.
    for e in elements:
        if e.get("kind") in {"chart", "table"}:
            matches = [s for s in slide.shapes if s.name == e.get("name")
                       and ((e["kind"] == "chart" and s.has_chart) or (e["kind"] == "table" and s.has_table))]
            if len(matches) == 1:
                e["id"] = str(matches[0].shape_id)
    # Layout exports also include placeholder prompt text and off-canvas master
    # palettes. Neither is visible slide content; never mistake them for evidence.
    for layer in page.layout.get("inheritedLayers", []):
        native = slide.slide_layout if layer.get("scope") == "layout" else slide.slide_layout.slide_master
        visible_ids = {str(s.shape_id) for s in native.shapes if not s.is_placeholder}
        for e in layer.get("elements", []):
            if str(e.get("id")) in visible_ids and _content(e):
                elements.append(dict(e, scope=layer.get("scope", "master")))
    return elements


def inspect_pages(payload: bytes, pages) -> list[VisualIssue]:
    deck = Presentation(io.BytesIO(payload))
    if len(pages) != len(deck.slides):
        return [VisualIssue(0, "PAGE_COUNT", "critical", message="Renderer omitted or added pages.")]
    issues = []
    ratio = deck.slide_width / deck.slide_height
    for number, (page, slide) in enumerate(zip(pages, deck.slides), 1):
        try:
            with Image.open(io.BytesIO(page.png)) as im:
                if im.format != "PNG" or im.width * im.height > 8_000_000 or min(im.size) < 100:
                    raise ValueError("invalid image dimensions")
                im.load()
                if abs(im.width / im.height / ratio - 1) > .015:
                    raise ValueError("incorrect aspect ratio")
                extrema = np.asarray(im.convert("RGB").resize((160, 90)), dtype=float).std(axis=(0, 1))
                if float(extrema.max()) < .6:
                    issues.append(VisualIssue(number, "BLANK_RENDER", "critical", message="Rendered page is blank or effectively uniform."))
                pixels = im.convert("RGB")
            layout = page.layout
            frame = layout["slide"]["frame"]
            width, height = float(frame["width"]), float(frame["height"])
            if (layout.get("unit") != "px" or width <= 0 or height <= 0
                    or not math.isfinite(width + height) or abs(width / height / ratio - 1) > .015):
                raise ValueError("invalid layout coordinates")
            elements = _elements(page, slide)
            for e in elements:
                box = e["bbox"]
                if len(box) != 4 or not all(isinstance(v, (int, float)) and math.isfinite(v) for v in box) or min(box[2:]) < 0:
                    raise ValueError("invalid element bounds")
        except (OSError, ValueError, KeyError, TypeError, Image.DecompressionBombError):
            issues.append(VisualIssue(number, "INVALID_RENDER", "critical", message="Unreadable PNG or invalid renderer layout."))
            continue
        body = [e for e in elements if e["scope"] == "slide"]
        ids = {str(e.get("id")) for e in body}
        for shape in slide.shapes:
            if (shape.has_chart or shape.has_table or (shape.has_text_frame and shape.text.strip())) and str(shape.shape_id) not in ids:
                issues.append(VisualIssue(number, "MISSING_RENDER_OBJECT", "critical", (str(shape.shape_id),), "A native text/chart/table object is missing from the render layout."))
        for e in body:
            x, y, w, h = e["bbox"]
            sid = str(e["id"])
            if x < -2 or y < -2 or x+w > width+2 or y+h > height+2:
                issues.append(VisualIssue(number, "OUT_OF_BOUNDS", "critical", (sid,), "Content extends beyond the actual slide canvas."))
            if e.get("kind") in {"chart", "table"} and w*h > 10000 and x >= 0 and y >= 0 and x+w <= width and y+h <= height:
                region = pixels.crop((round(x/width*pixels.width), round(y/height*pixels.height),
                                      round((x+w)/width*pixels.width), round((y+h)/height*pixels.height)))
                if float(np.asarray(region.resize((100, 60)), dtype=float).std(axis=(0, 1)).max()) < .6:
                    issues.append(VisualIssue(number, "BLANK_DATA_VISUAL", "critical", (sid,), "Chart/table region rendered as an effectively uniform image."))
            lines = e.get("textLayout", {}).get("lineCount", 0)
            sizes = [r["fontSize"] for p in e.get("paragraphs", []) for r in p.get("runs", []) if r.get("fontSize")]
            if lines and sizes and lines * max(sizes) > h + 3:
                issues.append(VisualIssue(number, "TEXT_FIT_ESTIMATE", "warning", (sid,), "Renderer line count suggests crowded text; inspect wrapping."))
            if not (e.get("text", "").lstrip().lower().startswith("source:") or e.get("name", "").startswith("qa:source")):
                continue
            blockers = [other for other in elements if other is not e and _intersection(e["bbox"], other["bbox"]) > 8
                        and (other.get("text", "").strip() or other.get("kind") in {"chart", "table"})]
            if blockers:
                issues.append(VisualIssue(number, "SOURCE_OVERLAP", "critical", (sid,), "Source footer overlaps other content or inherited template text."))
        # Chart/table collisions are never intentional nesting. Text/text box
        # intersections alone are uncertain (many authoring boxes have padding).
        for idx, a in enumerate(body):
            for b in body[idx+1:]:
                area = _intersection(a["bbox"], b["bbox"])
                denom = min(a["bbox"][2]*a["bbox"][3], b["bbox"][2]*b["bbox"][3])
                if denom > 0 and area / denom > .15 and a.get("kind") in {"chart", "table"} and b.get("kind") in {"chart", "table"}:
                    issues.append(VisualIssue(number, "DATA_VISUAL_OVERLAP", "critical", (str(a["id"]), str(b["id"])), "Native chart/table content overlaps."))
                elif (denom > 0 and area / denom > .3 and a.get("text", "").strip() and b.get("text", "").strip()
                      and a.get("kind") == b.get("kind") == "shape"):
                    issues.append(VisualIssue(number, "TEXT_BOX_OVERLAP", "warning", (str(a["id"]), str(b["id"])), "Text boxes overlap; padding and glyph bounds need visual review."))
    return list(dict.fromkeys(issues))


def repair_positions(payload: bytes, pages, issues: list[VisualIssue]) -> tuple[bytes, list[dict]]:
    """Move a source footer into free space or clamp a small overflow. No resizing."""
    deck = Presentation(io.BytesIO(payload))
    changes = {}
    actions = []
    for issue in issues:
        if issue.code not in {"SOURCE_OVERLAP", "OUT_OF_BOUNDS"} or not issue.shape_ids:
            continue
        page = pages[issue.slide-1]
        frame = page.layout["slide"]["frame"]
        width, height = frame["width"], frame["height"]
        elements = _elements(page, deck.slides[issue.slide-1])
        target = next((e for e in elements if e["scope"] == "slide" and str(e["id"]) == issue.shape_ids[0]), None)
        if not target or (issue.slide, issue.shape_ids[0]) in changes:
            continue
        x, y, w, h = target["bbox"]
        candidates = []
        if issue.code == "SOURCE_OVERLAP":
            # Only search the bottom quarter: sources must stay associated with
            # their page, not get moved into titles or data regions.
            candidates = [(x, y-offset) for offset in range(8, int(height*.25), 8)]
        elif w <= width-8 and h <= height-8:
            candidates = [(min(max(4, x), width-w-4), min(max(4, y), height-h-4))]
        for nx, ny in candidates:
            if min(nx, ny) < 0 or nx+w > width or ny+h > height or abs(nx-x)+abs(ny-y) > height*.25:
                continue
            padded = (nx-3, ny-3, w+6, h+6)
            blockers = [e for e in elements if e is not target and _intersection(padded, e["bbox"]) > 1]
            if blockers:
                continue
            changes[(issue.slide, issue.shape_ids[0])] = (round(nx/width*deck.slide_width), round(ny/height*deck.slide_height))
            actions.append({"slide": issue.slide, "shape_id": issue.shape_ids[0], "action": "translate", "reason": issue.code})
            target["bbox"] = [nx, ny, w, h]
            break
    if not changes:
        return payload, []
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(payload)) as archive, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dest:
        for entry in archive.infolist():
            data = archive.read(entry)
            for number in {key[0] for key in changes}:
                # Use actual slide-part names rather than assuming XML order.
                if entry.filename != str(deck.slides[number-1].part.partname).lstrip("/"):
                    continue
                root = etree.fromstring(data, parser=etree.XMLParser(resolve_entities=False, no_network=True))
                for (slide_no, sid), (x, y) in changes.items():
                    if slide_no != number:
                        continue
                    shape = root.xpath(".//p:spTree/*[.//p:cNvPr/@id=$sid]", namespaces=NS, sid=sid)
                    if len(shape) != 1:
                        continue
                    off = shape[0].xpath("./p:spPr/a:xfrm/a:off | ./p:xfrm/a:off", namespaces=NS)
                    if len(off) == 1:
                        off[0].set("x", str(x))
                        off[0].set("y", str(y))
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dest.writestr(entry, data)
    return output.getvalue(), actions


def verify_presentation(payload: bytes, *, renderer=None, max_repairs: int = 2, cache: dict | None = None) -> VerifiedPresentation:
    if not 0 <= max_repairs <= 2:
        raise ValueError("Visual repair budget must be between zero and two.")
    report = VisualQAReport(input_sha256=hashlib.sha256(payload).hexdigest())
    try:
        check_render_input(payload)
        renderer = renderer or configured_renderer()
        report.renderer = renderer.identity
        original_facts = package_digest(payload, exclude_positions=True)
        cache_key = (POLICY_VERSION, renderer.identity, max_repairs, package_digest(payload))
        if cache is not None and cache_key in cache:
            from copy import deepcopy
            hit = deepcopy(cache[cache_key])
            hit.report.cache_hit = True
            hit.report.input_sha256 = report.input_sha256
            return hit
        report.slide_count = len(Presentation(io.BytesIO(payload)).slides)
        if not 1 <= report.slide_count <= 150:
            raise RenderingError("Presentation exceeds the 150-slide rendering limit.")
        candidate = payload
        with tempfile.TemporaryDirectory(prefix="ada-ppt-qa-") as temporary:
            for attempt in range(max_repairs+1):
                directory = Path(temporary) / str(attempt)
                directory.mkdir()
                report.attempts += 1
                pages = renderer.render(candidate, directory)
                report.issues = inspect_pages(candidate, pages)
                report.history.append(report.issues.copy())
                report.facts_preserved = package_digest(candidate, exclude_positions=True) == original_facts
                if not report.facts_preserved:
                    raise RenderingError("Visual repair changed protected presentation content.")
                if not any(i.severity == "critical" for i in report.issues):
                    report.status = "passed_with_warnings" if report.issues else "passed"
                    report.output_sha256 = hashlib.sha256(candidate).hexdigest()
                    verified = VerifiedPresentation(candidate, report)
                    if cache is not None:
                        from copy import deepcopy
                        # Session-owned, one-entry cache. No global document cache.
                        cache.clear()
                        cache[cache_key] = deepcopy(verified)
                    return verified
                if attempt == max_repairs:
                    break
                repaired, actions = repair_positions(candidate, pages, report.issues)
                if not actions or repaired == candidate:
                    break
                if package_digest(repaired, exclude_positions=True) != original_facts:
                    raise RenderingError("Visual repair changed protected presentation content.")
                report.repairs.extend(actions)
                candidate = repaired
    except RenderingError as exc:
        report.issues.append(VisualIssue(0, "RENDERING_BLOCKED", "critical", message=str(exc)))
    except (OSError, ValueError, KeyError, TypeError, subprocess.TimeoutExpired) as exc:
        report.issues.append(VisualIssue(0, "VALIDATION_FAILED", "critical", message=f"Local rendering/validation could not complete ({type(exc).__name__})."))
    raise VisualQAError(report)
