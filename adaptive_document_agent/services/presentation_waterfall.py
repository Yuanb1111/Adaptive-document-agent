"""Exact, editable reconciliation bridge selected by the presentation model."""

from __future__ import annotations

from dataclasses import dataclass
import json

from adaptive_document_agent.document_model.period_semantic_validator import format_observation_period


@dataclass(frozen=True)
class Waterfall:
    observations: tuple
    levels: tuple[float, ...]


def waterfall_data(block, index) -> Waterfall:
    """Accept only a complete, same-table arithmetic reconciliation.

    Order is semantic and supplied by the model: opening value, signed
    components, closing value. Python checks arithmetic and scope only.
    """
    ids = block.observation_ids
    if not 3 <= len(ids) <= 6 or len(set(ids)) != len(ids):
        raise ValueError("A waterfall needs an opening value, one to four components and a closing value.")
    items = [index.get(oid) for oid in ids]
    if any(item is None or item.value is None or not item.evidence
           or item.validation_status != "valid" or item.anomaly_notes for item in items):
        raise ValueError("Waterfall components must be valid source observations.")
    first = items[0]
    scope = (first.period, first.period_basis, first.entity, first.currency,
             first.unit, first.raw_unit, first.unit_family, first.effective_table_id)
    if not scope[-1] or any((item.period, item.period_basis, item.entity, item.currency,
                             item.unit, item.raw_unit, item.unit_family, item.effective_table_id) != scope
                            for item in items[1:]):
        raise ValueError("Waterfall observations must share a source table, period and unit scope.")
    levels = [float(first.value)]
    for component in items[1:-1]:
        levels.append(levels[-1] + float(component.value))
    closing = float(items[-1].value)
    scale = max(1.0, *(abs(float(item.value)) for item in items))
    if abs(levels[-1] - closing) > max(1e-6, scale * 1e-9):
        raise ValueError("Waterfall components do not reconcile to the reported closing value.")
    if min([*levels, closing]) < 0:
        raise ValueError("Waterfall levels must stay nonnegative; use an exact-data table instead.")
    return Waterfall(tuple(items), tuple([*levels, closing]))


def render_waterfall(presentation, slide_plan, block, index):
    """Render a native Office stacked-column bridge with complete source rows."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt

    from .pptx_export import _source_footer, _text, FOURIER_BG_CARD, FOURIER_DARK, FOURIER_MUTED, FOURIER_PURPLE
    from .slide_compositor import _base
    from .presentation_header_overflow import visual_header, append_header_commentary
    from .presentation_style import FONT

    bridge = waterfall_data(block, index)
    items, levels = bridge.observations, bridge.levels
    header = visual_header(slide_plan)
    slide, top = _base(presentation, header.title, header.subtitle)
    slide.name = "evidence_waterfall"
    width = presentation.slide_width.inches - 1.1
    labels = [item.metric_original for item in items]
    data = CategoryChartData()
    data.categories = labels
    base, increase, decrease, total = [], [], [], []
    for offset, item in enumerate(items):
        if offset in {0, len(items) - 1}:
            base.append(0); increase.append(0); decrease.append(0)
            total.append(float(item.value))
        else:
            previous, current = levels[offset - 1], levels[offset]
            base.append(min(previous, current))
            increase.append(max(current - previous, 0))
            decrease.append(max(previous - current, 0))
            total.append(0)
    for name, values in (("Offset", base), ("Increase", increase),
                         ("Decrease", decrease), ("Reported total", total)):
        data.add_series(name, values)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED,
                                   Inches(.55), Inches(top + .42), Inches(width), Inches(1.76), data).chart
    # python-pptx can emit signed axis IDs; OpenXML expects unsigned integers,
    # and the offline renderer rejects negative IDs before rendering any page.
    from pptx.oxml.ns import qn
    for axis in chart._chartSpace.iter():
        if axis.tag in {qn("c:axId"), qn("c:crossAx")}:
            value = int(axis.get("val"))
            if value < 0:
                axis.set("val", str(value & 0xFFFFFFFF))
    slide.shapes[-1].name = "chart:waterfall"
    chart.has_title = False
    chart.has_legend = False
    chart.font.name = FONT
    colors = (FOURIER_BG_CARD, FOURIER_PURPLE, FOURIER_MUTED, FOURIER_DARK)
    for series, color in zip(chart.series, colors):
        from .pptx_export import _rgb
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = _rgb(color)
        series.format.line.fill.background()
    # Hide the positioning series while retaining an explicit approved RGB
    # under its alpha value for the generated-chart brand check.
    from pptx.oxml.xmlchemy import OxmlElement
    for rgb in chart.series[0]._element.iter(qn("a:srgbClr")):
        alpha = OxmlElement("a:alpha")
        alpha.set("val", "0")
        rgb.append(alpha)
    for n, (label, color) in enumerate((("Increase", FOURIER_PURPLE),
                                        ("Decrease", FOURIER_MUTED),
                                        ("Reported total", FOURIER_DARK))):
        _text(slide, "■ " + label, .55 + n * 1.65, top + .13, 1.6, .22,
              size=10, color=color).name = "waterfall:legend"
    chart.category_axis.tick_labels.font.name = FONT
    chart.category_axis.tick_labels.font.size = Pt(9)
    chart.value_axis.tick_labels.font.name = FONT
    chart.value_axis.tick_labels.font.size = Pt(9)
    # The exact signed reported components stay visible beneath the bridge.
    table_y = top + 2.38
    row_h = .34
    shape = slide.shapes.add_table(len(items) + 1, 3, Inches(.55), Inches(table_y),
                                   Inches(width), Inches(row_h * (len(items) + 1)))
    shape.name = "table:waterfall:" + ",".join(item.id for item in items)
    shape.table.columns[0].width = Inches(width * .48)
    shape.table.columns[1].width = Inches(width * .20)
    shape.table.columns[2].width = Inches(width * .32)
    for row in shape.table.rows:
        row.height = Inches(row_h)
    unit_label = " ".join(part for part in (items[0].currency, items[0].raw_unit or items[0].unit) if part)
    values = [("Step", "Period", f"Reported value ({unit_label})" if unit_label else "Reported value")]
    for item in items:
        values.append((item.metric_original, format_observation_period(item), item.raw_value or str(item.value)))
    for i, row in enumerate(values):
        for j, content in enumerate(row):
            cell = shape.table.cell(i, j)
            cell.text = content
            cell.text_frame.word_wrap = True
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = FONT
                paragraph.font.size = Pt(11 if i else 12)
                paragraph.font.bold = i == 0
    if table_y + row_h * (len(items) + 1) > presentation.slide_height.inches - 1.0:
        raise ValueError("Waterfall reconciliation exceeds readable slide capacity.")
    pages = sorted({e.page for item in items for e in item.evidence})
    _text(slide, _source_footer(pages), .55, presentation.slide_height.inches - .82,
          width, .20, size=9, color=FOURIER_MUTED)
    notes = {
        "ordered_reconciliation_ids": [item.id for item in items],
        "source_observations": [item.model_dump(mode="json") for item in items],
    }
    slide.notes_slide.notes_text_frame.text = json.dumps(notes, ensure_ascii=False)
    append_header_commentary(presentation, slide_plan, header, pages, notes)
    return slide
