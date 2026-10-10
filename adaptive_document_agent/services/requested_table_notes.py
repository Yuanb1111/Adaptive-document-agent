"""Visible, editable source footnotes with character-preserving pagination."""
import json
from collections import defaultdict

from .text_capacity import wrap_copy


def render_table_notes(presentation, record):
    from .pptx_export import _base_slide, _content_zone, _text, _source_footer, FOURIER_DARK, FOURIER_MUTED
    from .requested_tables import _footer_geometry
    table, page, section = record['table'], record['page'], record['section']
    if not table.raw_footnotes:
        return
    literal = '\n\n'.join(table.raw_footnotes)
    remaining = wrap_copy(literal, 11.4, 11)
    segment = 0
    while remaining:
        slide = _base_slide(presentation, 'Data Index: Source table notes', section.title)
        top, _ = _content_zone(slide)
        footer_top, footer_height, bottom = _footer_geometry(presentation, slide)
        capacity = max(1, int((bottom-top-.12)/.205))
        text = ''.join(remaining[:capacity])
        remaining = remaining[capacity:]
        shape = _text(slide, '', .45, top, 11.7, capacity*.205+.10, size=11, color=FOURIER_DARK)
        # Source annotations bypass audience-copy cleanup. Native paragraphs
        # preserve literal line breaks, spaces, signs and punctuation exactly.
        from pptx.util import Pt
        from .pptx_export import _rgb
        frame = shape.text_frame
        frame.clear()
        for index,line in enumerate(text.split('\n')):
            paragraph = frame.paragraphs[0] if index==0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name, paragraph.font.size = 'Arial', Pt(11)
            paragraph.font.color.rgb = _rgb(FOURIER_DARK)
        shape.name = f'customization:source_footnote:{table.table_id}:{segment}'
        for placeholder in slide.placeholders:
            if placeholder.placeholder_format.idx == 16:
                placeholder.name = 'customization:source_section'
        _text(slide, _source_footer([page.page_number]), .45, footer_top, 10.8, footer_height,
              size=9, color=FOURIER_MUTED).name = 'customization:source_footer'
        slide.notes_slide.notes_text_frame.text = json.dumps({'CUSTOM_SOURCE_NOTES_V1': {
            'table_id': table.table_id, 'page': page.page_number, 'segment': segment,
            'shape_name': shape.name, 'literal_footnotes': table.raw_footnotes}}, ensure_ascii=False)
        segment += 1


def verify_table_notes(presentation, records):
    expected = {r['table'].table_id: '\n\n'.join(r['table'].raw_footnotes)
                for r in records if r['table'].raw_footnotes}
    actual = defaultdict(dict)
    for slide in presentation.slides:
        shapes = {s.name: s for s in slide.shapes if s.name.startswith('customization:source_footnote:')}
        if not shapes:
            continue
        record = json.loads(slide.notes_slide.notes_text_frame.text)['CUSTOM_SOURCE_NOTES_V1']
        identifier, segment = record['table_id'], record['segment']
        if identifier not in expected or segment in actual[identifier] or record['shape_name'] not in shapes:
            raise ValueError('Invalid requested source footnote provenance')
        actual[identifier][segment] = shapes[record['shape_name']].text
    for identifier, literal in expected.items():
        segments = actual[identifier]
        if sorted(segments) != list(range(len(segments))) or ''.join(segments[i] for i in sorted(segments)) != literal:
            raise ValueError('Requested source footnotes were omitted or changed: ' + identifier)
