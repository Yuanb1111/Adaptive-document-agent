"""Render requested physical table fragments without numerical or topic filtering."""

import json
import re
from collections import defaultdict

from pptx.util import Inches

from .text_capacity import wrap_copy


def requested_table_catalog(result):
    requirements = result.profile.report_requirements
    if not requirements:
        return []
    targets = [(req, section) for req in requirements.items
               if req.kind == 'section_tables' and req.resolution == 'resolved'
               for section in req.sections]
    pages = {p.page_number: p for p in result.document.pages}
    records = {}
    for req, section in targets:
        for number in range(section.start_page, section.end_page + 1):
            page = pages.get(number)
            if page is None:
                continue
            for table in page.raw_tables:
                record = records.setdefault(table.table_id, {'table': table, 'page': page,
                    'section': section, 'requirement_ids': []})
                if req.id not in record['requirement_ids']:
                    record['requirement_ids'].append(req.id)
    return list(records.values())


def _grid(table):
    if table.raw_cells:
        raw = [[cell or '' for cell in row] for row in table.raw_cells]
        if table.raw_header_lines:
            # Aligned-text extraction retains header text separately from its
            # raw body grid. Repeat resolved source column context, never style
            # the first data row as if it were the header.
            width = max(len(row) for row in raw)
            headings = []
            for index in range(width):
                period = (table.column_periods[index] if index < len(table.column_periods) else None)
                label = table.headers[index] if index < len(table.headers) else ''
                if re.fullmatch(r'column_\d+', label):
                    label = ''  # Extractor placeholders are not source column meanings.
                if index == 0 and label in {'label', 'row_label'}:
                    label = 'Source row'
                headings.append('\n'.join(str(value) for value in (period, label) if value))
            return [headings, *raw]
        return raw
    return [list(table.headers), *[[cell or '' for cell in row.cells] for row in table.rows]]


def _row_height(cells, widths, font=11):
    return max(.32, max((len(wrap_copy(text, width - .16, font)) * .205 + .10
                         for text, width in zip(cells, widths)), default=.32))


def render_requested_tables(presentation, result):
    """Use editable tables, repeated source headers and explicit cell coordinates."""
    from .pptx_export import (_base_slide, _content_zone, _text, _source_footer, _cell_style,
                              FOURIER_PURPLE, FOURIER_DARK, FOURIER_MUTED, FOURIER_BG_CARD, WHITE)
    from adaptive_document_agent.extraction.table_reconstructor import TableReconstructor
    for ordinal, record in enumerate(requested_table_catalog(result), 1):
        table, page, section = record['table'], record['page'], record['section']
        raw = _grid(table)
        if not raw:
            raise ValueError('Requested source table has no retained cells: ' + table.table_id)
        width = max(len(row) for row in raw)
        raw = [row + [''] * (width - len(row)) for row in raw]
        header_count = (1 if table.raw_header_lines else
                        min(max(1, TableReconstructor.detect_header_row_count(raw)), max(1, len(raw) - 1)))
        groups = ([list(range(width))] if width <= 8 else
                  [[0, *range(start, min(start + 7, width))] for start in range(1, width, 7)])
        table_title = table.table_title or f'Source table {ordinal} (PDF page {page.page_number})'
        context = [table.unit_header] if table.unit_header else []
        # Copy standalone original header qualifications verbatim; no inferred
        # currency/scale or assignment of audit status to particular columns.
        context.extend(line.strip() for line in table.raw_header_lines
                       if line.strip().startswith('(') and line.strip().endswith(')'))
        if context:
            table_title += '\n' + ' | '.join(dict.fromkeys(context))
        label_height = len(wrap_copy(table_title, 11.6, 13)) * .22 + .05
        for columns in groups:
            first_slide = _base_slide(presentation, 'Data Index: Source tables', section.title)
            content_top, _ = _content_zone(first_slide)
            top = content_top + label_height + .10
            first_width = (3.4 if width >= 3 else 11.7 / max(1, width))
            widths = ([11.7] if len(columns) == 1 else
                      [first_width, *[(11.7 - first_width) / (len(columns) - 1)] * (len(columns) - 1)])
            headers = [[row[c] for c in columns] for row in raw[:header_count]]
            header_heights = [_row_height(row, widths) for row in headers]
            available = min(4.15, 6.25 - top) - sum(header_heights)
            if available < .4:
                raise ValueError('Requested table headers exceed readable page capacity: ' + table.table_id)
            # Every source row/column survives. Oversized cells continue across
            # pages at character-preserving wrap boundaries, never ellipses.
            rows = []
            for index in range(header_count, len(raw)):
                pieces = [wrap_copy(raw[index][c], w - .16, 11) for c, w in zip(columns, widths)]
                lines = max(1, int((available - .10) / .205))
                fragments = max((len(p) + lines - 1) // lines for p in pieces)
                for segment in range(fragments):
                    cells = [''.join(p[segment * lines:(segment + 1) * lines]) for p in pieces]
                    rows.append((index, segment, cells, _row_height(cells, widths)))
            batches, current, occupied = [], [], 0.0
            for row in rows:
                if current and occupied + row[3] > available:
                    batches.append(current)
                    current, occupied = [], 0.0
                current.append(row)
                occupied += row[3]
            batches.append(current)
            if len(batches) > 1 and len(batches[-1]) < 3:
                previous, last = batches[-2], batches[-1]
                while len(previous) > len(last) + 1 and sum(row[3] for row in last) + previous[-1][3] <= available:
                    last.insert(0, previous.pop())
            for batch_index, batch in enumerate(batches):
                title = 'Data Index: Source tables'
                slide = first_slide if batch_index == 0 else _base_slide(presentation, title, section.title)
                label = _text(slide, table_title, .45, content_top, 11.7, label_height, size=13, color=FOURIER_PURPLE)
                label.name = 'customization:table_label'
                data = [*headers, *[row[2] for row in batch]]
                heights = [*header_heights, *[row[3] for row in batch]]
                if top + sum(heights) > 6.35:
                    raise ValueError('Requested table cannot fit a readable page: ' + table.table_id)
                shape = slide.shapes.add_table(len(data), len(columns), Inches(.45), Inches(top),
                                               Inches(11.7), Inches(sum(heights)))
                shape.name = f'customization:source_table:{table.table_id}:{columns[0]}:{columns[-1]}:{batch_index}'
                for col, col_width in zip(shape.table.columns, widths):
                    col.width = Inches(col_width)
                coordinates = []
                for r, (cells, height) in enumerate(zip(data, heights)):
                    shape.table.rows[r].height = Inches(height)
                    source_row, segment = (r, 0) if r < header_count else batch[r - header_count][:2]
                    for c, text in enumerate(cells):
                        cell = shape.table.cell(r, c)
                        cell.text = text
                        _cell_style(cell, fill=FOURIER_PURPLE if r < header_count else FOURIER_BG_CARD,
                                    color=WHITE if r < header_count else FOURIER_DARK,
                                    bold=r < header_count, size=11)
                        coordinates.append([r, c, source_row, columns[c], segment])
                _text(slide, _source_footer([page.page_number]) + ' | Original header text and context in notes.',
                      .45, 6.43, 10.8, .18, size=9, color=FOURIER_MUTED).name = 'customization:source_footer'
                provenance = {'table_id': table.table_id, 'page': page.page_number, 'shape_name': shape.name,
                    'section_title': section.title,
                    'requirement_ids': record['requirement_ids'], 'columns': columns, 'cells': coordinates}
                slide.notes_slide.notes_text_frame.text = json.dumps({'CUSTOM_SOURCE_TABLES_V1': [{
                    **provenance, 'source_table': table.model_dump(mode='json'), 'source_page_text': page.text}]},
                    ensure_ascii=False, indent=2)
    pack_requested_tables(presentation)
    verify_requested_tables(presentation, result)


def pack_requested_tables(presentation):
    """Stack short complete table blocks, retaining their labels, grids and source contexts."""
    from .evidence_page_packing import _copy_shape
    from .pptx_export import _source_footer
    retained = []
    for slide in list(presentation.slides):
        tables = [s for s in slide.shapes if s.has_table and s.name.startswith('customization:source_table:')]
        if not tables:
            continue
        label = next(s for s in slide.shapes if s.name == 'customization:table_label')
        incoming = json.loads(slide.notes_slide.notes_text_frame.text)
        height = label.height.inches + .10 + sum(s.height.inches for s in tables)
        previous = None
        # A one-row source table may fit an earlier page in the same chapter,
        # even if the immediately preceding page is full. Keep its own label,
        # physical source citation and grid; never discard it to hide sparsity.
        for candidate in reversed(retained):
            old = json.loads(candidate.notes_slide.notes_text_frame.text)
            if {r['section_title'] for r in old['CUSTOM_SOURCE_TABLES_V1']} != {r['section_title'] for r in incoming['CUSTOM_SOURCE_TABLES_V1']}:
                break
            previous_tables = [s for s in candidate.shapes if s.has_table and s.name.startswith('customization:source_table:')]
            top = max(s.top.inches + s.height.inches for s in previous_tables) + .18
            if top + height <= 6.25:
                previous = candidate
                break
        if previous is None:
            retained.append(slide)
            continue
        label_copy = _copy_shape(previous, label)
        label_copy.top = Inches(top)
        top += label.height.inches + .10
        for table in tables:
            copied = _copy_shape(previous, table)
            copied.top = Inches(top)
            top += table.height.inches
        old['CUSTOM_SOURCE_TABLES_V1'].extend(incoming['CUSTOM_SOURCE_TABLES_V1'])
        previous.notes_slide.notes_text_frame.text = json.dumps(old, ensure_ascii=False, indent=2)
        footer = next(s for s in previous.shapes if s.name == 'customization:source_footer')
        footer.text = _source_footer({r['page'] for r in old['CUSTOM_SOURCE_TABLES_V1']}) + ' | Original header text and context in notes.'
        # Assigning .text creates a new run; preserve the approved footer style.
        from .pptx_export import FOURIER_MUTED, _rgb
        from pptx.util import Pt
        p = footer.text_frame.paragraphs[0]
        p.font.name, p.font.size, p.font.color.rgb = 'Arial', Pt(9), _rgb(FOURIER_MUTED)
        identifier = next(sid for sid in presentation.slides._sldIdLst if sid.id == slide.slide_id)
        presentation.part.drop_rel(identifier.rId)
        presentation.slides._sldIdLst.remove(identifier)
    presentation.part.rename_slide_parts(sid.rId for sid in presentation.slides._sldIdLst)


def verify_requested_tables(presentation, result):
    """Compare actual editable cells to every original coordinate, including duplicates."""
    records = requested_table_catalog(result)
    expected = {record['table'].table_id: _grid(record['table']) for record in records}
    actual = defaultdict(lambda: defaultdict(dict))
    slide_numbers = defaultdict(set)
    for number, slide in enumerate(presentation.slides, 1):
        shapes = [s for s in slide.shapes if s.has_table and s.name.startswith('customization:source_table:')]
        if not shapes:
            continue
        metadata = json.loads(slide.notes_slide.notes_text_frame.text)['CUSTOM_SOURCE_TABLES_V1']
        by_name = {s.name: s for s in shapes}
        for record in metadata:
            identifier = record['table_id']
            if identifier not in expected or record['shape_name'] not in by_name:
                raise ValueError('Unknown requested table in export')
            for r, c, sr, sc, segment in record['cells']:
                text = by_name[record['shape_name']].table.cell(r, c).text
                pieces = actual[identifier][sr, sc]
                if segment in pieces and pieces[segment] != text:
                    raise ValueError('Repeated requested table cell changed')
                pieces[segment] = text
            slide_numbers[identifier].add(number)
    for identifier, grid in expected.items():
        for r, row in enumerate(grid):
            for c, value in enumerate(row):
                pieces = actual[identifier].get((r, c), {})
                if not pieces or sorted(pieces) != list(range(len(pieces))) or ''.join(pieces[i] for i in sorted(pieces)) != value:
                    raise ValueError('Requested source cell was omitted or changed: ' + identifier + f' [{r},{c}]')
    return {identifier: sorted(numbers) for identifier, numbers in slide_numbers.items()}
