"""Native, editable hero chart layout using the application's existing template."""

import re

from adaptive_document_agent.models import ChartPlan

from .single_metric_analysis import SingleMetricAnalysis
from .presentation_labels import qualified_metric_name, qualify_heading


def add_single_metric_slide(presentation, plan: ChartPlan, analysis: SingleMetricAnalysis, *, title: str,
                            narrative: str, definition: dict | None = None, source_context: str = ""):
    # Reuse the established renderer and theme; no provider or document-specific layout.
    from .pptx_export import (
        _add_native_chart, _display_scale, _panel, _source_footer,
        _text, _unit_label, FOURIER_BG_CARD, FOURIER_DARK, FOURIER_MUTED, FOURIER_PURPLE,
    )

    original_title, original_narrative = title, narrative
    series = analysis.observations
    first, last = series[0], series[-1]
    metric_label = re.sub(r"%\s*of\b", "% of", qualified_metric_name(first), flags=re.I)
    title = re.sub(r"%\s*of\b", "% of", title, flags=re.I)
    narrative = re.sub(r"%\s*of\b", "% of", narrative, flags=re.I)
    if first.parent_section or first.dimensions.get("section"):
        qualified = metric_label
        if (first.parent_section or first.dimensions.get("section", "")).casefold() not in title.casefold():
            title = qualified
    title = qualify_heading(title, series)
    from .presentation_trajectory import scoped_direction_title
    title = scoped_direction_title(title, [plan], {item.id: item for item in series})
    scale, scale_label = _display_scale(series, max(abs(o.value) for o in series))
    unit = _unit_label(series, scale_label)
    if (first.unit_family == "currency" or first.unit == "currency") and not first.currency:
        unit = "Currency unspecified" + (f", {scale_label}" if scale_label else "")
    elif scale_label and (first.unit_family == "count" or first.unit == "count"):
        unit = f"{scale_label} of units"

    def value(number: float, *, signed: bool = False, level: bool = False) -> str:
        text = f"{number / scale:+,.1f}" if signed else f"{number / scale:,.1f}"
        return text + "%" if level and analysis.is_percentage else text

    if not narrative.strip():
        values = (f"{value(first.value, level=True)} to {value(last.value, level=True)}"
                  if analysis.is_percentage else f"{value(first.value)} to {value(last.value)} {unit}")
        narrative = f"{metric_label} moved from {values} between {first.period} and {last.period}."
    from .slide_compositor import _base, _lines
    if len(_lines(title, 8.91, 32)) > 2:
        # Use the existing metric label for the title role, retaining the full
        # analytical claim below it and the original narrative in notes.
        title, narrative = metric_label, title
    slide, top = _base(presentation, title, narrative)
    import json
    slide.notes_slide.notes_text_frame.text = json.dumps({
        "planned_title": original_title, "narrative": original_narrative,
        "source_ratio_definition": definition,
        "source_table_context": source_context,
    }, ensure_ascii=False)
    slide.name = "single_metric_hero"
    height = 6.25 - top
    from .text_capacity import wrap_copy
    caption = metric_label if unit == "%" and "%" in metric_label else f"{metric_label} ({unit})"
    # Only a unique, row-bound source footnote may name the denominator. Unknown
    # bases remain unknown; a percentage label alone cannot establish a formula.
    if definition:
        caption += f"\nDefinition: {definition['numerator']} / {definition['denominator']}"
    elif source_context:
        caption += "\n" + source_context
    caption_height = max(.24, len(wrap_copy(caption, 11.35, 10)) * 12.5 / 72 + .06)
    caption_shape = _text(slide, caption, 0.60, top, 11.35, caption_height,
                          size=10, color=FOURIER_MUTED)
    caption_shape.name = "single_metric:definition" if definition else "single_metric:measure"
    from pptx.util import Pt
    for paragraph in caption_shape.text_frame.paragraphs:
        paragraph.line_spacing = Pt(12.5)
    # Leave dedicated slots for four KPIs, annotations and evidence footer.
    chart_offset = caption_height + .02
    chart_height = height - 2.03 - (caption_height - .24)
    hero = plan.model_copy(update={"chart_type": "line" if len(series) >= 3 else "bar", "title": metric_label})
    _add_native_chart(slide, hero, series, (0.60, top + chart_offset, 11.35, chart_height))
    chart_shape = next(shape for shape in slide.shapes if shape.has_chart)
    chart_shape.name = "single_metric_hero_chart"

    if analysis.cagr is not None:
        rate_label, rate_value, rate_period = "CAGR", f"{analysis.cagr:+.1f}%", f"{first.period} to {last.period}"
    elif analysis.percentage_change is not None:
        rate_label, rate_value, rate_period = "Percentage change", f"{analysis.percentage_change:+.1f}%", f"{first.period} to {last.period}"
    else:
        rate_label, rate_value, rate_period = "Reported range", (
            value(analysis.trough.value, level=True) + ' to ' + value(analysis.peak.value, level=True)), 'Within displayed periods'
    cards = [
        ("Start value", value(first.value, level=True), first.period),
        ("End value", value(last.value, level=True), last.period),
        ("Change (pp)" if analysis.is_percentage else "Absolute change", value(analysis.absolute_change, signed=True), "percentage points" if analysis.is_percentage else unit),
        (rate_label, rate_value, rate_period),
    ]
    cards_top = top + chart_offset + chart_height + .12
    for i, (label, number, caption) in enumerate(cards):
        left = 0.60 + i * 2.88
        _panel(slide, left, cards_top, 2.70, 0.92, fill=FOURIER_BG_CARD)
        _text(slide, label, left + 0.12, cards_top + 0.07, 2.46, 0.20, size=10, color=FOURIER_MUTED)
        _text(slide, number, left + 0.12, cards_top + 0.29, 2.46, 0.32, size=20, bold=True, color=FOURIER_PURPLE)
        _text(slide, caption, left + 0.12, cards_top + 0.65, 2.46, 0.20, size=9, color=FOURIER_DARK)
    changes = []
    for change in analysis.changes[-3:]:
        if change.percentage_change is not None:
            detail = f"{change.percentage_change:+.1f}%"
        elif analysis.is_percentage:
            detail = f"{change.absolute_change:+.1f} pp"
        else:
            detail = f"{value(change.absolute_change, signed=True)} {unit}"
        changes.append(f"{change.end_period} {'YoY' if change.is_yoy else 'vs ' + change.start_period}: {detail}")
    _text(slide, "   /   ".join(changes), 0.60, cards_top + 1.02, 11.35, 0.26, size=10.5, color=FOURIER_DARK)
    callout = f"Peak: {analysis.peak.period} ({value(analysis.peak.value, level=True)}). Low: {analysis.trough.period} ({value(analysis.trough.value, level=True)})."
    if analysis.turning_periods:
        callout += " Turning point: " + ", ".join(analysis.turning_periods[:2]) + "."
    _text(slide, callout, 0.60, cards_top + 1.30, 11.35, 0.25, size=10, color=FOURIER_MUTED)
    pages = sorted({e.page for o in series for e in o.evidence})
    _text(slide, _source_footer(pages), 0.60, 6.27, 11.35, 0.22, size=9, color=FOURIER_MUTED)
    return slide
