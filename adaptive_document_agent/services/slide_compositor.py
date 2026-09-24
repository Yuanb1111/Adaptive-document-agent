"""Evidence-driven slide composition, independent of issuer, industry and document type.

The planner owns meaning and evidence selection. This module owns geometry only:
roles determine slots, content determines capacity, and overflow is continued,
never clipped or replaced with invented filler.
"""

from dataclasses import dataclass
import json
import re

from adaptive_document_agent.document_model import display_metric_name
from adaptive_document_agent.models import ChartPlan, PipelineResult, PresentationSlide

from .presentation_style import CHART_TITLE_PT, DARK, FONT, FOOTNOTE_PT, GUTTER, MUTED, PURPLE

COMMENTARY_PT = 18
COMMENTARY_LINE_PT = 22


@dataclass(frozen=True)
class Rect:
    x: float
    y: float
    w: float
    h: float

    def tuple(self):
        return self.x, self.y, self.w, self.h


@dataclass(frozen=True)
class CompositionGeometry:
    charts: list[Rect]
    support: Rect | None
    commentary: Rect | None
    footer: Rect


def compose_geometry(width: float, height: float, top: float, chart_count: int, *,
                     layout: str, has_support: bool, has_commentary: bool,
                     commentary_text: str = "") -> CompositionGeometry:
    """Partition the content zone into non-overlapping slots with a common baseline."""
    left, total = 0.55, width - 1.1
    # Reserve the template's copyright/page-number band separately from sources.
    bottom = height - 1.02
    footer = Rect(left, height - 0.82, total, 0.20)
    support = commentary = None
    chart_bottom = bottom
    if chart_count == 0:
        if has_support and layout == "kpi_band":
            support = Rect(left, top, total, 1.02)
            commentary = Rect(left, top + 1.02 + GUTTER, total, bottom - top - 1.02 - GUTTER) if has_commentary else None
        elif has_support and has_commentary:
            support = Rect(left, top, total * .40, bottom - top)
            commentary = Rect(left + total * .40 + GUTTER, top, total * .60 - GUTTER, bottom - top)
        elif has_support:
            support = Rect(left, top, total, bottom - top)
        elif has_commentary:
            commentary = Rect(left, top, total, bottom - top)
        return CompositionGeometry([], support, commentary, footer)
    if layout == "kpi_band" and has_support:
        support = Rect(left, top, total, 1.02)
        top += 1.02 + GUTTER
        if has_commentary:
            needed = len(_lines(commentary_text, total - .18, COMMENTARY_PT)) * COMMENTARY_LINE_PT / 72 + .23
            band_h = max(.72, min(needed, bottom - top - GUTTER - 1.8))
            commentary = Rect(left, bottom - band_h, total, band_h)
            chart_bottom = commentary.y - GUTTER
        cw = (total - GUTTER * (chart_count - 1)) / chart_count
        charts = [Rect(left + i * (cw + GUTTER), top, cw, chart_bottom - top) for i in range(chart_count)]
        if any(r.h < 1.8 - 1e-6 or r.w < 2.4 for r in charts):
            # Preserve content with the established overflow/continuation path.
            return compose_geometry(width, height, support.y, chart_count,
                layout="two_up", has_support=has_support, has_commentary=has_commentary,
                commentary_text=commentary_text)
        return CompositionGeometry(charts, support, commentary, footer)
    side = chart_count == 1 and (has_support or has_commentary)
    if side:
        cw = (total - GUTTER) * 0.64
        charts = [Rect(left, top, cw, bottom - top)]
        right = left + cw + GUTTER
        rw = total - cw - GUTTER
        if has_support and has_commentary:
            support = Rect(right, top, rw, (bottom - top) * 0.55)
            commentary = Rect(right, support.y + support.h + GUTTER, rw, bottom - support.y - support.h - GUTTER)
        elif has_support:
            support = Rect(right, top, rw, bottom - top)
        else:
            commentary = Rect(right, top, rw, bottom - top)
    else:
        if has_support or has_commentary:
            band_h = 1.22
            if has_commentary:
                text_width = total * .52 - GUTTER if has_support else total
                needed = len(_lines(commentary_text, text_width - .18, COMMENTARY_PT)) * COMMENTARY_LINE_PT / 72 + .23
                band_h = max(band_h, min(needed, bottom - top - GUTTER - 2.2))
            chart_bottom = bottom - band_h - GUTTER
            if has_support and has_commentary:
                support = Rect(left, bottom - band_h, total * 0.48, band_h)
                commentary = Rect(left + total * 0.48 + GUTTER, bottom - band_h, total * 0.52 - GUTTER, band_h)
            elif has_support:
                support = Rect(left, bottom - band_h, total, band_h)
            else:
                commentary = Rect(left, bottom - band_h, total, band_h)
        if chart_count == 2 and layout == "hero_plus_supporting":
            hero_w = (total - GUTTER) * 0.62
            charts = [Rect(left, top, hero_w, chart_bottom - top),
                      Rect(left + hero_w + GUTTER, top, total - hero_w - GUTTER, chart_bottom - top)]
        elif chart_count == 3 and layout == "hero_plus_supporting":
            hero_w = (total - GUTTER) * 0.62
            rh = (chart_bottom - top - GUTTER) / 2
            charts = [Rect(left, top, hero_w, chart_bottom - top),
                      Rect(left + hero_w + GUTTER, top, total - hero_w - GUTTER, rh),
                      Rect(left + hero_w + GUTTER, top + rh + GUTTER, total - hero_w - GUTTER, rh)]
        else:
            cw = (total - GUTTER * (chart_count - 1)) / max(chart_count, 1)
            charts = [Rect(left + i * (cw + GUTTER), top, cw, chart_bottom - top) for i in range(chart_count)]
    if any(r.h < 1.4 or r.w < 2.4 for r in charts):
        raise ValueError("Slide composition cannot fit readable charts; split the evidence into additional pages.")
    return CompositionGeometry(charts, support, commentary, footer)


def _lines(text: str, width: float, size: float) -> list[str]:
    from .text_capacity import wrap_copy
    return wrap_copy(text, width, size)


def _put_text(slide, text: str, rect: Rect, *, size=12, bold=False, color=DARK):
    from .pptx_export import _text
    shape = _text(slide, text, *rect.tuple(), size=size, bold=bold, color=color)
    shape.name = f"composed:text:{shape.shape_id}"
    return shape


def _put_commentary(slide, text: str, rect: Rect) -> str:
    capacity = max(1, int((rect.h - .10) * 72 / COMMENTARY_LINE_PT))
    # Partition the original string, preserving paragraphs and even long IDs.
    # Inserted line breaks would split phrases; joining wrapped words would
    # erase paragraph boundaries and could corrupt identifiers at a page break.
    lines = _lines(text, rect.w - .18, COMMENTARY_PT)
    from .pptx_export import _rule
    cut = sum(len(line) for line in lines[:capacity])
    if cut < len(text):
        # Prefer a complete paragraph/sentence over fragments such as a currency
        # prefix separated from its amount. Oversized single sentences still
        # make progress and retain every character on the following page.
        boundaries = [m.end() for m in re.finditer(r"\n+|(?<=[.!?。！？])\s+", text[:cut])]
        if boundaries:
            cut = boundaries[-1]
    _rule(slide, rect.x, rect.y + .03, .035, min(rect.h - .06, .62), PURPLE)
    body = _put_text(slide, text[:cut],
                     Rect(rect.x + .18, rect.y, rect.w - .18, rect.h), size=COMMENTARY_PT)
    from pptx.util import Pt
    for paragraph in body.text_frame.paragraphs:
        paragraph.line_spacing = Pt(COMMENTARY_LINE_PT)
    # Reflow on the next page, whose text column can be wider than this slot.
    return text[cut:]


def _base(presentation, title, message):
    from .pptx_export import _base_slide, _content_zone, _rule, _rgb
    from pptx.util import Inches, Pt
    slide = _base_slide(presentation, title, message)
    # Preserve complete planned copy. The legacy helper intentionally truncates
    # strings, which is inappropriate for the new composition contract.
    for ph in slide.placeholders:
        idx = ph.placeholder_format.idx
        if idx in (14, 15) and ph.has_text_frame:
            ph.left = Inches(.55)
            ph.top = Inches(.52)
            ph.width = Inches(presentation.slide_width.inches - 1.1)
            ph.text_frame.margin_left = ph.text_frame.margin_right = Inches(.02)
            ph.text_frame.margin_top = ph.text_frame.margin_bottom = Inches(.01)
            lines = _lines(title, ph.width.inches - 0.15, 24)
            if len(lines) > 3:
                raise ValueError("Presentation title exceeds readable capacity; shorten the planned title.")
            ph.text = title
            ph.height = Inches(max(0.45, len(lines) * 0.38))
            for p in ph.text_frame.paragraphs:
                p.font.size = Pt(24)
                p.font.bold = True
                p.font.color.rgb = _rgb(DARK)
        if idx == 16 and ph.has_text_frame:
            title_bottom = max((s.top.inches + s.height.inches for s in slide.placeholders if s.placeholder_format.idx in (14, 15)), default=1.0)
            ph.text = message
            ph.left = Inches(.55)
            ph.top = Inches(title_bottom + 0.20)
            ph.width = Inches(presentation.slide_width.inches - 1.1)
            lines = _lines(message, ph.width.inches - 0.15, 11)
            if len(lines) > 3:
                raise ValueError("Presentation subtitle exceeds readable capacity; move supporting details to commentary.")
            ph.height = Inches(max(0.24, len(lines) * 0.19))
            for p in ph.text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.color.rgb = _rgb(MUTED)
    title_bottom = max((s.top.inches + s.height.inches for s in slide.placeholders if s.placeholder_format.idx in (14, 15)), default=1.0)
    _rule(slide, .55, title_bottom + .07, 1.35, .035, PURPLE)
    slide.shapes[-1].name = "decoration:title_rule"
    slide._ada_colors = getattr(presentation, "_ada_colors", {})
    return slide, _content_zone(slide)[0]


def render_composed_slide(presentation, slide_plan: PresentationSlide, charts: list[ChartPlan],
                          result: PipelineResult, index) -> list:
    """Render all requested charts, explicit KPI/table facts and retained commentary."""
    from .pptx_export import _add_native_chart, _source_footer, _unit_label, _display_source_unit
    from adaptive_document_agent.document_model.metric_semantic_classifier import classify_metric, format_metric_display_value
    from adaptive_document_agent.document_model.period_semantic_validator import format_observation_period
    from pptx.util import Inches, Pt

    by_id = {c.id: c for c in charts}
    ordered = list(dict.fromkeys([cid for b in slide_plan.visual_blocks if b.role == "hero" for cid in b.chart_ids] + [c.id for c in charts]))
    charts = [by_id[cid] for cid in ordered]
    if slide_plan.theme_id and len(charts) == 3:
        from adaptive_document_agent.validation.layout_qa import _is_cramped_multi_chart_slide
        if _is_cramped_multi_chart_slide(slide_plan, by_id, index):
            rendered = []
            for part, group in enumerate((charts[:2], charts[2:])):
                cids = {c.id for c in group}
                blocks = [b.model_copy(update={"chart_ids": [cid for cid in b.chart_ids if cid in cids]})
                          for b in slide_plan.visual_blocks
                          if any(cid in cids for cid in b.chart_ids) or (not part and not b.chart_ids)]
                physical = slide_plan.model_copy(update={
                    "chart_ids": [c.id for c in group], "visual_blocks": blocks,
                    "title": slide_plan.title + (" (continued)" if part else ""),
                    "bullets": [] if part else slide_plan.bullets,
                    "insight_ids": [] if part else slide_plan.insight_ids,
                    "observation_ids": [] if part else slide_plan.observation_ids,
                    "layout": "two_up" if not part else "chart_plus_commentary"})
                rendered.extend(render_composed_slide(presentation, physical, group, result, index))
            return rendered
    chart_obs = {oid for c in charts for oid in c.observation_ids}
    explicit = [oid for b in slide_plan.visual_blocks if b.role in {"kpi", "table"} for oid in b.observation_ids]
    # References establish provenance, not a request to display every raw row.
    # Retain the legacy data-only fallback when no visual selection exists.
    legacy_data_layout = slide_plan.layout in {"chart_plus_kpis", "chart_with_data", "data_overview"}
    extra = ([oid for oid in slide_plan.observation_ids if oid not in chart_obs]
             if not slide_plan.visual_blocks and (not charts or legacy_data_layout) else [])
    support_ids = list(dict.fromkeys([*explicit, *extra]))
    support = [index.get(oid) for oid in support_ids if index.get(oid)]
    if any(o.value is None or not o.evidence or o.validation_status not in {"valid", "partially_valid"} for o in support):
        raise ValueError("Supporting KPI/table evidence is incomplete or invalid.")
    insight_ids = list(dict.fromkeys(slide_plan.insight_ids + [iid for b in slide_plan.visual_blocks for iid in b.insight_ids]))
    insight_map = {i.id: i for i in result.insights}
    visible_insights = [iid for b in slide_plan.visual_blocks if b.role == "commentary" for iid in b.insight_ids]
    if not slide_plan.bullets and not visible_insights:
        visible_insights = insight_ids
    commentary = list(dict.fromkeys([*slide_plan.bullets, *(insight_map[i].narrative for i in visible_insights if i in insight_map)]))
    text = "\n".join(t for t in commentary if t.strip() and t.strip() != slide_plan.message.strip())
    pages = sorted(set(slide_plan.source_pages) | {e.page for o in support for e in o.evidence}
                   | {e.page for c in charts for oid in [*c.observation_ids, *c.total_observation_ids] if index.get(oid) for e in index.get(oid).evidence}
                   | {e.page for iid in insight_ids if iid in insight_map for e in insight_map[iid].evidence})
    slide, top = _base(presentation, slide_plan.title, slide_plan.message)
    slide.name = f"composed_{slide_plan.layout}"
    reference_ids = set(slide_plan.observation_ids) | chart_obs | set(support_ids)
    reference_ids.update(oid for b in slide_plan.visual_blocks for oid in b.observation_ids)
    reference_ids.update(oid for c in charts for oid in c.total_observation_ids)
    # Keep exact values, qualifiers and page-level evidence recoverable without
    # turning supporting references into duplicate audience-facing pages.
    slide.notes_slide.notes_text_frame.text = json.dumps({
        "slide_id": slide_plan.id,
        "observations": [index.get(oid).model_dump(mode="json") for oid in sorted(reference_ids) if index.get(oid)],
        "insights": [insight_map[i].model_dump(mode="json") for i in insight_ids if i in insight_map],
        "source_pages": pages,
    }, ensure_ascii=False, indent=2)
    geometry = compose_geometry(presentation.slide_width.inches, presentation.slide_height.inches, top, len(charts),
        layout=slide_plan.layout, has_support=bool(support), has_commentary=bool(text), commentary_text=text)
    headings = []
    for chart in charts:
        heading = next((b.title for b in slide_plan.visual_blocks if b.chart_ids == [chart.id] and b.title), chart.title)
        values = [index.get(oid) for oid in chart.observation_ids if index.get(oid)]
        qualified = {display_metric_name(o) for o in values}
        if len(qualified) == 1 and not any(o.category_dimensions for o in values) and any(o.parent_section or o.dimensions.get("section") for o in values):
            heading = next(iter(qualified))
        headings.append(re.sub(r"(?i)^Adjusted for Adjusted\b", "Adjusted", heading))
    shared_heading_h = max([.28] + [len(_lines(t, r.w, CHART_TITLE_PT)) * .23
                                      for t, r in zip(headings, geometry.charts)])
    for chart, rect, title in zip(charts, geometry.charts, headings):
        values = [index.get(oid) for oid in chart.observation_ids if index.get(oid)]
        qualified = {display_metric_name(o) for o in values}
        # Retain an explicitly extracted parent even when the planned short
        # label names only the child. Do not turn grants into expense metrics.
        if len(qualified) == 1 and not any(o.category_dimensions for o in values) and any(o.parent_section or o.dimensions.get("section") for o in values):
            title = re.sub(r"(?i)^Adjusted for Adjusted\b", "Adjusted", next(iter(qualified)))
        title_lines = _lines(title, rect.w, CHART_TITLE_PT)
        if len(title_lines) > 3:
            raise ValueError(f"Chart title exceeds readable capacity: {chart.id}")
        heading_h = shared_heading_h
        _put_text(slide, title, Rect(rect.x, rect.y, rect.w, heading_h), size=CHART_TITLE_PT, bold=True)
        totals = [index.get(oid) for oid in chart.total_observation_ids if index.get(oid)]
        bounds = (rect.x, rect.y + heading_h + 0.22, rect.w, rect.h - heading_h - 0.22)
        if bounds[3] < 1.25:
            raise ValueError("Chart labels cannot fit; split the planned charts into additional slides.")
        scale, scale_label = _add_native_chart(slide, chart, values, bounds, compact=rect.w < 4.5, totals=totals)
        unit = "Share (%)" if chart.chart_type == "stacked_percent" else _unit_label(values, scale_label)
        if chart.chart_type == "line" and any(o.as_of_date or o.period_basis == "point_in_time" for o in values):
            unit += " | Dates shown as equally spaced categories"
        if chart.chart_type == "doughnut":
            unit = f"{values[0].period} | {unit}"
        _put_text(slide, unit, Rect(rect.x, rect.y + heading_h, rect.w, 0.20), size=FOOTNOTE_PT, color=MUTED)
        list(s for s in slide.shapes if s.has_chart)[-1].name = f"chart:{chart.id}"

    def draw_support(owner, items, rect, table=False):
        # Keep a compact period series together. If stacked KPI cards do not
        # fit, use a readable complete table in the SAME slot before paginating.
        horizontal = rect.w > 5.5
        capacity = min(4, max(1, int(rect.w / 2.1))) if horizontal else max(1, int(rect.h / 0.98))
        from .presentation_evidence import evidence_groups
        if table:
            # Collapse only exact corroborating values in the visible table.
            # Every original ID and source remains in the slide's notes above.
            displayed = []
            for group in evidence_groups(items):
                seen_values = set()
                for o in group:
                    key = (o.period, o.value, o.audited_status, o.id if not o.period else None)
                    if key not in seen_values:
                        seen_values.add(key)
                        displayed.append(o)
            items = displayed
        if not table and len(items) > capacity and len(evidence_groups(items)) == 1:
            name = display_metric_name(items[0])
            title_h = len(_lines(name, rect.w, 14)) * .24 + .12
            rows = []
            for o in items:
                semantic = classify_metric(name, unit=o.unit, raw_unit=o.raw_unit, value=o.value)
                value = format_metric_display_value(o.raw_value, o.value, semantic, raw_unit=_display_source_unit(o), currency=o.currency, compact=True)
                period = format_observation_period(o)
                height = max(.38, max(len(_lines(period, rect.w * .34 - .15, 14)), len(_lines(value, rect.w * .66 - .15, 14))) * .24 + .12)
                rows.append((period, value, height))
            if title_h + sum(row[2] for row in rows) <= rect.h:
                _put_text(owner, name, Rect(rect.x, rect.y, rect.w, title_h), size=14, bold=True)
                shape = owner.shapes.add_table(len(rows), 2, Inches(rect.x), Inches(rect.y + title_h), Inches(rect.w), Inches(sum(row[2] for row in rows)))
                shape.name = "table:" + ",".join(o.id for o in items)
                shape.table.columns[0].width = Inches(rect.w * .34)
                shape.table.columns[1].width = Inches(rect.w * .66)
                for i, (period, value, height) in enumerate(rows):
                    shape.table.rows[i].height = Inches(height)
                    for j, content in enumerate((period, value)):
                        cell = shape.table.cell(i, j)
                        cell.text = content
                        for p in cell.text_frame.paragraphs:
                            p.font.name, p.font.size = FONT, Pt(14)
                return []
        if table:
            # A short single-series support band shows ALL periods horizontally,
            # rather than losing intermediate changes or creating a near-empty continuation.
            if horizontal and 1 < len(items) <= 6 and len(evidence_groups(items)) == 1:
                name = display_metric_name(items[0])
                name = re.sub(r"(?i)^Adjusted for Adjusted\b", "Adjusted", name)
                cells = [["Metric", *[format_observation_period(o) for o in items]], [name]]
                for o in items:
                    semantic = classify_metric(name, unit=o.unit, raw_unit=o.raw_unit, value=o.value)
                    cells[1].append(format_metric_display_value(o.raw_value, o.value, semantic,
                        raw_unit=_display_source_unit(o), currency=o.currency, compact=True))
                widths = [rect.w * .34] + [rect.w * .66 / len(items)] * len(items)
                heights = [max(.4, max(len(_lines(v, w - .15, 14)) for v, w in zip(row, widths)) * .24 + .1)
                           for row in cells]
                if sum(heights) <= rect.h:
                    shape = owner.shapes.add_table(2, len(items) + 1, Inches(rect.x), Inches(rect.y), Inches(rect.w), Inches(sum(heights)))
                    shape.name = "table:" + ",".join(o.id for o in items)
                    for j, w in enumerate(widths):
                        shape.table.columns[j].width = Inches(w)
                    for i, row in enumerate(cells):
                        shape.table.rows[i].height = Inches(heights[i])
                        for j, value in enumerate(row):
                            cell = shape.table.cell(i, j)
                            cell.text = value
                            for p in cell.text_frame.paragraphs:
                                p.font.name, p.font.size = FONT, Pt(14)
                    return []
            rows, heights, shown = [], [0.4], []
            for o in items:
                name = display_metric_name(o)
                semantic = classify_metric(name, unit=o.unit, raw_unit=o.raw_unit, value=o.value)
                display = format_metric_display_value(o.raw_value, o.value, semantic, raw_unit=_display_source_unit(o), currency=o.currency, compact=True)
                period_label = format_observation_period(o)
                if not period_label:
                    column = next((e.column_label for e in o.evidence if e.column_label), "")
                    period_label = "Unspecified" + (f" ({column})" if column else "")
                cells = (name, period_label, display)
                row_h = max(0.4, max(len(_lines(v, rect.w * w - 0.15, 14)) for v, w in zip(cells, (0.50, 0.20, 0.30))) * 0.24 + 0.10)
                if sum(heights) + row_h > rect.h:
                    break
                rows.append(cells)
                heights.append(row_h)
                shown.append(o)
            if not shown:
                if rect.h < 2:
                    return items
                raise ValueError("Exact-data row exceeds readable page capacity; shorten its display label without changing evidence.")
            shape = owner.shapes.add_table(len(shown) + 1, 3, Inches(rect.x), Inches(rect.y), Inches(rect.w), Inches(sum(heights)))
            shape.name = "table:" + ",".join(o.id for o in shown)
            table_obj = shape.table
            table_obj.columns[0].width = Inches(rect.w * 0.50)
            table_obj.columns[1].width = Inches(rect.w * 0.20)
            table_obj.columns[2].width = Inches(rect.w * 0.30)
            for j, label in enumerate(("Metric", "Period", "Value")):
                table_obj.cell(0, j).text = label
            for i, cells in enumerate(rows, 1):
                for j, label in enumerate(cells):
                    table_obj.cell(i, j).text = label
            for row, h in zip(table_obj.rows, heights):
                row.height = Inches(h)
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        p.font.name = FONT
                        p.font.size = Pt(14)
            return items[len(shown):]
        shown = items[:capacity]
        for i, o in enumerate(shown):
            w = (rect.w - GUTTER * (len(shown) - 1)) / len(shown) if horizontal else rect.w
            x = rect.x + i * (w + GUTTER) if horizontal else rect.x
            y = rect.y if horizontal else rect.y + i * 0.98
            name = display_metric_name(o)
            semantic = classify_metric(name, unit=o.unit, raw_unit=o.raw_unit, value=o.value)
            value = format_metric_display_value(o.raw_value, o.value, semantic, raw_unit=_display_source_unit(o), currency=o.currency, compact=True)
            if len(_lines(name, w - .14, 11)) > 2 or len(_lines(value, w - .14, 20)) > 1:
                return items[i:]  # Preserve long-label facts on a continuation table.
            from .pptx_export import _rule
            from .presentation_style import semantic_color
            accent = getattr(presentation, "_ada_colors", {}).get(name, semantic_color(name))
            _rule(owner, x, y, .035, .88, accent)
            _put_text(owner, name, Rect(x + .14, y, w - .14, 0.35), size=11, color=MUTED)
            shape = _put_text(owner, value, Rect(x + .14, y + 0.37, w - .14, 0.34), size=20, bold=True, color=accent)
            if shape is not None:
                shape.name = f"kpi:{o.id}"
            _put_text(owner, format_observation_period(o), Rect(x + .14, y + 0.74, w - .14, 0.18), size=FOOTNOTE_PT, color=MUTED)
        return items[capacity:]

    table_ids = {oid for b in slide_plan.visual_blocks if b.role == "table" for oid in b.observation_ids}
    kpi_ids = {oid for b in slide_plan.visual_blocks if b.role == "kpi" for oid in b.observation_ids}
    support_start = len(slide.shapes)
    if support and table_ids and kpi_ids:
        kpis = [o for o in support if o.id in kpi_ids]
        rows = [o for o in support if o.id in table_ids or o.id not in kpi_ids]
        rect = geometry.support
        if rect.h >= 2.5:
            kpi_h = 0.98 if rect.w > 5.5 else min(1.96, 0.98 * len(kpis))
            overflow = draw_support(slide, kpis, Rect(rect.x, rect.y, rect.w, kpi_h))
            overflow += draw_support(slide, rows, Rect(rect.x, rect.y + kpi_h + GUTTER, rect.w, rect.h - kpi_h - GUTTER), table=True)
        else:
            # Keep the requested KPI role on the main page; continue exact-data
            # rows separately instead of collapsing both roles into one table.
            overflow = draw_support(slide, kpis, rect) + rows
    else:
        overflow = draw_support(slide, support, geometry.support, table=bool(table_ids)) if support else []
    commentary_rect = geometry.commentary
    side_commentary = bool(charts) and all(
        commentary_rect is not None and (commentary_rect.x >= c.x + c.w or commentary_rect.x + commentary_rect.w <= c.x)
        for c in geometry.charts)
    if support and text and (not charts or side_commentary) and commentary_rect.y > geometry.support.y + geometry.support.h:
        # Reclaim unused support height before adding a prose-only continuation.
        used_bottom = max((s.top.inches + s.height.inches for s in list(slide.shapes)[support_start:]), default=geometry.support.y)
        new_top = min(commentary_rect.y, used_bottom + GUTTER)
        commentary_rect = Rect(commentary_rect.x, new_top, commentary_rect.w,
                               commentary_rect.y + commentary_rect.h - new_top)
    remaining = _put_commentary(slide, text, commentary_rect) if text else ""
    visible_observations = support + [index.get(oid) for oid in chart_obs if index.get(oid)]
    footnote = _source_footer(pages)
    if any("*" in format_observation_period(o) for o in visible_observations):
        footnote += " | * Unaudited"
    _put_text(slide, footnote, geometry.footer, size=FOOTNOTE_PT, color=MUTED)
    slides = [slide]
    while overflow or remaining:
        continuation, ctop = _base(presentation, slide_plan.title + " (continued)", "")
        continuation.name = "composed_continuation"
        bottom = presentation.slide_height.inches - 1.02
        full = Rect(0.55, ctop, presentation.slide_width.inches - 1.1, bottom - ctop)
        if overflow:
            overflow = draw_support(continuation, overflow, full, table=True)
            used_bottom = max((s.top.inches + s.height.inches for s in continuation.shapes if s.name.startswith("table:")), default=ctop)
            text_top = used_bottom + GUTTER
            if not overflow and remaining and bottom - text_top >= .75:
                remaining = _put_commentary(continuation, remaining, Rect(full.x, text_top, full.w, bottom - text_top))
        elif remaining:
            remaining = _put_commentary(continuation, remaining, full)
        _put_text(continuation, footnote, geometry.footer, size=FOOTNOTE_PT, color=MUTED)
        slides.append(continuation)
    return slides


def validate_composed_geometry(presentation) -> None:
    """Block actual out-of-bounds or overlapping composition objects before save.

    Template background shapes intentionally overlap. Only owned content objects
    enter this check; rendered visual review remains a separate quality gate.
    """
    from .qa_reporter import CriticalQAError
    tolerance = 0.025 * 914400
    for slide in presentation.slides:
        if not slide.name.startswith("composed_"):
            continue
        shapes = [s for s in slide.shapes if s.name.startswith(("composed:", "chart:", "kpi:", "table:"))]
        for shape in shapes:
            if (shape.left < -tolerance or shape.top < -tolerance
                    or shape.left + shape.width > presentation.slide_width + tolerance
                    or shape.top + shape.height > presentation.slide_height + tolerance):
                raise CriticalQAError(f"Composed slide content is outside the page: {shape.name}")
        for i, a in enumerate(shapes):
            for b in shapes[i + 1:]:
                overlap_w = min(a.left + a.width, b.left + b.width) - max(a.left, b.left)
                overlap_h = min(a.top + a.height, b.top + b.height) - max(a.top, b.top)
                if overlap_w > tolerance and overlap_h > tolerance:
                    raise CriticalQAError(f"Composed slide content overlaps: {a.name}, {b.name}")
