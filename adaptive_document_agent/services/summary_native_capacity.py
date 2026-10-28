"""Measure introduction candidates using the same renderer and template as export."""
from functools import lru_cache
from io import BytesIO
import json
from pathlib import Path


@lru_cache(maxsize=4)
def _template_bytes(path, revision, size):
    return Path(path).read_bytes()


@lru_cache(maxsize=512)
def _measure(path, revision, size, title, encoded):
    from pptx import Presentation
    from adaptive_document_agent.models.summary_display import SummaryDisplay
    from .presentation_brief import BriefItem, render_profile
    deck = Presentation(BytesIO(_template_bytes(path, revision, size)))
    for slide in list(deck.slides._sldIdLst):
        deck.part.drop_rel(slide.rId)
        deck.slides._sldIdLst.remove(slide)
    items = [BriefItem(row['label'], row['text'], [], display=SummaryDisplay.model_validate(row['display']))
             for row in json.loads(encoded)]
    try:
        rendered = render_profile(deck, title, items)
    except ValueError:
        # A continuation header can overflow even when the opening header
        # fits. Let the existing bounded layout repair shorten/split its copy.
        return 0, False
    complete = ''.join(shape.text for slide in rendered for shape in slide.shapes if shape.name == 'brief:body')
    from .display_copy_tokens import display_text
    # PowerPoint stores an authored newline as a soft line break. Accept only
    # that existing display equivalence, not changed words or removed spaces.
    return len(rendered), display_text(complete) == display_text(''.join(item.text for item in items))


def native_summary_capacity(page):
    """No model, file export or approval: exact native count plus copy preservation."""
    from .pptx_export import _resolve_template_path
    from .summary_display_copy import summary_display
    path = _resolve_template_path()
    revision = path.stat()
    encoded = json.dumps([dict(label=item.label, text=item.text,
        display=summary_display(item, item.text).model_dump(mode='json')) for item in page.items],
        ensure_ascii=False, sort_keys=True)
    return _measure(str(path), revision.st_mtime_ns, revision.st_size, page.title, encoded)
