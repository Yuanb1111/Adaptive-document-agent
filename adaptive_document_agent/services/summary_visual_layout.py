"""Native summary layouts that rearrange complete reviewed copy, never its meaning.

The writer can select semantic structure during its existing authoring request.
Older accepted drafts receive only lossless punctuation-based paragraph breaks.
No metric, date, caption, relationship or graphical assertion is inferred here.
"""
import re

from pptx.util import Pt

from adaptive_document_agent.models.summary_display import SummaryDisplay
from .fourier_brand import WHITE
from .presentation_brief import BODY_PT, HEADING_PT, LINE_HEIGHT_FACTOR, BriefItem, _copy_height, _heading_height


def summary_brief_items(result, page):
    """Reuse approved/localized copy and carry optional writer-selected structure."""
    from .localized_summary_copy import summary_item_copy
    items = []
    for index, item in enumerate(page.items):
        text = summary_item_copy(result, page, index)
        # Anchors address the exact writer copy. A translation has its own
        # reviewed wording, so it can only receive mechanical paragraph breaks.
        display = item.display if text == item.text and item.display else SummaryDisplay()
        items.append(BriefItem(item.label, text, item.source_pages, display=display))
    return items


def display_segments(text, display=None):
    """Return exact contiguous text slices; malformed anchors cannot change copy."""
    anchors = display.break_after if display else []
    ends, start = [], 0
    for anchor in anchors:
        if (not anchor or len(anchor) > 160 or text.count(anchor) != 1
                or (offset := text.find(anchor, start)) < 0):
            ends = []
            break
        end = offset + len(anchor)
        if end < len(text) and not text[end].isspace():
            ends = []
            break
        ends.append(end)
        start = end
    if not ends:
        # Sentence boundaries provide whitespace only, not semantic labels.
        # Keep parenthetical qualifications attached to the same assertion.
        depth = 0
        for match in re.finditer(r'[()]|[.!?。！？](?:\s+|(?=[\u3400-\u9fff]))', text):
            token = match.group()
            if token == '(':
                depth += 1
            elif token == ')':
                depth = max(0, depth - 1)
            elif not depth:
                previous = text[:match.start()].rsplit(None, 1)[-1:]
                # Compact initials and abbreviations retain their sentence.
                if previous and re.fullmatch(r'(?:[A-Z]|Mr|Mrs|Ms|Dr|Prof|No|vs)', previous[0]):
                    continue
                ends.append(match.end())
    ends = sorted(set([0, *ends, len(text)]))
    pieces = [text[left:right] for left, right in zip(ends, ends[1:]) if text[left:right]]
    # A trailing whitespace-only piece belongs to the preceding assertion.
    if len(pieces) > 1 and not pieces[-1].strip():
        pieces[-2] += pieces.pop()
    return pieces or [text]


def _column_geometry(parts, width, capacity, *, gaps=(.12, .08, .04)):
    """Measure all breaks before drawing, retaining author reading order."""
    heights = [_copy_height(text, width, BODY_PT) + .02 for text in parts]
    for gap in gaps:
        used = sum(heights) + gap * max(0, len(parts) - 1)
        if used <= capacity + 1e-6:
            return heights, used, gap
    return None


def _columns(parts, width, capacity, *, gaps=(.12, .08, .04)):
    gap = .44
    col_width = (width - gap) / 2
    choices = []
    for split in range(1, len(parts)):
        left = _column_geometry(parts[:split], col_width, capacity, gaps=gaps)
        right = _column_geometry(parts[split:], col_width, capacity, gaps=gaps)
        if left and right:
            choices.append((max(left[1], right[1]), abs(left[1] - right[1]), split, left, right))
    if not choices:
        return None
    _, _, split, left, right = min(choices)
    return [(0, col_width, parts[:split], left[0], left[2]),
            (col_width + gap, col_width, parts[split:], right[0], right[2])]


def _render_pair_panels(slide, items, top, width, capacity):
    """Group two independent authored items visually, without linking their meanings."""
    from .pptx_export import _text, _rule, FOURIER_DARK, FOURIER_PURPLE
    gap, padding = .36, .20
    panel_width = (width - gap) / 2
    text_width = panel_width - padding * 2 - .04
    plans = []
    for item in items:
        heading_height = _heading_height(item.title, text_width)
        body_capacity = capacity - padding * 2 - heading_height - .10
        parts = display_segments(item.text, item.display)
        geometry = _column_geometry(parts, text_width, body_capacity)
        if geometry is None:
            return False
        plans.append((item, heading_height, parts, geometry))
    for index, (item, heading_height, parts, geometry) in enumerate(plans):
        x = .65 + index * (panel_width + gap)
        # Capacity is a fit limit, not the height of a content card. Each
        # independent item keeps its own measured copy and bottom padding.
        panel_height = padding * 2 + heading_height + .10 + geometry[1]
        _rule(slide, x, top, panel_width, panel_height, WHITE)
        slide.shapes[-1].name = 'summary:panel'
        _rule(slide, x, top, .035, panel_height, FOURIER_PURPLE)
        slide.shapes[-1].name = 'summary:accent'
        y = top + padding
        heading = _text(slide, item.title, x + padding + .04, y, text_width, heading_height,
                        size=HEADING_PT, bold=True, color=FOURIER_PURPLE)
        heading.name = 'brief:heading'
        heading.text_frame.paragraphs[0].line_spacing = Pt(HEADING_PT * LINE_HEIGHT_FACTOR)
        y += heading_height + .10
        heights, _, gap_y = geometry
        for text, height in zip(parts, heights):
            body = _text(slide, text, x + padding + .04, y, text_width, height,
                         size=BODY_PT, color=FOURIER_DARK)
            body.name = 'brief:body'
            for paragraph in body.text_frame.paragraphs:
                paragraph.line_spacing = Pt(BODY_PT * LINE_HEIGHT_FACTOR)
            y += height + gap_y
    return True


def render_structured_profile_group(presentation, title, items, *, notes=''):
    """Keep each item's supported structure local within a multi-item page.

    Separate source assertions never acquire a shared chronology or relationship
    because they happen to occupy the same slide. If the complete native groups
    do not fit, the caller retains its established complete-copy pagination.
    """
    if len(items) == 1:
        return render_structured_profile(presentation, title, items[0], notes=notes)
    from .slide_compositor import _base
    from .pptx_export import _source_footer, _text, _rule, FOURIER_DARK, FOURIER_MUTED, FOURIER_PURPLE

    slide, top = _base(presentation, title, '')
    width = presentation.slide_width.inches - 1.3
    capacity = presentation.slide_height.inches - 1.28 - top
    if len(items) == 2 and _render_pair_panels(slide, items, top, width, capacity):
        pages = sorted({page for item in items for page in item.pages})
        _text(slide, _source_footer(pages), .55, presentation.slide_height.inches - .82,
              presentation.slide_width.inches - 1.1, .20, size=9, color=FOURIER_MUTED)
        full_copy = '\n\n'.join(item.title + '\n' + item.text for item in items)
        slide.notes_slide.notes_text_frame.text = '\n\n'.join(value for value in (notes, full_copy) if value)
        return slide
    selected = None
    for paragraph_gap, item_gap in ((.12, .20), (.08, .14), (.04, .10)):
        plans, total = [], 0
        for item in items:
            parts = display_segments(item.text, item.display)
            layout = item.display.layout if item.display else 'paragraphs'
            heading = _heading_height(item.title, width)
            columns = (_columns(parts, width, capacity, gaps=(paragraph_gap,))
                       if layout in {'comparison', 'relationship', 'key_metrics'} else None)
            if columns is None:
                inset = .34 if layout == 'timeline' else 0
                single = _column_geometry(parts, width - inset, capacity, gaps=(paragraph_gap,))
                if single is None:
                    plans = []
                    break
                columns = [(inset, width - inset, parts, single[0], single[2])]
            body_height = max(sum(heights) + gap * max(0, len(copy) - 1)
                              for _, _, copy, heights, gap in columns)
            height = heading + (.08 if heading else 0) + body_height
            plans.append((item, layout, heading, columns, height))
            total += height
        if plans and total + item_gap * (len(plans) - 1) <= capacity + 1e-6:
            selected = plans, item_gap
            break
    if selected is None:
        slide_id = presentation.slides._sldIdLst[-1]
        presentation.part.drop_rel(slide_id.rId)
        presentation.slides._sldIdLst.remove(slide_id)
        return None
    plans, item_gap = selected
    cursor = top
    for item, layout, heading_height, columns, height in plans:
        if item.title:
            heading = _text(slide, item.title, .65, cursor, width, heading_height,
                            size=HEADING_PT, bold=True, color=FOURIER_PURPLE)
            heading.name = 'brief:heading'
            heading.text_frame.paragraphs[0].line_spacing = Pt(HEADING_PT * LINE_HEIGHT_FACTOR)
        body_top = cursor + heading_height + (.08 if heading_height else 0)
        for x, column_width, copy, heights, gap_y in columns:
            y = body_top
            for text, body_height in zip(copy, heights):
                if layout == 'timeline':
                    _rule(slide, .69, y + .04, .035, max(.12, body_height - .08), FOURIER_PURPLE)
                body = _text(slide, text, .65 + x, y, column_width, body_height,
                             size=BODY_PT, color=FOURIER_DARK)
                body.name = 'brief:body'
                for paragraph in body.text_frame.paragraphs:
                    paragraph.line_spacing = Pt(BODY_PT * LINE_HEIGHT_FACTOR)
                y += body_height + gap_y
        cursor += height + item_gap
    pages = sorted({page for item in items for page in item.pages})
    _text(slide, _source_footer(pages), .55, presentation.slide_height.inches - .82,
          presentation.slide_width.inches - 1.1, .20, size=9, color=FOURIER_MUTED)
    full_copy = '\n\n'.join(item.title + '\n' + item.text for item in items)
    slide.notes_slide.notes_text_frame.text = '\n\n'.join(value for value in (notes, full_copy) if value)
    return slide


def render_structured_profile(presentation, title, item, *, notes=''):
    """Render one complete authored item as editable, readable text groups.

    The same geometry guard used for all structures rejects arrangements that
    would need clipping, smaller fonts or an additional unbudgeted slide. The
    caller then uses its established complete-copy pagination unchanged.
    """
    from .slide_compositor import _base
    from .pptx_export import _source_footer, _text, FOURIER_DARK, FOURIER_MUTED, FOURIER_PURPLE

    parts = display_segments(item.text, item.display)
    slide, top = _base(presentation, title, '')
    width = presentation.slide_width.inches - 1.3
    heading_height = _heading_height(item.title, width)
    y = top + heading_height + .12
    capacity = presentation.slide_height.inches - 1.28 - y
    layout = item.display.layout if item.display else 'paragraphs'
    # Punctuation alone never assigns chronology, causal links or categories.
    # Explicit semantic structures only influence arrangement of existing copy.
    use_columns = layout in {'comparison', 'relationship', 'key_metrics'} or (
        layout == 'paragraphs' and len(parts) >= 3 and len(item.text) >= 600)
    columns = _columns(parts, width, capacity) if use_columns else None
    if columns is None:
        single_width = width - .34 if layout == 'timeline' else width
        single = _column_geometry(parts, single_width, capacity)
        if single:
            columns = [(.34 if layout == 'timeline' else 0, single_width, parts, single[0], single[2])]
    if columns is None:
        slide_id = presentation.slides._sldIdLst[-1]
        presentation.part.drop_rel(slide_id.rId)
        presentation.slides._sldIdLst.remove(slide_id)
        return None
    if item.title:
        heading = _text(slide, item.title, .65, top, width, heading_height,
                        size=HEADING_PT, bold=True, color=FOURIER_PURPLE)
        heading.name = 'brief:heading'
        heading.text_frame.paragraphs[0].line_spacing = Pt(HEADING_PT * LINE_HEIGHT_FACTOR)
    for x, column_width, copy, heights, gap_y in columns:
        cursor = y
        for text, height in zip(copy, heights):
            if layout == 'timeline':
                # The model, independently reviewed above, owns chronology.
                # A plain editable rule marks the sequence without inventing
                # a calendar label or extracting a guessed milestone.
                from .pptx_export import _rule
                _rule(slide, .69, cursor + .04, .035, max(.12, height - .08), FOURIER_PURPLE)
            body = _text(slide, text, .65 + x, cursor, column_width, height,
                         size=BODY_PT, color=FOURIER_DARK)
            body.name = 'brief:body'
            for paragraph in body.text_frame.paragraphs:
                paragraph.line_spacing = Pt(BODY_PT * LINE_HEIGHT_FACTOR)
            cursor += height + gap_y
    _text(slide, _source_footer(item.pages), .55, presentation.slide_height.inches - .82,
          presentation.slide_width.inches - 1.1, .20, size=9, color=FOURIER_MUTED)
    slide.notes_slide.notes_text_frame.text = '\n\n'.join(value for value in (
        notes, item.title + '\n' + item.text + '\n' + _source_footer(item.pages)) if value)
    slide._ada_summary_layout = layout
    return slide
