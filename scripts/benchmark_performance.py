"""Offline repeatable microbenchmarks, not an end-to-end prospectus SLA.

Run from the repository root: python -m scripts.benchmark_performance
Requires development dependencies. No API keys, external calls or user PDF.
The default renderer is a test stub. Use --real-renderer with a configured
local renderer to measure actual verified export of the synthetic fixture.
"""

import json
import argparse
import io
from statistics import median
from time import perf_counter, sleep
from types import SimpleNamespace

from adaptive_document_agent.agent.document_discovery import ChunkDiscovery, DocumentDiscovery
from adaptive_document_agent.services.export import export_pptx_with_report
from adaptive_document_agent.utils.chunking import DocumentChunk
from tests.ppt_render_stub import LocalRenderStub
from tests.test_pptx_export import _result


def benchmark_preflight(payload: bytes) -> dict[str, float]:
    """Compare wrapper reuse on identical input while verifying every slide."""
    from pptx import Presentation
    from adaptive_document_agent.services.ppt_preflight import PresentationPreflight

    class FreshWrappers(PresentationPreflight):
        def _shapes(self, slide):
            return slide.shapes

        def _placeholders(self, slide):
            return slide.placeholders

    timings = {"fresh_wrappers": [], "reused_wrappers": []}
    baseline = None
    for run in range(11):
        for label, checker in (("fresh_wrappers", FreshWrappers),
                               ("reused_wrappers", PresentationPreflight)):
            deck = Presentation(io.BytesIO(payload))
            started = perf_counter()
            issues = checker(deck).validate_and_sanitize()
            elapsed = perf_counter() - started
            if run:
                timings[label].append(elapsed)
            output = ([vars(issue) for issue in issues],
                      [slide._element.xml for slide in deck.slides],
                      [shape.chart._element.xml for slide in deck.slides
                       for shape in slide.shapes if shape.has_chart])
            if baseline is None:
                baseline = output
            assert output == baseline, "Preflight output changed with wrapper reuse"
    return {f"preflight_{label}_median_ms": round(median(values) * 1000, 1)
            for label, values in timings.items()}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--real-renderer", action="store_true")
    args = parser.parse_args()
    chunks = [DocumentChunk(chunk_id=str(i), start_page=i, end_page=i,
                            text=f"[PAGE {i}] Evidence", estimated_tokens=10) for i in range(1, 13)]

    def simulated_request(chunk):
        sleep(.04)
        return ChunkDiscovery(summary=chunk.text)

    measured = {}
    baseline = None
    for workers in (1, 3):
        discovery = DocumentDiscovery(SimpleNamespace(discovery_workers=workers))
        discovery._discover_chunk = simulated_request
        started = perf_counter()
        output = discovery._discover_chunks(chunks, lambda _: None)
        measured[f"simulated_discovery_{workers}_workers_ms"] = round((perf_counter() - started) * 1000, 1)
        if baseline is None:
            baseline = output
        assert output == baseline

    from adaptive_document_agent.services.presentation_rendering import configured_renderer
    renderer = configured_renderer() if args.real_renderer else LocalRenderStub()
    measured["renderer"] = type(renderer).__name__
    measured["synthetic_fixture_not_real_prospectus"] = True
    result, build_cache, visual_cache = _result(), {}, {}
    payload = None
    for label in ("cold", "warm"):
        started = perf_counter()
        verified = export_pptx_with_report(result, renderer=renderer, build_cache=build_cache, visual_cache=visual_cache)
        measured[f"native_ppt_{label}_ms"] = round((perf_counter() - started) * 1000, 1)
        measured[f"{label}_build_reused"] = verified.build_cache_hit
        measured[f"{label}_visual_qa_reused"] = verified.report.cache_hit
        if payload is not None:
            assert verified.payload == payload
        payload = verified.payload
    if not args.real_renderer:
        measured["renderer_calls"] = renderer.calls
    measured.update(benchmark_preflight(payload))
    print(json.dumps(measured, indent=2))


if __name__ == "__main__":
    main()
