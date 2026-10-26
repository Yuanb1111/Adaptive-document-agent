"""Replay introduction pagination offline; measured choices are NOT semantic approval.

Usage: python scripts/replay_summary_layout.py path/to/analysis_data.json
Reads the private diagnostic without copying or changing it. No model calls.
"""
import argparse
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from adaptive_document_agent.agent.compact_summary import SummarySynthesisDraft, compact_input, expand_compact
from adaptive_document_agent.agent.summary_authored_coverage import authored_coverage_errors
from adaptive_document_agent.agent.summary_layout_choices import layout_choices
from adaptive_document_agent.agent.summary_layout_feedback import compression_binding_errors, selected_meaning_source
from adaptive_document_agent.agent.summary_theme_layout import pack_theme_pages
from adaptive_document_agent.models.summary import SummaryReview
from adaptive_document_agent.services.compact_summary import compact_layout_errors


def replay(path):
    raw = json.loads(Path(path).read_text(encoding='utf-8'))
    audit = json.loads(next(row['message'] for row in reversed(raw['validation_warnings'])
                           if row['code'] == 'company_introduction_summary_audit'))
    review = SummaryReview.model_validate(audit)
    _, catalog = compact_input(review)
    result = []
    for plan in review.plan_audits:
        for attempt in plan.get('layout_attempts', []):
            if 'draft' not in attempt:
                continue
            draft = SummarySynthesisDraft.model_validate(attempt['draft'])
            before = pack_theme_pages(expand_compact(draft, catalog, review))
            choices = layout_choices(draft, catalog, review, expand_compact)
            proposed = draft.model_copy(deep=True)
            for row in choices.values():
                proposed.pages[row['index']] = row['page']
            after = pack_theme_pages(expand_compact(proposed, catalog, review))
            errors = compression_binding_errors(proposed, draft, catalog)
            errors += authored_coverage_errors(proposed, selected_meaning_source(review), catalog, {})
            errors += compact_layout_errors(after)
            result.append(dict(attempt=attempt['attempt'], original_pages=len(before),
                candidate_pages=len(after), choices={key: dict(before=row['original_pages'],
                    after=row['inline_pages']) for key, row in choices.items()},
                validation_errors=errors))
    return dict(pipeline_version=raw['pipeline_version'], budget=review.max_intro_slides,
                attempts=result, semantic_approval=False, export_verified=False)


def replay_repair_path(path):
    """Exercise the actual repair controller with recorded copy and a MOCK choice.

    Unlike candidate enumeration, this fails when the selection branch is bypassed.
    Choosing inline here tests routing and preservation, not model acceptance.
    """
    from adaptive_document_agent.agent.summary_layout_repair import paginate_draft
    from adaptive_document_agent.agent.summary_writer_copy import address_only_copy
    from adaptive_document_agent.services.llm import LLMGateway, LLMSettings, MockLLMClient, ProviderName, PrivacyMode
    raw = json.loads(Path(path).read_text(encoding='utf-8'))
    audit = json.loads(next(row['message'] for row in reversed(raw['validation_warnings'])
                           if row['code'] == 'company_introduction_summary_audit'))
    review = SummaryReview.model_validate(audit)
    _, catalog = compact_input(review)
    recorded = next(plan['layout_attempts'] for plan in reversed(review.plan_audits)
                    if len(plan.get('layout_attempts', [])) >= 2)
    initial = SummarySynthesisDraft.model_validate(recorded[0]['draft'])
    first_rewrite = SummarySynthesisDraft.model_validate(recorded[1]['draft'])
    choices = layout_choices(first_rewrite, catalog, review, expand_compact)
    if not choices:
        raise ValueError('Recorded first rewrite has no feasible full-copy choice')
    client = MockLLMClient([
        dict(pages=address_only_copy(first_rewrite.model_dump(mode='json')['pages'])),
        {key: 'inline' for key in choices}])
    gateway = LLMGateway(client, LLMSettings(provider=ProviderName.MOCK, privacy_mode=PrivacyMode.LOCAL_ONLY))
    repair_audit = {}
    final, pages = paginate_draft(gateway, initial, catalog, review, [], repair_audit,
        expand_compact, SummarySynthesisDraft)
    errors = compression_binding_errors(final, initial, catalog)
    errors += authored_coverage_errors(final, selected_meaning_source(review), catalog, {})
    errors += compact_layout_errors(pages)
    from pptx import Presentation
    from adaptive_document_agent.models.summary_display import SummaryDisplay
    from adaptive_document_agent.services.pptx_export import _resolve_template_path
    from adaptive_document_agent.services.presentation_brief import BriefItem, render_profile
    deck = Presentation(_resolve_template_path())
    rendered = []
    for page in pages:
        rendered.extend(render_profile(deck, page.title, [BriefItem(item.label, item.text,
            item.source_pages, display=item.display or SummaryDisplay()) for item in page.items]))
    return dict(pipeline_version=raw['pipeline_version'], budget=review.max_intro_slides,
        operations=[row['operation'] for row in gateway.usage],
        physical_page_history=[r['layout_feedback']['score']['body_slides'] for r in repair_audit['layout_attempts']],
        first_rewrite_progress=repair_audit['layout_attempts'][1]['compression_progress'],
        native_generated_pages=len(rendered), validation_errors=errors,
        selection_response='mock_inline_choices', semantic_approval=False, export_verified=False)


def replay_recorded_final(path):
    """Measure and natively generate final recorded copy without any model choice."""
    from adaptive_document_agent.agent.summary_layout_repair import paginate_draft
    from adaptive_document_agent.models.summary import SummarySlidePage
    raw = json.loads(Path(path).read_text(encoding='utf-8'))
    audit = json.loads(next(row['message'] for row in reversed(raw['validation_warnings'])
        if row['code'] == 'company_introduction_summary_audit'))
    review = SummaryReview.model_validate(audit)
    _, catalog = compact_input(review)
    recorded = next(plan['layout_attempts'] for plan in reversed(review.plan_audits)
                    if plan.get('layout_attempts'))
    initial = SummarySynthesisDraft.model_validate(recorded[0]['draft'])
    draft = SummarySynthesisDraft.model_validate(recorded[-1]['draft'])
    snapshot = draft.model_dump(mode='json')
    class NoModel:
        def generate_structured(self, *args, **kwargs):
            raise AssertionError('Recorded-final replay cannot invent a new rewrite or model choice.')
    measured = {}
    final, pages = paginate_draft(NoModel(), draft, catalog, review, [], measured,
        expand_compact, SummarySynthesisDraft, allow_rewrite=False)
    assert final.model_dump(mode='json') == snapshot
    pages = [SummarySlidePage.model_validate_json(page.model_dump_json()) for page in pages]
    errors = compression_binding_errors(final, initial, catalog)
    errors += authored_coverage_errors(final, selected_meaning_source(review), catalog, {})
    errors += compact_layout_errors(pages)
    from pptx import Presentation
    from adaptive_document_agent.models.summary_display import SummaryDisplay
    from adaptive_document_agent.services.pptx_export import _resolve_template_path
    from adaptive_document_agent.services.presentation_brief import BriefItem, render_profile
    deck, rendered = Presentation(_resolve_template_path()), []
    for page in pages:
        rendered.extend(render_profile(deck, page.title, [BriefItem(item.label, item.text,
            item.source_pages, display=item.display or SummaryDisplay()) for item in page.items]))
    visible = ''.join(shape.text for slide in rendered for shape in slide.shapes if shape.name == 'brief:body')
    expected = ''.join(item.text for page in final.pages for item in page.items)
    assert visible == expected, 'Native generation must preserve complete authored text.'
    assert len(rendered) == len(pages), 'Native generation must agree with measured pagination.'
    return dict(pipeline_version=raw['pipeline_version'], budget=review.max_intro_slides,
        recorded_physical_pages=recorded[-1]['layout_feedback']['score']['body_slides'],
        measured_physical_pages=len(pages), native_generated_pages=len(rendered),
        complete_body_text_preserved=True, validation_errors=errors,
        model_calls=0, semantic_approval=False, export_verified=False)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('analysis_json')
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument('--repair-path', action='store_true', help='Exercise the real controller with a mock choice')
    mode.add_argument('--recorded-final', action='store_true', help='Natively generate the final recorded draft without model calls')
    arguments = parser.parse_args()
    action = replay_recorded_final if arguments.recorded_final else replay_repair_path if arguments.repair_path else replay
    print(json.dumps(action(arguments.analysis_json), indent=2))
