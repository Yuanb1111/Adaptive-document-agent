"""Deterministic rendering-boundary fake, NOT a real visual rendering test.

Legacy export tests inspect editable OOXML. They explicitly opt into this fake
so the offline suite does not require a separately installed native renderer.
P2 tests exercise the real gate with malformed outputs and bounded repairs.
"""

import io

from PIL import Image
from pptx import Presentation

from adaptive_document_agent.services.presentation_rendering import RenderedPage


class LocalRenderStub:
    identity = "test-only-native-boxes"

    def __init__(self):
        self.calls = 0
        self.directories = []

    def render(self, payload, directory):
        self.calls += 1
        self.directories.append(directory)
        deck = Presentation(io.BytesIO(payload))
        width = deck.slide_width / 9525
        height = deck.slide_height / 9525
        im = Image.new("RGB", (round(width), round(height)), "white")
        # Synthetic pixels only: provide variation inside every data-visual box.
        for x in range(0, round(width), 30):
            im.paste("black", (x, 0, x+2, round(height)))
        stream = io.BytesIO()
        im.save(stream, format="PNG")

        def elements(shapes):
            return [{"id": str(s.shape_id), "scope": "slide", "name": s.name,
                     "kind": "chart" if s.has_chart else "table" if s.has_table else "shape",
                     "text": s.text if s.has_text_frame else "",
                     "bbox": [s.left/9525, s.top/9525, s.width/9525, s.height/9525]}
                    for s in shapes if s.has_chart or s.has_table or (s.has_text_frame and s.text.strip())]

        return [RenderedPage(stream.getvalue(), {"unit": "px", "slide": {"frame": {"width": width, "height": height}},
                "elements": elements(slide.shapes), "inheritedLayers": [
                    {"scope": "layout", "elements": elements([s for s in slide.slide_layout.shapes if not s.is_placeholder])},
                    {"scope": "master", "elements": elements([s for s in slide.slide_layout.slide_master.shapes if not s.is_placeholder])}]})
                for slide in deck.slides]
