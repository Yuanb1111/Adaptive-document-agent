"""Physical source-index exemptions never exempt audience analysis or introductions."""
from io import BytesIO
import json

import pytest
from pptx import Presentation

from adaptive_document_agent.models import PresentationPlan, PresentationSlide
from adaptive_document_agent.services.presentation_requirements import apply_presentation_budget, fixed_presentation_budget
from adaptive_document_agent.services.presentation_body_budget import slide_counts, place_data_index_after_body
from adaptive_document_agent.services.slide_budgets import require_exported_budgets
from tests.test_compact_summary_v111 import fixture


def deck_with_index(body=20, index=3):
    result, _, _ = fixture(1)
    result.profile.report_requirements = apply_presentation_budget(None, '', fixed_presentation_budget())
    result.presentation_plan = PresentationPlan(title='Deck', slides=[
        PresentationSlide(id='analysis', slide_type='analysis', title='Data Index'),
        PresentationSlide(id='intro', slide_type='company_overview', title='Introduction')])
    deck = Presentation()
    for n in range(body + index):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 1000, 1000).text = f'Page {n + 1}'
        if n >= body:
            slide.element.cSld.set('name', 'ada:data-index')
    return result, deck


def roundtrip(deck):
    stream = BytesIO()
    deck.save(stream)
    return Presentation(stream)


def test_twenty_body_pages_plus_source_index_pass_after_saved_ppt_roundtrip():
    result, deck = deck_with_index()
    deck = roundtrip(deck)
    require_exported_budgets(deck, result)
    assert slide_counts(deck, result) == {'total': 23, 'body': 20, 'data_index': 3}
    result.presentation_slide_counts = slide_counts(deck, result)
    assert type(result).model_validate_json(result.model_dump_json()).presentation_slide_counts == result.presentation_slide_counts


@pytest.mark.parametrize('tag', ['', 'ada:appendix', 'ada:data-index'])
def test_twenty_first_body_page_cannot_hide_as_appendix_or_data_index(tag):
    result, deck = deck_with_index(body=21)
    deck.slides[20].element.cSld.set('name', tag)
    result.presentation_export_trace = [dict(planned_slide_id='analysis', slide_numbers=[20, 21])]
    with pytest.raises(ValueError, match='report body has 21'):
        require_exported_budgets(roundtrip(deck), result)


def test_all_eleven_analytical_continuations_count_even_with_source_index_tags():
    result, deck = deck_with_index(body=19)
    for slide in list(deck.slides)[8:19]:
        slide.element.cSld.set('name', 'ada:data-index')
    result.presentation_export_trace = [dict(planned_slide_id='analysis', slide_numbers=list(range(9, 20)))]
    with pytest.raises(ValueError, match='analysis body has 11'):
        require_exported_budgets(roundtrip(deck), result)


def test_sixth_introduction_page_counts_even_when_marked_as_source_index():
    from adaptive_document_agent.services.summary_export import _PREFIX
    result, deck = deck_with_index()
    for n, slide in enumerate(list(deck.slides)[:6]):
        slide.element.cSld.set('name', 'ada:data-index')
        slide.shapes[0].element.xpath('.//p:cNvPr')[0].set('descr', _PREFIX + json.dumps({'page_id': str(n)}))
    with pytest.raises(ValueError, match='introduction body has 6'):
        require_exported_budgets(roundtrip(deck), result)


def test_body_only_cap_does_not_trigger_paid_refinement_for_a_large_index(monkeypatch):
    from adaptive_document_agent.agent import presentation_budget_refinement as refinement
    from tests.test_physical_budget_refinement_v130 import source
    result, _, _, _ = source()
    result.profile.report_requirements = apply_presentation_budget(None, '', fixed_presentation_budget())
    _, deck = deck_with_index(body=20, index=12)
    result.presentation_export_trace = []
    monkeypatch.setattr(refinement, 'measure', lambda current: (deck, slide_counts(deck, current)))
    class Client:
        def generate_structured(self, *args, **kwargs):
            raise AssertionError('Source-index length does not justify omitting analytical evidence')
    assert not refinement.refine_presentation_budget(Client(), result)


def test_diagnostics_report_body_and_index_only_when_verified_counts_match():
    from adaptive_document_agent.services.export_diagnostics import export_diagnostics
    from adaptive_document_agent.services.presentation_visual_qa import VisualQAReport
    result, deck = deck_with_index()
    result.presentation_slide_counts = slide_counts(deck, result)
    report = export_diagnostics(result, visual_report=VisualQAReport(status='passed', slide_count=23))
    assert (report['physical_slides'], report['body_slides'], report['data_index_slides']) == (23, 20, 3)
    stale = export_diagnostics(result, visual_report=VisualQAReport(status='passed', slide_count=24))
    assert stale['body_slides'] is None and stale['data_index_slides'] is None
