"""Source comparisons are literal, editable and cannot omit conditional cases."""
import pytest
from pptx import Presentation
from pptx.util import Inches
from adaptive_document_agent.models.executive_brief import ExecutiveBrief, BriefComparisonTable
from adaptive_document_agent.services.executive_brief import validate_executive_brief, brief_items
from adaptive_document_agent.services.presentation_summary import render_complete_summary
from tests.test_executive_brief import result_for, payload


def fixture():
    text = 'Response time is estimated at 44 months without expansion, 60 months with partial expansion and 89 months with full expansion.'
    result = result_for([text])
    data = payload('Response time estimates depend on expansion assumptions.', label='Response scenarios')
    data['items'][0]['evidence'][0]['text'] = text
    data['items'][0]['comparison_table'] = {'headers': ['Assumption', 'Estimated response time'],
        'rows': [['without expansion', '44 months'], ['with partial expansion', '60 months'],
                 ['with full expansion', '89 months']]}
    brief = ExecutiveBrief.model_validate(data)
    return result, brief


def test_source_scenarios_are_editable_with_all_values_and_qualifiers():
    result, brief = fixture()
    before = result.model_dump(), brief.model_dump()
    assert not validate_executive_brief(brief, result)
    result.executive_brief = brief
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slides = render_complete_summary(deck, brief.title, brief_items(result))
    assert len(slides) == 1
    table = next(s.table for s in slides[0].shapes if s.has_table)
    assert [[cell.text for cell in row.cells] for row in table.rows] == [brief.items[0].comparison_table.headers,
                                                                     *brief.items[0].comparison_table.rows]
    assert '44 months' in slides[0].notes_slide.notes_text_frame.text
    assert 'estimates' in ' '.join(s.text for s in slides[0].shapes if s.has_text_frame)
    result.executive_brief = None
    assert (result.model_dump(), brief.model_dump()) == before


@pytest.mark.parametrize('fault', ['value', 'unit', 'case', 'invented_assumption'])
def test_table_cannot_invent_values_units_or_lose_conditional_cases(fault):
    result, brief = fixture()
    table = brief.items[0].comparison_table
    if fault == 'value': table.rows[0][1] = '45 months'
    if fault == 'unit': table.rows[0][1] = '44 years'
    if fault == 'case': table.rows.pop(1)
    if fault == 'invented_assumption': table.rows[0][0] = 'without recruitment'
    assert validate_executive_brief(brief, result)


def test_table_requires_rectangular_bounded_cells():
    with pytest.raises(ValueError):
        BriefComparisonTable(headers=['Case', 'Outcome'], rows=[['A']])


def test_literal_accounting_cells_keep_signs_without_magnitude_annotations():
    from tests.test_brief_quantity_representation import magnitude_fixture
    result, brief = magnitude_fixture()
    item = brief.items[0]
    item.text = 'The reported source row retains signed values.'
    item.quantity_representations = []
    item.comparison_table = BriefComparisonTable(headers=['Measure', 'USD in thousands'],
                                               rows=[['Net loss', '(100)'], ['Net loss', '(200)']])
    assert not validate_executive_brief(brief, result)
    item.comparison_table.rows[0][1] = '100'
    assert validate_executive_brief(brief, result)

def test_table_does_not_isolate_the_last_ordinary_finding():
    from adaptive_document_agent.services.presentation_brief import BriefItem
    from tests.test_presentation_brief import blank_deck
    table = BriefComparisonTable(headers=['Case', 'Estimate'], rows=[['Base case', '44 months']])
    first = BriefItem('First finding', 'Complete source-supported finding and qualification. ' * 3, [1])
    last = BriefItem('Last finding', 'Final source-supported finding and qualification. ' * 3, [3])
    pages = render_complete_summary(blank_deck(), 'Summary', [first,
        BriefItem('Scenario', 'Conditional estimate.', [2], table=table), last], single_column=True)
    assert len(pages) == 2
    copy = '\n'.join(s.text for s in pages[0].shapes if s.has_text_frame)
    assert first.text in copy and last.text in copy
    assert any(s.has_table for s in pages[1].shapes)

def test_complete_literal_duration_cell_can_restore_unit_from_explicit_header():
    from adaptive_document_agent.services.brief_table_units import literal_table_units
    result, brief = fixture()
    item = brief.items[0]
    item.comparison_table.headers[1] = 'Estimated response time (months)'
    item.comparison_table.rows[1][1] = '60'
    before = brief.model_dump()
    assert not validate_executive_brief(brief, result)
    table = literal_table_units(item, {1: result.document.pages[0].text})
    assert table.rows[1][1] == '60 months'
    assert brief.model_dump() == before
    item.comparison_table.rows[1][1] = '61'
    assert validate_executive_brief(brief, result)
