"""Synthetic relational evidence: rate denominators and conditional alternatives."""
import pytest
from pptx import Presentation
from pptx.util import Inches

from adaptive_document_agent.models import DocumentPage, DocumentProfile, ParsedDocument, PipelineResult
from adaptive_document_agent.models.executive_brief import ExecutiveBrief
from adaptive_document_agent.services.executive_brief import brief_items, validate_executive_brief
from adaptive_document_agent.services.presentation_summary import render_complete_summary


def _result(source, summary, definition='', *, extra_pages=None):
    texts = {1: definition or 'Unrelated background.', 2: source, **(extra_pages or {})}
    result = PipelineResult(document=ParsedDocument(document_id='synthetic', sha256='a'*64,
        safe_filename='synthetic.pdf', page_count=max(texts),
        pages=[DocumentPage(page_number=p, text=t) for p, t in texts.items()]),
        profile=DocumentProfile(document_purpose='Explain supported findings'))
    result.executive_brief = ExecutiveBrief.model_validate({'title': 'Findings', 'items': [{
        'label': 'Capacity', 'text': summary, 'evidence': [{'page': 2, 'text': source}]}]})
    return result


@pytest.mark.parametrize('metric,cadence,amount,unit', [
    ('resource consumption rate', 'monthly', '8.4', 'million'),
    ('processing rate', 'hourly', '72', 'units'),
    ('patient arrival rate', 'weekly', '160', 'patients'),
])
def test_adjacent_explicit_definition_restores_the_rate_denominator(metric, cadence, amount, unit):
    definition = f'Our {metric} refers to the average {cadence} quantity used.'
    source = f'Assuming the {metric} remains {amount} {unit}, the current reserve is sufficient.'
    result = _result(source, source, definition)
    before = result.model_dump_json()
    shown = brief_items(result)[0]
    assert f'{cadence} {metric}' in shown.text
    assert shown.pages == [1, 2]
    assert result.model_dump_json() == before


@pytest.mark.parametrize('fraction,unit,low,middle,high', [
    ('25%', 'weeks', '7.5', '9.2', '22'),
    ('40%', 'months', '5', '8', '14'),
    ('60%', 'days', '2.4', '4.1', '9.8'),
])
def test_same_sentence_keeps_every_conditional_case_and_its_allocation(fraction, unit, low, middle, high):
    source = (f'Assuming the processing rate remains 24 units, the reserve lasts approximately {low} {unit} '
        f'or, if we take into account {fraction} of the backup supply '
        f'(namely, the portion allocated for emergency operations), approximately {middle} {unit} '
        f'or, if we take into account the backup supply, approximately {high} {unit}.')
    summary = (f'Assuming the processing rate remains 24 units, the reserve lasts approximately {low} {unit} '
               f'or approximately {high} {unit} if the backup supply is included.')
    result = _result(source, summary, 'Our processing rate refers to the average hourly output.')
    shown = brief_items(result)[0]
    assert f'{middle} {unit}' in shown.text
    assert fraction in shown.text and 'allocated for emergency operations' in shown.text
    assert 'all' in shown.text or 'full' in shown.text
    assert 'hourly processing rate' in shown.text
    assert not validate_executive_brief(result.executive_brief, result)


def test_explicit_alternatives_with_values_before_assumptions_preserve_all_cases():
    source = ('Estimated coverage is 6 weeks without the reserve, 10 weeks with 30% of the reserve '
              'allocated to maintenance, or 19 weeks with all the reserve.')
    result = _result(source, 'Estimated coverage is 6 weeks without the reserve or 19 weeks with the reserve.')
    shown = brief_items(result)[0]
    assert '10 weeks with 30% of the reserve allocated to maintenance' in shown.text
    assert '19 weeks with all the reserve' in shown.text


@pytest.mark.parametrize('definition', [
    'Our processing rate refers to the average daily output. Our processing rate refers to the average weekly output.',
    'Our arrival rate refers to the average weekly output.',
    'The processing rate might be monthly or quarterly.',
    '',
])
def test_ambiguous_unrelated_or_absent_definitions_cannot_supply_a_unit(definition):
    source = 'The processing rate was 24 units.'
    result = _result(source, source, definition)
    assert brief_items(result)[0].text == source
    assert brief_items(result)[0].pages == [2]


def test_nonadjacent_definition_does_not_supply_a_unit():
    source = 'The processing rate was 24 units.'
    result = _result(source, source, extra_pages={7: 'Our processing rate refers to the average daily output.'})
    assert brief_items(result)[0].text == source


def test_nearby_but_uncited_scenarios_are_not_added():
    source = 'Coverage is 7 weeks without reserve and 18 weeks with reserve.'
    result = _result(source, source,
        'Coverage is 9 weeks if 20% of the reserve is used, or 22 weeks if all the reserve is used.')
    assert brief_items(result)[0].text == source


def test_units_and_assumptions_already_present_are_idempotent():
    source = 'Coverage is 5 weeks without reserve, 8 weeks with 40% of reserve, or 14 weeks with all reserve.'
    result = _result(source, source)
    assert brief_items(result)[0].text == source


def test_long_relational_copy_stays_visible_and_preserves_summary_line_spacing():
    source = ('Assuming the processing rate remains 24 units, estimated reserve coverage is 6 weeks '
        'or, if we take into account 30% of the backup supply (namely, the portion allocated for '
        'maintenance and emergency operations during scheduled interruptions), approximately 10 weeks '
        'or, if we take into account the backup supply, approximately 19 weeks.')
    result = _result(source, 'Reserve coverage is 6 weeks or 19 weeks with backup supply.')
    items = brief_items(result)
    presentation = Presentation(); presentation.slide_width = Inches(13.333); presentation.slide_height = Inches(7.5)
    render_complete_summary(presentation, 'Findings', items, single_column=True)
    visible = ''.join(shape.text for slide in presentation.slides for shape in slide.shapes
                      if shape.has_text_frame and shape.name == 'brief:body')
    assert '10 weeks' in visible and '30%' in visible
    assert items[0].text == visible
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.name == 'brief:body':
                paragraph = shape.text_frame.paragraphs[0]
                assert paragraph.line_spacing.pt == paragraph.font.size.pt * 1.25


def test_ratio_scenarios_keep_intermediate_case_and_its_assumption():
    source = ('Estimated throughput is 1.2 times without upgrades, 1.7 times with partial upgrades '
              'under the pilot plan, or 2.6 times with all upgrades.')
    result = _result(source, 'Estimated throughput is 1.2 times without upgrades or 2.6 times with upgrades.')
    assert '1.7 times with partial upgrades under the pilot plan' in brief_items(result)[0].text


def test_same_numbers_for_a_different_subject_do_not_authorize_scenario_replacement():
    source = ('Estimated coverage is 6 weeks without the reserve, 10 weeks with 30% of the reserve '
              'allocated to maintenance, or 19 weeks with all the reserve.')
    summary = 'The study had 6 weeks of recruitment and 19 weeks of follow-up.'
    result = _result(source, summary)
    assert brief_items(result)[0].text == summary


def test_cadence_cannot_be_added_when_the_quote_budget_is_full():
    source = 'The processing rate was 24 units.'
    result = _result(source, source, 'Our processing rate refers to the average daily output.')
    result.executive_brief.items[0].evidence *= 4
    assert brief_items(result)[0].text == source


def test_contiguous_definition_evidence_is_retained_with_its_source_page():
    from adaptive_document_agent.services.brief_context import preserve_brief_context
    source = 'The processing rate was 24 units.'
    definition = 'Our processing rate refers to the average daily output.'
    result = _result(source, source, definition)
    copy = preserve_brief_context(result.executive_brief.items[0], {1: definition, 2: source})
    assert copy.evidence[-1].page == 1 and copy.evidence[-1].text == definition
    assert len(copy.evidence) == 2


def test_partial_operand_does_not_turn_a_different_scenario_into_allocation_of_the_whole():
    source = ('Estimated coverage is 6 weeks or, if we take into account 30% of the backup supply, '
              'approximately 10 weeks or, if we take into account the emergency supply, approximately 19 weeks.')
    result = _result(source, 'Estimated coverage is 6 weeks or 19 weeks with emergency supply.')
    copy = brief_items(result)[0].text
    assert '10 weeks' in copy
    assert 'all' not in copy


def test_existing_different_explicit_denominator_is_not_silently_overwritten():
    source = 'The processing rate was 24 units per week.'
    result = _result(source, source, 'Our processing rate refers to the average daily output.')
    assert brief_items(result)[0].text == source


def test_full_cited_sentence_keeps_original_extra_facts_from_another_quote():
    source = ('Estimated coverage is 6 weeks without the reserve, 10 weeks with 30% of the reserve '
              'allocated to maintenance, or 19 weeks with all the reserve. Staffing was 84 people.')
    summary = 'Estimated coverage is 6 weeks or 19 weeks, while staffing was 84 people.'
    result = _result(source, summary)
    copy = brief_items(result)[0].text
    assert '84 people' in copy
    assert '10 weeks' in copy


def test_writer_receives_adjacent_literal_rate_definition_without_a_new_model_call():
    from adaptive_document_agent.agent.executive_brief import ExecutiveBriefWriter
    from tests.test_executive_brief import gateway, payload
    source = 'The processing rate was 24 units.'
    definition = 'Our processing rate refers to the average daily output.'
    result = _result(source, source, definition, extra_pages={p: 'Background context.' for p in range(3, 13)})
    g, client = gateway([{'pages': [2]}, payload(source, page=2)])
    ExecutiveBriefWriter(g).generate(result)
    assert len(client.calls) == 2
    assert definition in client.calls[1][1]['content']
    assert 'rate denominator' in client.calls[1][0]['content']
    assert 'intermediate' in client.calls[1][0]['content']


def test_alternative_denominators_in_a_single_definition_remain_ambiguous():
    source = 'The processing rate was 24 units.'
    result = _result(source, source, 'Our processing rate refers to the average daily or weekly output.')
    assert brief_items(result)[0].text == source


def test_decimal_notation_does_not_hide_a_conditional_case():
    source = ('Estimated coverage is 6.0 weeks without reserve, 10.0 weeks with partial reserve, '
              'or 19.0 weeks with all reserve.')
    result = _result(source, 'Estimated coverage is 6 weeks without reserve or 19 weeks with reserve.')
    assert '10.0 weeks with partial reserve' in brief_items(result)[0].text


def test_explicit_quoted_cadence_cannot_be_overridden_by_a_conflicting_definition():
    source = 'The weekly processing rate was 24 units.'
    result = _result(source, 'The processing rate was 24 units.',
                     'Our processing rate refers to the average daily output.')
    assert 'daily' not in brief_items(result)[0].text


def test_literal_quoted_rate_denominator_survives_shortening_without_a_separate_definition():
    source = 'The average weekly processing rate was 24 units.'
    result = _result(source, 'The average processing rate was 24 units.')
    assert brief_items(result)[0].text == source


def test_literal_scenario_copy_still_makes_the_partial_versus_whole_assumption_explicit():
    source = ('Estimated coverage is 6 weeks or, if we take into account 30% of the backup supply, '
              'approximately 10 weeks or, if we take into account the backup supply, approximately 19 weeks.')
    result = _result(source, source)
    assert 'take into account all of the backup supply' in brief_items(result)[0].text


def test_explicit_quoted_postfix_denominator_blocks_conflicting_adjacent_definition():
    source = 'The processing rate was 24 units per week.'
    summary = 'The processing rate was 24 units.'
    result = _result(source, summary, 'Our processing rate refers to the average daily output.')
    assert 'daily' not in brief_items(result)[0].text


@pytest.mark.parametrize('denominator,cadence', [('per hour', 'hourly'), ('per month', 'monthly'), ('/week', 'weekly')])
def test_literal_postfix_denominator_can_restore_the_same_source_bound_rate(denominator, cadence):
    source = f'The processing rate was 24 units {denominator}.'
    result = _result(source, 'The processing rate was 24 units.')
    assert f'{cadence} processing rate' in brief_items(result)[0].text


def test_postfix_denominator_for_a_different_measure_does_not_bind_to_a_rate():
    source = 'The processing rate was 24 units, and deliveries were 8 units per week.'
    result = _result(source, 'The processing rate was 24 units.')
    assert brief_items(result)[0].text == 'The processing rate was 24 units.'


def test_postfix_denominator_after_another_measure_in_the_same_clause_is_not_borrowed():
    source = 'The processing rate was 24 units and deliveries were 8 units per week.'
    result = _result(source, 'The processing rate was 24 units.')
    assert brief_items(result)[0].text == 'The processing rate was 24 units.'


def test_literal_independent_sentence_is_not_replaced_by_another_subjects_scenarios():
    summary = 'Staffing coverage was 6 weeks in the pilot, increasing to 19 weeks in the full rollout.'
    source = (summary + ' Estimated battery coverage is 6 weeks without the reserve, '
              '10 weeks with partial reserve, or 19 weeks with all the reserve.')
    assert brief_items(_result(source, summary))[0].text == summary


@pytest.mark.parametrize('summary', [
    'Staffing coverage is 6 weeks without reserve or 19 weeks with reserve.',
    'Staffing coverage for emergency operations is 6 weeks without reserve or 19 weeks with reserve.',
    'Estimated staffing coverage is 6 weeks in the pilot or 19 weeks in full rollout.',
])
def test_generic_or_nonsubject_shared_words_cannot_bind_different_scenarios(summary):
    source = ('Estimated battery coverage for emergency operations is 6 weeks without reserve, '
              '10 weeks with partial reserve, or 19 weeks with all reserve.')
    assert brief_items(_result(source, summary))[0].text == summary


def test_attributed_dated_subject_keeps_the_complete_conditional_comparison():
    source = ('Assuming a processing rate of 24 units, we estimate that our reserves as of March 12, 2028 '
              'will support operations for 6 weeks or, if we take into account 30% of the backup supply, '
              '10 weeks or, if we take into account the backup supply, 19 weeks.')
    summary = ('Assuming a processing rate of 24 units, the team estimated that its reserves as of March 12, 2028 '
               'could support operations for 6 weeks or 19 weeks with backup supply.')
    copy = brief_items(_result(source, summary))[0].text
    assert '10 weeks' in copy and '30%' in copy and 'March 12, 2028' in copy
    assert 'all of the backup supply' in copy


@pytest.mark.parametrize('prefix', ['Assuming the', 'Assuming that the', 'If the', 'Given the', 'Given that the', 'Provided that the'])
def test_conditional_rate_predicate_keeps_its_postfix_denominator_in_conflict_checks(prefix):
    source = f'{prefix} processing rate remains 24 units per week, the reserve is sufficient.'
    summary = f'{prefix} processing rate remains 24 units, the reserve is sufficient.'
    result = _result(source, summary, 'Our processing rate refers to the average daily output.')
    assert 'daily' not in brief_items(result)[0].text


def test_subject_name_containing_shared_rate_suffix_does_not_supply_denominator():
    source = 'The battery processing rate remains 24 units per week.'
    result = _result(source, 'The processing rate remains 24 units.',
                     'Our processing rate refers to the average daily output.')
    # A different explicitly named subject is not treated as the same rate.
    assert brief_items(result)[0].text == 'The processing rate remains 24 units.'


@pytest.mark.parametrize('connector', [', and ', '; ', ' but '])
def test_restoring_cases_keeps_a_separately_supported_qualitative_qualification(connector):
    source = ('Estimated coverage is 6 weeks without the reserve, 10 weeks with 30% of the reserve '
              'allocated to maintenance, or 19 weeks with all the reserve. The estimate excludes overtime demand.')
    summary = 'Estimated coverage is 6 weeks or 19 weeks with reserves' + connector + 'the estimate excludes overtime demand.'
    copy = brief_items(_result(source, summary))[0].text
    assert '10 weeks' in copy
    assert 'excludes overtime demand' in copy


@pytest.mark.parametrize('summary', [
    'Estimated coverage is 6 weeks without the reserve, excluding overtime demand, or 19 weeks with reserves.',
    'Estimated coverage, excluding overtime demand, is 6 weeks or 19 weeks with reserves.',
    'Estimated coverage is 6 weeks or 19 weeks with reserves, subject to the same demand limit.',
])
def test_authored_qualifications_anywhere_survive_scenario_completion(summary):
    source = ('Estimated coverage is 6 weeks without the reserve, 10 weeks with 30% of the reserve '
              'allocated to maintenance, or 19 weeks with all the reserve. '
              'The estimate excludes overtime demand and is subject to the same demand limit.')
    copy = brief_items(_result(source, summary))[0].text
    assert summary in copy
