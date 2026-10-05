"""Actual heading/support geometry paginates without dropping selected evidence."""

import json

import pytest
from pptx import Presentation

from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import ChartPlan, PresentationSlide, PresentationVisualBlock
from adaptive_document_agent.services.pptx_export import _resolve_template_path
from adaptive_document_agent.services.slide_compositor import (
    _render_composed_slide, render_composed_slide, validate_composed_geometry,
)
from tests.test_p0_composition import observation, result_for


def _deck():
    deck = Presentation(str(_resolve_template_path()))
    for sid in list(deck.slides._sldIdLst):
        deck.part.drop_rel(sid.rId)
        deck.slides._sldIdLst.remove(sid)
    return deck


@pytest.mark.parametrize("themed", [False, True])
@pytest.mark.parametrize("metrics", [
    ("Loss for the year/period", "Annual research and development expenditure",
     "Annual research and development expenditure ratio %"),
    ("Production for the year/period", "Annual production and maintenance expenditure",
     "Annual production and maintenance expenditure ratio %"),
])
def test_measured_chart_capacity_continues_all_panels_and_exact_table(themed, metrics):
    obs, charts = [], []
    for number, metric in enumerate(metrics):
        series = [observation(f"m{number}-{year}", metric, value, f"FY{year}",
                              unit="percent" if number == 2 else "currency")
                  for year, value in zip((2021, 2022, 2023), (34, 27, 28) if number == 2 else (40, 52, 103))]
        for item in series:
            item.validation_status = "valid"
        obs.extend(series)
        charts.append(ChartPlan(id=f"chart-{number}", title=metric + " — Reported Values",
                                question="How do the reported values compare?",
                                chart_type="line", observation_ids=[o.id for o in series], source_pages=[3]))
    slide = PresentationSlide(id="capacity", slide_type="analysis", layout="three_up",
        title="Reported operating performance from FY2021 to FY2023 with annual expenditure alongside its reported ratio to annual total operating expenditure.",
        section_title="Annual operating activity with rising expenditure and lower intensity",
        message="How did activity and annual expenditure compare across FY2021-FY2023?",
        theme_id="activity" if themed else "", chart_ids=[c.id for c in charts], source_pages=[3],
        visual_blocks=[PresentationVisualBlock(role="table", observation_ids=charts[0].observation_ids)])
    result = result_for(obs, charts, slide)
    index = DocumentIndex(obs)
    snapshot = result.model_dump()
    # This is the actual capacity failure from the full export, with valid data
    # and short enough titles to pass the earlier three-panel heuristic.
    with pytest.raises(ValueError, match="Chart labels cannot fit"):
        _render_composed_slide(_deck(), slide, charts, result, index)
    deck = _deck()
    sentinel = deck.slides.add_slide(deck.slide_layouts[0])
    rendered = render_composed_slide(deck, slide, charts, result, index)
    assert deck.slides[0] is sentinel
    assert len(rendered) == 2 and len(deck.slides) == 3
    native = [shape for page in rendered for shape in page.shapes if shape.has_chart]
    assert [s.name for s in native] == [f"chart:{c.id}" for c in charts]
    assert [list(s.chart.series[0].values) for s in native] == [[40, 52, 103], [40, 52, 103], [34, 27, 28]]
    tables = [shape.table for page in rendered for shape in page.shapes if shape.has_table]
    assert len(tables) == 1
    assert all(str(value) in " ".join(c.text for row in tables[0].rows for c in row.cells)
               for value in (40, 52, 103))
    assert all("Source:" in " ".join(s.text for s in page.shapes if s.has_text_frame) for page in rendered)
    notes = [json.loads(page.notes_slide.notes_text_frame.text) for page in rendered]
    assert {o["id"] for note in notes for o in note["observations"]} == {o.id for o in obs}
    assert result.model_dump() == snapshot
    validate_composed_geometry(deck)


def test_invalid_support_evidence_is_not_retried_or_hidden():
    from tests.test_p0_composition import paired_result
    result = paired_result()
    result.observations[-1].evidence = []
    deck = _deck()
    with pytest.raises(ValueError, match="evidence is incomplete"):
        render_composed_slide(deck, result.presentation_plan.slides[3], result.charts,
                              result, DocumentIndex(result.observations))
    assert not len(deck.slides)


def test_later_panel_failure_rolls_back_native_charts_and_preserves_hero_order():
    metrics = ["Output", "Annual qualified operating investment across regional facilities and service operations for the reported period including planned maintenance and technical support activity"]
    obs = [observation(f"m{i}-{year}", metric, value, f"FY{year}", unit="count")
           for i, metric in enumerate(metrics) for year, value in ((2024, 10), (2025, 20))]
    charts = [ChartPlan(id=f"chart-{i}", title=metric, question="How did activity compare?",
                        chart_type="bar", observation_ids=[o.id for o in obs if o.metric_original == metric])
              for i, metric in enumerate(metrics)]
    slide = PresentationSlide(id="later-panel", slide_type="analysis", title="Annual reported activity",
        layout="two_up", chart_ids=["chart-1", "chart-0"], source_pages=[3],
        visual_blocks=[PresentationVisualBlock(role="hero", chart_ids=["chart-0"])])
    result = result_for(obs, charts, slide)
    index = DocumentIndex(obs)
    failed = _deck()
    with pytest.raises(ValueError, match="title exceeds readable capacity"):
        _render_composed_slide(failed, slide, charts, result, index)
    assert sum(s.has_chart for p in failed.slides for s in p.shapes) == 1
    deck = _deck()
    pages = render_composed_slide(deck, slide, charts, result, index)
    assert len(pages) == len(deck.slides) == 2
    assert [s.name for p in pages for s in p.shapes if s.has_chart] == ["chart:chart-0", "chart:chart-1"]
    assert all(list(s.chart.series[0].values) == [10, 20] for p in pages for s in p.shapes if s.has_chart)
    validate_composed_geometry(deck)


def test_unfit_single_chart_still_blocks_and_removes_partial_page():
    item = observation("long", "Annual " + "qualified operating measure " * 30, 40, unit="count")
    chart = ChartPlan(id="long-chart", title=item.metric_original, question="What is reported?",
                      chart_type="bar", observation_ids=[item.id])
    slide = PresentationSlide(id="single", slide_type="analysis", title="Reported activity", chart_ids=[chart.id])
    result = result_for([item], [chart], slide)
    deck = _deck()
    with pytest.raises(ValueError, match="title exceeds readable capacity"):
        render_composed_slide(deck, slide, [chart], result, DocumentIndex([item]))
    assert not len(deck.slides)
