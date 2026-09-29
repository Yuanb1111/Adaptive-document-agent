"""Closing copy stays with its linked facts without guessing semantic links."""

from copy import deepcopy

from pptx import Presentation
from pptx.util import Inches

from adaptive_document_agent.models import (
    AnalysisResult, DocumentProfile, Insight, Observation, ParsedDocument,
    PipelineResult, PresentationSlide, SourceEvidence,
)
from adaptive_document_agent.services.presentation_closing import render_closing


def _deck():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    return deck


def _result():
    observations = []
    for metric, unit, periods, values, page in (
        ("Number of reserve units", "units", ("2022-12-31", "2024-05-31"), (80, 50), 20),
        ("Active sites", "units", ("FY2022", "FY2024"), (10, 18), 5),
        ("Recovery rate", "percent", ("FY2022", "FY2024"), (22, 31), 6),
    ):
        for index, (period, value) in enumerate(zip(periods, values)):
            observations.append(Observation(
                id=f"{metric}-{index}", metric_original=metric, value=value,
                raw_value=f"{value}.00", unit=unit, raw_unit=unit, period=period,
                period_type="point_in_time" if page == 20 else "fiscal_year",
                as_of_date=period if page == 20 else None,
                period_basis="point_in_time" if page == 20 else "FY",
                audited_status="unaudited" if page == 20 and index == 1 else "unknown",
                evidence=[SourceEvidence(page=page, text=f"{value}.00", row_label=metric,
                                         extraction_method="digital_table", confidence=.99)],
                confidence=.99,
            ))
    analyses = [AnalysisResult(
        task_id=identity, title=identity, result_type="calculated_result", result={},
        input_observation_ids=[o.id for o in observations if (o.source_page == 20) == point], confidence=.99,
    ) for identity, point in (("reservoir", True), ("sites", False))]
    insights = [
        Insight(id="reservoir", title="Equipment reserves declined", narrative="Reported equipment reserves declined.",
                kind="interpretation", implication="Fewer reserve units limit the available buffer.",
                watch_item="Monitor the next reported reserve count.", result_ids=["reservoir"],
                evidence=[observations[0].evidence[0]], confidence=.99),
        Insight(id="sites", title="More sites and a higher recovery rate", narrative="Both reported measures increased.",
                kind="interpretation", implication="The network expanded alongside an improved recovery rate.",
                result_ids=["sites"], evidence=[observations[2].evidence[0], observations[4].evidence[0]], confidence=.99),
    ]
    result = PipelineResult(
        document=ParsedDocument(document_id="generic-survey", sha256="generic", safe_filename="source.pdf", page_count=20),
        profile=DocumentProfile(), observations=observations, analysis_results=analyses, insights=insights,
    )
    plan = PresentationSlide(id="closing", slide_type="risks", title="Conclusions and Watch Items",
                             insight_ids=[i.id for i in insights],
                             bullets=[insights[0].implication, insights[1].implication, insights[0].watch_item])
    return result, plan


def _copy(slide):
    return "\n".join(s.text for s in slide.shapes if s.has_text_frame)


def _tables(slide):
    return [[[c.text for c in row.cells] for row in s.table.rows] for s in slide.shapes if s.has_table]


def test_independent_conclusions_and_their_evidence_share_two_compatible_pages():
    result, plan = _result()
    before, original_plan = result.model_dump(), plan.model_dump()
    deck = _deck()
    slides = render_closing(deck, result, plan)
    assert len(slides) == len(deck.slides) == 2
    assert result.insights[0].implication in _copy(slides[0])
    assert result.insights[0].watch_item in _copy(slides[0])
    assert result.insights[1].implication in _copy(slides[1])
    assert result.insights[1].implication not in _copy(slides[0])
    assert result.insights[0].watch_item not in _copy(slides[1])
    assert _tables(slides[0]) == [[
        ["Reported values (units)", "31 Dec 2022", "31 May 2024*"], ["Number of reserve units", "80.00", "50.00"],
    ]]
    annual_tables = _tables(slides[1])
    assert len(annual_tables) == 2
    assert all(table[0][1:] == ["FY2022", "FY2024"] for table in annual_tables)
    assert {table[0][0] for table in annual_tables} == {"Reported values (units)", "Reported values (%)"}
    assert "p. 20" in _copy(slides[0]) and "* Unaudited" in _copy(slides[0])
    assert "p. 5-6" in _copy(slides[1])
    for slide in slides:
        table_names = [shape.name for shape in slide.shapes if shape.has_table]
        assert len(table_names) == len(set(table_names))
        notes = slide.notes_slide.notes_text_frame.text
        assert all(o.id in notes and o.raw_value in notes for o in result.observations)
        assert all(shape.top + shape.height < deck.slide_height - Inches(1.0)
                   for shape in slide.shapes if shape.has_table or shape.name == "closing:body")
    assert result.model_dump() == before and plan.model_dump() == original_plan


def test_a_conclusion_without_watch_item_still_shares_its_evidence():
    result, plan = _result()
    plan.insight_ids = ["sites"]
    plan.bullets = [result.insights[1].implication]
    slides = render_closing(_deck(), result, plan)
    assert len(slides) == 1 and len(_tables(slides[0])) == 2
    assert plan.bullets[0] in _copy(slides[0])
    assert "Watch items" not in _copy(slides[0])


def test_short_findings_with_identical_periods_and_units_share_one_closing_page():
    result, plan = _result()
    for observation in result.observations[:2]:
        observation.period = observation.period.replace("2022-12-31", "FY2022").replace("2024-05-31", "FY2024")
        observation.period_type, observation.period_basis = "fiscal_year", "FY"
        observation.as_of_date = None
        observation.audited_status = "unknown"
    plan.insight_ids = []
    plan.bullets = ["The first reported measure changed.", "The second reported measure changed."]
    plan.bullet_observation_ids = [[o.id for o in result.observations[:2]],
                                   [o.id for o in result.observations[2:4]]]
    plan.observation_ids = [oid for group in plan.bullet_observation_ids for oid in group]
    slides = render_closing(_deck(), result, plan)
    assert len(slides) == 1
    assert all(bullet in _copy(slides[0]) for bullet in plan.bullets)
    assert len(_tables(slides[0])) == 1
    assert {row[0] for row in _tables(slides[0])[0][1:]} == {"Number of reserve units", "Active sites"}
    assert all(o.raw_value in slides[0].notes_slide.notes_text_frame.text for o in result.observations[:4])


def test_explicit_bullet_inputs_can_bind_a_reworded_conclusion():
    result, plan = _result()
    plan.insight_ids = []
    plan.observation_ids = [o.id for o in result.observations[:2]]
    plan.bullets = ["A retained model finding with explicitly selected source records."]
    plan.bullet_observation_ids = [list(plan.observation_ids)]
    slides = render_closing(_deck(), result, plan)
    assert len(slides) == 1
    assert plan.bullets[0] in _copy(slides[0]) and _tables(slides[0])


def test_shared_source_pages_do_not_assign_an_unlinked_bullet_to_evidence():
    result, plan = _result()
    plan.bullets = ["A separate retained observation with no explicit input link."]
    plan.source_pages = [20]
    slides = render_closing(_deck(), result, plan)
    text_slide = next(slide for slide in slides if plan.bullets[0] in _copy(slide))
    assert not _tables(text_slide)
    assert any(_tables(slide) for slide in slides)


def test_incompatible_periods_within_one_claim_keep_separate_tables():
    result, plan = _result()
    plan.insight_ids = ["sites"]
    plan.bullets = [result.insights[1].implication]
    for observation in result.observations[4:]:
        observation.period = observation.period.replace("FY", "6M")
        observation.period_type, observation.period_basis = "interim_flow", "6M"
    before = result.model_dump()
    slides = render_closing(_deck(), result, plan)
    assert not _tables(slides[0])
    tables = [table for slide in slides for table in _tables(slide)]
    assert tables
    assert all(not (any("FY" in cell for cell in table[0]) and any("6M" in cell for cell in table[0])) for table in tables)
    assert result.model_dump() == before


def test_conflicting_raw_values_remain_visible_and_are_not_joined_into_endpoints():
    result, plan = _result()
    conflict = result.observations[2].model_copy(deep=True, update={"id": "conflict", "value": 99, "raw_value": "99.00"})
    result.observations.append(conflict)
    result.analysis_results[1].input_observation_ids.append(conflict.id)
    before = deepcopy(result)
    slides = render_closing(_deck(), result, plan)
    all_cells = [cell for slide in slides for table in _tables(slide) for row in table for cell in row]
    assert "99.00" in all_cells and "10.00" in all_cells
    assert result == before


def test_long_linked_copy_falls_back_without_leaking_trial_slides_or_losing_copy():
    result, plan = _result()
    long_copy = "The evidence supports this limited finding with the stated uncertainty. " * 24
    result.insights[1].implication = long_copy
    plan.bullets[1] = long_copy
    deck = _deck()
    slides = render_closing(deck, result, plan)
    assert len(slides) == len(deck.slides)
    assert len({slide.part.partname for slide in slides}) == len(slides)
    bodies = "".join(shape.text for slide in slides for shape in slide.shapes if shape.name == "closing:body")
    assert long_copy.strip() in bodies
    assert all(any(o.id in slide.notes_slide.notes_text_frame.text for slide in slides) for o in result.observations)


def test_extra_explicit_facts_remain_visible_after_linked_bundles():
    result, plan = _result()
    extra = result.observations[2].model_copy(deep=True, update={"id": "extra", "metric_original": "Extra disclosed measure", "value": 123})
    extra.evidence[0].row_label = "Extra disclosed measure"
    result.observations.append(extra)
    plan.observation_ids = [extra.id]
    slides = render_closing(_deck(), result, plan)
    assert len(slides) == 3
    assert any("Extra disclosed measure" in row for table in _tables(slides[-1]) for row in table)
