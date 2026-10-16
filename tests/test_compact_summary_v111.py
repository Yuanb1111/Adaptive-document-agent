"""Hard slide budget, genuine part coverage, and source-preserving compact rendering."""
from io import BytesIO
import json

import pytest
from pptx import Presentation
from tests.summary_author_helpers import authored_test_draft

from adaptive_document_agent.agent.compact_summary import generate_compact_summary,CompactSummaryDraft,CompactSummaryReview
from adaptive_document_agent.models import PipelineResult,ParsedDocument,DocumentPage,DocumentProfile,PresentationPlan,PresentationSlide
from adaptive_document_agent.models.presentation import CompanyProfile
from adaptive_document_agent.models.summary import SummarySourceBlock,SummaryPart,SummaryFact,SummaryReview,SummaryPageRange
from adaptive_document_agent.services.pptx_export import BUNDLED_TEMPLATE_PATH,_add_company_at_a_glance,_appendix_display_value
from adaptive_document_agent.services.summary_validation import presentation_errors
from adaptive_document_agent.services.summary_export import verify_summary_export


def fixture(count=70):
    names=[chr(65+i//26)+chr(65+i%26) for i in range(count)]
    lines=[f'Team {name} supports scheduling and requires staff training before service delivery.\n' for name in names]
    text=''.join(lines)
    result=PipelineResult(document=ParsedDocument(document_id='compact',sha256='fixture',safe_filename='fixture.pdf',
        page_count=1,pages=[DocumentPage(page_number=1,text=text)]),profile=DocumentProfile())
    review=SummaryReview(document_id='compact',document_sha256='fixture',source_pages=[1],max_intro_slides=5,
        scope_ranges=[SummaryPageRange(start_page=1,end_page=1,reason='Complete introductory source.')],
        source_blocks=[SummarySourceBlock(id='block',page=1,char_start=0,char_end=len(text),text=text)],
        parts=[SummaryPart(id=f'part-{i}',title='Scheduling',role='content',block_id='block',start_line=i+1,end_line=i+1,
            source_page=1,source_text=line,reading_note='Service depends on staff training.',
            facts=[SummaryFact(label='Scheduling',text=line.strip(),source_quote=line.strip(),source_pages=[1])])
            for i,line in enumerate(lines)])
    pages=[]
    for offset in range(0,count,14):
        group=names[offset:offset+14]
        # Explicit fixture synthesis, retaining the shared source qualification.
        copy='Teams '+', '.join(group)+' support scheduling; staff training is required before service delivery.'
        pages.append(dict(title='Scheduling and delivery',items=[dict(label='Service prerequisites',text=copy,
                     fact_ids=[f'f{i}' for i in range(offset,min(count,offset+14))])]))
    return result,review,dict(pages=pages)


class Reviewer:
    def __init__(self,draft,*,reject=None,identity=None):self.draft=draft;self.reject=reject;self.identity=identity;self.calls=[]
    def generate_structured(self,messages,schema,**kwargs):
        self.calls.append(schema.__name__)
        if issubclass(schema, CompactSummaryDraft):return schema.model_validate(authored_test_draft(self.draft, messages))
        if schema.__name__ == 'SummaryLayoutPages':
            return schema.model_validate({'pages': authored_test_draft(self.draft, messages)['pages']})
        if schema.__name__ == 'SummarySynthesisRevision':
            payload=json.loads(messages[1]['content'].split('\n',1)[1].rsplit('\n',1)[0])
            replacements, additions = [], []
            candidate = authored_test_draft(self.draft, messages)
            for p, page in enumerate(candidate['pages']):
                existing = payload['editable_groups'][p]['items']
                for i, item in enumerate(page['items']):
                    if i >= len(existing): additions.append(dict(group_id=f'g{p}', item=item))
                    elif any(existing[i][k] != item[k] for k in ('label','text','fact_ids','meaning_refs')):
                        replacements.append(dict(item_id=f'g{p}_i{i}', replacement=item))
            return schema(replacements=replacements, additions=additions)
        payload=json.loads(messages[-1]['content'].split('\n',1)[1].rsplit('\n',1)[0])
        verdicts = [dict(part_id=p['id'],accepted=p['id']!=self.reject,
            reason='Training dependency is retained.' if p['id']!=self.reject else 'Source meaning is absent.')
            for p in payload['source']['parts']]
        for verdict, part in zip(verdicts, payload['source']['parts']):
            if 'required_meanings' in part:
                text = ' '.join(it['text'] for pg in payload['pages'] for it in pg['items'] if part['id'] in it['part_ids'])
                verdict['coverage'] = [dict(meaning_index=i, covered=True, summary_quote=text,
                    reason='The source condition is present in assigned copy.') for i in range(len(part['required_meanings']))]
        return schema(opening_slide_id=(payload.get('opening_target') or {}).get('slide_id'),
            opening_item_texts=(payload.get('opening_target') or {}).get('item_texts'), opening_accepted=True,
            opening_reason='The opening explains the scheduling service and training prerequisite.',
            identity_accepted=self.identity, parts=verdicts)


def test_seventy_source_parts_fit_five_visible_slides_with_exact_bindings():
    result,review,draft=fixture()
    company=generate_compact_summary(Reviewer(draft),result,review,CompanyProfile(summary_review=review))
    review.status='complete'
    assert len(company.summary_pages)==5 and not presentation_errors(review,company.summary_pages,result)
    result.presentation_plan=PresentationPlan(title='Source deck',company=company)
    deck=Presentation(str(BUNDLED_TEMPLATE_PATH))
    for ident in list(deck.slides._sldIdLst):deck.part.drop_rel(ident.rId);deck.slides._sldIdLst.remove(ident)
    _add_company_at_a_glance(deck,result,PresentationSlide(id='company',slide_type='company_overview',title='Introduction'))
    assert len(deck.slides)==5
    used=verify_summary_export(deck,result)
    assert len(used)==70 and {n for ns in used.values() for n in ns}=={1,2,3,4,5}
    stream=BytesIO();deck.save(stream)
    assert verify_summary_export(Presentation(stream),PipelineResult.model_validate_json(result.model_dump_json()))==used
    company.summary_pages[0].items[0].text+=' Missing condition.'
    assert any('independent' in e for e in presentation_errors(review,company.summary_pages,result))


def test_citing_all_ids_cannot_override_failed_part_meaning_review():
    result,review,draft=fixture()
    client=Reviewer(draft,reject='part-69')
    with pytest.raises(ValueError,match='Source meaning is absent'):
        generate_compact_summary(client,result,review,CompanyProfile(summary_review=review))
    assert client.calls==['CompactSummaryDraft','CompactSummaryReview','SummarySynthesisRevision']
    assert review.plan_audits[-1]['revision_failed'] is True


def test_omitted_part_and_changed_fact_bindings_fail_closed():
    result,review,draft=fixture()
    draft['pages'][-1]['items'][0]['fact_ids'].pop()
    with pytest.raises(ValueError,match='Every substantive Summary part'):
        generate_compact_summary(Reviewer(draft),result,review,CompanyProfile(summary_review=review))
    result,review,draft=fixture()
    company=generate_compact_summary(Reviewer(draft),result,review,CompanyProfile(summary_review=review))
    company.summary_pages[0].items[0].source_refs[0].fact_index=9
    assert any('unknown source fact' in e for e in presentation_errors(review,company.summary_pages,result))


def test_six_page_model_plan_is_rejected_instead_of_silently_truncated():
    _,_,draft=fixture()
    draft['pages'].append(draft['pages'][0])
    with pytest.raises(ValueError):CompactSummaryDraft.model_validate(draft)


def test_numeric_claim_still_requires_its_exact_source_currency_value_and_scale():
    result,review,draft=fixture()
    draft['pages'][0]['items'][0]['text']='The teams generated RMB999 million revenue.'
    with pytest.raises(ValueError,match='unsupported numeric claims'):
        generate_compact_summary(Reviewer(draft),result,review,CompanyProfile(summary_review=review))


def test_literal_other_entity_is_not_accepted_as_company_identity():
    result,review,draft=fixture()
    draft.update(name='Team AA',name_quote=review.parts[0].source_text.strip(),name_page=1)
    company=CompanyProfile(summary_review=review)
    with pytest.raises(ValueError,match='identity was not independently accepted'):
        generate_compact_summary(Reviewer(draft,identity=False),result,review,company)
    assert company.name=='' and company.identity_state=='UNRESOLVED'


def test_real_adjusted_loss_tie_uses_consistent_half_up_rounding():
    from adaptive_document_agent.models import Observation
    value=Observation(id='value',metric_original='Adjusted net loss',value=-44675000,raw_value='(44,675)',
                      currency='RMB',unit_scale=1000,unit='currency',raw_unit='RMB in thousands',confidence=.9)
    from adaptive_document_agent.document_model.metric_semantic_classifier import classify_metric
    assert _appendix_display_value(value,classify_metric('Adjusted net loss',unit='currency',raw_unit='RMB in thousands'))=='-44.68'
