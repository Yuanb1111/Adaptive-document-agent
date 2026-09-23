"""Regression tests for generic Company Overview pipeline improvements.

Validates:
1. Company discovery finds profile outside financial analysis pages.
2. Incomplete product fragments and corruptions are rejected.
3. Missing industry is not invented or guessed.
4. Missing stock code is not invented or guessed.
5. Resolved vs unresolved company name detection and slide title alignment.
6. Company Overview vs Data Quality consistency (no contradictions).
7. Fallback behavior: only supported cards are rendered, no generic filler.
"""

from __future__ import annotations

import pytest

from adaptive_document_agent.models import (
    CompanyProfile,
    DocumentPage,
    DocumentProfile,
    ParsedDocument,
    PipelineResult,
    PresentationPlan,
    PresentationSlide,
    ReportPlan,
    ValidationIssue,
)
from adaptive_document_agent.services.company_discovery import CompanyProfileDiscovery
from adaptive_document_agent.services.company_extractor import (
    extract_structured_company_fields,
    is_company_identity_resolved,
    is_generic_name,
    validate_business_model,
    validate_company_name,
    validate_geography,
    validate_industry,
    validate_market_position,
    validate_product,
    validate_stock_code,
)
from adaptive_document_agent.services.qa_reporter import (
    run_comprehensive_qa,
    sanitize_company_identity_contradictions,
)


@pytest.mark.parametrize("hq_pages", [[103], []])
def test_cover_name_and_issuer_bound_hq_exclude_competitor_and_investor(hq_pages):
    pages = [
        DocumentPage(page_number=1, text="GLOBAL OFFERING\nEXAMPLE MOTION CORP LTD\nStock Code: 4321"),
        DocumentPage(page_number=103, text=("INDUSTRY OVERVIEW\nTop 10 market players\n"
            "Company A\nEstablished in 2005, headquartered in Denmark.\n"
            "Our Company\nEstablished in 2015, headquartered in Shenzhen, China.\n"
            "Company B\nEstablished in 2008, headquartered in Shanghai, China.")),
        DocumentPage(page_number=151, text="HISTORY AND CORPORATE STRUCTURE\nGroup Co., Ltd.\nInvestor details."),
    ]
    company = CompanyProfile(name="Group Co., Ltd.", headquarters="Denmark",
                             field_source_pages={"name": [151], "headquarters": hq_pages})
    result = _make_dummy_pipeline_result(pages, company)
    recovered = extract_structured_company_fields(company, result)
    assert recovered.name == "EXAMPLE MOTION CORP LTD"
    assert recovered.field_source_pages["name"] == [1]
    assert recovered.headquarters == "Shenzhen, China"
    assert recovered.field_source_pages["headquarters"] == [103]
    assert recovered.market_position == ""


def test_suffix_fragment_and_listing_heading_are_not_identity_values():
    from adaptive_document_agent.services.company_extractor import validate_headquarters, validate_listing_market
    assert validate_company_name("Group Co., Ltd.") == ""
    assert validate_headquarters("Ed In Denmark") == ""
    assert validate_listing_market("For Listing On The Stock Exchange") == ""
    result = _make_dummy_pipeline_result([
        DocumentPage(page_number=1, text="Prospectus"),
        DocumentPage(page_number=20, text="APPLICATION FOR LISTING ON THE STOCK EXCHANGE"),
        DocumentPage(page_number=151, text="HISTORY AND CORPORATE STRUCTURE\nInvestor Holdings Limited"),
    ])
    recovered = extract_structured_company_fields(CompanyProfile(), result)
    assert recovered.identity_state == "UNRESOLVED"
    assert recovered.name == "" and recovered.listing_market == ""


def test_multiple_cover_entities_do_not_choose_the_first_legal_name():
    result = _make_dummy_pipeline_result([
        DocumentPage(page_number=1, text="Global Offering\nSponsor Holdings Limited\nIssuer Holdings Limited"),
    ])
    recovered = extract_structured_company_fields(CompanyProfile(name="Unnamed robotics company"), result)
    assert recovered.identity_state == "UNRESOLVED"
    assert recovered.name == ""


def test_explicit_application_areas_are_cited_without_inference():
    result = _make_dummy_pipeline_result([
        DocumentPage(page_number=1, text="GLOBAL OFFERING\nACME MOTION CORP LTD"),
        DocumentPage(page_number=2, text=("BUSINESS OVERVIEW\nWe build automation equipment with "
            "use cases in manufacturing, retail, and healthcare settings. "
            "Our sales are through distributors.")),
    ])
    recovered = extract_structured_company_fields(CompanyProfile(), result)
    assert recovered.application_areas == ["manufacturing", "retail", "healthcare settings"]
    assert recovered.field_source_pages["application_areas"] == [2]

    unsupported = CompanyProfile(application_areas=["aerospace"],
        field_source_pages={"application_areas": [2]})
    assert "aerospace" not in extract_structured_company_fields(unsupported, result).application_areas


@pytest.mark.parametrize("listing_text", [
    "(payable in full on application in Hong\nKong dollars, subject to refund)\nNominal value : RMB1 per share",
    "Payable on application in Hong Kong dollars, subject to refund.",
])
def test_listing_application_is_not_a_business_use_case(listing_text: str):
    result = _make_dummy_pipeline_result([
        DocumentPage(page_number=1, text="GLOBAL OFFERING\nACME MOTION CORP LTD\n" + listing_text),
        DocumentPage(page_number=2, text="BUSINESS OVERVIEW\nThe company sells motion equipment."),
    ])
    recovered = extract_structured_company_fields(CompanyProfile(), result)
    assert recovered.application_areas == []
    assert "application_areas" not in recovered.field_source_pages


def test_products_listed_in_issuer_prose_keep_source_and_drop_price_qualifier():
    result = _make_dummy_pipeline_result([
        DocumentPage(page_number=1, text="ACME DRINKS LIMITED"),
        DocumentPage(page_number=2, text=(
            "BUSINESS\nOVERVIEW\nWe provide value-for-money products to consumers, "
            "including fruit drinks, tea drinks, ice\ncream and coffee, typically "
            "priced around one dollar per item."
        )),
    ])
    recovered = extract_structured_company_fields(CompanyProfile(), result)
    assert recovered.products == ["fruit drinks", "tea drinks", "ice cream", "coffee"]
    assert recovered.field_source_pages["products"] == [2]


def test_explicit_issuer_model_in_prose_is_cited():
    result = _make_dummy_pipeline_result([
        DocumentPage(page_number=1, text="ACME DRINKS LIMITED"),
        DocumentPage(page_number=2, text=(
            "BUSINESS\nOVERVIEW\nWe serve consumers with freshly-made drinks. "
            "Through a franchise model, we have cultivated a store network."
        )),
    ])
    recovered = extract_structured_company_fields(CompanyProfile(), result)
    assert recovered.business_model == "Franchise model"
    assert recovered.field_source_pages["business_model"] == [2]


def _make_dummy_pipeline_result(
    pages: list[DocumentPage],
    company: CompanyProfile | None = None,
    data_quality_notes: list[str] | None = None,
    analysis_page_ranges: list[tuple[int, int]] | None = None,
) -> PipelineResult:
    doc = ParsedDocument(
        document_id="doc_test_123",
        sha256="0123456789abcdef0123456789abcdef0123456789abcdef0123456789abcdef",
        safe_filename="sample_test.pdf",
        pages=pages,
        page_count=len(pages),
    )
    profile = DocumentProfile(
        document_type="Prospectus",
        document_summary="Summary of document operations.",
        document_purpose="Corporate intelligence overview",
        overview_title=company.name if company and company.name else "Document Overview",
        document_summary_pages=[1],
        analysis_page_ranges=analysis_page_ranges or [],
        data_quality_notes=data_quality_notes or [],
    )
    comp = company or CompanyProfile()
    slide_title = "Company at a Glance" if is_company_identity_resolved(comp) else "Document at a Glance"
    slides = [
        PresentationSlide(id="cover", slide_type="cover", title="Presentation Title"),
        PresentationSlide(id="slide_company_overview", slide_type="company_overview", title=slide_title),
        PresentationSlide(id="exec_sum", slide_type="executive_summary", title="Executive Summary"),
        PresentationSlide(id="quality", slide_type="data_quality", title="Data Quality & Scope"),
        PresentationSlide(id="appendix", slide_type="appendix", title="Appendix"),
    ]
    plan = PresentationPlan(
        title="Corporate Overview Deck",
        company=comp,
        slides=slides,
    )
    return PipelineResult(
        document=doc,
        profile=profile,
        observations=[],
        presentation_plan=plan,
        report_plan=ReportPlan(title="Report", sections=[]),
    )


# =============================================================================
# 1. Company discovery outside financial analysis pages
# =============================================================================

def test_company_discovery_finds_profile_outside_financial_pages() -> None:
    """Company Overview must discover business/products/listing facts on pages outside financial scope."""
    p1 = DocumentPage(page_number=1, text="Global Offering Prospectus\nAcme Biotech Holdings Limited\nStock Code: 9988.HK\nHong Kong Stock Exchange")
    p2 = DocumentPage(page_number=2, text="Table of Contents\n1. Business Overview ... Page 15\n2. Financial Information ... Page 100")
    # Business overview pages (pages 15-16)
    p15 = DocumentPage(
        page_number=15,
        text=(
            "BUSINESS OVERVIEW\n"
            "We are a commercial-stage biopharmaceutical provider operating in the Biotechnology space. "
            "Our core products include AB-101 Oncology Antibody and AB-202 Cellular Therapy. "
            "We operate a direct sales model and distributor network sales to tertiary hospitals."
        ),
    )
    p16 = DocumentPage(
        page_number=16,
        text=(
            "MARKETS AND COMPETITIVE POSITION\n"
            "Our principal markets include Mainland China, Hong Kong, and North America. "
            "The Company holds a 24.5% market share in targeted oncology therapies and is ranked No. 2 nationally. "
            "Headquarters: Shanghai, China. Reporting Currency: RMB. Track record period covers FY2021 to FY2023."
        ),
    )
    # Financial statements pages (pages 100-101)
    p100 = DocumentPage(
        page_number=100,
        text=(
            "FINANCIAL INFORMATION\n"
            "Consolidated statement of financial position\n"
            "Revenue: 100,000, 120,000, 150,000\n"
            "Gross profit: 45,000, 55,000, 70,000\n"
            "The Group's operations continued to expand."
        ),
    )
    p101 = DocumentPage(
        page_number=101,
        text=(
            "Notes to consolidated financial statements\n"
            "Operating profit and trade receivables disclosures."
        ),
    )

    all_pages = [p1, p2, p15, p16, p100, p101]
    # Set financial scope strictly to pages 100-101
    result = _make_dummy_pipeline_result(
        all_pages,
        company=CompanyProfile(),
        analysis_page_ranges=[(100, 101)],
    )

    # 1. Test dedicated discovery ranking
    discovered_pages = CompanyProfileDiscovery.rank_profile_pages(result.document, result.profile, max_pages=6)
    assert 1 in discovered_pages  # Cover
    assert 15 in discovered_pages  # Business Overview
    assert 16 in discovered_pages  # Markets & Competitive Position

    # 2. Extract structured fields
    enriched = extract_structured_company_fields(result.presentation_plan.company, result)

    assert enriched.name == "Acme Biotech Holdings Limited"
    assert enriched.identity_state == "RESOLVED"
    assert enriched.stock_code == "9988.HK"
    assert enriched.listing_market == "HKEX" or "Hong Kong" in enriched.listing_market or "Main Board" in enriched.listing_market
    assert enriched.industry == "Biotechnology"
    assert any("AB-101 Oncology Antibody" in p for p in enriched.products)
    assert any("AB-202 Cellular Therapy" in p for p in enriched.products)
    assert "Direct sales model" in enriched.business_model or "Distributor network sales" in enriched.business_model
    assert "Mainland China" in enriched.geographies
    assert "24.5%" in enriched.market_position and ("market share" in enriched.market_position.lower() or "ranked no. 2" in enriched.market_position.lower())
    assert "Shanghai" in enriched.headquarters
    assert enriched.reporting_currency == "RMB"
    assert "FY2021" in enriched.track_record_period

    # Verify per-field source pages are recorded
    assert 1 in enriched.field_source_pages.get("name", []) or 15 in enriched.field_source_pages.get("name", [])
    assert 15 in enriched.field_source_pages.get("products", [])
    assert 16 in enriched.field_source_pages.get("geographies", []) or 16 in enriched.field_source_pages.get("market_position", [])


def test_profile_discovery_keeps_business_overview_continuation_ahead_of_risk_hits() -> None:
    pages = [
        DocumentPage(page_number=1, text="ACME ROBOTICS LIMITED\nStock Code: 1234"),
        DocumentPage(page_number=10, text=(
            "OVERVIEW\nWe are a developer and manufacturer of collaborative robots. "
            "Our products cover industrial automation and education."
        )),
        DocumentPage(page_number=11, text="Our product portfolio includes Series A and Series B robots."),
        DocumentPage(page_number=12, text="Our sales network serves customers across several markets."),
        DocumentPage(page_number=40, text=(
            "Risk factors\nWe cannot assure you that distributors will buy our products. "
            "This could materially and adversely affect our business."
        )),
    ]
    result = _make_dummy_pipeline_result(pages)
    selected = CompanyProfileDiscovery.rank_profile_pages(result.document, result.profile, max_pages=4)
    assert selected == [1, 10, 11, 12]


def test_profile_discovery_keeps_readable_cover_and_consumer_business_overview() -> None:
    pages = [
        DocumentPage(page_number=1, text=""),
        DocumentPage(page_number=2, text="GLOBAL OFFERING\nACME DRINKS LIMITED\nStock code: 1234"),
        DocumentPage(page_number=10, text="SUMMARY\nOur store network serves consumers."),
        DocumentPage(page_number=100, text=(
            "BUSINESS\nOVERVIEW\nWe are a drinks company providing freshly-made tea and coffee "
            "to consumers. We have stores in multiple countries."
        )),
        DocumentPage(page_number=101, text="Our brand sells drinks through franchised stores."),
        DocumentPage(page_number=102, text="We operate a network of franchisees."),
        DocumentPage(page_number=200, text="FINANCIAL INFORMATION\nRevenue and profit."),
    ]
    result = _make_dummy_pipeline_result(pages)
    selected = CompanyProfileDiscovery.rank_profile_pages(result.document, result.profile, max_pages=5)
    assert selected == [2, 100, 101, 102]


def test_prospectus_distribution_is_not_a_company_product() -> None:
    assert validate_product("printed copies of this prospectus") == ""
    result = _make_dummy_pipeline_result([
        DocumentPage(page_number=1, text="ACME ROBOTICS LIMITED"),
        DocumentPage(page_number=2, text="The company provides printed copies of this prospectus on request."),
    ])
    recovered = extract_structured_company_fields(CompanyProfile(name="ACME ROBOTICS LIMITED"), result)
    assert recovered.products == []


def test_issuer_product_series_are_recovered_without_forecast_period() -> None:
    result = _make_dummy_pipeline_result([
        DocumentPage(page_number=1, text="ACME ROBOTICS LIMITED"),
        DocumentPage(page_number=10, text=(
            "OVERVIEW\nWe are a company that specializes in the development, manufacturing\n"
            "and commercialization of collaborative robots. "
            "Our products cover industrial automation."
        )),
        DocumentPage(page_number=11, text=(
            "OUR PRODUCTS\nCR Series\nOur CR Series includes collaborative robot models.\n"
            "Nova Series\nThe Nova Series features lightweight robot models."
        )),
        DocumentPage(page_number=12, text="Industry revenue is forecast to grow from 2023 to 2028."),
    ])
    company = CompanyProfile(
        name="ACME ROBOTICS LIMITED",
        products=["printed copies of this prospectus"],
        track_record_period="2023 to 2028",
        field_source_pages={"products": [1], "track_record_period": [12]},
    )
    recovered = extract_structured_company_fields(company, result)
    assert recovered.products == ["CR Series", "Nova Series"]
    assert recovered.one_line_description == "Development, manufacturing and commercialization of collaborative robots"
    assert recovered.field_source_pages["products"] == [11]
    assert recovered.track_record_period == ""


# =============================================================================
# 2. Incomplete product fragments and corruptions
# =============================================================================

def test_incomplete_product_fragments_rejected() -> None:
    """Validator must reject broken fragments such as 'd excerpts' while retaining valid products."""
    # Corrupted / fragment cases
    assert validate_product("d excerpts") == ""
    assert validate_product("s products") == ""
    assert validate_product("ing solutions") == ""
    assert validate_product("ed systems") == ""
    assert validate_product("and services") == ""
    assert validate_product("including ") == ""
    assert validate_product("excerpts") == ""
    assert validate_product("products") == ""
    assert validate_product("n/a") == ""
    assert validate_product("Table of Contents") == ""
    assert validate_product("Page 45") == ""
    assert validate_product("abc") == ""  # Too short

    # Valid products with surrounding artifacts cleaned
    assert validate_product("• Cloud ERP Software Suite") == "Cloud ERP Software Suite"
    assert validate_product("- Industrial Robotic Arms -") == "Industrial Robotic Arms"
    assert validate_product("provides High-Speed Optical Transceivers") == "High-Speed Optical Transceivers"
    assert validate_product("Next-Gen Sequencing Reagents, and") == "Next-Gen Sequencing Reagents"

    # End-to-end extraction from messy text
    messy_text = (
        "BUSINESS OVERVIEW\n"
        "Our main offerings include d excerpts, s offerings, Next-Gen Sequencing Reagents, "
        "and High-Speed Optical Transceivers, consisting of and products."
    )
    p = DocumentPage(page_number=5, text=messy_text)
    res = _make_dummy_pipeline_result([p])
    extracted = extract_structured_company_fields(CompanyProfile(), res)

    assert "d excerpts" not in extracted.products
    assert "s offerings" not in extracted.products
    assert "and products" not in extracted.products
    assert "Next-Gen Sequencing Reagents" in extracted.products
    assert "High-Speed Optical Transceivers" in extracted.products


# =============================================================================
# 3. Missing industry not invented
# =============================================================================

def test_missing_industry_not_invented() -> None:
    """When industry is absent from source pages, it must remain empty and not be guessed."""
    p = DocumentPage(
        page_number=1,
        text=(
            "COMPANY OVERVIEW\n"
            "Vertex Innovations Ltd.\n"
            "We provide proprietary data processing platforms and client analytics. "
            "Headquarters: Singapore. Track record: FY2021-FY2023."
        ),
    )
    res = _make_dummy_pipeline_result([p])
    extracted = extract_structured_company_fields(CompanyProfile(), res)

    assert extracted.name == "Vertex Innovations Ltd."
    # Industry was never mentioned in the text
    assert extracted.industry == ""

    # Ensure field validator also rejects generic words
    assert validate_industry("Industry") == ""
    assert validate_industry("Sector") == ""
    assert validate_industry("Business") == ""
    assert validate_industry("General") == ""
    assert validate_industry("n/a") == ""


# =============================================================================
# 4. Missing stock code not invented
# =============================================================================

def test_missing_stock_code_not_invented() -> None:
    """When company is unlisted or lacks a ticker, stock code must remain empty without dummy fallbacks."""
    p = DocumentPage(
        page_number=1,
        text=(
            "PRIVATE COMPANY PROFILE\n"
            "Apex Logistics Global Pte. Ltd.\n"
            "Operating across Southeast Asia. Core products include Freight Forwarding and Cold Chain Express."
        ),
    )
    res = _make_dummy_pipeline_result([p])
    extracted = extract_structured_company_fields(CompanyProfile(), res)

    assert extracted.name == "Apex Logistics Global Pte. Ltd."
    assert extracted.stock_code == ""
    assert extracted.listing_market == ""

    # Ensure stock code validator rejects non-code strings
    assert validate_stock_code("None") == ""
    assert validate_stock_code("N/A") == ""
    assert validate_stock_code("Unknown") == ""
    assert validate_stock_code("Private Company") == ""
    assert validate_stock_code("1234.HK") == "1234.HK"
    assert validate_stock_code("600519") == "600519"
    assert validate_stock_code("AAPL") == "AAPL"


# =============================================================================
# 5. Resolved vs unresolved company name
# =============================================================================

def test_resolved_vs_unresolved_company_name() -> None:
    """Generic or missing company names must be UNRESOLVED; real entity names must be RESOLVED."""
    # Generic names must be identified as generic
    assert is_generic_name("") is True
    assert is_generic_name("company overview") is True
    assert is_generic_name("Document Overview") is True
    assert is_generic_name("Unnamed Issuer") is True
    assert is_generic_name("Company not identified") is True
    assert is_generic_name("The Company") is True
    assert is_generic_name("The Group") is True
    assert is_generic_name("Issuer") is True

    # Valid company names
    assert is_generic_name("Baidu, Inc.") is False
    assert is_generic_name("Tencent Holdings Limited") is False
    assert is_generic_name("Innovent Biologics, Inc.") is False

    # Resolution helper
    unresolved_comp = CompanyProfile(name="Document Overview", identity_state="UNRESOLVED")
    assert is_company_identity_resolved(unresolved_comp) is False

    generic_comp = CompanyProfile(name="Unnamed Issuer", identity_state="RESOLVED")
    assert is_company_identity_resolved(generic_comp) is False

    resolved_comp = CompanyProfile(name="Acme Tech Co., Ltd.", identity_state="RESOLVED")
    assert is_company_identity_resolved(resolved_comp) is True


# =============================================================================
# 6. Company Overview vs Data Quality consistency
# =============================================================================

def test_company_overview_vs_data_quality_consistency_resolved() -> None:
    """When company is resolved, Data Quality notes must not contradict by stating issuer is unnamed."""
    p = DocumentPage(page_number=1, text="Annual Report\nHorizon Robotics Inc.\nArtificial Intelligence Solutions")
    comp = CompanyProfile(name="Horizon Robotics Inc.", identity_state="RESOLVED")
    res = _make_dummy_pipeline_result(
        [p],
        company=comp,
        data_quality_notes=[
            "The issuer/company name is not stated in the document summary, while many pages refer only to the Company.",
            "Table 3 on page 12 has missing column headers.",
        ],
    )

    fixes = sanitize_company_identity_contradictions(res)
    assert len(fixes) > 0

    # Contradiction must be eliminated from Data Quality notes
    for note in res.profile.data_quality_notes:
        assert "issuer/company name is not stated" not in note.casefold()
        assert "unnamed issuer" not in note.casefold()

    # Legitimate non-contradictory data quality notes must be retained
    assert any("missing column headers" in note for note in res.profile.data_quality_notes)


def test_company_overview_vs_data_quality_consistency_unresolved() -> None:
    """When company is unresolved, both Company Overview and Data Quality must reflect unresolved status."""
    p = DocumentPage(page_number=1, text="General Research Report\nIndustry statistical analysis across 2020-2023.")
    comp = CompanyProfile(name="", identity_state="UNRESOLVED")
    res = _make_dummy_pipeline_result(
        [p],
        company=comp,
        data_quality_notes=["Statistical sampling limitations noted on page 4."],
    )

    sanitize_company_identity_contradictions(res)

    # Slide title aligned to Document at a Glance
    overview_slide = next(s for s in res.presentation_plan.slides if s.slide_type == "company_overview")
    assert overview_slide.title == "Document at a Glance"

    # Data quality notes must clearly record that issuer is not stated
    assert any("issuer name is not stated" in note.casefold() for note in res.profile.data_quality_notes)


# =============================================================================
# 7. Fallback behavior: only supported cards rendered
# =============================================================================

def test_fallback_behavior_only_supported_cards_in_company_overview() -> None:
    """When only some company fields exist, no generic filler bullets should be forced into PPTX."""
    from pptx import Presentation
    from adaptive_document_agent.services.pptx_export import _add_company_at_a_glance

    prs = Presentation()
    prs.slide_width = 12192000  # 13.333 inches (16:9 widescreen)
    prs.slide_height = 6858000   # 7.5 inches

    # Partial company profile: only Name, Industry, and Products (no business model, no markets, no listing)
    comp = CompanyProfile(
        name="Precision Instruments Ltd.",
        industry="Industrial Automation",
        reporting_currency="USD",
        products=["Sensor Calibration Kit", "Digital Flow Meter"],
        identity_state="RESOLVED",
        source_pages=[1],
    )
    p = DocumentPage(page_number=1, text="Precision Instruments Ltd.\nIndustrial Automation\nUSD")
    res = _make_dummy_pipeline_result([p], company=comp)
    slide_plan = res.presentation_plan.slides[1]

    _add_company_at_a_glance(prs, res, slide_plan)

    slide = prs.slides[0]
    slide_text = " ".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)

    # Verify real content is present
    assert "Precision Instruments Ltd." in slide_text
    assert "Industrial Automation" in slide_text
    assert "Sensor Calibration Kit" in slide_text

    # Verify generic filler phrases are NOT present
    assert "Core commercial activities and reported operating model." not in slide_text
    assert "Multi-regional market footprint with commercial presence." not in slide_text
    assert "Document-grounded analysis of reported business operations." not in slide_text
    assert "Key product and service lines as disclosed in documentation." not in slide_text


def test_field_validators_comprehensive() -> None:
    """Test all field validators against corruptions, fragments, and valid inputs."""
    # validate_company_name
    assert validate_company_name("d excerpts") == ""
    assert validate_company_name("and ABC Corp") == "ABC Corp"
    assert validate_company_name("Acme Corp (Stock Code: 1234") == "Acme Corp"
    assert validate_company_name("12345") == ""

    # validate_business_model
    assert validate_business_model("d excerpts") == ""
    assert validate_business_model("business model is B2B SaaS enterprise subscription") == "B2B SaaS enterprise subscription"

    # validate_geography
    assert validate_geography("d excerpts") == ""
    assert validate_geography("markets") == ""
    assert validate_geography("North America") == "North America"

    # validate_market_position
    assert validate_market_position("d excerpts") == ""
    assert validate_market_position("General market participant") == ""  # lacks rank/share/leading
    assert validate_market_position("Ranked #1 provider in China") == "Ranked #1 provider in China"
    assert validate_market_position("Holds 32% market share") == "Holds 32% market share"
