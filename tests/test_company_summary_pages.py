"""Summary retrieval, evidence boundaries and two physical introductory pages."""

import pytest
from pptx import Presentation
from pptx.util import Inches

from adaptive_document_agent.models import DocumentPage, DocumentProfile, ParsedDocument, PipelineResult, PresentationPlan, PresentationSlide
from adaptive_document_agent.models.presentation import CompanyProfile, CompanySummaryPage, CompanySummaryItem
from adaptive_document_agent.services.company_summary import summary_excerpts, validate_summary
from adaptive_document_agent.services.pptx_export import _add_company_at_a_glance


def sample(topic="OUR SERVICES", claim="We provide payroll and scheduling services."):
    result = PipelineResult(
        document=ParsedDocument(document_id="summary", sha256="abc", safe_filename="source.pdf", page_count=40,
            pages=[DocumentPage(page_number=2, text="CONTENTS\nSUMMARY ... 5"),
                   DocumentPage(page_number=5, text="SUMMARY\nOVERVIEW\nWe operate a business software platform."),
                   DocumentPage(page_number=18, text=f"SUMMARY\n{topic}\n{claim}"),
                   DocumentPage(page_number=30, text="RISK FACTORS\nOther material")]),
        profile=DocumentProfile(),
    )
    company = CompanyProfile(
        summary_overview=CompanySummaryPage(title="Company at a Glance", source_section="OVERVIEW",
            items=[CompanySummaryItem(label="Company", text="We operate a business software platform.",
                source_quote="We operate a business software platform.", source_pages=[5])]),
        summary_business=CompanySummaryPage(title="Products and services", source_section=topic,
            items=[CompanySummaryItem(label="Offerings", text=claim, source_quote=claim, source_pages=[18])]),
    )
    result.presentation_plan = PresentationPlan(title="Summary", company=company)
    return result


@pytest.mark.parametrize("topic,claim", [
    ("OURCOLLABORATIVE ROBOTPRODUCTS", "We manufacture industrial and educational robots."),
    ("OUR SERVICES", "We provide payroll and scheduling services."),
    ("平台及运营", "我们提供企业采购平台及配送服务。"),
])
def test_summary_retrieves_later_subsections_without_industry_assumptions(topic, claim):
    result = sample(topic, claim)
    excerpts = summary_excerpts(result.document)
    assert {5, 18} <= {p["page"] for p in excerpts}
    assert topic in next(p["text"] for p in excerpts if p["page"] == 18)
    assert not validate_summary(result.presentation_plan.company, result)


def test_summary_rejects_missing_quotes_outside_pages_and_invented_values():
    result = sample()
    item = result.presentation_plan.company.summary_business.items[0]
    item.text = "We offer 99 services."
    assert any("numeric" in error for error in validate_summary(result.presentation_plan.company, result))
    item.source_pages = [99]
    errors = validate_summary(result.presentation_plan.company, result)
    assert any("candidates" in error for error in errors)
    assert any("quote" in error for error in errors)


def test_renders_two_separate_company_pages_with_source_notes():
    result = sample()
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    _add_company_at_a_glance(deck, result, PresentationSlide(id="company", slide_type="company_overview", title="Company at a Glance"))
    assert len(deck.slides) == 2
    text = [" ".join(s.text for s in slide.shapes if s.has_text_frame) for slide in deck.slides]
    assert "software platform" in text[0]
    assert "payroll" not in text[0]
    assert "payroll" in text[1]
    assert "OUR SERVICES" in deck.slides[1].notes_slide.notes_text_frame.text


@pytest.mark.parametrize("heading", ["OUR GROUP", "Investment highlights", "业务概览", "", "The enterprise today"])
def test_different_or_missing_headings_do_not_exclude_supported_introduction(heading):
    result = sample()
    result.document.pages[1].text = f"{heading}\nWe operate a business software platform."
    result.document.pages[2].text = "WHAT WE OFFER\nWe provide payroll and scheduling services."
    assert not validate_summary(result.presentation_plan.company, result)
    assert {5, 18} <= {p["page"] for p in summary_excerpts(result.document, result.profile)}


def test_semantically_discovered_late_intro_and_continuation_are_retrieved():
    result = sample()
    result.document.page_count = 150
    result.document.pages[1].page_number = 100
    result.document.pages[1].text = "The enterprise today\nWe operate a business software platform."
    result.document.pages[2].page_number = 101
    result.document.pages[2].text = "Our offer\nWe provide payroll and scheduling services."
    result.profile.document_summary_pages = [100]
    result.presentation_plan.company.summary_overview.items[0].source_pages = [100]
    result.presentation_plan.company.summary_business.items[0].source_pages = [101]
    assert not validate_summary(result.presentation_plan.company, result)
