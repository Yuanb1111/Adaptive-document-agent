"""Complete reading, source ownership, editorial decisions and physical pages."""
import json
from io import BytesIO
from concurrent.futures import CancelledError
from threading import Barrier, Event, Lock
from time import sleep

import pytest
from pptx import Presentation
from pptx.util import Inches

from adaptive_document_agent.agent.company_introduction import (
    IntroductionDraft, IntroductionPages, ensure_company_introduction, prepare_company_introduction,
)
from adaptive_document_agent.agent.summary_reader import ReadFact, ReadPart, SummaryReadBatch
from adaptive_document_agent.models import DocumentPage, DocumentProfile, ParsedDocument, PipelineResult, PresentationPlan, PresentationSlide
from adaptive_document_agent.models.summary import SummaryFact, SummaryPageRange, SummaryPartDecision, SummarySlideItem, SummarySlidePage
from adaptive_document_agent.services.company_summary import validate_summary
from adaptive_document_agent.services.llm import LLMGateway, LLMSettings, MockLLMClient, PrivacyMode, ProviderName
from adaptive_document_agent.services.llm.base import LLMResponse
from adaptive_document_agent.services.pptx_export import _add_company_at_a_glance
from adaptive_document_agent.services.summary_source import source_blocks, summary_scope
from adaptive_document_agent.services.summary_validation import reading_errors


def source(pages=6):
    first = ('Cedar provides scheduling software for manufacturing teams, with applications '
             'covering shift planning, employee availability and on-site production coordination.')
    second = ('Delivery depends on customer integration and staff training; the document '
              'describes these prerequisites without guaranteeing a completion date or customer adoption.')
    return PipelineResult(document=ParsedDocument(document_id='complete', sha256='source-hash',
        safe_filename='summary.pdf', page_count=pages, pages=[DocumentPage(page_number=i,
            text=f'SUMMARY\nBusiness operations\n{first}\nDelivery constraints\n{second}\n')
            for i in range(1, pages + 1)]), profile=DocumentProfile())


def evidence_payload(messages):
    text = next(m['content'] for m in messages if m['content'].startswith('<UNTRUSTED_DOCUMENT_CONTENT>'))
    return json.loads(text.split('\n', 1)[1].rsplit('\n', 1)[0])


class SummaryClient(MockLLMClient):
    supports_concurrent_requests = True

    def __init__(self, *, on_read=None, edit_read=None, edit_plan=None, omit=False):
        super().__init__()
        self.on_read, self.edit_read, self.edit_plan, self.omit = on_read, edit_read, edit_plan, omit
        self.operations, self.read_inputs, self.plan_inputs = [], [], []
        self.lock = Lock()

    def generate_structured(self, messages, response_model, **kwargs):
        operation = response_model.__name__
        with self.lock:
            self.operations.append(operation)
        payload = evidence_payload(messages)
        if operation == 'IntroductionPages':
            value = IntroductionPages(pages=[], summary_ranges=[SummaryPageRange(
                start_page=payload['page_previews'][0]['page'],
                end_page=payload['page_previews'][-1]['page'], reason='Complete introductory section')])
        elif operation == 'SummaryReadBatch':
            self.read_inputs.append(payload)
            if self.on_read:
                self.on_read(payload)
            parts = []
            for block in payload:
                lines = block['lines']
                cursor = 0
                while cursor < len(lines):
                    index, line = lines[cursor]
                    is_fact = line.startswith(('Cedar provides', 'Delivery depends'))
                    end = cursor + 1
                    if not is_fact:
                        while end < len(lines) and not lines[end][1].startswith(('Cedar provides', 'Delivery depends')):
                            end += 1
                    facts = [ReadFact(label='Operations' if line.startswith('Cedar') else 'Delivery',
                        text=line.strip(), quote_start_line=index, quote_end_line=index)] if is_fact else []
                    parts.append(ReadPart(block_id=block['block_id'], start_line=index, end_line=lines[end-1][0],
                        title='Operations' if line.startswith('Cedar') else 'Delivery' if is_fact else 'Source headings',
                        role='content' if is_fact else 'layout', reading_note='Read the complete assigned source lines.', facts=facts))
                    cursor = end
            value = SummaryReadBatch(parts=parts)
            if self.edit_read:
                self.edit_read(value, payload, messages)
        elif operation in {'IntroductionDraft', 'SummaryEditorialDraft'}:
            self.plan_inputs.append(payload)
            # Equivalent passages on repeated source pages remain read/audited;
            # the editor chooses unique, materially different source claims.
            selected, seen = [], set()
            for part in payload['parts']:
                for fact in part['facts']:
                    key = (fact['text'], part['source_page'])
                    if key not in seen:
                        seen.add(key)
                        selected.append((part['id'], fact))
            pages, used = [], set()
            if not self.omit:
                for start in range(0, len(selected), 4):
                    group = selected[start:start+4]
                    items = []
                    for pid, fact in group:
                        items.append(SummarySlideItem(**fact, part_ids=[pid]))
                        used.add(pid)
                    titles = ['Operations and delivery', 'Customer requirements', 'Service coordination']
                    pages.append(SummarySlidePage(id=f'intro-{start}', title=titles[(start // 4) % 3],
                        source_section='Business operations; Delivery constraints', items=items))
            value = IntroductionDraft(summary_pages=pages, summary_decisions=[SummaryPartDecision(
                part_id=part['id'], decision='include' if part['id'] in used else 'omit',
                reason='Supports the introductory operations narrative.' if part['id'] in used else
                       'Running headings add no independent substantive finding.' if part['role'] == 'layout' else
                       'Reserved for later source-specific analysis.') for part in payload['parts']])
            if self.edit_plan:
                self.edit_plan(value, payload, messages)
            if operation == 'SummaryEditorialDraft':
                facts = {f['id']: (part['id'], f) for part in payload['parts'] for f in part['facts']}
                value = response_model(summary_decisions=value.summary_decisions, summary_pages=[{
                    **page.model_dump(exclude={'items'}), 'items': [{'label': item.label, 'text': item.text,
                        'fact_id': next((fid for fid, (pid, fact) in facts.items()
                            if item.part_ids == [pid] and item.source_quote == fact['source_quote']), 'unknown')}
                        for item in page.items]} for page in value.summary_pages])
        else:
            raise AssertionError(operation)
        return value, LLMResponse(text=value.model_dump_json())


def gateway(client=None, *, local=False, workers=2):
    return LLMGateway(client or SummaryClient(), LLMSettings(provider=ProviderName.OLLAMA if local else ProviderName.DEEPSEEK,
        model='fixture', privacy_mode=PrivacyMode.LOCAL_ONLY if local else PrivacyMode.CLOUD,
        discovery_workers=workers))


def distinct_source(pages=6):
    result = source(pages)
    # Distinct source assertions avoid redundant editorial findings. No invented
    # numeric page labels enter the narrative or its evidence.
    subjects = ['manufacturing', 'warehouse', 'retail', 'hospital', 'transport', 'hospitality']
    for page, subject in zip(result.document.pages, subjects * (pages // 6 + 1)):
        page.text = page.text.replace('manufacturing', subject).replace('customer integration', f'{subject} customer integration')
    return result


def generate(result=None, client=None, **kwargs):
    result = result or distinct_source()
    plan = PresentationPlan(title='Source deck', slides=[PresentationSlide(id='company',
        slide_type='company_overview', title='Company overview')])
    ensure_company_introduction(gateway(client, **kwargs), result, plan)
    result.presentation_plan = plan
    return result


def test_all_parts_read_then_merged_into_dynamic_pages_with_omission_reasons():
    client = SummaryClient()
    result = generate(client=client)
    review = result.presentation_plan.company.summary_review
    assert review.status == 'complete' and review.source_pages == list(range(1, 7))
    assert len([p for p in review.parts if p.role == 'content']) == 12
    assert len(result.presentation_plan.company.summary_pages) == 3
    assert {p.id for p in review.parts} == {d.part_id for d in review.decisions}
    assert all(d.reason for d in review.decisions)
    assert any(d.decision == 'omit' for d in review.decisions)
    assert not validate_summary(result.presentation_plan.company, result)
    assert any(w.code == 'company_introduction_summary_audit' for w in result.validation_warnings)
    from adaptive_document_agent.services.llm.reasoning_policy import request_reasoning_policy
    settings = LLMSettings(provider=ProviderName.DEEPSEEK, model='deepseek-v4-pro')
    for operation in ('SummaryReadBatch', 'IntroductionDraft'):
        assert request_reasoning_policy(settings, stage='presentation', operation=operation,
                                        model=settings.model)['intent'] == 'preserve'


def test_full_scope_includes_more_than_twenty_header_pages_and_long_page_tail():
    result = source(25)
    result.document.pages[-1].text = 'SUMMARY\n' + '\n' * 7100 + result.document.pages[-1].text[8:]
    pages = summary_scope(result, [])
    blocks = source_blocks(result, pages)
    assert pages == list(range(1, 26))
    assert ''.join(b.text for b in blocks if b.page == 25) == result.document.pages[-1].text
    client = SummaryClient(omit=True)
    generate(result, client)
    from adaptive_document_agent.models.summary import SummaryReview
    review = SummaryReview.model_validate_json(next(w.message for w in result.validation_warnings
        if w.code == 'company_introduction_summary_audit'))
    assert not reading_errors(review, result)
    assert any(p.source_page == 25 and p.facts for p in review.parts)
    assert review.status == 'complete'


@pytest.mark.parametrize('mutation,fragment', [
    ('drop_part', 'skip'), ('overlap', 'overlap'), ('bad_quote', 'outside its assigned part'),
    ('reversed_span', 'outside its assigned part'),
])
def test_reader_rejects_gaps_overlaps_and_unbound_claims_and_retains_failed_audit(mutation, fragment):
    def edit(response, payload, messages):
        part = next(p for p in response.parts if p.facts)
        if mutation == 'drop_part':
            response.parts.remove(part)
        elif mutation == 'overlap':
            response.parts.append(part.model_copy(deep=True))
        elif mutation == 'bad_quote':
            part.facts[0].quote_end_line = 999
        elif mutation == 'reversed_span':
            part.facts[0].quote_start_line = part.end_line + 1
        elif mutation == 'wrong_unit':
            part.facts[0].text = 'Cedar revenue was RMB 30 million.'
        else:
            part.facts[0].text = 'Cedar has 999 customers.'
    result = distinct_source(1)
    client = SummaryClient(edit_read=edit)
    with pytest.raises(ValueError):
        generate(result, client)
    audit = json.loads(next(w.message for w in result.validation_warnings if w.code == 'company_introduction_summary_audit'))
    assert audit['status'] == 'incomplete'
    assert len(audit['read_audits'][0]['attempts']) == 2
    assert any(fragment in e for a in audit['read_audits'][0]['attempts'] for e in a['errors'])
    assert 'IntroductionDraft' not in client.operations


def test_reader_retries_one_schema_failure_with_original_source_and_keeps_diagnostics():
    from adaptive_document_agent.services.llm.exceptions import LLMStructuredOutputError
    class InvalidFirstRead(SummaryClient):
        failed = False
        def generate_structured(self, messages, response_model, **kwargs):
            if response_model is SummaryReadBatch and not self.failed:
                self.failed = True
                response = LLMResponse(text='{"parts": [{"facts": "invalid"}]}')
                raise LLMStructuredOutputError('schema validation failed', response=response) from ValueError('facts must be a list')
            return super().generate_structured(messages, response_model, **kwargs)
    result = distinct_source(1)
    generate(result, InvalidFirstRead())
    review = result.presentation_plan.company.summary_review
    assert review.status == 'complete'
    attempts = review.read_audits[0]['attempts']
    assert len(attempts) == 2 and attempts[0]['validation_detail'] == 'facts must be a list'
    assert 'invalid' in attempts[0]['raw_response']
    assert all(f.source_pages == [part.source_page] and f.source_quote in part.source_text
               for part in review.parts for f in part.facts)


def test_reading_accepts_exact_page_annotations_without_lending_numbers_to_facts():
    def edit(response, payload, messages):
        for part in response.parts:
            if part.role == 'layout':
                part.title = f"Page {payload[0]['page']} source headings"
    result = generate(distinct_source(1), SummaryClient(edit_read=edit))
    assert result.presentation_plan.company.summary_review.status == 'complete'
    review = result.presentation_plan.company.summary_review.model_copy(deep=True)
    fact = next(f for part in review.parts for f in part.facts)
    fact.text += ' It has 1 product.'
    assert any('unsupported numeric' in error for error in reading_errors(review, result))


def test_reading_preserves_a_qualified_fact_between_220_and_280_characters():
    from adaptive_document_agent.agent.summary_reader import ReadFact
    text = ('Services depend on successful customer integration and staff training; '
            'the source describes the delivery requirements without guaranteeing adoption, '
            'an implementation date, a financial return or continuing purchases by existing customers.')
    assert 220 < len(text) <= 280
    fact = ReadFact(label='Delivery conditions', text=text, quote_start_line=1, quote_end_line=4)
    assert fact.text == text


def test_long_qualified_reading_fact_keeps_semantic_repair_available():
    from adaptive_document_agent.agent.summary_reader import _read_batch
    from adaptive_document_agent.models.summary import SummaryReview
    text = ('The organisation offers 3 services, subject to customer integration, staff training '
            'and continuing access to licensed software; delivery conditions depend on the '
            'customer environment and the source does not guarantee adoption, a completion '
            'date, a financial return, future purchases or any improvement in operating results.')
    assert 280 < len(text) <= 600
    result = source(1)
    result.document.pages[0].text = text.replace('3 services', 'several services') + '\nIt offers 3 services.\n'
    blocks = source_blocks(result, [1])
    calls = []
    class Reader:
        def generate_structured(self, messages, model, **kwargs):
            calls.append(messages)
            return SummaryReadBatch(parts=[ReadPart(block_id=blocks[0].id,
                start_line=1, end_line=2, title='Delivery', role='content',
                reading_note='Delivery depends on customer prerequisites.', facts=[ReadFact(
                    label='Delivery', text=text, quote_start_line=1,
                    quote_end_line=1 if len(calls) == 1 else 2)])])
    outcome = _read_batch(Reader(), result, blocks, Event())
    assert outcome.error is None and len(calls) == 2
    assert outcome.audit['attempts'][0]['errors'] == ["Delivery: unsupported numeric claims ['3']."]
    assert 'quote line endpoints' in calls[1][-1]['content']
    assert outcome.parts[0].facts[0].text == text
    assert outcome.audit['attempts'][1]['errors'] == []
    with pytest.raises(ValueError):
        SummarySlideItem(text=text, part_ids=['part'], source_pages=[1], source_quote=text)


def test_failed_introduction_does_not_render_an_evidence_limitation_placeholder():
    from adaptive_document_agent.models import ValidationIssue
    from tests.test_presentation_brief import blank_deck
    result = source(1)
    result.presentation_plan = PresentationPlan(title='Findings')
    result.validation_warnings.append(ValidationIssue(code='company_introduction_unavailable',
        stage='presentation', severity='warning', message='Source verification failed'))
    deck = blank_deck()
    deck.slides.add_slide(deck.slide_layouts[0])
    _add_company_at_a_glance(deck, result, PresentationSlide(id='intro', slide_type='company_overview', title='Overview'))
    assert len(deck.slides) == 1
    assert 'source verification failed' in deck.slides[0].notes_slide.notes_text_frame.text


@pytest.mark.parametrize('mutation', ['sparse', 'missing_decision', 'false_include', 'unknown_part', 'new_quote'])
def test_editor_rejects_bad_draft_and_recovers_complete_validated_reading(mutation):
    def edit(draft, payload, messages):
        if mutation == 'sparse':
            for item in draft.summary_pages[0].items:
                item.text = 'Cedar provides software.'
        elif mutation == 'missing_decision':
            draft.summary_decisions.pop()
        elif mutation == 'false_include':
            next(d for d in draft.summary_decisions if d.decision == 'omit').decision = 'include'
        elif mutation == 'unknown_part':
            draft.summary_pages[0].items[0].part_ids = ['invented-part']
        else:
            draft.summary_pages[0].items[0].source_quote = 'An invented source quote for this item.'
    result = distinct_source(1)
    client = SummaryClient(edit_plan=edit)
    generate(result, client)
    audit = json.loads(next(w.message for w in result.validation_warnings if w.code == 'company_introduction_summary_audit'))
    assert len(audit['plan_audits']) == 3 and audit['status'] == 'complete'
    assert audit['plan_audits'][-1]['reading_layout'] is True
    company = result.presentation_plan.company
    assert {p.id for p in company.summary_review.parts if p.role=='content'} == {
        pid for page in company.summary_pages for item in page.items for pid in item.part_ids}
    assert all(item.reading_fact_copy for page in company.summary_pages for item in page.items)


def test_sparse_editor_repair_cannot_omit_substantive_parts():
    def edit(draft, payload, messages):
        if len(messages) == 2:
            for item in draft.summary_pages[0].items:
                item.text = 'Cedar provides software.'
        else:
            draft.summary_pages = []
            for decision in draft.summary_decisions:
                decision.decision = 'omit'
                decision.reason = 'Retained source facts do not merit a separate introductory slide.'
    result = generate(distinct_source(1), SummaryClient(edit_plan=edit))
    assert result.presentation_plan.company.summary_review.status == 'complete'
    assert all(d.decision=='include' for d in result.presentation_plan.company.summary_review.decisions
               if d.part_id in {p.id for p in result.presentation_plan.company.summary_review.parts if p.role=='content'})


def test_parallel_batches_overlap_and_keep_source_order(monkeypatch):
    from adaptive_document_agent.services import summary_source
    monkeypatch.setattr(summary_source, 'BATCH_CHARACTERS', 500)
    barrier = Barrier(2)
    client = SummaryClient(on_read=lambda payload: barrier.wait(5))
    result = generate(client=client, workers=3)
    review = result.presentation_plan.company.summary_review
    assert len(review.read_audits) == 6
    assert [p.source_page for p in review.parts] == sorted(p.source_page for p in review.parts)


def test_local_only_reading_is_serial_and_does_not_enable_cloud_light_policy(monkeypatch):
    from adaptive_document_agent.services import summary_source
    monkeypatch.setattr(summary_source, 'BATCH_CHARACTERS', 500)
    active, peak = 0, 0
    lock = Lock()
    def read(payload):
        nonlocal active, peak
        with lock:
            active += 1
            peak = max(peak, active)
        sleep(.005)
        with lock:
            active -= 1
    client = SummaryClient(on_read=read)
    result = generate(client=client, local=True, workers=4)
    assert peak == 1 and len(result.presentation_plan.company.summary_review.read_audits) == 6


def test_cancellation_prevents_later_batches_and_editor(monkeypatch):
    from adaptive_document_agent.agent.summary_reader import read_summary
    from adaptive_document_agent.models.summary import SummaryReview
    from adaptive_document_agent.services import summary_source
    monkeypatch.setattr(summary_source, 'BATCH_CHARACTERS', 500)
    result = distinct_source()
    signal = Event()
    client = SummaryClient(on_read=lambda payload: signal.set())
    review = SummaryReview(document_id=result.document.document_id, document_sha256=result.document.sha256,
        source_pages=list(range(1, 7)), source_blocks=source_blocks(result, list(range(1, 7))))
    with pytest.raises(CancelledError):
        read_summary(gateway(client, local=True), result, review, cancelled=signal)
    assert client.operations == ['SummaryReadBatch']


def test_scope_model_can_select_renamed_intro_and_rejects_unreadable_pages():
    result = distinct_source(2)
    for page in result.document.pages:
        page.text = page.text.replace('SUMMARY', 'THE ENTERPRISE TODAY')
    assert summary_scope(result, []) == []
    span = SummaryPageRange(start_page=1, end_page=2, reason='Introductory overview under a different name')
    assert summary_scope(result, [span]) == [1, 2]
    result.document.pages[-1].text = ''
    with pytest.raises(ValueError, match='OCR'):
        summary_scope(result, [span])


def test_source_tampering_and_removed_entire_page_are_detected():
    result = generate()
    review = result.presentation_plan.company.summary_review.model_copy(deep=True)
    review.source_pages.pop()
    review.source_blocks = [b for b in review.source_blocks if b.page != 6]
    review.parts = [p for p in review.parts if p.source_page != 6]
    assert any('complete source scope' in e for e in reading_errors(review, result))
    review = result.presentation_plan.company.summary_review.model_copy(deep=True)
    review.parts[1].source_text = 'Changed raw text'
    assert any('changed source text' in e for e in reading_errors(review, result))


def test_prepared_introduction_validates_latest_source_before_attachment():
    result = distinct_source(1)
    with prepare_company_introduction(gateway(SummaryClient(), local=True), result) as prepared:
        plan = PresentationPlan(title='Source deck')
        prepared(plan, result)
        result.document.pages[0].text += '\nChanged source'
        with pytest.raises(ValueError, match='no longer matches'):
            prepared(PresentationPlan(title='Another deck'), result)


def test_physical_pages_retain_every_selected_fact_citations_and_full_decisions():
    result = generate()
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    _add_company_at_a_glance(deck, result, result.presentation_plan.slides[0])
    assert len(deck.slides) == 3
    for slide, page in zip(deck.slides, result.presentation_plan.company.summary_pages):
        text = ' '.join(shape.text for shape in slide.shapes if shape.has_text_frame)
        assert all(item.text in text for item in page.items)
        assert slide._ada_section_label == page.title
        assert sum(len(shape.text) for shape in slide.shapes if shape.name == 'brief:body') >= 180
        assert 'Source: Document disclosures' in text
    assert 'summary-reading-v2-complete-content' in deck.slides[0].notes_slide.notes_text_frame.text
    assert 'Reserved for later' not in ' '.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)


@pytest.mark.parametrize('claim', ['Revenue was USD 30 million in 2024.',
                                   'Revenue was RMB 30 billion in 2024.',
                                   'Revenue was RMB 30 million in 2025.'])
def test_reading_fact_checks_currency_scale_and_period(claim):
    from adaptive_document_agent.services.summary_validation import fact_errors
    result = distinct_source(1)
    quote = 'Revenue was RMB 30 million in 2024.'
    result.document.pages[0].text += quote
    fact = SummaryFact(label='Revenue', text=claim, source_quote=quote, source_pages=[1])
    assert fact_errors(fact, result, quote, 1)


def test_budget_excess_fails_without_sampling_or_claiming_complete(monkeypatch):
    from adaptive_document_agent.services import summary_source
    monkeypatch.setattr(summary_source, 'MAX_SUMMARY_CHARACTERS', 20)
    result = distinct_source(1)
    client = SummaryClient()
    with pytest.raises(ValueError, match='no source was sampled'):
        generate(result, client)
    assert client.operations == ['IntroductionPages']
    audit = json.loads(next(w.message for w in result.validation_warnings if w.code == 'company_introduction_summary_audit'))
    assert audit['status'] == 'incomplete' and audit['source_pages'] == [1]


def complete_deck(result):
    company = result.presentation_plan.company
    result.presentation_plan = PresentationPlan(title='Source deck', company=company, slides=[
        PresentationSlide(id=kind, slide_type=kind, title=title, section_title=title)
        for kind, title in [('cover', 'Source deck'), ('company_overview', 'Company overview'),
                            ('executive_summary', 'Executive summary'), ('data_quality', 'Scope'),
                            ('appendix', 'Data index')]])
    return result


@pytest.mark.parametrize('omit', [False, True])
def test_full_export_agenda_matches_dynamic_introduction_and_all_omissions(omit):
    from adaptive_document_agent.services.pptx_export import build_presentation
    result = complete_deck(generate(client=SummaryClient(omit=omit)))
    payload = build_presentation(result)
    deck = Presentation(BytesIO(payload))
    contents = next(slide for slide in deck.slides if any(shape.name.startswith('contents:entry:') for shape in slide.shapes))
    text = ' '.join(shape.text for shape in contents.shapes if shape.has_text_frame)
    for page in result.presentation_plan.company.summary_pages:
        assert page.title in text
        assert any(page.items[0].text in ' '.join(s.text for s in slide.shapes if s.has_text_frame) for slide in deck.slides)
    assert 'Company overview' not in text
    if omit:
        from adaptive_document_agent.services.summary_export import verify_summary_export
        assert verify_summary_export(deck,result)


def test_layout_repair_does_not_replace_complete_agent_reading_with_heuristic_fields():
    from adaptive_document_agent.validation.layout_qa import validate_presentation_layout
    result = generate()
    before = result.presentation_plan.company.model_dump()
    issues = validate_presentation_layout(result)
    assert result.presentation_plan.company.model_dump() == before
    assert not any(issue.code == 'company_overview_structured_repaired' for issue in issues)


def test_reader_repairs_unknown_block_once_before_editing_and_keeps_both_attempts():
    def edit(response, payload, messages):
        if len(messages) == 2:
            response.parts[0].block_id = 'unknown-source-block'
    client = SummaryClient(edit_read=edit)
    result = generate(distinct_source(1), client)
    review = result.presentation_plan.company.summary_review
    assert client.operations == ['IntroductionPages', 'SummaryReadBatch', 'SummaryReadBatch', 'SummaryEditorialDraft']
    attempts = review.read_audits[0]['attempts']
    assert 'unknown source block' in ' '.join(attempts[0]['errors'])
    assert attempts[1]['errors'] == [] and review.status == 'complete'


def test_empty_review_cannot_claim_complete_reading():
    from adaptive_document_agent.models.summary import SummaryReview
    result = distinct_source(1)
    result.document.pages[0].text = result.document.pages[0].text.replace('SUMMARY', 'Generic document')
    review = SummaryReview(document_id=result.document.document_id, document_sha256=result.document.sha256,
        source_pages=[], status='complete')
    assert any('requires supplied' in error for error in reading_errors(review, result))
