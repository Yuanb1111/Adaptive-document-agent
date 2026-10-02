"""Selected analytical composition survives claim repair and optional enrichment."""

import copy

import pytest

from adaptive_document_agent.models import (
    ChartPlan, DocumentPage, ExtractedTable, PresentationSlide, PresentationTheme,
    PresentationTopic, PresentationTopicSelection, PresentationVisualBlock,
)
from adaptive_document_agent.services.presentation_claim_evidence import prepare_presentation_claims, visible_observation_ids
from adaptive_document_agent.services.presentation_share_claims import share_claims
from tests.test_p0_composition import observation, result_for


def composition_result(*, labels=("Long-range sensors", "Short-range sensors", "Other"),
                       measure="revenue", periods=("FY2025", "FY2026")):
    observations = []
    shares = ((25, 55), (65, 30), (10, 15))
    for label, values in zip(labels, shares):
        for period, value in zip(periods, values):
            item = observation(f"share-{label}-{period}", f"{label}: % of Total", value, period, unit="percent")
            item.table_id = "mix-source"
            item.dimensions = {"table_context": measure, "column_role": "percentage"}
            item.evidence[0].table_id = "mix-source"
            item.evidence[0].row_label = label
            item.evidence[0].column_label = "% of Total"
            item.validation_status = "valid"
            observations.append(item)
    chart = ChartPlan(id="chosen-mix", title="Reported category mix", chart_type="stacked_percent",
                      series_dimension="source_row", composition_table_id="mix-source",
                      question="How does the mix compare?", observation_ids=[o.id for o in observations],
                      source_pages=[3])
    amounts = []
    for label, values in zip(labels[:2], ((100, 140), (150, 90))):
        for period, value in zip(periods, values):
            item = observation(f"units-{label}-{period}", "Units shipped", value, period, unit="count")
            item.dimensions = {"category": label}
            item.category_dimensions = {"category": label}
            item.evidence[0].row_label = label
            item.evidence[0].column_label = "Units shipped"
            item.evidence[0].table_id = item.table_id = "units-source"
            item.validation_status = "valid"
            amounts.append(item)
    unit_charts = [ChartPlan(id=f"units-{i}", title=label, chart_type="line",
                            question="How did shipments change?", source_pages=[3],
                            observation_ids=[o.id for o in amounts if o.dimensions["category"] == label])
                   for i, label in enumerate(labels[:2])]
    question = "How did category mix and shipments change?"
    title = f"{labels[0]} gained {measure} share as {labels[1]} share fell."
    singular = labels[0][:-3] + "y" if labels[0].endswith("ies") else labels[0][:-1] if labels[0].endswith("s") else labels[0]
    takeaway = f"{singular} {measure} share rose while {labels[1]} share fell."
    slide = PresentationSlide(id="mix-analysis", slide_type="analysis", title="Reported measures",
        section_title=title, message=question, theme_id="mix-topic", layout="three_up",
        visual_blocks=[PresentationVisualBlock(role="matrix", matrix_dimension="category",
                                              observation_ids=[o.id for o in amounts])],
        analytical_question="How do the reported values vary across the cited periods?",
        selection_reason="The retained records support these reported values; the broader comparison was omitted.",
        source_pages=[3])
    result = result_for([*observations, *amounts], [chart, *unit_charts], slide)
    result.document.pages = [DocumentPage(page_number=3, text="Reported mix and units", tables=[ExtractedTable(
        table_id="mix-source", page=3, raw_header_lines=[
            f"The following table sets forth a breakdown of our {measure} by category for the years indicated."],
    )])]
    result.presentation_plan.themes = [PresentationTheme(id="mix-topic", title=title, question=question,
        rationale="Selected mix and operating context.", chart_ids=[c.id for c in result.charts],
        observation_ids=[o.id for o in result.observations], source_pages=[3],
        caveats=["Shipment series cover selected categories only; composition covers all categories."])]
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(
        id="mix-topic", title=title, question=question, rationale="Selected mix and operating context.",
        takeaway=takeaway, series_ids=["selected-composition", "selected-units-a", "selected-units-b"],
        caveats=result.presentation_plan.themes[0].caveats)])
    result.presentation_plan.coverage_notes = ["Individual category share redundant with selected composition."]
    return result


@pytest.mark.parametrize("labels,measure", [
    (("Long-range sensors", "Short-range sensors", "Other"), "revenue"),
    (("Enterprise accounts", "Consumer accounts", "Partners"), "customers"),
    (("Coastal facilities", "Inland facilities", "Remote facilities"), "production"),
])
def test_supported_selected_composition_returns_to_analysis_without_mutating_facts(labels, measure):
    result = composition_result(labels=labels, measure=measure)
    selected = next(s for s in result.presentation_plan.slides if s.id == "mix-analysis")
    selected.chart_ids = ["units-0", "units-1"]
    selected.visual_blocks = []
    before = copy.deepcopy(result.observations)
    prepare_presentation_claims(result)
    charts = {c.id: c for c in result.charts}
    visible = {oid for slide in result.presentation_plan.slides if slide.slide_type == "analysis"
               for oid in visible_observation_ids(slide, charts)}
    assert set(charts["chosen-mix"].observation_ids) <= visible
    selected = next(s for s in result.presentation_plan.slides if s.id == "mix-analysis")
    assert selected.title == result.presentation_topics.topics[0].takeaway
    assert result.observations == before
    snapshot = result.model_dump()
    prepare_presentation_claims(result)
    assert result.model_dump() == snapshot


def test_opposed_share_predicates_split_at_as_and_keep_explicit_denominator():
    result = composition_result()
    claims = share_claims(result.observations, result.presentation_topics.topics[0].title)
    assert len(claims) == 2
    assert claims[0].denominator == "revenue"
    claims = share_claims(result.observations, result.presentation_topics.topics[0].takeaway)
    assert claims[0].denominator == "revenue"


@pytest.mark.parametrize("mutation", ["missing", "invalid", "wrong_table", "conflict", "wrong_basis"])
def test_unproved_selected_composition_is_withheld_and_coverage_is_current(mutation):
    result = composition_result()
    item = result.observations[0]
    if mutation == "missing":
        result.observations.remove(item)
    elif mutation == "invalid":
        item.validation_status = "invalid"
    elif mutation == "wrong_table":
        item.evidence[0].table_id = "different-source"
    elif mutation == "conflict":
        item.value = 90
    else:
        item.period_basis = "6M"
    raw = copy.deepcopy(result.observations)
    prepare_presentation_claims(result)
    charts = {c.id: c for c in result.charts}
    visible = {oid for s in result.presentation_plan.slides if s.slide_type == "analysis"
               for oid in visible_observation_ids(s, charts)}
    assert not set(charts["chosen-mix"].observation_ids) <= visible
    coverage = " ".join(result.presentation_plan.coverage_notes + result.presentation_plan.themes[0].caveats)
    assert "covers all" not in coverage and "redundant with selected composition" not in coverage
    assert "not displayed" in coverage
    assert result.observations == raw


def test_valid_composition_kept_when_directional_claim_is_false():
    result = composition_result()
    topic = result.presentation_topics.topics[0]
    topic.takeaway = topic.takeaway.replace("rose", "fell")
    prepare_presentation_claims(result)
    slide = next(s for s in result.presentation_plan.slides if s.id == "mix-analysis")
    assert any("chosen-mix" in s.chart_ids for s in result.presentation_plan.slides)
    assert slide.title != topic.takeaway


def test_optional_subset_visual_does_not_discard_selected_composition():
    from adaptive_document_agent.agent.presentation_visual_enrichment import enrich_selected_plan, VisualSelection, VisualSuggestion

    result = composition_result()
    slide = next(s for s in result.presentation_plan.slides if s.id == "mix-analysis")
    slide.title = "Category mix and shipments"
    slide.chart_ids = [c.id for c in result.charts]
    slide.visual_blocks = []
    result.presentation_plan.slides.insert(-2, PresentationSlide(
        id="other-analysis", slide_type="analysis", title="Shipments", message="Reported shipments",
        section_title="Shipments", theme_id="mix-topic", chart_ids=["units-0"], source_pages=[3],
        analytical_question="How did shipments change?", selection_reason="Reported operating measure."))
    before = result.presentation_plan.model_dump()
    class Gateway:
        def generate_structured(self, *args, **kwargs):
            return VisualSelection(suggestions=[VisualSuggestion(
                topic_id="mix-topic", kind="matrix", matrix_dimension="category",
                observation_ids=[o.id for o in result.observations if o.unit == "count"])])
    enriched = enrich_selected_plan(Gateway(), result, result.presentation_plan)
    assert enriched.model_dump() == before


def test_wrong_named_denominator_does_not_hide_valid_composition_or_revive_claim():
    result = composition_result()
    topic = result.presentation_topics.topics[0]
    topic.takeaway = topic.takeaway.replace("revenue share", "unit share")
    prepare_presentation_claims(result)
    slide = next(s for s in result.presentation_plan.slides if s.id == "mix-analysis")
    assert any("chosen-mix" in s.chart_ids for s in result.presentation_plan.slides)
    assert slide.title != topic.takeaway


def test_singular_alias_collision_is_not_a_unique_share_subject():
    result = composition_result(labels=("Premium plans", "Premium plan", "Other"))
    assert share_claims(result.observations, "Premium plan revenue share rose.") is None


def test_fresh_compilation_retains_mix_when_takeaway_needs_source_scope_recovery():
    from adaptive_document_agent.agent.chart_planner import ChartPlanner
    from adaptive_document_agent.agent.presentation_topic_selector import series_directory
    from adaptive_document_agent.agent.topic_plan_compiler import compile_topic_plan
    from adaptive_document_agent.document_model import DocumentIndex

    result = composition_result()
    directory, lookup = series_directory(result)
    chosen = [row["id"] for row in directory if row["visual_kind"] == "stacked_percent"]
    chosen += [row["id"] for row in directory if row["metric"] == "Units shipped" and row["visual_kind"] == "series"]
    result.presentation_topics.topics[0].series_ids = chosen
    result.presentation_plan = None
    result.charts = ChartPlanner().plan([], [], DocumentIndex(result.observations),
        requested_series=[lookup[sid] for sid in chosen], only_requested=True)
    plan = compile_topic_plan(result)
    charts = {c.id: c for c in result.charts}
    visible = {oid for slide in plan.slides if slide.slide_type == "analysis"
               for oid in visible_observation_ids(slide, charts)}
    assert {o.id for o in result.observations} <= visible
    assert any("Displayed composition" in note for note in plan.themes[0].caveats)


def test_restored_composition_renders_all_categories_and_exact_native_values():
    from pptx import Presentation
    from pptx.util import Inches
    from adaptive_document_agent.document_model import DocumentIndex
    from adaptive_document_agent.services.slide_compositor import render_composed_slide

    result = composition_result()
    prepare_presentation_claims(result)
    slide = next(s for s in result.presentation_plan.slides if "chosen-mix" in s.chart_ids)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    chart = next(c for c in result.charts if c.id == "chosen-mix")
    render_composed_slide(presentation, slide, [chart], result, DocumentIndex(result.observations))
    native = next(shape.chart for page in presentation.slides for shape in page.shapes if shape.has_chart)
    expected = {label: [o.value / 100 for o in result.observations
                        if o.unit == "percent" and o.evidence[0].row_label == label]
                for label in ("Long-range sensors", "Short-range sensors", "Other")}
    assert {series.name: list(series.values) for series in native.series} == expected
    assert [category.label for category in native.plots[0].categories] == ["FY2025", "FY2026"]


def test_scope_recovery_realigns_summary_duplicates_without_changing_source_facts():
    from adaptive_document_agent.agent.presentation_topic_selector import series_directory
    from adaptive_document_agent.agent.topic_plan_compiler import compile_topic_plan
    from adaptive_document_agent.validation.claim_validator import ClaimValidator
    from tests.test_topic_scope_recovery_regressions import _synthetic_result

    result = _synthetic_result("metric_alias")
    volume = [o for o in result.observations if o.metric_original == "Units"]
    for item in volume:
        item.dimensions["table_context"] = "Operating review"
    duplicates = [o.model_copy(deep=True) for o in volume[:2]]
    for item in duplicates:
        item.id = "duplicate-" + item.id
        item.dimensions["table_context"] = "Introductory summary"
    result.observations.extend(duplicates)
    _, lookup = series_directory(result)
    result.presentation_topics.topics[1].series_ids = [sid for sid, members in lookup.items()
        if all(o.metric_original == "Units" for o in members)]
    before = copy.deepcopy(result.observations)
    plan = compile_topic_plan(result)
    assert not ClaimValidator().validate_plan(plan, result.observations, result.charts)
    summary = next(s for s in plan.slides if s.slide_type == "executive_summary")
    assert set(summary.bullet_observation_ids[1]) == {o.id for o in volume}
    assert result.observations == before
    assert summary.bullets[1] == result.presentation_topics.topics[1].takeaway


def test_coverage_reconciliation_preserves_independent_composition_qualifications():
    result = composition_result()
    caveat = "Composition includes source estimates; category definitions changed during the reported period."
    result.presentation_plan.themes[0].caveats.append(caveat)
    prepare_presentation_claims(result)
    qualifications = " ".join(result.presentation_plan.themes[0].caveats)
    assert "Composition includes source estimates" in qualifications
    assert "category definitions changed" in qualifications


def test_cached_selected_mix_rechecks_strict_claim_scope_before_export(local_render_stub):
    from io import BytesIO
    from pptx import Presentation
    from adaptive_document_agent.agent.presentation_topic_selector import series_directory
    from adaptive_document_agent.services.export import export_pptx
    from adaptive_document_agent.validation.claim_validator import ClaimValidator

    result = composition_result()
    result.presentation_plan.planning_origin = "topic_compilation"
    selected = next(s for s in result.presentation_plan.slides if s.id == "mix-analysis")
    selected.chart_ids = ["units-0", "units-1"]
    selected.visual_blocks = []
    directory, _ = series_directory(result)
    result.presentation_topics.topics[0].series_ids = [row["id"] for row in directory
        if row["visual_kind"] == "stacked_percent"
        or row["metric"] == "Units shipped" and row["visual_kind"] == "series"]
    before = copy.deepcopy(result.observations)
    deck = Presentation(BytesIO(export_pptx(result)))
    selected = next(s for s in result.presentation_plan.slides if s.id == "mix-analysis")
    assert selected.title == result.presentation_topics.topics[0].question
    assert not ClaimValidator().validate_plan(result.presentation_plan, result.observations, result.charts)
    assert any(shape.has_chart for page in deck.slides for shape in page.shapes)
    assert result.observations == before
    snapshot = result.model_dump()
    prepare_presentation_claims(result)
    assert result.model_dump() == snapshot


def test_partial_stacked_amount_view_does_not_claim_full_source_population():
    result = composition_result()
    chart = next(c for c in result.charts if c.id == "chosen-mix")
    chart.chart_type = "stacked_bar"
    chart.observation_ids = [oid for oid in chart.observation_ids if "Other" not in oid]
    prepare_presentation_claims(result)
    coverage = " ".join(result.presentation_plan.themes[0].caveats)
    assert "complete validated source-category" not in coverage
    assert "validated selected-category matrix" in coverage


def test_restored_mix_preserves_category_matrix_and_all_visible_row_identities():
    from pptx import Presentation
    from pptx.util import Inches
    from adaptive_document_agent.services.presentation_matrix import render_matrix

    result = composition_result()
    original = next(s for s in result.presentation_plan.slides if s.id == "mix-analysis")
    original_block = original.visual_blocks[0].model_dump()
    prepare_presentation_claims(result)
    selected = next(s for s in result.presentation_plan.slides if s.id == "mix-analysis")
    assert selected.visual_blocks[0].model_dump() == original_block
    assert not selected.chart_ids
    composition_page = next(s for s in result.presentation_plan.slides if "chosen-mix" in s.chart_ids)
    assert composition_page.id != selected.id
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    render_matrix(presentation, selected, selected.visual_blocks[0], {o.id: o for o in result.observations})
    tables = [shape.table for page in presentation.slides for shape in page.shapes if shape.has_table]
    assert len(tables) == 1
    table = tables[0]
    visible = " ".join(cell.text for row in table.rows for cell in row.cells)
    for label in ("Long-range sensors", "Short-range sensors", "FY2025", "FY2026"):
        assert label in visible
    assert len(table.rows) == 3 and len(table.columns) == 3
    snapshot = result.model_dump()
    prepare_presentation_claims(result)
    assert result.model_dump() == snapshot
