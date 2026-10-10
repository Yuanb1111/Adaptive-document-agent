"""Regressions for scoped intent and complete, visible introductory coverage."""
import json
from io import BytesIO
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from adaptive_document_agent.agent.report_localization import CopyTranslations, prepare_report_language
from adaptive_document_agent.models.customization import ReportRequirement
from adaptive_document_agent.services.customization_checks import check_requirements
from adaptive_document_agent.services.executive_brief import restore_percentage_symbols
from adaptive_document_agent.services.report_language import apply_report_language, audience_copy
from adaptive_document_agent.services.summary_export import verify_summary_export
from tests.test_complete_summary_reading import SummaryClient, distinct_source, generate, complete_deck
from tests.test_report_customization import table_result, presentation


def test_chinese_body_and_original_tables_are_compatible_and_independently_verified(monkeypatch):
    from adaptive_document_agent.services.requested_tables import render_requested_tables
    from adaptive_document_agent.services import pptx_export
    result, deck = table_result(), presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1)).text = 'Company overview'
    render_requested_tables(deck, result)
    result.profile.report_requirements.items.extend([
        ReportRequirement(id='body', kind='output_language', request_quote='正文中文',
            description='正文中文', language='zh', language_scope='body'),
        ReportRequirement(id='raw', kind='output_language', request_quote='原表原文',
            description='原表原文', language='original', language_scope='source_tables')])
    stream = BytesIO()
    deck.save(stream)
    monkeypatch.setattr(pptx_export, 'presentation_for_copy', lambda *a, **k: deck)
    seen = []
    def call(messages, model, **kwargs):
        seen.append(model)
        rows = json.loads(messages[-1]['content'].split('\n', 1)[1].rsplit('\n', 1)[0])
        if model is CopyTranslations:
            return model(items=[{'id': row['id'], 'text': '中文：' + row['text']} for row in rows])
        return model(items=[{'id': row['id'], 'accepted': True, 'reason': 'Equivalent checked display copy'} for row in rows])
    prepare_report_language(SimpleNamespace(generate_structured=call), result)
    assert seen and all(c.status != 'ambiguous' for c in result.customization_report)
    apply_report_language(deck, result)
    check_requirements(result, deck, export_verified=True)
    checks = {c.requirement_id: c for c in result.customization_report}
    assert checks['body'].status == checks['raw'].status == 'satisfied'


def test_chart_display_labels_translate_without_mutating_original_workbook():
    result, deck = table_result(), presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    data = CategoryChartData()
    data.categories = ['Domestic', 'Overseas']
    data.add_series('Revenue', [10, 20])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(8), Inches(4), data).chart
    workbook = chart.part.chart_workbook.xlsx_part.blob
    assert {'Domestic', 'Overseas', 'Revenue'} <= set(audience_copy(deck))
    result.profile.report_requirements.copy_translations = {'Domestic': '国内', 'Overseas': '海外', 'Revenue': '收入'}
    apply_report_language(deck, result)
    assert {'国内', '海外', '收入'} <= set(audience_copy(deck))
    assert chart.part.chart_workbook.xlsx_part.blob == workbook
    assert list(chart.series[0].values) == [10, 20]


def test_all_requested_tables_keep_the_reference_style_and_literal_context():
    from adaptive_document_agent.services.requested_tables import render_requested_tables, verify_requested_tables
    from adaptive_document_agent.services.pptx_export import FOURIER_PURPLE, FOURIER_DARK, WHITE
    result, deck = table_result(), presentation()
    render_requested_tables(deck, result)
    labels = [s for slide in deck.slides for s in slide.shapes if s.name == 'customization:table_label']
    sections = [s for slide in deck.slides for s in slide.shapes if s.name == 'customization:source_section']
    before = [s.text for s in [*labels, *sections]]
    result.profile.report_requirements.copy_translations = {text: 'changed source context' for text in before}
    apply_report_language(deck, result)
    assert [s.text for s in [*labels, *sections]] == before
    grids = [s.table for slide in deck.slides for s in slide.shapes
             if s.has_table and s.name.startswith('customization:source_table:')]
    assert grids and labels and sections
    for table in grids:
        for cell in table.rows[0].cells:
            assert str(cell.fill.fore_color.rgb) == FOURIER_PURPLE
            assert str(cell.text_frame.paragraphs[0].font.color.rgb) == WHITE
            assert cell.text_frame.paragraphs[0].font.bold
        for cell in table.rows[1].cells:
            assert str(cell.text_frame.paragraphs[0].font.color.rgb) == FOURIER_DARK
    assert verify_requested_tables(deck, result)


@pytest.mark.parametrize('text', ['2024年上半年收入为8.8%', '2024年度收入', 'For 2024 H1'])
def test_period_headers_never_restore_a_percent_sign_on_a_year(text):
    assert restore_percentage_symbols(text, ['2021 2022 2023 2024\n% of Total\n6.5 8.8']) == text
    assert restore_percentage_symbols('Share was 8.8', ['Share was 8.8%']) == 'Share was 8.8%'


def test_more_than_eight_introduction_pages_cover_every_content_part():
    result = distinct_source(20)
    for page in result.document.pages:
        page.text = page.text.replace('Cedar provides', f'Cedar provides {chr(64 + page.page_number)}')
        page.text = page.text.replace('Delivery depends', f'Delivery depends {chr(96 + page.page_number)}')
    client = SummaryClient()
    generated = generate(result, client)
    company = generated.presentation_plan.company
    assert len(company.summary_pages) > 8 and len(client.plan_inputs) > 1
    content = {p.id for p in company.summary_review.parts if p.role == 'content'}
    covered = {pid for page in company.summary_pages for item in page.items for pid in item.part_ids}
    assert covered == content
    assert company.summary_review.status == 'complete'


def test_export_rejects_a_summary_fact_retained_only_in_notes():
    from adaptive_document_agent.services.pptx_export import build_presentation
    result = complete_deck(generate())
    deck = Presentation(BytesIO(build_presentation(result)))
    assert verify_summary_export(deck, result)
    shape = next(s for slide in deck.slides for s in slide.shapes if s.name == 'brief:body')
    shape.text = ''
    with pytest.raises(ValueError, match='missing from visible'):
        verify_summary_export(deck, result)
