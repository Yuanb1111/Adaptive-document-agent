"""Failures observed in the customized/default v104 runs, with exact source gates."""
import json

import pytest

from adaptive_document_agent.agent.summary_editorial import SummaryEditorialDraft
from adaptive_document_agent.services.llm.base import LLMResponse
from adaptive_document_agent.services.llm.exceptions import LLMStructuredOutputError
from adaptive_document_agent.validation.presentation_plan_validator import PresentationPlanValidator
from tests.test_complete_summary_reading import SummaryClient, distinct_source, generate


def test_pdf_year_row_does_not_borrow_percent_from_next_header_line():
    numbers = PresentationPlanValidator._numbers('2021 2022 2023 2023 2024\n\n% of Total\n33%-44%')
    assert {'2024', '33%', '44%'} <= numbers
    assert '2024%' not in numbers and '-44%' not in numbers
    assert PresentationPlanValidator._numbers('Margin 44 %; loss -44%; change - 44') == {'44%', '-44%', '-44'}


def test_complete_summary_accepts_supported_year_annotation_but_rejects_invented_year():
    result = distinct_source(1)
    result.document.pages[0].text = result.document.pages[0].text.replace(
        'SUMMARY\n', 'SUMMARY\n2021 2022 2023 2023 2024\n% of Total\n')

    def edit(response, payload, messages):
        next(p for p in response.parts if p.role == 'layout').reading_note = 'Periods include 2024.'

    assert generate(result, SummaryClient(edit_read=edit)).presentation_plan.company.summary_review.status == 'complete'
    bad = distinct_source(1)
    with pytest.raises(ValueError, match='unsupported numbers'):
        generate(bad, SummaryClient(edit_read=edit))


def test_intro_schema_failure_gets_one_source_bound_repair_and_keeps_diagnostic():
    class InvalidFirstPlan(SummaryClient):
        failed = False

        def generate_structured(self, messages, response_model, **kwargs):
            if response_model is SummaryEditorialDraft and not self.failed:
                self.failed = True
                raise LLMStructuredOutputError('schema validation failed',
                    response=LLMResponse(text='{"summary_pages": "invalid"}')) from ValueError('summary_pages must be a list')
            return super().generate_structured(messages, response_model, **kwargs)

    result = generate(distinct_source(1), InvalidFirstPlan())
    review = result.presentation_plan.company.summary_review
    assert review.status == 'complete'
    assert len(review.plan_audits) == 2
    assert review.plan_audits[0]['validation_detail'] == 'summary_pages must be a list'
    assert 'invalid' in review.plan_audits[0]['raw_response']
    content = {p.id for p in review.parts if p.role == 'content'}
    assert content == {pid for page in result.presentation_plan.company.summary_pages for item in page.items for pid in item.part_ids}


def test_repeated_intro_schema_failure_remains_incomplete_with_both_attempts():
    class InvalidPlan(SummaryClient):
        def generate_structured(self, messages, response_model, **kwargs):
            if response_model is SummaryEditorialDraft:
                raise LLMStructuredOutputError('schema validation failed', response=LLMResponse(text='{}'))
            return super().generate_structured(messages, response_model, **kwargs)

    result = distinct_source(1)
    with pytest.raises(LLMStructuredOutputError):
        generate(result, InvalidPlan())
    review = json.loads(next(w.message for w in result.validation_warnings if w.code == 'company_introduction_summary_audit'))
    assert review['status'] == 'incomplete' and len(review['plan_audits']) == 2


def test_combined_stock_date_headers_do_not_invent_an_annual_date_for_an_interim_year():
    from adaptive_document_agent.services.brief_table_evidence import _header_dates
    from tests.test_brief_table_context import table_result
    result, _ = table_result()
    table = result.document.pages[0].tables[0]
    table.raw_header_lines = ['As of December 31,', 'June 30, October 31,', '2021 2022 2023 2024 2024']
    table.column_periods = [None, '2021-12-31', '2022-12-31', '2023-12-31', '2024-06-30', '2024-10-31']
    source = '\n'.join(table.raw_header_lines)
    dates = _header_dates(table, source)
    assert ('june', '30', '2024') in dates and ('october', '31', '2024') in dates
    assert ('december', '31', '2023') in dates and ('december', '31', '2024') not in dates


def test_meaning_review_accepts_literal_unitful_endpoints_but_rejects_rescaling():
    from adaptive_document_agent.agent.brief_meaning_review import MeaningComparison, comparison_errors
    from adaptive_document_agent.models.executive_brief import ExecutiveBriefItem, BriefQuote
    text = 'The rate rose from RMB3.5 million to RMB20.6 million.'
    item = ExecutiveBriefItem(label='Rate', text=text, evidence=[BriefQuote(page=1, text=text)])
    comparison = MeaningComparison(claim=text, start_value='RMB3.5 million', end_value='RMB20.6 million',
        direction='increase', start_context=text, end_context=text)
    assert not comparison_errors(comparison, item)
    assert comparison_errors(comparison.model_copy(update={'end_value':'RMB20.6 billion'}), item)
    assert comparison_errors(comparison.model_copy(update={'direction':'decrease'}), item)


def test_native_repair_options_supply_exact_annotations_for_loss_magnitudes():
    from adaptive_document_agent.agent.brief_native_quantities import native_quantity_options
    from tests.test_brief_table_context import table_result
    result, brief = table_result()
    table = result.document.pages[0].tables[0]
    table.rows[0].cells = ['Net loss', '(1,250)', '(1,475)']
    quote = 'Net loss (1,250) (1,475)'
    result.document.pages[0].text = '\n'.join(table.raw_header_lines) + '\n' + quote
    brief.items[0].evidence[0].text = quote
    options = native_quantity_options(brief.model_dump(mode='json'), {0:[]}, result,
                                     {1:result.document.pages[0].text})['item_0']
    option = next(o for o in options if o['signed_coefficient'] == '-1475')
    assert option['magnitude_option']['quantity_representation'] == {
        'quantity_text':'US$ 1475 thousand', 'source_value':'-1475', 'representation':'absolute_magnitude'}


def test_copy_inventory_is_independent_of_export_preflight_and_raw_table_layout(monkeypatch):
    from adaptive_document_agent.services import pptx_export, requested_tables
    from adaptive_document_agent.services.report_language import audience_copy
    from tests.test_report_customization import table_result
    from tests.test_complete_summary_reading import complete_deck
    result = complete_deck(generate(distinct_source(1)))
    result.profile.report_requirements = table_result().profile.report_requirements

    def blocked(*args, **kwargs):
        raise ValueError('Unrelated table/export gate is blocked')

    monkeypatch.setattr(PresentationPlanValidator, 'validate', blocked)
    monkeypatch.setattr(requested_tables, 'render_requested_tables', blocked)
    inventory = pptx_export.presentation_for_copy(result)
    assert 'Operations and delivery' in audience_copy(inventory)
    with pytest.raises(ValueError, match='Unrelated'):
        pptx_export.build_presentation(result)


def test_localization_repairs_only_failed_copy_and_preserves_verified_translations(monkeypatch):
    from types import SimpleNamespace
    from pptx.util import Inches, Pt
    from adaptive_document_agent.agent.report_localization import CopyTranslations, prepare_report_language
    from adaptive_document_agent.services import pptx_export
    from adaptive_document_agent.services.report_language import apply_report_language
    from adaptive_document_agent.models.customization import ReportRequirement, ReportRequirements
    from tests.test_report_customization import presentation, table_result
    deck = presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    frames = []
    for index, text in enumerate(['Revenue RMB 10 million in FY2025', 'Debt stable']):
        shape = slide.shapes.add_textbox(Inches(1), Inches(2 + index), Inches(8), Inches(.6))
        shape.text = text
        shape.text_frame.paragraphs[0].font.size = Pt(18)
        frames.append(shape)
    monkeypatch.setattr(pptx_export, 'presentation_for_copy', lambda *args: deck)
    result = table_result()
    result.profile.report_requirements = ReportRequirements(items=[ReportRequirement(id='zh',
        kind='output_language', request_quote='中文', description='中文正文', language='zh')])
    calls = []

    def call(messages, model, **kwargs):
        payload = next(json.loads(m['content'].split('\n', 1)[1].rsplit('\n', 1)[0])
                       for m in messages if m['content'].startswith('<UNTRUSTED_DOCUMENT_CONTENT>'))
        calls.append((model.__name__, [row['id'] for row in payload]))
        if model is CopyTranslations:
            return model(items=[{'id': row['id'], 'text': 'FY2025收入RMB 10 million' if row['id'] == 0 else
                '负债稳定' if len(calls) > 2 else '负债稳定' * 100} for row in payload])
        return model(items=[{'id': row['id'], 'accepted': True, 'reason':'Faithful target-language prose'} for row in payload])

    prepare_report_language(SimpleNamespace(generate_structured=call), result)
    assert calls == [('CopyTranslations',[0,1]), ('TranslationReview',[0,1]),
                     ('CopyTranslations',[1]), ('TranslationReview',[1])]
    assert result.customization_report[-1].status == 'planned'
    apply_report_language(deck, result)
    assert [s.text for s in frames] == ['FY2025收入RMB 10 million', '负债稳定']
    audit = json.loads(next(w.message for w in result.validation_warnings if w.code == 'report_localization_audit'))
    assert audit['accepted_count'] == 2 and not audit['failed_ids']
    assert 'readable native text box' in str(audit['attempts'][0]['errors'])


def test_localization_copy_discovery_failure_is_actionable_in_exported_audit(monkeypatch):
    from types import SimpleNamespace
    from adaptive_document_agent.agent.report_localization import prepare_report_language
    from adaptive_document_agent.services import pptx_export
    from adaptive_document_agent.models.customization import ReportRequirement, ReportRequirements
    from tests.test_report_customization import table_result
    result = table_result()
    result.profile.report_requirements = ReportRequirements(items=[ReportRequirement(id='zh',
        kind='output_language', request_quote='中文', description='中文正文', language='zh')])

    def blocked(*args):
        raise ValueError('Missing readable source for topic_scope')

    monkeypatch.setattr(pptx_export, 'presentation_for_copy', blocked)
    prepare_report_language(SimpleNamespace(), result)
    assert result.customization_report[-1].status == 'partial'
    audit = json.loads(next(w.message for w in result.validation_warnings if w.code == 'report_localization_audit'))
    assert audit['attempts'][0]['stage'] == 'copy_discovery'
    assert audit['attempts'][0]['error'] == 'Missing readable source for topic_scope'


def test_source_table_copy_is_immutable_through_every_native_sanitizer():
    from pptx.util import Inches
    from adaptive_document_agent.services.ppt_preflight import PresentationPreflight
    from tests.test_report_customization import presentation
    deck = presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    label = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    label.name = 'customization:table_label'
    label.text = 'Source None  parser debug {note} (RMBinthousands)'
    shape = slide.shapes.add_table(2,2,Inches(1),Inches(2),Inches(8),Inches(2))
    shape.name = 'customization:source_table:literal'
    shape.table.cell(0,0).text = 'Revenue 1,200%'
    shape.table.cell(0,1).text = 'RMBinthousands'
    before = label.text, [[c.text for c in row.cells] for row in shape.table.rows]
    PresentationPreflight(deck).validate_and_sanitize()
    assert (label.text, [[c.text for c in row.cells] for row in shape.table.rows]) == before


def test_appendix_footer_reserves_the_actual_template_copyright_box():
    from pptx import Presentation
    from adaptive_document_agent.services.pptx_export import _resolve_template_path
    from adaptive_document_agent.services.requested_tables import render_requested_tables
    from tests.test_report_customization import table_result
    deck = Presentation(_resolve_template_path())
    render_requested_tables(deck,table_result())
    checked = False
    for slide in deck.slides:
        for footer in [s for s in slide.shapes if s.name=='customization:source_footer']:
            for layer in (slide.slide_layout,slide.slide_layout.slide_master):
                for shape in layer.shapes:
                    if (not shape.is_placeholder and shape.has_text_frame and shape.text.strip()
                        and deck.slide_height.inches*.75 < shape.top.inches < deck.slide_height.inches
                        and shape.left < footer.left+footer.width and shape.left+shape.width > footer.left):
                        assert footer.top.inches+footer.height.inches <= shape.top.inches-.09
                        checked = True
    assert checked


def test_summary_density_counts_wide_characters_without_padding_chinese_copy():
    from adaptive_document_agent.models.summary import SummaryPart, SummaryFact, SummaryReview, SummarySlidePage
    from adaptive_document_agent.services.summary_validation import presentation_errors
    result = distinct_source(1)
    texts = [
        '系统提供排班与人员可用性管理功能，适用于制造团队的现场生产协调；功能范围以文件披露为准，不补充未披露能力。',
        '项目交付依赖客户系统集成与员工培训，文件说明了这些实施条件；没有承诺完成日期，也没有保证客户最终采用方案。']
    result.document.pages[0].text = '\n'.join(texts)
    assert 90 <= sum(map(len,texts)) < 180
    parts = [SummaryPart(id=str(i),block_id='source',source_page=1,start_line=i+1,end_line=i+1,
        title=label,role='content',reading_note='完整来源内容',source_text=text,
        facts=[SummaryFact(label=label,text=text,source_quote=text,source_pages=[1])])
        for i,(label,text) in enumerate(zip(['业务功能','交付条件'],texts))]
    review = SummaryReview(document_id=result.document.document_id,document_sha256=result.document.sha256,
        source_pages=[1],parts=parts,decisions=[{'part_id':p.id,'decision':'include','reason':'包含完整实质内容'} for p in parts])
    page = SummarySlidePage(id='intro',title='业务与交付',source_section='Summary',items=[{
        **p.facts[0].model_dump(),'part_ids':[p.id]} for p in parts])
    assert not presentation_errors(review,[page],result,validate_reading=False)


def test_long_chinese_introduction_paginates_with_visible_binding_for_every_part():
    from io import BytesIO
    from pptx import Presentation
    from adaptive_document_agent.models.presentation import CompanyProfile
    from adaptive_document_agent.models.summary import SummaryPart, SummaryFact, SummaryReview, SummarySlidePage
    from adaptive_document_agent.services.summary_source import source_blocks
    from adaptive_document_agent.services.summary_validation import presentation_errors
    from adaptive_document_agent.services.pptx_export import build_presentation
    from adaptive_document_agent.services.summary_export import verify_summary_export
    from tests.test_complete_summary_reading import complete_deck
    result = complete_deck(generate(distinct_source(1)))
    # Deliberately dense, literal source fixture: each whole item fits, the
    # four-item logical page needs native continuation at normal type sizes.
    texts = [label + '。' + ('文件披露了功能范围和实施条件，并说明交付依赖系统集成及人员培训，不保证完成日期或最终采用。' * 3)
             for label in ['系统功能','产品应用','平台部署','服务交付']]
    assert all(len(t) <= 180 for t in texts)
    result.document.pages[0].text = 'SUMMARY\n' + '\n'.join(texts) + '\n'
    block = source_blocks(result,[1])[0]
    parts = [SummaryPart(id='heading',block_id=block.id,start_line=1,end_line=1,source_page=1,
        source_text='SUMMARY\n',title='Summary',role='heading',reading_note='Standalone source heading')]
    parts += [SummaryPart(id=str(i),block_id=block.id,start_line=i+2,end_line=i+2,source_page=1,
        source_text=text+'\n',title=label,role='content',reading_note='完整实质内容',
        facts=[SummaryFact(label=label,text=text,source_quote=text,source_pages=[1])])
        for i,(label,text) in enumerate(zip(['系统功能','产品应用','平台部署','服务交付'],texts))]
    review = SummaryReview(document_id=result.document.document_id,document_sha256=result.document.sha256,
        source_pages=[1],source_blocks=[block],parts=parts,status='complete',decisions=[{
            'part_id':p.id,'decision':'include' if p.role=='content' else 'omit','reason':'已完整归纳或独立标题'} for p in parts])
    page = SummarySlidePage(id='intro',title='公司业务与交付',source_section='Summary',items=[{
        **p.facts[0].model_dump(),'part_ids':[p.id]} for p in parts if p.role=='content'])
    assert not presentation_errors(review,[page],result)
    result.presentation_plan.company = CompanyProfile(summary_review=review,summary_pages=[page],source_pages=[1])
    deck = Presentation(BytesIO(build_presentation(result)))
    mapping = verify_summary_export(deck,result)
    assert set(mapping) == {p.id for p in parts if p.role=='content'}
    assert len({n for numbers in mapping.values() for n in numbers}) > 1
