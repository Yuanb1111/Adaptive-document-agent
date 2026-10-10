"""Production failures in v105: source reading, localization and native roles."""
from io import BytesIO

import pytest
from pptx import Presentation
from pptx.util import Inches

from adaptive_document_agent.models.summary import SummaryFact
from adaptive_document_agent.models.executive_brief import ExecutiveBriefItem
from adaptive_document_agent.services.summary_validation import fact_errors
from adaptive_document_agent.services.ppt_preflight import PresentationPreflight
from adaptive_document_agent.services.pptx_export import _add_thank_you_slide, _resolve_template_path
from tests.test_complete_summary_reading import distinct_source


def test_full_reading_fact_limit_is_not_the_shorter_brief_display_limit():
    result = distinct_source(1)
    text = 'Delivery depends on integration, staff training and source conditions. ' * 8
    assert 550 < len(text) <= 600
    result.document.pages[0].text = text
    fact = SummaryFact(label='Delivery', text=text, source_quote=text, source_pages=[1])
    assert fact_errors(fact, result, text, 1) == []
    with pytest.raises(ValueError):
        ExecutiveBriefItem(label='Delivery', text=text, evidence=[{'page':1,'text':text}])
    fact.text = text + ' Unsupported 987654.'
    assert any('unsupported numeric' in e for e in fact_errors(fact, result, text, 1))


def test_localized_official_closing_role_survives_save_and_is_not_empty():
    deck = Presentation(_resolve_template_path())
    _add_thank_you_slide(deck)
    deck.slides[-1].shapes[-1].text = '谢谢'
    output = BytesIO(); deck.save(output)
    restored = Presentation(BytesIO(output.getvalue()))
    check = PresentationPreflight(restored)
    check._check_empty_slide(len(restored.slides)-1,restored.slides[-1])
    check._check_thank_you_slide()
    assert not check.issues


def test_short_unmarked_content_is_still_rejected_as_empty():
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text = '谢谢'
    check=PresentationPreflight(deck);check._check_empty_slide(1,slide)
    assert any(i.code=='empty_slide' for i in check.issues)


@pytest.mark.parametrize('source,target', [
    ('FY2021 to FY2023: +112.4 RMB millions','2021财年至2023财年：+112.4百万元人民币'),
    ('31 Dec 2021 to 31 Oct 2024*: +78.6 RMB millions','2021年12月31日至2024年10月31日*：+78.6百万元人民币'),
    ('6M2023 to 6M2024: -6.3 pp','2023年上半年至2024年上半年：-6.3个百分点'),
    ('Share (%)','占比（%）'),
    ('Four-axis cobots: % of revenue','四轴协作机器人：收入占比'),
    ('Revenue FY2023 RMB 10m','收入FY2023 10百万元人民币'),
    ('Gross profit margin: FY2021 50.5%; FY2022 40.8%; FY2023 43.5%.','毛利率：2021财年50.5%；2022财年40.8%；2023财年43.5%。'),
])
def test_faithful_chinese_display_spellings_preserve_numeric_scope(source,target):
    from adaptive_document_agent.services.report_language import translation_errors
    assert not translation_errors(source,target)


@pytest.mark.parametrize('source,target', [
    ('6M2023 to 6M2024: -6.3 pp','2023年上半年至2024年下半年：-6.3个百分点'),
    ('FY2023 10 RMB millions','FY2023 10十亿元人民币'),
    ('FY2023 10 RMB millions','FY2023 1千万元人民币'),
    ('31 Oct 2024: -7.76 RMB millions','2024年6月30日：-7.76百万元人民币'),
    ('FY2023 10 RMB millions','FY2023 10百万美元'),
    ('FY2023 -10 RMB millions','FY2023 10百万元人民币'),
    ('FY2023 10 and 10 RMB millions','FY2023 10百万元人民币'),
    ('FY2023 43.5%','FY2023 43.5'),
    ('FY2023 RMB 10 million; USD 20 million','FY2023 20百万元人民币；10百万美元'),
    ('6M2023* 43.5%','6M2023 43.5%'),
])
def test_localization_still_rejects_changed_period_currency_scale_sign_or_repeated_value(source,target):
    from adaptive_document_agent.services.report_language import translation_errors
    assert translation_errors(source,target)


def test_localized_line_breaks_are_native_and_keep_font_properties():
    from pptx.util import Pt
    from adaptive_document_agent.models.customization import ReportRequirements
    from adaptive_document_agent.services.report_language import apply_report_language, audience_copy
    deck = Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
    shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(2))
    shape.text='Operating flow\vInvesting flow'
    shape.text_frame.paragraphs[0].runs[0].font.size=Pt(18)
    result=distinct_source(1)
    result.profile.report_requirements=ReportRequirements(copy_translations={
        'Operating flow\nInvesting flow':'经营现金流\v投资现金流'})
    assert 'Operating flow\nInvesting flow' in audience_copy(deck)
    apply_report_language(deck,result)
    assert shape.text == '经营现金流\n投资现金流'
    assert all(p.runs[0].font.size.pt==18 for p in shape.text_frame.paragraphs)
    assert all(len(p._p.xpath('./a:pPr'))==1 for p in shape.text_frame.paragraphs)
    assert '_x000B_' not in shape._element.xml


def test_single_complete_row_requires_independent_year_geometry_for_sparse_alignment():
    from adaptive_document_agent.extraction.borderless_layout import SourceLine, header_supported_anchors, align_sparse_values
    line=SourceLine('2021 2022 2023',[(0,4,100,120),(5,9,200,220),(10,14,300,320)])
    anchors=header_supported_anchors([[(100,125),(200,225),(300,325)]],line,3)
    assert align_sparse_values(['29.3%'],[(300,325)],anchors,3)==[None,None,'29.3%']
    assert header_supported_anchors([[(30,55),(80,105),(130,155)]],line,3) is None
    assert header_supported_anchors([[(100,125),(200,225),(300,325)]],SourceLine(line.text),3) is None


def test_dot_leader_section_does_not_split_table_and_unlabelled_subtotal_stays_raw_only():
    from adaptive_document_agent.extraction.borderless_table_extractor import BorderlessTableExtractor
    class Page:
        def __init__(self):
            self.words=[]
            lines=[('Year ended December 31,',20,40),('2021',40,100),('2022',40,200),('2023',40,300),
                   ('USD in thousands',60,40),('Main assets',80,20),('10',80,100),('20',80,200),('30',80,300),
                   ('Other assets',100,20),('15',100,100),('25',100,200),('35',100,300),
                   ('Equity attributable to owners . . . .',120,20),
                   ('Capital',140,20),('10',140,100),('20',140,200),('30',140,300),
                   ('15',160,100),('25',160,200),('35',160,300),
                   ('Total equity',180,20),('25',180,100),('45',180,200),('65',180,300)]
            for text,y,x in lines:
                self.words.append({'text':text,'x0':x,'x1':x+20,'top':y,'bottom':y+10})
        def extract_text(self,**kwargs): return ''
        def extract_words(self,**kwargs): return self.words
    tables=BorderlessTableExtractor().extract(Page(),1)
    assert len(tables)==1
    table=tables[0]
    assert ['', '15','25','35'] in table.raw_cells
    assert all(r.cells[0] for r in table.rows)
    assert table.column_periods[1:]==['FY2021','FY2022','FY2023']


def test_exporter_capabilities_are_verified_in_native_ppt_and_detect_style_damage():
    from adaptive_document_agent.models.customization import ReportRequirement
    from adaptive_document_agent.services.customization_checks import check_requirements
    from adaptive_document_agent.services.requested_tables import render_requested_tables
    from tests.test_report_customization import table_result,presentation
    result=table_result();deck=presentation()
    result.profile.report_requirements.items.extend([
        ReportRequirement(id='style',request_quote='紫色表头',description='Editable purple source tables',kind='table_presentation'),
        ReportRequirement(id='cite',request_quote='来源页码',description='Page citations',kind='source_citations')])
    check_requirements(result)
    assert all(c.status=='planned' for c in result.customization_report if c.requirement_id in {'style','cite'})
    render_requested_tables(deck,result);check_requirements(result,deck,export_verified=True)
    assert all(c.status=='satisfied' for c in result.customization_report if c.requirement_id in {'style','cite'})
    next(s for s in deck.slides[0].shapes if s.has_table).table.cell(0,0).text_frame.paragraphs[0].font.bold=False
    check_requirements(result,deck,export_verified=True)
    assert next(c for c in result.customization_report if c.requirement_id=='style').status=='not_met'


def test_flow_calendar_dates_bind_to_quoted_duration_headers_without_cartesian_year_borrowing():
    from adaptive_document_agent.services.brief_table_evidence import _header_dates
    from tests.test_brief_table_context import table_result
    result,_=table_result();table=result.document.pages[0].tables[0]
    table.raw_header_lines=['Year ended December 31, Six months ended June 30,','2021 2022 2023 2023 2024']
    table.column_periods=[None,'FY2021','FY2022','FY2023','6M2023','6M2024']
    dates=_header_dates(table,'\n'.join(table.raw_header_lines))
    assert ('june','30','2024') in dates and ('december','31','2023') in dates
    assert ('december','31','2024') not in dates


def test_compact_quantity_format_preserves_value_unit_and_chinese_prose():
    from adaptive_document_agent.services.display_copy_tokens import compact_display_units
    from adaptive_document_agent.services.report_language import translation_errors,translatable
    source='总负债：人民币87.1百万元；现金流变化：-78.6百万元人民币。'
    compact=compact_display_units(source)
    assert compact=='总负债：RMB 87.1m；现金流变化：RMB -78.6m。'
    assert not translation_errors(source,compact)
    assert not translatable('-RMB 157.7m')
    assert translatable('Revenue decreased to RMB 157.7m')
