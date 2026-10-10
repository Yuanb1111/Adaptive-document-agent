"""Complete reading, faithful Chinese copy and source-only grid recovery."""
from threading import Event
import json

import pytest

from adaptive_document_agent.services.report_language import translation_errors
from adaptive_document_agent.extraction.source_grid_recovery import recover_source_grids
from adaptive_document_agent.agent.summary_reader import _read_batch
from adaptive_document_agent.models.summary import SummarySlideItem
from tests.test_complete_summary_reading import SummaryClient, distinct_source, gateway, generate
from adaptive_document_agent.services.summary_source import source_blocks


@pytest.mark.parametrize('source,target', [
    ('-RMB 51.6m in 6M2023*','6M2023*为人民币-51.6百万元'),
    ('-9.2% vs 6M2023; prior 4,918 units','较6M2023下降9.2%；前期4,918台'),
    ('+9.6% vs 6M2023*; prior RMB109.9m','较6M2023*增长9.6%；前期人民币109.9百万元'),
    ('FY and 6M period bases are presented separately.','FY和6M期间基准分别列示。'),
    ('The October 2024 balance sheet date is unaudited.','2024年10月资产负债表日未经审核。'),
    ('RMB -6M','人民币-6百万元'),
    ('EUR 6M','6百万欧元'),
])
def test_faithful_translation_is_not_rejected_by_display_representation(source,target):
    assert not translation_errors(source,target)


@pytest.mark.parametrize('source,target', [
    ('-9.2% vs 6M2023','较6M2023增长9.2%'),
    ('The October 2024 date is unaudited.','2024年10月31日未经审核。'),
    ('RMB 6M','6M期间'),
    ('-RMB 51.6m','人民币51.6百万元'),
])
def test_wrong_direction_invented_date_and_changed_money_remain_rejected(source,target):
    assert translation_errors(source,target)


@pytest.mark.parametrize('bad', ['Cedar has 999 customers.', 'Cedar revenue was RMB 30 million.'])
def test_unverified_reading_paraphrase_is_replaced_by_exact_selected_source_not_published(bad):
    def edit(response, payload, messages):
        next(p for p in response.parts if p.facts).facts[0].text = bad
    result = generate(distinct_source(1), SummaryClient(edit_read=edit))
    review = result.presentation_plan.company.summary_review
    assert review.status == 'complete'
    extracts = review.read_audits[0]['attempts'][1]['literal_extracts']
    assert extracts[0]['rejected_paraphrase'] == bad and extracts[0]['errors']
    assert all(bad not in item.text for page in result.presentation_plan.company.summary_pages for item in page.items)


def test_long_internal_extract_does_not_expand_published_slide_limit():
    text = 'Delivery depends on integration and customer training. ' * 20
    result = distinct_source(1);result.document.pages[0].text = text
    blocks = source_blocks(result,[1])
    class Reader:
        def generate_structured(self,messages,model,**kwargs):
            return model.model_validate({'parts':[{'block_id':blocks[0].id,'start_line':1,'end_line':1,
                'title':'Delivery','role':'content','reading_note':'Delivery qualifications.',
                'facts':[{'label':'Delivery','text':text,'quote_start_line':1,'quote_end_line':1}]}]})
    outcome = _read_batch(Reader(), result, blocks, Event())
    assert outcome.error is None and outcome.parts[0].facts[0].text == text
    with pytest.raises(ValueError):
        SummarySlideItem(text=text,label='Delivery',source_pages=[1],source_quote=text,part_ids=['part'])


class WordPage:
    def __init__(self,rows):
        self.words = []
        for y,row in enumerate(rows):
            for text,left,right in row:
                self.words.append(dict(text=text,x0=left,x1=right,top=y*15,bottom=y*15+10))
    def extract_words(self,**kwargs):
        return self.words


def test_qualitative_source_table_needs_repeated_geometry_not_years():
    rows = [[('Category A',20,100),('Shorter of remaining lease term and useful life',160,420)],
            [('Category B',20,95),('3.17% to 4.75%',160,250)],
            [('Category C',20,110),('19% to 32%',160,245)]]
    tables = recover_source_grids(WordPage(rows),1,[])
    assert len(tables)==1
    assert tables[0].raw_cells == [[c[0] for c in row] for row in rows]
    assert not tables[0].rows and not tables[0].column_periods and not tables[0].default_currency
    # Repeated header tiers and prose without physical boundaries are rejected.
    assert not recover_source_grids(WordPage([[('Gross',100,125),('Gross',200,225),('Gross',300,325)],
        [('profit',100,125),('profit',200,225),('profit',300,325)],
        [('margin',100,125),('margin',200,225),('margin',300,325)]]),1,[])


def test_price_in_scenario_label_is_not_a_sixth_data_value():
    rows = [[('Column A',200,235),('Column B',280,315)],
            [('Price scenario',20,100)],
            [('USD18.80 per unit',25,150),('323,948',200,235),('2.60',295,315)],
            [('Price scenario',20,100)],
            [('USD19.80 per unit',25,150),('323,948',200,235),('2.70',295,315)],
            [('Price scenario',20,100)],
            [('USD20.80 per unit',25,150),('323,948',200,235),('2.80',295,315)]]
    table, = recover_source_grids(WordPage(rows),1,[])
    assert table.raw_cells[0] == ['Price scenario USD18.80 per unit','323,948','2.60']
    assert table.raw_header_cells == [['','Column A','Column B']]


def test_model_selects_oversized_brief_then_all_selected_claims_are_validated():
    from adaptive_document_agent.agent.executive_brief import ExecutiveBriefWriter
    from tests.test_executive_brief import result_for, gateway as brief_gateway, payload
    findings = [f'Measure {i} was reported.' for i in range(8)]
    original = {'title':'Findings','items':[payload(t)['items'][0] for t in findings]}
    result = result_for([' '.join(findings)])
    g,client = brief_gateway([original,{'indices':[6,1,4]}])
    brief = ExecutiveBriefWriter(g).generate(result)
    assert [i.text for i in brief.items] == [findings[i] for i in [6,1,4]]
    audit = json.loads(next(w.message for w in result.validation_warnings if w.code=='executive_brief_repair_audit'))
    assert audit['original'] == original and audit['selection']['indices']==[6,1,4]
    assert len(client.calls)==2


def test_numeric_labels_recovered_from_mistaken_header_context_extend_existing_fragment():
    from adaptive_document_agent.models.table import ExtractedTable
    page = WordPage([[('Within 1 year',20,100),('10',200,220),('20',300,320)],
                     [('1 to 2 years',20,100),('11',200,220),('21',300,320)],
                     [('Over 2 years',20,100),('12',200,220),('22',300,320)]])
    existing = ExtractedTable(table_id='original',page=1,bbox=(20,50,320,70),
        raw_cells=[['Total','33','63']],raw_header_lines=[
            'Within 1 year 10 20','1 to 2 years 11 21','Over 2 years 12 22'])
    assert recover_source_grids(page,1,[existing]) == []
    assert existing.raw_cells == [['Within 1 year','10','20'],['1 to 2 years','11','21'],
                                 ['Over 2 years','12','22'],['Total','33','63']]


def test_post_inventory_chart_wrap_rebinds_only_exact_whitespace_equivalent_copy():
    from io import BytesIO
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    from adaptive_document_agent.models.customization import ReportRequirements
    from adaptive_document_agent.services.report_language import apply_report_language, chart_display_strings
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
    data=CategoryChartData();data.categories=['FY2023'];data.add_series('Adjusted net loss \n(non-IFRS \nmeasure)',[3])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(1),Inches(1),Inches(8),Inches(4),data)
    result=distinct_source(1);result.profile.report_requirements=ReportRequirements(copy_translations={
        'Adjusted net loss (non-IFRS measure)':'调整后净亏损（非IFRS指标）'})
    apply_report_language(deck,result)
    assert any(e.text=='调整后净亏损（非IFRS指标）' for e in chart_display_strings(deck))
