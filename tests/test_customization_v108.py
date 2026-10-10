"""Source-scoped regression cases from the v107 acceptance defects."""
import json
from io import BytesIO

import pytest
from pptx import Presentation

from adaptive_document_agent.extraction.borderless_table_extractor import BorderlessTableExtractor
from adaptive_document_agent.services.report_language import translation_errors, audience_copy
from adaptive_document_agent.services.requested_tables import render_requested_tables, verify_requested_tables
from adaptive_document_agent.agent.summary_editorial import (
    SummaryEditorialDraft, editorial_parts, expand_editorial, split_editorial_copy,
)
from adaptive_document_agent.agent.company_introduction import IntroductionDraft
from adaptive_document_agent.services.summary_validation import presentation_errors
from tests.test_complete_summary_reading import generate, distinct_source
from tests.test_borderless_alignment_regressions import PositionedPage
from tests.test_report_customization import table_result, presentation


def test_long_editorial_copy_preserves_complete_claim_and_missing_continuation_is_rejected():
    result = generate(distinct_source(1))
    review = result.presentation_plan.company.summary_review
    part = next(p for p in review.parts if p.role=='content')
    _,catalog = editorial_parts([part])
    fact_id = next(iter(catalog))
    text = part.facts[0].text+' '+part.facts[0].text
    draft = SummaryEditorialDraft(summary_pages=[dict(id='long',title='Operations',source_section='Operations',
        items=[dict(fact_id=fact_id,label=part.facts[0].label,text=text)])])
    expanded = expand_editorial(draft,catalog,IntroductionDraft)
    items = [item for p in expanded.summary_pages for item in p.items]
    assert ''.join(item.text for item in items)==text and all(len(item.text)<=180 for item in items)
    assert not any('continuation' in e for e in presentation_errors(review,expanded.summary_pages,result,validate_reading=False))
    expanded.summary_pages[-1].items.pop()
    assert 'Invalid Summary continuation provenance' in presentation_errors(review,expanded.summary_pages,result,validate_reading=False)


def test_layout_split_does_not_drop_words_punctuation_or_signed_amounts():
    text = ('A qualification for an evidence-bound estimate; USD -1,234.50 thousand. '*20).strip()
    chunks = split_editorial_copy(text)
    assert ''.join(chunks)==text and all(len(c)<=180 for c in chunks)


def test_source_bound_intro_continuations_survive_preflight_and_saved_deck():
    from adaptive_document_agent.services.ppt_preflight import PresentationPreflight
    from adaptive_document_agent.services.pptx_export import _text, FOURIER_DARK
    from pptx.util import Inches
    deck=presentation();slide=deck.slides.add_slide(deck.slide_layouts[0])
    texts=["Revenue (RMB'000) was 1,234, with", 'customers who depend on completion of', 'the staged acceptance procedure.']
    for index,text in enumerate(texts):
        shape=_text(slide,text,.6,1.5+index,10,.5,size=14,color=FOURIER_DARK)
        shape.name='brief:body'
        shape.element.xpath('.//p:cNvPr')[0].set('descr','ADA_SUMMARY_ITEM_V2:'+json.dumps({'item_index':index}))
    PresentationPreflight(deck).validate_and_sanitize()
    out=BytesIO();deck.save(out)
    restored=Presentation(BytesIO(out.getvalue()))
    assert [s.text for s in restored.slides[0].shapes if s.name=='brief:body']==texts


def test_identical_continuation_wording_binds_each_distinct_source_item_once():
    from adaptive_document_agent.services.summary_export import bind_summary_pages
    from adaptive_document_agent.models.summary import SummarySlidePage,SummarySlideItem
    from adaptive_document_agent.services.presentation_brief import render_profile,BriefItem
    deck=presentation()
    items=[SummarySlideItem(label=label,text='respectively (USD in thousands).',source_quote=quote,
        source_pages=[1],part_ids=[label]) for label,quote in [('assets','Assets were 10 and 20.'),('equity','Equity was 30 and 40.')]]
    page=SummarySlidePage(id='repeated',title='Balances',source_section='Balances',items=items)
    slides=render_profile(deck,page.title,[BriefItem(i.label,i.text,i.source_pages) for i in items])
    bind_summary_pages(slides,page)
    bindings=[json.loads(s.element.xpath('.//p:cNvPr')[0].get('descr').split(':',1)[1])
              for slide in slides for s in slide.shapes if s.name=='brief:body']
    assert [b['part_ids'] for b in bindings]==[['assets'],['equity']]


class RatiosPage(PositionedPage):
    def __init__(self):
        self.words=[]
        self.row('Year ended December 31,',10,265)
        for year,center in zip(('2020','2021','2022'),(295,415,535)):
            self.word(year,30,center)
        self.row('Liquidity:',50,30)
        for y,label,values in [(70,'Current ratio(2)',('2.9','N/A','2.8')),
                               (90,'Quick ratio(3)',('2.5','N/A','1.8'))]:
            self.row(label,y,30)
            for right,value in zip((305,425,545),values):
                self.word(value,y,right-len(value)*2)
        for right,value in zip((305,425,545),('5.4','—','4.6')):
            self.word(value,110,right-len(value)*2)


def test_na_rows_and_final_unlabelled_subtotal_survive_literal_grid():
    table, = BorderlessTableExtractor().extract(RatiosPage(),1)
    assert table.raw_cells[-3:] == [['Current ratio(2)','2.9','N/A','2.8'],
        ['Quick ratio(3)','2.5','N/A','1.8'],['','5.4','—','4.6']]
    assert table.bbox[3]>=110
    assert all(h not in {'Purchase','of','items'} for h in table.headers)


def test_wrapped_left_column_label_is_not_repeated_data_column_header():
    page=RatiosPage()
    page.words=[]
    page.row('Year ended December 31,',10,265)
    for year,center in zip(('2020','2021','2022'),(295,415,535)):
        page.word(year,30,center)
    for text,y in [('Purchase',45),('of',57),('items',69),('of',81)]:
        page.row(text,y,30)
    page.row('property, plant and equipment',93,30)
    for right,value in zip((305,425,545),('10','20','30')):
        page.word(value,93,right-len(value)*2)
    page.row('Other commitments',115,30)
    for right,value in zip((305,425,545),('15','25','35')):
        page.word(value,115,right-len(value)*2)
    table,=BorderlessTableExtractor().extract(page,1)
    assert all(h not in {'Purchase','of','items'} for h in table.headers)
    assert table.column_periods[1:]==['FY2020','FY2021','FY2022']


def test_localization_inventory_includes_unit_words_and_excludes_literal_source_notes():
    from adaptive_document_agent.services.pptx_export import _text,FOURIER_DARK
    deck=presentation();slide=deck.slides.add_slide(deck.slide_layouts[0])
    for index,text in enumerate(['RMB millions','Source validation limit','Revenue was RMB 10 million.']):
        _text(slide,text,.6,1.5+index,10,.5,size=14,color=FOURIER_DARK)
    note=_text(slide,'(1) Original financial note.',.6,5,10,.5,size=14,color=FOURIER_DARK)
    note.name='customization:source_footnote:table:0'
    assert set(audience_copy(deck))=={'RMB millions','Source validation limit','Revenue was RMB 10 million.'}


def test_quality_export_trace_uses_rendered_scope_instead_of_stale_planned_pages(monkeypatch):
    from types import SimpleNamespace
    from adaptive_document_agent.services.presentation_export_trace import record_section
    from adaptive_document_agent.services.presentation_brief import BriefItem
    from adaptive_document_agent.services import presentation_scope
    result=table_result()
    slide_plan=SimpleNamespace(id='quality',slide_type='data_quality',chart_ids=[],visual_blocks=[],
        observation_ids=[],source_pages=[99],theme_id='',section_title='Limits',title='Limits')
    monkeypatch.setattr(presentation_scope,'scope_items',lambda r:[BriefItem('Boundary','Literal source caveat.',[1])])
    deck=presentation()
    with record_section(deck,slide_plan,result):
        deck.slides.add_slide(deck.slide_layouts[0])
    assert result.presentation_export_trace[-1]['source_pages']==[1]


def test_long_presentation_batches_preserve_every_slide_in_source_order(tmp_path):
    from adaptive_document_agent.services.presentation_render_batches import render_pages
    from adaptive_document_agent.services.presentation_rendering import RenderedPage,RenderingError
    deck=presentation()
    for index in range(175):
        slide=deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes[0].text=f'Source page {index+1}'
    stream=BytesIO();deck.save(stream)
    class Renderer:
        def render(self,payload,directory):
            chunk=Presentation(BytesIO(payload))
            assert len(chunk.slides)<=100
            return [RenderedPage(b'png',{'text':s.shapes[0].text}) for s in chunk.slides]
    pages=render_pages(Renderer(),stream.getvalue(),tmp_path)
    assert [p.layout['text'] for p in pages]==[f'Source page {n}' for n in range(1,176)]
    assert len(Presentation(BytesIO(stream.getvalue())).slides)==175
    class Missing(Renderer):
        def render(self,payload,directory):
            return super().render(payload,directory)[:-1]
    missing=tmp_path/'missing';missing.mkdir()
    with pytest.raises(RenderingError,match='omitted pages'):
        render_pages(Missing(),stream.getvalue(),missing)


def test_visible_source_footnotes_are_paginated_and_verified_after_reload():
    result=table_result()
    table=next(t for p in result.document.pages for t in p.raw_tables)
    table.raw_footnotes=['(1) The exchange assumption is USD1.00 to EUR0.92526. '
                         'This illustrative assumption is not a promise of convertibility. '*25,
                        '(2) Values exclude subsequent transactions.']
    deck=presentation()
    render_requested_tables(deck,result)
    out=BytesIO();deck.save(out)
    restored=Presentation(BytesIO(out.getvalue()))
    assert verify_requested_tables(restored,result)
    notes=[s.text for slide in restored.slides for s in slide.shapes
           if s.name.startswith('customization:source_footnote:')]
    assert len(notes)>1 and ''.join(notes)=='\n\n'.join(table.raw_footnotes)
    assert not any('exchange assumption' in s for s in audience_copy(restored))
    shape=next(s for slide in restored.slides for s in slide.shapes if s.name.startswith('customization:source_footnote:'))
    shape.text=shape.text.replace('0.92526','0.92500')
    with pytest.raises(ValueError,match='footnotes were omitted or changed'):
        verify_requested_tables(restored,result)


def test_page_list_punctuation_is_not_a_numeric_spelling_and_debt_is_not_liabilities():
    assert not translation_errors('Source: Document disclosures (p. 31, 33, 47)', '来源：文件披露（第31、33、47页）')
    assert translation_errors('Source: Document disclosures (p. 31, 33, 47)', '来源：文件披露（第31、34、47页）')
    assert not translation_errors('Total indebtedness RMB 87.1m','总债务人民币87.1百万元')
    assert translation_errors('Total indebtedness RMB 87.1m','总负债人民币87.1百万元')
    assert not translation_errors('Total liabilities RMB 342.2m','总负债人民币342.2百万元')


def test_literal_shared_narrative_date_supports_chinese_month_day_without_extra_number():
    from adaptive_document_agent.models.summary import SummaryFact
    from adaptive_document_agent.services.summary_validation import fact_errors
    from tests.test_executive_brief import result_for
    text='The six months ended June 30, 2023 and 2024 had 12 and 14 active customers, respectively.'
    result=result_for([text])
    fact=SummaryFact(label='Customers',text='截至2023年及2024年6月30日的六个月，活跃客户分别为12及14。',source_quote=text,source_pages=[1])
    assert not fact_errors(fact,result,text,1)
    fact.text=fact.text.replace('6月30日','6月29日')
    assert fact_errors(fact,result,text,1)


def test_explicit_percentage_magnitude_keeps_source_sign_without_authorizing_unannotated_copy():
    from tests.test_brief_table_context import table_result
    from adaptive_document_agent.services.executive_brief import validate_executive_brief
    from adaptive_document_agent.models.executive_brief import BriefQuantityRepresentation
    result,brief=table_result()
    table=result.document.pages[0].tables[0]
    table.column_types=['label','percentage','percentage']
    table.headers=['label','2023 %','2024 %']
    table.raw_header_lines=['USD in thousands','2023 2024','% of Revenue']
    table.rows[0].cells=['Loss','(20.0)','(30.8)']
    result.document.pages[0].text='\n'.join(table.raw_header_lines)+'\nLoss (20.0) (30.8)'
    brief.items[0].evidence[0].text='Loss (20.0) (30.8)'
    brief.items[0].text='Loss represented 30.8%.'
    assert validate_executive_brief(brief,result)
    brief.items[0].quantity_representations=[BriefQuantityRepresentation(quantity_text='30.8%',source_value='-30.8',representation='absolute_magnitude')]
    assert not validate_executive_brief(brief,result)


def test_distinct_literal_unit_header_scopes_repeated_loss_rows_to_quoted_grid():
    from tests.test_brief_table_context import table_result
    from adaptive_document_agent.models.table import TableRow
    from adaptive_document_agent.models.executive_brief import BriefQuantityRepresentation
    from adaptive_document_agent.services.executive_brief import validate_executive_brief
    result,brief=table_result()
    table=result.document.pages[0].tables[0]
    table.headers=['label','Amount','% of revenue','Amount','% of revenue']
    table.column_types=['label','amount','percentage','amount','percentage']
    table.column_currencies=[None,'USD',None,'USD',None]
    table.column_scales=[None,1000,1,1000,1]
    table.rows=[TableRow(cells=['Loss','(100)','(10.0)','(200)','(20.0)'],page=1)]
    table.raw_header_lines=['2023 2024','(USDinthousands,exceptforpercentages)']
    other=table.model_copy(deep=True);other.table_id='reconciliation'
    other.rows[0].cells=['Loss','(100)','(200)']
    other.column_types=['label','amount','amount'];other.column_currencies=[None,'USD','USD'];other.column_scales=[None,1000,1000]
    other.raw_header_lines=['2023 2024','(USD in thousands)']
    first='2023 2024\n(USD in thousands, except for percentages)\nLoss (100) (10.0) (200) (20.0)'
    result.document.pages[0].tables.append(other)
    result.document.pages[0].text=first+'\n2023 2024\n(USD in thousands)\nLoss (100) (200)'
    brief.items[0].evidence[0].text=first
    brief.items[0].text='Loss was USD 200 thousand.'
    brief.items[0].quantity_representations=[BriefQuantityRepresentation(quantity_text='USD 200 thousand',source_value='-200',representation='absolute_magnitude')]
    assert not validate_executive_brief(brief,result)


def test_actual_pdf_font_evidence_preserves_footnotes_without_following_body():
    from reportlab.pdfgen.canvas import Canvas
    from adaptive_document_agent.extraction.table_extractor import TableExtractor
    stream=BytesIO();canvas=Canvas(stream,pagesize=(600,800))
    canvas.setFont('Helvetica',10)
    canvas.drawString(30,760,'Document context for a comparison table and its definitions.')
    canvas.drawString(250,720,'Year ended December 31,')
    for year,x in zip(('2020','2021','2022'),(280,380,480)):canvas.drawString(x,700,year)
    for label,y,values in [('Measure A',680,('10','20','30')),('Measure B',660,('15','25','35'))]:
        canvas.drawString(30,y,label)
        for x,value in zip((290,390,490),values):canvas.drawRightString(x,y,value)
    canvas.setFont('Helvetica',8)
    canvas.drawString(30,635,'(1) The figures exclude the stated illustrative adjustment.')
    canvas.drawString(50,624,'No representation of future results is made.')
    canvas.setFont('Helvetica',10)
    canvas.drawString(30,596,'Following narrative is not a table annotation.')
    canvas.save()
    extractor=TableExtractor();extractor.extract(stream.getvalue())
    table,=extractor.raw_tables_by_page[1]
    assert table.raw_footnotes==['(1) The figures exclude the stated illustrative adjustment.\nNo representation of future results is made.']
