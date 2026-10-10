"""Complete claims, bounded review and observable timing regressions."""
import json
from io import BytesIO
from threading import Barrier
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx.util import Inches

from adaptive_document_agent.agent.brief_meaning_review import meaning_validator
from adaptive_document_agent.agent.report_localization import prepare_report_language
from adaptive_document_agent.models.customization import ReportRequirement, ReportRequirements
from adaptive_document_agent.models.executive_brief import ExecutiveBrief
from adaptive_document_agent.services import report_language
from adaptive_document_agent.services.report_language import translation_errors
from adaptive_document_agent.services.translation_atoms import protect_quantities, restore_quantities
from tests.test_executive_brief import gateway, result_for
from tests.test_report_customization import presentation


@pytest.mark.parametrize('source,target', [
    ('As of December 31, 2021, 2022 and 2023', '截至2021年、2022年及2023年12月31日'),
    ('Six months ended June 30, 2023 and 2024', '截至2023年及2024年6月30日止六个月'),
    ('Loss (20.0%) and revenue 100', '亏损（20.0%）与收入100'),
    ('Year ended December 31: 2021, 2023', '截至12月31日止年度：2021、2023'),
])
def test_shared_dates_and_chinese_parentheses_keep_every_literal_coordinate(source, target):
    assert not translation_errors(source, target)
    assert translation_errors(source, target.replace('2023', '2022').replace('20.0', '21.0').replace('100', '101'))


def test_month_precision_and_negative_sign_cannot_be_erased():
    assert translation_errors('June 2024 revenue 10', '2024年6月30日收入10')
    assert translation_errors('Loss (20.0%)', '亏损20.0%')
    assert translation_errors('As of December 31, 2021, 2022 and 2023', '截至2023年12月31日')
    assert translation_errors('Year ended December 31', '截至2024年12月31日止年度')


@pytest.mark.parametrize('source', ['RMB39.1 million', 'HK$42.2 million', 'US$1,039.5 million',
                                   'RMB -1,234.50 thousand'])
def test_protected_quantities_keep_source_coefficient_and_localize_units(source):
    masked, atoms = protect_quantities(source, 'zh')
    translated = restore_quantities(masked, atoms)
    assert atoms and not translation_errors(source, translated)
    assert not any(word in translated for word in ('million', 'thousand'))
    with pytest.raises(ValueError, match='protected quantity'):
        restore_quantities(masked + masked, atoms)
    with pytest.raises(ValueError, match='protected quantity'):
        restore_quantities('', atoms)


def test_adjacent_chinese_prose_does_not_hide_explicit_monies_or_weaken_scale_validation():
    source = '研发开支由2021年的RMB46.9 million增加至2023年的RMB70.5 million。'
    masked, atoms = protect_quantities(source, 'zh')
    assert len(atoms) == 2
    text = restore_quantities(masked, atoms)
    assert not translation_errors(source, text) and 'million' not in text
    assert translation_errors(source, text.replace('46.9百万元','4690万元'))
    _, atoms = protect_quantities('RMB10 millionaires', 'zh')
    assert not atoms


def test_false_reviewer_approval_cannot_keep_english_monetary_unit_words():
    from adaptive_document_agent.agent.report_localization import _translate_batch
    g, _ = gateway([dict(items=[dict(id=0,text='收入RMB10 million。')]),
                   dict(items=[dict(id=0,accepted=True,reason='Incorrect fixture approval')])])
    accepted, failed, audit = _translate_batch(g,[dict(id=0,text='Revenue RMB10 million.')],'zh',presentation())
    assert not accepted and failed == [0] and 'untranslated' in str(audit['errors'])


def test_geometry_inventory_is_built_once_for_many_translations(monkeypatch):
    deck = presentation(); slide = deck.slides.add_slide(deck.slide_layouts[0])
    for i in range(20):
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(1)).text = f'Claim {i}'
    original = report_language.text_frames
    calls = []
    def frames(deck):
        calls.append(1)
        yield from original(deck)
    monkeypatch.setattr(report_language, 'text_frames', frames)
    for i in range(20):
        assert not report_language.copy_fit_errors(deck, f'Claim {i}', f'事实{i}')
    assert len(calls) == 1


def test_native_pagination_wrappers_compose_only_exact_checked_base_copy():
    deck = presentation(); slide = deck.slides.add_slide(deck.slide_layouts[0])
    heading = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(1))
    heading.name = 'brief:heading'; heading.text = '02  Product range'
    title = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(10), Inches(1))
    title.text = 'Company overview (continued)'
    result = result_for(['Company source.'])
    result.profile.report_requirements = ReportRequirements(items=[ReportRequirement(id='zh',
        kind='output_language',request_quote='中文',description='中文',language='zh')],
        copy_translations={'Product range':'产品范围','Company overview':'公司介绍'})
    report_language.apply_report_language(deck,result)
    assert heading.text == '02  产品范围' and title.text == '公司介绍（续）'
    assert result.profile.report_requirements.copy_translations['02  Product range'] == '02  产品范围'
    assert translation_errors('02 Product range','03 产品范围')


def test_whole_continuations_localize_before_native_layout_and_keep_source_bindings():
    from adaptive_document_agent.agent.summary_editorial import expand_editorial, editorial_parts, SummaryEditorialDraft
    from adaptive_document_agent.agent.company_introduction import IntroductionDraft
    from adaptive_document_agent.services.localized_summary_copy import (
        summary_copy_units, retain_summary_translations, verify_summary_translations)
    from adaptive_document_agent.services.pptx_export import _add_company_at_a_glance
    from adaptive_document_agent.services.summary_export import verify_summary_export
    from tests.test_complete_summary_reading import generate, distinct_source
    result = generate(distinct_source(1))
    company = result.presentation_plan.company
    part = next(p for p in company.summary_review.parts if p.role == 'content')
    _, catalog = editorial_parts([part]); fid = next(iter(catalog))
    text = part.facts[0].text * 3
    draft = SummaryEditorialDraft(summary_pages=[dict(id='long', title='Operations', source_section='Operations',
        items=[dict(fact_id=fid, label=part.facts[0].label, text=text)])])
    expanded = expand_editorial(draft, catalog, IntroductionDraft)
    # Keep the other substantive part; this test exercises every source locator.
    company.summary_pages = [*expanded.summary_pages, company.summary_pages[0].model_copy(
        update={'id':'other', 'items':company.summary_pages[0].items[1:]})]
    result.profile.report_requirements = ReportRequirements(items=[ReportRequirement(id='zh',
        kind='output_language', request_quote='中文', description='中文', language='zh')])
    units = summary_copy_units(result)
    assert text in {u['source'] for u in units}
    accepted = {u['source']: u['source'] for u in units}
    result.profile.report_requirements.copy_translations = accepted
    retain_summary_translations(result, units, accepted)
    verify_summary_translations(result)
    deck = presentation()
    _add_company_at_a_glance(deck, result, result.presentation_plan.slides[0])
    out = BytesIO(); deck.save(out)
    assert verify_summary_export(Presentation(BytesIO(out.getvalue())), result)
    key = next(iter(result.profile.report_requirements.summary_copy_translations))
    result.profile.report_requirements.summary_copy_translations[key] += ' fabricated'
    with pytest.raises(ValueError, match='complete checked claim'):
        verify_summary_translations(result)


def test_localized_claim_layout_prefers_clause_boundaries_over_cutting_names_and_amounts():
    from adaptive_document_agent.services.localized_summary_copy import retain_summary_translations
    result = result_for(['Source.'])
    result.profile.report_requirements = ReportRequirements()
    text = 'Magician Lite 系列覆盖制造应用；人民币1,234.50百万元，2024年6月30日。'
    units = [{'source':text,'slots':[(i,f'page:{i}') for i in range(3)]}]
    retain_summary_translations(result,units,{text:text})
    pieces = list(result.profile.report_requirements.summary_copy_translations.values())
    assert ''.join(pieces)==text
    for literal in ('Magician Lite','1,234.50','2024年6月30日'):
        assert any(literal in piece for piece in pieces)


def test_meaning_references_are_item_scoped_and_still_check_arithmetic():
    source = 'Share was 2% in 2023. Share was 3% in 2024.'
    result = result_for([source])
    brief = ExecutiveBrief(title='Shares', items=[dict(label='Share', text='Share rose from 2% to 3%.',
        evidence=[dict(page=1, text=source)])])
    verdict = dict(index=0, accepted=True, reason='Exact comparison', numeric_comparison=True,
        comparisons=[dict(claim=brief.items[0].text, start_value='2%', end_value='3%', direction='increase',
                          start_ref='i0_q0', end_ref='i0_q0', start_duration_months=12, end_duration_months=12)])
    g, _ = gateway([dict(items=[verdict])])
    assert not meaning_validator(g, result=result)(brief)
    verdict['comparisons'][0]['start_ref'] = 'i1_q0'
    g, _ = gateway([dict(items=[verdict])])
    assert 'outside' in str(meaning_validator(g, result=result)(brief))
    verdict['comparisons'][0]['start_ref'] = 'i0_q0'
    verdict['comparisons'][0]['direction'] = 'decrease'
    g, _ = gateway([dict(items=[verdict])])
    assert 'direction' in str(meaning_validator(g, result=result)(brief))


def test_meaning_format_failure_has_one_correction_and_memoizes_only_failed_claims():
    source = 'Share was 2% in 2023. Share was 3% in 2024.'
    result = result_for([source])
    brief = ExecutiveBrief(title='Shares', items=[dict(label='Share', text='Share rose from 2% to 3%.',
        evidence=[dict(page=1, text=source)])])
    g, client = gateway(['invalid JSON', 'invalid JSON'])
    validate = meaning_validator(g, result=result)
    assert 'unavailable' in str(validate(brief))
    assert 'unavailable' in str(validate(brief)) and len(client.calls) == 2
    assert len([w for w in result.validation_warnings if w.code.endswith('format_audit')]) == 2
    assert 'executive_brief_meaning_review' in result.stage_details_ms


def test_failed_patch_meaning_does_not_erase_verified_findings_or_relax_topic_gate():
    from adaptive_document_agent.agent.brief_item_repair import _salvage
    from adaptive_document_agent.models.executive_brief import ExecutiveBriefItem
    result = result_for(['Customers 10. Customers 20. Customers 30.'])
    items = [ExecutiveBriefItem(label='Customers', text=f'Customers {n}.',
             evidence=[dict(page=1,text=f'Customers {n}.')]) for n in (10,20,30)]
    values = [i.model_dump() for i in items]
    kwargs = dict(result=result, excerpts={1:result.document.pages[0].text},
                  semantic_validator=lambda brief: {2:['Review unavailable']})
    retained, discarded = _salvage(values, {0:items[0],1:items[1]}, 'Customers', topics=[], **kwargs)
    assert retained.items == items[:2] and discarded == [2]
    retained, _ = _salvage(values, {0:items[0],1:items[1]}, 'Customers',
                           topics=[('A',[1]),('B',[2]),('C',[3])], **kwargs)
    assert retained is None


def test_localization_progress_and_subtimings_survive_result_serialization(monkeypatch):
    from adaptive_document_agent.services import pptx_export
    from adaptive_document_agent.models import PipelineResult
    result = result_for(['Revenue RMB 10 million in FY2025'])
    result.profile.report_requirements = ReportRequirements(items=[ReportRequirement(id='zh',
        kind='output_language', request_quote='中文', description='中文', language='zh')])
    deck = presentation(); slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(1)).text = result.document.pages[0].text
    monkeypatch.setattr(pptx_export, 'presentation_for_copy', lambda _: deck)
    g, _ = gateway([dict(items=[dict(id=0, text='FY2025收入⟦Q0⟧')]),
                   dict(items=[dict(id=0, accepted=True, reason='Faithful Chinese prose')])])
    updates = []
    prepare_report_language(g, result, progress=updates.append)
    assert result.customization_report[0].status == 'planned'
    assert any('(1/1;' in update for update in updates)
    stored = PipelineResult.model_validate_json(result.model_dump_json())
    assert {'localization_copy_inventory', 'localization_translation_requests', 'localization_review_requests',
            'localization_copy_fit', 'localization_validation'} <= stored.stage_details_ms.keys()


def test_customization_review_wall_time_is_in_exported_pipeline_result(monkeypatch):
    import pymupdf
    from adaptive_document_agent.agent import orchestrator, report_requirements, document_discovery
    from adaptive_document_agent.models import DocumentProfile, PipelineResult
    doc = pymupdf.open(); doc.new_page().insert_text((72,72), 'Plain source narrative.')
    raw = doc.tobytes(); doc.close()
    monkeypatch.setattr(document_discovery.DocumentDiscovery, 'discover', lambda *a, **k: DocumentProfile())
    monkeypatch.setattr(report_requirements, 'interpret_requirements', lambda *a, **k: ReportRequirements(
        items=[ReportRequirement(id='detail', kind='content_detail', request_quote='Detail', description='Detail')]))
    result = orchestrator.DocumentOrchestrator().analyse_pdf(raw, analysis_focus='Detail')
    stored = PipelineResult.model_validate_json(result.model_dump_json())
    assert {'customization_review', 'background_drain', 'source_chart_inventory'} <= stored.timings_ms.keys()
    assert {'content_requirement_review', 'report_localization', 'customization_checks'} <= stored.stage_details_ms.keys()
    assert sum(stored.timings_ms.values()) <= stored.pipeline_total_ms


def test_editorial_batches_run_concurrently_and_merge_in_source_order(monkeypatch):
    from adaptive_document_agent.agent import summary_presentation
    from tests.test_complete_summary_reading import generate, distinct_source
    original = summary_presentation._generate_summary_batch
    rendezvous = Barrier(2, timeout=20)
    monkeypatch.setattr(summary_presentation, 'planning_batches',
        lambda parts: [parts[:len(parts)//2], parts[len(parts)//2:]])
    def edit(*args, **kwargs):
        rendezvous.wait()
        return original(*args, **kwargs)
    monkeypatch.setattr(summary_presentation, '_generate_summary_batch', edit)
    result = generate(distinct_source(2), workers=2)
    review = result.presentation_plan.company.summary_review
    assert review.status == 'complete'
    assert {'reading', 'editorial', 'editorial_batch_1', 'editorial_batch_2'} <= review.timings_ms.keys()
    assert [p.source_page for p in review.parts if p.role == 'content'] == [1,1,2,2]


def test_failed_translation_requests_are_timed_and_already_chinese_copy_is_immutable(monkeypatch):
    from adaptive_document_agent.services import pptx_export
    result = result_for(['Source prose.'])
    result.profile.report_requirements = ReportRequirements(items=[ReportRequirement(id='zh',
        kind='output_language', request_quote='中文', description='中文', language='zh')])
    deck = presentation(); slide = deck.slides.add_slide(deck.slide_layouts[0])
    for text in ('收入为100千元。', 'Source prose.'):
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(1)).text = text
    monkeypatch.setattr(pptx_export, 'presentation_for_copy', lambda _: deck)
    g, client = gateway(['invalid JSON', 'invalid JSON'])
    prepare_report_language(g, result)
    assert result.profile.report_requirements.copy_translations == {'收入为100千元。': '收入为100千元。'}
    assert len(client.calls) == 2 and 'localization_translation_requests' in result.stage_details_ms
    audit = json.loads(next(w.message for w in result.validation_warnings if w.code=='report_localization_audit'))
    assert audit['unchanged_target_language_count'] == 1 and audit['failed_ids']


def test_one_invalid_protected_quantity_does_not_displace_verified_translation():
    from adaptive_document_agent.agent.report_localization import _translate_batch
    deck = presentation()
    g, _ = gateway([dict(items=[dict(id=0, text='收入⟦Q0⟧'), dict(id=1, text='收入1000万元')]),
                   dict(items=[dict(id=0, accepted=True, reason='Exact protected value'),
                               dict(id=1, accepted=True, reason='Incorrect fixture approval')])])
    accepted, failed, _ = _translate_batch(g, [dict(id=0,text='Revenue RMB10 million'),
                                             dict(id=1,text='Revenue RMB20 million')], 'zh', deck)
    assert accepted == {'Revenue RMB10 million': '收入人民币10百万元'} and failed == [1]


@pytest.mark.parametrize('setting', ['reduced', 'provider_default'])
def test_summary_editorial_reasoning_policy_preserves_reading_and_meaning_review(setting):
    from adaptive_document_agent.services.llm import LLMSettings, ProviderName
    from adaptive_document_agent.services.llm.reasoning_policy import request_reasoning_policy
    settings = LLMSettings(provider=ProviderName.DEEPSEEK, simple_task_reasoning=setting)
    policy = request_reasoning_policy(settings, stage='presentation', operation='SummaryEditorialDraft', model='deepseek-flash')
    assert bool(policy['options']) == (setting == 'reduced')
    for stage, operation in [('presentation', 'SummaryReadBatch'), ('report', 'BriefMeaningReview')]:
        assert not request_reasoning_policy(settings, stage=stage, operation=operation, model='deepseek-flash')['options']
