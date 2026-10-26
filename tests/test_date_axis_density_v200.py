"""Date labels reserve physical width without changing reported time spacing."""
from datetime import date
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt
import pytest

from adaptive_document_agent.services.presentation_axes import style_date_axis
from tests.test_presentation_editorial_review import blank_deck


@pytest.mark.parametrize('width,font,expected_min', [(11.7, 11, 6), (6, 11, 12), (11.7, 24, 12)])
def test_tick_density_accounts_for_date_format_and_font(width, font, expected_min):
    deck = blank_deck()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    dates = [date(2021, 12, 31), date(2022, 12, 31), date(2023, 12, 31), date(2024, 6, 30)]
    data = CategoryChartData()
    data.categories = dates
    data.add_series('Reported quantity', [71, 132, 142, 155])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(.5), Inches(1),
                                  Inches(width), Inches(3), data).chart
    chart.category_axis.tick_labels.font.size = Pt(font)
    before = list(chart.series[0].values)
    style_date_axis(chart, dates, width)
    assert int(chart.category_axis._element.xpath('c:majorUnit')[0].get('val')) >= expected_min
    assert chart.category_axis._element.xpath('c:majorTimeUnit')[0].get('val') == 'months'
    assert chart.category_axis._element.scaling.xpath('c:min')[0].get('val') == str(
        (dates[0] - date(1899, 12, 30)).days)
    assert list(chart.series[0].values) == before
