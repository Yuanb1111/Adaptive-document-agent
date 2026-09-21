from __future__ import annotations

import io
from pptx import Presentation
import pytest

pytestmark = pytest.mark.usefixtures("local_render_stub")

from adaptive_document_agent.document_model.metric_semantic_classifier import (
    classify_metric,
    format_metric_change,
    format_metric_display_value,
    sanitize_metric_label,
)
from adaptive_document_agent.document_model.period_semantic_validator import format_period_label
from adaptive_document_agent.models import (
    ChartPlan,
    CompanyFact,
    CompanyProfile,
    DocumentProfile,
    Observation,
    ParsedDocument,
    PipelineResult,
    PresentationPlan,
    PresentationSlide,
    SourceEvidence,
)
from adaptive_document_agent.services.export import export_pptx
from adaptive_document_agent.services.ppt_preflight import PresentationPreflight
from adaptive_document_agent.services.pptx_export import _classify_financial_theme, _sanitize_investor_narrative
from adaptive_document_agent.validation.presentation_plan_validator import PresentationPlanValidator


def _build_test_pipeline_result(
    *,
    company_name: str = "Pioneer Robotics Ltd.",
    has_facts: bool = True,
    slide_title: str = "Current Ratio Declined from 0.4x to 0.2x",
) -> PipelineResult:
    ev1 = SourceEvidence(page=45, text="1200", extraction_method="digital_table", confidence=0.95)
    ev2 = SourceEvidence(page=48, text="350", extraction_method="digital_table", confidence=0.92)

    obs = [
        Observation(
            id="rev_2022",
            metric_original="Revenue",
            value=1000.0,
            raw_value="1,000",
            raw_unit="RMB '000",
            currency="RMB",
            unit="currency",
            period="FY2022",
            evidence=[ev1],
            confidence=0.95,
        ),
        Observation(
            id="rev_2023",
            metric_original="Revenue",
            value=1500.0,
            raw_value="1,500",
            raw_unit="RMB '000",
            currency="RMB",
            unit="currency",
            period="FY2023",
            evidence=[ev1],
            confidence=0.95,
        ),
        Observation(
            id="cr_2022",
            metric_original="Current ratio",
            value=0.4,
            raw_value="0.4",
            raw_unit="",
            unit="multiple",
            period="31 Dec 2022",
            evidence=[ev2],
            confidence=0.92,
        ),
        Observation(
            id="cr_2023",
            metric_original="Current ratio",
            value=0.2,
            raw_value="0.2",
            raw_unit="",
            unit="multiple",
            period="31 Dec 2023",
            evidence=[ev2],
            confidence=0.92,
        ),
        Observation(
            id="cr_2024_interim",
            metric_original="Current ratio",
            value=0.2,
            raw_value="0.2",
            raw_unit="",
            unit="multiple",
            period="31 May 2025*",
            evidence=[ev2],
            confidence=0.90,
        ),
    ]

    c1 = ChartPlan(
        id="chart_rev",
        title="Revenue Trajectory",
        chart_type="bar",
        question="How did revenue grow?",
        observation_ids=["rev_2022", "rev_2023"],
        source_pages=[45],
    )
    c2 = ChartPlan(
        id="chart_cr",
        title="Current Ratio Contraction",
        chart_type="line",
        question="How did liquidity evolve?",
        observation_ids=["cr_2022", "cr_2023", "cr_2024_interim"],
        source_pages=[48],
    )

    facts = [
        CompanyFact(label="Scale", value="Leading manufacturer of mobile robots", source_pages=[1]),
        CompanyFact(label="Clients", value="Global active enterprise client network", source_pages=[2]),
    ] if has_facts else []

    plan = PresentationPlan(
        title="Institutional Financial Assessment",
        company=CompanyProfile(
            name=company_name,
            one_line_description="Leading developer of autonomous mobile robotics.",
            business_model="Turnkey robotics hardware and software solutions.",
            key_facts=facts,
            stock_code="9988.HK",
            offering_type="Main Board IPO",
            reporting_currency="RMB",
            source_pages=[1, 2],
        ),
        slides=[
            PresentationSlide(id="cover", slide_type="cover", title="Institutional Financial Assessment"),
            PresentationSlide(id="overview", slide_type="company_overview", title="Company at a Glance"),
            PresentationSlide(
                id="summary",
                slide_type="executive_summary",
                title="Executive Summary",
                observation_ids=["rev_2022", "rev_2023", "cr_2022", "cr_2023"],
                source_pages=[45, 48],
                bullets=[
                    "Revenue expanded from RMB 1,000k to RMB 1,500k driven by warehouse robotics demand.",
                    "Current ratio contracted from 0.4x to 0.2x reflecting working capital movements.",
                    "Top-line growth trajectory was maintained across the track record period.",
                    "Disclosures provide grounded basis for liquidity and performance assessment.",
                ],
            ),
            PresentationSlide(
                id="analysis_multichart",
                slide_type="analysis",
                title=slide_title,
                section_id="liquidity_analysis",
                section_title="Liquidity & Working Capital",
                message="Current ratio weakened while top-line growth continued.",
                chart_ids=["chart_rev", "chart_cr"],
                source_pages=[45, 48],
            ),
            PresentationSlide(id="quality", slide_type="data_quality", title="Data Quality & Methodology"),
            PresentationSlide(id="appendix", slide_type="appendix", title="Key data appendix"),
        ],
    )

    return PipelineResult(
        document=ParsedDocument(document_id="doc_pioneer", sha256="hash123", safe_filename="prospectus.pdf", page_count=350),
        profile=DocumentProfile(document_type="Prospectus", document_summary="Pioneer Robotics Prospectus"),
        observations=obs,
        charts=[c1, c2],
        presentation_plan=plan,
    )


def test_standard_deck_order_and_thank_you_slide() -> None:
    result = _build_test_pipeline_result()
    deck = Presentation(io.BytesIO(export_pptx(result)))

    slide_titles = [
        next((s.text.strip() for s in slide.shapes if s.has_text_frame and s.text.strip()), "")
        for slide in deck.slides
    ]

    assert slide_titles[0] == "Institutional Financial Assessment"
    assert slide_titles[1] == "Contents"
    assert slide_titles[2] == "Company at a Glance"
    assert slide_titles[3] == "Executive Summary"

    contents_slide = deck.slides[1]
    contents_text = " ".join(s.text for s in contents_slide.shapes if s.has_text_frame)
    assert "Liquidity & Working Capital" in contents_text
    assert "Thank You" not in contents_text
    assert "THANK YOU" not in contents_text

    thank_you_slide = deck.slides[-1]
    ty_text = " ".join(s.text for s in thank_you_slide.shapes if s.has_text_frame)
    assert "THANK YOU" in ty_text

    cover_slide = deck.slides[0]
    cover_nums = [s.text for s in cover_slide.shapes if s.has_text_frame and s.text.strip().isdigit()]
    assert len(cover_nums) == 0

    ty_nums = [s.text for s in thank_you_slide.shapes if s.has_text_frame and s.text.strip().isdigit()]
    assert len(ty_nums) == 0


def test_company_at_a_glance_empty_facts_full_width() -> None:
    result = _build_test_pipeline_result(has_facts=False)
    deck = Presentation(io.BytesIO(export_pptx(result)))

    glance_slide = deck.slides[2]
    glance_text = " ".join(s.text for s in glance_slide.shapes if s.has_text_frame)

    assert "KEY FACTS" not in glance_text
    assert "Pioneer Robotics Ltd." in glance_text
    assert "BUSINESS FOCUS" in glance_text
    assert "Stock Code: 9988.HK" in glance_text
    assert "Offering: Main Board IPO" in glance_text
    assert "Currency: RMB" in glance_text

    panels = [s for s in glance_slide.shapes if not getattr(s, "text", "").strip() and not s.has_table and not s.has_chart]
    full_width_panels = [p for p in panels if abs(p.width.inches - 11.70) < 0.15]
    assert len(full_width_panels) >= 1


def test_company_at_a_glance_missing_identity_fallback() -> None:
    result = _build_test_pipeline_result(company_name="", has_facts=False)
    deck = Presentation(io.BytesIO(export_pptx(result)))

    glance_slide = deck.slides[2]
    glance_text = " ".join(s.text for s in glance_slide.shapes if s.has_text_frame)
    assert "Document at a Glance" in glance_text


def test_multi_chart_slide_has_distinct_kpi_card_titles() -> None:
    result = _build_test_pipeline_result(slide_title="Working Capital & Liquidity Divergence")
    deck = Presentation(io.BytesIO(export_pptx(result)))

    analysis_slides = [
        s for s in deck.slides
        if any(shape.has_chart for shape in s.shapes)
    ]
    assert len(analysis_slides) >= 1
    slide = analysis_slides[0]

    charts = [s for s in slide.shapes if s.has_chart]
    assert len(charts) == 2

    slide_title = next(s.text.strip() for s in slide.shapes if s.has_text_frame and s.text.strip())
    card_title_shapes = [
        s for s in slide.shapes
        if s.has_text_frame and s.text.strip() and s.text.strip() != slide_title and ("Revenue" in s.text or "Current" in s.text)
    ]
    card_titles = [s.text.strip() for s in card_title_shapes]
    assert len(set(card_titles)) == len(card_titles)
    for title in card_titles:
        assert title != slide_title


def test_tesla_style_wide_appendix_financial_table() -> None:
    result = _build_test_pipeline_result()
    deck = Presentation(io.BytesIO(export_pptx(result)))

    # Flow and balance-sheet periods are deliberately split into compatible
    # appendix tables. Inspect every appendix table before the Thank You slide.
    appendix_slides = list(deck.slides)[:-1]
    tables = [shape.table for slide in appendix_slides for shape in slide.shapes if shape.has_table]
    assert tables
    table = tables[0]

    headers = [table.cell(0, col).text for col in range(len(table.columns))]
    assert headers[0] == "Financial Metric"
    assert headers[1] == "Unit"
    assert any("2022" in h for h in headers[2:])
    assert any("2023" in h for h in headers[2:])

    row_0_texts = [
        table.cell(r, 0).text
        for table in tables
        for r in range(len(table.rows))
    ]
    assert any("FINANCIAL PERFORMANCE" in t for t in row_0_texts)
    assert any("LIQUIDITY" in t for t in row_0_texts)

    slide_text = " ".join(
        shape.text
        for slide in appendix_slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "Source: Document disclosures (p." in slide_text
    assert "* Unaudited" in slide_text
    assert "Complete reported dataset available in accompanying CSV export." in slide_text


def test_negative_value_language_sanitization() -> None:
    sem_liab = classify_metric("Total liabilities", value=200.0, raw_unit="RMB", unit="currency")
    text_liab = format_metric_change(150.0, 200.0, 1.0, sem_liab)
    assert "widened" in text_liab.lower() or "increased" in text_liab.lower()

    text_narrowed = format_metric_change(200.0, 150.0, 1.0, sem_liab)
    assert "decreased" in text_narrowed.lower() or "narrowed" in text_narrowed.lower() or "contracted" in text_narrowed.lower()

    narrative_in = "Net loss widened by -12.5% while revenue increased by -5.2% and net profit was -30.0m."
    sanitized = _sanitize_investor_narrative(narrative_in)
    assert "widened by -" not in sanitized
    assert "increased by -" not in sanitized
    assert "net profit was -" not in sanitized
    assert "net loss narrowed by 12.5%" in sanitized
    assert "decreased by 5.2%" in sanitized
    assert "net loss was 30.0m" in sanitized


def test_ppt_preflight_checks() -> None:
    result = _build_test_pipeline_result()
    pptx_bytes = export_pptx(result)
    deck = Presentation(io.BytesIO(pptx_bytes))

    preflight = PresentationPreflight(deck)
    issues = preflight.validate_and_sanitize()
    errors = [i for i in issues if i.severity == "error"]
    assert len(errors) == 0
