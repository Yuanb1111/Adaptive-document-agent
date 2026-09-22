"""End-to-end regression tests proving that the final PPTX output and presentation layers
use the unified FinancialMovementFormatter and canonical period formatting, render negative charts visibly,
and paginate Key Findings without row overlapping.
"""

from __future__ import annotations

import io
import xml.etree.ElementTree as ET
import zipfile

from pptx import Presentation
from pptx.util import Inches

from adaptive_document_agent.document_model.period_semantic_validator import format_canonical_period, format_observation_period
from adaptive_document_agent.models import (
    ChartPlan,
    DocumentProfile,
    Observation,
    ParsedDocument,
    PipelineResult,
    PresentationPlan,
    PresentationSlide,
    SourceEvidence,
)
from adaptive_document_agent.services.export import export_pptx
from adaptive_document_agent.services.movement_formatter import FinancialMovementFormatter


# ---------------------------------------------------------------------------
# 1. FinancialMovementFormatter Semantic Tests
# ---------------------------------------------------------------------------

def test_movement_rd_expense_decreased() -> None:
    """R&D expense -437 -> -282 must say 'decreased by RMB155m', NEVER 'Loss narrowed'."""
    res = FinancialMovementFormatter.format_movement("R&D expenses", -437_000_000.0, -282_000_000.0)
    assert "decreased" in res.casefold()
    assert "loss" not in res.casefold()
    assert "155m" in res.casefold()


def test_movement_selling_expense_decreased_slightly() -> None:
    """Selling & marketing expense -456 -> -446 must say 'decreased slightly' or 'decreased', NEVER 'Loss narrowed'."""
    res = FinancialMovementFormatter.format_movement("Selling & marketing expenses", -456_000_000.0, -446_000_000.0)
    assert "decreased" in res.casefold()
    assert "loss" not in res.casefold()


def test_movement_operating_cash_flow_narrowed() -> None:
    """Operating cash flow -649 -> -108 must say 'Operating cash outflow narrowed by RMB541m'."""
    res = FinancialMovementFormatter.format_movement("Operating cash flow", -649_000_000.0, -108_000_000.0)
    assert "outflow narrowed" in res.casefold()
    assert "541m" in res.casefold()
    assert "loss" not in res.casefold()


def test_movement_net_current_liabilities_widened() -> None:
    """Net current liabilities -4.47 -> -6.62 must say 'Net current liabilities widened by RMB2.15bn', NOT 'decreased'."""
    res = FinancialMovementFormatter.format_movement("Net current liabilities", -4_470_000_000.0, -6_620_000_000.0)
    assert "widened" in res.casefold()
    assert "decreased" not in res.casefold()
    assert "2.15bn" in res.casefold()


def test_movement_loss_narrowed() -> None:
    """Loss -1.57 -> -0.83 must say 'Loss narrowed by RMB0.74bn'."""
    res_scaled = FinancialMovementFormatter.format_movement("Loss for the year", -1.57, -0.83, scale=1_000_000_000)
    assert "loss narrowed" in res_scaled.casefold()
    assert "0.74bn" in res_scaled.casefold()

    res_base = FinancialMovementFormatter.format_movement("Loss for the year", -1_570_000_000.0, -830_000_000.0)
    assert "loss narrowed" in res_base.casefold()
    assert "0.74bn" in res_base.casefold()


# ---------------------------------------------------------------------------
# 2. Canonical Period Display Tests
# ---------------------------------------------------------------------------

def test_canonical_period_balance_sheet() -> None:
    """30 Apr 2025 point-in-time date on balance sheet must render as '30 Apr 2025*'."""
    obs = Observation(
        id="obs_bs_1",
        metric_original="Net current liabilities",
        raw_value="-6620000",
        confidence=0.95,
        period="30 April 2025",
        period_type="balance_sheet_date",
        audited_status="unaudited",
    )
    assert format_canonical_period(obs) == "30 Apr 2025*"


def test_canonical_period_string_with_flag() -> None:
    """format_canonical_period with raw string and is_balance_sheet=True."""
    assert format_canonical_period("2025-04-30", is_balance_sheet=True, is_unaudited=True) == "30 Apr 2025*"
    assert format_canonical_period("FY2024") == "FY2024"


# ---------------------------------------------------------------------------
# 3. Negative Chart Rendering in PPTX Export
# ---------------------------------------------------------------------------

def test_negative_chart_rendering_end_to_end() -> None:
    """Negative-only series [-1.57, -1.13, -0.83] must produce visible bars below zero with signed data labels."""
    from adaptive_document_agent.services.pptx_export import _add_native_chart

    prs = Presentation(r"adaptive_document_agent/templates/FOURIER Light Version Template EN_251217.pptx")
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    obs_list = [
        Observation(id="o1", metric_original="Loss for the year", value=-1.57, raw_value="-1.57", period="FY2022", period_type="fiscal_year", confidence=0.9),
        Observation(id="o2", metric_original="Loss for the year", value=-1.13, raw_value="-1.13", period="FY2023", period_type="fiscal_year", confidence=0.9),
        Observation(id="o3", metric_original="Loss for the year", value=-0.83, raw_value="-0.83", period="FY2024", period_type="fiscal_year", confidence=0.9),
    ]
    chart_plan = ChartPlan(
        id="c_loss",
        title="Loss for the year",
        chart_type="bar",
        question="What is the net loss trajectory?",
        observation_ids=["o1", "o2", "o3"],
    )

    _add_native_chart(slide, chart_plan, obs_list, (1.0, 1.0, 6.0, 4.0))

    # Save to memory and inspect the chart XML
    bio = io.BytesIO()
    prs.save(bio)
    bio.seek(0)

    with zipfile.ZipFile(bio) as z:
        chart_xmls = [z.read(n).decode("utf-8") for n in z.namelist() if "charts/chart" in n]
        assert len(chart_xmls) >= 1

        # Newly added chart is the last one in the package
        xml_str = chart_xmls[-1]
        # 1. maximum_scale anchored at 0.0
        assert '<c:max val="0.0"/>' in xml_str
        # 2. minimum_scale with headroom (below -1.57)
        assert '<c:min val="-2.' in xml_str or '<c:min val="-1.9' in xml_str or '<c:min val="-2.0' in xml_str
        # 3. Signed 2-decimal format code
        assert 'formatCode="0.00;-0.00;0.00"' in xml_str
        # 4. Data points are negative in the cached series data
        assert "<c:v>-1.57</c:v>" in xml_str
        assert "<c:v>-1.13</c:v>" in xml_str
        assert "<c:v>-0.83</c:v>" in xml_str
        # 5. Category tick labels at low position
        assert '<c:tickLblPos val="low"/>' in xml_str


# ---------------------------------------------------------------------------
# 4. Key Findings Pagination Layout Tests
# ---------------------------------------------------------------------------

def test_key_findings_multi_slide_pagination() -> None:
    """Retain findings at readable size, with at most three complete items per page."""
    from adaptive_document_agent.document_model import DocumentIndex
    from adaptive_document_agent.services.pptx_export import _add_findings_slide

    prs = Presentation(r"adaptive_document_agent/templates/FOURIER Light Version Template EN_251217.pptx")
    initial_slide_count = len(prs.slides)

    # Create 8 distinct charts to generate 8 findings
    obs = []
    charts = []
    for i in range(1, 9):
        o1 = Observation(id=f"o_{i}_1", metric_original=f"Metric {i}", value=10.0 * i, raw_value=str(10 * i), period="FY2023", period_type="fiscal_year", confidence=0.9, evidence=[SourceEvidence(page=i, text="t", extraction_method="heuristic", confidence=0.9)])
        o2 = Observation(id=f"o_{i}_2", metric_original=f"Metric {i}", value=20.0 * i, raw_value=str(20 * i), period="FY2024", period_type="fiscal_year", confidence=0.9, evidence=[SourceEvidence(page=i, text="t", extraction_method="heuristic", confidence=0.9)])
        obs.extend([o1, o2])
        charts.append(ChartPlan(id=f"c_{i}", title=f"Metric {i}", chart_type="bar", question="q", observation_ids=[o1.id, o2.id]))

    index = DocumentIndex(obs)
    dummy_result = PipelineResult(
        document=ParsedDocument(document_id="d1", sha256="s1", safe_filename="test.pdf", page_count=10),
        profile=DocumentProfile(document_type="Prospectus", document_summary="Summary"),
        observations=obs,
        charts=charts,
    )

    _add_findings_slide(prs, dummy_result, charts, index)

    # Eight complete findings use three pages, not truncated five-card pages.
    added_slides = len(prs.slides) - initial_slide_count
    assert added_slides == 3, f"Expected 3 slides for 8 findings, got {added_slides}"

    slide_1 = prs.slides[initial_slide_count]
    slide_2 = prs.slides[initial_slide_count + 1]

    # Verify slide titles indicate pagination
    s1_text = " ".join(s.text for s in slide_1.shapes if s.has_text_frame)
    s2_text = " ".join(s.text for s in slide_2.shapes if s.has_text_frame)
    assert "Key findings (1/3)" in s1_text
    assert "Key findings (2/3)" in s2_text
    all_text = " ".join(s.text for slide in list(prs.slides)[initial_slide_count:] for s in slide.shapes if s.has_text_frame)
    assert all(f"Metric {i}" in all_text for i in range(1, 9))
