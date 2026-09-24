"""Regression coverage for multi-metric repairs and insight provenance."""

import pytest

from adaptive_document_agent.models import AnalysisResult, Insight, PresentationPlan, PresentationSlide
from adaptive_document_agent.validation.claim_validator import ClaimValidator, repair_presentation_plan
from adaptive_document_agent.validation.presentation_provenance import insight_inputs
from adaptive_document_agent.services.qa_reporter import run_comprehensive_qa
from test_company_summary_pages import sample
from test_presentation_review_regressions import observations


def margin_data():
    groups = [observations('Gross profit', [50, 60, 70], 'currency'),
              observations('Gross profit: %of Revenue', [50.5, 40.8, 43.5], 'percent'),
              observations('Cost of sales', [50, 80, 100], 'currency')]
    for i, group in enumerate(groups):
        for o in group:
            o.id = f'{i}_{o.id}'
    return [o for group in groups for o in group]


@pytest.mark.parametrize('direction', ['declined', 'increased'])
def test_margin_direction_is_not_assigned_to_neighbouring_cost_metric(direction):
    obs = margin_data()
    title = f'Gross profit increased while gross margin {direction} and cost of sales rose.'
    plan = PresentationPlan(title='Analysis', slides=[PresentationSlide(
        id='margin', slide_type='analysis', title=title, observation_ids=[o.id for o in obs])])
    repaired, _ = repair_presentation_plan(plan, obs)
    assert 'gross margin declined' in repaired.slides[0].title
    assert 'cost of sales rose' in repaired.slides[0].title
    assert not ClaimValidator().validate_plan(repaired, obs)
    before = repaired.model_dump()
    repair_presentation_plan(repaired, obs)
    assert repaired.model_dump() == before


def test_repair_rolls_back_slide_when_a_new_metric_error_is_introduced(monkeypatch):
    from adaptive_document_agent.validation import claim_validator
    obs = margin_data()
    original = 'Gross profit increased while gross margin increased and cost of sales rose.'
    plan = PresentationPlan(title='Analysis', slides=[PresentationSlide(id='margin',
        slide_type='analysis', title=original, observation_ids=[o.id for o in obs])])
    real_repair = claim_validator.repair_presentation_plan_from_issues
    def faulty_repair(plan, issues):
        repaired, messages = real_repair(plan, issues)
        repaired.slides[0].title = repaired.slides[0].title.replace('cost of sales rose', 'cost of sales declined')
        return repaired, messages
    monkeypatch.setattr(claim_validator, 'repair_presentation_plan_from_issues', faulty_repair)
    repaired, messages = repair_presentation_plan(plan, obs)
    assert repaired.slides[0].title == original
    assert not messages


def linked_result():
    result = sample()
    selected = observations('Current liabilities', [100, 200], 'currency')
    other = observations('Current liabilities', [500, 400], 'currency')
    for o in selected:
        o.dimensions = {'table_context': 'Selected statement'}
    for o in other:
        o.id = 'other_' + o.id
        o.dimensions = {'table_context': 'Different scope'}
    result.observations = selected + other
    result.analysis_results = [AnalysisResult(task_id='task', title='Trend',
        input_observation_ids=[o.id for o in selected])]
    result.insights = [Insight(id='insight', title='Current liabilities increased',
        narrative='Current liabilities increased.', kind='calculated_result', result_ids=['task'])]
    result.presentation_plan.slides = [PresentationSlide(id='summary', slide_type='executive_summary',
        title='Executive summary', bullets=['Current liabilities increased'], insight_ids=['insight'])]
    return result


def test_summary_uses_exact_calculation_inputs_not_whole_document():
    result = linked_result()
    mapping = insight_inputs(result)
    assert mapping == {'insight': ['o0', 'o1']}
    assert not ClaimValidator().validate_plan(result.presentation_plan, result.observations,
        insight_observation_ids=mapping)
    before = result.presentation_plan.model_dump()
    repair_presentation_plan(result.presentation_plan, result.observations, insight_observation_ids=mapping)
    assert result.presentation_plan.model_dump() == before
    result.presentation_plan.slides[0].bullets = ['Current liabilities decreased']
    assert any(i.code == 'directional_contradiction' for i in ClaimValidator().validate_plan(
        result.presentation_plan, result.observations, insight_observation_ids=mapping))


def test_provenance_excludes_missing_inputs_and_unrelated_tasks():
    result = linked_result()
    result.analysis_results[0].input_observation_ids += ['missing', 'o0']
    result.analysis_results.append(AnalysisResult(task_id='other', title='Other', input_observation_ids=['other_o0']))
    assert insight_inputs(result) == {'insight': ['o0', 'o1']}


def test_read_only_qa_does_not_mutate_saved_result():
    result = linked_result()
    before = result.model_dump()
    run_comprehensive_qa(result, auto_repair=False)
    assert result.model_dump() == before


def test_summary_bullets_keep_independent_series_scopes():
    result = linked_result()
    slide = result.presentation_plan.slides[0]
    slide.bullets = ['Current liabilities increased', 'Current liabilities decreased']
    slide.observation_ids = [o.id for o in result.observations]
    slide.bullet_observation_ids = [['o0', 'o1'], ['other_o0', 'other_o1']]
    assert not ClaimValidator().validate_slide(slide, result.observations)
    slide.bullets[1] = 'Current liabilities increased'
    issues = ClaimValidator().validate_slide(slide, result.observations)
    assert issues and all(i.bullet_index == 1 for i in issues)


def test_valid_existing_introduction_also_shortens_cover():
    from adaptive_document_agent.agent.company_introduction import ensure_company_introduction
    result = sample()
    result.presentation_plan.company.name = 'Example Company'
    result.presentation_plan.company.identity_state = 'RESOLVED'
    result.presentation_plan.slides = [PresentationSlide(id='cover', slide_type='cover', title='An overly verbose title')]
    ensure_company_introduction(None, result, result.presentation_plan)
    assert result.presentation_plan.slides[0].title == 'Example Company'


def test_plan_numeric_validation_follows_insight_input_dates():
    from adaptive_document_agent.agent.presentation_plan_recovery import PresentationPlanRecovery
    from adaptive_document_agent.validation.presentation_plan_validator import PresentationPlanValidator
    from tests.test_pptx_export import _result
    result = _result()
    result.observations[0].period = '2021-12-31'
    result.observations[-1].period = '2024-09-30'
    result.analysis_results = [AnalysisResult(task_id='source', title='Trend',
        input_observation_ids=[o.id for o in result.observations])]
    result.insights[0].result_ids = ['source']
    plan = PresentationPlanRecovery().fallback(result)
    summary = next(s for s in plan.slides if s.slide_type == 'executive_summary')
    summary.bullets = ['Revenue increased from 2021-12-31 to 2024-09-30']
    PresentationPlanValidator().validate(plan, result)
    summary.bullets = ['Revenue increased from 2021-12-31 to 2099-09-30']
    with pytest.raises(ValueError, match='unsupported numeric'):
        PresentationPlanValidator().validate(plan, result)


def test_header_percent_is_supported_without_authorizing_invented_percent():
    from adaptive_document_agent.agent.presentation_plan_recovery import PresentationPlanRecovery
    from adaptive_document_agent.validation.presentation_plan_validator import PresentationPlanValidator
    from tests.test_pptx_export import _result
    result = _result()
    plan = PresentationPlanRecovery().fallback(result)
    item = result.observations[0]
    item.raw_value = '46.8'
    item.value = 46.8
    item.unit = 'percent'
    summary = next(s for s in plan.slides if s.slide_type == 'executive_summary')
    summary.observation_ids = [item.id]
    summary.bullets = ['Reported percentage: 46.8%']
    PresentationPlanValidator().validate(plan, result)
    summary.bullets = ['Reported percentage: 99.9%']
    with pytest.raises(ValueError, match='unsupported numeric'):
        PresentationPlanValidator().validate(plan, result)


def test_topic_summary_represents_distinct_model_selected_questions():
    from adaptive_document_agent.agent.presentation_plan_recovery import PresentationPlanRecovery
    from adaptive_document_agent.agent.presentation_topic_selector import series_directory
    from adaptive_document_agent.models import PresentationTopic, PresentationTopicSelection
    from tests.test_pptx_export import _result
    result = _result()
    cash = [o.model_copy(deep=True) for o in result.observations]
    for o in cash:
        o.id = 'cash_' + o.id
        o.metric_original = 'Cash'
    result.observations += cash
    result.analysis_results = [AnalysisResult(task_id='revenue', title='Revenue',
        input_observation_ids=[o.id for o in result.observations if not o.id.startswith('cash_')]),
        AnalysisResult(task_id='cash', title='Cash', input_observation_ids=[o.id for o in cash])]
    result.insights[0].result_ids = ['revenue']
    duplicate = result.insights[0].model_copy(update={'id': 'duplicate', 'importance': .95})
    other = result.insights[0].model_copy(update={'id': 'cash', 'title': 'Cash increased',
        'narrative': 'Cash increased.', 'result_ids': ['cash'], 'importance': .2})
    result.insights += [duplicate, other]
    directory, _ = series_directory(result)
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(
        id=f't{i}', title='Reported movement', question='How did the reported values change?',
        rationale='Comparable values', series_ids=[entry['id']]) for i, entry in enumerate(directory)])
    plan = PresentationPlanRecovery().from_selected_topics(result)
    summary = next(s for s in plan.slides if s.slide_type == 'executive_summary')
    assert set(summary.insight_ids) == {'duplicate', 'cash'}
    for topic in result.presentation_topics.topics:
        topic.takeaway = 'Reported values increased across the observed periods.'
    plan = PresentationPlanRecovery().from_selected_topics(result)
    summary = next(s for s in plan.slides if s.slide_type == 'executive_summary')
    assert summary.bullets == [t.takeaway for t in result.presentation_topics.topics]
    assert set(summary.observation_ids) == {o.id for o in result.observations}
    assert not summary.insight_ids
    assert len(summary.bullet_observation_ids) == len(summary.bullets)


def test_topic_summary_renders_repaired_text_not_old_narrative():
    from pptx import Presentation
    from pptx.util import Inches
    from adaptive_document_agent.document_model import DocumentIndex
    from adaptive_document_agent.services.pptx_export import _add_planned_summary
    result = linked_result()
    result.presentation_plan.planning_origin = 'topic_recovery'
    result.insights[0].narrative = 'Incorrect old narrative'
    summary = result.presentation_plan.slides[0]
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    _add_planned_summary(deck, result, summary, DocumentIndex(result.observations))
    text = ' '.join(s.text for s in deck.slides[0].shapes if s.has_text_frame)
    assert 'Current liabilities increased' in text
    assert 'Incorrect old narrative' not in text
    assert 'Incorrect old narrative' in deck.slides[0].notes_slide.notes_text_frame.text
