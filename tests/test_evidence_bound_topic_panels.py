"""Synthetic source-bound regressions; no private source documents or IDs."""
import json

import pytest

from adaptive_document_agent.models import (Observation, SourceEvidence, PresentationSlide,
    PresentationTopic, PresentationTopicSelection, PresentationTheme, ValidationIssue,
    DocumentPage, ExtractedTable, TableRow)
from adaptive_document_agent.agent.presentation_topic_selector import series_directory, PresentationTopicSelector
from adaptive_document_agent.agent.presentation_plan_recovery import PresentationPlanRecovery
from adaptive_document_agent.agent.topic_selection_repair import retain_valid_topics
from adaptive_document_agent.document_model.series import display_metric_name
from adaptive_document_agent.document_model.topic_matcher import is_positive_topic_mismatch
from adaptive_document_agent.services.executive_brief import restore_percentage_symbols
from adaptive_document_agent.services.presentation_scope import scope_items
from adaptive_document_agent.services.pptx_export import _unit_label
from adaptive_document_agent.validation.presentation_topic_relations import TopicRelationValidator
from adaptive_document_agent.validation.presentation_parallel_series import validate_parallel_series
from tests.test_pptx_export import _result


def observation(label, value, row, *, period='FY2024', unit='percent'):
    return Observation(id=f'{row}-{period}', metric_original=label, value=value, raw_value=str(value),
        unit=unit, unit_family='percentage' if unit == 'percent' else unit,
        raw_unit='%' if unit == 'percent' else 'RMB', currency=None if unit == 'percent' else 'CNY',
        unit_scale=1, period=period, period_basis='FY', row_id=row, column_id=1,
        table_id='synthetic-table', confidence=.99,
        evidence=[SourceEvidence(page=234, table_id='synthetic-table', row_label=label,
            text=str(value), extraction_method='digital_table', confidence=.99)])


def relation_result(labels, values, *, title, rationale, source='', unit='percent'):
    result = _result()
    result.observations = [observation(label, value, i, unit=unit) for i, (label, value) in enumerate(zip(labels, values))]
    result.document.pages = [DocumentPage(page_number=234, text=source, tables=[ExtractedTable(
        table_id='synthetic-table', page=234,
        rows=[TableRow(page=234, cells=[label, str(value)]) for label, value in zip(labels, values)])])]
    _, lookup = series_directory(result)
    topic = PresentationTopic(id='source-topic', title=title, question=title+'?', rationale=rationale,
                              series_ids=list(lookup))
    result.presentation_topics = PresentationTopicSelection(topics=[topic])
    slide = PresentationSlide(id='source-slide', slide_type='analysis', title=title,
        theme_id=topic.id, analytical_question=topic.question, selection_reason=rationale,
        observation_ids=[o.id for o in result.observations], source_pages=[234])
    result.presentation_plan = PresentationPlanRecovery().fallback(result, validate=False)
    result.presentation_plan.themes = [PresentationTheme(id=topic.id, title=title,
        question=topic.question, rationale=rationale, observation_ids=slide.observation_ids, source_pages=[234])]
    return result, slide


def test_declared_cost_share_has_signed_source_subtotal_proof():
    result, slide = relation_result(['Revenue', 'Cost of sales', 'Gross profit'], [100, -63, 37],
        title='Gross profitability improvement', rationale='Compare gross profit and cost of sales.')
    cost = result.observations[1]
    assert is_positive_topic_mismatch(cost, slide)
    validator = TopicRelationValidator(result)
    assert validator.relation(cost, slide).kind == 'source_adjacent_subtotal'
    assert not validator.mismatch(cost, slide)
    # The chart's own caption still has the original strict matcher.
    caption = PresentationSlide(id='caption', slide_type='analysis', title='Gross profit margin')
    assert is_positive_topic_mismatch(cost, caption)


@pytest.mark.parametrize('change', ['arithmetic', 'unit', 'scale', 'period', 'entity', 'category', 'table', 'alignment', 'undeclared', 'unselected'])
def test_cost_support_never_globally_allows_expenses(change):
    result, slide = relation_result(['Revenue', 'Cost of sales', 'Gross profit'], [100, -63, 37],
        title='Gross profitability improvement', rationale='Compare gross profit and cost of sales.')
    cost = result.observations[1]
    if change == 'arithmetic': cost.value = -60
    elif change == 'unit': cost.raw_unit = 'RMB'
    elif change == 'scale': cost.unit_scale = 1000
    elif change == 'period': cost.period = '6M2024'
    elif change == 'entity': cost.entity = 'Other entity'
    elif change == 'category': cost.category_dimensions = {'business': 'Other segment'}
    elif change == 'table': cost.evidence[0].table_id = 'unrelated-table'; cost.table_id = 'unrelated-table'
    elif change == 'alignment': result.document.pages[0].tables[0].rows[1].alignment_status = 'ambiguous'
    elif change == 'undeclared': result.presentation_topics.topics[0].rationale = 'Compare gross profit.'
    else:
        _, lookup = series_directory(result)
        result.presentation_topics.topics[0].series_ids = [sid for sid, items in lookup.items() if cost.id not in {o.id for o in items}]
    assert TopicRelationValidator(result).mismatch(cost, slide)


def test_explicit_source_adjustment_definition_supports_loss_comparison():
    result, slide = relation_result(['Reported loss', 'Share-based payments expense', 'Adjusted net loss'], [-80, 12, -60],
        title='Reported loss widened while adjusted net loss narrowed',
        rationale='Compare reported loss, adjusted net loss and share-based payments expense.',
        source='We define adjusted net loss as reported loss adjusted for share-based payments expense and listing expenses.', unit='currency')
    relation = TopicRelationValidator(result).relation(result.observations[1], slide)
    assert relation and relation.kind == 'source_definition'
    assert 'listing expenses' in relation.source_quote
    result.document.pages[0].text = 'Share-based payments expense and adjusted net loss are listed in this table.'
    assert TopicRelationValidator(result).mismatch(result.observations[1], slide)


@pytest.mark.parametrize('source', [
    'Adjusted net loss is defined separately from share-based payments expense.',
    'Adjusted net loss excludes share-based payments expense.',
    'Adjusted net loss does not include share-based payments expense.',
])
def test_citation_cooccurrence_or_exclusion_is_not_support_proof(source):
    result, slide = relation_result(['Reported loss', 'Share-based payments expense', 'Adjusted net loss'], [-80, 12, -60],
        title='Reported loss widened while adjusted net loss narrowed',
        rationale='Compare reported loss, adjusted net loss and share-based payments expense.', source=source, unit='currency')
    assert TopicRelationValidator(result).mismatch(result.observations[1], slide)


def test_supported_adjustment_role_cannot_hide_a_missing_bridge_component():
    result, slide = relation_result(['Reported loss', 'Share-based payments expense', 'Adjusted net loss'], [-80, 12, -60],
        title='Reported loss widened while adjusted net loss narrowed',
        rationale='Compare reported loss, adjusted net loss and share-based payments expense.',
        source='Adjusted net loss consists of reported loss adjusted for share-based payments expense and listing expenses.', unit='currency')
    assert TopicRelationValidator(result).relation(result.observations[1], slide)
    from adaptive_document_agent.models import PresentationVisualBlock
    from adaptive_document_agent.services.presentation_waterfall import waterfall_data
    with pytest.raises(ValueError, match='do not reconcile'):
        waterfall_data(PresentationVisualBlock(role='waterfall', observation_ids=[o.id for o in result.observations]),
                       {o.id: o for o in result.observations})


@pytest.mark.parametrize('change', ['unit', 'currency', 'period', 'conflict', 'invalid'])
def test_parallel_panels_keep_conflict_guards(change):
    items = [observation('Revenue', 100, 0, unit='currency'), observation('Revenue', 110, 1, period='FY2025', unit='currency')]
    if change == 'unit': items[1].unit_scale = 1000
    elif change == 'currency': items[1].currency = 'USD'
    elif change == 'period': items[1].period = '6M2025'
    elif change == 'conflict': items[1].period = 'FY2024'
    else: items[1].validation_status = 'ambiguous'
    with pytest.raises(ValueError): validate_parallel_series([items])


def test_more_than_three_valid_series_paginate_with_distinct_page_scope():
    result = _result()
    result.observations = []
    result.charts = []
    labels = ['Revenue', 'Gross profit', 'Research and development expenses', 'Administrative expenses']
    for row, label in enumerate(labels):
        for year in (2023, 2024, 2025):
            result.observations.append(observation(label, (row + 1) * year, row, period=f'FY{year}', unit='currency'))
    _, lookup = series_directory(result)
    # Duplicate literal copies are harmless in independent panels, but cannot
    # be represented as a falsely complete matrix.
    duplicate = result.observations[0].model_copy(deep=True)
    duplicate.id = 'same-value-copy'
    result.observations.append(duplicate)
    _, lookup = series_directory(result)
    topic = PresentationTopic(id='cost-growth', title='Costs and revenue', question='How did costs and revenue move?',
        rationale='Compare revenue, gross profit and operating expenses.', series_ids=list(lookup))
    result.presentation_topics = PresentationTopicSelection(topics=[topic])
    PresentationTopicSelector._validate(result.presentation_topics, lookup)
    from adaptive_document_agent.agent.chart_planner import ChartPlanner
    from adaptive_document_agent.document_model import DocumentIndex
    result.charts = ChartPlanner().plan([], [], DocumentIndex(result.observations), requested_series=list(lookup.values()), only_requested=True)
    plan = PresentationPlanRecovery().from_selected_topics(result)
    slides = [s for s in plan.slides if s.slide_type == 'analysis']
    assert len([s for s in slides if s.chart_ids]) == 2
    assert {s.theme_id for s in slides} == {topic.id}
    assert all(len(s.chart_ids) <= 3 and len(s.observation_ids) <= 40 for s in slides)
    assert {cid for s in slides for cid in s.chart_ids} == {c.id for c in result.charts}
    assert len({frozenset(s.chart_ids) for s in slides if s.chart_ids}) == 2
    type(plan).model_validate(plan.model_dump())


def test_repair_budget_is_per_topic_even_when_combined_payload_is_large():
    result = _result()
    records = result.observations
    for o in records: o.evidence[0].text = 'source context ' * 400
    lookup = {'first': records, 'second': records}
    drafts = [PresentationTopic(id=key, title='Revenue 999', question='How did revenue move?',
        rationale='Comparable reported revenue.', series_ids=[key]) for key in lookup]
    class Gateway:
        calls = []
        def generate_structured(self, messages, response_model, **kwargs):
            self.calls.append(messages)
            topic = drafts[len(self.calls) - 1].model_copy(update={'title': 'Revenue movement'})
            return response_model(topics=[topic])
    gateway = Gateway()
    retained = retain_valid_topics(PresentationTopicSelection(topics=drafts), lookup, None, result,
                                  gateway, PresentationTopicSelector._validate)
    assert len(gateway.calls) == 2
    assert len(retained.topics) == 2


def test_count_scaled_units_and_source_bound_child_label():
    item = observation('Sales volume (units)', 12000, 0, unit='count')
    item.currency = None; item.raw_unit = 'units'; item.parent_section = 'Average selling price'
    assert _unit_label([item], 'thousand') == 'thousand units'
    assert display_metric_name(item) == 'Sales volume (units)'
    item.parent_section = 'Industrial products'
    assert display_metric_name(item) == 'Industrial products: Sales volume (units)'
    item.metric_original = '(RMB per unit)'; item.evidence[0].row_label = '(RMB per unit)'
    item.parent_section = 'Average selling price'; item.unit = 'currency'; item.raw_unit = 'RMB per unit'
    assert 'Average selling price' in display_metric_name(item)


def test_scaled_count_chart_labels_preserve_whole_units():
    from pptx import Presentation
    from adaptive_document_agent.models import ChartPlan
    from adaptive_document_agent.services.pptx_export import _add_native_chart
    values = [observation('Sales volume (units)', value, 0, period=f'FY{year}', unit='count')
              for year, value in ((2023, 17), (2024, 11234), (2025, 35467))]
    for o in values: o.currency = None; o.raw_unit = 'units'
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    plan = ChartPlan(id='volume-chart', title='Sales volume', question='How did sales volume change?', chart_type='line',
        observation_ids=[o.id for o in values], source_pages=[234], show_data_labels=True)
    scale, label = _add_native_chart(slide, plan, values, (.5, .5, 7, 4), compact=False)
    chart = next(s.chart for s in slide.shapes if s.has_chart)
    assert scale == 1000 and _unit_label(values, label) == 'thousands units'
    assert list(chart.series[0].values) == [0.017, 11.234, 35.467]
    assert '0.000' in chart.plots[0].data_labels.number_format


def test_support_relation_does_not_relax_a_chart_own_caption():
    from adaptive_document_agent.models import ChartPlan
    from adaptive_document_agent.services.qa_reporter import run_comprehensive_qa
    result, slide = relation_result(['Revenue', 'Cost of sales', 'Gross profit'], [100, -63, 37],
        title='Gross profitability improvement', rationale='Compare gross profit and cost of sales.')
    chart = ChartPlan(id='wrong-caption', title='Gross profit margin', question='How did gross margin change?', chart_type='bar',
        observation_ids=[result.observations[1].id], source_pages=[234])
    result.charts = [chart]; slide.chart_ids = [chart.id]
    result.presentation_plan.slides[3] = slide
    assert not TopicRelationValidator(result).mismatch(result.observations[1], slide)
    qa = run_comprehensive_qa(result, auto_repair=False)
    assert any(issue.code == 'chart_title_data_mismatch' for issue in qa.critical_errors)


@pytest.mark.parametrize('copies', [1, 2])
def test_cached_audit_recovery_requires_unique_original_topic_ids(copies):
    from adaptive_document_agent.agent.presentation_topic_audit_recovery import recover_audited_topics
    result = _result(); result.charts = []
    original = list(result.observations)
    for item in original:
        gross = item.model_copy(deep=True)
        gross.id = 'gross-' + gross.id; gross.metric_original = 'Gross profit'
        gross.value *= .6; gross.raw_value = str(gross.value / 1000)
        result.observations.append(gross)
    _, lookup = series_directory(result)
    revenue = next(sid for sid, items in lookup.items() if items[0].metric_original == 'Revenue')
    gross = next(sid for sid, items in lookup.items() if items[0].metric_original == 'Gross profit')
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(id='revenue',
        title='Revenue movement', question='How did revenue change?', rationale='Comparable reported revenue.', series_ids=[revenue])])
    draft = PresentationTopic(id='gross', title='Gross profit movement', question='How did gross profit change?',
                             rationale='Comparable reported gross profit.', series_ids=[gross])
    for _ in range(copies):
        result.validation_warnings.append(ValidationIssue(code='presentation_topic_validation', stage='presentation',
            message=json.dumps({'phase': 'initial', 'error': 'Prior compilation limit.', 'topic': draft.model_dump()})))
    assert recover_audited_topics(result) == (copies == 1)
    if copies == 1:
        assert {t.id for t in result.presentation_plan.themes} == {'revenue', 'gross'}


def test_percent_symbols_use_only_own_literal_unambiguous_source():
    quote = 'Overseas margin was 41.4% versus domestic margin of 22.1% in 2024.'
    text = restore_percentage_symbols('Margins were 41.4 versus 22.1 in 2024.', [quote])
    assert text == 'Margins were 41.4% versus 22.1% in 2024.'
    assert restore_percentage_symbols(text, [quote]) == text
    assert restore_percentage_symbols('There were 41.4 units.', [quote]) == 'There were 41.4 units.'
    assert restore_percentage_symbols('Value 41.4.', [quote, 'Price was RMB 41.4.']) == 'Value 41.4.'


@pytest.mark.parametrize('change', [None, 'column', 'cell', 'alignment', 'competing_unit'])
def test_bare_table_percentages_require_own_quoted_resolved_source_column(change):
    from adaptive_document_agent.services.executive_brief import brief_items
    from adaptive_document_agent.models.executive_brief import ExecutiveBrief
    result, _ = relation_result(['Gross margin'], [37], title='Gross profitability', rationale='Compare gross margin.')
    page = result.document.pages[0]
    page.text = 'Gross margin 37 in 2024.'
    table = page.tables[0]; table.column_types = ['label', 'percentage']
    result.executive_brief = ExecutiveBrief.model_validate({'title': 'Results', 'items': [{
        'label': 'Margin', 'text': 'Gross margin was 37 in 2024.',
        'evidence': [{'page': 234, 'text': page.text}]}]})
    if change == 'column': table.column_types[1] = 'amount'
    elif change == 'cell': table.rows[0].cells[1] = '38'
    elif change == 'alignment': table.rows[0].alignment_status = 'ambiguous'
    elif change == 'competing_unit':
        result.observations.append(observation('Revenue', 37, 1, unit='currency'))
    assert ('37%' in brief_items(result)[0].text) == (change is None)


def test_many_series_continue_beyond_physical_page_quota_without_omitting_topics():
    result = _result(); result.observations = []; result.charts = []
    for row in range(48):
        label = 'Revenue measure ' + chr(65 + row // 26) + chr(65 + row % 26)
        for year in (2023, 2024, 2025):
            result.observations.append(observation(label, (row + 1) * year, row, period=f'FY{year}', unit='currency'))
    _, lookup = series_directory(result); series = list(lookup)
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(
        id=f'scope-{letter}', title='Revenue measures', question='How did revenue measures compare?',
        rationale='Compare the selected source revenue measures.', series_ids=ids)
        for letter, ids in zip(('a', 'b'), (series[:24], series[24:]))])
    from adaptive_document_agent.agent.chart_planner import ChartPlanner
    from adaptive_document_agent.document_model import DocumentIndex
    result.charts = ChartPlanner().plan([], [], DocumentIndex(result.observations), requested_series=list(lookup.values()), only_requested=True)
    plan = PresentationPlanRecovery().from_selected_topics(result)
    assert len(plan.slides) > 18
    assert {t.id for t in plan.themes} == {'scope-a', 'scope-b'}
    assert {cid for s in plan.slides for cid in s.chart_ids} == {c.id for c in result.charts}
    type(plan).model_validate(plan.model_dump())


def test_final_scope_lists_dropped_topics_without_audit_json_and_hides_restored_drafts():
    result = _result()
    result.presentation_plan = PresentationPlanRecovery().fallback(result, validate=False)
    topic = PresentationTopic(id='omitted', title='Regional profitability', question='What changed?',
        rationale='Compare comparable regions.', series_ids=['source-series'])
    result.validation_warnings.append(ValidationIssue(code='presentation_topic_validation', stage='presentation',
        message=json.dumps({'phase': 'initial', 'error': 'Conflicting source cells.', 'topic': topic.model_dump()})))
    items = scope_items(result)
    assert any('Not covered: Regional profitability' == i.title for i in items)
    assert not any('series_ids' in i.text for i in items)
    result.presentation_plan.themes = [PresentationTheme(id=topic.id, title=topic.title,
        question=topic.question, rationale=topic.rationale)]
    assert not any(i.title.startswith('Not covered:') for i in scope_items(result))
