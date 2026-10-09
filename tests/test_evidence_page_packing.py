"""Evidence packing changes geometry only, across arbitrary period groups."""

from pptx import Presentation
from pptx.util import Inches, Pt

from adaptive_document_agent.services.evidence_page_packing import pack_evidence_pages


def _page(deck, periods, labels, page):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    rows = [["Metric", "Unit", *periods], ["REPORTED MEASURES", "", *[""] * len(periods)]]
    rows.extend([label, "units", *[str(page * 100 + i) for i in range(len(periods))]] for label in labels)
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(.45), Inches(1.45),
                                  Inches(11.7), Inches(.38 * len(rows)))
    shape.name = "evidence:packable"
    shape.table.columns[0].width = Inches(3.8)
    for column in list(shape.table.columns)[1:]:
        column.width = Inches(7.9 / (len(rows[0]) - 1))
    for index, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = shape.table.cell(index, j)
            cell.text = text
            cell.margin_top = cell.margin_bottom = Inches(.04)
            cell.margin_left = cell.margin_right = Inches(.08)
            cell.text_frame.paragraphs[0].font.size = Pt(10.5)
    source = slide.shapes.add_textbox(Inches(.45), Inches(6.22), Inches(11.7), Inches(.25))
    source.name = "evidence:footer"
    source.text = f"Source: p. {page}"
    source.text_frame.paragraphs[0].font.size = Pt(9)
    slide.notes_slide.notes_text_frame.text = f"original records from p. {page}: {rows!r}"
    return slide


def _tables(slide):
    return [shape for shape in slide.shapes if shape.has_table]


def _cells(shape):
    return [[cell.text for cell in row.cells] for row in shape.table.rows]


def test_three_sparse_tables_fit_without_joining_annual_and_interim_headers():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(12.6), Inches(7.1)
    periods = ["31 Dec 2030", "31 Dec 2031", "31 Dec 2032", "30 Apr 2033*"]
    slides = [_page(deck, periods, ["Allocated capacity"], 3),
              _page(deck, periods[:3], ["Available capacity"], 7),
              _page(deck, periods, ["Stored units", "Transferred units", "Committed units", "Total units"], 11)]
    original = [_cells(_tables(slide)[0]) for slide in slides]
    notes = [slide.notes_slide.notes_text_frame.text for slide in slides]

    packed = pack_evidence_pages(deck, slides)

    assert len(packed) == len(deck.slides) == 1
    tables = _tables(packed[0])
    assert [_cells(shape) for shape in tables] == original
    assert [len(shape.table.columns) for shape in tables] == [6, 5, 6]
    assert all(item in packed[0].notes_slide.notes_text_frame.text for item in notes)
    footer = next(shape for shape in packed[0].shapes if shape.name == "evidence:footer")
    assert [paragraph.text for paragraph in footer.text_frame.paragraphs] == [
        "Source: p. 3", "Source: p. 7", "Source: p. 11"]
    assert all(tables[i].top + tables[i].height < tables[i + 1].top for i in range(2))
    assert tables[-1].top + tables[-1].height < footer.top
    assert all(cell.text_frame.paragraphs[0].font.size == Pt(10.5)
               for table in tables for row in table.table.rows for cell in row.cells)


def test_fitting_prefix_moves_without_dropping_remaining_table_or_sources():
    from copy import deepcopy
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(12.6), Inches(7.1)
    first = _page(deck, ['FY2030'], ['First measure'] * 4, 3)
    second = _page(deck, ['FY2031'], ['Second measure'], 7)
    third = _page(deck, ['FY2032'], ['Third measure'] * 6, 11)
    # Put the latter two complete tables on one page; only its prefix fits
    # after the first table. Keep exact source records on both output pages.
    second.shapes._spTree.insert_element_before(deepcopy(_tables(third)[0]._element), 'p:extLst')
    _tables(second)[-1].top = Inches(3.0)
    originals = [_cells(t) for s in (first, second) for t in _tables(s)]
    footer = next(s for s in first.shapes if s.name == 'evidence:footer')
    footer.text = 'Source: Document disclosures (p. 3) | * Unaudited'
    footer = next(s for s in second.shapes if s.name == 'evidence:footer')
    footer.text = 'Source: Document disclosures (p. 7, 11) | * Unaudited'
    packed = pack_evidence_pages(deck, [first, second])
    assert len(packed) == 2
    assert [len(_tables(s)) for s in packed] == [2, 1]
    assert [_cells(t) for s in packed for t in _tables(s)] == originals
    text = next(s.text for s in packed[0].shapes if s.name == 'evidence:footer')
    assert text == 'Source: Document disclosures (p. 3, 7, 11) | * Unaudited'


def test_wrapped_rows_keep_separate_pages_when_compaction_would_clip_text():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(12.6), Inches(7.1)
    long_label = "Programme capacity with explicit scope and geographic qualifications " * 4
    slides = [_page(deck, ["FY2030", "FY2031"], [long_label] * 5, page) for page in (2, 9)]
    original = [_cells(_tables(slide)[0]) for slide in slides]

    assert len(pack_evidence_pages(deck, slides)) == 2
    assert [_cells(_tables(slide)[0]) for slide in slides] == original


def test_first_fit_uses_earlier_space_without_splitting_or_rewriting_later_table():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(12.6), Inches(7.1)
    slides = [_page(deck, ['FY2030'], ['First'], 3),
              _page(deck, ['FY2031'], ['Complete row'] * 9, 7),
              _page(deck, ['6M2032'], ['Final'], 11)]
    original = [_cells(_tables(s)[0]) for s in slides]
    packed = pack_evidence_pages(deck, slides, first_fit=True)
    assert len(packed) == 2
    assert [_cells(s) for s in _tables(packed[0])] == [original[0], original[2]]
    assert _cells(_tables(packed[1])[0]) == original[1]
    assert 'p. 11' in packed[0].notes_slide.notes_text_frame.text


def test_packing_retains_conflicting_values_and_distinct_conventions():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(12.6), Inches(7.1)
    slides = [_page(deck, ["FY2030", "FY2031"], ["Capacity"], page) for page in (2, 9)]
    for slide, text in zip(slides, ["Units are reported in thousands.", "Negative values are signed source values."]):
        shape = slide.shapes.add_textbox(Inches(.45), Inches(5.90), Inches(11.7), Inches(.28))
        shape.name = "evidence:convention"
        shape.text = text
        shape.text_frame.paragraphs[0].font.size = Pt(10)

    packed = pack_evidence_pages(deck, slides)

    assert len(packed) == 1
    tables = _tables(packed[0])
    assert tables[0].table.cell(2, 2).text == "200"
    assert tables[1].table.cell(2, 2).text == "900"
    convention = next(shape for shape in packed[0].shapes if shape.name == "evidence:convention")
    assert "thousands" in convention.text and "signed source values" in convention.text
    assert tables[-1].top + tables[-1].height < convention.top
