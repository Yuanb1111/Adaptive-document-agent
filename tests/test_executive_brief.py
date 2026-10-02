"""Editorial briefs generalise across documents and preserve item-level evidence."""
from contextlib import nullcontext
import json
import pytest
from pptx import Presentation
from pptx.util import Inches
from adaptive_document_agent.agent.executive_brief import ExecutiveBriefWriter
from adaptive_document_agent.models import DocumentPage, DocumentProfile, ParsedDocument, PipelineResult, PresentationPlan
from adaptive_document_agent.models.executive_brief import ExecutiveBrief
from adaptive_document_agent.services.executive_brief import brief_items, display_brief, validate_executive_brief
from adaptive_document_agent.services.llm import LLMGateway
from adaptive_document_agent.services.llm.config import LLMSettings, ProviderName
from adaptive_document_agent.services.llm.mock import MockLLMClient
from adaptive_document_agent.services.pptx_export import _add_planned_summary
from adaptive_document_agent.models import PresentationSlide
from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.ui.overview import render, _literal


def result_for(texts):
    return PipelineResult(document=ParsedDocument(document_id='brief',sha256='a'*64,
        safe_filename='unknown.pdf',page_count=len(texts),
        pages=[DocumentPage(page_number=i,text=t) for i,t in enumerate(texts,1)]),
        profile=DocumentProfile(document_purpose='Explain material disclosed facts'))


def payload(text, page=1, label='Performance'):
    return {'title':'Key takeaways','items':[{'label':label,'text':text,
                                            'evidence':[{'page':page,'text':text}]}]}


def gateway(responses):
    client=MockLLMClient(responses)
    return LLMGateway(client, LLMSettings(provider=ProviderName.MOCK,model='mock')),client


@pytest.mark.parametrize('source,label', [
    ('Revenue was USD 4.6 billion in 2025. The reported loss includes non-cash revaluation charges.', 'Results'),
    ('Employee participation was 82% in 2025. Temporary workers were excluded from the survey.', 'Coverage'),
    ('The trial enrolled 240 adults. Its findings do not establish long-term effects.', 'Study limits'),
])
def test_writer_keeps_material_numbers_and_qualifiers_for_unknown_documents(source,label):
    r=result_for([source]);g,client=gateway([payload(source,label=label)])
    r.executive_brief=ExecutiveBriefWriter(g).generate(r)
    assert brief_items(r)[0].text==source
    assert not validate_executive_brief(r.executive_brief,r)
    assert len(client.calls)==1
    assert 'untrusted' in client.calls[0][0]['content']
    assert 'non-cash' in client.calls[0][0]['content']


@pytest.mark.parametrize('change', ['page','quote','value','currency','scale'])
def test_cached_brief_cannot_borrow_numbers_or_change_currency_and_scale(change):
    source='Revenue was USD 4.6 billion in 2025.'
    r=result_for([source,'Costs were USD 8.1 billion in 2025.'])
    brief=ExecutiveBrief.model_validate(payload(source))
    item=brief.items[0]
    if change=='page':item.evidence[0].page=2
    elif change=='quote':item.evidence[0].text='Revenue was USD 99 billion in 2025.'
    elif change=='value':item.text=item.text.replace('4.6','8.1')
    elif change=='currency':item.text=item.text.replace('USD','RMB')
    else:item.text=item.text.replace('billion','million')
    assert validate_executive_brief(brief,r)
    r.executive_brief=brief
    with pytest.raises(ValueError):brief_items(r)
    from adaptive_document_agent.validation.presentation_data_qa import validate_presentation_data
    assert any(i.code=='invalid_executive_brief' for i in validate_presentation_data(r))


def test_writer_repairs_only_once_without_silently_accepting_invalid_copy():
    source='Revenue was USD 4.6 billion in 2025.'
    bad=payload(source);bad['items'][0]['text']='Revenue was USD 99 billion in 2025.'
    g,client=gateway([bad,{'item_0': payload(source)['items'][0]}])
    assert ExecutiveBriefWriter(g).generate(result_for([source])).items[0].text==source
    assert len(client.calls)==2
    g,client=gateway([bad,{'item_0': bad['items'][0]}])
    with pytest.raises(ValueError,match='evidence checks'):
        ExecutiveBriefWriter(g).generate(result_for([source]))
    assert len(client.calls)==2


def test_schema_invalid_brief_without_identifiable_items_fails_without_regeneration():
    source = 'Revenue was USD 4.6 billion in 2025.'
    result = result_for([source + ' Context.' * 800])
    g, client = gateway([{}, payload(source)])
    with pytest.raises(ValueError, match='nonempty original items'):
        ExecutiveBriefWriter(g).generate(result)
    assert len(client.calls) == 1
    audit = json.loads(result.validation_warnings[-1].message)
    assert audit['outcome'] == 'rejected'
    assert audit['original_response'] == '{}'


def test_page_selection_can_find_a_late_narrative_constraint_without_a_chart():
    source='Customers may cancel purchases without a long-term commitment.'
    r=result_for(['Background context.']*11+[source])
    g,client=gateway([{'pages':[12]},payload(source,page=12,label='Contracts')])
    brief=ExecutiveBriefWriter(g).generate(r)
    assert brief.items[0].evidence[0].page==12
    assert len(client.calls)==2
    assert '12' in client.calls[0][1]['content']
    assert 'Background context.' not in client.calls[1][1]['content']


def test_brief_retrieves_leading_selected_topic_pages_and_repairs_missing_coverage():
    texts = ['General background information.'] * 9 + [
        'First selected measure was 8 units.', 'Second selected measure was 12 units.',
        'Third selected measure was 16 units.',
    ]
    result = result_for(texts)
    result.presentation_plan = PresentationPlan(title='Review', slides=[
        PresentationSlide(id=f'topic-{page}', slide_type='analysis', title=f'Topic {page}',
                          section_title=f'Selected measure {page}', source_pages=[page])
        for page in (10, 11, 12)
    ])
    first = payload(texts[9], page=10, label='First measure')
    second = {'title': 'Key takeaways', 'items': [
        first['items'][0], payload(texts[10], page=11, label='Second measure')['items'][0],
    ]}
    g, client = gateway([{'pages': [1]}, first, {'additions': [second['items'][1]]}])
    brief = ExecutiveBriefWriter(g).generate(result)
    assert len(brief.items) == 2
    assert len(client.calls) == 3
    request = client.calls[1][1]['content']
    assert all(f'"page": {page}' in request for page in (10, 11, 12))


class WebRecorder:
    def __init__(self):self.copy=[]
    def markdown(self,text):self.copy.append(text)
    def caption(self,text):self.copy.append(text)
    def subheader(self,text):self.copy.append(text)
    def write(self,text):self.copy.append(text)
    def warning(self,text):self.copy.append(text)
    def columns(self,n):return [self]*n
    def metric(self,*args):pass
    def expander(self,*args,**kwargs):return nullcontext()


def test_web_and_native_slide_render_identical_brief_and_keep_source_notes():
    source='Revenue was USD 4.6 billion in 2025. Costs were USD 7.3 billion.'
    r=result_for([source]);r.executive_brief=ExecutiveBrief.model_validate(payload(source,label='Scale and costs'))
    web=WebRecorder();render(web,r)
    assert '**Scale and costs** — '+_literal(source) in web.copy
    ppt=Presentation();ppt.slide_width,ppt.slide_height=Inches(13.333),Inches(7.5)
    _add_planned_summary(ppt,r,PresentationSlide(id='summary',slide_type='executive_summary',
        title='Old summary',bullets=['Old chart title']),DocumentIndex([]))
    copy='\n'.join(s.text for slide in ppt.slides for s in slide.shapes if s.has_text_frame)
    assert source in copy and 'Scale and costs' in copy
    assert 'Old chart title' not in copy
    assert source in ppt.slides[0].notes_slide.notes_text_frame.text
    assert r.model_validate_json(r.model_dump_json()).executive_brief==r.executive_brief


def test_web_brief_displays_document_links_as_literal_text():
    text = "![tracking](https://invalid.example/image.png)"
    assert _literal(text).startswith(r"\!\[tracking\]\(")


def test_chinese_amounts_keep_currency_and_magnitude():
    source = "2025年收入46亿美元，相关费用为73亿元人民币。"
    r = result_for([source])
    brief = ExecutiveBrief.model_validate(payload(source))
    assert not validate_executive_brief(brief, r)
    for changed in (source.replace("46亿美元", "46亿港元"), source.replace("46亿美元", "46万美元")):
        brief.items[0].text = changed
        assert validate_executive_brief(brief, r)


def test_invalid_item_does_not_discard_three_independently_verified_brief_items():
    source = 'Revenue was 12. Costs were 8. Cash was 4.'
    r = result_for([source])
    items = [
        {'label': label, 'text': sentence, 'evidence': [{'page': 1, 'text': sentence}]}
        for label, sentence in [('Revenue', 'Revenue was 12.'),
                                ('Costs', 'Costs were 8.'),
                                ('Cash', 'Cash was 4.')]
    ]
    items.append({'label': 'Invented', 'text': 'Debt was 99.',
                  'evidence': [{'page': 1, 'text': 'Revenue was 12.'}]})
    g, client = gateway([{'title': 'Highlights', 'items': items}, {'item_3': items[3]}])
    brief = ExecutiveBriefWriter(g).generate(r)
    assert [item.label for item in brief.items] == ['Revenue', 'Costs', 'Cash']
    assert not validate_executive_brief(brief, r)
    assert len(client.calls) == 2


def test_missing_editorial_brief_uses_selected_sourced_charts_in_web_and_ppt():
    from tests.test_pptx_export import _result
    from adaptive_document_agent.models import PresentationPlan

    r = _result()
    original = r.model_dump()
    from adaptive_document_agent.models import ValidationIssue
    r.validation_warnings.append(ValidationIssue(code='executive_brief_unavailable',
        message='Source checks rejected the final briefing.', stage='report'))
    r.presentation_plan = PresentationPlan(title='Review', slides=[
        PresentationSlide(id='summary', slide_type='executive_summary', title='Executive Summary'),
        PresentationSlide(id='revenue', slide_type='analysis', title='Revenue',
                          message='How did revenue change?', chart_ids=['chart-1']),
    ])
    title, items = display_brief(r)
    assert title == 'Executive Summary'
    assert len(items) == 1 and items[0].pages == [234]
    assert 'FY2023 RMB 267.0m' in items[0].text
    assert 'FY2025 RMB 521.7m' in items[0].text
    web = WebRecorder(); render(web, r)
    assert any(_literal(items[0].text) in copy for copy in web.copy)
    ppt = Presentation(); ppt.slide_width, ppt.slide_height = Inches(13.333), Inches(7.5)
    _add_planned_summary(ppt, r, r.presentation_plan.slides[0], DocumentIndex(r.observations))
    visible = ' '.join(s.text for slide in ppt.slides for s in slide.shapes if s.has_text_frame)
    assert items[0].text in visible
    assert [item.model_dump() for item in r.observations] == original['observations']
