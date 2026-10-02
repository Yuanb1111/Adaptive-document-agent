"""Synthetic evidence boundaries for deterministic fallback chart grouping."""

import pytest

from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import ChartPlan, Observation, SourceEvidence
from adaptive_document_agent.services.pptx_export import (
    _chart_group_title,
    _charts_belong_together,
    _group_chart_plans,
)


def _series(key, metric, *, page=2, context="", table=None, unit="count",
            canonical=None, periods=("FY2018", "FY2019"), title=None):
    values = [
        Observation(
            id=f"{key}-{position}", metric_original=metric, metric_canonical=canonical,
            value=10 + position, raw_value=str(10 + position), unit=unit,
            period=period, dimensions={"table_context": context} if context else {},
            evidence=[SourceEvidence(page=page, table_id=table, row_label=metric,
                                     extraction_method="synthetic", confidence=1)],
            confidence=1,
        )
        for position, period in enumerate(periods)
    ]
    chart = ChartPlan(id=key, title=title or metric, question="What was reported?",
                      chart_type="bar", observation_ids=[item.id for item in values],
                      source_pages=[page])
    return chart, values


def _together(left, right):
    return _charts_belong_together(left[0], right[0], DocumentIndex(left[1] + right[1]))


@pytest.mark.parametrize("context", ["FINANCIAL INFORMATION", "Annual review", "Selected measures"])
def test_repeated_section_heading_does_not_group_unrelated_metrics(context):
    inputs = [
        _series("retention", "Customer retention", page=2, context=context, unit="percent"),
        _series("delay", "Delivery delay", page=7, context=context, unit="days"),
        _series("volume", "Sales volume", page=11, context=context),
    ]
    index = DocumentIndex([item for _, values in inputs for item in values])
    groups = _group_chart_plans([chart for chart, _ in inputs], index)
    assert [[chart.id for chart in group] for group in groups] == [[chart.id] for chart, _ in inputs]


def test_specific_context_still_needs_source_or_metric_relationship():
    assert not _together(
        _series("delay", "Delivery delay", page=2, context="Service delivery"),
        _series("staff", "Staff headcount", page=7, context="Service delivery"),
    )


def test_same_page_different_tables_cannot_turn_section_into_relationship():
    assert not _together(
        _series("delay", "Delivery delay", context="Annual review", table="delays"),
        _series("staff", "Staff headcount", context="Annual review", table="staff"),
    )



def test_shared_table_and_generic_heading_do_not_group_unrelated_measures():
    assert not _together(
        _series("delay", "Delivery delay", context="FINANCIAL INFORMATION", table="summary"),
        _series("staff", "Staff headcount", context="FINANCIAL INFORMATION", table="summary"),
    )


def test_source_table_display_heading_is_not_a_retained_table_identity():
    left = _series("delay", "Delivery delay", page=2, context="Service delivery")
    right = _series("staff", "Staff headcount", page=7, context="Service delivery")
    for item in left[1] + right[1]:
        item.source_table = "Service delivery"
    assert not _together(left, right)

def test_source_table_context_preserves_structurally_related_measures_across_pages():
    assert _together(
        _series("inventory", "Inventory balance", page=2, context="Working capital", table="working-capital", unit="currency"),
        _series("receivables", "Trade receivables", page=3, context="Working capital", table="working-capital", unit="currency"),
    )


def test_source_table_snapshot_can_accompany_its_longer_trend():
    assert _together(
        _series("inventory", "Inventory balance", context="Working capital", table="working-capital",
                periods=("FY2016", "FY2017", "FY2018", "FY2019")),
        _series("receivables", "Trade receivables", context="Working capital", table="working-capital",
                periods=("FY2019",)),
    )


def test_source_table_context_does_not_bridge_disjoint_periods():
    assert not _together(
        _series("inventory", "Inventory balance", context="Working capital", table="working-capital", periods=("FY2017",)),
        _series("receivables", "Trade receivables", context="Working capital", table="working-capital", periods=("FY2019",)),
    )


def test_same_concrete_metric_can_group_across_source_pages():
    assert _together(_series("first", "Response duration", page=2, unit="days"),
                     _series("second", "Response duration", page=7, unit="days"))


def test_broad_canonical_metric_does_not_merge_distinct_qualified_measures():
    assert not _together(
        _series("amount", "Revenue", canonical="revenue", unit="currency"),
        _series("share", "Revenue: percentage of total", canonical="revenue", unit="percent"),
    )


def test_generic_metric_identity_remains_bound_to_its_source_table():
    assert not _together(
        _series("first", "Others", table="products", page=2),
        _series("second", "Others", table="regions", page=7),
    )


def test_first_metric_match_does_not_hide_other_plotted_metrics():
    left, left_values = _series("duration", "Response duration", page=2, unit="days")
    _, additional = _series("staff", "Staff headcount", page=2)
    left.observation_ids.extend(item.id for item in additional)
    assert not _together((left, left_values + additional), _series("other", "Response duration", page=7, unit="days"))


def test_chart_title_cannot_invent_metric_pairing():
    assert not _together(
        _series("delay", "Delivery delay", page=2, title="Gross margin"),
        _series("staff", "Staff headcount", page=7, title="Cost of sales"),
    )


def test_chart_title_cannot_invent_shared_concrete_terms_on_same_page():
    assert not _together(
        _series("delay", "Delivery delay", title="Service productivity"),
        _series("staff", "Staff headcount", title="Service productivity"),
    )


@pytest.mark.parametrize("left,right", [
    ("Gross profit margin", "Cost of sales: share of revenue"),
    ("Research and development expenses: share of revenue", "Selling and distribution expenses: share of revenue"),
    ("Sales volume", "Average selling price (ASP)"),
])
def test_known_complementary_metrics_survive_uninformative_chart_titles(left, right):
    assert _together(_series("first", left, page=2, title="Reported values"),
                     _series("second", right, page=7, title="Reported values"))


def test_concrete_metric_terms_on_shared_page_preserve_grouping():
    assert _together(_series("east", "Eastern warehouse utilization"),
                     _series("west", "Western warehouse utilization"))


def test_group_title_preserves_all_complete_metric_labels_over_old_length_limit():
    inputs = [_series(str(i), name, context="Working capital", table="working-capital")
              for i, name in enumerate(("Inventory held in regional warehouses",
                                       "Trade receivables from international customers",
                                       "Contract advances to third party suppliers"))]
    title = _chart_group_title([chart for chart, _ in inputs],
                               DocumentIndex([item for _, values in inputs for item in values]))
    for chart, _ in inputs:
        assert chart.title in title
    assert len(title) > 72


def test_group_title_deduplicates_same_metric_without_replacing_it_with_context():
    left = _series("east", "Customer retention", context="Annual review")
    right = _series("west", "Customer retention", context="Annual review")
    assert _chart_group_title([left[0], right[0]], DocumentIndex(left[1] + right[1])) == "Customer retention"


def test_group_title_removes_source_footnote_without_mutating_evidence():
    chart, values = _series("retention", "Customer retention(2)")
    assert _chart_group_title([chart], DocumentIndex(values)) == "Customer retention"
    assert all(item.metric_original == "Customer retention(2)" for item in values)


def test_fallback_paginates_before_complete_titles_are_shortened():
    import io
    from pptx import Presentation
    from adaptive_document_agent.agent.presentation_plan_recovery import PresentationPlanRecovery
    from adaptive_document_agent.models import DocumentProfile, ParsedDocument, PipelineResult
    from adaptive_document_agent.services.pptx_export import build_presentation
    names=('Inventory held in regional warehouses',
           'Trade receivables from international customers',
           'Contract advances to third party suppliers')
    inputs=[_series(str(i), name, context='Working capital', table='working-capital')
            for i,name in enumerate(names)]
    result=PipelineResult(document=ParsedDocument(document_id='synthetic',sha256='synthetic',
        safe_filename='study.pdf',page_count=20), profile=DocumentProfile(document_type='Document'),
        observations=[o for _,items in inputs for o in items], charts=[c for c,_ in inputs])
    result.presentation_plan=PresentationPlanRecovery().fallback(result)
    slides=[s for s in result.presentation_plan.slides if s.slide_type=='analysis']
    assert len(slides)>1
    for name in names:
        assert any(name in slide.title for slide in slides)
    payload=build_presentation(result)
    deck=Presentation(io.BytesIO(payload))
    titles=[shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame]
    assert all(any(name in text for text in titles) for name in names)
