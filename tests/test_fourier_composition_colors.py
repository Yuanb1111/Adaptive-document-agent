"""Fourier native chart styling must preserve editable, evidenced data."""

from io import BytesIO
from zipfile import ZipFile

import pytest
from pptx import Presentation

from adaptive_document_agent.models import ChartPlan, Observation, SourceEvidence
from adaptive_document_agent.services.composition_renderer import add_composition_chart
from adaptive_document_agent.services.fourier_brand import (
    ALL_CHART_COLORS, AMBER, BORDER, CYAN, MUTED, PRIMARY, TECH_BLUE, TEXT, WHITE,
    label_color,
)


def _chart(kind, *, colors=None, categories=("Alpha", "Beta", "Gamma")):
    periods = ("FY2024",) if kind == "doughnut" else ("FY2023", "FY2024")
    observations = [Observation(
        id=f"{period}-{category}", metric_original="Output share", value=100 / len(categories),
        raw_value=f"{100 / len(categories)}%", unit="percent", unit_family="percentage",
        period=period, dimensions={"channel": category}, category_dimensions={"channel": category},
        confidence=0.95, evidence=[SourceEvidence(page=2, extraction_method="test", confidence=0.95)],
    ) for period in periods for category in categories]
    plan = ChartPlan(id="composition", title="Output share", question="Reported output distribution",
                     chart_type=kind, series_dimension="channel", show_data_labels=True,
                     observation_ids=[item.id for item in observations], source_pages=[2])
    original = [item.model_dump() for item in observations]
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    if colors is not None:
        slide._ada_colors = dict(colors)
    add_composition_chart(slide, plan, observations, (0.5, 0.5, 8, 4))
    assert [item.model_dump() for item in observations] == original
    output = BytesIO()
    deck.save(output)
    restored = Presentation(BytesIO(output.getvalue()))
    native = next(shape.chart for shape in restored.slides[0].shapes if shape.has_chart)
    return native, output.getvalue()


def _fills(chart, kind):
    items = chart.series[0].points if kind == "doughnut" else chart.series
    return [str(item.format.fill.fore_color.rgb) for item in items]


@pytest.mark.parametrize("kind", ["stacked_bar", "stacked_percent", "doughnut"])
def test_composition_brand_colors_and_label_contrast_survive_roundtrip(kind):
    colors = {"Alpha": PRIMARY, "Beta": AMBER, "Gamma": CYAN}
    chart, output = _chart(kind, colors=colors)
    assert _fills(chart, kind) == [PRIMARY, AMBER, CYAN]
    assert chart.chart_style is None
    assert chart.font.name == chart.legend.font.name == "Arial"
    assert str(chart.legend.font.color.rgb) == MUTED
    assert chart.plots[0].has_data_labels
    for series in chart.series:
        assert str(series.format.fill.fore_color.rgb) in ALL_CHART_COLORS
        assert series._element.xpath("c:spPr/a:ln/a:noFill")
    if kind == "doughnut":
        labels = [point.data_label for point in chart.series[0].points]
        assert len(chart.series[0]._element.xpath("c:dPt/c:spPr/a:ln/a:noFill")) == 3
        for label in labels:
            assert label._dLbl.xpath("c:showVal")[0].get("val") == "0"
            assert label._dLbl.xpath("c:showPercent")[0].get("val") == "1"
            assert not label.has_text_frame  # PowerPoint still derives labels from the data.
        assert chart.series[0].data_labels.show_percentage
        assert not chart.series[0].data_labels.show_value
        assert list(chart.series[0].values) == pytest.approx([100 / 3] * 3)
    else:
        labels = [series.data_labels for series in chart.series]
        assert all(label.show_value and not label.show_percentage for label in labels)
        for axis in (chart.category_axis, chart.value_axis):
            assert axis.tick_labels.font.name == "Arial"
            assert str(axis.tick_labels.font.color.rgb) == MUTED
            assert str(axis.format.line.color.rgb) == BORDER
        expected = 1 / 3 if kind == "stacked_percent" else 100 / 3
        assert all(list(series.values) == pytest.approx([expected, expected]) for series in chart.series)
    assert [str(label.font.color.rgb) for label in labels] == [WHITE, TEXT, TEXT]
    assert all(label.font.name == "Arial" for label in labels)
    with ZipFile(BytesIO(output)) as package:
        embedded = [name for name in package.namelist() if name.startswith("ppt/embeddings/")]
        assert len(embedded) == 1
        with ZipFile(BytesIO(package.read(embedded[0]))) as workbook:
            assert "xl/worksheets/sheet1.xml" in workbook.namelist()


@pytest.mark.parametrize("kind", ["stacked_bar", "stacked_percent", "doughnut"])
def test_invalid_legacy_color_map_is_repaired_without_duplicate_series_colors(kind):
    chart, _ = _chart(kind, colors={name: "5B21B6" for name in ("Alpha", "Beta", "Gamma")})
    fills = _fills(chart, kind)
    assert len(set(fills)) == 3
    assert set(fills) <= set(ALL_CHART_COLORS)
    assert "5B21B6" not in fills


@pytest.mark.parametrize("kind", ["stacked_bar", "stacked_percent", "doughnut"])
def test_canonical_deck_assignments_are_not_recolored_by_chart_subset(kind):
    colors = {"Alpha": AMBER, "Beta": PRIMARY, "Gamma": CYAN,
              "Delta": TECH_BLUE, "Unrelated": PRIMARY}
    full, _ = _chart(kind, colors=colors)
    subset, _ = _chart(kind, colors=colors, categories=("Beta", "Delta", "Gamma"))
    assert _fills(full, kind) == [AMBER, PRIMARY, CYAN]
    assert _fills(subset, kind) == [PRIMARY, TECH_BLUE, CYAN]


def test_eight_category_composition_keeps_every_category_with_approved_colors():
    chart, _ = _chart("doughnut", categories=tuple(f"Category {index}" for index in range(8)))
    fills = _fills(chart, "doughnut")
    assert len(set(fills)) == 8
    assert set(fills) <= set(ALL_CHART_COLORS)
    assert list(chart.series[0].values) == [12.5] * 8
    assert [str(point.data_label.font.color.rgb) for point in chart.series[0].points] == [
        label_color(color) for color in fills
    ]


def test_foreign_color_is_not_allowed_to_leak_through_a_saved_map():
    chart, _ = _chart("stacked_bar", colors={"Alpha": "123456", "Beta": "#f4b923", "Gamma": "ab74ff"})
    fills = _fills(chart, "stacked_bar")
    assert set(fills) <= set(ALL_CHART_COLORS)
    assert len(set(fills)) == 3
