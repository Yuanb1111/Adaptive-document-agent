"""Long branded headings must leave room for complete supporting evidence."""

from io import BytesIO
import json

from pptx import Presentation

from adaptive_document_agent.models import ChartPlan, PresentationSlide, PresentationVisualBlock
from adaptive_document_agent.services.pptx_export import build_presentation
from adaptive_document_agent.services.presentation_conventions import signed_expense_note
from tests.test_p0_composition import observation, result_for


def test_expense_convention_does_not_split_a_small_support_series():
    periods = ("FY2021", "FY2022", "FY2023")
    observations, charts = [], []
    for identifier, metric in (("selling", "Selling and distribution expenses"),
                               ("research", "Research and development expenses")):
        values = [observation(f"{identifier}-{i}", metric, -30 - i, period, unit="percent")
                  for i, period in enumerate(periods)]
        observations.extend(values)
        charts.append(ChartPlan(id=identifier, title=metric, question="Reported expense ratios",
                                chart_type="bar", observation_ids=[o.id for o in values], source_pages=[3]))
    support = [observation(f"loss-{i}", "Loss margin", -40 - i, period, unit="percent")
               for i, period in enumerate(periods)]
    observations.extend(support)
    slide_plan = PresentationSlide(id="expense", slide_type="analysis", layout="two_up",
        title="Selling and research expense ratios remained significant relative to revenue, and the loss margin widened.",
        section_title="Operating expense ratios and loss margin",
        message="How did the reported expense ratios and loss margins change?",
        chart_ids=[c.id for c in charts], source_pages=[3],
        visual_blocks=[PresentationVisualBlock(role="table", observation_ids=[o.id for o in support])])
    result = result_for(observations, charts, slide_plan)
    original = [o.model_dump() for o in observations]
    output = Presentation(BytesIO(build_presentation(result)))
    slides = [s for s in output.slides if s.name.startswith("composed_")]
    assert len(slides) == 1
    slide = slides[0]
    tables = [shape for shape in slide.shapes if shape.has_table]
    cells = {cell.text for shape in tables for row in shape.table.rows for cell in row.cells}
    assert set(periods) <= cells
    note = next(shape for shape in slide.shapes
                if shape.has_text_frame and shape.text == signed_expense_note(observations))
    assert note.text_frame.paragraphs[0].font.size.pt == 12
    assert note.width.inches > 10
    assert all(shape.top + shape.height <= note.top for shape in tables)
    assert [o.model_dump() for o in observations] == original


def test_long_composed_subtitle_moves_complete_following_sentence_to_commentary():
    observations = [
        observation(f"revenue-{period}", "Revenue", value, period)
        for period, value in (("FY2021", 174.3), ("FY2023", 286.7))
    ]
    observations += [
        observation(f"cost-{period}", "Cost of sales", value, period)
        for period, value in (("FY2021", 100.0), ("FY2023", 170.0))
    ]
    charts = [ChartPlan(
        id=metric, title=metric, question="Reported values", chart_type="bar",
        observation_ids=[item.id for item in observations if item.metric_original == metric],
        source_pages=[3],
    ) for metric in ("Revenue", "Cost of sales")]
    first = ("Revenue rose from 174.3 in FY2021 to 286.7 in FY2023, "
             "with an increasing trend and no turning points.")
    second = "The expansion occurred alongside higher cost of sales and operating expenditure."
    message = first + " " + second
    slide_plan = PresentationSlide(
        id="revenue", slide_type="analysis", layout="two_up",
        title="Revenue increased over FY2021–FY2023", section_title="Revenue",
        message=message, chart_ids=[chart.id for chart in charts], source_pages=[3],
    )
    deck = Presentation(BytesIO(build_presentation(result_for(observations, charts, slide_plan))))
    slide = next(item for item in deck.slides if item.name.startswith("composed_"))
    subtitle = next(ph.text for ph in slide.placeholders if ph.placeholder_format.idx == 16)
    commentary = "\n".join(shape.text for shape in slide.shapes if shape.name.startswith("composed:text:"))
    assert subtitle == first
    assert second in commentary
    assert json.loads(slide.notes_slide.notes_text_frame.text)["analytical_question"] == message
