"""Check branded output against the preserved official template and source data."""

from hashlib import sha256
from io import BytesIO
import json

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from adaptive_document_agent.models import ChartPlan
from adaptive_document_agent.services.fourier_brand import ALL_CHART_COLORS, PRIMARY, TEXT
from adaptive_document_agent.services.pptx_export import BUNDLED_TEMPLATE_PATH, build_presentation
from adaptive_document_agent.services.presentation_brand_qa import validate_generated_brand
from tests.test_p0_composition import observation, paired_result, result_for


def _artwork(presentation):
    def pictures(shapes):
        found = []
        for shape in shapes:
            if hasattr(shape, "image"):
                found.append((shape.name, shape.left, shape.top, shape.width, shape.height, shape.rotation,
                              shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom,
                              sha256(shape.image.blob).hexdigest()))
            if hasattr(shape, "shapes"):
                found.extend(pictures(shape.shapes))
        return found

    return {
        "masters": [pictures(master.shapes) for master in presentation.slide_masters],
        "layouts": [pictures(layout.shapes) for layout in presentation.slide_layouts],
    }


@pytest.fixture(scope="module")
def branded_output():
    result = paired_result()
    result.charts[0].title = "Net energy balance"
    result.charts[1].chart_type = "line"
    analysis = result.presentation_plan.slides[3]
    analysis.title = "Operating measures"
    analysis.message = "Comparable reported energy balances and service activity"
    analysis.bullets = ["The chart retains each disclosed period and its original signed balance."]
    for item in result.observations:
        if item.id.startswith("a"):
            item.metric_original = "Net energy balance"
            item.value = -item.value
            item.raw_value = str(item.value)
            item.evidence[0].text = item.raw_value
            item.evidence[0].row_label = item.metric_original
    before = result.model_dump()
    template_hash = sha256(BUNDLED_TEMPLATE_PATH.read_bytes()).hexdigest()
    reference = Presentation(str(BUNDLED_TEMPLATE_PATH))
    output = Presentation(BytesIO(build_presentation(result, BUNDLED_TEMPLATE_PATH)))
    return result, before, reference, output, template_hash


def test_official_template_artwork_and_file_are_unchanged(branded_output):
    _, _, reference, output, template_hash = branded_output
    expected = _artwork(reference)
    assert any(expected["masters"]) or any(expected["layouts"])
    assert _artwork(output) == expected
    assert output.slide_width == reference.slide_width
    assert output.slide_height == reference.slide_height
    assert sha256(BUNDLED_TEMPLATE_PATH.read_bytes()).hexdigest() == template_hash


def test_exported_content_headers_match_template_typography(branded_output):
    _, _, _, output, _ = branded_output
    slide = next(slide for slide in output.slides if slide.name.startswith("composed_"))
    for text, size, color in (
        ("Operating measures", 32, TEXT),
        ("Comparable reported energy balances and service activity", 18, PRIMARY),
    ):
        shape = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text.startswith(text))
        paragraph = shape.text_frame.paragraphs[0]
        assert paragraph.font.name == "Arial"
        assert paragraph.font.size.pt == size
        assert str(paragraph.font.color.rgb) == color
    validate_generated_brand(output)


def test_signed_bars_and_line_markers_use_explicit_approved_colors(branded_output):
    _, _, _, output, _ = branded_output
    charts = {shape.name: shape.chart for slide in output.slides for shape in slide.shapes if shape.has_chart}
    bar, line = charts["chart:a"], charts["chart:b"]
    assert bar.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    assert list(bar.series[0].values) == [-100, -120, -140]
    assert bar.series[0].invert_if_negative is False
    assert line.chart_type == XL_CHART_TYPE.LINE_MARKERS
    assert list(line.series[0].values) == [60, 70, 80]
    for chart in (bar, line):
        series = chart.series[0]
        color = str(series.format.fill.fore_color.rgb)
        assert color in ALL_CHART_COLORS
        assert str(series.format.line.color.rgb) == color
    assert str(line.series[0].marker.format.fill.fore_color.rgb) == str(line.series[0].format.fill.fore_color.rgb)
    assert str(line.series[0].marker.format.line.color.rgb) == str(line.series[0].format.fill.fore_color.rgb)


def test_brand_export_preserves_source_evidence_and_analytical_models(branded_output):
    result, before, _, _, _ = branded_output
    after = result.model_dump()
    # The existing presentation compiler fills company overview source pages.
    # Styling must leave the source, analysis and all raw evidence untouched.
    for key in before.keys() - {"presentation_plan", "presentation_export_trace"}:
        assert after[key] == before[key], key
    assert after['presentation_export_trace']
    assert all(item['mapping_basis'].startswith('renderer_input_scope') for item in after['presentation_export_trace'])
    assert after["presentation_plan"]["slides"] == before["presentation_plan"]["slides"]


def test_full_pie_export_uses_distinct_brand_colors_without_changing_values():
    observations = [observation(f"share-{name}", "Output share", value, unit="percent", category=name)
                    for name, value in (("Alpha", 40), ("Beta", 35), ("Gamma", 25))]
    chart = ChartPlan(id="parts", title="Output share", question="Disclosed category shares",
                      chart_type="pie", available_chart_types=["pie", "table"], x_dimension="segment",
                      observation_ids=[item.id for item in observations], source_pages=[3])
    result = result_for(observations, [chart])
    original = [item.model_dump() for item in observations]
    output = Presentation(BytesIO(build_presentation(result, BUNDLED_TEMPLATE_PATH)))
    pie = next(shape.chart for slide in output.slides for shape in slide.shapes if shape.has_chart)
    assert pie.chart_type == XL_CHART_TYPE.PIE
    assert list(pie.series[0].values) == [40, 35, 25]
    colors = [str(point.format.fill.fore_color.rgb) for point in pie.series[0].points]
    assert len(set(colors)) == 3
    assert set(colors) <= set(ALL_CHART_COLORS)
    assert [item.model_dump() for item in observations] == original
    validate_generated_brand(output)


@pytest.mark.parametrize("chart_type", ["bar", "line"])
def test_overlapping_category_subsets_keep_same_native_series_color(chart_type):
    observations = [observation(f"{period}-{category}", "Output", value + offset, period,
                                unit="count", category=category)
                    for offset, period in enumerate(("FY2023", "FY2024"))
                    for category, value in (("East", 40), ("West", 50), ("Zulu", 60))]
    charts = [ChartPlan(id=identifier, title="Output", question="Disclosed category output",
                        chart_type=chart_type, series_dimension="segment",
                        observation_ids=[item.id for item in observations
                                         if item.dimensions["segment"] in categories], source_pages=[3])
              for identifier, categories in (("a", ("East", "West")), ("b", ("West", "Zulu")))]
    result = result_for(observations, charts)
    output = Presentation(BytesIO(build_presentation(result, BUNDLED_TEMPLATE_PATH)))
    native = {shape.name: {series.name: str(series.format.fill.fore_color.rgb)
                          for series in shape.chart.series}
              for slide in output.slides for shape in slide.shapes if shape.has_chart}
    common = native["chart:a"].keys() & native["chart:b"].keys()
    assert len(common) == 1
    for name in common:
        assert native["chart:a"][name] == native["chart:b"][name]
    assert all(len(set(colors.values())) == 2 for colors in native.values())


@pytest.mark.parametrize("claim,placement", [
    (
        "Operating activity expanded across the comparable reporting periods, with shipment "
        "and service observations supporting the scale of the change",
        "subtitle",
    ),
    (
        "Operating activity expanded across the comparable reporting periods, while the retained "
        "shipment and service observations provide separate evidence for the scale and timing of this change",
        "commentary",
    ),
])
def test_long_claim_keeps_full_text_under_semantic_section_heading(claim, placement):
    result = paired_result()
    slide_plan = result.presentation_plan.slides[3]
    slide_plan.title = claim
    slide_plan.section_title = "Operating activity"
    slide_plan.message = "How did shipments and service activity change across the reported periods?"
    before_slides = [slide.model_dump() for slide in result.presentation_plan.slides]
    before_evidence = [item.model_dump() for item in result.observations]
    output = Presentation(BytesIO(build_presentation(result, BUNDLED_TEMPLATE_PATH)))
    slide = next(slide for slide in output.slides if slide.name.startswith("composed_"))
    title = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text == slide_plan.section_title)
    assert title.text_frame.paragraphs[0].font.size.pt == 32
    assert str(title.text_frame.paragraphs[0].font.color.rgb) == TEXT
    assert title.text_frame.paragraphs[0]._p.pPr.get("marL") == "0"
    assert title.text_frame.paragraphs[0]._p.pPr.get("indent") == "0"
    subtitle = next(shape for shape in slide.placeholders if shape.placeholder_format.idx == 16)
    if placement == "subtitle":
        assert subtitle.text == claim
        assert subtitle.text_frame.paragraphs[0].font.size.pt == 18
        assert str(subtitle.text_frame.paragraphs[0].font.color.rgb) == PRIMARY
    else:
        assert subtitle.text == ""
        commentary = next(shape for shape in slide.shapes
                          if shape.has_text_frame and claim in shape.text)
        assert commentary.name.startswith("composed:text:")
        assert commentary.text_frame.paragraphs[0].font.size.pt == 16
        assert str(commentary.text_frame.paragraphs[0].font.color.rgb) == TEXT
        assert "Both measures provide complementary evidence of operating activity." in commentary.text
    notes = json.loads(slide.notes_slide.notes_text_frame.text)
    assert notes["analytical_question"] == slide_plan.message
    assert notes["planned_title"] == slide_plan.title
    assert [slide.model_dump() for slide in result.presentation_plan.slides] == before_slides
    assert [item.model_dump() for item in result.observations] == before_evidence


def test_long_selected_topic_heading_keeps_complete_claim_visible():
    result = paired_result()
    slide_plan = result.presentation_plan.slides[3]
    claim = (
        "The reported operating activity expanded across the comparable annual periods, while shipment "
        "and service observations show the scale of the change and retain their separate reporting definitions"
    )
    slide_plan.title = claim
    slide_plan.section_title = claim
    snapshot = slide_plan.model_dump()
    deck = Presentation(BytesIO(build_presentation(result, BUNDLED_TEMPLATE_PATH)))
    slides = [s for s in deck.slides if s.name.startswith("composed_")]
    assert any(shape.text == "Units shipped and Service hours" for s in slides for shape in s.shapes if shape.has_text_frame)
    assert all(shape.text != "Reported measures" for s in slides for shape in s.shapes if shape.has_text_frame)
    visible = "\n".join(shape.text for s in slides for shape in s.shapes if shape.has_text_frame)
    assert claim in visible
    assert slide_plan.model_dump() == snapshot
