"""Comprehensive test suite for generic improvements across Adaptive Document Agent.

Covers:
1. Amount vs percentage classification (monetary items never coerced to %)
2. Unit-aware movement wording (days -> days, ratio -> x, margin -> pp, currency -> currency)
3. Full trajectory semantics (peak/trough detection across series, e.g. 30 -> 140 -> 260 -> 4.5)
4. Company profile validation and cleanup (rejecting dangling connectors, incomplete periods)
5. Company Overview vs Data Quality consistency (resolving external identity citations)
6. Title quality polishing (sanitizing mechanical 'Trajectory' suffixes)
7. PPT preflight QA checks (impossible percentages, wrong units, empty slides)
"""

from __future__ import annotations

import pytest
from decimal import Decimal

from adaptive_document_agent.models import (
    CanonicalFact,
    CompanyProfile,
    DocumentProfile,
    Observation,
    ParsedDocument,
    PipelineResult,
    PresentationPlan,
    PresentationSlide,
    SourceEvidence,
)
from adaptive_document_agent.document_model.metric_semantic_classifier import (
    classify_metric,
    format_metric_display_value,
    MetricSemantic,
)
from adaptive_document_agent.services.company_extractor import (
    _clean_text_fragment,
    validate_company_name,
    validate_product,
    validate_track_record_period,
    validate_business_model,
)
from adaptive_document_agent.services.language_qa import (
    clean_metric_label,
    polish_slide_title,
)
from adaptive_document_agent.services.movement_formatter import (
    analyze_trajectory,
    format_movement_headline,
    format_movement_narrative,
)
from adaptive_document_agent.services.financial_normalizer import FinancialNormalizer
from adaptive_document_agent.services.qa_reporter import (
    sanitize_company_identity_contradictions,
)
from adaptive_document_agent.agent.insight_generator import InsightGenerator
from adaptive_document_agent.models import AnalysisResult
from adaptive_document_agent.validation.claim_validator import (
    MetricSemanticFamily,
    classify_metric_semantic_family,
)


# -----------------------------------------------------------------------------
# 1. Amount vs Percentage Classification
# -----------------------------------------------------------------------------

def test_monetary_metrics_never_coerced_to_percentage():
    """Verify monetary statement metrics are never formatted as percentage."""
    # Classify Revenue: should be currency
    sem = classify_metric("Revenue")
    assert sem.is_currency is True
    assert sem.is_percentage is False

    # format_metric_display_value should guard against formatting with % even if is_percentage is mistakenly set
    spurious_pct_semantic = MetricSemantic(
        metric_type="percentage",
        unit_family="percentage",
        semantic_type="monetary_amount",
        clean_name="Revenue",
        display_unit="%",
        is_currency=False,
        is_percentage=True,
        is_multiple=False,
    )
    formatted = format_metric_display_value(
        "10,350,986",
        10350986.0,
        spurious_pct_semantic,
        currency="RMB",
    )
    assert "%" not in formatted
    assert "10,350,986" in formatted


def test_metric_semantic_family_classification():
    """Verify directional families remain distinct from unit classification."""
    assert classify_metric_semantic_family("Revenue") == MetricSemanticFamily.GENERIC
    assert classify_metric_semantic_family("Net Profit") == MetricSemanticFamily.PROFIT_LOSS
    assert classify_metric_semantic_family("Gross Margin") == MetricSemanticFamily.RATIO
    assert classify_metric_semantic_family("Current Ratio") == MetricSemanticFamily.MULTIPLE
    assert classify_metric_semantic_family("Inventory Turnover Days") == MetricSemanticFamily.DAYS
    assert classify_metric_semantic_family("Trade Receivables Turnover Days") == MetricSemanticFamily.DAYS


def test_financial_normalizer_intrinsic_days_precedes_inventory_amount_semantics():
    """A days metric containing a balance-sheet keyword must remain a days metric."""
    obs = Observation(
        id="inventory-days",
        metric_original="Inventory turnover days",
        value=31.2,
        raw_value="31.2",
        unit="percent",
        currency="RMB",
        period="FY2024",
        confidence=0.9,
    )
    normalized = FinancialNormalizer.normalize_observation(obs)
    assert normalized.unit == "days"
    assert normalized.unit_family == "days"
    assert normalized.currency is None


# -----------------------------------------------------------------------------
# 2. Unit-Aware Movement Wording
# -----------------------------------------------------------------------------

def test_days_metric_movement_uses_days_never_pp():
    """Turnover days movement must be formatted as 'days', never 'pp' or '%'."""
    headline = format_movement_headline(
        "Inventory Turnover Days",
        start_val=30.0,
        end_val=45.0,
        unit="days",
        start_period="FY2021",
        end_period="FY2022",
    )
    assert "days" in headline.lower()
    assert "pp" not in headline.lower()
    assert "%" not in headline

    narrative = format_movement_narrative(
        "Inventory Turnover Days",
        start_val=45.0,
        end_val=35.0,
        unit="days",
        start_period="FY2021",
        end_period="FY2022",
    )
    assert "10.0 days" in narrative
    assert "pp" not in narrative


def test_ratio_multiple_metric_movement_uses_x_never_pp():
    """Ratios like Current Ratio must use 'x', never 'pp'."""
    narrative = format_movement_narrative(
        "Current Ratio",
        start_val=1.2,
        end_val=1.5,
        unit="ratio",
        start_period="FY2021",
        end_period="FY2022",
    )
    assert "0.30x" in narrative
    assert "pp" not in narrative
    assert "%" not in narrative


def test_margin_metric_movement_uses_pp():
    """Margins like Gross Margin must use 'pp' for percentage-point changes."""
    narrative = format_movement_narrative(
        "Gross Margin",
        start_val=25.0,
        end_val=30.0,
        unit="percent",
        start_period="FY2021",
        end_period="FY2022",
    )
    assert "5.0 pp" in narrative


def test_currency_metric_movement_uses_currency_amount():
    """Monetary metrics must express change as currency amounts, not raw % without context."""
    narrative = format_movement_narrative(
        "Revenue",
        start_val=100.0,
        end_val=120.0,
        unit="RMB million",
        start_period="FY2021",
        end_period="FY2022",
    )
    assert "RMB20m" in narrative
    assert "+20.0%" in narrative


# -----------------------------------------------------------------------------
# 3. Full Trajectory Semantics
# -----------------------------------------------------------------------------

def test_analyze_trajectory_detects_peak_and_sharp_fall():
    """Series 30 -> 140 -> 260 -> 4.5 must detect the peak and sharp contraction."""
    vals = [30.0, 140.0, 260.0, 4.5]
    periods = ["FY2020", "FY2021", "FY2022", "FY2023"]

    info = analyze_trajectory(vals, periods)
    assert info["shape"] == "rose_then_fell_sharply"
    assert info["peak_val"] == 260.0
    assert info["peak_period"] == "FY2022"
    assert "peaked at 260.0" in info["description"].lower()


def test_format_movement_narrative_with_full_series():
    """Narrative formatting with intermediate values includes trajectory context."""
    vals = [30.0, 140.0, 260.0, 4.5]
    periods = ["FY2020", "FY2021", "FY2022", "FY2023"]

    narrative = format_movement_narrative(
        "Operating Cash Flow",
        start_val=30.0,
        end_val=4.5,
        unit="RMB million",
        start_period="FY2020",
        end_period="FY2023",
        values=vals,
        periods=periods,
    )
    assert "peaked" in narrative.lower()
    assert "FY2022" in narrative


def test_analyze_trajectory_detects_trough_and_plateau_shift():
    trough = analyze_trajectory([100.0, 40.0, 85.0], ["FY2022", "FY2023", "FY2024"])
    assert trough["shape"] == "fell_then_partially_recovered"
    assert trough["trough_period"] == "FY2023"

    plateau = analyze_trajectory([100.0, 101.0, 125.0], ["FY2022", "FY2023", "FY2024"])
    assert plateau["shape"] == "stable_then_increasing"


def test_deterministic_insight_uses_full_compare_period_trajectory():
    evidence = SourceEvidence(page=7, extraction_method="digital_table", confidence=0.9)
    result = AnalysisResult(
        task_id="cash-flow-series",
        title="Operating cash flow",
        result=[
            {"period": "FY2020", "value": 30.0},
            {"period": "FY2021", "value": 140.0},
            {"period": "FY2022", "value": 260.0},
            {"period": "FY2023", "value": 4.5},
        ],
        evidence=[evidence],
        confidence=0.9,
    )
    insight = InsightGenerator().generate([result])[0]
    assert "peaked at 260.0 in FY2022" in insight.movement
    assert "falling sharply" in insight.movement


# -----------------------------------------------------------------------------
# 4. Company Profile Cleanup
# -----------------------------------------------------------------------------

def test_clean_text_fragment_strips_dangling_connectors():
    """Dangling prepositions and conjunctions must be stripped."""
    assert _clean_text_fragment("leading supplier of fresh ingredients to") == "leading supplier of fresh ingredients"
    assert _clean_text_fragment("operates a global network and") == "operates a global network"
    assert _clean_text_fragment("specialized products including") == "specialized products"
    assert _clean_text_fragment("financial results from") == "financial results"


def test_validate_track_record_period_rejects_incomplete_ranges():
    """Incomplete track record period ending in 'to' or 'from' must be rejected."""
    assert validate_track_record_period("FY2021 to") == ""
    assert validate_track_record_period("from FY2021") == ""
    assert validate_track_record_period("2021 -") == ""
    assert validate_track_record_period("FY2021 to FY2023") == "FY2021 to FY2023"
    assert validate_track_record_period("2020 - 2023") == "2020 - 2023"


def test_validate_product_cleans_hanging_connectors():
    """Products ending in hanging connectors must have connectors stripped."""
    assert validate_product("Fruit drinks and teas and") == "Fruit drinks and teas"
    assert validate_product("Ice cream products including") == "Ice cream products"


def test_company_narrative_validators_reject_broken_prefixes():
    assert validate_business_model("d excerpts from the issuer profile") == ""


# -----------------------------------------------------------------------------
# 5. Company Overview vs Data Quality Consistency
# -----------------------------------------------------------------------------

def test_reconcile_company_identity_outside_financial_range():
    """When company identity is resolved outside core analysis pages, note is updated accurately."""
    profile = DocumentProfile(
        title="Prospectus",
        document_type="prospectus",
        analysis_page_ranges=[(210, 245)],
        data_quality_notes=["Issuer name is not stated in financial tables; document analysed as general corporate document."],
    )
    company = CompanyProfile(
        name="Global Beverage Holdings Limited",
        identity_state="RESOLVED",
        source_pages=[1, 5],
        field_source_pages={"name": [1, 5]},
    )
    plan = PresentationPlan(
        title="Financial Overview",
        company=company,
        slides=[],
    )
    result = PipelineResult(
        document=ParsedDocument(
            document_id="doc-1",
            sha256="abc",
            safe_filename="prospectus.pdf",
            page_count=0,
        ),
        profile=profile,
        extracted_facts=[],
        canonical_facts=[],
        presentation_plan=plan,
    )

    fixes = sanitize_company_identity_contradictions(result)
    assert any(f.code == "company_identity_reconciled" for f in fixes)
    assert any("verified from p. 1, p. 5 outside the financial analysis range" in n for n in result.profile.data_quality_notes)
    assert not any("issuer name is not stated" in n.lower() for n in result.profile.data_quality_notes)


# -----------------------------------------------------------------------------
# 6. Title Quality Polishing
# -----------------------------------------------------------------------------

def test_polish_slide_title_replaces_mechanical_trajectory():
    """Short noun-phrase titles ending in 'Trajectory' must become 'Overview'."""
    assert polish_slide_title("Revenue Trajectory") == "Revenue Overview"
    assert polish_slide_title("Gross Profit Trajectory") == "Gross Profit Overview"
    assert polish_slide_title("Selling and Marketing Expenses Rose, With Trajectory") == "Selling and Marketing Expenses Rose"
    assert polish_slide_title("Trajectory") == "Overview"


# -----------------------------------------------------------------------------
# 7. PPT Preflight QA Checks
# -----------------------------------------------------------------------------

def test_preflight_impossible_percentages_and_units():
    """PresentationPreflight catches impossible percentages and wrong units."""
    from pptx import Presentation
    from pptx.util import Inches
    from adaptive_document_agent.services.ppt_preflight import PresentationPreflight

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_layout = prs.slide_layouts[6]  # Blank slide
    slide = prs.slides.add_slide(slide_layout)

    # Add text box with impossible percentage and days with pp
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    tf = txBox.text_frame
    p1 = tf.paragraphs[0]
    p1.text = "Revenue: 10,350,986%"
    p2 = tf.add_paragraph()
    p2.text = "Inventory turnover days decreased by -5.0 pp"

    # Add official thank you slide to avoid missing thank you warning
    ty_slide = prs.slides.add_slide(slide_layout)
    ty_box = ty_slide.shapes.add_textbox(Inches(2), Inches(2), Inches(6), Inches(2))
    ty_box.text_frame.text = "Thank you for your attention"

    preflight = PresentationPreflight(prs)
    issues = preflight.validate_and_sanitize()

    # Check that issues were detected and sanitized
    issue_codes = {issue.code for issue in issues}
    assert "impossible_percentage" in issue_codes
    assert "invalid_days_change" in issue_codes

    # Verify in-place sanitization
    assert "10,350,986%" not in p1.text
    assert "Revenue: 10,350,986" in p1.text
    assert "pp" not in p2.text
    assert "days" in p2.text


def test_preflight_flags_title_direction_that_conflicts_with_chart_data():
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    from adaptive_document_agent.services.ppt_preflight import PresentationPreflight

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Revenue Grew"
    data = ChartData()
    data.categories = ["FY2023", "FY2024"]
    data.add_series("Revenue", (120.0, 90.0))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(4), data)
    issues = PresentationPreflight(prs).validate_and_sanitize()
    assert any(issue.code == "title_data_misalignment" for issue in issues)
