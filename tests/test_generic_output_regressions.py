"""Output regressions use synthetic metrics, not issuer-specific rules."""
import pytest
from pptx import Presentation
from pptx.util import Inches
from adaptive_document_agent.services.ppt_preflight import PresentationPreflight
from adaptive_document_agent.validation.claim_validator import associate_clause_directions, extract_metric_aliases, repair_presentation_plan
from adaptive_document_agent.models import PresentationPlan, PresentationSlide
from tests.test_p0_composition import observation

@pytest.mark.parametrize("word", ["fulfillment", "installment", "debugging", "JSONata"])
def test_internal_word_filter_preserves_embedded_substrings(word):
    deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[6])
    box=slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text=f"{word} LLM json"
    PresentationPreflight(deck)._check_banned_phrases(0, slide)
    assert box.text == word


def test_unknown_conjunct_does_not_borrow_known_subject():
    assocs=associate_clause_directions("Units shipped increased and processing latency declined", {"Units shipped":["Units shipped"]})
    assert assocs == [("Units shipped", "increased")]


def test_average_price_alias_and_independent_predicates():
    obs=[observation(f"{p}{i}", name, value, f"FY{2020+i}") for p,name,values in
         [("v","Sales volume",(300,900)),("p","ASP",(80,50))] for i,value in enumerate(values)]
    text="Sales volume rose and average selling price declined."
    slide=PresentationSlide(id="s",slide_type="analysis",title=text,observation_ids=[o.id for o in obs])
    plan=PresentationPlan(title="Review",slides=[slide])
    assert not repair_presentation_plan(plan,obs)[1]
    assert slide.title == text
    slide.title="Sales volume declined and average selling price increased."
    repair_presentation_plan(plan,obs)
    assert slide.title == "Sales volume increased and average selling price decreased."


def test_mixed_half_year_header_and_wrapped_qualifiers_keep_all_columns():
    from types import SimpleNamespace
    from adaptive_document_agent.extraction.borderless_table_extractor import BorderlessTableExtractor
    source="""Allocation schedule    2028H2    2029    2030    2031
(in USD millions)
Engineering investment
    (Platform design and
    equipment
    renewal)     12    24    36    36
Operations investment
    (Process
    maintenance)    5    10    15    15
"""
    tables=BorderlessTableExtractor().extract(SimpleNamespace(text=source), 7)
    rows=[r for t in tables for r in t.rows if any(c and c[0].isdigit() for c in r.cells[1:])]
    assert rows[0].cells[0] == "Engineering investment (Platform design and equipment renewal)"
    assert rows[1].cells[0] == "Operations investment (Process maintenance)"
    assert rows[0].cells[1:] == ["12", "24", "36", "36"]
    assert all(t.column_periods == [None,"2H2028","2029","2030","2031"] for t in tables)


def test_unknown_period_cells_and_raw_unit_variants_are_not_merged():
    from adaptive_document_agent.document_model.builder import DocumentModelBuilder
    first = observation("a", "Allocated capacity", 36, None)
    second = first.model_copy(deep=True, update={"id": "b"})
    assert len(DocumentModelBuilder().build([first, second]).observations) == 2
    first.period = second.period = "FY2029"
    second.raw_unit = "count"
    first.raw_unit = "units"
    assert len(DocumentModelBuilder().build([first, second]).observations) == 2


def test_dedup_retains_both_sources_without_mutating_input():
    from adaptive_document_agent.document_model.builder import DocumentModelBuilder
    first = observation("a", "Allocated capacity", 36, "FY2029")
    second = first.model_copy(deep=True, update={"id": "b", "confidence": .99})
    second.evidence[0].page = first.evidence[0].page + 1
    merged = DocumentModelBuilder().build([first, second]).observations
    assert len(merged) == 1 and len(merged[0].evidence) == 2
    assert len(first.evidence) == len(second.evidence) == 1


@pytest.mark.parametrize("name", [
    "Research and development expenses (product iteration)",
    "Administrative expenses (operations support)",
    "R&D expenses (migration programme)",
])
def test_ratio_substrings_do_not_convert_currency_to_percentage(name):
    from adaptive_document_agent.document_model.metric_semantic_classifier import classify_metric, sanitize_metric_label
    from adaptive_document_agent.services.financial_normalizer import FinancialNormalizer
    item = observation("expense", name, 45100000, "FY2030", unit="currency")
    item.raw_value, item.raw_unit, item.unit_scale, item.currency = "45.1", "USD millions", 1000000, "USD"
    FinancialNormalizer.normalize_observation(item)
    assert item.value == 45100000 and item.unit_family == "currency"
    assert classify_metric(name, value=item.value, raw_unit=item.raw_unit, unit=item.unit).is_currency
    assert sanitize_metric_label(name) == name


def test_small_evidence_pages_keep_separate_headers_sources_and_notes():
    from adaptive_document_agent.services.evidence_page_packing import pack_evidence_pages
    deck = Presentation(); slides = []
    for period, value, page in [("FY2028", "100", 3), ("2H2028", "36", 9)]:
        slide = deck.slides.add_slide(deck.slide_layouts[6]); slides.append(slide)
        table = slide.shapes.add_table(2, 2, Inches(.5), Inches(1), Inches(8), Inches(1))
        table.name = "evidence:packable"
        table.table.cell(0, 1).text = period
        table.table.cell(1, 1).text = value
        footer = slide.shapes.add_textbox(Inches(.5), Inches(6), Inches(8), Inches(.3))
        footer.name = "evidence:footer"; footer.text = f"Source: p. {page}"
        slide.notes_slide.notes_text_frame.text = f"raw={value}; page={page}"
    packed = pack_evidence_pages(deck, slides)
    assert len(deck.slides) == len(packed) == 1
    tables = [s for s in deck.slides[0].shapes if s.has_table]
    assert [s.table.cell(0, 1).text for s in tables] == ["FY2028", "2H2028"]
    assert tables[1].top > tables[0].top + tables[0].height
    assert "raw=100" in packed[0].notes_slide.notes_text_frame.text
    assert "raw=36" in packed[0].notes_slide.notes_text_frame.text
    assert len({s.shape_id for s in packed[0].shapes}) == len(packed[0].shapes)


def test_complete_multi_metric_schedule_uses_one_matrix_without_losing_cells():
    from adaptive_document_agent.models import PresentationVisualBlock
    from adaptive_document_agent.document_model import DocumentIndex
    from adaptive_document_agent.services.slide_compositor import render_composed_slide
    from tests.test_pptx_export import _result
    result = _result()
    result.observations = [observation(f"m{m}-{year}", f"Programme {m} allocation", 100*m+year,
                                      f"FY{2020+year}", unit="count")
                           for m in range(1, 4) for year in range(1, 6)]
    ids = [o.id for o in result.observations]
    plan = PresentationSlide(id="schedule", slide_type="analysis", title="Reported allocation schedule",
        layout="data_overview", observation_ids=ids,
        visual_blocks=[PresentationVisualBlock(role="table", observation_ids=ids[:12]),
                       PresentationVisualBlock(role="table", observation_ids=ids[12:])])
    deck = Presentation(); deck.slide_width = Inches(12.6); deck.slide_height = Inches(7.1)
    slides = render_composed_slide(deck, plan, [], result, DocumentIndex(result.observations))
    assert len(slides) == 1
    table = next(s.table for s in slides[0].shapes if s.has_table)
    assert len(table.rows) == 4 and len(table.columns) == 6
    assert all(oid in slides[0].notes_slide.notes_text_frame.text for oid in ids)


def test_packing_middle_slides_then_appending_does_not_duplicate_package_parts():
    import io, zipfile
    from adaptive_document_agent.services.evidence_page_packing import pack_evidence_pages
    deck = Presentation(); slides = []
    for i in range(4):
        slide = deck.slides.add_slide(deck.slide_layouts[6]); slides.append(slide)
        shape = slide.shapes.add_table(2, 2, Inches(.5), Inches(1), Inches(8), Inches(1))
        shape.name = "evidence:packable"
        shape.table.cell(0, 0).text = f"Raw record {i}"
    assert len(pack_evidence_pages(deck, slides)) == 1
    deck.slides.add_slide(deck.slide_layouts[6])
    stream = io.BytesIO(); deck.save(stream)
    with zipfile.ZipFile(stream) as archive:
        assert len(archive.namelist()) == len(set(archive.namelist()))
    reopened = Presentation(io.BytesIO(stream.getvalue()))
    assert len(reopened.slides) == 2
    labels = [s.table.cell(0, 0).text for slide in reopened.slides for s in slide.shapes if s.has_table]
    assert labels == [f"Raw record {i}" for i in range(4)]
