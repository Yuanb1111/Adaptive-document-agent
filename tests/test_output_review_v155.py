"""Regressions for measured panels, inherited effects and complete source navigation."""
import io

import pytest
from pptx import Presentation

from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import ChartPlan, PresentationPlan, PresentationSlide
from adaptive_document_agent.models.summary_display import SummaryDisplay
from adaptive_document_agent.services.presentation_brief import BriefItem, render_profile
from adaptive_document_agent.services.presentation_data_index import index_entries
from adaptive_document_agent.services.pptx_export import _base_slide, _rule
from tests.test_presentation_brief import blank_deck, result_fixture


def test_independent_panels_end_at_their_own_copy_and_have_no_theme_reflection():
    deck = blank_deck()
    short = 'The source describes the stated scope.'
    long = 'The source describes the stated scope and its limitations. ' * 3
    slide, = render_profile(deck, 'Source scope', [
        BriefItem('First scope', short, [1], display=SummaryDisplay()),
        BriefItem('Second scope', long, [2], display=SummaryDisplay()),
    ])
    panels = [s for s in slide.shapes if s.name == 'summary:panel']
    accents = [s for s in slide.shapes if s.name == 'summary:accent']
    assert len(panels) == len(accents) == 2
    assert panels[0].height < panels[1].height
    for panel, accent in zip(panels, accents):
        content = [s for s in slide.shapes if s.name in {'brief:body', 'brief:heading'}
                   and panel.left < s.left < panel.left + panel.width]
        copy_bottom = max(s.top.inches + s.height.inches for s in content)
        assert panel.top.inches + panel.height.inches == pytest.approx(copy_bottom + .20, abs=1e-5)
        assert (accent.top, accent.height) == (panel.top, panel.height)
        assert panel.top.inches + panel.height.inches < deck.slide_height.inches - 1.28
    stream = io.BytesIO()
    deck.save(stream)
    saved = Presentation(io.BytesIO(stream.getvalue()))
    for shape in saved.slides[0].shapes:
        if shape.name in {'summary:panel', 'summary:accent'}:
            assert shape.element.xpath('./p:style/a:effectRef')[0].get('idx') == '0'
            effects, = shape.element.xpath('./p:spPr/a:effectLst')
            assert len(effects) == 0


def test_unused_subtitle_placeholder_is_removed_but_real_subtitle_survives():
    deck = blank_deck()
    empty = _base_slide(deck, 'Source scope')
    assert not any(s.placeholder_format.idx == 16 for s in empty.placeholders)
    populated = _base_slide(deck, 'Source scope', 'Reported values')
    assert next(s for s in populated.placeholders if s.placeholder_format.idx == 16).text == 'Reported values'
    _rule(empty, 1, 2, .035, 1, '7823FF')
    assert empty.shapes[-1].element.xpath('./p:style/a:effectRef')[0].get('idx') == '0'


def test_fallback_index_names_all_charts_and_orders_periods_without_rewriting_source():
    result = result_fixture()
    charts = result.charts[:2]
    result.presentation_plan = PresentationPlan(title='Review', slides=[PresentationSlide(
        id='pair', slide_type='analysis', title='Both measures', section_title=charts[0].title,
        chart_ids=[c.id for c in charts])])
    ids = [oid for c in charts for oid in reversed(c.observation_ids)]
    result.presentation_export_trace = [dict(planned_slide_id='pair', slide_numbers=[4],
        observation_ids=ids, chart_ids=[c.id for c in charts], source_pages=[1, 2])]
    before = [o.model_dump() for o in result.observations]
    entry, = index_entries(result, DocumentIndex(result.observations))
    assert all(c.title in entry['topic'] for c in charts)
    assert entry['periods'] == ['FY2021', 'FY2022', 'FY2023']
    assert entry['observation_ids'] == ids
    assert [o.model_dump() for o in result.observations] == before
    result.presentation_export_trace[0]['chart_ids'] = [charts[0].id]
    result.presentation_export_trace.append(dict(planned_slide_id='pair', slide_numbers=[5],
        observation_ids=ids, chart_ids=[charts[1].id], source_pages=[2]))
    entry, = index_entries(result, DocumentIndex(result.observations))
    assert all(c.title in entry['topic'] for c in charts)
    assert entry['slides'] == [4, 5]
    result.presentation_export_trace.pop()
    result.charts = [ChartPlan(id='composition', chart_type='stacked_percent',
        title='Category share of output', question='Reported mix', observation_ids=ids)]
    result.presentation_export_trace[0]['chart_ids'] = ['composition']
    entry, = index_entries(result, DocumentIndex(result.observations))
    assert entry['topic'] == 'Category share of output'


def test_neutral_question_does_not_replace_long_metric_heading_or_repeat_as_commentary():
    from adaptive_document_agent.services.pptx_export import build_presentation
    from tests.test_p0_composition import paired_result

    result = paired_result()
    slide_plan = result.presentation_plan.slides[3]
    question = 'How do the cited measures compare?'
    slide_plan.title = slide_plan.section_title = slide_plan.message = question
    slide_plan.bullets = []
    for chart in result.charts:
        chart.title = 'Net energy from/(used in) operating activities'
    before = slide_plan.model_dump()
    output = Presentation(io.BytesIO(build_presentation(result)))
    slide = next(s for s in output.slides if s.name.startswith('composed_'))
    visible = [s.text for s in slide.shapes if s.has_text_frame]
    assert question not in visible
    assert any('Net energy' in text for text in visible)
    assert slide_plan.model_dump() == before
