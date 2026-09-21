"""Generic composition contracts across shapes of evidence, not named issuers."""

import io

import pytest

pytestmark = pytest.mark.usefixtures("local_render_stub")
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from adaptive_document_agent.agent.chart_planner import ChartPlanner
from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import (
    ChartPlan, DocumentProfile, Observation, ParsedDocument, PipelineResult,
    PresentationPlan, PresentationSlide, PresentationVisualBlock, SourceEvidence,
)
from adaptive_document_agent.services.composition_data import composition_data
from adaptive_document_agent.services.export import export_pptx
from adaptive_document_agent.services.pptx_export import _add_native_chart, _planned_chart_requests
from adaptive_document_agent.services.presentation_style import deck_color_map
from adaptive_document_agent.services.qa_reporter import CriticalQAError, run_comprehensive_qa
from adaptive_document_agent.services.slide_compositor import compose_geometry, render_composed_slide, validate_composed_geometry
from adaptive_document_agent.ui.charts import render_chart


def observation(oid, metric, value, period="FY2024", *, unit="currency", category=None):
    return Observation(id=oid, metric_original=metric, value=value, raw_value=str(value), period=period,
        unit=unit, unit_family="percentage" if unit == "percent" else unit,
        currency="USD" if unit == "currency" else None, confidence=0.95,
        category_dimensions={"segment": category} if category else {},
        dimensions={"segment": category} if category else {},
        evidence=[SourceEvidence(page=3, text=str(value), row_label=metric, extraction_method="test", confidence=0.95)])


def matrix(kind="stacked_percent", metric="Revenue share", categories=("Online", "Stores")):
    periods = ("FY2023", "FY2024") if kind != "doughnut" else ("FY2024",)
    obs = [observation(f"m-{p}-{i}", metric, value, p, unit="percent", category=category)
           for p in periods for i, (category, value) in enumerate(zip(categories, (40, 60)))]
    chart = ChartPlan(id="mix", title=metric, question="Reported composition", chart_type=kind,
        available_chart_types=[kind, "table"], series_dimension="segment",
        observation_ids=[o.id for o in obs], source_pages=[3])
    return obs, chart


def result_for(obs, charts, slide=None):
    if slide is None:
        slide = PresentationSlide(id="analysis", slide_type="analysis", title="Reported performance",
            message="Comparable reported observations", section_title="Performance", layout="two_up",
            chart_ids=[c.id for c in charts], source_pages=[3])
    return PipelineResult(document=ParsedDocument(document_id="generic", sha256="a"*64, safe_filename="generic.pdf", page_count=5),
        profile=DocumentProfile(document_purpose="Review disclosed evidence"), observations=obs, charts=charts,
        presentation_plan=PresentationPlan(title="Review", slides=[
            PresentationSlide(id="cover", slide_type="cover", title="Review"),
            PresentationSlide(id="company", slide_type="company_overview", title="Document overview"),
            PresentationSlide(id="summary", slide_type="executive_summary", title="Summary"),
            slide,
            PresentationSlide(id="quality", slide_type="data_quality", title="Data quality"),
            PresentationSlide(id="appendix", slide_type="appendix", title="Source data"),
        ]))


@pytest.mark.parametrize("metric,categories", [
    ("Revenue share", ("Subscriptions", "Services")),
    ("Production share", ("Plant East", "Plant West")),
    ("Customer share", ("Enterprise", "Consumer")),
])
def test_composition_discovery_depends_on_evidence_not_company(metric, categories):
    obs, _ = matrix(metric=metric, categories=categories)
    charts = ChartPlanner().plan([], [], DocumentIndex(obs))
    assert any(c.chart_type == "stacked_percent" and set(c.observation_ids) == {o.id for o in obs} for c in charts)


@pytest.mark.parametrize("mutation", ["missing", "duplicate", "currency", "period", "negative", "zero", "partial", "evidence", "margin", "entity", "total_category", "basis"])
def test_invalid_composition_never_normalises_missing_or_incompatible_data(mutation):
    obs, chart = matrix()
    if mutation == "missing":
        obs.pop()
        chart.observation_ids.pop()
    elif mutation == "duplicate":
        obs.append(obs[-1].model_copy(update={"id": "duplicate"}))
        chart.observation_ids.append("duplicate")
    elif mutation == "currency":
        obs[-1].currency = "EUR"
    elif mutation == "period":
        obs[-1].period = "6M2024"
    elif mutation == "negative":
        obs[-1].value = -20
    elif mutation == "zero":
        for o in obs:
            o.value = 0
    elif mutation == "partial":
        obs[-1].value = 20
    elif mutation == "evidence":
        obs[-1].evidence = []
    elif mutation == "margin":
        for o in obs:
            o.metric_original = "Gross profit margin"
    elif mutation == "entity":
        obs[-1].entity = "Different entity"
    elif mutation == "basis":
        obs[-1].ifrs_status = "ADJUSTED"
    else:
        obs[-1].category_dimensions["segment"] = "Total"
    with pytest.raises(ValueError):
        composition_data(chart, obs)


def test_amount_percent_stack_requires_reconciled_source_totals():
    obs, chart = matrix(metric="Revenue")
    for o in obs:
        o.unit = o.unit_family = "currency"
        o.currency = "USD"
    with pytest.raises(ValueError, match="total observations"):
        composition_data(chart, obs)
    totals = [observation(f"total-{p}", "Total revenue", 100, p) for p in ("FY2023", "FY2024")]
    chart.total_observation_ids = [o.id for o in totals]
    original = [o.model_dump() for o in obs]
    assert composition_data(chart, obs, totals).values == [[40, 40], [60, 60]]
    assert [o.model_dump() for o in obs] == original
    totals[-1].value = 120
    with pytest.raises(ValueError, match="reconcile"):
        composition_data(chart, obs, totals)


@pytest.mark.parametrize("kind,expected", [
    ("stacked_percent", XL_CHART_TYPE.COLUMN_STACKED_100),
    ("stacked_bar", XL_CHART_TYPE.COLUMN_STACKED),
    ("doughnut", XL_CHART_TYPE.DOUGHNUT),
])
def test_native_and_ui_charts_share_validated_data(kind, expected):
    obs, chart = matrix(kind)
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    _add_native_chart(slide, chart, obs, (0.5, 0.5, 7, 4))
    native = next(s.chart for s in slide.shapes if s.has_chart)
    assert native.chart_type == expected
    ui = render_chart(chart, DocumentIndex(obs))
    if kind == "stacked_percent":
        assert list(native.series[0].values) == [0.4, 0.4]
        assert list(ui.data[0].y) == [40, 40]
        assert native.value_axis.maximum_scale == 1
    elif kind == "doughnut":
        assert list(native.series[0].values) == list(ui.data[0].values) == [40, 60]
        assert ui.data[0].hole == pytest.approx(0.62)
    else:
        assert list(native.series[0].values) == list(ui.data[0].y) == [40, 40]


def test_invalid_composition_is_a_critical_export_blocker():
    obs, chart = matrix()
    obs[-1].value = 20
    result = result_for(obs, [chart])
    assert any(i.code == "invalid_composition_chart" for i in run_comprehensive_qa(result).critical_errors)
    with pytest.raises(CriticalQAError):
        export_pptx(result)


@pytest.mark.parametrize("count,layout", [(1, "chart_plus_commentary"), (2, "two_up"), (2, "hero_plus_supporting"), (3, "three_up"), (3, "hero_plus_supporting")])
def test_geometry_slots_fit_without_overlap(count, layout):
    geometry = compose_geometry(13.33, 7.5, 1.45, count, layout=layout, has_support=True, has_commentary=True)
    rects = [*geometry.charts, geometry.support, geometry.commentary, geometry.footer]
    for a in rects:
        assert a.x >= 0 and a.y >= 0 and a.x + a.w <= 13.33 and a.y + a.h <= 7.5
    for i, a in enumerate(rects):
        for b in rects[i + 1:]:
            assert min(a.x+a.w, b.x+b.w) <= max(a.x,b.x)+1e-6 or min(a.y+a.h,b.y+b.h) <= max(a.y,b.y)+1e-6
    if layout == "hero_plus_supporting":
        assert geometry.charts[0].w > geometry.charts[1].w


def paired_result(layout="hero_plus_supporting", long_text=False):
    obs = [observation(f"a{i}", "Units shipped", v, f"FY{2022+i}", unit="count") for i, v in enumerate((100, 120, 140))]
    obs += [observation(f"b{i}", "Service hours", v, f"FY{2022+i}", unit="count") for i, v in enumerate((60, 70, 80))]
    obs += [observation("kpi", "Customer share", 40, unit="percent")]
    charts = [ChartPlan(id=c, title=t, chart_type="bar", available_chart_types=["bar", "line"], question="Reported periods",
        observation_ids=[f"{c}{i}" for i in range(3)], source_pages=[3]) for c, t in (("a", "Units shipped"), ("b", "Service hours"))]
    slide = PresentationSlide(id="analysis", slide_type="analysis", title="Operating activity expanded",
        message="Shipments and service activity increased across comparable periods.", section_title="Operations", layout=layout,
        chart_ids=["b", "a"], source_pages=[3], bullets=[("Supported operational commentary. " * 90 + "FINAL DETAIL PRESERVED") if long_text else "Both measures provide complementary evidence of operating activity."],
        visual_blocks=[PresentationVisualBlock(role="hero", chart_ids=["a"]),
            PresentationVisualBlock(role="supporting", chart_ids=["b"]),
            PresentationVisualBlock(role="kpi", observation_ids=["kpi"])])
    return result_for(obs, charts, slide)


def test_roles_drive_hero_kpi_commentary_and_source_placement():
    result = paired_result()
    deck = Presentation(io.BytesIO(export_pptx(result)))
    slide = next(s for s in deck.slides if s.name.startswith("composed_"))
    shapes = {s.name: s for s in slide.shapes}
    assert shapes["chart:a"].width > shapes["chart:b"].width
    assert "kpi:kpi" in shapes
    text = "\n".join(s.text for s in slide.shapes if s.has_text_frame)
    assert "Both measures provide complementary" in text and "Source:" in text and "40%" in text
    assert sum(s.has_chart for s in slide.shapes) == 2
    validate_composed_geometry(deck)


def test_long_commentary_continues_without_dropping_evidence_or_repeating_charts():
    result = paired_result(long_text=True)
    deck = Presentation(io.BytesIO(export_pptx(result)))
    assert any(s.name == "composed_continuation" for s in deck.slides)
    text = " ".join(s.text for sl in deck.slides for s in sl.shapes if s.has_text_frame)
    assert "FINAL DETAIL PRESERVED" in text
    assert sum(s.has_chart for sl in deck.slides for s in sl.shapes) == 2


def test_explicit_table_keeps_all_rows_and_long_labels():
    result = paired_result()
    support = [observation(f"support{i}", "Detailed operating capacity " + "qualified "*8 + chr(65+i), 10+i, unit="count") for i in range(8)]
    result.observations.extend(support)
    result.presentation_plan.slides[3].visual_blocks[-1] = PresentationVisualBlock(role="table", observation_ids=[o.id for o in support])
    deck = Presentation(io.BytesIO(export_pptx(result)))
    cells = [cell.text for sl in deck.slides if sl.name.startswith("composed_") for shape in sl.shapes if shape.has_table for row in shape.table.rows for cell in row.cells]
    assert all(o.metric_original in cells for o in support)


def test_explicit_chart_type_override_is_not_lost_by_deduplication():
    result = paired_result()
    slide = result.presentation_plan.slides[3]
    slide.visual_blocks[0].chart_type = "line"
    assert dict(_planned_chart_requests(slide))["a"] == "line"
    deck = Presentation(io.BytesIO(export_pptx(result)))
    assert next(s.chart.chart_type for sl in deck.slides for s in sl.shapes if s.name == "chart:a") == XL_CHART_TYPE.LINE_MARKERS


def test_colors_are_stable_when_input_order_changes():
    assert deck_color_map(["West", "East", "Digital"]) == deck_color_map(["Digital", "West", "East"])
    result = paired_result()
    def colors(result):
        deck = Presentation(io.BytesIO(export_pptx(result)))
        return {shape.name: str(shape.chart.series[0].format.fill.fore_color.rgb) for slide in deck.slides for shape in slide.shapes if shape.name.startswith("chart:")}
    expected = colors(result)
    result.charts.reverse()
    result.observations.reverse()
    assert colors(result) == expected


def test_actual_geometry_gate_blocks_moved_chart():
    result = paired_result()
    deck = Presentation(io.BytesIO(export_pptx(result)))
    chart = next(s for sl in deck.slides for s in sl.shapes if s.has_chart)
    chart.left = deck.slide_width
    with pytest.raises(CriticalQAError, match="outside"):
        validate_composed_geometry(deck)


def test_summary_can_render_selected_charts_and_exact_kpis():
    result = paired_result()
    summary = result.presentation_plan.slides[2]
    summary.chart_ids = ["a", "b"]
    summary.visual_blocks = [PresentationVisualBlock(role="kpi", observation_ids=["kpi"])]
    summary.source_pages = [3]
    summary.layout = "two_up"
    deck = Presentation(io.BytesIO(export_pptx(result)))
    assert sum(s.has_chart for s in deck.slides[3].shapes) == 2
    assert any(s.name == "kpi:kpi" for s in deck.slides[3].shapes)


@pytest.mark.parametrize("metric,unit", [("Revenue", "currency"), ("Energy produced", "count")])
def test_ordinary_amount_composition_is_discovered_without_percentage_conversion(metric, unit):
    obs, _ = matrix("stacked_bar", metric=metric)
    for o in obs:
        o.unit = o.unit_family = unit
        o.currency = "USD" if unit == "currency" else None
    original = [o.model_dump() for o in obs]
    charts = ChartPlanner().plan([], [], DocumentIndex(obs))
    assert any(c.chart_type == "stacked_bar" for c in charts)
    assert not any(c.chart_type == "stacked_percent" for c in charts)
    assert [o.model_dump() for o in obs] == original


@pytest.mark.parametrize("kind", ["stacked_percent", "doughnut"])
def test_constant_composition_and_category_dimensions_only_survive_export(kind):
    obs, chart = matrix(kind)
    for o in obs:
        o.value = 50
        o.raw_value = "50%"
        o.dimensions = {}
    deck = Presentation(io.BytesIO(export_pptx(result_for(obs, [chart]))))
    assert any(s.has_chart for sl in deck.slides for s in sl.shapes)
    if kind == "doughnut":
        slide = next(sl for sl in deck.slides if sl.name.startswith("composed_"))
        assert "FY2024" in " ".join(s.text for s in slide.shapes if s.has_text_frame)


def test_amount_share_export_preserves_denominators_in_appendix():
    obs, chart = matrix(metric="Revenue")
    for o in obs:
        o.unit = o.unit_family = "currency"
        o.currency = "USD"
    totals = [observation(f"total-{p}", "Total revenue", 100, p) for p in ("FY2023", "FY2024")]
    chart.total_observation_ids = [o.id for o in totals]
    result = result_for(obs + totals, [chart])
    deck = Presentation(io.BytesIO(export_pptx(result)))
    native = next(s.chart for sl in deck.slides for s in sl.shapes if s.has_chart)
    assert list(native.series[0].values) == [0.4, 0.4]
    from adaptive_document_agent.services.pptx_export import _appendix_observations
    assert set(chart.total_observation_ids) <= {o.id for o in _appendix_observations(result, [chart])}
    assert [o.value for o in obs] == [40, 60, 40, 60]


def test_full_title_and_source_safe_band_are_preserved():
    result = paired_result()
    title = "Supporting explanations remain available on continuation pages for operating activity"
    result.presentation_plan.slides[3].title = title
    deck = Presentation(io.BytesIO(export_pptx(result)))
    slide = next(sl for sl in deck.slides if sl.name.startswith("composed_"))
    assert title in [s.text for s in slide.shapes if s.has_text_frame]
    source = next(s for s in slide.shapes if s.has_text_frame and s.text.startswith("Source:"))
    assert source.top.inches + source.height.inches <= deck.slide_height.inches - 0.60


def test_block_title_cannot_relabel_cash_as_revenue():
    obs = [observation(f"cash-{i}", "Cash and cash equivalents", 100+i, f"FY{2023+i}") for i in range(2)]
    chart = ChartPlan(id="cash", title="Cash and cash equivalents", question="Cash balances", chart_type="bar",
        observation_ids=[o.id for o in obs], source_pages=[3])
    result = result_for(obs, [chart])
    result.presentation_plan.slides[3].visual_blocks = [PresentationVisualBlock(role="hero", chart_ids=[chart.id], title="Revenue")]
    with pytest.raises(CriticalQAError):
        export_pptx(result)


def test_invalid_summary_composition_is_critical_too():
    obs, chart = matrix()
    result = result_for(obs, [chart])
    result.presentation_plan.slides[2].chart_ids = [chart.id]
    obs[-1].value = 20
    assert any(i.code == "invalid_composition_chart" and i.slide_id == "summary"
               for i in run_comprehensive_qa(result).critical_errors)


def test_split_trims_cross_page_blocks_and_keeps_support_once():
    from adaptive_document_agent.validation.layout_qa import validate_presentation_layout
    result = paired_result()
    chart = result.charts[0].model_copy(update={"id": "c", "title": "Units shipped across comparable reporting periods"})
    result.charts.append(chart)
    slide = result.presentation_plan.slides[3]
    slide.chart_ids = ["a", "b", "c"]
    slide.visual_blocks = [PresentationVisualBlock(role="hero", chart_ids=["a", "c"]),
                           PresentationVisualBlock(role="supporting", chart_ids=["b"]),
                           PresentationVisualBlock(role="kpi", observation_ids=["kpi"])]
    validate_presentation_layout(result, auto_repair=True)
    parts = [s for s in result.presentation_plan.slides if s.slide_type == "analysis"]
    assert len(parts) == 2
    assert all(set(b.chart_ids) <= set(s.chart_ids) for s in parts for b in s.visual_blocks)
    assert sum("kpi" in b.observation_ids for s in parts for b in s.visual_blocks) == 1
    assert not parts[1].bullets


def test_too_many_categories_fail_instead_of_becoming_unreadable():
    obs, chart = matrix("doughnut")
    obs = [obs[0].model_copy(update={"id": f"part-{i}", "value": 100/9,
            "category_dimensions": {"segment": f"category-{i}"}}) for i in range(9)]
    chart.observation_ids = [o.id for o in obs]
    with pytest.raises(ValueError, match="eight"):
        composition_data(chart, obs)


def test_composition_review_table_retains_category_labels():
    from adaptive_document_agent.ui.charts import chart_rows
    obs, chart = matrix()
    for o in obs:
        o.dimensions = {}
    rows = chart_rows(chart, DocumentIndex(obs))
    assert {r["Series"] for r in rows} == {"Online", "Stores"}
    assert {r["Label"] for r in rows} == {"FY2023", "FY2024"}


def test_three_chart_hero_with_kpis_splits_before_labels_become_cramped():
    result = paired_result()
    result.charts.append(result.charts[0].model_copy(update={"id": "c"}))
    slide = result.presentation_plan.slides[3]
    slide.chart_ids.append("c")
    slide.visual_blocks.append(PresentationVisualBlock(role="supporting", chart_ids=["c"]))
    deck = Presentation(io.BytesIO(export_pptx(result)))
    assert sum(s.has_chart for sl in deck.slides for s in sl.shapes) == 3
    assert all(sum(s.has_chart for s in sl.shapes) <= 2 for sl in deck.slides)
    assert sum(s.name == "kpi:kpi" for sl in deck.slides for s in sl.shapes) == 1


def test_kpi_and_table_roles_remain_distinct_on_same_planned_slide():
    result = paired_result()
    result.presentation_plan.slides[3].visual_blocks.append(
        PresentationVisualBlock(role="table", observation_ids=["a0", "a1"]))
    deck = Presentation(io.BytesIO(export_pptx(result)))
    assert any(s.name == "kpi:kpi" for sl in deck.slides for s in sl.shapes)
    table_ids = {oid for sl in deck.slides for s in sl.shapes if s.name.startswith("table:")
                 for oid in s.name.removeprefix("table:").split(",")}
    assert {"a0", "a1"} <= table_ids
