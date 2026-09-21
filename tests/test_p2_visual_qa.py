import io
import json
from pathlib import Path
import subprocess
import zipfile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from adaptive_document_agent.services import presentation_visual_qa as qa
from adaptive_document_agent.services.presentation_rendering import ArtifactRenderer, RenderedPage, RenderingError, check_render_input, configured_renderer
from tests.ppt_render_stub import LocalRenderStub


def deck_bytes(*, x=1, y=1, source=False, chart=False, width=10, height=6):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(width), Inches(height)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(3), Inches(.4))
    shape.text = "Source: pages 1–2" if source else "Revenue 10,350,986 CNY"
    if source:
        footer = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(2), Inches(.3))
        footer.text = "Copyright"
    if chart:
        data = CategoryChartData()
        data.categories = ["FY2023", "FY2024"]
        data.add_series("Revenue", [100, 140])
        chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1), Inches(4), Inches(3), data)
        # python-pptx's stock chart template emits signed IDs; ECMA-376 axis IDs
        # are unsigned. Application exports already sanitize these in preflight.
        for tag in ("axId", "crossAx"):
            for node in chart_shape.chart._chartSpace.xpath(f".//c:{tag}"):
                node.set("val", str(int(node.get("val")) % (2**32)))
    stream = io.BytesIO()
    deck.save(stream)
    return stream.getvalue()


def test_pass_keeps_all_bytes_and_native_evidence():
    raw = deck_bytes(chart=True)
    renderer = LocalRenderStub()
    result = qa.verify_presentation(raw, renderer=renderer)
    assert result.payload == raw
    assert result.report.status == "passed" and result.report.facts_preserved
    assert renderer.calls == 1
    assert all(not p.exists() for p in renderer.directories)


@pytest.mark.parametrize("width,height", [(10, 6), (10, 10), (7.5, 10), (13.33, 7.5)])
def test_uses_actual_canvas_not_widescreen_constants(width, height):
    assert qa.verify_presentation(deck_bytes(width=width, height=height), renderer=LocalRenderStub()).report.status == "passed"


@pytest.mark.parametrize("x,y", [(-.1, 1), (7.1, 1), (1, -.1), (1, 5.7)])
def test_small_overflow_is_repaired_and_rerendered_without_fact_change(x, y):
    raw = deck_bytes(x=x, y=y, chart=False)
    renderer = LocalRenderStub()
    verified = qa.verify_presentation(raw, renderer=renderer)
    assert renderer.calls == 2
    assert verified.report.repairs
    assert qa.package_digest(raw, exclude_positions=True) == qa.package_digest(verified.payload, exclude_positions=True)


def test_source_overlap_moves_footer_and_preserves_workbook():
    raw = deck_bytes(source=True, y=5.3, chart=True)
    verified = qa.verify_presentation(raw, renderer=LocalRenderStub())
    assert verified.report.attempts == 2
    assert verified.report.repairs[0]["reason"] == "SOURCE_OVERLAP"
    with zipfile.ZipFile(io.BytesIO(raw)) as before, zipfile.ZipFile(io.BytesIO(verified.payload)) as after:
        for name in before.namelist():
            if not name.startswith("ppt/slides/slide"):
                assert before.read(name) == after.read(name)


def test_native_chart_translation_preserves_graphic_frame_content():
    deck = Presentation(io.BytesIO(deck_bytes(chart=True)))
    deck.slides[0].shapes[1].left = Inches(6.1)
    stream = io.BytesIO()
    deck.save(stream)
    raw = stream.getvalue()
    verified = qa.verify_presentation(raw, renderer=LocalRenderStub())
    assert verified.report.attempts == 2
    assert verified.report.facts_preserved
    chart = Presentation(io.BytesIO(verified.payload)).slides[0].shapes[1].chart
    assert list(chart.series[0].values) == [100, 140]


def test_text_fit_estimates_are_disclosed_not_silent_or_factual_rewrites():
    class Crowded(LocalRenderStub):
        def render(self, raw, directory):
            pages = super().render(raw, directory)
            text = pages[0].layout["elements"][0]
            text["textLayout"] = {"lineCount": 10}
            text["paragraphs"] = [{"runs": [{"fontSize": 20}]}]
            return pages
    raw = deck_bytes()
    result = qa.verify_presentation(raw, renderer=Crowded())
    assert result.report.status == "passed_with_warnings" and result.payload == raw
    assert result.report.issues[0].code == "TEXT_FIT_ESTIMATE"


def test_large_overflow_blocks_instead_of_rewriting_or_shrinking():
    with pytest.raises(qa.VisualQAError) as exc:
        qa.verify_presentation(deck_bytes(x=40), renderer=LocalRenderStub())
    assert exc.value.report.attempts == 1
    assert exc.value.report.issues[0].code == "OUT_OF_BOUNDS"


@pytest.mark.parametrize("budget", [-1, 3])
def test_bounded_budget(budget):
    with pytest.raises(ValueError):
        qa.verify_presentation(deck_bytes(), renderer=LocalRenderStub(), max_repairs=budget)


def test_no_repair_budget_blocks_overlap():
    with pytest.raises(qa.VisualQAError) as exc:
        qa.verify_presentation(deck_bytes(source=True, y=5.3), renderer=LocalRenderStub(), max_repairs=0)
    assert exc.value.report.attempts == 1


def test_modified_text_is_rejected_before_rerender(monkeypatch):
    def corrupt(raw, pages, issues):
        deck = Presentation(io.BytesIO(raw))
        deck.slides[0].shapes[0].text = "Fabricated 42%"
        stream = io.BytesIO()
        deck.save(stream)
        return stream.getvalue(), [{"action": "bad_repair"}]
    monkeypatch.setattr(qa, "repair_positions", corrupt)
    renderer = LocalRenderStub()
    with pytest.raises(qa.VisualQAError, match="protected"):
        qa.verify_presentation(deck_bytes(x=-.1), renderer=renderer)
    assert renderer.calls == 1


@pytest.mark.parametrize("failure,code", [("count", "PAGE_COUNT"), ("corrupt", "INVALID_RENDER"),
    ("blank", "BLANK_RENDER"), ("ratio", "INVALID_RENDER"), ("missing", "MISSING_RENDER_OBJECT"), ("nan", "INVALID_RENDER")])
def test_bad_renderer_output_never_passes(failure, code):
    class Broken(LocalRenderStub):
        def render(self, raw, directory):
            pages = super().render(raw, directory)
            if failure == "count":
                return []
            if failure == "missing":
                pages[0].layout["elements"] = []
            elif failure == "nan":
                pages[0].layout["elements"][0]["bbox"][0] = float("nan")
            elif failure in {"blank", "ratio", "corrupt"}:
                stream = io.BytesIO()
                Image.new("RGB", (800, 480) if failure == "blank" else (800, 800), "white").save(stream, format="PNG")
                pages[0] = RenderedPage(b"broken" if failure == "corrupt" else stream.getvalue(), pages[0].layout)
            return pages
    with pytest.raises(qa.VisualQAError) as exc:
        qa.verify_presentation(deck_bytes(), renderer=Broken())
    assert code in [i.code for i in exc.value.report.issues]


def test_unavailable_renderer_blocks_no_cloud_fallback(monkeypatch):
    monkeypatch.delenv("PPTX_QA_NODE", raising=False)
    monkeypatch.delenv("PPTX_QA_ARTIFACT_MODULE", raising=False)
    with pytest.raises(qa.VisualQAError, match="no cloud fallback") as exc:
        qa.verify_presentation(deck_bytes())
    assert exc.value.report.attempts == 0
    assert not exc.value.report.facts_preserved


def test_failures_not_cached_and_temporary_files_cleaned():
    class Failed(LocalRenderStub):
        def render(self, raw, directory):
            self.directories.append(directory)
            (directory / "sensitive.pptx").write_bytes(raw)
            raise RenderingError("Timed out")
    cache, renderer = {}, Failed()
    with pytest.raises(qa.VisualQAError):
        qa.verify_presentation(deck_bytes(), renderer=renderer, cache=cache)
    assert not cache and all(not p.exists() for p in renderer.directories)


def test_session_cache_reuses_only_same_package_runtime_and_policy():
    raw, cache, renderer = deck_bytes(), {}, LocalRenderStub()
    qa.verify_presentation(raw, renderer=renderer, cache=cache)
    hit = qa.verify_presentation(raw, renderer=renderer, cache=cache)
    assert renderer.calls == 1 and hit.report.cache_hit
    hit.report.issues.append(qa.VisualIssue(0, "MUTATED", "critical"))
    assert not qa.verify_presentation(raw, renderer=renderer, cache=cache).report.issues
    qa.verify_presentation(deck_bytes(x=2), renderer=renderer, cache=cache)
    assert renderer.calls == 2 and len(cache) == 1
    renderer.identity = "updated-runtime"
    qa.verify_presentation(raw, renderer=renderer, cache=cache)
    assert renderer.calls == 3


@pytest.mark.parametrize("part,data", [("ppt/vbaProject.bin", b"macro"), ("ppt/embeddings/oleObject1.bin", b"ole"),
    ("ppt/slides/_rels/slide1.xml.rels", b'<Relationships><Relationship TargetMode="External" Type="image" Target="https://example.com/private.png"/></Relationships>')])
def test_active_and_external_content_rejected_before_render(part, data):
    raw = io.BytesIO()
    with zipfile.ZipFile(raw, "w") as package:
        package.writestr(part, data)
    with pytest.raises(RenderingError):
        check_render_input(raw.getvalue())


def test_adapter_timeout_uses_argument_list_and_no_shell(tmp_path, monkeypatch):
    node, module = tmp_path / "node.exe", tmp_path / "module.mjs"
    node.touch()
    module.touch()
    renderer = ArtifactRenderer(str(node), str(module), timeout=.1)
    workers = []
    class TimeoutWorker:
        def __init__(self, args, **kwargs):
            assert isinstance(args, list) and kwargs["shell"] is False
            self.killed = False
            workers.append(self)
        def poll(self):
            return -9 if self.killed else None
        def kill(self):
            self.killed = True
        def wait(self, timeout):
            return -9
    monkeypatch.setattr(subprocess, "Popen", TimeoutWorker)
    with pytest.raises(RenderingError, match="timed out"):
        renderer.render(deck_bytes(), tmp_path)
    assert workers[0].killed


@pytest.mark.parametrize("exit_code,manifest", [(1, None), (0, None), (0, {"complete": False, "pages": 1}), (0, {"complete": True, "pages": 1})])
def test_adapter_rejects_failed_or_partial_output(tmp_path, monkeypatch, exit_code, manifest):
    node, module = tmp_path / "node.exe", tmp_path / "module.mjs"
    node.touch()
    module.touch()
    renderer = ArtifactRenderer(str(node), str(module))
    class FinishedWorker:
        def __init__(self, *a, **k):
            if manifest:
                (tmp_path / "complete.json").write_text(json.dumps(manifest))
        def poll(self):
            return exit_code
        def wait(self, timeout):
            return exit_code
    monkeypatch.setattr(subprocess, "Popen", FinishedWorker)
    with pytest.raises(RenderingError):
        renderer.render(deck_bytes(), tmp_path)


def test_importer_renumbered_chart_ids_match_unique_native_names():
    class Renumbering(LocalRenderStub):
        def render(self, raw, directory):
            pages = super().render(raw, directory)
            for e in pages[0].layout["elements"]:
                if e["kind"] == "chart":
                    e["id"] = "9999"
            return pages
    assert qa.verify_presentation(deck_bytes(chart=True), renderer=Renumbering()).report.facts_preserved


def test_blank_chart_region_blocks_even_with_nonblank_header():
    class BlankChart(LocalRenderStub):
        def render(self, raw, directory):
            page = super().render(raw, directory)[0]
            image = Image.new("RGB", (960, 576), "white")
            image.paste("black", (10, 10, 50, 50))
            stream = io.BytesIO()
            image.save(stream, format="PNG")
            return [RenderedPage(stream.getvalue(), page.layout)]
    with pytest.raises(qa.VisualQAError, match="BLANK_DATA_VISUAL"):
        qa.verify_presentation(deck_bytes(chart=True), renderer=BlankChart())


def test_two_repair_attempts_are_a_hard_limit(monkeypatch):
    def unhelpful(raw, pages, issues):
        # Valid position-only change, but does not fix the out-of-bounds content.
        from lxml import etree
        source, target = io.BytesIO(raw), io.BytesIO()
        with zipfile.ZipFile(source) as old, zipfile.ZipFile(target, "w") as new:
            for item in old.infolist():
                content = old.read(item)
                if item.filename == "ppt/slides/slide1.xml":
                    root = etree.fromstring(content)
                    off = root.xpath(".//a:xfrm/a:off", namespaces=qa.NS)[0]
                    off.set("x", str(int(off.get("x"))+1))
                    content = etree.tostring(root)
                new.writestr(item, content)
        return target.getvalue(), [{"action": "translate"}]
    monkeypatch.setattr(qa, "repair_positions", unhelpful)
    renderer = LocalRenderStub()
    with pytest.raises(qa.VisualQAError):
        qa.verify_presentation(deck_bytes(x=40), renderer=renderer)
    assert renderer.calls == 3


def test_external_links_inside_chart_workbook_are_rejected():
    workbook, ppt = io.BytesIO(), io.BytesIO()
    with zipfile.ZipFile(workbook, "w") as zipfile_out:
        zipfile_out.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")
    with zipfile.ZipFile(ppt, "w") as zipfile_out:
        zipfile_out.writestr("ppt/embeddings/chart.xlsx", workbook.getvalue())
    with pytest.raises(RenderingError):
        check_render_input(ppt.getvalue())


def test_failure_message_does_not_include_document_or_local_path(monkeypatch):
    class Broken(LocalRenderStub):
        def render(self, payload, directory):
            raise OSError("Private document phrase / private path")
    with pytest.raises(qa.VisualQAError) as exc:
        qa.verify_presentation(deck_bytes(), renderer=Broken())
    assert "Private" not in str(exc.value)
    assert "VALIDATION_FAILED" in str(exc.value)


def test_financial_critical_blocks_before_render_and_force_does_not_skip_visual(monkeypatch):
    from adaptive_document_agent.services.export import export_pptx
    from tests.test_pptx_export import _result
    monkeypatch.delenv("PPTX_QA_NODE", raising=False)
    monkeypatch.delenv("PPTX_QA_ARTIFACT_MODULE", raising=False)
    with pytest.raises(qa.VisualQAError):
        export_pptx(_result(), force=True)


@pytest.mark.skipif(not __import__("os").environ.get("PPTX_QA_INTEGRATION"), reason="Opt-in real local rendering integration")
def test_real_configured_renderer_and_repair():
    renderer = configured_renderer()
    raw = deck_bytes(source=True, y=5.3, chart=True)
    verified = qa.verify_presentation(raw, renderer=renderer)
    assert verified.report.attempts == 2 and verified.report.facts_preserved
