"""Actual native cost feedback cannot change source data or evade omission review."""
from copy import deepcopy
import pytest
from pptx import Presentation

from adaptive_document_agent.agent import presentation_budget_refinement as budget
from adaptive_document_agent.models import PresentationPlan, PresentationSlide, PresentationTopicSelection
from adaptive_document_agent.models.customization import PresentationBudget
from adaptive_document_agent.services.presentation_requirements import apply_presentation_budget
from tests.test_topic_coverage_review import _fixture, _topic


def source():
    result, _, lookup, annual, interim, adverse = _fixture()
    original = PresentationTopicSelection(topics=[_topic(annual),_topic(interim,'recent'),_topic(adverse,'constraint')])
    result.presentation_topics = original
    result.charts = []
    from adaptive_document_agent.agent.topic_plan_compiler import compile_topic_plan
    result.presentation_plan = compile_topic_plan(result)
    result.profile.report_requirements = apply_presentation_budget(None,'',PresentationBudget(body_slides=20,introduction_slides=5))
    selected = PresentationTopicSelection(topics=[_topic(interim,'recent'),_topic(adverse,'constraint')])
    draft = budget.PhysicalBudgetSelection(selection=selected,decisions=[
        dict(series_id=annual,include=False,reason='Older history remains in source records; the recent reversal and adverse outcome are retained.'),
        dict(series_id=interim,include=True,reason='Latest comparable decline qualifies historical growth.'),
        dict(series_id=adverse,include=True,reason='Adverse operating outcome is material counterevidence.')])
    return result, lookup, original, draft


@pytest.mark.parametrize('mutation',['unknown_topic','unknown_series','missing_decision','duplicate_decision','false_include','blank_reason'])
def test_source_subset_and_complete_series_decisions_are_mandatory(mutation):
    result,lookup,original,draft = source()
    if mutation=='unknown_topic':draft.selection.topics[0].id='new-document-specific-topic'
    elif mutation=='unknown_series':draft.selection.topics[0].series_ids=['unseen-source']
    elif mutation=='missing_decision':draft.decisions.pop()
    elif mutation=='duplicate_decision':draft.decisions[-1]=draft.decisions[0]
    elif mutation=='false_include':draft.decisions[0].include=True
    else:draft.decisions[0].reason=' '
    with pytest.raises(ValueError,match='every original series'):
        budget.validate_subset(draft,original,lookup,result)


def measurement(result, total, analysis):
    deck=Presentation()
    for n in range(total):
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        if n>=total-2:slide.element.cSld.set('name','ada:conclusion')
    identifier = next(s.id for s in result.presentation_plan.slides if s.slide_type=='analysis')
    result.presentation_export_trace=[dict(planned_slide_id=identifier,slide_numbers=list(range(5,5+analysis)))]
    return deck,dict(total=total,introduction=3,analysis=analysis,topic_pages=[])


@pytest.mark.parametrize('accepted',[True,False])
def test_measured_cost_refinement_requires_review_and_refreshes_only_final_brief(monkeypatch,accepted):
    result,lookup,original,draft=source()
    raw=[o.model_dump() for o in result.observations]
    old_plan=result.presentation_plan.model_copy(deep=True)
    counts=[]
    def measured(current):
        n=23 if len(current.presentation_topics.topics)==3 else 18
        counts.append(n)
        return measurement(current,n,12 if n==23 else 8)
    monkeypatch.setattr(budget,'measure',measured)
    from adaptive_document_agent.agent import topic_plan_compiler,executive_brief
    brief_calls=[]
    monkeypatch.setattr(executive_brief.ExecutiveBriefWriter,'generate',lambda self,current: brief_calls.append(current.presentation_topics.model_dump()) or None)
    class Client:
        schemas=[]
        def generate_structured(self,messages,schema,**kwargs):
            self.schemas.append(schema.__name__)
            if schema is budget.PhysicalBudgetSelection:return draft
            return schema(accepted=accepted,reason='The latest reversal and material adverse measure remain.' if accepted else 'A required source condition is missing.')
    client=Client()
    if accepted:
        assert budget.refine_presentation_budget(client,result)
        assert counts==[23,18,18] and len(brief_calls)==1
        assert result.presentation_topics==draft.selection
        assert result.validation_warnings[-1].code=='presentation_physical_budget_review'
    else:
        with pytest.raises(ValueError,match='independently rejected'):
            budget.refine_presentation_budget(client,result)
        assert result.presentation_topics==original and not brief_calls
        assert result.pipeline_state=='partial' and result.retry_from=='reporting'
        assert client.schemas==['PhysicalBudgetSelection','PhysicalBudgetReview']*2
    assert [o.model_dump() for o in result.observations]==raw


def test_raw_source_mutation_cannot_be_applied_even_after_model_approval(monkeypatch):
    result,_,original,draft=source()
    raw=deepcopy(result.observations)
    monkeypatch.setattr(budget,'measure',lambda current:measurement(current,23,12))
    from adaptive_document_agent.agent import topic_plan_compiler
    compile_plan=topic_plan_compiler.compile_topic_plan
    def corrupted(current):
        plan=compile_plan(current)
        current.observations[0].raw_value='999'
        return plan
    monkeypatch.setattr(topic_plan_compiler,'compile_topic_plan',corrupted)
    class Client:
        def generate_structured(self,messages,schema,**kwargs):
            return draft if schema is budget.PhysicalBudgetSelection else schema(accepted=True,reason='Reviewed selection.')
    with pytest.raises(ValueError,match='changed retained extraction'):
        budget.refine_presentation_budget(Client(),result)
    assert result.observations==raw and result.presentation_topics==original


def test_fitting_or_body_only_plans_need_no_paid_budget_selection(monkeypatch):
    result,_,_,_=source()
    monkeypatch.setattr(budget,'measure',lambda current:measurement(current,18,8))
    class Client:
        def generate_structured(self,*args,**kwargs):raise AssertionError('No paid budget request justified')
    assert not budget.refine_presentation_budget(Client(),result)


def test_source_wording_correction_keeps_scope_and_does_not_consume_omission_review(monkeypatch):
    result,_,_,draft=source()
    bad=draft.model_copy(deep=True)
    bad.selection.topics[0].rationale='The supported measure is 999 units.'
    monkeypatch.setattr(budget,'measure',lambda current:measurement(current,
        23 if len(current.presentation_topics.topics)==3 else 18,
        12 if len(current.presentation_topics.topics)==3 else 8))
    from adaptive_document_agent.agent import executive_brief
    monkeypatch.setattr(executive_brief.ExecutiveBriefWriter,'generate',lambda *args:None)
    class Client:
        calls=[]
        def generate_structured(self,messages,schema,**kwargs):
            self.calls.append(schema.__name__)
            if schema is budget.PhysicalBudgetSelection:return bad
            if schema is PresentationTopicSelection:
                return PresentationTopicSelection(topics=[draft.selection.topics[0]])
            return schema(accepted=True,reason='The corrected copy preserves recent reversal and adverse context.')
    client=Client()
    assert budget.refine_presentation_budget(client,result)
    assert client.calls==['PhysicalBudgetSelection','PresentationTopicSelection','PhysicalBudgetReview']
    assert result.presentation_topics==draft.selection


def test_brief_source_failure_cannot_request_a_different_topic_selection(monkeypatch):
    result,_,original,draft=source()
    monkeypatch.setattr(budget,'measure',lambda current:measurement(current,
        23 if len(current.presentation_topics.topics)==3 else 18,
        12 if len(current.presentation_topics.topics)==3 else 8))
    from adaptive_document_agent.agent import executive_brief
    def failed(*args):raise ValueError('Source accounting sign remains unsupported')
    monkeypatch.setattr(executive_brief.ExecutiveBriefWriter,'generate',failed)
    class Client:
        calls=[]
        def generate_structured(self,messages,schema,**kwargs):
            self.calls.append(schema.__name__)
            return draft if schema is budget.PhysicalBudgetSelection else schema(accepted=True,reason='Reviewed retained counterevidence.')
    client=Client()
    with pytest.raises(budget.FinalBudgetBriefError,match='accounting sign'):
        budget.refine_presentation_budget(client,result)
    assert client.calls==['PhysicalBudgetSelection','PhysicalBudgetReview']
    assert result.presentation_topics==original and result.pipeline_state=='partial'
    result.profile.report_requirements=apply_presentation_budget(None,'',PresentationBudget(body_slides=20,introduction_slides=5,slide_scope='body'))
    assert not budget.refine_presentation_budget(Client(),result)
