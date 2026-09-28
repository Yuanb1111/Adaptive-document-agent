"""Evidence and layout regressions, using synthetic inputs only."""

from copy import deepcopy

from pptx import Presentation
from pptx.util import Inches

from adaptive_document_agent.models import AnalysisResult, Insight, PresentationSlide
from adaptive_document_agent.services.appendix_layout import paginate_themes
from adaptive_document_agent.services.closing_evidence import closing_evidence
from adaptive_document_agent.services.presentation_closing import render_closing
from adaptive_document_agent.services.presentation_conventions import signed_expense_note
from adaptive_document_agent.services.pptx_export import _add_evidence_table_slides
from tests.test_company_summary_pages import sample
from tests.test_presentation_review_regressions import observations


def deck():
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    return presentation


def test_appendix_balances_rows_without_reordering_or_dropping_them():
    themes = [(f"Topic {n}", {f"{n}-{i}": {"value": i} for i in range(count)})
              for n, count in enumerate((4, 3, 3, 3))]
    pages = paginate_themes(themes)
    sizes = [sum(1 + len(entries) for _, entries in page) for page in pages]
    assert len(pages) == 2 and max(sizes) <= 10
    assert max(sizes) - min(sizes) <= 2
    assert [(theme, key, value) for page in pages for theme, entries in page for key, value in entries.items()] == [
        (theme, key, value) for theme, entries in themes for key, value in entries.items()]


def test_appendix_keeps_raw_unit_variants_even_if_normalized_values_match():
    result = sample()
    result.presentation_plan = None
    result.observations = []
    for context, scale in (("Source A", 1), ("Source B", 1000)):
        for item in observations("Operating loss", [-41000, -52000], "currency"):
            item.id = context + item.id
            item.raw_value = str(item.value / scale)
            item.raw_unit = "USD" if scale == 1 else "USD thousand"
            item.unit_scale, item.currency = scale, "USD"
            item.dimensions["table_context"] = context
            result.observations.append(item)
    original = deepcopy(result.observations)
    presentation = deck()
    _add_evidence_table_slides(presentation, result, [])
    labels = [row.cells[0].text for slide in presentation.slides for shape in slide.shapes
              if shape.has_table for row in shape.table.rows]
    assert sum("Operating loss" in label for label in labels) == 2
    assert result.observations == original


def test_appendix_separates_annual_and_interim_periods():
    result = sample()
    result.presentation_plan = None
    result.observations = observations("Revenue", [100, 120, 70], "currency")
    result.observations[-1].period = "6M2024"
    presentation = deck()
    _add_evidence_table_slides(presentation, result, [])
    assert len(presentation.slides) == 1
    tables = [shape.table for slide in presentation.slides for shape in slide.shapes if shape.has_table]
    assert len(tables) == 2
    for table in tables:
        headers = [cell.text for cell in table.rows[0].cells]
        assert not (any("FY" in h for h in headers) and any("6M" in h for h in headers))


def test_signed_expense_note_explains_magnitude_without_changing_values():
    expenses = observations("Selling expenses: %of Revenue", [-15, -20], "percent")
    original = deepcopy(expenses)
    assert "larger absolute ratio" in signed_expense_note(expenses)
    assert expenses == original
    assert not signed_expense_note(observations("Gross profit margin", [-15, -20], "percent"))
    assert not signed_expense_note(observations("Selling expenses", [-15, -20], "currency"))


def closing_result():
    result = sample()
    result.observations = observations("Inventories", [100, 150, 180], "currency")
    for o in result.observations:
        o.currency, o.raw_unit, o.unit_scale = "USD", "USD", 1
        o.evidence[0].row_label = "Inventories"
    result.analysis_results = [AnalysisResult(task_id="inventory-analysis", title="Inventory movement",
        result_type="calculated_result", result={"start_value": 100, "end_value": 180},
        input_observation_ids=[o.id for o in result.observations], confidence=.99)]
    result.insights = [Insight(id="inventory-insight", kind="interpretation", title="Inventory movement", narrative="Inventory balances increased.",
        implication="Higher inventories may require additional working capital.",
        watch_item="Monitor inventory balances in future disclosures.",
        result_ids=["inventory-analysis"], evidence=result.observations[0].evidence, confidence=.9)]
    plan = PresentationSlide(id="conclusions", slide_type="risks", title="Conclusions and Watch Items",
        insight_ids=["inventory-insight"], bullets=[result.insights[0].implication, result.insights[0].watch_item])
    return result, plan


def test_closing_renders_linked_values_and_preserves_all_raw_records():
    result, plan = closing_result()
    original = result.model_dump()
    tables, records = closing_evidence(result, plan)
    assert len(tables) == 1 and len(records) == 3
    presentation = deck()
    render_closing(presentation, result, plan)
    cells = [cell.text for slide in presentation.slides for shape in slide.shapes if shape.has_table
             for row in shape.table.rows for cell in row.cells]
    assert "Inventories" in cells and "0.0001" in cells and "0.00018" in cells
    assert "Reported values (US$ million)" in cells
    notes = "\n".join(slide.notes_slide.notes_text_frame.text for slide in presentation.slides)
    assert all(o.id in notes for o in records)
    assert result.model_dump() == original


def test_closing_never_retrieves_unlinked_facts_or_merges_conflicting_values():
    result, plan = closing_result()
    result.observations.append(result.observations[0].model_copy(update={"id": "unlinked", "value": 999}))
    _, records = closing_evidence(result, plan)
    assert "unlinked" not in {o.id for o in records}
    result.analysis_results[0].input_observation_ids.append("unlinked")
    tables, _ = closing_evidence(result, plan)
    values = [value for _, rows in tables for _, cells, _ in rows for value in cells]
    assert "0.0001" in values and "0.000999" in values
