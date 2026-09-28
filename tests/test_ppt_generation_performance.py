"""Concurrency and reuse must preserve privacy, evidence and validation results."""

import io
from threading import Event, get_ident

import pytest
from pptx import Presentation
from pptx.util import Inches

from adaptive_document_agent.agent import company_introduction as introduction
from adaptive_document_agent.models import PresentationPlan, PresentationSlide
from adaptive_document_agent.services.llm import LLMGateway, LLMSettings, MockLLMClient, PrivacyMode, ProviderName
from adaptive_document_agent.services.ppt_preflight import PresentationPreflight
from adaptive_document_agent.services import presentation_visual_qa as visual_qa
from tests.ppt_render_stub import LocalRenderStub
from tests.test_company_summary_pages import sample
from tests.test_p2_visual_qa import deck_bytes


def gateway(*, local=False, concurrent=True):
    client = MockLLMClient()
    client.supports_concurrent_requests = concurrent
    settings = LLMSettings(
        provider=ProviderName.OLLAMA if local else ProviderName.DEEPSEEK,
        model="test", privacy_mode=PrivacyMode.LOCAL_ONLY if local else PrivacyMode.CLOUD,
    )
    return LLMGateway(client, settings)


@pytest.mark.parametrize("local,concurrent", [(True, True), (False, False)])
def test_local_and_stateful_clients_do_not_start_background_requests(monkeypatch, local, concurrent):
    calls = []
    monkeypatch.setattr(introduction, "ensure_company_introduction", lambda *args: calls.append(get_ident()))
    with introduction.prepare_company_introduction(gateway(local=local, concurrent=concurrent), sample()) as attach:
        assert calls == []
        attach(PresentationPlan(title="Deck"))
    assert calls == [get_ident()]


def test_cloud_intro_overlaps_planning_and_attaches_only_after_completion(monkeypatch):
    result = sample()
    expected = result.presentation_plan.company.model_copy(deep=True)
    started, release = Event(), Event()
    caller = get_ident()
    source_name = result.document.safe_filename

    def generate(client, snapshot, draft):
        assert get_ident() != caller
        started.set()
        assert release.wait(3)
        assert snapshot.document.safe_filename == source_name
        draft.company = expected

    monkeypatch.setattr(introduction, "ensure_company_introduction", generate)
    # Also represents a plan produced by the existing deterministic recovery.
    plan = PresentationPlan(title="Recovered deck", slides=[PresentationSlide(id="cover", slide_type="cover", title="Deck")])
    with introduction.prepare_company_introduction(gateway(), result) as attach:
        try:
            assert started.wait(3), "Introduction must start before slide planning finishes"
            result.document.safe_filename = "Changed by caller"
            assert plan.company.summary_business is None
        finally:
            release.set()
        attach(plan)
    assert plan.company == expected
    assert plan.company is not expected


def test_background_validation_failure_reaches_existing_failure_handler(monkeypatch):
    def fail(*args):
        raise ValueError("Unsupported introduction evidence")
    monkeypatch.setattr(introduction, "ensure_company_introduction", fail)
    with introduction.prepare_company_introduction(gateway(), sample()) as attach:
        with pytest.raises(ValueError, match="Unsupported introduction"):
            attach(PresentationPlan(title="Deck"))


def test_prepared_intro_is_revalidated_if_source_changes_before_attach(monkeypatch):
    result = sample()
    expected = result.presentation_plan.company.model_copy(deep=True)
    def generate(client, snapshot, draft):
        draft.company = expected
    monkeypatch.setattr(introduction, "ensure_company_introduction", generate)
    with introduction.prepare_company_introduction(gateway(), result) as attach:
        result.document.pages[1].text = "Different source evidence"
        with pytest.raises(ValueError, match="no longer matches"):
            attach(PresentationPlan(title="Deck"))


def test_valid_existing_intro_survives_failure_of_optional_background_draft(monkeypatch):
    result = sample()
    expected = result.presentation_plan.company.model_copy(deep=True)
    def fail(*args):
        raise ValueError("Unavailable optional draft")
    monkeypatch.setattr(introduction, "ensure_company_introduction", fail)
    with introduction.prepare_company_introduction(gateway(), result) as attach:
        attach(result.presentation_plan)
    assert result.presentation_plan.company == expected


def test_preflight_reuse_matches_fresh_wrappers_and_refreshes_after_edits():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(0), Inches(.3))
    shape.text = "debug RMBinthousands"
    raw = io.BytesIO()
    deck.save(raw)

    class FreshWrappers(PresentationPreflight):
        def _shapes(self, slide):
            return slide.shapes
        def _placeholders(self, slide):
            return slide.placeholders

    fresh_deck, reused_deck = Presentation(io.BytesIO(raw.getvalue())), Presentation(io.BytesIO(raw.getvalue()))
    fresh = FreshWrappers(fresh_deck).validate_and_sanitize()
    preflight = PresentationPreflight(reused_deck)
    reused = preflight.validate_and_sanitize()
    assert [vars(issue) for issue in reused] == [vars(issue) for issue in fresh]
    assert reused_deck.slides[0]._element.xml == fresh_deck.slides[0]._element.xml
    assert preflight._shape_snapshots == {}
    new_shape = reused_deck.slides[0].shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(.3))
    new_shape.text = "parser"
    assert any(issue.check == "banned_phrase" for issue in preflight.validate_and_sanitize())


def test_visual_cache_hit_avoids_protected_content_reparse_but_not_identity_check(monkeypatch):
    payload = deck_bytes()
    renderer, cache = LocalRenderStub(), {}
    visual_qa.verify_presentation(payload, renderer=renderer, cache=cache)
    calls = []
    digest = visual_qa.package_digest
    def track(data, *, exclude_positions=False):
        calls.append(exclude_positions)
        return digest(data, exclude_positions=exclude_positions)
    monkeypatch.setattr(visual_qa, "package_digest", track)
    verified = visual_qa.verify_presentation(payload, renderer=renderer, cache=cache)
    assert calls == [False]
    assert verified.report.cache_hit and renderer.calls == 1
    assert verified.payload == payload
