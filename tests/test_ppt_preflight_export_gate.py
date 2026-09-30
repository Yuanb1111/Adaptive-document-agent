"""Native preflight findings survive export, caches, and later QA failures."""

import io
from unittest.mock import Mock

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from adaptive_document_agent.services import export, pptx_export
from adaptive_document_agent.services.export_diagnostics import export_diagnostics
from adaptive_document_agent.services.ppt_preflight import PreflightIssue
from adaptive_document_agent.services.presentation_preflight_report import PreflightQAError, PreflightReport
from adaptive_document_agent.services.presentation_visual_qa import VisualQAError, VisualQAReport
from adaptive_document_agent.services.qa_reporter import CriticalQAError
from tests.ppt_render_stub import LocalRenderStub
from tests.test_pptx_export import _result


def inject_finding(monkeypatch, kind):
    """Exercise real preflight checkers on native shapes in the actual template."""
    original = pptx_export._add_thank_you_slide

    def inject(deck):
        slide = deck.slides[0]
        if kind == "empty_slide":
            deck.slides.add_slide(next(layout for layout in deck.slide_layouts if not len(layout.placeholders)))
        elif kind in {"impossible_percentage_table", "visual_block_overlap"}:
            table = slide.shapes.add_table(1, 1, Inches(1), Inches(5.1), Inches(4), Inches(.5)).table
            table.cell(0, 0).text = "Revenue: 1200%" if kind.endswith("table") else "Output"
            if kind == "visual_block_overlap":
                slide.shapes.add_table(1, 1, Inches(1.1), Inches(5.1), Inches(4), Inches(.5))
        elif kind == "title_data_misalignment":
            data = CategoryChartData()
            data.categories = ["2024", "2025"]
            data.add_series("Output", [20, 10])
            chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(5), Inches(3), data).chart
            chart.has_title = True
            chart.chart_title.text_frame.text = "Output increased"
        else:
            text = {"impossible_percentage": "Revenue: 1200%", "truncated_text_fragment": "ing broken fragment",
                    "raw_unit_token_sanitized": "Revenue in USDmillions"}[kind]
            slide.shapes.add_textbox(Inches(1), Inches(5.1), Inches(5), Inches(.5)).text = text
        original(deck)

    monkeypatch.setattr(pptx_export, "_add_thank_you_slide", inject)


@pytest.mark.parametrize("kind", ["impossible_percentage", "impossible_percentage_table",
                                  "empty_slide", "visual_block_overlap", "title_data_misalignment"])
def test_actual_error_families_block_before_save_and_keep_slide_context(monkeypatch, kind):
    inject_finding(monkeypatch, kind)
    save = Mock(side_effect=AssertionError("An invalid deck must not be saved"))
    monkeypatch.setattr("pptx.presentation.Presentation.save", save)
    with pytest.raises(PreflightQAError) as caught:
        pptx_export.build_presentation(_result())
    assert isinstance(caught.value, CriticalQAError)
    report = caught.value.preflight_report
    assert report.status == "failed"
    code = kind.removesuffix("_table")
    assert any(issue.code == code for issue in report.errors)
    assert all(issue["slide"] >= 1 for issue in report.to_dict()["issues"])
    assert "Slide " in str(caught.value)
    save.assert_not_called()


@pytest.mark.parametrize("force", [False, True])
def test_public_export_cannot_force_preflight_errors_or_cache_invalid_bytes(monkeypatch, force):
    inject_finding(monkeypatch, "impossible_percentage")
    renderer, cache = LocalRenderStub(), {}
    result = _result()
    with pytest.raises(PreflightQAError) as caught:
        export.export_pptx_with_report(result, renderer=renderer, build_cache=cache, force=force)
    assert renderer.calls == 0 and not cache
    diagnostic = export_diagnostics(result, caught.value)
    assert diagnostic["export_error"]["stage"] == "preflight"
    assert diagnostic["is_export_blocked"]
    assert diagnostic["financial_qa"] == caught.value.financial_report.model_dump()
    assert any(issue["code"] == "impossible_percentage" and issue["slide_id"] == "1" for issue in diagnostic["critical_errors"])
    assert diagnostic["preflight_qa"]["issues"]


def test_repaired_warnings_survive_fresh_and_cached_exports_without_accumulation(monkeypatch):
    inject_finding(monkeypatch, "raw_unit_token_sanitized")
    builder = Mock(wraps=export.build_presentation)
    monkeypatch.setattr(export, "build_presentation", builder)
    result, cache, visual_cache = _result(), {}, {}
    renderer = LocalRenderStub()
    first = export.export_pptx_with_report(result, renderer=renderer, build_cache=cache, visual_cache=visual_cache)
    report = first.preflight_report.to_dict()
    assert report["status"] == "passed_with_warnings"
    assert not report["is_export_blocked"]
    assert any(i["code"] == "raw_unit_token_sanitized" for i in report["issues"])
    deck = Presentation(io.BytesIO(first.payload))
    assert any("USD million" in shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
    # Caller mutation must not corrupt the stored native report or visual cache.
    first.preflight_report.issues.clear()
    for _ in range(2):
        repeated = export.export_pptx_with_report(result, renderer=renderer, build_cache=cache, visual_cache=visual_cache)
        assert repeated.preflight_report.to_dict() == report
        assert repeated.build_cache_hit and repeated.report.cache_hit
        assert repeated.payload == first.payload
    assert builder.call_count == renderer.calls == 1
    fresh = export.export_pptx_with_report(result, renderer=renderer)
    assert fresh.preflight_report.to_dict() == report


def test_direct_builder_bytes_contract_and_optional_report(monkeypatch):
    inject_finding(monkeypatch, "raw_unit_token_sanitized")
    report = PreflightReport()
    payload = pptx_export.build_presentation(_result(), preflight_report=report)
    assert isinstance(payload, bytes) and payload.startswith(b"PK")
    assert report.status == "passed_with_warnings"
    assert isinstance(pptx_export.build_presentation(_result()), bytes)


def test_severities_are_preserved_in_complete_diagnostics_without_promoting_warnings():
    report = PreflightReport([
        PreflightIssue(0, "repair", "Formatting was repaired", "warning"),
        PreflightIssue(1, "detail", "Review detail", "info"),
    ])
    diagnostic = export_diagnostics(_result(), preflight_report=report)
    assert not diagnostic["is_export_blocked"]
    assert any(i["code"] == "repair" and i["severity"] == "WARNING" for i in diagnostic["warnings"])
    assert any(i["code"] == "detail" and i["severity"] == "INFO" for i in diagnostic["info"])
    assert not any(i["code"] in {"repair", "detail"} for i in diagnostic["critical_errors"])


def test_later_visual_failure_retains_preflight_findings(monkeypatch):
    inject_finding(monkeypatch, "raw_unit_token_sanitized")
    failure = VisualQAError(VisualQAReport())
    monkeypatch.setattr("adaptive_document_agent.services.presentation_visual_qa.verify_presentation", Mock(side_effect=failure))
    cache = {}
    with pytest.raises(VisualQAError) as caught:
        export.export_pptx_with_report(_result(), build_cache=cache)
    assert not cache
    diagnostic = export_diagnostics(_result(), caught.value, caught.value.report)
    assert diagnostic["export_error"]["stage"] == "rendering"
    assert any(i["code"] == "raw_unit_token_sanitized" for i in diagnostic["preflight_qa"]["issues"])


def test_bytes_only_build_cache_cannot_discard_sanitization_history(monkeypatch):
    inject_finding(monkeypatch, "raw_unit_token_sanitized")
    result, cache = _result(), {}
    export.export_pptx_with_report(result, renderer=LocalRenderStub(), build_cache=cache)
    key, value = next(iter(cache.items()))
    cache[key] = value[0]  # Previous cache value contract.
    repeated = export.export_pptx_with_report(result, renderer=LocalRenderStub(), build_cache=cache)
    assert not repeated.build_cache_hit
    assert any(i.code == "raw_unit_token_sanitized" for i in repeated.preflight_report.issues)


def test_error_report_in_cache_still_blocks_before_rendering(monkeypatch):
    result, cache = _result(), {}
    export.export_pptx_with_report(result, renderer=LocalRenderStub(), build_cache=cache)
    key, (payload, _) = next(iter(cache.items()))
    cache[key] = (payload, PreflightReport([PreflightIssue(1, "invalid", "Invalid slide", "error")]))
    renderer = LocalRenderStub()
    with pytest.raises(PreflightQAError):
        export.export_pptx_with_report(result, renderer=renderer, build_cache=cache, force=True)
    assert not renderer.calls


@pytest.mark.parametrize("failure", [ValueError("Invalid generated brand"), RuntimeError("Save failed")])
def test_later_generation_failure_retains_preflight_without_changing_exception_type(monkeypatch, failure):
    inject_finding(monkeypatch, "raw_unit_token_sanitized")
    monkeypatch.setattr("adaptive_document_agent.services.presentation_brand_qa.validate_generated_brand", Mock(side_effect=failure))
    with pytest.raises(type(failure)) as caught:
        export.export_pptx_with_report(_result(), renderer=LocalRenderStub())
    assert caught.value is failure
    diagnostic = export_diagnostics(_result(), caught.value)
    assert diagnostic["export_error"]["stage"] == "generation"
    assert any(i["code"] == "raw_unit_token_sanitized" for i in diagnostic["preflight_qa"]["issues"])


def test_failure_before_preflight_is_not_reported_as_a_pass(monkeypatch):
    assert PreflightReport().status == "not_run"
    monkeypatch.setattr(pptx_export, "_add_thank_you_slide", Mock(side_effect=CriticalQAError("Build failed")))
    with pytest.raises(CriticalQAError) as caught:
        export.export_pptx_with_report(_result(), renderer=LocalRenderStub())
    assert caught.value.preflight_report is None
    assert export_diagnostics(_result(), caught.value)["preflight_qa"] is None


@pytest.mark.parametrize("in_table", [False, True])
@pytest.mark.parametrize("percentage", ["1,100%", "999999%"])
def test_percentage_magnitude_alone_never_proves_an_invalid_claim(monkeypatch, in_table, percentage):
    original = pptx_export._add_thank_you_slide

    def inject(deck):
        slide = deck.slides[0]
        text = f"Output growth: {percentage}"
        if in_table:
            table = slide.shapes.add_table(1, 1, Inches(1), Inches(5.1), Inches(5), Inches(.5)).table
            table.cell(0, 0).text = text
        else:
            slide.shapes.add_textbox(Inches(1), Inches(5.1), Inches(5), Inches(.5)).text = text
        original(deck)

    monkeypatch.setattr(pptx_export, "_add_thank_you_slide", inject)
    verified = export.export_pptx_with_report(_result(), renderer=LocalRenderStub())
    findings = [issue for issue in verified.preflight_report.issues if issue.code == "extreme_percentage"]
    assert findings and all(issue.severity == "warning" for issue in findings)
    assert not verified.preflight_report.errors
    assert verified.payload.startswith(b"PK")


@pytest.mark.parametrize("text", ["a company serving customers worldwide", "x axis represents output", "ing broken fragment"])
def test_lexical_fragment_heuristics_warn_without_rejecting_valid_articles_or_variables(monkeypatch, text):
    original = pptx_export._add_thank_you_slide
    def inject(deck):
        deck.slides[0].shapes.add_textbox(Inches(1), Inches(5.1), Inches(5), Inches(.5)).text = text
        original(deck)
    monkeypatch.setattr(pptx_export, "_add_thank_you_slide", inject)
    verified = export.export_pptx_with_report(_result(), renderer=LocalRenderStub())
    findings = [i for i in verified.preflight_report.issues if i.code == "truncated_text_fragment"]
    assert findings and all(i.severity == "warning" for i in findings)
    deck = Presentation(io.BytesIO(verified.payload))
    assert any(shape.text == text for shape in deck.slides[0].shapes if shape.has_text_frame)


@pytest.mark.parametrize("title,categories,values", [
    ("Revenue increased versus last year", ["Large segment", "Small segment"], [120, 10]),
    ("Costs improved", ["FY2024", "FY2025"], [120, 90]),
    ("Revenue increased", ["FY2025", "FY2024"], [120, 90]),
    ("Revenue grew while costs fell", ["FY2024", "FY2025"], [90, 120]),
    ("Losses increased", ["FY2024", "FY2025"], [-90, -120]),
    ("Revenue increased", ["3 months ended 30 June 2024", "6 months ended 30 June 2025"], [120, 90]),
])
def test_direction_guesses_require_comparable_time_axis_and_unambiguous_semantics(title, categories, values):
    from adaptive_document_agent.services.ppt_preflight import PresentationPreflight
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = CategoryChartData()
    data.categories = categories
    data.add_series("Measure", values)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = title
    findings = [i for i in PresentationPreflight(deck).validate_and_sanitize() if i.code == "title_data_misalignment"]
    assert findings and all(i.severity == "warning" for i in findings)
