"""PowerPoint export remains editable and presentation-ready."""

import io
from pathlib import Path
import zipfile

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from adaptive_document_agent.models import (
    ChartPlan,
    DocumentProfile,
    Insight,
    Observation,
    ParsedDocument,
    PipelineResult,
    PresentationPlan,
    PresentationSlide,
    PresentationVisualBlock,
    CompanyProfile,
    ReportPlan,
    SourceEvidence,
    ValidationIssue,
)
from adaptive_document_agent.services.export import export_pptx
from adaptive_document_agent.services.pptx_export import _series_rows


def _result() -> PipelineResult:
    evidence = SourceEvidence(page=234, text="267,025", table_id="income", extraction_method="digital_table", confidence=0.9)
    observations = [
        Observation(
            id=f"revenue-{year}",
            metric_original="Revenue",
            value=value,
            raw_value=f"{value / 1_000:,.0f}",
            unit="currency",
            raw_unit="RMB in thousands",
            unit_scale=1_000,
            currency="CNY",
            period=f"FY{year}",
            dimensions={"period_basis": "FY", "table_context": "RESULTS OF OPERATIONS"},
            evidence=[evidence],
            confidence=0.9,
        )
        for year, value in ((2023, 267_025_000), (2024, 325_257_000), (2025, 521_747_000))
    ]
    return PipelineResult(
        document=ParsedDocument(document_id="doc", sha256="abc", safe_filename="prospectus.pdf", page_count=300),
        profile=DocumentProfile(
            document_type="Prospectus",
            document_purpose="Review historical financial performance for investor assessment.",
            overview_title="Company overview",
            document_summary="The company develops automation products for industrial customers and sells them across several markets.",
            document_summary_pages=[8, 9],
            metrics=["Revenue"],
            analysis_page_ranges=[(227, 263)],
            analysis_focus="Financial information",
            data_quality_notes=["Values use RMB in thousands in the source table."],
        ),
        observations=observations,
        insights=[
            Insight(
                id="insight-1",
                title="Revenue increased across the reported period",
                narrative="Revenue rose from RMB267.0 million in 2023 to RMB521.7 million in 2025.",
                kind="reported_fact",
                importance=0.9,
                confidence=0.9,
                evidence=[evidence],
            )
        ],
        report_plan=ReportPlan(title="Financial information analysis"),
        charts=[
            ChartPlan(
                id="chart-1",
                title="Revenue — Reported Values",
                chart_type="line",
                available_chart_types=["line", "bar", "area", "table"],
                question="How did revenue change across the reported years?",
                observation_ids=[item.id for item in observations],
                source_pages=[234],
                x_metric="revenue",
                y_metric="revenue",
                x_axis_title="Period",
                y_axis_title="Revenue (CNY)",
            )
        ],
        validation_warnings=[ValidationIssue(code="source_scale", message="Source values are reported in RMB thousands.", stage="report")],
    )


def test_pptx_export_contains_editable_chart_and_table() -> None:
    payload = export_pptx(_result())

    assert payload.startswith(b"PK")
    deck = Presentation(io.BytesIO(payload))
    assert len(deck.slides) >= 9
    assert any(shape.has_chart for slide in deck.slides for shape in slide.shapes)
    assert any(shape.has_table for slide in deck.slides for shape in slide.shapes)
    assert any("Company overview" in shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    all_text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert "Contents" in all_text
    assert "Thematic analysis" in all_text

    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        names = archive.namelist()
        assert any(name.startswith("ppt/charts/chart") for name in names)
        assert any(name.startswith("ppt/embeddings/") for name in names)
        chart_xml = b"".join(archive.read(name) for name in names if name.startswith("ppt/charts/chart"))
        assert b'axId val="-' not in chart_xml
        assert b'crossAx val="-' not in chart_xml


def test_pptx_export_renders_validated_ai_story_plan() -> None:
    result = _result()
    result.presentation_plan = PresentationPlan(
        title="AI planned review",
        company=CompanyProfile(
            name="Example Automation",
            one_line_description="A source-described automation products company.",
            document_type="Prospectus",
            source_pages=[8],
        ),
        slides=[
            PresentationSlide(id="cover", slide_type="cover", title="AI planned review", message="Evidence-led analysis"),
            PresentationSlide(id="overview", slide_type="company_overview", title="Company at a Glance", source_pages=[8]),
            PresentationSlide(
                id="summary",
                slide_type="executive_summary",
                title="Executive Summary",
                bullets=["Revenue increased across the retained periods."],
                insight_ids=["insight-1"],
                source_pages=[234],
            ),
            PresentationSlide(
                id="analysis",
                slide_type="analysis",
                title="Revenue growth accelerated in the latest period",
                section_title="Financial Performance",
                message="The retained revenue series supports the direction and magnitude shown.",
                chart_ids=["chart-1"],
                source_pages=[234],
            ),
            PresentationSlide(id="quality", slide_type="data_quality", title="Data quality and methodology"),
            PresentationSlide(id="appendix", slide_type="appendix", title="Source data"),
        ],
    )

    deck = Presentation(io.BytesIO(export_pptx(result)))
    titles = [
        next((shape.text.strip() for shape in slide.shapes if shape.has_text_frame and shape.text.strip()), "")
        for slide in deck.slides
    ]

    assert titles[:5] == [
        "AI planned review",
        "Contents",
        "Company at a Glance",
        "Executive Summary",
        "Revenue growth accelerated in the latest period",
    ]
    all_text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert "Example Automation" in all_text
    assert "Revenue growth accelerated in the latest period" in all_text
    assert any(
        shape.has_text_frame and shape.text.strip() == "Revenue growth accelerated in the latest period"
        for shape in deck.slides[4].shapes
    )
    contents_text = "\n".join(shape.text for shape in deck.slides[1].shapes if shape.has_text_frame)
    assert "Financial Performance" in contents_text
    assert "Revenue growth accelerated" not in contents_text


def test_planned_pptx_places_multiple_editable_charts_on_one_message_slide() -> None:
    result = _result()
    evidence = result.observations[0].evidence
    cash = [
        Observation(
            id=f"cash-{year}", metric_original="Cash balance", value=value, raw_value=str(value),
            unit="currency", currency="CNY", period=f"FY{year}", evidence=evidence, confidence=0.9,
        )
        for year, value in ((2023, 80_000_000.0), (2024, 95_000_000.0), (2025, 120_000_000.0))
    ]
    result.observations.extend(cash)
    result.charts.append(
        ChartPlan(
            id="cash-chart", title="Cash balance", chart_type="bar", available_chart_types=["bar", "line", "table"],
            question="How did cash change?", observation_ids=[item.id for item in cash], source_pages=[234],
        )
    )
    result.presentation_plan = PresentationPlan(
        title="AI planned review",
        company=CompanyProfile(name="Example Automation", one_line_description="Automation company.", source_pages=[8]),
        slides=[
            PresentationSlide(id="cover", slide_type="cover", title="AI planned review"),
            PresentationSlide(id="overview", slide_type="company_overview", title="Company at a Glance", source_pages=[8]),
            PresentationSlide(id="summary", slide_type="executive_summary", title="Executive Summary", insight_ids=["insight-1"], source_pages=[234]),
            PresentationSlide(
                id="analysis", slide_type="analysis", title="Growth and liquidity moved together",
                section_title="Financial Performance", message="Revenue and cash both increased across the retained periods.",
                layout="hero_plus_supporting", source_pages=[234],
                visual_blocks=[
                    PresentationVisualBlock(role="hero", chart_ids=["chart-1"]),
                    PresentationVisualBlock(role="supporting", chart_ids=["cash-chart"]),
                ],
            ),
            PresentationSlide(id="quality", slide_type="data_quality", title="Data quality"),
            PresentationSlide(id="appendix", slide_type="appendix", title="Source data"),
        ],
    )

    deck = Presentation(io.BytesIO(export_pptx(result)))
    chart_counts = [sum(1 for shape in slide.shapes if shape.has_chart) for slide in deck.slides]

    assert chart_counts.count(2) == 1
    assert sum(chart_counts) == 2


def test_pptx_contents_order_matches_generated_sections() -> None:
    deck = Presentation(io.BytesIO(export_pptx(_result())))
    contents_index = next(
        index
        for index, slide in enumerate(deck.slides)
        if any(shape.has_text_frame and shape.text.strip() == "Contents" for shape in slide.shapes)
    )
    findings_index = next(
        index
        for index, slide in enumerate(deck.slides)
        if index > contents_index
        and any(shape.has_text_frame and shape.text.strip() == "Key findings" for shape in slide.shapes)
    )
    thematic_index = next(
        index
        for index, slide in enumerate(deck.slides)
        if index > contents_index
        and any(shape.has_text_frame and shape.text.strip() == "Thematic analysis" for shape in slide.shapes)
    )

    assert findings_index < thematic_index


def test_series_rows_preserve_non_axis_business_dimensions() -> None:
    evidence = SourceEvidence(page=5, text="10", extraction_method="digital_table", confidence=0.9)
    observations = [
        Observation(
            id=f"segment-{segment}-{year}-{ordinal}",
            metric_original="Revenue",
            value=value,
            raw_value=str(value),
            period=f"FY{year}",
            dimensions={"segment": segment, "table_context": "Revenue by segment"},
            evidence=[evidence],
            confidence=confidence,
        )
        for segment, year, ordinal, value, confidence in (
            ("Automation", 2024, 1, 10.0, 0.9),
            ("Services", 2024, 1, 20.0, 0.9),
            ("Automation", 2025, 1, 15.0, 0.9),
            ("Services", 2025, 1, 25.0, 0.9),
            ("Automation", 2024, 2, 999.0, 0.2),
        )
    ]
    plan = ChartPlan(
        id="segment-chart",
        title="Revenue by segment",
        chart_type="line",
        question="How did each segment change?",
        observation_ids=[item.id for item in observations],
        source_pages=[5],
    )

    rows = _series_rows(plan, observations)

    assert len(rows) == 4
    assert {series for _, series, _ in rows} == {
        "Revenue (Segment: Automation)",
        "Revenue (Segment: Services)",
    }
    assert ("FY2024", "Revenue (Segment: Automation)", 10.0) in rows


def test_pptx_export_skips_constant_period_chart() -> None:
    result = _result()
    for observation in result.observations:
        observation.value = 100.0
        observation.raw_value = "100"

    payload = export_pptx(result)
    deck = Presentation(io.BytesIO(payload))

    assert not any(shape.has_chart for slide in deck.slides for shape in slide.shapes)


def test_pptx_export_keeps_the_planned_chart_mix_editable() -> None:
    result = _result()
    base = result.charts[0]
    result.charts = [
        base.model_copy(update={"id": f"chart-{chart_type}", "chart_type": chart_type, "title": f"Revenue {chart_type}"})
        for chart_type in ("line", "bar", "area")
    ]

    deck = Presentation(io.BytesIO(export_pptx(result)))
    chart_types = [shape.chart.chart_type for slide in deck.slides for shape in slide.shapes if shape.has_chart]

    assert chart_types == [XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.AREA]


def test_pptx_appendix_paginates_all_chart_observations_and_cleans_units() -> None:
    result = _result()
    evidence = result.observations[0].evidence
    result.observations = [
        Observation(
            id=f"revenue-{year}",
            metric_original="Revenue",
            value=float(year),
            raw_value=str(year),
            raw_unit="RMB\ufffd\ufffd000",
            currency="CNY",
            unit="currency",
            period=f"FY{year}",
            evidence=evidence,
            confidence=0.9,
        )
        for year in range(2003, 2022)
    ]
    result.charts[0] = result.charts[0].model_copy(update={"observation_ids": [item.id for item in result.observations]})

    deck = Presentation(io.BytesIO(export_pptx(result)))
    appendix_slides = [
        slide
        for slide in deck.slides
        if any(shape.has_text_frame and "Key data appendix" in shape.text for shape in slide.shapes)
    ]
    assert len(appendix_slides) == 2
    appendix_values = [
        cell.text
        for slide in appendix_slides
        for shape in slide.shapes
        if shape.has_table
        for row in shape.table.rows
        for cell in row.cells
    ]
    assert all(str(year) in appendix_values for year in range(2003, 2022))
    assert "RMB '000" in appendix_values


def test_pptx_filters_junk_metrics_and_derives_findings_from_valid_charts() -> None:
    result = _result()
    result.insights = []
    evidence = result.observations[0].evidence
    junk = [
        Observation(
            id=f"junk-{year}", metric_original="at", value=value, raw_value=str(value), unit="percent",
            period=str(year), evidence=evidence, confidence=0.9,
        )
        for year, value in ((2023, 31.7), (2024, 36.9))
    ]
    result.observations.extend(junk)
    result.charts.append(
        ChartPlan(
            id="junk-chart", title="FINANCIAL INFORMATION - at", chart_type="bar",
            question="invalid fragment", observation_ids=[item.id for item in junk], source_pages=[234],
        )
    )

    deck = Presentation(io.BytesIO(export_pptx(result)))
    all_text = "\n".join(
        shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame
    )

    assert "FINANCIAL INFORMATION - at" not in all_text
    assert "No validated analytical findings were produced" not in all_text
    assert "Revenue increased from" in all_text
    assert "\nat\n" not in f"\n{all_text}\n"


def test_pptx_groups_related_charts_into_one_story_slide() -> None:
    result = _result()
    evidence = result.observations[0].evidence
    cash = [
        Observation(
            id=f"cash-{year}", metric_original="Cash balance", value=value, raw_value=str(value),
            unit="currency", currency="CNY", period=f"FY{year}",
            dimensions={"period_basis": "FY", "table_context": "Financial position"},
            evidence=evidence, confidence=0.9,
        )
        for year, value in ((2023, 80.0), (2024, 95.0), (2025, 120.0))
    ]
    result.observations.extend(cash)
    result.charts.append(
        ChartPlan(
            id="cash-chart", title="Cash balance", chart_type="bar",
            question="How did cash change?", observation_ids=[item.id for item in cash], source_pages=[234],
        )
    )
    # Give the original chart the same discovered table context. The grouping
    # is generic and does not depend on financial metric names.
    for observation in result.observations[:3]:
        observation.dimensions["table_context"] = "Financial position"

    deck = Presentation(io.BytesIO(export_pptx(result)))
    chart_counts = [sum(1 for shape in slide.shapes if shape.has_chart) for slide in deck.slides]

    assert 2 in chart_counts
    grouped_slide = deck.slides[chart_counts.index(2)]
    grouped_text = "\n".join(shape.text for shape in grouped_slide.shapes if shape.has_text_frame)
    assert "Revenue" in grouped_text
    assert "Cash balance" in grouped_text


def test_pptx_keeps_unrelated_charts_on_separate_slides() -> None:
    result = _result()
    evidence = result.observations[0].evidence
    other = [
        Observation(
            id=f"other-{year}", metric_original="Distinct measure", value=value, raw_value=str(value),
            unit="percent", period=f"FY{year}", dimensions={"table_context": "Separate source table"},
            evidence=evidence, confidence=0.9,
        )
        for year, value in ((2023, 10.0), (2024, 20.0), (2025, 30.0))
    ]
    result.observations.extend(other)
    result.charts.append(
        ChartPlan(
            id="other-chart", title="Distinct measure", chart_type="line",
            question="How did the measure change?", observation_ids=[item.id for item in other], source_pages=[250],
        )
    )

    deck = Presentation(io.BytesIO(export_pptx(result)))
    chart_counts = [sum(1 for shape in slide.shapes if shape.has_chart) for slide in deck.slides]

    assert chart_counts.count(1) == 2
    assert 2 not in chart_counts


def test_fourier_template_dimensions_and_layouts() -> None:
    payload = export_pptx(_result())
    deck = Presentation(io.BytesIO(payload))

    # Assert 12.60in x 7.09in dimensions from FOURIER template
    assert deck.slide_width.inches == pytest.approx(12.60, abs=0.01)
    assert deck.slide_height.inches == pytest.approx(7.09, abs=0.01)

    # Check that template layout names are used
    layout_names = [slide.slide_layout.name for slide in deck.slides]
    assert any("封面" in name for name in layout_names)
    assert any("Single-line title" in name or "Two-line" in name for name in layout_names)
    assert any("短文本" in name for name in layout_names)


def test_pptx_export_fails_cleanly_when_template_missing(monkeypatch: pytest.MonkeyPatch) -> None:
    non_existent = r"C:\path\to\non_existent_template_file.pptx"
    with pytest.raises(FileNotFoundError, match="Required PowerPoint template file not found"):
        export_pptx(_result(), template_path=non_existent)

    monkeypatch.setenv("PPTX_TEMPLATE_PATH", non_existent)
    with pytest.raises(FileNotFoundError, match="Required PowerPoint template file not found"):
        export_pptx(_result())


def test_pptx_export_fails_cleanly_when_template_corrupted(tmp_path: Path) -> None:
    corrupted_file = tmp_path / "corrupted_template.pptx"
    corrupted_file.write_bytes(b"not a valid pptx content")

    with pytest.raises(ValueError, match="Failed to load PowerPoint template"):
        export_pptx(_result(), template_path=corrupted_file)


def test_bundled_template_exists_and_loads_by_default(monkeypatch: pytest.MonkeyPatch) -> None:
    from adaptive_document_agent.services.pptx_export import BUNDLED_TEMPLATE_PATH
    assert BUNDLED_TEMPLATE_PATH.exists()
    assert BUNDLED_TEMPLATE_PATH.is_file()
    monkeypatch.delenv("PPTX_TEMPLATE_PATH", raising=False)
    payload = export_pptx(_result())
    assert payload.startswith(b"PK")


