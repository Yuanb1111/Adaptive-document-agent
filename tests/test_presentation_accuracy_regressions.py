"""Generic contracts reproduced from exported-deck review, with synthetic evidence."""

import io

import pytest
from pptx import Presentation

from adaptive_document_agent.document_model import DocumentIndex, paired_observations
from adaptive_document_agent.extraction.borderless_layout import SourceLine
from adaptive_document_agent.extraction.borderless_table_extractor import BorderlessTableExtractor
from adaptive_document_agent.extraction.period_header_geometry import geometric_periods
from adaptive_document_agent.models import ChartPlan, Observation, SourceEvidence, PipelineResult, ParsedDocument, DocumentProfile
from adaptive_document_agent.services.composition_candidates import reported_composition_charts
from adaptive_document_agent.services.composition_data import composition_data, is_aggregate_category
from adaptive_document_agent.services.movement_formatter import format_movement_narrative
from adaptive_document_agent.services.pptx_export import build_presentation, _source_footer


def obs(oid, value, period="FY2024", metric="Resources", category=None):
    return Observation(id=oid, metric_original=metric, value=value, raw_value=str(value),
        period=period, unit="currency", currency="USD", confidence=.95,
        category_dimensions={"component": category} if category else {},
        evidence=[SourceEvidence(page=3, text=str(value), row_label=metric, extraction_method="test", confidence=.95)])


@pytest.mark.parametrize("label", ["Total current assets", "TOTAL REVENUE", "Subtotal resources", "Grand total", "合计资源"])
def test_qualified_totals_are_never_components(label):
    values = [obs(f"{p}-{c}", v, p, category=c) for p in ("FY2023", "FY2024")
              for c, v in (("East", 40), ("West", 60), (label, 100))]
    plan = ChartPlan(id="c", title="Resources", question="Composition", chart_type="stacked_bar",
                     series_dimension="component", observation_ids=[o.id for o in values])
    with pytest.raises(ValueError, match="Aggregate totals"):
        composition_data(plan, values)
    charts = reported_composition_charts(DocumentIndex(values))
    assert len(charts) == 1
    chosen = charts[0]
    assert len(chosen.observation_ids) == 4
    assert len(chosen.total_observation_ids) == 2
    components = [o for o in values if o.id in chosen.observation_ids]
    totals = [o for o in values if o.id in chosen.total_observation_ids]
    assert composition_data(chosen, components, totals).values == [[40, 40], [60, 60]]
    totals[0].value = 110
    with pytest.raises(ValueError, match="reconcile"):
        composition_data(chosen, components, totals)


def test_similar_spelling_and_equal_values_are_not_inferred_totals():
    assert not is_aggregate_category("TotalEnergies")
    values = [obs(f"{p}-{c}", v, p, category=c) for p in ("FY2023", "FY2024")
              for c, v in (("East", 40), ("West", 60), ("Central", 100))]
    assert len(reported_composition_charts(DocumentIndex(values))[0].observation_ids) == 6


def line(*phrases):
    text, spans = "", []
    for phrase, center in phrases:
        start = len(text)
        text += phrase + " "
        spans.append((start, start+len(phrase), center-len(phrase)*2, center+len(phrase)*2))
    return SourceLine(text.rstrip(), spans)


@pytest.mark.parametrize("years", [("2021", "2022", "2023", "2024"), ("2018", "2019", "2020", "2020")])
def test_split_superheaders_use_geometry_not_duplicate_year_heuristic(years):
    year_line = line(*zip(years, (200, 270, 340, 410)))
    headers = [line(("Six months", 410)), line(("Year ended December 31,", 270)), line(("ended June 30,", 410))]
    assert geometric_periods(headers, year_line, 4) == [*(f"FY{y}" for y in years[:3]), f"6M{years[3]}"]
    assert geometric_periods(headers, year_line, 8) == [v for v in [*(f"FY{y}" for y in years[:3]), f"6M{years[3]}"] for _ in range(2)]


def test_snapshot_dates_and_missing_geometry():
    years = line(*zip(("2021", "2022", "2023", "2024"), (200, 270, 340, 410)))
    headers = [line(("As of December 31,", 270)), line(("As of June 30,", 410))]
    assert geometric_periods(headers, years, 4) == ["2021-12-31", "2022-12-31", "2023-12-31", "2024-06-30"]
    assert geometric_periods(headers, SourceLine("2021 2022 2023 2024"), 4) is None
    raw = ["Six months", "Year ended December 31,", "ended June 30,", "2021 2022 2023 2024"]
    assert BorderlessTableExtractor._expand_periods(raw[-1].split(), 4, raw, 3) == [None] * 4
    assert BorderlessTableExtractor._context_label(["Inventory", *raw], 4) == "Inventory"


@pytest.mark.parametrize("metric", ["Inventory turnover days", "Operating cash flow", "Trade receivables turnover days"])
def test_explicit_flow_period_not_converted_to_snapshot_by_metric_keywords(metric):
    from adaptive_document_agent.document_model.period_semantic_validator import format_observation_period
    value = obs("flow", 53, "6M2024", metric)
    value.period_type = "interim_flow"
    assert format_observation_period(value) == "6M2024"


@pytest.mark.parametrize("values", [[6.4e6, -157.7e6], [6.4e6, 15e6, -157.7e6], [-20e6, -50e6, -10e6]])
def test_cash_narrative_retains_negative_source_levels(values):
    periods = [f"FY{2021+i}" for i in range(len(values))]
    text = format_movement_narrative("Operating cash flow", start_val=values[0], end_val=values[-1],
        currency="USD", unit="currency", values=values, periods=periods,
        start_period=periods[0], end_period=periods[-1])
    assert ("-US$ 157.7m" if values[-1] == -157.7e6 else "-US$ 10m") in text


def test_correlation_rejects_pooling_annual_and_interim_samples():
    values = [obs(f"{m}-{p}", i+1, p, metric=m) for m in ("Volume", "Output")
              for i, p in enumerate(("FY2021", "FY2022", "FY2023", "6M2023", "6M2024"))]
    assert paired_observations(values, "volume", "output") == []
    for o in values:
        o.period = f"FY{2010 + int(o.value)}"
    assert len(paired_observations(values, "volume", "output")) == 5


def test_fallback_preserves_long_copy_and_does_not_advertise_unplotted_series():
    values = [obs("a", 10, "FY2023", "Cost of sales"), obs("b", 20, "FY2024", "Cost of sales")]
    chart = ChartPlan(id="c", title="Cost of sales and Revenue", question="Has multiple periods",
        chart_type="bar", observation_ids=[o.id for o in values], source_pages=[3])
    summary = "The source describes retained operations and their documented limitations. " * 45 + "COMPLETE_END_MARKER"
    result = PipelineResult(document=ParsedDocument(document_id="test", sha256="a"*64, safe_filename="test.pdf", page_count=3),
        profile=DocumentProfile(document_summary=summary, overview_title="Operations overview"), observations=values, charts=[chart])
    result.report_plan.title = "Operations and financial information analysis for the reporting periods (Pages 1 to 3)"
    deck = Presentation(io.BytesIO(build_presentation(result)))
    text = "\n".join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
    assert "COMPLETE_END_MARKER" in text
    assert result.report_plan.title in text
    assert "Cost of sales and Revenue" not in text
    assert text.index("Key findings", text.index("Contents") + len("Contents")) < text.index("Analysis at a glance")
    assert sum(s.has_chart for slide in deck.slides for s in slide.shapes) == 1
    assert _source_footer([3, 4, 5, 8, 9]) == "Source: Document disclosures (p. 3-5, 8-9)"


def test_narrative_pagination_preserves_every_word_and_caveat():
    from adaptive_document_agent.services.pptx_export import _add_text_pages, BUNDLED_TEMPLATE_PATH
    deck = Presentation(str(BUNDLED_TEMPLATE_PATH))
    body = ("The amount was negative, and the interim period is not comparable with a full year.\n\n" * 80) + "FINAL CAVEAT"
    _add_text_pages(deck, "Scope and limitations", body, [1, 2, 3])
    shown = " ".join(s.text for slide in deck.slides for s in slide.shapes if s.name == "narrative:body")
    assert " ".join(shown.split()) == " ".join(body.split())


def test_composition_colors_unique_and_thin_segments_unlabelled():
    from adaptive_document_agent.services.composition_renderer import add_composition_chart
    values = [obs(f"{p}-{c}", v, p, category=c) for p in ("FY2023", "FY2024")
              for c, v in (("East", 40), ("West", 59), ("Small", 1))]
    plan = ChartPlan(id="mix", title="Resources", question="Composition", chart_type="stacked_bar",
        series_dimension="component", observation_ids=[o.id for o in values], show_data_labels=True)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide._ada_colors = {c: "5B21B6" for c in ("East", "West", "Small")}
    add_composition_chart(slide, plan, values, (.5, .5, 8, 4))
    chart = next(s.chart for s in slide.shapes if s.has_chart)
    assert len({str(s.format.fill.fore_color.rgb) for s in chart.series}) == 3
    assert not chart.plots[0].has_data_labels
    assert sorted(sum(s.values) for s in chart.series) == [2, 80, 118]
