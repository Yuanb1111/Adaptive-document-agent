from io import BytesIO
import pytest
from pptx import Presentation
from adaptive_document_agent.models import PresentationPlan, PresentationSlide
from adaptive_document_agent.models.customization import ReportRequirement, ReportRequirements
from adaptive_document_agent.services.presentation_body_budget import place_data_index_after_body
from tests.test_compact_summary_v111 import fixture


@pytest.mark.parametrize('scope,overview,limit', [
    ('body', True, 3), ('all', True, 3),
    ('body', False, 3), ('body', True, 6)])
def test_only_source_index_moves_and_analysis_always_remains_body(scope, overview, limit):
    result, _, _ = fixture(1)
    result.profile.report_requirements = ReportRequirements(items=[ReportRequirement(
        id='length', kind='slide_limit', request_quote='Body limit', description='Body limit',
        max_slides=limit, slide_scope=scope)])
    result.presentation_plan = PresentationPlan(title='Report', slides=[
        PresentationSlide(id='overview', slide_type='analysis', title='Overview', theme_id='theme',
                          slide_role='overview' if overview else 'drivers'),
        PresentationSlide(id='detail', slide_type='analysis', title='Detail', theme_id='theme', slide_role='deep_dive')])
    result.presentation_export_trace = [
        dict(planned_slide_id='overview', topic_id='theme', slide_numbers=[2], observation_ids=['a'], chart_ids=[], source_pages=[1]),
        dict(planned_slide_id='detail', topic_id='theme', slide_numbers=[1,4], observation_ids=['b','c'], chart_ids=[], source_pages=[2])]
    deck = Presentation()
    for text in ['Detail A','Overview','Cover','Detail B','Existing appendix','Closing']:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 1000, 1000).text = text
        if text == 'Existing appendix': slide.element.cSld.set('name', 'ada:data-index')
    before = {s.slide_id: s.shapes[0].element.xml for s in deck.slides}
    place_data_index_after_body(deck, result)
    assert {s.slide_id: s.shapes[0].element.xml for s in deck.slides} == before
    assert len(deck.slides) == 6
    assert sum(s.element.cSld.get('name') != 'ada:data-index' for s in deck.slides) == 5
    assert [s.shapes[0].text for s in deck.slides] == [
        'Detail A', 'Overview', 'Cover', 'Detail B', 'Closing', 'Existing appendix']
    assert result.presentation_export_trace[1]['observation_ids'] == ['b','c']
    for record in result.presentation_export_trace:
        for number in record['slide_numbers']:
            assert ('Overview' if record['planned_slide_id']=='overview' else 'Detail') in deck.slides[number-1].shapes[0].text
    stream=BytesIO(); deck.save(stream)
    reopened=Presentation(stream)
    assert sum(s.element.cSld.get('name')=='ada:data-index' for s in reopened.slides)==1
