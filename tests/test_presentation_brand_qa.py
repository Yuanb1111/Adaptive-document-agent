"""Brand checks must catch generated colour drift without modifying evidence."""

import io

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

from adaptive_document_agent.services.fourier_brand import PRIMARY, TEXT
from adaptive_document_agent.services.presentation_brand_qa import validate_generated_brand


def _chart_deck(*, grouped=False, chart_type=XL_CHART_TYPE.COLUMN_CLUSTERED):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["FY2023", "FY2024"]
    data.add_series("Reported loss", [-25.5, -90.25])
    shapes = slide.shapes.add_group_shape().shapes if grouped else slide.shapes
    chart = shapes.add_chart(chart_type, Inches(1), Inches(1),
                             Inches(6), Inches(3), data).chart
    series = chart.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = RGBColor.from_string(PRIMARY)
    series.format.line.color.rgb = RGBColor.from_string(PRIMARY)
    if chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED:
        series.invert_if_negative = False
    return deck, slide, chart


def test_approved_negative_series_and_workbook_are_not_modified():
    deck, slide, chart = _chart_deck()
    textbox = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(.4))
    textbox.text = "Source: document disclosures (p. 3)"
    textbox.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(TEXT)
    before_slide = slide._element.xml
    before_chart = chart._chartSpace.xml
    before_workbook = chart.part.chart_workbook.xlsx_part.blob
    assert validate_generated_brand(deck) is None
    assert slide._element.xml == before_slide
    assert chart._chartSpace.xml == before_chart
    assert chart.part.chart_workbook.xlsx_part.blob == before_workbook
    assert list(chart.series[0].values) == [-25.5, -90.25]
    assert chart.series[0].invert_if_negative is False


@pytest.mark.parametrize("target", ["text", "series", "point", "marker", "axis"])
def test_forged_rgb_in_generated_content_is_rejected(target):
    deck, slide, chart = _chart_deck(chart_type=XL_CHART_TYPE.LINE_MARKERS
                                   if target == "marker" else XL_CHART_TYPE.COLUMN_CLUSTERED)
    forged = RGBColor.from_string("5B21B6")
    if target == "text":
        box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(.4))
        box.text = "Heading"
        box.text_frame.paragraphs[0].font.color.rgb = forged
    elif target == "series":
        chart.series[0].format.fill.fore_color.rgb = forged
    elif target == "point":
        point = chart.series[0].points[0]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = forged
    elif target == "marker":
        chart.series[0].marker.format.line.color.rgb = forged
    else:
        chart.value_axis.tick_labels.font.color.rgb = forged
    with pytest.raises(ValueError, match="unsupported generated brand colour #5B21B6"):
        validate_generated_brand(deck)


@pytest.mark.parametrize("replacement", [None, "scheme", "no_fill"])
def test_series_fill_cannot_be_left_to_office_automatic_colours(replacement):
    deck, _, chart = _chart_deck()
    properties = chart.series[0]._element.find(qn("c:spPr"))
    properties.remove(properties.find(qn("a:solidFill")))
    if replacement == "scheme":
        fill = OxmlElement("a:solidFill")
        color = OxmlElement("a:schemeClr")
        color.set("val", "accent1")
        fill.append(color)
        properties.insert(0, fill)
    elif replacement == "no_fill":
        properties.insert(0, OxmlElement("a:noFill"))
    # Explicit point colours do not substitute for an explicit series default.
    for point in chart.series[0].points:
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor.from_string(PRIMARY)
    with pytest.raises(ValueError, match="explicit approved solid series fill"):
        validate_generated_brand(deck)


def test_series_line_cannot_be_left_to_office_automatic_colours():
    deck, _, chart = _chart_deck()
    properties = chart.series[0]._element.find(qn("c:spPr"))
    properties.remove(properties.find(qn("a:ln")))
    with pytest.raises(ValueError, match="explicit approved series line"):
        validate_generated_brand(deck)


def test_explicit_no_fill_outline_is_allowed():
    deck, _, chart = _chart_deck()
    chart.series[0].format.line.fill.background()
    validate_generated_brand(deck)


def test_charts_inside_groups_are_checked():
    deck, _, chart = _chart_deck(grouped=True)
    chart.series[0].format.fill.fore_color.rgb = RGBColor.from_string("123456")
    with pytest.raises(ValueError, match="Slide 1, chart .*#123456"):
        validate_generated_brand(deck)


def test_template_master_layout_and_picture_artwork_are_retained():
    from PIL import Image

    deck, slide, _ = _chart_deck()
    for owner in (slide.slide_layout, slide.slide_layout.slide_master):
        # Template gradients may intentionally contain colours outside the
        # generated content palette; their XML is not rewritten or rejected.
        gradient = OxmlElement("a:gradFill")
        stops = OxmlElement("a:gsLst")
        stop = OxmlElement("a:gs")
        stop.set("pos", "0")
        color = OxmlElement("a:srgbClr")
        color.set("val", "9467FF")
        stop.append(color)
        stops.append(stop)
        gradient.append(stops)
        owner._element.append(gradient)
    image = io.BytesIO()
    Image.new("RGB", (2, 2), color="#123456").save(image, format="PNG")
    image.seek(0)
    picture = slide.shapes.add_picture(image, Inches(8), Inches(1), Inches(1), Inches(1))
    # A picture colour effect belongs to that source artwork as well.
    effect = OxmlElement("a:duotone")
    for value in ("123456", "654321"):
        color = OxmlElement("a:srgbClr")
        color.set("val", value)
        effect.append(color)
    picture._element.find(qn("p:blipFill")).find(qn("a:blip")).append(effect)
    originals = [node.xml for node in (slide._element, slide.slide_layout._element,
                                      slide.slide_layout.slide_master._element)]
    blob = picture.image.blob
    validate_generated_brand(deck)
    assert originals == [node.xml for node in (slide._element, slide.slide_layout._element,
                                               slide.slide_layout.slide_master._element)]
    assert picture.image.blob == blob
