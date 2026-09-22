"""Opening slides stay short without erasing evidence or inventing takeaways."""

from pptx import Presentation

from adaptive_document_agent.agent.insight_generator import InsightGenerator, InsightList
from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import (
    AnalysisResult, ChartPlan, DocumentProfile, Insight, Observation,
    ParsedDocument, PipelineResult, PresentationSlide, SourceEvidence,
)
from adaptive_document_agent.services.pptx_export import (
    BUNDLED_TEMPLATE_PATH, _add_document_overview, _add_findings_slide, _add_planned_summary,
)
from adaptive_document_agent.services.presentation_brief import BriefItem, render_brief


def result_fixture():
    values = []
    charts = []
    results = []
    insights = []
    for i, metric in enumerate(("Revenue", "Operating cash flow", "Net loss", "Inventory")):
        ids = []
        for year, value in ((2021, 10_000_000), (2022, 15_000_000), (2023, 20_000_000)):
            ev = SourceEvidence(page=i + 1, text=f"{metric} {year}: {value / 1000:g}",
                                extraction_method="digital_table", confidence=.99)
            oid = f"m{i}-{year}"
            ids.append(oid)
            values.append(Observation(id=oid, metric_original=metric, value=value,
                raw_value=str(value / 1000), currency="RMB", unit="currency", unit_scale=1000,
                raw_unit="RMB thousands", period=f"FY{year}", period_type="fiscal_year",
                evidence=[ev], confidence=.99))
        charts.append(ChartPlan(id=f"c{i}", title=metric, chart_type="line", question="Movement",
                                observation_ids=ids, source_pages=[i + 1]))
        results.append(AnalysisResult(task_id=f"a{i}", title=metric, result={"slope": 5000000},
                                      input_observation_ids=ids, evidence=[v.evidence[0] for v in values[-3:]]))
        insights.append(Insight(id=f"i{i}", title=f"{metric} trend",
            narrative="The supplied trend calculation reports slope 5000000, start_value 10000000, end_value 20000000.",
            kind="calculated_result", result_ids=[f"a{i}"], evidence=[ev], importance=1-i/10))
    return PipelineResult(document=ParsedDocument(document_id="d", sha256="x", safe_filename="test.pdf", page_count=4),
        profile=DocumentProfile(overview_title="Operations overview", document_summary="Operations span several markets. " * 100,
            overview_points=["The source describes operations in several markets.",
                             "Annual comparisons cover 2021 to 2023.",
                             "Reported amounts use RMB thousands."],
            document_summary_pages=[1], important_sections=[f"Section {i}" for i in range(80)]),
        observations=values, charts=charts, analysis_results=results, insights=insights)


def blank_deck():
    deck = Presentation(str(BUNDLED_TEMPLATE_PATH))
    for item in list(deck.slides._sldIdLst):
        deck.part.drop_rel(item.rId)
        deck.slides._sldIdLst.remove(item)
    return deck


def visible(slide):
    return "\n".join(s.text for s in slide.shapes if s.has_text_frame)


def test_long_overview_and_80_sections_produce_one_structured_page():
    result = result_fixture()
    original = result.model_dump_json()
    deck = blank_deck()
    _add_document_overview(deck, result)
    assert len(deck.slides) == 1
    slide = deck.slides[0]
    assert "Section 79" not in visible(slide)
    assert "continued" not in visible(slide)
    assert all(p in visible(slide) for p in result.profile.overview_points)
    assert result.profile.document_summary in slide.notes_slide.notes_text_frame.text
    assert "Section 79" in slide.notes_slide.notes_text_frame.text
    assert result.model_dump_json() == original


def test_legacy_overview_is_labelled_excerpt_and_preserves_complete_context():
    result = result_fixture()
    result.profile.overview_points = []
    result.profile.document_summary += "Important limitation: scope excludes overseas entities."
    deck = blank_deck()
    _add_document_overview(deck, result)
    assert len(deck.slides) == 1
    assert "full context in speaker notes" in visible(deck.slides[0])
    assert result.profile.document_summary in deck.slides[0].notes_slide.notes_text_frame.text


def test_overview_keeps_long_opening_subject_with_adaptive_row_height():
    result = result_fixture()
    result.profile.overview_points = []
    subject = "The manufacturer develops collaborative equipment for industrial customers, " + (
        "with production and sales activities across the disclosed markets, " * 4) + "as described in the source."
    result.profile.document_summary = subject + " The period is annual."
    deck = blank_deck()
    _add_document_overview(deck, result)
    assert subject in visible(deck.slides[0])
    assert "The period is annual." in visible(deck.slides[0])


def test_technical_findings_use_actual_linked_inputs_and_preserve_full_notes():
    result = result_fixture()
    original = result.model_dump_json()
    deck = blank_deck()
    _add_findings_slide(deck, result, result.charts, DocumentIndex(result.observations))
    assert len(deck.slides) == 1
    slide = deck.slides[0]
    text = visible(slide)
    assert "slope" not in text and "start_value" not in text
    assert "FY2021" in text and "FY2023" in text and "RMB" in text
    assert "Revenue" in text and "Inventory" not in text
    assert len([s for s in slide.shapes if s.name == "brief:body"]) == 3
    notes = slide.notes_slide.notes_text_frame.text
    assert "Inventory" in notes and "start_value" in notes
    assert result.model_dump_json() == original


def test_unlinked_debug_is_not_rewritten_using_an_unrelated_metric():
    result = result_fixture()
    result.insights[0].result_ids = ["absent"]
    deck = blank_deck()
    _add_findings_slide(deck, result, [], DocumentIndex(result.observations))
    assert "Revenue" not in visible(deck.slides[0])
    assert "Revenue" in deck.slides[0].notes_slide.notes_text_frame.text


def test_complete_caveat_stays_with_claim_and_oversize_copy_is_not_clipped():
    deck = blank_deck()
    caveat = "Output rose, but this excludes overseas entities."
    long = "Long qualification " * 200 + "FINAL CAVEAT"
    render_brief(deck, "Findings", [BriefItem("", caveat, [2]), BriefItem("", long, [3])])
    slide = deck.slides[0]
    assert caveat in visible(slide)
    assert "Long qualification" not in visible(slide)
    assert long in slide.notes_slide.notes_text_frame.text


def test_chinese_brief_preserves_complete_text_at_readable_size():
    deck = blank_deck()
    text = "报告披露三个年度的经营情况，比较范围不包括未披露业务。"
    render_brief(deck, "关键发现", [BriefItem("经营情况", text, [1])])
    assert text in visible(deck.slides[0])
    bodies = [s for s in deck.slides[0].shapes if s.name == "brief:body"]
    assert bodies
    assert all(p.font.size.pt == 18 for s in bodies for p in s.text_frame.paragraphs)


def test_planned_text_summary_is_also_one_page_with_all_bullets_in_notes():
    result = result_fixture()
    plan = PresentationSlide(id="s", slide_type="executive_summary", title="Summary",
        bullets=[f"Finding {i}: The disclosed scope remains unchanged." for i in range(5)], source_pages=[1])
    deck = blank_deck()
    _add_planned_summary(deck, result, plan, DocumentIndex(result.observations))
    assert len(deck.slides) == 1
    assert "Finding 4" not in visible(deck.slides[0])
    assert "Finding 4" in deck.slides[0].notes_slide.notes_text_frame.text


def test_insight_gateway_receives_normalized_units_and_all_input_periods():
    result = result_fixture()
    captured = {}
    class Gateway:
        def generate_structured(self, messages, schema, **kwargs):
            captured["messages"] = messages
            return InsightList(insights=[result.insights[0]])
    InsightGenerator(Gateway()).generate(result.analysis_results, result.observations)
    payload = captured["messages"][-1]["content"]
    assert '"period": "FY2023"' in payload
    assert '"raw_unit": "RMB thousands"' in payload
    assert '"value": 10000000.0' in payload
    assert '"unit_scale": 1000.0' in payload
