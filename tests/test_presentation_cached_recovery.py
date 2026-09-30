"""Failed planning caches can recover without losing source fields or QA gates."""

import io
import json

import pytest
from pptx import Presentation

from adaptive_document_agent.agent.presentation_plan_recovery import PresentationPlanRecovery
from adaptive_document_agent.agent.presentation_topic_selector import series_directory
from adaptive_document_agent.agent.topic_plan_compiler import compile_topic_plan
from adaptive_document_agent.models import PresentationTopic, PresentationTopicSelection
from adaptive_document_agent.services.export import export_pptx
from adaptive_document_agent.services.qa_reporter import CriticalQAError
from adaptive_document_agent.validation.claim_validator import ClaimValidator
from tests.test_pptx_export import _result
from tests.test_scoped_presentation_recovery import _linked_result
from tests.test_segmented_direction_claims import _series, _slide


def _selected(result, **changes):
    sid = series_directory(result)[0][0]["id"]
    topic = PresentationTopic(id="movement", title="Reported movement", question="How did the measure move?",
        rationale="Comparable source values support this question.", series_ids=[sid])
    result.presentation_topics = PresentationTopicSelection(topics=[topic.model_copy(update=changes)])
    result.presentation_plan = None
    return result


def test_grouped_source_value_cannot_swallow_period_end_year_during_compilation():
    result = _result()
    for item, year in zip(result.observations, (2023, 2024, 2025)):
        item.period = f"{year}-10-31"
        item.period_type = "point_in_time"
        item.period_basis = "point_in_time"
        item.dimensions = {}
    _selected(result, rationale="Period-end observations through 2025-10-31 are retained as reported.")
    raw = result.model_dump(include={"observations", "insights", "presentation_topics", "llm_usage"})
    plan = compile_topic_plan(result)
    slide = next(s for s in plan.slides if s.slide_type == "analysis")
    assert slide.selection_reason == result.presentation_topics.topics[0].rationale
    assert result.model_dump(include=set(raw)) == raw


@pytest.mark.parametrize("preposition,date_kind", [("in", "fiscal_year"), ("at", "point_in_time")])
def test_reversed_value_endpoints_resolve_explicit_source_periods(preposition, date_kind):
    observations = _series("Revenue", [100, 120, 150])
    if date_kind == "point_in_time":
        for o in observations:
            o.period = o.period.removeprefix("FY") + "-12-31"
            o.period_type = date_kind
    start, end = observations[0].period, observations[-1].period
    text = f"Revenue increased to RMB150 {preposition} {end} from RMB100 {preposition} {start}."
    assert not ClaimValidator().validate_slide(_slide(text, observations), observations)
    bad = text.replace(end, "2027-12-31" if date_kind == "point_in_time" else "FY2027")
    assert any(i.code == "direction_scope_ambiguous"
               for i in ClaimValidator().validate_slide(_slide(bad, observations), observations))
    wrong = text.replace("increased", "declined")
    assert any(i.code == "directional_contradiction"
               for i in ClaimValidator().validate_slide(_slide(wrong, observations), observations))


def test_unbound_optional_closing_claim_cannot_block_selected_questions_or_borrow_dates():
    result = _linked_result()
    insight = result.insights[0]
    insight.kind = "calculated_result"
    insight.narrative = "Revenue increased to RMB521.7 million from RMB267.0 million in FY2023."
    insight.watch_item = "Monitor the next reported revenue."
    _selected(result)
    raw = result.model_dump(include={"observations", "insights", "analysis_results"})
    plan = compile_topic_plan(result)
    closing = next(s for s in plan.slides if s.slide_type == "risks")
    assert closing.bullets == [insight.watch_item]
    assert closing.bullet_observation_ids == [[o.id for o in result.observations]]
    audit = json.loads(next(i.message for i in result.validation_warnings
                           if i.code == "presentation_closing_claim_withheld"))
    assert audit["statement"] == insight.narrative
    assert "explicit start and end" in " ".join(audit["errors"])
    assert plan.editorial_status == "needs_review"
    assert result.model_dump(include=set(raw)) == raw


def test_public_export_recovers_missing_plan_without_model_or_evidence_changes(local_render_stub):
    result = _selected(_result())
    raw = result.model_dump(include={"observations", "insights", "presentation_topics", "llm_usage"})
    deck = Presentation(io.BytesIO(export_pptx(result)))
    assert result.presentation_plan.planning_origin == "topic_compilation"
    assert any(shape.has_chart for slide in deck.slides for shape in slide.shapes)
    assert result.model_dump(include=set(raw)) == raw
    snapshot = result.model_dump()
    assert not PresentationPlanRecovery().recover_missing_plan(result)
    assert result.model_dump() == snapshot


def test_invalid_cached_selection_stays_blocked_and_recovery_is_transactional(local_render_stub):
    result = _selected(_result(), title="Revenue reached 999999")
    raw = result.model_dump(include={"observations", "insights", "charts", "presentation_topics", "llm_usage"})
    with pytest.raises(CriticalQAError, match="validated presentation plan"):
        export_pptx(result)
    assert result.presentation_plan is None
    assert result.model_dump(include=set(raw)) == raw
    assert any(i.code == "presentation_cached_recovery_failed" for i in result.validation_warnings)
