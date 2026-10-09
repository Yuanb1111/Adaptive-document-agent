"""Comparative copy requires current identity and complete, visible source facts."""

import copy

import pytest

from adaptive_document_agent.models import (
    ChartPlan, CompanyProfile, DocumentPage, Observation, PresentationPlan,
    PresentationSlide, PresentationTheme, SourceEvidence,
)
from adaptive_document_agent.services.presentation_claim_evidence import prepare_presentation_claims
from adaptive_document_agent.services.presentation_identity import presentation_quality_notes
from tests.test_pptx_export import _result


def _observation(row, year, value, *, share=True, table="sales", state="valid"):
    unit = "percent" if share else "currency"
    return Observation(
        id=f"{table}-{row}-{year}-{unit}", metric_original=row + (": % of Total" if share else ""),
        value=value, raw_value=str(value), unit=unit, currency=None if share else "USD",
        period=f"FY{year}", period_basis="FY", period_type="fiscal_year", table_id=table,
        dimensions={"column_role": "percentage" if share else "amount"},
        validation_status=state, semantic_type="ratio_share" if share else "monetary_amount",
        evidence=[SourceEvidence(page=2, table_id=table, row_label=row,
                                 column_label="% of Total" if share else "Amount", text=str(value),
                                 extraction_method="digital_table", confidence=.9)], confidence=.9,
    )


def _sample(*, rank=False):
    result = _result()
    result.validation_warnings = []
    result.document.pages = [DocumentPage(page_number=1, text="Northstar Limited"), DocumentPage(page_number=2, text="Reported table")]
    shares = [_observation(row, year, value) for row, values in
              (("Enterprise", (30, 60)), ("Consumer", (70, 40)))
              for year, value in zip((2022, 2023), values)]
    amounts = [_observation("Enterprise", year, value, share=False)
               for year, value in ((2022, 300), (2023, 720))]
    result.observations = [*shares, *amounts]
    target = shares[:2] if rank else amounts
    result.charts = [ChartPlan(id="selected", title="Selected measure", chart_type="bar",
                              question="What changed?", observation_ids=[item.id for item in target], source_pages=[2])]
    title = "Enterprise became the largest category." if rank else "Enterprise increased its share."
    slide = PresentationSlide(id="mix", slide_type="analysis", title=title,
                              section_title="Customer mix", chart_ids=["selected"], theme_id="mix",
                              message="Reported sales mix", analytical_question="How did the mix change?",
                              selection_reason="The model selected the sales mix question.", source_pages=[2])
    result.presentation_plan = PresentationPlan(title="Review", company=CompanyProfile(
        name="Northstar Limited", identity_state="RESOLVED", source_pages=[1], field_source_pages={"name": [1]}),
        themes=[PresentationTheme(id="mix", title="Customer mix", question="How did the mix change?",
                                  rationale="Reported category information", chart_ids=["selected"],
                                  observation_ids=[item.id for item in target], source_pages=[2])],
        slides=[PresentationSlide(id="summary", slide_type="executive_summary", title="Summary", bullets=[title]), slide])
    return result


def test_ranking_chart_contains_complete_same_source_comparison_and_keeps_raw():
    result = _sample(rank=True)
    before = copy.deepcopy(result.observations)
    original_chart = result.charts[0].model_dump()
    prepare_presentation_claims(result)
    slide = result.presentation_plan.slides[-1]
    assert "largest" in slide.title
    chart = next(chart for chart in result.charts if chart.id == slide.chart_ids[0])
    assert set(chart.observation_ids) == {item.id for item in result.observations if item.unit == "percent"}
    assert {item.id for item in result.observations if item.unit == "percent"} <= set(result.presentation_plan.themes[0].observation_ids)
    assert result.observations == before and result.charts[0].model_dump() == original_chart
    snapshot = result.model_dump()
    prepare_presentation_claims(result)
    assert result.model_dump() == snapshot


def test_narrowed_comparison_does_not_leave_a_period_only_title():
    result = _sample(rank=True)
    result.observations = result.observations[:2] + result.observations[4:]
    slide = result.presentation_plan.slides[-1]
    slide.section_title = "Enterprise became the largest category, FY2022–FY2023"
    before = [o.model_dump() for o in result.observations]
    prepare_presentation_claims(result)
    assert slide.title == "Enterprise: % of Total"
    assert slide.section_title == slide.title
    assert [o.model_dump() for o in result.observations] == before


@pytest.mark.parametrize("failure", ["partial", "conflict", "wrong_table", "wrong_basis", "unreliable", "not_largest"])
def test_unproved_ranking_is_narrowed_in_slide_and_summary(failure):
    result = _sample(rank=True)
    if failure == "partial":
        result.observations = result.observations[:2] + result.observations[4:]
    elif failure == "conflict":
        result.observations.append(result.observations[3].model_copy(update={"id": "conflict", "value": 41, "raw_value": "41"}))
    elif failure == "wrong_table":
        for item in result.observations[2:4]:
            item.table_id = "another_table"
    elif failure == "wrong_basis":
        for item in result.observations[2:4]:
            item.period_basis = "6M"
    elif failure == "unreliable":
        result.observations[3].validation_status = "invalid"
    else:
        result.observations[1].value, result.observations[1].raw_value = 40, "40"
        result.observations[3].value, result.observations[3].raw_value = 60, "60"
    prepare_presentation_claims(result)
    assert result.presentation_plan.slides[-1].title == "Customer mix"
    assert result.presentation_plan.slides[0].bullets == ["Customer mix"]
    assert len(result.charts) == 1


def test_share_claim_adds_reported_shares_from_exact_selected_row_and_table():
    result = _sample()
    result.observations.extend([_observation("Enterprise", 2022, 90, table="unrelated"),
                                _observation("Enterprise", 2023, 95, table="unrelated")])
    before = copy.deepcopy(result.observations)
    prepare_presentation_claims(result)
    slide = result.presentation_plan.slides[-1]
    assert slide.visual_blocks[0].observation_ids == [item.id for item in result.observations[:2]]
    assert result.observations == before
    snapshot = result.model_dump()
    prepare_presentation_claims(result)
    assert result.model_dump() == snapshot


def test_amount_growth_does_not_prove_share_growth():
    result = _sample()
    result.observations[0].value, result.observations[0].raw_value = 60, "60"
    result.observations[1].value, result.observations[1].raw_value = 40, "40"
    prepare_presentation_claims(result)
    assert result.presentation_plan.slides[-1].title == "Customer mix"
    assert not result.presentation_plan.slides[-1].visual_blocks


def test_explicit_ratio_footnote_preserves_source_backed_possessive_share() -> None:
    result = _sample()
    title = ("A&B expenditure rose from FY2021 to FY2022 while its share of "
             "total expenditure fell.")
    rows = (
        ("Annual alpha and beta expenditure", "currency", (40, 50), "amount"),
        ("Annual total expenditure", "currency", (100, 200), "total"),
        ("Annual alpha and beta expenditure ratio", "percent", (40.0, 25.0), "ratio"),
    )
    result.observations = [Observation(
        id=f"{kind}-{year}", metric_original=row, value=value,
        raw_value=f"{value:.1f}%" if unit == "percent" else str(value),
        unit=unit, period=f"FY{year}", period_basis="FY", period_type="fiscal_year",
        table_id=f"{kind}-table", validation_status="valid",
        evidence=[SourceEvidence(page=2, table_id=f"{kind}-table", row_label=row,
                                 text=str(value), extraction_method="digital_table", confidence=.9)],
        confidence=.9,
    ) for row, unit, values, kind in rows for year, value in zip((2021, 2022), values)]
    result.document.pages[1].text = (
        "Annual alpha and beta expenditure ratio(1) 40.0% 25.0%\n"
        "(1) Calculated by dividing annual alpha and beta expenditure by annual total expenditure."
    )
    result.charts = [ChartPlan(
        id=kind, title=row, chart_type="line", question="How did this measure change?",
        observation_ids=[f"{kind}-{year}" for year in (2021, 2022)], source_pages=[2],
    ) for row, _, _, kind in rows]
    ids = [item.id for item in result.observations]
    summary = PresentationSlide(id="summary", slide_type="executive_summary", title="Summary",
                                bullets=[title], bullet_observation_ids=[ids], observation_ids=ids)
    analysis = PresentationSlide(id="ratio", slide_type="analysis", title=title,
                                 section_title="Expenditure and intensity", chart_ids=[c.id for c in result.charts])
    result.presentation_plan = PresentationPlan(title="Expenditure", slides=[summary, analysis])
    raw = copy.deepcopy(result.observations)
    prepare_presentation_claims(result)
    assert analysis.title == title
    assert summary.bullets == [title]
    assert result.observations == raw
    snapshot = result.model_dump()
    prepare_presentation_claims(result)
    assert result.model_dump() == snapshot

    for change in ("wrong_direction", "wrong_denominator", "wrong_value",
                   "incompatible_currency", "missing_definition"):
        candidate = copy.deepcopy(result)
        selected = candidate.presentation_plan.slides[-1]
        bullet = candidate.presentation_plan.slides[0]
        if change == "wrong_direction":
            selected.title = selected.title.replace("fell", "rose")
            bullet.bullets = [selected.title]
        elif change == "wrong_denominator":
            selected.title = selected.title.replace("total expenditure", "revenue")
            bullet.bullets = [selected.title]
        elif change == "wrong_value":
            ratio = next(item for item in candidate.observations if item.id == "ratio-2022")
            ratio.value, ratio.raw_value = 26.0, "26.0%"
        elif change == "incompatible_currency":
            denominator = next(item for item in candidate.observations if item.id == "total-2022")
            denominator.currency = "USD"
        else:
            candidate.document.pages[1].text = "Annual alpha and beta expenditure ratio 40.0% 25.0%"
        prepare_presentation_claims(candidate)
        assert candidate.presentation_plan.slides[-1].title == "Expenditure and intensity"
        assert candidate.presentation_plan.slides[0].bullets == ["Expenditure and intensity"]


def test_unsupported_share_keeps_independent_topic_measure_instead_of_generic_title():
    result = _sample()
    slide = result.presentation_plan.slides[-1]
    slide.section_title = "Enterprise amount, Revenue Share"
    slide.message = "Its share increased."
    prepare_presentation_claims(result)
    assert slide.title == "Enterprise amount"
    assert slide.message == "Reported values for Enterprise amount across the cited periods."
    assert result.presentation_plan.slides[0].bullets == ["Enterprise amount"]


def test_share_direction_binds_to_its_clause_subject_not_first_named_category():
    result = _sample()
    consumers = [_observation("Consumer", year, value, share=False) for year, value in ((2022, 700), (2023, 480))]
    result.observations.extend(consumers)
    slide = result.presentation_plan.slides[-1]
    slide.observation_ids = [item.id for item in consumers]
    slide.title = "Enterprise revenue grew while Consumer increased its share."
    prepare_presentation_claims(result)
    assert slide.title == "Customer mix"
    assert not slide.visual_blocks


def test_ranking_does_not_bypass_contradictory_share_direction():
    result = _sample(rank=True)
    slide = result.presentation_plan.slides[-1]
    slide.title = "Enterprise became the largest category while its share declined."
    prepare_presentation_claims(result)
    assert slide.title == "Customer mix"
    assert len(result.charts) == 1


def test_neutral_share_question_does_not_remove_an_evidenced_ranking():
    result = _sample(rank=True)
    slide = result.presentation_plan.slides[-1]
    slide.message = "How did the category's revenue share and volume evolve?"
    prepare_presentation_claims(result)
    assert "largest" in slide.title
    assert len(result.charts) == 2


def test_narrowing_claim_preserves_independent_qualifications():
    result = _sample(rank=True)
    result.observations = result.observations[:2]
    slide = result.presentation_plan.slides[-1]
    slide.message = "Amounts are preliminary and unaudited; no causal inference is supported."
    original = slide.message
    prepare_presentation_claims(result)
    assert slide.title == "Customer mix"
    assert slide.message == original


def test_missing_shares_do_not_invent_denominator_or_use_unrelated_table():
    result = _sample()
    result.observations = result.observations[4:] + [
        _observation("Enterprise", 2022, 30, table="unrelated"),
        _observation("Enterprise", 2023, 60, table="unrelated")]
    prepare_presentation_claims(result)
    assert result.presentation_plan.slides[-1].title == "Customer mix"


@pytest.mark.parametrize("resolved", [True, False])
def test_identity_warning_is_removed_only_after_retained_name_evidence(resolved):
    result = _sample()
    original = "The company's name is not stated in the supplied summaries, so its identity cannot be confirmed from this material alone."
    result.profile.data_quality_notes = [original, "Source page labels may be inconsistent.",
                                         "Issuer name is not stated in supplied source pages; document analysed as general corporate document."]
    result.presentation_plan.coverage_notes = [original]
    result.presentation_plan.themes[0].caveats = [original]
    if not resolved:
        result.presentation_plan.company.field_source_pages["name"] = [999]
    notes = presentation_quality_notes(result)
    prepare_presentation_claims(result)
    assert result.profile.data_quality_notes[0] == original
    assert (original not in notes) is resolved
    assert "Source page labels may be inconsistent." in notes
    assert (not result.presentation_plan.coverage_notes) is resolved
    assert (not result.presentation_plan.themes[0].caveats) is resolved


def test_identity_filter_preserves_unrelated_sentence_in_same_note():
    result = _sample()
    result.profile.data_quality_notes = ["Company name is unknown. Revenue values conflict between two sources."]
    assert presentation_quality_notes(result) == ["Revenue values conflict between two sources."]


def test_sourced_name_does_not_remove_a_conflicting_identity_warning():
    result = _sample()
    note = "Company identity cannot be confirmed because source names conflict."
    result.profile.data_quality_notes = [note]
    assert presentation_quality_notes(result) == [note]


def test_identity_cleanup_keeps_same_note_risk_and_other_entities():
    result = _sample()
    result.profile.data_quality_notes = [
        "The company's name is unknown; Revenue values conflict between two sources.",
        "Supplier company name is not stated in the source material.",
        "Company name is unknown but auditor qualifications remain unresolved.",
    ]
    assert presentation_quality_notes(result) == [
        "Revenue values conflict between two sources.",
        "Supplier company name is not stated in the source material.",
        "auditor qualifications remain unresolved.",
    ]


def test_expense_rates_summing_to_one_hundred_do_not_prove_common_denominator():
    result = _sample(rank=True)
    for item in result.observations:
        if item.unit == "percent":
            item.evidence[0].column_label = "R&D / segment revenue (%)"
    prepare_presentation_claims(result)
    assert result.presentation_plan.slides[-1].title == "Customer mix"
    assert len(result.charts) == 1


def test_unrelated_quality_copy_keeps_punctuation_and_contrast():
    result = _sample()
    note = "Revenue increased, but losses widened; no causal inference is supported."
    result.profile.data_quality_notes = [note]
    assert presentation_quality_notes(result) == [note]


def test_full_qa_retains_raw_discovery_and_validation_notes():
    from adaptive_document_agent.models import ValidationIssue
    from adaptive_document_agent.services.qa_reporter import run_comprehensive_qa
    result = _sample()
    result.profile.data_quality_notes = [
        "Issuer name is not stated in supplied source pages; document analysed as general corporate document.",
        "Company identity cannot be confirmed because source names conflict.",
    ]
    result.validation_warnings = [ValidationIssue(code="identity_context", severity="warning",
        stage="discovery", message="Issuer name is not stated; source values conflict.")]
    notes, warnings = copy.deepcopy(result.profile.data_quality_notes), copy.deepcopy(result.validation_warnings)
    run_comprehensive_qa(result)
    assert result.profile.data_quality_notes == notes
    assert result.validation_warnings == warnings
    displayed = presentation_quality_notes(result)
    assert "Company identity cannot be confirmed because source names conflict." in displayed
    assert any("source values conflict" in note for note in displayed)


def test_narrowed_message_and_summary_keep_audit_qualification():
    result = _sample(rank=True)
    result.observations = result.observations[:2]
    slide = result.presentation_plan.slides[-1]
    slide.title = "Customer mix"
    slide.message = "Enterprise became the largest category; amounts are preliminary and unaudited."
    result.presentation_plan.slides[0].bullets = [slide.message]
    prepare_presentation_claims(result)
    assert slide.message == "amounts are preliminary and unaudited."
    assert result.presentation_plan.slides[0].bullets == [slide.message]


def test_compact_category_comparison_keeps_values_and_readable_labels():
    from pptx import Presentation
    from adaptive_document_agent.services.pptx_export import _add_native_chart, _series_rows
    result = _sample(rank=True)
    values = [item for item in result.observations if item.unit == "percent"]
    values += [_observation("Services", year, value) for year,value in ((2022,10),(2023,10))]
    plan = ChartPlan(id="comparison",title="Reported category shares",chart_type="bar",
                     observation_ids=[item.id for item in values],question="How do categories compare?")
    rows = _series_rows(plan, values)
    assert all("Column Role" not in name for _,name,_ in rows)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _add_native_chart(slide,plan,values,(.5,1,5.5,3))
    chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
    assert [list(series.values) for series in chart.series] == [[30,60],[70,40],[10,10]]
    assert 9 <= chart.plots[0].data_labels.font.size.pt <= 11
