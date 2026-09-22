"""Analytical quality, rendering intent and artwork isolation regressions."""

import io

import pytest
from PIL import Image
from pptx import Presentation

from adaptive_document_agent.agent.presentation_planner import PresentationPlanner
from adaptive_document_agent.agent.presentation_plan_recovery import PresentationPlanRecovery
from adaptive_document_agent.models import PresentationVisualBlock
from adaptive_document_agent.services.export import _build_cache_key
from adaptive_document_agent.services.pptx_export import build_presentation
from adaptive_document_agent.services.presentation_artwork import validate_artwork
from adaptive_document_agent.services.presentation_editorial import review_presentation, distinct_findings
from adaptive_document_agent.services.single_metric_slides import enrich_single_metric_slides
from adaptive_document_agent.services.slide_compositor import compose_geometry, validate_composed_geometry
from tests.test_p1_theme_planning import gateway_for, themed_result


def artwork_bytes():
    # Test input only, not a generated user-facing visual asset.
    output = io.BytesIO()
    Image.new("RGB", (320, 200), (90, 40, 150)).save(output, format="PNG")
    return output.getvalue()


def test_weak_but_safe_plan_gets_one_editorial_repair():
    result = themed_result()
    weak = result.presentation_plan.model_copy(deep=True)
    weak.slides[3].message = "Evidence-backed comparison of retained reported values."
    gateway, client = gateway_for(weak.model_dump(mode="json"), result.presentation_plan.model_dump(mode="json"))
    plan = PresentationPlanner(gateway).plan(result)
    assert len(client.calls) == 2
    assert "generic subtitle" in client.calls[1][2]["content"]
    assert plan.planning_origin == "repaired" and plan.editorial_status == "ready"


def test_model_cannot_certify_a_weak_plan_and_retry_is_bounded():
    result = themed_result()
    weak = result.presentation_plan.model_copy(deep=True)
    weak.slides[3].message = "Evidence-backed comparison of retained reported values."
    weak.editorial_status = "ready"
    gateway, client = gateway_for(weak.model_dump(mode="json"), weak.model_dump(mode="json"))
    plan = PresentationPlanner(gateway).plan(result)
    assert len(client.calls) == 2
    assert plan.editorial_status == "needs_review"
    assert plan.editorial_notes


def test_evidence_only_recovery_is_permanently_labelled():
    result = themed_result()
    before = [o.model_dump() for o in result.observations]
    plan = PresentationPlanRecovery().fallback(result)
    assert plan.planning_origin == "fallback" and plan.editorial_status == "degraded"
    assert "presentation_degraded" in {f.code for f in review_presentation(plan, result)}
    assert before == [o.model_dump() for o in result.observations]


def test_failed_editorial_revision_retains_safe_original_instead_of_fallback():
    result = themed_result()
    weak = result.presentation_plan.model_copy(deep=True)
    weak.slides[3].message = "Evidence-backed comparison of retained reported values."
    unsafe = weak.model_copy(deep=True)
    unsafe.slides[3].title = "Revenue increased by 99999%"
    gateway, client = gateway_for(weak.model_dump(mode="json"), unsafe.model_dump(mode="json"))
    plan = PresentationPlanner(gateway).plan(result)
    assert len(client.calls) == 2
    assert plan.slides[3].title == weak.slides[3].title
    assert plan.planning_origin == "model" and plan.editorial_status == "needs_review"
    assert "could not be retained" in plan.editorial_notes[-1]


def test_single_chart_theme_keeps_planned_bar_and_composition():
    result = themed_result()
    slide = result.presentation_plan.slides[3]
    slide.chart_ids = ["a"]
    slide.observation_ids = []
    slide.visual_blocks = [PresentationVisualBlock(role="hero", chart_ids=["a"], chart_type="bar")]
    slide.bullets = []
    slide.layout = "chart_plus_commentary"
    assert enrich_single_metric_slides(result) == []
    deck = Presentation(io.BytesIO(build_presentation(result)))
    rendered = next(s for s in deck.slides if s.name.startswith("composed_"))
    assert slide.layout == "chart_plus_commentary"
    assert not any(s.name == "single_metric_hero" for s in deck.slides)
    assert "Start value" not in " ".join(s.text for s in rendered.shapes if s.has_text_frame)
    assert any(s.name == "chart:a" for s in rendered.shapes)


@pytest.mark.parametrize("charts", [0, 1, 2, 3])
def test_kpi_band_preserves_nonoverlapping_top_chart_and_takeaway_slots(charts):
    geometry = compose_geometry(12.8, 7.2, 1.5, charts, layout="kpi_band", has_support=True, has_commentary=True)
    assert geometry.support.y == 1.5
    for chart in geometry.charts:
        assert chart.y > geometry.support.y + geometry.support.h
        assert chart.y + chart.h < geometry.commentary.y
    assert geometry.commentary.y + geometry.commentary.h < geometry.footer.y


def test_kpi_overview_renders_selected_values_even_when_they_also_appear_in_charts():
    result = themed_result()
    slide = result.presentation_plan.slides[3]
    slide.layout = "kpi_band"
    slide.visual_blocks[-1].observation_ids = ["a2", "b2"]
    deck = Presentation(io.BytesIO(build_presentation(result)))
    rendered = next(s for s in deck.slides if s.name == "composed_kpi_band")
    shapes = {s.name: s for s in rendered.shapes}
    assert shapes["kpi:a2"].top < shapes["chart:a"].top
    assert shapes["kpi:b2"].top < shapes["chart:b"].top
    validate_composed_geometry(deck)


def test_summary_without_charts_keeps_explicit_kpis():
    result = themed_result()
    summary = result.presentation_plan.slides[2]
    summary.layout = "kpi_band"
    summary.source_pages = [3]
    summary.visual_blocks = [PresentationVisualBlock(role="kpi", observation_ids=["a2"])]
    summary.bullets = ["Operating activity expanded."]
    deck = Presentation(io.BytesIO(build_presentation(result)))
    assert any(s.name == "kpi:a2" for sl in list(deck.slides)[:4] for s in sl.shapes)


def test_artwork_is_local_optional_and_invalidates_native_cache():
    result = themed_result()
    picture = artwork_bytes()
    before = result.model_dump()
    assert _build_cache_key(result, "template") != _build_cache_key(result, "template", picture)
    deck = Presentation(io.BytesIO(build_presentation(result, artwork=picture)))
    assert deck.slides[0].name == "picture_cover"
    assert any(sl.name == "picture_profile" for sl in deck.slides)
    pictures = [s for sl in deck.slides for s in sl.shapes if s.name == "user_supplied_illustration"]
    assert len(pictures) == 2
    assert all(abs(s.width / s.height - 1.6) < .001 for s in pictures)
    assert all(s.image.blob == picture for s in pictures)
    assert result.observations == type(result).model_validate(before).observations
    assert "artwork" not in result.model_dump_json()


@pytest.mark.parametrize("payload", [b"", b"not a picture", b"x" * 8_000_001, b'<svg xmlns="http://www.w3.org/2000/svg"/>'], ids=["empty", "invalid", "oversized", "svg"])
def test_artwork_rejects_invalid_or_active_content(payload):
    with pytest.raises(ValueError):
        validate_artwork(payload)


def test_summary_deduplication_keeps_explanation_and_material_caveats():
    assert distinct_findings([("", "Revenue CAGR"), ("Revenue CAGR", "Comparable annual periods."),
                             ("Revenue CAGR", "Comparable annual periods."),
                             ("Revenue CAGR", "Interim observations are excluded.")]) == [
        ("Revenue CAGR", "Comparable annual periods."), ("Revenue CAGR", "Interim observations are excluded.")]
    contradictory = [("Change", "+5%"), ("Change", "-5%"), ("Change", "0.5%"), ("Change", "0,5%")]
    assert distinct_findings(contradictory) == contradictory


def test_commentary_pagination_preserves_paragraphs_and_unbroken_identifiers():
    from adaptive_document_agent.services.slide_compositor import _put_commentary, Rect
    from adaptive_document_agent.services.pptx_export import _base_slide
    deck = Presentation()
    text = "First paragraph.\n\n" + "abcdefghij" * 40 + "\nFINAL DETAIL PRESERVED"
    remaining, retained = text, []
    while remaining:
        slide = _base_slide(deck, "Preview")
        remaining = _put_commentary(slide, remaining, Rect(.5, 1, 5, 1))
        retained.append(next(s.text for s in slide.shapes if s.name.startswith("composed:")))
    # python-pptx exposes paragraph-internal soft breaks as vertical tabs.
    assert "".join(retained).replace("\v", "\n") == text
