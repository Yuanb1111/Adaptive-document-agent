"""Regression coverage for the observed briefing and presentation defects."""

from copy import deepcopy
from datetime import date
import io

from pptx import Presentation

from adaptive_document_agent.models import PresentationSlide
from adaptive_document_agent.models.executive_brief import ExecutiveBrief
from adaptive_document_agent.services.brief_money_display import readable_money
from adaptive_document_agent.services.presentation_closing import render_closing
from adaptive_document_agent.services.presentation_summary import render_complete_summary
from adaptive_document_agent.services.presentation_brief import BriefItem
from tests.test_executive_brief import payload, result_for
from tests.test_presentation_brief import blank_deck


def test_money_display_is_exact_idempotent_and_preserves_rates():
    text = 'Revenue RMB 286,749 thousand; loss USD -103,281 thousand; ASP RMB 47.1 thousand/unit.'
    shown = readable_money(text)
    assert shown == 'Revenue RMB 286.749 million; loss USD -103.281 million; ASP RMB 47.1 thousand/unit.'
    assert readable_money(shown) == shown
    assert readable_money('USD +12,300 thousand') == 'USD +12.3 million'
    assert readable_money('USD 12,300 thousand per employee') == 'USD 12,300 thousand per employee'


def test_closing_does_not_repeat_editorial_summary_and_preserves_prior_plan():
    result = result_for(['Revenue was USD 12 million.'])
    result.executive_brief = ExecutiveBrief.model_validate(payload('Revenue was USD 12 million.'))
    deck = blank_deck()
    render_complete_summary(deck, 'Summary', [BriefItem('Scale', 'Revenue was USD 12 million.', [1])])
    plan = PresentationSlide(id='end', slide_type='risks', title='Old closing', bullets=['Prior draft'])
    before = deepcopy(result.model_dump())
    assert render_closing(deck, result, plan) == []
    assert len(deck.slides) == 1
    assert 'Prior draft' in deck.slides[0].notes_slide.notes_text_frame.text
    assert result.model_dump() == before


def test_summary_uses_one_font_across_continuations_and_retains_every_item():
    items = [BriefItem(f'Finding {i}', ('A complete qualified source finding. ' * (6 + i)), [i + 1])
             for i in range(5)]
    pages = render_complete_summary(blank_deck(), 'Summary', items, single_column=True)
    bodies = [shape for page in pages for shape in page.shapes if shape.name == 'brief:body']
    assert [shape.text for shape in bodies] == [item.text for item in items]
    assert len({s.text_frame.paragraphs[0].font.size.pt for s in bodies}) == 1
    assert all(s.top.inches + s.height.inches < 6.05 for s in bodies)


def test_source_dates_remain_visible_without_changing_calendar_spacing():
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    from adaptive_document_agent.services.presentation_axes import style_date_axis, label_source_dates

    deck = blank_deck()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    dates = [date(2023, 12, 31), date(2024, 6, 30), date(2024, 10, 31)]
    data = CategoryChartData()
    data.categories = dates
    data.add_series('Cash', [110.32, 73.033, 81.324])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS,
        Inches(1), Inches(1), Inches(6), Inches(4), data).chart
    style_date_axis(chart, dates, 6)
    label_source_dates(chart, dates)
    stream = io.BytesIO()
    deck.save(stream)
    chart = next(s.chart for s in Presentation(io.BytesIO(stream.getvalue())).slides[0].shapes if s.has_chart)
    assert chart._chartSpace.xpath('.//c:dateAx')
    assert list(chart.series[0].values) == [110.32, 73.033, 81.324]
    assert '30 Jun 2024' in chart.series[0].points[1].data_label.text_frame.text
    assert '81.32' in chart.series[0].points[2].data_label.text_frame.text


def test_share_denominator_is_bound_to_exact_source_table():
    from adaptive_document_agent.services.presentation_labels import source_share_heading
    from tests.test_presentation_claim_evidence import _sample
    result = _sample(rank=True)
    share = next(o for o in result.observations if o.unit == 'percent')
    original = share.model_dump()
    label = 'Enterprise: % of Total'
    assert source_share_heading(label, [share], {share.effective_table_id: 'revenue'}) == 'Enterprise: % of revenue'
    assert source_share_heading(label, [share], {'another_table': 'revenue'}) == label
    from adaptive_document_agent.services.presentation_labels import source_fact_caption
    caption = label + ': FY2023 40.0%.'
    assert source_fact_caption(caption, [share], {share.effective_table_id: 'revenue'}) == (
        'Enterprise: % of revenue: FY2023 40.0%.')
    unrelated = 'Another category: % of Total: FY2023 40.0%.'
    assert source_fact_caption(unrelated, [share], {share.effective_table_id: 'revenue'}) == unrelated
    assert share.model_dump() == original


def test_appendix_uses_chart_precision_instead_of_three_significant_digits():
    from adaptive_document_agent.services.pptx_export import _appendix_display_value
    from adaptive_document_agent.document_model.metric_semantic_classifier import classify_metric
    from tests.test_pptx_export import _result
    item = _result().observations[0].model_copy(update={'value': -103281000, 'raw_value': '(103,281)',
                                                      'unit_scale': 1000})
    semantic = classify_metric(item.metric_original, unit=item.unit, raw_unit=item.raw_unit, value=item.value)
    assert _appendix_display_value(item, semantic) == '-103.28'
