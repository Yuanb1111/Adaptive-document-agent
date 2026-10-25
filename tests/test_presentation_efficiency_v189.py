"""Lossless context factoring and measured same-topic commentary packing."""

from copy import deepcopy
import json

import pytest
from pptx.util import Inches

from adaptive_document_agent.agent.coverage_review_context import compact_context, encode
from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import ChartPlan
from adaptive_document_agent.services.composition_commentary_packing import pack_commentary_pages
from adaptive_document_agent.services.llm import LLMSettings, ProviderName
from adaptive_document_agent.services.llm.reasoning_policy import request_reasoning_policy
from adaptive_document_agent.services.slide_compositor import (
    Rect, _base, _put_commentary, _put_text, render_composed_slide, validate_composed_geometry,
)
from tests.test_bounded_topic_coverage import context_fixture
from tests.test_chart_capacity_pagination import _deck
from tests.test_p0_composition import observation, result_for


def restore_cells(wire):
    restored = deepcopy(wire)
    dictionaries = restored.pop('series_value_dictionaries', {})
    restored.pop('series_cell_encoding', None)
    for row in restored['all_extracted_series']:
        for i, name in enumerate(restored['series_columns']):
            if name in dictionaries:
                row[i] = dictionaries[name][row[i]]
    for name, field in [('source_row_qualifications', 'notes'), ('ratio_definitions', 'definitions')]:
        dictionary = restored.pop(name + '_dictionary', None)
        restored.pop(name + '_encoding', None)
        if dictionary is not None:
            restored[name] = [dict(series_id=row['series_id'], **{field: dictionary[row[field + '_ref']]})
                              for row in restored[name]]
    return restored


def test_wire_roundtrip_preserves_all_points_qualifications_types_and_input():
    *_, payload = context_fixture(extra=20)
    source_note = {'page': 3, 'table_id': 'owned-table', 'text': 'A literal scoped condition. ' * 40}
    ids = [row[0] for row in payload['all_extracted_series']]
    payload['source_row_qualifications'] = [dict(series_id=sid, notes=[source_note]) for sid in ids]
    payload['ratio_definitions'] = [dict(series_id=sid, definitions=[{'text': 'A / B', 'page': 3}]) for sid in ids]
    context = compact_context(payload)
    for n, row in enumerate(context['all_extracted_series']):
        row['future_metadata'] = [None, False, 0, '0', [], {}][n % 6]
    before = deepcopy(context)
    plain = encode(context)
    packed = encode(context, factor_cells=True)
    wire = json.loads(packed)
    assert wire['series_value_dictionaries'] and wire['source_row_qualifications_dictionary']
    assert 'id' not in wire['series_value_dictionaries']
    assert restore_cells(wire) == json.loads(plain)
    assert len(packed) < len(plain) * .65
    assert context == before
    assert encode(wire, factor_cells=True) == packed


def test_small_context_keeps_literal_cells_without_dictionary_overhead():
    context = {'series_columns': ['id', 'metric', 'reported_points'],
               'all_extracted_series': [['series', 'Metric', [['FY2024', '13', 13.0]]]]}
    assert encode(context, factor_cells=True) == encode(context)


@pytest.mark.parametrize('operation', ['IntroductionAudienceReview', 'SummaryLayoutPages', 'PhysicalBudgetSelection'])
def test_editorial_policy_keeps_provider_optout_and_exact_route_guards(operation):
    settings = LLMSettings(provider=ProviderName.OPENAI, model='gpt-6-luna')
    def policy(stage='presentation'):
        return request_reasoning_policy(settings, stage=stage, operation=operation, model=settings.model)
    assert policy()['options'] == {'reasoning_effort': 'low'}
    assert policy('insight')['options'] == {}
    settings.simple_task_reasoning = 'provider_default'
    assert policy()['options'] == {}
    settings.simple_task_reasoning = 'reduced'
    settings.provider = ProviderName.OPENAI_COMPATIBLE
    settings.base_url = 'http://127.0.0.1:8000/v1'
    assert policy()['options'] == {} and policy()['status'] == 'unsupported'


def packing_case(copy='Reported observations retain their source scope.', subtitle=''):
    values = [observation(f'o-{year}', 'Output', value, f'FY{year}', unit='units')
              for year, value in [(2022, 20), (2023, 32), (2024, 28)]]
    chart = ChartPlan(id='output', title='Output', question='Reported output', chart_type='bar',
                      observation_ids=[o.id for o in values], source_pages=[3])
    result = result_for(values, [chart])
    index = DocumentIndex(values)
    deck = _deck()
    target = render_composed_slide(deck, result.presentation_plan.slides[3], [chart], result, index)[0]
    donor, top = _base(deck, 'Reported performance (continued)', subtitle)
    donor.name = 'composed_continuation'
    assert not _put_commentary(donor, copy, Rect(.55, top, deck.slide_width.inches - 1.1, 4.5))
    citation = next(s.text for s in target.shapes if s.has_text_frame and s.text.startswith('Source:'))
    _put_text(donor, citation, Rect(.55, deck.slide_height.inches - .82, deck.slide_width.inches - 1.1, .2))
    donor.notes_slide.notes_text_frame.text = 'Exact original evidence record'
    return deck, target, donor, [chart], index, result


def test_short_continuation_packs_without_losing_copy_evidence_or_chart_values():
    deck, target, donor, charts, index, result = packing_case()
    before = result.model_dump()
    original_values = [list(series.values) for series in next(s.chart for s in target.shapes if s.has_chart).series]
    assert pack_commentary_pages(deck, [target, donor], charts, index) == [target]
    assert len(deck.slides) == 1
    assert 'Reported observations retain their source scope.' in '\n'.join(s.text for s in target.shapes if s.has_text_frame)
    notes = json.loads(target.notes_slide.notes_text_frame.text)
    assert notes['packed_commentary_pages'][0]['original_notes'] == 'Exact original evidence record'
    assert [list(series.values) for series in next(s.chart for s in target.shapes if s.has_chart).series] == original_values
    assert min(s.height.inches for s in target.shapes if s.has_chart) >= 1.8
    assert result.model_dump() == before
    validate_composed_geometry(deck)


@pytest.mark.parametrize('block', ['different_source', 'missing_source', 'no_plot_space', 'qualified_subtitle', 'geometry_failure'])
def test_unfit_or_unsafe_candidate_retains_original_page_and_geometry(monkeypatch, block):
    deck, target, donor, charts, index, _ = packing_case(
        subtitle='Unaudited source columns only' if block == 'qualified_subtitle' else '')
    if block == 'different_source':
        next(s for s in donor.shapes if s.has_text_frame and s.text.startswith('Source:')).text = 'Source: Other page 8'
    elif block == 'missing_source':
        for slide in (donor, target):
            for shape in list(slide.shapes):
                if shape.has_text_frame and shape.text.startswith('Source:'):
                    shape.element.getparent().remove(shape.element)
    elif block == 'no_plot_space':
        chart = next(s for s in target.shapes if s.has_chart)
        chart.top = Inches(deck.slide_height.inches - 1.02 - 1.8)
        chart.height = Inches(1.8)
    elif block == 'geometry_failure':
        from adaptive_document_agent.services import slide_compositor
        def fail(*a):
            raise ValueError('Rejected geometry')
        monkeypatch.setattr(slide_compositor, 'validate_composed_geometry', fail)
    original = [(s.element.xml, s.notes_slide.notes_text_frame.text) for s in deck.slides]
    assert pack_commentary_pages(deck, [target, donor], charts, index) == [target, donor]
    assert [(s.element.xml, s.notes_slide.notes_text_frame.text) for s in deck.slides] == original
