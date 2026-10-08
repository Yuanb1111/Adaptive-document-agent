import io,json
from pptx import Presentation
from adaptive_document_agent.agent.source_coverage import section_ledger
from adaptive_document_agent.models import AnalysisPageRange
from adaptive_document_agent.services.export import export_pptx
from tests.test_p1_theme_planning import themed_result


def test_actual_native_slide_mapping_survives_serialized_notes(local_render_stub):
    result=themed_result()
    result.profile.source_coverage=section_ledger(result.document,[AnalysisPageRange(title='Complete',start_page=1,end_page=5,reason='Test scope')])
    deck=Presentation(io.BytesIO(export_pptx(result)))
    record=next(r for r in result.presentation_export_trace if r['topic_id']=='operations')
    assert {'a0','a1','a2','b0','b1','b2','kpi'}<=set(record['observation_ids'])
    for number in record['slide_numbers']:
        notes=deck.slides[number-1].notes_slide.notes_text_frame.text
        try:persisted=json.loads(notes)['ADA_EXPORT_TRACE_V1']
        except ValueError:persisted=json.loads(notes.split('ADA_EXPORT_TRACE_V1\n')[-1])
        assert persisted['topic_id']=='operations' and persisted['physical_slide']==number
    section=result.profile.source_coverage.sections[0]
    assert section.export_mapping_status=='recorded_renderer_scope'
    assert section.exported_topic_ids==['operations'] and set(record['slide_numbers'])<=set(section.exported_slide_numbers)
