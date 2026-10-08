"""Resolved identities and complete directory labels survive presentation export."""

import json

import pytest
from pptx import Presentation

from adaptive_document_agent.models import DocumentPage, PresentationSlide
from adaptive_document_agent.services.pptx_export import (
    BUNDLED_TEMPLATE_PATH, _add_company_at_a_glance, _add_cover, _add_planned_contents,
)
from tests.test_company_summary_pages import sample


def _deck():
    deck = Presentation(str(BUNDLED_TEMPLATE_PATH))
    for slide_id in list(deck.slides._sldIdLst):
        deck.part.drop_rel(slide_id.rId)
        deck.slides._sldIdLst.remove(slide_id)
    return deck


def _resolved_result(name="Northstar Instruments Limited"):
    result = sample()
    company = result.presentation_plan.company
    company.name = name
    company.identity_state = "RESOLVED"
    company.field_source_pages["name"] = [1]
    company.source_pages = [1, 5, 18]
    result.document.pages.insert(0, DocumentPage(page_number=1, text=name))
    return result


def _text(slide):
    return "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)


@pytest.mark.parametrize("name,topic", [
    ("Northstar Instruments Limited", "Equipment Operations Review"),
    ("Harbor Software Corporation", "Platform Services Annual Review"),
])
def test_cover_displays_resolved_company_and_retains_document_scope(name, topic):
    result = _resolved_result(name)
    before = result.model_dump()
    title = topic + ": Operational performance, working capital and longer term investment plans"
    deck = _deck()
    _add_cover(deck, result, title=title, purpose="Review the reported operations and their limitations.")
    visible = _text(deck.slides[0])
    assert name in visible and topic in visible
    notes = json.loads(deck.slides[0].notes_slide.notes_text_frame.text)
    assert notes["planned_title"] == title
    assert notes["company_identity"] == {"name": name, "source_pages": [1]}
    assert result.model_dump() == before


def test_cover_does_not_repeat_company_already_in_planned_title():
    result = _resolved_result()
    name = result.presentation_plan.company.name
    deck = _deck()
    _add_cover(deck, result, title=name + " Review", purpose="Document analysis")
    assert _text(deck.slides[0]).count(name) == 1


def test_unresolved_identity_is_not_promoted_to_cover_or_summary():
    result = _resolved_result()
    result.presentation_plan.company.identity_state = "UNRESOLVED"
    name = result.presentation_plan.company.name
    deck = _deck()
    _add_cover(deck, result, title="Operational Review", purpose="Document analysis")
    _add_company_at_a_glance(deck, result, PresentationSlide(
        id="company", slide_type="company_overview", title="Company overview"))
    assert name not in _text(deck.slides[0])
    assert name not in _text(deck.slides[1])


def test_summary_introduction_keeps_company_identity_citation_and_source_copy():
    result = _resolved_result()
    before = result.model_dump()
    deck = _deck()
    _add_company_at_a_glance(deck, result, PresentationSlide(
        id="company", slide_type="company_overview", title="Document at a Glance"))
    assert len(deck.slides) == 2
    assert result.presentation_plan.company.name in _text(deck.slides[0])
    assert "Source: Document disclosures (p. 1, 5)" in _text(deck.slides[0])
    assert "Source: Document disclosures (p. 18)" in _text(deck.slides[1])
    assert "We operate a business software platform." in _text(deck.slides[0])
    assert "We provide payroll and scheduling services." in _text(deck.slides[1])
    assert "Company identity: Northstar Instruments Limited" in deck.slides[0].notes_slide.notes_text_frame.text
    assert result.model_dump() == before


def _contents_entries(deck):
    return [shape for slide in deck.slides for shape in slide.shapes if shape.name.startswith("contents:entry:")]


def test_contents_keeps_complete_financial_terms_and_colon_qualifications():
    labels = ["Liquidity, leverage and operating cash flow",
              "Product mix and average selling price",
              "Regional performance and working capital: comparable annual periods",
              "Regional performance and working capital: separate interim periods"]
    plans = [PresentationSlide(id=str(index), slide_type="analysis", title="A supported finding",
                               section_title=label) for index, label in enumerate(labels)]
    before = [item.model_dump() for item in plans]
    deck = _deck()
    _add_planned_contents(deck, plans)
    assert [shape.text for shape in _contents_entries(deck)] == labels
    assert [item.model_dump() for item in plans] == before


def test_contents_fits_one_page_at_a_readable_size_without_dropping_entries():
    labels = [f"Section {index:02d}: Comparable operational performance, customer activity and service capacity"
              for index in range(1, 19)]
    plans = [PresentationSlide(id=str(index), slide_type="analysis", title="Reported evidence",
                               section_title=label) for index, label in enumerate(labels)]
    deck = _deck()
    _add_planned_contents(deck, plans)
    entries = _contents_entries(deck)
    assert len(deck.slides) == 1
    assert all(label in deck.slides[0].notes_slide.notes_text_frame.text for label in labels)
    assert [shape.name for shape in entries] == [f"contents:entry:{index}" for index in range(1, 19)]
    assert all(shape.text_frame.paragraphs[0].font.size.pt >= 12 for shape in entries)
    assert all(shape.top.inches + shape.height.inches <= deck.slide_height.inches - 1.0 + .001
               for shape in entries)
    assert "Contents (continued)" not in _text(deck.slides[0])


def test_exceptionally_long_contents_label_keeps_full_text_in_notes_on_one_page():
    label = "Operational context: " + "retained evidence and qualified comparisons " * 70 + "FINAL_SECTION_WORD"
    deck = _deck()
    _add_planned_contents(deck, [PresentationSlide(id="long", slide_type="analysis", title="Reported evidence",
                                                  section_title=label)])
    entries = _contents_entries(deck)
    assert len(deck.slides) == len(entries) == 1
    assert entries[0].text.endswith("…")
    assert label in deck.slides[0].notes_slide.notes_text_frame.text
    assert all(shape.name == "contents:entry:1" for shape in entries)
