"""Key figures use selected, compatible source series across document types."""

from dataclasses import replace

import pytest
from pptx import Presentation
from pptx.util import Inches

from adaptive_document_agent.models import (
    ChartPlan, DocumentProfile, Observation, ParsedDocument, PipelineResult,
    PresentationPlan, PresentationSlide, SourceEvidence,
)
from adaptive_document_agent.services.presentation_key_figures import (
    render_key_figures, select_key_figures,
)


def _result():
    observations, charts, slides = [], [], []
    for n in range(4):
        name = f'Reported measure {n + 1}'
        ids = []
        for year, value in ((2022, 10 + n), (2023, 20 + n)):
            identifier = f'measure-{n}-{year}'
            ids.append(identifier)
            observations.append(Observation(
                id=identifier, metric_original=name, value=value, raw_value=str(value),
                unit='count', raw_unit='units', period=f'FY{year}', period_basis='FY',
                period_type='fiscal_year', table_id=f'table-{n}', validation_status='valid',
                evidence=[SourceEvidence(page=n + 2, text=str(value), row_label=name,
                                         extraction_method='digital_table', confidence=.99)],
                confidence=.99,
            ))
        charts.append(ChartPlan(id=f'chart-{n}', title=name, chart_type='line',
                                question='What changed?', observation_ids=ids))
        slides.append(PresentationSlide(id=f'slide-{n}', slide_type='analysis', title=name,
                                        section_title=f'Topic {n + 1}', chart_ids=[f'chart-{n}']))
    return PipelineResult(
        document=ParsedDocument(document_id='generic', sha256='x' * 64,
                                safe_filename='source.pdf', page_count=8),
        profile=DocumentProfile(), observations=observations, charts=charts,
        presentation_plan=PresentationPlan(title='Review', slides=slides),
    )


def test_key_figures_keep_selected_source_values_periods_and_notes():
    result = _result()
    before = result.model_dump()
    figures = select_key_figures(result, result.charts)
    assert len(figures) == 4
    assert figures[0].label == 'Reported measure 1'
    assert figures[0].value.startswith('20')
    assert '+100.0% vs FY2022' in figures[0].comparison
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = render_key_figures(deck, figures)
    visible = ' '.join(shape.text for shape in slide.shapes if shape.has_text_frame)
    assert 'FY2023' in visible
    assert all(o.id in slide.notes_slide.notes_text_frame.text for o in result.observations)
    assert result.model_dump() == before


def test_key_figures_reject_mixed_full_year_and_interim_series():
    result = _result()
    result.observations[1].period = '6M2023'
    result.observations[1].period_basis = '6M'
    result.observations[1].period_type = 'interim_flow'
    assert select_key_figures(result, result.charts) == []


def test_optional_key_figures_never_block_on_an_unreadable_metric_label():
    result = _result()
    for item in result.observations[:2]:
        item.metric_original = 'A complete and unusually long reported measure with several mandatory definition qualifiers ' * 3
    figures = select_key_figures(result, result.charts)
    assert figures == []


@pytest.mark.parametrize("unaudited_index", [0, 1])
def test_visible_current_or_prior_unaudited_period_has_explained_marker(unaudited_index):
    result = _result()
    result.observations[unaudited_index].audited_status = "unaudited"
    before = result.model_dump()
    figures = select_key_figures(result, result.charts)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = render_key_figures(deck, figures)
    visible = '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
    assert f'FY{2022 + unaudited_index}*' in visible
    assert visible.count('* Unaudited') == 1
    assert result.model_dump() == before


@pytest.mark.parametrize("star_location", ["none", "label", "hidden_prior"])
def test_unaudited_note_requires_a_star_in_a_visible_period(star_location):
    result = _result()
    figures = select_key_figures(result, result.charts)
    if star_location == "label":
        figures[0] = replace(figures[0], label=figures[0].label + '*')
    elif star_location == "hidden_prior":
        figures[0] = replace(figures[0], prior_period='FY2020*')
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = render_key_figures(deck, figures)
    visible = '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
    assert '* Unaudited' not in visible


def test_same_selected_source_measure_uses_latest_comparable_view():
    result = _result()
    originals = result.observations[:2]
    interim = []
    for item, year, value in zip(originals, [2023, 2024], [8, 12]):
        interim.append(item.model_copy(deep=True, update={
            'id': f'interim-{year}', 'period': f'6M{year}', 'period_type': 'interim_flow',
            'period_basis': '6M', 'raw_value': str(value), 'value': value, 'audited_status': 'unaudited'}))
    result.observations.extend(interim)
    chart = result.charts[0].model_copy(update={'id': 'recent-view', 'observation_ids': [o.id for o in interim]})
    result.charts.append(chart)
    result.presentation_plan.slides.append(PresentationSlide(id='recent-page', slide_type='analysis',
                                                           title='Recent comparison', chart_ids=[chart.id]))
    before = result.model_dump()
    figures = select_key_figures(result, result.charts)
    assert figures[0].observation_ids == tuple(o.id for o in interim)
    assert '+50.0%' in figures[0].comparison
    assert '6M2023' in figures[0].comparison
    assert len(figures) == 4
    assert result.model_dump() == before
    # A different definition must never replace the selected reported outcome.
    for item in interim:
        item.ifrs_status = 'ADJUSTED'
    assert select_key_figures(result, result.charts)[0].observation_ids == tuple(o.id for o in originals)
