"""One language policy and final display coverage, independent of raw evidence."""
from types import SimpleNamespace
import json

import pytest
from pptx.util import Inches

from adaptive_document_agent.models import DocumentProfile
from adaptive_document_agent.models.customization import ReportRequirement, ReportRequirements
from adaptive_document_agent.services.presentation_language import resolve_language, language_messages, require_consistent_language
from adaptive_document_agent.services.report_language import apply_report_language, audience_copy
from tests.test_report_customization import table_result, presentation


def requirement(language, **kwargs):
    return ReportRequirement(id='language', kind='output_language', request_quote=language,
                             description=language, language=language, **kwargs)


@pytest.mark.parametrize('source,target', [('English','english'),('中文','chinese'),('French','french')])
def test_source_language_default_without_user_customization(source, target):
    profile = DocumentProfile(language=source)
    assert resolve_language(profile).language == target
    assert resolve_language(profile).origin == 'source'
    assert target in language_messages(profile)[0]['content']
    assert profile.report_requirements is None


@pytest.mark.parametrize('language,target', [('中文','chinese'),('English','english'),('French','french')])
def test_user_language_wins_over_source(language, target):
    profile = DocumentProfile(language='Japanese', report_requirements=ReportRequirements(items=[requirement(language)]))
    assert resolve_language(profile).language == target
    assert resolve_language(profile).origin == 'user'


def test_aliases_are_not_conflicts_and_original_means_discovered_source():
    profile = DocumentProfile(language='English', report_requirements=ReportRequirements(items=[
        requirement('en'), requirement('English'), requirement('original')]))
    assert resolve_language(profile).language == 'english'
    profile.report_requirements.items.append(requirement('Chinese', language_scope='source_tables'))
    assert resolve_language(profile).language == 'english'


@pytest.mark.parametrize('items', [[requirement('English'), requirement('Chinese')],
                                  [requirement('Chinese', resolution='ambiguous')]])
def test_conflicting_or_unresolved_requests_cannot_fall_back(items):
    with pytest.raises(ValueError):
        resolve_language(DocumentProfile(language='English', report_requirements=ReportRequirements(items=items)))


def test_page_budget_instruction_does_not_set_the_output_language():
    from adaptive_document_agent.agent.report_requirements import instruction_messages
    profile = DocumentProfile(language='English', report_requirements=ReportRequirements(
        original_request='介绍最多五页', items=[ReportRequirement(id='budget', kind='slide_limit',
            request_quote='介绍最多五页', description='介绍最多五页', max_slides=5)]))
    messages = instruction_messages(profile, purpose='complete introductory synthesis')
    assert 'english' in messages[0]['content']
    assert resolve_language(profile).origin == 'source'


def deck_result(texts):
    result, deck = table_result(), presentation()
    result.profile.language = 'English'
    result.profile.report_requirements = ReportRequirements()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    for text in texts:
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(1)).text = text
    return result, deck


def test_actual_display_requires_complete_review_and_current_target():
    result, deck = deck_result(['公司介绍','Revenue'])
    with pytest.raises(ValueError, match='language verification required'):
        require_consistent_language(deck, result)
    req = result.profile.report_requirements
    req.translation_language = 'english'
    req.copy_translations = {'公司介绍':'Company overview','Revenue':'Revenue'}
    # A prepared map alone is not a successful application.
    with pytest.raises(ValueError, match='coverage incomplete'):
        require_consistent_language(deck, result)
    apply_report_language(deck, result)
    require_consistent_language(deck, result)
    req.items = [requirement('Chinese')]
    with pytest.raises(ValueError, match='language verification required'):
        require_consistent_language(deck, result)


def test_unlocalized_intro_binding_is_not_a_language_exemption():
    result, deck = deck_result(['中文介绍'])
    shape = deck.slides[0].shapes[-1]
    shape.name = 'brief:body'
    shape.element.xpath('.//p:cNvPr')[0].set('descr', 'ADA_SUMMARY_ITEM_V2:' + json.dumps(
        {'page_id':'intro','item_index':0}))
    result.profile.report_requirements.translation_language = 'english'
    result.profile.report_requirements.copy_translations = {'Other':'Other'}
    with pytest.raises(ValueError, match='coverage incomplete'):
        require_consistent_language(deck, result)


def test_preflight_unit_spelling_is_equivalent_but_changed_numbers_are_not():
    result, deck = deck_result(["RMB '000/unitforASP — revenue", 'Revenue 11'])
    req = result.profile.report_requirements
    req.translation_language = 'english'
    req.copy_translations = {'unit':'RMBinthousands/unitforASP — revenue','value':'Revenue 10'}
    with pytest.raises(ValueError, match='Revenue 11'):
        require_consistent_language(deck, result)
    req.copy_translations['value'] = 'Revenue 11'
    require_consistent_language(deck, result)


def test_source_fallback_localizes_without_fabricating_user_requirements(monkeypatch):
    from adaptive_document_agent.agent.report_localization import prepare_report_language, CopyTranslations, TranslationReview
    from adaptive_document_agent.services import pptx_export
    result, deck = deck_result(['公司介绍','Source: Document disclosures (p. 2)'])
    result.profile.report_requirements = None
    monkeypatch.setattr(pptx_export, 'presentation_for_copy', lambda _: deck)
    calls = []
    def generate(messages, model, **kwargs):
        calls.append(model)
        assert any('english' in m['content'] for m in messages)
        if model is CopyTranslations:
            return model(items=[{'id':0,'text':'Company overview'},
                                {'id':1,'text':'Source: Document disclosures (p. 2)'}])
        return model(items=[{'id':i,'accepted':True,'reason':'Faithful target-language prose'} for i in (0,1)])
    prepare_report_language(SimpleNamespace(generate_structured=generate), result)
    assert calls == [CopyTranslations, TranslationReview]
    req = result.profile.report_requirements
    assert req.items == [] and req.original_request == ''
    assert req.translation_language == 'english' and req.translation_language_origin == 'source'
    apply_report_language(deck, result)
    require_consistent_language(deck, result)
    assert 'Source: Document disclosures (p. 2)' in audience_copy(deck)


def test_missing_gateway_cannot_certify_default_language(monkeypatch):
    from adaptive_document_agent.agent.report_localization import prepare_report_language
    result, deck = deck_result(['公司介绍'])
    prepare_report_language(None, result)
    assert not result.profile.report_requirements.copy_translations
    with pytest.raises(ValueError, match='language verification required'):
        require_consistent_language(deck, result)


def test_english_quantity_translation_preserves_coefficient_and_scale():
    from adaptive_document_agent.services.translation_atoms import protect_quantities, restore_quantities
    from adaptive_document_agent.services.report_language import translation_errors
    source = '收入466.6百万美元。'
    text, atoms = protect_quantities(source, 'english')
    translated = restore_quantities('Revenue '+next(iter(atoms)), atoms)
    assert translated == 'Revenue USD 466.6 million'
    assert not translation_errors(source, translated)


def test_physical_speed_is_not_a_monetary_million_suffix():
    from adaptive_document_agent.services.report_language import translation_errors
    assert not translation_errors('速度1米/秒，标准0.25米/秒。', 'Speed 1 m/s; standard 0.25 m/s.')
    assert translation_errors('收入USD 1 million', 'Revenue USD 1 billion')


def test_localized_continuation_uses_reviewed_base_and_marker():
    result, deck = deck_result(['Industry overview (continued)', '01  Market position'])
    req = result.profile.report_requirements
    req.translation_language = 'english'
    req.copy_translations = {'行业概览':'Industry overview', '市场地位':'Market position',
                             '(continued)':'(continued)'}
    apply_report_language(deck, result)
    # The native ordinal requires the actual heading role.
    deck.slides[0].shapes[-1].name = 'brief:heading'
    apply_report_language(deck, result)
    require_consistent_language(deck, result)


def test_translated_sections_reflow_as_complete_contiguous_groups():
    from pptx import Presentation
    from pptx.util import Inches
    from adaptive_document_agent.models.summary_display import SummaryDisplay
    from adaptive_document_agent.services.presentation_brief import BriefItem, render_profile
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(12.5986), Inches(7.0868)
    texts = [
        'The disclosed products support customers across the reported operating settings. ' * 7,
        'The company designs and develops its technology, retaining the stated operating limitations. ' * 7,
        'The platform supports data collection and learning within its disclosed scope. ' * 3,
    ]
    items = [BriefItem(label, text, [i+1], display=SummaryDisplay()) for i, (label, text) in
             enumerate(zip(['Product scope', 'Technology capabilities', 'Platform'], texts))]
    slides = render_profile(deck, 'Industry overview and technology', items)
    assert len(slides) == 2
    bodies = [shape for slide in slides for shape in slide.shapes if shape.name == 'brief:body']
    assert ''.join(shape.text for shape in bodies) == ''.join(texts)
    assert all(shape.top.inches + shape.height.inches <= deck.slide_height.inches - 1.28 + .001
               for shape in bodies)
    assert all(p.font.size.pt == 16 for shape in bodies for p in shape.text_frame.paragraphs)


def test_name_context_is_literal_page_bound_and_invalidates_stale_translation_cache():
    from adaptive_document_agent.services.localization_name_evidence import name_evidence, context_fingerprint
    page = 'Current Example Ltd\n示例有限公司\nRegistered entity.'
    result = SimpleNamespace(
        presentation_plan=SimpleNamespace(company=SimpleNamespace(
            name='示例有限公司', source_pages=[2], field_source_pages={})),
        document=SimpleNamespace(pages=[SimpleNamespace(page_number=1, text='示例有限公司 Unrelated context'),
                                        SimpleNamespace(page_number=2, text=page)]))
    evidence = name_evidence(result)
    assert evidence == [{'page':2, 'literal_name':'示例有限公司', 'quote':page}]
    assert context_fingerprint('english', evidence) != context_fingerprint('chinese', evidence)
    assert context_fingerprint('english', evidence) != context_fingerprint('english', [])
