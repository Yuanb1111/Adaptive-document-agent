"""Source indexes precede closing artwork; narrow fact tables keep readable labels."""
from io import BytesIO
import json

import pytest
from pptx import Presentation

from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import ChartPlan, PresentationSlide, PresentationVisualBlock
from adaptive_document_agent.services.presentation_roles import is_closing_slide
from adaptive_document_agent.services.presentation_body_budget import data_index_slide_numbers, slide_counts
from adaptive_document_agent.services.slide_compositor import render_composed_slide, validate_composed_geometry
from tests.test_p0_composition import observation, result_for
from tests.test_presentation_brief import blank_deck


@pytest.mark.parametrize('index_pages', [0, 1, 3])
def test_native_export_keeps_all_source_pages_before_final_closing(monkeypatch, index_pages):
    from adaptive_document_agent.services import pptx_export
    from adaptive_document_agent.services.presentation_export_trace import record_section
    result = result_for([], [])
    body = result.presentation_plan.slides[3]
    appendix = result.presentation_plan.slides[-1]
    def build(deck, current):
        with record_section(deck, appendix, current):
            for i in range(index_pages):
                slide = pptx_export._base_slide(deck, f'Index {i}', '')
                slide.element.cSld.set('name', 'ada:data-index')
        with record_section(deck, body, current):
            pptx_export._base_slide(deck, 'Evidence body', '')
    monkeypatch.setattr(pptx_export, '_build_planned_presentation', build)
    deck = pptx_export._compose_presentation(result, source_tables=False, validate_plan=False)
    stream = BytesIO()
    deck.save(stream)
    deck = Presentation(stream)
    assert is_closing_slide(deck.slides[-1])
    assert sum(is_closing_slide(s) for s in deck.slides) == 1
    assert data_index_slide_numbers(deck, result) == list(range(2, index_pages + 2))
    assert slide_counts(deck, result) == dict(total=index_pages + 2, body=2, data_index=index_pages)
    for record in result.presentation_export_trace:
        if record['planned_slide_id'] == body.id:
            assert record['slide_numbers'] == [1]
        else:
            assert record['slide_numbers'] == list(range(2, index_pages + 2))
        for n in record['slide_numbers']:
            trace = json.loads(deck.slides[n-1].notes_slide.notes_text_frame.text.split('ADA_EXPORT_TRACE_V1\n')[-1])
            assert trace['physical_slide'] == n


@pytest.mark.parametrize('period', ['FY2023', None])
def test_single_support_fact_uses_full_width_label_without_losing_unknown_period(period):
    points = [observation(f'point-{y}', 'Annual expenditure', 40 + y, f'FY{y}') for y in (2022, 2023)]
    label = 'Total research and development expenditure ratio %'
    fact = observation('support', label, 29.3, period, unit='percent')
    chart = ChartPlan(id='annual', title='Annual expenditure', question='Reported spending', chart_type='line',
                      observation_ids=[o.id for o in points], source_pages=[3])
    plan = PresentationSlide(id='analysis', slide_type='analysis', title='Expenditure', layout='single',
        chart_ids=[chart.id], source_pages=[3], visual_blocks=[
            PresentationVisualBlock(role='hero', chart_ids=[chart.id]),
            PresentationVisualBlock(role='table', observation_ids=[fact.id])])
    result = result_for([*points, fact], [chart], plan)
    before = result.model_dump()
    deck = blank_deck()
    slides = render_composed_slide(deck, plan, [chart], result, DocumentIndex(result.observations))
    assert len(slides) == 1
    shape = next(s for s in slides[0].shapes if s.has_table)
    assert len(shape.table.columns) == 2
    heading = next(s for s in slides[0].shapes if s.has_text_frame and s.text == label)
    assert abs(heading.width - shape.width) <= 1  # Integer EMU rounding.
    assert heading.top + heading.height <= shape.top + 1
    assert shape.table.cell(1, 0).text == (period or 'Unspecified')
    assert shape.table.cell(1, 1).text == '29.3%'
    assert fact.id in slides[0].notes_slide.notes_text_frame.text
    assert result.model_dump() == before
    validate_composed_geometry(deck)
