"""Budget whole introductions before thematic expansion and count actual pages."""
from io import BytesIO

from pptx import Presentation

from adaptive_document_agent.models.customization import ReportRequirement, ReportRequirements
from adaptive_document_agent.services.customization_checks import check_requirements
from tests.test_compact_summary_v111 import fixture, Reviewer


def test_global_synthesis_keeps_configured_reasoning_defaults():
    from adaptive_document_agent.services.llm import LLMSettings
    from adaptive_document_agent.services.llm.reasoning_policy import request_reasoning_policy
    settings = LLMSettings(simple_task_reasoning='reduced')
    for operation in ('SummarySynthesisDraft', 'SummarySynthesisReview'):
        policy = request_reasoning_policy(settings, stage='presentation', operation=operation,
                                          model='deepseek-flash')
        assert policy['intent'] == 'preserve'
        assert policy['options'] == {}


def test_targeted_revision_is_bounded_and_honors_small_task_reasoning():
    from adaptive_document_agent.services.llm import LLMSettings
    from adaptive_document_agent.services.llm.reasoning_policy import request_reasoning_policy
    settings = LLMSettings(provider='deepseek', model='deepseek-flash', simple_task_reasoning='reduced')
    policy = request_reasoning_policy(settings, stage='presentation', operation='SummarySynthesisRevision',
                                      model='deepseek-flash')
    assert policy['output_token_limit'] == 8000
    assert policy['options'] == {'extra_body': {'thinking': {'type': 'disabled'}}}
    configured = settings.model_copy(update={'simple_task_reasoning': 'provider_default'})
    policy = request_reasoning_policy(configured, stage='presentation', operation='SummarySynthesisRevision',
                                      model='deepseek-flash')
    assert policy['options'] == {}


def test_requested_budget_synthesizes_all_parts_in_one_context(monkeypatch):
    from adaptive_document_agent.agent import summary_presentation as module
    result, reading, draft = fixture()
    result.profile.report_requirements = ReportRequirements(items=[ReportRequirement(
        id='length', kind='slide_limit', request_quote='正文最多20页',
        description='正文20页，附录另计', max_slides=20, slide_scope='body')])

    def read(gateway, result, review, **kwargs):
        review.source_blocks = reading.source_blocks
        review.parts = reading.parts
    monkeypatch.setattr(module, 'read_summary', read)
    client = Reviewer(draft)
    company = module.generate_summary_presentation(client, result, [1], object,
                                                   scope_ranges=reading.scope_ranges)
    assert client.calls == ['SummarySynthesisDraft', 'SummarySynthesisReview']
    assert len(company.summary_pages) == 5
    assert company.summary_review.status == 'complete'
    assert {pid for p in company.summary_pages for i in p.items for pid in i.part_ids} == {
        p.id for p in reading.parts}


def test_body_budget_excludes_only_marked_appendices_after_ppt_roundtrip():
    result, _, _ = fixture(1)
    result.profile.report_requirements = ReportRequirements(items=[
        ReportRequirement(id='body', kind='slide_limit', request_quote='Body 2',
                          description='Body limit', max_slides=2, slide_scope='body'),
        ReportRequirement(id='all', kind='slide_limit', request_quote='Total 2',
                          description='Total limit', max_slides=2)])
    deck = Presentation()
    for index in range(5):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        if index > 1:
            slide.element.cSld.set('name', 'ada:appendix')
    stream = BytesIO()
    deck.save(stream)
    check_requirements(result, Presentation(stream), export_verified=True)
    checks = {c.requirement_id:c for c in result.customization_report}
    assert checks['body'].status == 'satisfied'
    assert 'Actual count: 2' in checks['body'].message
    assert checks['all'].status == 'not_met'
    assert 'Actual count: 5' in checks['all'].message


def test_unmarked_pages_cannot_disappear_from_body_budget():
    result, _, _ = fixture(1)
    result.profile.report_requirements = ReportRequirements(items=[ReportRequirement(
        id='body', kind='slide_limit', request_quote='Body 1', description='Body limit',
        max_slides=1, slide_scope='body')])
    deck = Presentation()
    for _ in range(2):
        deck.slides.add_slide(deck.slide_layouts[6])
    check_requirements(result, deck)
    assert result.customization_report[0].status == 'not_met'


def test_supporting_appendix_retains_every_source_part_in_native_slides():
    from adaptive_document_agent.agent.compact_summary import generate_compact_summary, SummarySynthesisDraft
    from adaptive_document_agent.models import PresentationPlan, PresentationSlide
    from adaptive_document_agent.models.presentation import CompanyProfile
    from adaptive_document_agent.services.pptx_export import BUNDLED_TEMPLATE_PATH, _add_company_at_a_glance
    from adaptive_document_agent.services.summary_export import verify_summary_export
    result, review, draft = fixture()
    review.max_intro_slides = 3
    for page in draft['pages'][3:]:
        page['placement'] = 'appendix'
    result.profile.report_requirements = ReportRequirements(items=[ReportRequirement(
        id='body', kind='slide_limit', request_quote='Body 20, appendix separate',
        description='Body limit', max_slides=20, slide_scope='body')])
    company = generate_compact_summary(Reviewer(draft), result, review,
        CompanyProfile(summary_review=review), draft_schema=SummarySynthesisDraft)
    review.status = 'complete'
    result.presentation_plan = PresentationPlan(title='Report', company=company)
    deck = Presentation(str(BUNDLED_TEMPLATE_PATH))
    for identifier in list(deck.slides._sldIdLst):
        deck.part.drop_rel(identifier.rId)
        deck.slides._sldIdLst.remove(identifier)
    slide_plan = PresentationSlide(id='intro', slide_type='company_overview', title='Introduction')
    _add_company_at_a_glance(deck, result, slide_plan)
    assert len(deck.slides) == 3
    import pytest
    with pytest.raises(ValueError, match='does not cover'):
        verify_summary_export(deck, result)
    _add_company_at_a_glance(deck, result, slide_plan, placement='appendix')
    stream = BytesIO()
    deck.save(stream)
    reopened = Presentation(stream)
    assert len(verify_summary_export(reopened, result)) == 70
    assert sum(s.element.cSld.get('name') == 'ada:appendix' for s in reopened.slides) == 2


def test_author_cannot_exempt_content_from_an_all_slides_limit():
    import pytest
    from adaptive_document_agent.agent.compact_summary import generate_compact_summary, SummarySynthesisDraft
    from adaptive_document_agent.models.presentation import CompanyProfile
    result, review, draft = fixture()
    draft['pages'][-1]['placement'] = 'appendix'
    result.profile.report_requirements = ReportRequirements(items=[ReportRequirement(
        id='all', kind='slide_limit', request_quote='20 slides total', description='Total limit', max_slides=20)])
    client = Reviewer(draft)
    with pytest.raises(ValueError, match='explicit body-only'):
        generate_compact_summary(client, result, review, CompanyProfile(summary_review=review),
                                 draft_schema=SummarySynthesisDraft)
    assert client.calls == ['SummarySynthesisDraft', 'SummarySynthesisDraft']


def test_physical_budget_repair_does_not_spend_semantic_content_repair():
    import copy
    from adaptive_document_agent.agent.compact_summary import (
        generate_compact_summary, SummarySynthesisDraft, SummarySynthesisReview)
    from adaptive_document_agent.models.presentation import CompanyProfile
    result, review, initial = fixture()
    review.max_intro_slides = 3
    packed = copy.deepcopy(initial)
    packed['pages'] = [dict(title='Scheduling and delivery', items=[
        item for page in initial['pages'][start:start+2] for item in page['items']])
        for start in range(0,5,2)]
    corrected = copy.deepcopy(packed)
    corrected['pages'][-1]['items'][0]['text'] += ' Delivery depends on this training.'

    class Repair(Reviewer):
        drafts = [initial, packed, corrected]
        authored = 0
        reviews = 0
        def generate_structured(self, messages, schema, **kwargs):
            if schema.__name__ == 'SummarySynthesisRevision':
                return schema(replacements=[dict(item_id='g2_i0',
                    replacement=corrected['pages'][2]['items'][0])])
            if schema is SummarySynthesisDraft:
                self.draft = self.drafts[self.authored]
                self.authored += 1
            else:
                assert schema is SummarySynthesisReview
                self.reject = 'part-69' if self.reviews == 0 else None
                self.reviews += 1
            return super().generate_structured(messages, schema, **kwargs)
    client = Repair(initial)
    company = generate_compact_summary(client,result,review,CompanyProfile(summary_review=review),
                                       draft_schema=SummarySynthesisDraft)
    assert len(company.summary_pages) == 3
    assert client.authored == 2 and client.reviews == 2
    assert len(review.plan_audits) == 2
    assert [a['status'] for a in review.plan_audits[0]['layout_attempts']] == ['overflow', 'success']


def test_source_correction_leaves_one_revision_for_independent_meaning_review():
    import copy
    from adaptive_document_agent.agent.compact_summary import generate_compact_summary, CompactSummaryDraft
    from adaptive_document_agent.models.presentation import CompanyProfile
    result, review, complete = fixture()
    missing = copy.deepcopy(complete)
    missing['pages'][-1]['items'][0]['fact_ids'].pop()
    final = copy.deepcopy(complete)
    final['pages'][-1]['items'][0]['text'] += ' Delivery depends on this training.'
    class Repair(Reviewer):
        authored = 0
        reviewed = 0
        def generate_structured(self, messages, schema, **kwargs):
            if schema is CompactSummaryDraft:
                self.draft = [missing, complete, final][self.authored]
                self.authored += 1
            else:
                self.reject = 'part-69' if self.reviewed == 0 else None
                self.reviewed += 1
            return super().generate_structured(messages, schema, **kwargs)
    client = Repair(missing)
    company = generate_compact_summary(client, result, review, CompanyProfile(summary_review=review))
    assert len(company.summary_pages) == 5
    assert client.authored == 3 and client.reviewed == 2
    assert [a.get('content_failure_phase') for a in review.plan_audits] == ['source', 'meaning', None]


def test_meaning_patch_can_correct_new_evidence_binding_after_prior_source_repair():
    import copy
    from adaptive_document_agent.agent.compact_summary import generate_compact_summary, CompactSummaryDraft
    from adaptive_document_agent.models.presentation import CompanyProfile
    result, review, complete = fixture()
    missing = copy.deepcopy(complete)
    missing['pages'][-1]['items'][0]['fact_ids'].pop()
    unsupported = copy.deepcopy(complete)
    unsupported['pages'][-1]['items'][0]['text'] += ' Revenue was RMB999 million.'
    final = copy.deepcopy(complete)
    final['pages'][-1]['items'][0]['text'] += ' Delivery depends on staff training.'
    class Repair(Reviewer):
        authored = 0
        reviewed = 0
        def generate_structured(self, messages, schema, **kwargs):
            if schema is CompactSummaryDraft:
                self.draft = [missing, complete, unsupported, final][self.authored]
                self.authored += 1
            else:
                self.reject = 'part-69' if self.reviewed == 0 else None
                self.reviewed += 1
            return super().generate_structured(messages, schema, **kwargs)
    client = Repair(missing)
    company = generate_compact_summary(client, result, review, CompanyProfile(summary_review=review))
    assert len(company.summary_pages) == 5
    assert client.authored == 4 and client.reviewed == 2
    assert [a.get('content_failure_phase') for a in review.plan_audits] == ['source', 'meaning', 'source', None]


def test_second_meaning_correction_requires_strictly_fewer_rejected_parts():
    import copy
    import pytest
    from adaptive_document_agent.agent.compact_summary import generate_compact_summary, CompactSummaryDraft
    from adaptive_document_agent.models.presentation import CompanyProfile
    for improves in [True, False]:
        result, review, initial = fixture()
        class Repair(Reviewer):
            authored = 0
            reviewed = 0
            def generate_structured(self, messages, schema, **kwargs):
                if schema is CompactSummaryDraft:
                    self.draft = copy.deepcopy(initial)
                    self.draft['pages'][0]['items'][0]['text'] += ' Training is required.' * self.authored
                    self.authored += 1
                    return super().generate_structured(messages, schema, **kwargs)
                response = super().generate_structured(messages, schema, **kwargs)
                failed = ({'part-0', 'part-1'} if self.reviewed == 0 or not improves else
                          {'part-2'} if self.reviewed == 1 else set())
                for part in response.parts:
                    part.accepted = part.part_id not in failed
                self.reviewed += 1
                return response
        client = Repair(initial)
        if improves:
            generate_compact_summary(client, result, review, CompanyProfile(summary_review=review))
            assert client.authored == client.reviewed == 3
        else:
            with pytest.raises(ValueError, match='unavailable'):
                generate_compact_summary(client, result, review, CompanyProfile(summary_review=review))
            assert client.authored == client.reviewed == 2
