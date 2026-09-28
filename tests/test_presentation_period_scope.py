"""Broad period claims follow visible evidence without inventing new trends."""

import copy
import json

import pytest

from adaptive_document_agent.models import ChartPlan, Observation, PresentationPlan, PresentationSlide, SourceEvidence
from adaptive_document_agent.services.presentation_period_scope import prepare_presentation_period_scope, review_presentation_period_scope
from tests.test_pptx_export import _result


def _item(metric, period, value, *, table="operations", page=10):
    interim = period.startswith("6M")
    return Observation(id=f"{metric}-{period}", metric_original=metric, value=value, raw_value=str(value),
        unit="count", period=period, period_basis="6M" if interim else "FY",
        period_type="interim_flow" if interim else "fiscal_year", table_id=table, confidence=.9,
        evidence=[SourceEvidence(page=page, table_id=table, row_label=metric,
                                 text=str(value), extraction_method="digital_table", confidence=.9)])


def _sample(extra=True):
    result = _result()
    result.validation_warnings = []
    annual = [_item(metric, f"FY{year}", value) for metric in ("Defects", "Repairs")
              for year, value in ((2021, 10), (2022, 20), (2023, 30))]
    interim = [_item(metric, f"6M{year}", value) for metric in ("Defects", "Repairs")
               for year, value in ((2023, 15), (2024, 10))]
    result.observations = annual + (interim if extra else [])
    result.charts = [ChartPlan(id=metric, title=metric, chart_type="bar", question="What changed?",
                              observation_ids=[o.id for o in annual if o.metric_original == metric], source_pages=[10])
                     for metric in ("Defects", "Repairs")]
    result.presentation_plan = PresentationPlan(title="Operations", editorial_status="ready", slides=[
        PresentationSlide(id="movements", slide_type="analysis", title="Defects and repairs both increased over the track record period.",
                          message="How did these measures change during the entire reporting period?",
                          section_title="Operating measures", chart_ids=["Defects", "Repairs"], layout="two_up", source_pages=[10])])
    return result


def test_annual_scope_is_narrowed_and_interim_remains_separate_with_raw_audit():
    result = _sample()
    originals = copy.deepcopy(result.observations), copy.deepcopy(result.charts)
    original_title = result.presentation_plan.slides[0].title
    assert prepare_presentation_period_scope(result) == ["movements"]
    slide = result.presentation_plan.slides[0]
    assert slide.title == "Defects and repairs both increased over FY2021–FY2023."
    assert "6M2023 and 6M2024" in slide.bullets[0]
    assert "see source pages" in slide.bullets[0] and slide.source_pages == [10]
    assert "outside the displayed comparison" in slide.bullets[0]
    assert "decreas" not in slide.bullets[0]  # Never create a new semantic conclusion.
    assert set(slide.bullet_observation_ids[0]) == {o.id for o in result.observations if o.period_basis == "6M"}
    assert (result.observations, result.charts) == originals
    audit = json.loads(result.validation_warnings[-1].message)
    assert audit["original"]["title"] == original_title
    assert len(audit["additional_observation_ids"]) == 4
    snapshot = result.model_dump()
    assert prepare_presentation_period_scope(result) == []
    assert result.model_dump() == snapshot


def test_same_full_series_does_not_narrow_valid_original_copy():
    result = _sample(extra=False)
    original = result.model_dump()
    assert prepare_presentation_period_scope(result) == []
    assert result.model_dump() == original


@pytest.mark.parametrize("period,basis,ptype", [(None, "FY", "fiscal_year"), ("Unknown", "", "generic"),
                                               ("FY2021", "6M", "fiscal_year"), ("FY2021", "FY", "point_in_time"),
                                               ("2021", "generic", "generic")])
def test_unknown_or_conflicting_period_is_not_guessed(period, basis, ptype):
    result = _sample()
    result.observations[0].period = period
    result.observations[0].period_basis = basis
    result.observations[0].period_type = ptype
    title = result.presentation_plan.slides[0].title
    assert prepare_presentation_period_scope(result) == []
    assert result.presentation_plan.slides[0].title == title
    assert result.presentation_plan.editorial_status == "needs_review"
    assert result.validation_warnings[-1].code == "presentation_period_scope_unresolved"


def test_different_chart_periods_do_not_become_one_implied_comparison():
    result = _sample()
    result.charts[1].observation_ids = [o.id for o in result.observations if o.metric_original == "Repairs" and o.period_basis == "6M"]
    assert prepare_presentation_period_scope(result) == ["movements"]
    slide = result.presentation_plan.slides[0]
    assert "over the respective displayed periods" in slide.title
    assert "Defects: FY2021–FY2023" in slide.bullets[0]
    assert "Repairs: 6M2023 and 6M2024" in slide.bullets[0]
    assert result.presentation_plan.editorial_status == "needs_review"
    assert all("FY" in oid for oid in result.charts[0].observation_ids)
    assert all("6M" in oid for oid in result.charts[1].observation_ids)


@pytest.mark.parametrize("difference", ["table", "row", "currency", "unit", "entity", "dimension", "ifrs", "invalid"])
def test_same_page_or_similar_metric_is_not_enough_for_companion_scope(difference):
    result = _sample()
    for item in result.observations[6:]:
        if difference == "table":
            item.table_id = "another-table"
        elif difference == "row":
            item.evidence[0].row_label = "Unrelated row"
        elif difference == "currency":
            item.currency = "USD"
        elif difference == "unit":
            item.unit = "percent"
        elif difference == "entity":
            item.entity = "Other entity"
        elif difference == "dimension":
            item.dimensions = {"product": "other"}
        elif difference == "ifrs":
            item.ifrs_status = "ADJUSTED"
        else:
            item.validation_status = "invalid"
    assert prepare_presentation_period_scope(result) == []
    assert not result.presentation_plan.slides[0].bullets


def test_incomplete_annual_sequence_does_not_invent_intermediate_years():
    result = _sample()
    for chart in result.charts:
        chart.observation_ids = [oid for oid in chart.observation_ids if "2022" not in oid]
    prepare_presentation_period_scope(result)
    assert "over FY2021 and FY2023" in result.presentation_plan.slides[0].title
    assert "FY2022" in result.presentation_plan.slides[0].bullets[0]


def test_reference_limit_does_not_silently_drop_extra_period_evidence():
    result = _sample(extra=False)
    extras = [_item("Defects", f"6M{year}", 10) for year in range(1970, 2020)]
    result.observations.extend(extras)
    assert prepare_presentation_period_scope(result) == []
    assert result.presentation_plan.editorial_status == "needs_review"
    issue = result.validation_warnings[-1]
    assert issue.code == "presentation_period_scope_unresolved"
    assert {o.id for o in extras} <= set(issue.related_ids)
    # Still serializes through the bounded public model, rather than creating
    # a cached result that cannot be loaded on the next run.
    PresentationPlan.model_validate(result.presentation_plan.model_dump())


def test_summary_exact_copy_is_scoped_idempotently_even_after_cached_rebuild():
    result = _sample()
    original = result.presentation_plan.slides[0].title
    summary = PresentationSlide(id="summary", slide_type="executive_summary", title="Summary",
                                bullets=[original, "Similar but different wording over the track record period."])
    result.presentation_plan.slides.insert(0, summary)
    assert prepare_presentation_period_scope(result) == ["movements", "summary"]
    assert summary.bullets[0] == result.presentation_plan.slides[1].title
    assert summary.bullets[1] == "Similar but different wording over the track record period."
    snapshot = result.model_dump()
    assert prepare_presentation_period_scope(result) == []
    assert result.model_dump() == snapshot
    summary.bullets[0] = original
    assert prepare_presentation_period_scope(result) == ["summary"]
    assert result.model_dump() == snapshot


def test_summary_does_not_reuse_a_stale_scope_after_analysis_copy_changed():
    result = _sample()
    original = result.presentation_plan.slides[0].title
    prepare_presentation_period_scope(result)
    result.presentation_plan.slides[0].title = "A different model conclusion."
    summary = PresentationSlide(id="summary", slide_type="executive_summary", title="Summary", bullets=[original])
    result.presentation_plan.slides.append(summary)
    assert prepare_presentation_period_scope(result) == []
    assert summary.bullets == [original]


def test_live_scope_review_is_read_only_and_clears_after_evidence_is_fixed():
    result = _sample(extra=False)
    result.observations[0].period = None
    prepare_presentation_period_scope(result)
    snapshot = result.model_dump()
    assert review_presentation_period_scope(result)[0][0] == "movements"
    assert result.model_dump() == snapshot
    result.observations[0].period = "FY2021"
    assert review_presentation_period_scope(result) == []
    assert any(w.code == "presentation_period_scope_unresolved" for w in result.validation_warnings)


def test_live_scope_review_detects_mixed_panels_and_clears_after_copy_changes():
    result = _sample()
    result.charts[1].observation_ids = [o.id for o in result.observations if o.metric_original == "Repairs" and o.period_basis == "6M"]
    prepare_presentation_period_scope(result)
    assert review_presentation_period_scope(result)
    result.presentation_plan.slides[0].title = "Separately reported measures"
    result.presentation_plan.slides[0].message = "Annual defects and interim repairs are separate comparisons."
    assert review_presentation_period_scope(result) == []


def test_coverage_is_visible_but_does_not_add_interim_values_to_annual_charts():
    from pptx import Presentation
    from adaptive_document_agent.document_model import DocumentIndex
    from adaptive_document_agent.services.pptx_export import BUNDLED_TEMPLATE_PATH
    from adaptive_document_agent.services.slide_compositor import render_composed_slide
    result = _sample()
    prepare_presentation_period_scope(result)
    deck = Presentation(str(BUNDLED_TEMPLATE_PATH))
    slides = render_composed_slide(deck, result.presentation_plan.slides[0], result.charts,
                                   result, DocumentIndex(result.observations))
    visible = "\n".join(shape.text for slide in slides for shape in slide.shapes if shape.has_text_frame)
    assert "FY2021–FY2023" in visible and "6M2023 and 6M2024" in visible
    charts = [shape.chart for slide in slides for shape in slide.shapes if shape.has_chart]
    assert len(charts) == 2
    assert all(list(chart.series[0].values) == [10, 20, 30] for chart in charts)
    notes = slides[0].notes_slide.notes_text_frame.text
    assert "Defects-6M2024" in notes
