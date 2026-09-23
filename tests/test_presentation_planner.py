import io
import pytest

from adaptive_document_agent.agent.presentation_plan_recovery import PresentationPlanRecovery
from adaptive_document_agent.agent.presentation_plan_repairer import PresentationPlanRepairer
from adaptive_document_agent.agent.presentation_planner import PresentationPlanner
from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import (
    AnalysisResult,
    AnalysisTask,
    ChartPlan,
    CompanyFact,
    CompanyProfile,
    DocumentProfile,
    Insight,
    Observation,
    ParsedDocument,
    PipelineResult,
    PresentationPlan,
    PresentationSlide,
    PresentationVisualBlock,
    SourceEvidence,
)
from adaptive_document_agent.services.llm import LLMGateway, LLMSettings, MockLLMClient, ProviderName
from adaptive_document_agent.validation.presentation_plan_validator import PresentationPlanValidator
from tests.test_pptx_export import _result


def test_presentation_plan_validator_rejects_unknown_fact_references() -> None:
    result = _result()
    plan = PresentationPlan(
        title="Invalid plan",
        slides=[
            PresentationSlide(id="cover", slide_type="cover", title="Review"),
            PresentationSlide(id="overview", slide_type="company_overview", title="Company at a Glance"),
            PresentationSlide(id="summary", slide_type="executive_summary", title="Executive Summary"),
            PresentationSlide(
                id="analysis",
                slide_type="analysis",
                title="Unsupported conclusion",
                section_title="Financial Performance",
                message="This message has no retained source.",
                observation_ids=["missing-observation"],
                source_pages=[234],
            ),
            PresentationSlide(id="quality", slide_type="data_quality", title="Data quality"),
            PresentationSlide(id="appendix", slide_type="appendix", title="Source data"),
        ],
    )

    with pytest.raises(ValueError, match="unknown observations"):
        PresentationPlanValidator().validate(plan, result)


def test_presentation_plan_validator_rejects_invented_numeric_claims() -> None:
    result = _result()
    plan = PresentationPlan(
        title="Invalid numeric claim",
        slides=[
            PresentationSlide(id="cover", slide_type="cover", title="Review"),
            PresentationSlide(id="overview", slide_type="company_overview", title="Company at a Glance"),
            PresentationSlide(
                id="summary",
                slide_type="executive_summary",
                title="Revenue increased by 999%",
                insight_ids=["insight-1"],
                source_pages=[234],
            ),
            PresentationSlide(id="quality", slide_type="data_quality", title="Data quality"),
            PresentationSlide(id="appendix", slide_type="appendix", title="Source data"),
        ],
    )

    with pytest.raises(ValueError, match="unsupported numeric claims"):
        PresentationPlanValidator().validate(plan, result)


def test_presentation_plan_validator_accepts_company_excerpt_pages_and_grouped_numbers() -> None:
    result = _result()
    plan = PresentationPlan(
        title="Supported plan",
        company=CompanyProfile(
            name="Example Automation",
            one_line_description="Source-supported company description.",
            source_pages=[8],
        ),
        slides=[
            PresentationSlide(id="cover", slide_type="cover", title="Review"),
            PresentationSlide(
                id="overview",
                slide_type="company_overview",
                title="Company at a Glance",
                observation_ids=["revenue-2023"],
                source_pages=[8],
            ),
            PresentationSlide(
                id="summary",
                slide_type="executive_summary",
                title="Revenue reached 267 025 in the first retained period",
                observation_ids=["revenue-2023"],
                source_pages=[234],
            ),
            PresentationSlide(
                id="analysis",
                slide_type="analysis",
                title="Revenue Performance",
                section_title="Revenue",
                message="Revenue reached 267 025 in the first retained period.",
                chart_ids=["chart-1"],
                source_pages=[234],
            ),
            PresentationSlide(id="quality", slide_type="data_quality", title="Data quality"),
            PresentationSlide(id="appendix", slide_type="appendix", title="Source data"),
        ],
    )

    assert PresentationPlanValidator().validate(plan, result) is plan
    assert PresentationPlanValidator._numbers("659 336 034; 754,257,120; 1’234") == {
        "659336034",
        "754257120",
        "1234",
    }


def test_presentation_planner_repairs_one_semantically_invalid_plan() -> None:
    result = _result()
    invalid = {
        "title": "Review",
        "slides": [
            {"id": "cover", "slide_type": "cover", "title": "Review"},
            {"id": "overview", "slide_type": "company_overview", "title": "Company at a Glance"},
            {
                "id": "summary",
                "slide_type": "executive_summary",
                "title": "Revenue increased by 999%",
                "insight_ids": ["insight-1"],
                "source_pages": [234],
            },
            {"id": "quality", "slide_type": "data_quality", "title": "Data quality"},
            {"id": "appendix", "slide_type": "appendix", "title": "Source data"},
        ],
    }
    repaired = {
        **invalid,
        "slides": [
            *invalid["slides"][:2],
            {
                "id": "summary",
                "slide_type": "executive_summary",
                "title": "Revenue increased across the retained period",
                "insight_ids": ["insight-1"],
                "source_pages": [234],
            },
            *invalid["slides"][3:],
        ],
    }
    client = MockLLMClient([invalid, repaired])
    gateway = LLMGateway(client, LLMSettings(provider=ProviderName.MOCK, model="mock"))

    plan = PresentationPlanner(gateway).plan(result)

    assert plan.slides[2].title == "Revenue increased across the retained period"
    assert len(client.calls) == 2
    assert "failed deterministic validation" in client.calls[1][2]["content"]


def test_presentation_plan_repairer_cleans_unsupported_references_and_numbers() -> None:
    result = _result()
    plan = PresentationPlan(
        title="Dirty AI plan",
        slides=[
            PresentationSlide(id="cover", slide_type="cover", title="Review"),
            PresentationSlide(id="overview", slide_type="company_overview", title="Company at a Glance"),
            PresentationSlide(
                id="summary",
                slide_type="executive_summary",
                title="Revenue was reported. Invented figure 999,999,999.",
                message="Valid statement. Another unsupported 888,888 number.",
                bullets=["Supported bullet", "Invented 777,777 claim"],
                observation_ids=["revenue-2023", "nonexistent-obs-id"],
                chart_ids=["chart-1", "nonexistent-chart-id"],
                insight_ids=["insight-1", "nonexistent-insight-id"],
                source_pages=[234, 999],
            ),
            PresentationSlide(
                id="empty_analysis",
                slide_type="analysis",
                title="Empty",
                observation_ids=["nonexistent-only"],
            ),
            PresentationSlide(
                id="quality",
                slide_type="data_quality",
                title="Data Quality",
                section_title="",
            ),
            PresentationSlide(id="appendix", slide_type="appendix", title="Source Data"),
        ],
    )
    repaired = PresentationPlanRepairer().repair(plan, result)
    assert PresentationPlanValidator().validate(repaired, result) is repaired
    summary_slide = next(s for s in repaired.slides if s.slide_type == "executive_summary")
    assert 999 not in summary_slide.source_pages
    assert "nonexistent-obs-id" not in summary_slide.observation_ids
    assert "nonexistent-chart-id" not in summary_slide.chart_ids
    assert not any(s.id == "empty_analysis" for s in repaired.slides)
    quality_slide = next(s for s in repaired.slides if s.slide_type == "data_quality")
    assert quality_slide.section_title == "Data quality"


def test_presentation_plan_repairer_merges_duplicates_and_limits_charts() -> None:
    result = _result()
    plan = PresentationPlan(
        title="Duplicates and extra charts",
        company=CompanyProfile(
            name="Test Corp",
            products=["Widget A", "Widget B", "widget a"],
            segments=["Hardware", "Software", "hardware"],
            geographies=["Asia", "asia", "Europe"],
            key_facts=[
                CompanyFact(label="CEO", value="Alice", source_pages=[8]),
                CompanyFact(label="ceo", value="Alice", source_pages=[8]),
            ],
            source_pages=[8],
        ),
        slides=[
            PresentationSlide(id="cover", slide_type="cover", title="Review"),
            PresentationSlide(id="overview", slide_type="company_overview", title="Company at a Glance"),
            PresentationSlide(id="summary", slide_type="executive_summary", title="Executive Summary"),
            PresentationSlide(
                id="analysis_1",
                slide_type="analysis",
                title="Analysis 1",
                section_title="Financials",
                message="Evidence analysis.",
                chart_ids=["chart-1", "chart-1"],
                visual_blocks=[
                    PresentationVisualBlock(role="hero", title="Block 1", chart_ids=["chart-1"], chart_type="table"),
                    PresentationVisualBlock(role="supporting", title="Block 2", chart_ids=["chart-1"], chart_type="scatter"),
                ],
                source_pages=[234],
            ),
            PresentationSlide(
                id="risks",
                slide_type="risks",
                title="Risks",
                section_title="Key risks",
                message="Key risk observation.",
                bullets=["Risk A", "risk a"],
                observation_ids=["revenue-2023"],
                source_pages=[234],
            ),
            PresentationSlide(id="quality", slide_type="data_quality", title="Data Quality"),
            PresentationSlide(id="appendix", slide_type="appendix", title="Source Data"),
        ],
    )
    repaired = PresentationPlanRepairer().repair(plan, result)
    assert len(repaired.company.products) == 2
    assert len(repaired.company.segments) == 2
    assert len(repaired.company.geographies) == 2
    assert len(repaired.company.key_facts) == 1
    analysis_slide = next(s for s in repaired.slides if s.id == "analysis_1")
    assert analysis_slide.visual_blocks[0].chart_type is None
    assert len(analysis_slide.chart_ids) <= 3
    risk_slide = next(s for s in repaired.slides if s.slide_type == "risks")
    assert len(risk_slide.bullets) == 1


def test_presentation_plan_recovery_fallback_builds_complete_modern_structure() -> None:
    result = _result()
    fallback_plan = PresentationPlanRecovery().fallback(result)
    slide_types = [s.slide_type for s in fallback_plan.slides]
    assert slide_types[0] == "cover"
    assert slide_types[1] == "company_overview"
    assert slide_types[2] == "executive_summary"
    assert "analysis" in slide_types
    assert slide_types[-2] == "data_quality"
    assert slide_types[-1] == "appendix"
    assert fallback_plan.company.name == ""
    assert fallback_plan.company.identity_state == "UNRESOLVED"


def test_fallback_omits_risk_already_selected_for_summary() -> None:
    result = _result()
    result.insights[0].kind = "risk"
    plan = PresentationPlanRecovery().fallback(result)
    assert plan.planning_origin == "fallback"
    assert all(slide.slide_type != "risks" for slide in plan.slides)
    PresentationPlanValidator().validate(plan, result)


def test_chart_planner_prioritizes_core_metrics_over_peripheral_items() -> None:
    from adaptive_document_agent.agent.chart_planner import ChartPlanner

    obs_core = [
        Observation(
            id=f"rev_{year}",
            metric_original="Revenue",
            metric_canonical="revenue",
            value=float(year * 100),
            raw_value=str(year * 100),
            period=str(year),
            confidence=0.95,
            evidence=[SourceEvidence(page=1, extraction_method="table", confidence=0.95)],
        )
        for year in (2022, 2023, 2024, 2025)
    ]
    obs_peripheral = [
        Observation(
            id=f"tax_{year}",
            metric_original="Prepaid income tax",
            metric_canonical="prepaid_income_tax",
            value=float(year * 2),
            raw_value=str(year * 2),
            period=str(year),
            confidence=0.9,
            evidence=[SourceEvidence(page=2, extraction_method="table", confidence=0.9)],
        )
        for year in (2023, 2024)
    ]
    index = DocumentIndex([*obs_core, *obs_peripheral])
    tasks = [
        AnalysisTask(
            id="task_tax",
            title="Prepaid tax change",
            description="Prepaid tax change",
            expected_output="percentage",
            analysis_type="percentage_change",
            reason="Check tax",
            required_metrics=["prepaid_income_tax"],
            observation_query={"observation_ids": [o.id for o in obs_peripheral]},
        ),
        AnalysisTask(
            id="task_rev",
            title="Revenue change",
            description="Revenue change",
            expected_output="percentage",
            analysis_type="percentage_change",
            reason="Check revenue",
            required_metrics=["revenue"],
            observation_query={"observation_ids": [o.id for o in obs_core]},
        ),
    ]
    results = [
        AnalysisResult(task_id="task_tax", title="Prepaid tax change", result=0.1, evidence=[obs_peripheral[0].evidence[0]]),
        AnalysisResult(task_id="task_rev", title="Revenue change", result=0.2, evidence=[obs_core[0].evidence[0]]),
    ]
    charts = ChartPlanner().plan(tasks, results, index, analysis_focus="revenue trend", maximum=2)
    assert len(charts) == 2
    assert charts[0].analysis_task_id == "task_rev"


def test_extract_presentation_errors_formats_clean_bullets() -> None:
    from adaptive_document_agent.ui.technical import _extract_presentation_errors

    raw_error = (
        "The AI presentation plan could not be retained. Reason: Presentation plan validation failed. "
        "Initial reason: Invalid presentation plan: slide_analysis_cash cites unsupported page 525; "
        "unsupported numeric claim: 1,257,120; chart_type area is unavailable; section_title is missing. "
        "AI repair reason: repair error. Deterministic repair reason: none"
    )
    errors = _extract_presentation_errors(raw_error)
    assert any("slide_analysis_cash cites unsupported page 525" in e for e in errors)
    assert any("unsupported numeric claim: 1,257,120" in e for e in errors)
    assert any("chart_type area is unavailable" in e for e in errors)
    assert any("section_title is missing" in e for e in errors)


def test_presentation_plan_repairer_preserves_usable_charts_when_ai_slides_have_hallucinated_ids() -> None:
    result = _result()
    # Simulate an AI plan with invalid / hallucinated chart and observation IDs
    plan = PresentationPlan(
        title="AI Plan With Hallucinated IDs",
        slides=[
            PresentationSlide(id="cover", slide_type="cover", title="Review"),
            PresentationSlide(id="overview", slide_type="company_overview", title="Company at a Glance"),
            PresentationSlide(id="summary", slide_type="executive_summary", title="Executive Summary"),
            PresentationSlide(
                id="slide_analysis_bad",
                slide_type="analysis",
                title="Revenue Overview",
                section_title="Financial Performance",
                message="Revenue analysis statement.",
                chart_ids=["chart_hallucinated_999"],
                observation_ids=["obs_hallucinated_888"],
            ),
            PresentationSlide(id="quality", slide_type="data_quality", title="Data Quality"),
            PresentationSlide(id="appendix", slide_type="appendix", title="Appendix"),
        ],
    )

    repaired = PresentationPlanRepairer().repair(plan, result)
    assert PresentationPlanValidator().validate(repaired, result) is repaired

    # Ensure the presentation retains analysis slides and charts
    analysis_slides = [s for s in repaired.slides if s.slide_type == "analysis"]
    assert len(analysis_slides) >= 1
    total_charts = [
        cid
        for s in analysis_slides
        for cid in (*s.chart_ids, *(c for b in s.visual_blocks for c in b.chart_ids))
    ]
    assert "chart-1" in total_charts


def test_presentation_plan_repairer_backfills_empty_plan() -> None:
    result = _result()
    plan = PresentationPlan(title="Empty AI Plan", slides=[])

    repaired = PresentationPlanRepairer().repair(plan, result)

    assert PresentationPlanValidator().validate(repaired, result) is repaired
    assert [slide.slide_type for slide in repaired.slides[:3]] == [
        "cover",
        "company_overview",
        "executive_summary",
    ]
    analysis_slides = [slide for slide in repaired.slides if slide.slide_type == "analysis"]
    assert analysis_slides
    assert any("chart-1" in slide.chart_ids for slide in analysis_slides)


def test_presentation_plan_repairer_backfills_after_all_ai_slides_are_filtered() -> None:
    result = _result()
    plan = PresentationPlan(
        title="Filtered AI Plan",
        slides=[
            PresentationSlide(
                id="unsupported_analysis",
                slide_type="analysis",
                title="Hallucinated Analysis",
                chart_ids=["missing-chart"],
                observation_ids=["missing-observation"],
                insight_ids=["missing-insight"],
            )
        ],
    )

    repaired = PresentationPlanRepairer().repair(plan, result)

    assert PresentationPlanValidator().validate(repaired, result) is repaired
    assert all(slide.id != "unsupported_analysis" for slide in repaired.slides)
    analysis_slides = [slide for slide in repaired.slides if slide.slide_type == "analysis"]
    assert analysis_slides
    assert any("chart-1" in slide.chart_ids for slide in analysis_slides)


def test_validator_rejects_plan_without_analysis_slide_when_charts_exist() -> None:
    result = _result()
    plan = PresentationPlan(
        title="No Analysis Plan",
        slides=[
            PresentationSlide(id="cover", slide_type="cover", title="Review"),
            PresentationSlide(id="overview", slide_type="company_overview", title="Company at a Glance", source_pages=[8]),
            PresentationSlide(id="summary", slide_type="executive_summary", title="Executive Summary", source_pages=[234]),
            PresentationSlide(id="quality", slide_type="data_quality", title="Data Quality", source_pages=[8]),
            PresentationSlide(id="appendix", slide_type="appendix", title="Appendix"),
        ],
    )
    with pytest.raises(ValueError, match="at least one analysis slide"):
        PresentationPlanValidator().validate(plan, result)


def test_validator_rejects_unsupported_company_numbers_and_repairer_prunes_them() -> None:
    result = _result()
    bad_plan = PresentationPlan(
        title="Invalid Company Claim",
        company=CompanyProfile(
            name="Example",
            one_line_description="Revenue was 999 trillion in 2023.",
            source_pages=[8],
            key_facts=[
                CompanyFact(label="Revenue", value="267 025", source_pages=[234]),
                CompanyFact(label="Invented", value="999 trillion", source_pages=[8]),
            ],
        ),
        slides=[
            PresentationSlide(id="cover", slide_type="cover", title="Review"),
            PresentationSlide(id="overview", slide_type="company_overview", title="Company at a Glance", source_pages=[8]),
            PresentationSlide(id="summary", slide_type="executive_summary", title="Executive Summary", source_pages=[234]),
            PresentationSlide(
                id="analysis",
                slide_type="analysis",
                title="Revenue Performance",
                section_title="Revenue",
                message="Evidence-backed comparison.",
                chart_ids=["chart-1"],
                source_pages=[234],
            ),
            PresentationSlide(id="quality", slide_type="data_quality", title="Data Quality", source_pages=[8]),
            PresentationSlide(id="appendix", slide_type="appendix", title="Appendix"),
        ],
    )

    with pytest.raises(ValueError, match="unsupported numeric claims"):
        PresentationPlanValidator().validate(bad_plan, result)

    repaired = PresentationPlanRepairer().repair(bad_plan, result)
    # The unsupported sentence "Revenue was 999 trillion in 2023." must be pruned
    assert "999" not in repaired.company.one_line_description
    # The unsupported fact with 999 trillion must be pruned
    fact_values = [f.value for f in repaired.company.key_facts]
    assert not any("999" in v for v in fact_values)
    # The valid fact with 267 025 must be retained
    assert any("267 025" in v for v in fact_values)
    assert PresentationPlanValidator().validate(repaired, result) is repaired


def test_chart_ranking_boosts_deposit_when_in_focus() -> None:
    evidence = SourceEvidence(page=1, text="100", extraction_method="table", confidence=0.9)
    obs_deposit = [
        Observation(id=f"dep_{y}", metric_original="Customer deposit", value=100.0 * y, raw_value=str(100 * y), period=f"202{y}", evidence=[evidence], confidence=0.9)
        for y in (1, 2, 3)
    ]
    obs_tax = [
        Observation(id=f"tax_{y}", metric_original="Prepaid tax", value=50.0 * y, raw_value=str(50 * y), period=f"202{y}", evidence=[evidence], confidence=0.9)
        for y in (1, 2, 3)
    ]
    deposit_chart = ChartPlan(
        id="chart_deposit",
        title="Customer deposit trend",
        question="How did customer deposits change?",
        chart_type="line",
        observation_ids=[o.id for o in obs_deposit],
        source_pages=[1],
        analysis_task_id="task_deposit",
    )
    tax_chart = ChartPlan(
        id="chart_tax",
        title="Prepaid tax trend",
        question="How did prepaid taxes change?",
        chart_type="line",
        observation_ids=[o.id for o in obs_tax],
        source_pages=[1],
        analysis_task_id="task_tax",
    )
    res = PipelineResult(
        document=ParsedDocument(document_id="doc", sha256="abc", safe_filename="doc.pdf", page_count=10),
        profile=DocumentProfile(
            document_type="Banking Report",
            document_purpose="Evaluate customer deposit growth and liquidity.",
            analysis_focus="customer deposit trends",
        ),
        observations=[*obs_deposit, *obs_tax],
        analysis_plan=[
            AnalysisTask(id="task_deposit", title="Deposit analysis", description="Deposit growth", analysis_type="trend", reason="analyze deposit trend", expected_output="growth", priority=1),
            AnalysisTask(id="task_tax", title="Tax analysis", description="Tax analysis", analysis_type="trend", reason="analyze tax trend", expected_output="growth", priority=1),
        ],
        insights=[
            Insight(id="ins_deposit", title="Customer deposit expansion", narrative="Customer deposit expanded rapidly.", kind="interpretation", importance=0.9, confidence=0.9, evidence=[evidence]),
        ],
        charts=[tax_chart, deposit_chart],
    )
    ranked = PresentationPlanRecovery()._rank_charts(res)
    assert ranked[0].id == "chart_deposit"


def test_summary_filters_slope_and_intercept_artifacts() -> None:
    evidence = SourceEvidence(page=1, text="100", extraction_method="table", confidence=0.9)
    res = PipelineResult(
        document=ParsedDocument(document_id="doc", sha256="abc", safe_filename="doc.pdf", page_count=5),
        profile=DocumentProfile(document_type="Report"),
        observations=[
            Observation(id="obs_1", metric_original="Revenue", value=100.0, raw_value="100", period="2023", evidence=[evidence], confidence=0.9),
            Observation(id="obs_2", metric_original="Revenue", value=150.0, raw_value="150", period="2024", evidence=[evidence], confidence=0.9),
        ],
        insights=[
            Insight(id="ins_calc", title="Linear trend slope: 50.0, intercept: 100.0", narrative="slope is 50, turning points: 0", kind="calculated_result", importance=0.99, confidence=0.99, evidence=[evidence]),
            Insight(id="ins_biz", title="Revenue expanded significantly", narrative="Operational performance was solid.", kind="interpretation", importance=0.8, confidence=0.8, evidence=[evidence]),
        ],
    )
    summary = PresentationPlanRecovery()._summary_slide(res)
    for bullet in summary.bullets:
        assert "slope" not in bullet.casefold()
        assert "intercept" not in bullet.casefold()
        assert "turning point" not in bullet.casefold()
    assert any("expanded significantly" in b.casefold() for b in summary.bullets)


def test_appendix_limited_to_one_page_when_no_charts_in_body() -> None:
    from adaptive_document_agent.services.pptx_export import build_presentation
    from pptx import Presentation

    evidence = SourceEvidence(page=1, text="100", extraction_method="table", confidence=0.9)
    res = PipelineResult(
        document=ParsedDocument(document_id="doc", sha256="abc", safe_filename="doc.pdf", page_count=5),
        profile=DocumentProfile(document_type="Report", document_summary="A short report."),
        observations=[
            Observation(id=f"obs_{i}", metric_original=f"Metric {i}", value=10.0 * i, raw_value=str(10 * i), period="2023", evidence=[evidence], confidence=0.9)
            for i in range(25)
        ],
        presentation_plan=PresentationPlan(
            title="Deck Without Charts",
            slides=[
                PresentationSlide(id="cover", slide_type="cover", title="Cover"),
                PresentationSlide(id="overview", slide_type="company_overview", title="Overview", source_pages=[1]),
                PresentationSlide(id="summary", slide_type="executive_summary", title="Summary", source_pages=[1]),
                PresentationSlide(
                    id="analysis",
                    slide_type="analysis",
                    title="Reported Data",
                    section_title="Data",
                    message="Evidence overview.",
                    observation_ids=["obs_1", "obs_2"],
                    source_pages=[1],
                ),
                PresentationSlide(id="quality", slide_type="data_quality", title="Data Quality", source_pages=[1]),
                PresentationSlide(id="appendix", slide_type="appendix", title="Appendix"),
            ],
        ),
    )
    data = build_presentation(res)
    prs = Presentation(io.BytesIO(data))
    # Count appendix slides
    appendix_slides = [
        s for s in prs.slides
        if any("representative retained evidence" in shape.text.casefold() for shape in s.shapes if shape.has_text_frame)
    ]
    # Appendix must be limited to <= 1 page of representative evidence
    assert len(appendix_slides) <= 1
    assert any(w.code == "no_body_charts" for w in res.validation_warnings)
