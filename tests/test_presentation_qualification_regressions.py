"""Regressions for source audit scope, scenario assumptions and stale closings."""
import pytest
from pptx import Presentation
from pptx.util import Inches
from adaptive_document_agent.models import PresentationSlide
from adaptive_document_agent.models.executive_brief import ExecutiveBrief
from adaptive_document_agent.services.presentation_audit_scope import _repair
from adaptive_document_agent.services.executive_brief import brief_items
from adaptive_document_agent.services.presentation_closing import render_closing
from adaptive_document_agent.services.presentation_summary import render_complete_summary
from adaptive_document_agent.agent.coverage_representation import representation_links
from tests.test_topic_coverage_review import _observation
from tests.test_executive_brief import result_for, payload


def deck():
    d = Presentation()
    d.slide_width, d.slide_height = Inches(12.6), Inches(7.087)
    return d


def test_unknown_period_is_never_declared_audited():
    prior = _observation('Output', '6M2023', 90, basis='6M')
    latest = _observation('Output', '6M2024', 70, basis='6M')
    prior.audited_status, latest.audited_status = 'unaudited', 'unknown'
    before = [o.model_dump() for o in [prior, latest]]
    assert _repair('6M2024 is unaudited.', [prior, latest]) == 'Source marks 6M2023 as unaudited.'
    assert _repair('6M2023 is unaudited.', [prior, latest]) == '6M2023 is unaudited.'
    assert _repair('6M2024 is unaudited. Estimates remain uncertain.', [prior, latest]) == 'Source marks 6M2023 as unaudited. Estimates remain uncertain.'
    assert _repair('6M2024 is unaudited.', [latest]) == 'Audit status is not established for the periods in this statement.'
    assert [o.model_dump() for o in [prior, latest]] == before


@pytest.mark.parametrize('different', ['period', 'raw_value', 'unit', 'audited_status', 'ifrs_status', 'category'])
def test_corroborating_view_requires_the_complete_same_fact(different):
    selected = _observation('Output', 'FY2023', 70)
    duplicate = selected.model_copy(deep=True, update={'id': 'other-page'})
    duplicate.evidence[0].page = 300
    assert representation_links([duplicate], [selected]) == {'other-page': [selected.id]}
    if different == 'category':
        duplicate.category_dimensions = {'product': 'other'}
    else:
        setattr(duplicate, different, {'period': '6M2023', 'raw_value': '71', 'unit': 'percent',
                                     'audited_status': 'unaudited', 'ifrs_status': 'ADJUSTED'}[different])
    assert representation_links([duplicate], [selected]) is None


def test_comparison_restores_adjacent_shared_assumption_without_definition_leak():
    definition = 'Response rate is defined as average monthly operating activity.'
    conditions = 'assuming no optional expansion is exercised and assuming a rate of USD19.80 per unit.'
    cases = 'Assuming activity of USD20.6 million, response time is approximately 10.6 months without expansion, 13.9 months with partial expansion and 43.0 months with full expansion.'
    r = result_for([definition, conditions + ' ' + cases])
    data = payload('Response times depend on the activity and expansion assumptions.', label='Response scenarios')
    data['items'][0]['evidence'] = [{'page': 1, 'text': definition}, {'page': 2, 'text': cases}]
    data['items'][0]['comparison_table'] = {'headers': ['Activity', 'No expansion', 'Partial expansion', 'Full expansion'],
        'rows': [['USD20.6 million', '10.6 months', '13.9 months', '43.0 months']]}
    r.executive_brief = ExecutiveBrief.model_validate(data)
    raw = r.model_dump()
    items = brief_items(r)
    assert items[0].conditions == conditions
    slides = render_complete_summary(deck(), 'Summary', items)
    assert len(slides) == 1
    assert conditions in '\n'.join(s.text for s in slides[0].shapes if s.has_text_frame)
    assert r.model_dump() == raw


def test_closing_preserves_current_summary_once_and_keeps_stale_question_in_notes():
    source = 'Reported loss widened while adjusted loss narrowed in the matched period.'
    r = result_for([source])
    r.executive_brief = ExecutiveBrief.model_validate(payload(source, label='Loss comparison'))
    plan = PresentationSlide(id='closing', slide_type='risks', title='Closing',
        bullets=['Check whether adjusted loss follows reported loss.'], source_pages=[1])
    raw = r.model_dump(), plan.model_dump()
    presentation = deck()
    slides = render_complete_summary(presentation, 'Summary', brief_items(r))
    assert render_closing(presentation, r, plan) == []
    visible = '\n'.join(s.text for p in slides for s in p.shapes if s.has_text_frame)
    assert visible.count(source) == 1
    assert plan.bullets[0] not in visible
    assert 'Source' in visible
    assert plan.bullets[0] in slides[-1].notes_slide.notes_text_frame.text
    assert (r.model_dump(), plan.model_dump()) == raw


def test_closing_cannot_use_a_brief_with_invented_values():
    r = result_for(['Output was 70 units.'])
    r.executive_brief = ExecutiveBrief.model_validate(payload('Output was 71 units.'))
    with pytest.raises(ValueError):
        render_closing(deck(), r, PresentationSlide(id='closing', slide_type='risks', title='Closing'))


def test_incomplete_coverage_remains_visible_in_the_delivered_deck(monkeypatch):
    monkeypatch.setattr('adaptive_document_agent.agent.topic_coverage_review.coverage_review_pending', lambda r: True)
    source = 'Observed output increased in the matched period.'
    r = result_for([source])
    r.executive_brief = ExecutiveBrief.model_validate(payload(source))
    slides = render_closing(deck(), r, PresentationSlide(id='closing', slide_type='risks', title='Closing'))
    visible = '\n'.join(s.text for p in slides for s in p.shapes if s.has_text_frame)
    assert 'Coverage review incomplete' in visible
    assert 'Material omissions may remain' in visible
