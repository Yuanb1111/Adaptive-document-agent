"""Generic regressions from the September presentation review; no issuer fixtures."""

import pytest

from adaptive_document_agent.agent.company_introduction import (
    IntroductionDraft, IntroductionPages, ensure_company_introduction,
)
from adaptive_document_agent.models import Observation, PresentationSlide, SourceEvidence
from adaptive_document_agent.document_model.topic_matcher import is_positive_topic_mismatch
from adaptive_document_agent.validation.claim_validator import ClaimValidator, extract_metric_aliases
from test_company_summary_pages import sample


def observations(name, values, unit="days"):
    return [Observation(id=f"o{i}", metric_original=name, value=v, raw_value=str(v),
        period=f"FY{2022+i}", unit=unit, validation_status="valid",
        evidence=[SourceEvidence(page=1, text=str(v), extraction_method="digital_table", confidence=.99)],
        confidence=.99) for i, v in enumerate(values)]


def test_inventory_inflection_does_not_hide_wrong_direction_or_alias_revenue():
    obs = observations("Inventories turnover days", [300, 290, 250])
    slide = PresentationSlide(id="working", slide_type="analysis",
        title="Inventory turnover days increased", observation_ids=[o.id for o in obs])
    issues = ClaimValidator().validate_slide(slide, obs)
    assert any(i.code == "directional_contradiction" for i in issues)
    assert "revenue" not in extract_metric_aliases("Inventories turnover days")


def test_plural_losses_with_rd_are_not_an_unrelated_metric():
    obs = observations("Loss for the year/period", [-100, -130], "currency")
    slide = PresentationSlide(id="losses", slide_type="analysis",
        title="Reported and adjusted losses widened, while R&D expense ratio remained substantial.")
    assert not is_positive_topic_mismatch(obs[0], slide)
    unrelated = observations("Revenue", [100, 130], "currency")[0]
    assert is_positive_topic_mismatch(unrelated, PresentationSlide(
        id="rd", slide_type="analysis", title="Research and development expenses"))


class Gateway:
    def __init__(self, responses):
        self.responses = iter(responses)
        self.calls = []

    def generate_structured(self, messages, schema, **kwargs):
        self.calls.append((messages, schema, kwargs))
        return next(self.responses)


def test_recovery_extracts_two_pages_independently_and_removes_heuristic_fields():
    result = sample()
    original = result.presentation_plan.company.model_copy(deep=True)
    result.presentation_plan.company.summary_overview = None
    result.presentation_plan.company.summary_business = None
    result.presentation_plan.company.customer_types = ["according to a consultant"]
    gateway = Gateway([IntroductionPages(pages=[5,18]), IntroductionDraft(
        overview=original.summary_overview, business=original.summary_business)])
    ensure_company_introduction(gateway, result, result.presentation_plan)
    assert result.presentation_plan.company.summary_overview == original.summary_overview
    assert result.presentation_plan.company.summary_business == original.summary_business
    assert not result.presentation_plan.company.customer_types
    assert all(call[2]["stage"] == "presentation" for call in gateway.calls)
    ensure_company_introduction(gateway, result, result.presentation_plan)
    assert len(gateway.calls) == 2


def test_introduction_rejects_invented_quotation_after_one_repair():
    result = sample()
    original = result.presentation_plan.company.model_copy(deep=True)
    original.summary_business.items[0].source_quote = "An invented product quotation"
    result.presentation_plan.company.summary_overview = None
    result.presentation_plan.company.summary_business = None
    draft = IntroductionDraft(overview=original.summary_overview, business=original.summary_business)
    gateway = Gateway([IntroductionPages(pages=[5,18]), draft, draft])
    with pytest.raises(ValueError, match="could not be verified"):
        ensure_company_introduction(gateway, result, result.presentation_plan)
    assert result.presentation_plan.company.summary_overview is None
    assert len(gateway.calls) == 3


def test_introduction_rejects_unsupplied_selection_without_extracting():
    result = sample()
    result.presentation_plan.company.summary_overview = None
    gateway = Gateway([IntroductionPages(pages=[999])])
    with pytest.raises(ValueError, match="No supported"):
        ensure_company_introduction(gateway, result, result.presentation_plan)


def test_appendix_keeps_unknown_period_cells_separate():
    from pptx import Presentation
    from pptx.util import Inches
    from adaptive_document_agent.services.pptx_export import _add_evidence_table_slides
    result = sample()
    result.observations = observations("Component cost change", [-4.8, -1.5], "percent")
    result.presentation_plan = None
    for i, o in enumerate(result.observations):
        o.period = None
        o.evidence[0].column_label = f"column_{i+2}"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    _add_evidence_table_slides(deck, result, [])
    rows = [cell.text for slide in deck.slides for shape in slide.shapes if shape.has_table
            for row in shape.table.rows for cell in row.cells]
    assert any("column_2; period unspecified" in text for text in rows)
    assert any("column_3; period unspecified" in text for text in rows)
    assert any("4.8" in text for text in rows)
    assert any("1.5" in text for text in rows)
