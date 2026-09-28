"""Audience-facing chart labels retain the dimensions of the plotted measure."""

import pytest
from pptx import Presentation

from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import ChartPlan, PresentationSlide
from adaptive_document_agent.services.presentation_labels import qualified_metric_name, qualify_heading
from adaptive_document_agent.services.pptx_export import BUNDLED_TEMPLATE_PATH, _presentation_chart_title, _series_rows
from adaptive_document_agent.services.single_metric_analysis import single_metric_analysis
from adaptive_document_agent.services.single_metric_renderer import add_single_metric_slide
from adaptive_document_agent.services.slide_compositor import render_composed_slide
from tests.test_single_metric_enrichment import result_for, series


def _deck():
    deck = Presentation(str(BUNDLED_TEMPLATE_PATH))
    for entry in list(deck.slides._sldIdLst):
        deck.part.drop_rel(entry.rId)
        deck.slides._sldIdLst.remove(entry)
    return deck


def _scoped():
    observations = series((16, 24, 23), metric="Gross profit margin", unit="percent")
    for obs in observations:
        obs.category_dimensions = {"channel": "Direct sales"}
        obs.dimensions = {"table_context": "Sales network", "period_basis": "FY", "column_role": "percentage"}
    return observations


@pytest.mark.parametrize(("metric", "dimension", "value"), [
    ("Gross profit margin", "channel", "Direct sales"),
    ("Response time", "region", "Western district"),
    ("Satisfaction", "department", "Engineering"),
])
def test_scope_labels_are_generic_and_preserve_original_records(metric, dimension, value):
    obs = series(metric=metric)[0]
    obs.dimensions = {"table_context": "Original source heading", "column_role": "amount"}
    obs.category_dimensions = {dimension: value}
    before = obs.model_dump()
    assert qualified_metric_name(obs) == f"{metric} ({value})"
    assert obs.model_dump() == before
    assert qualify_heading(f"{value} {metric}", [obs]) == f"{value} {metric}"


def test_multiple_categories_are_not_mislabeled_as_one_category():
    observations = _scoped()
    observations[-1].category_dimensions = {"channel": "Distributors"}
    assert qualify_heading("Gross profit margin", observations) == "Gross profit margin"
    title = _presentation_chart_title("Margin", observations)
    assert "Direct sales" in title and "Distributors" in title
    rows = _series_rows(ChartPlan(id="chart", title="Margin", question="", chart_type="line"), observations)
    assert "Direct sales" in rows[0][1] and "Distributors" in rows[-1][1]


@pytest.mark.parametrize("kind", ["hero", "composed"])
def test_scoped_single_series_is_visible_without_a_legend(kind):
    observations = _scoped()
    before = [o.model_dump() for o in observations]
    chart = ChartPlan(id="chart", title="Gross profit margin", question="", chart_type="line",
                      observation_ids=[o.id for o in observations], source_pages=[3])
    deck = _deck()
    if kind == "hero":
        add_single_metric_slide(deck, chart, single_metric_analysis(observations),
                                title="Gross profit margin", narrative="")
    else:
        plan = PresentationSlide(id="slide", slide_type="analysis", title="Margin review",
                                 chart_ids=[chart.id], source_pages=[3])
        result = result_for(observations)
        result.charts = [chart]
        render_composed_slide(deck, plan, [chart], result, DocumentIndex(observations))
    slide = deck.slides[0]
    text = "\n".join(s.text for s in slide.shapes if s.has_text_frame)
    assert "Gross profit margin (Direct sales)" in text
    assert "Sales network" not in text and "column_role" not in text
    native = next(s.chart for s in slide.shapes if s.has_chart)
    assert not native.has_legend
    assert list(native.series[0].values) == [16, 24, 23]
    assert [o.model_dump() for o in observations] == before
