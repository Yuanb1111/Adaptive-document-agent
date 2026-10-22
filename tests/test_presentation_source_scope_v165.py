"""Units follow plotted columns; unresolved scope stays accurate and visible."""
import io
import json

import pytest
from pptx import Presentation

from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import PresentationSlide, ValidationIssue
from adaptive_document_agent.models.executive_brief import ExecutiveBrief
from adaptive_document_agent.services.composition_candidates import presentation_compositions
from adaptive_document_agent.services.presentation_closing import editorial_closing_items, render_closing
from adaptive_document_agent.services.presentation_labels import composition_heading
from adaptive_document_agent.services.presentation_sparse_pages import fold_sparse_text_pages
from adaptive_document_agent.services.presentation_summary import render_complete_summary
from adaptive_document_agent.services.presentation_brief import BriefItem
from adaptive_document_agent.services.slide_compositor import render_composed_slide
from tests.test_brief_source_checks import sample
from tests.test_executive_brief import payload
from tests.test_p0_composition import result_for
from tests.test_presentation_scope_labels import _deck
from tests.test_source_row_composition import source_matrix


@pytest.mark.parametrize('context', ['RMBinthousands/unitforASP', 'kg per batch', 'USD millions', 'Source table 7'])
def test_raw_table_unit_context_cannot_become_a_percentage_chart_subject(context):
    observations = source_matrix()
    for item in observations:
        item.dimensions['table_context'] = context
    before = [item.model_dump() for item in observations]
    charts = presentation_compositions(DocumentIndex(observations))
    assert all(context not in chart.title for chart in charts)
    chart = next(c for c in charts if c.chart_type == 'stacked_percent')
    # A saved analysis must also be repaired without rebuilding raw facts.
    chart.title = f'{context} — % of Total'
    original = chart.model_dump()
    plan = PresentationSlide(id='mix', slide_type='analysis', title='Reported category mix',
                             chart_ids=[chart.id], source_pages=[3], layout='single')
    deck = _deck()
    render_composed_slide(deck, plan, [chart], result_for(observations, [chart], plan),
                          DocumentIndex(observations), preserve_charts=True)
    visible = '\n'.join(s.text for s in deck.slides[0].shapes if s.has_text_frame)
    assert context not in visible
    assert 'Category share of total' in visible and 'Share (%)' in visible
    native = next(s.chart for s in deck.slides[0].shapes if s.has_chart)
    assert sorted(tuple(s.values) for s in native.series) == [(0.1, 0.1), (0.2, 0.2), (0.3, 0.3), (0.4, 0.4)]
    assert chart.model_dump() == original
    assert [item.model_dump() for item in observations] == before
    assert context in deck.slides[0].notes_slide.notes_text_frame.text


def test_source_column_repair_preserves_snapshot_period_and_authored_qualification():
    observations = source_matrix()
    chart = next(c for c in presentation_compositions(DocumentIndex(observations)) if c.chart_type == 'pie')
    selected = [o for o in observations if o.id in chart.observation_ids]
    legacy = 'Reported sales mix — % of Total (FY2024)'
    chart.title = legacy
    repaired = composition_heading(legacy, chart, selected)
    assert 'Reported sales mix' not in repaired and '(FY2024)' in repaired
    authored = legacy + '; excludes discontinued operations'
    assert composition_heading(authored, chart, selected) == authored


def test_scope_page_does_not_migrate_into_introduction_even_after_roundtrip(monkeypatch):
    monkeypatch.setattr('adaptive_document_agent.agent.topic_coverage_review.coverage_review_pending', lambda r: True)
    result, _ = sample()
    result.executive_brief = ExecutiveBrief.model_validate(payload('Revenue was USD 12 million.'))
    result.validation_warnings.append(ValidationIssue(code='presentation_topic_coverage_unresolved',
        stage='report', severity='warning', message='Retained unresolved coverage record'))
    before = result.model_dump()
    deck = _deck()
    render_complete_summary(deck, 'Introduction', [BriefItem('Company', 'Source-backed company introduction.', [1])])
    slides = render_closing(deck, result, PresentationSlide(id='closing', slide_type='risks', title='Closing'))
    assert len(slides) == 1
    scope = slides[0]
    assert scope.element.cSld.get('name') == 'ada:scope_limits'
    visible = '\n'.join(s.text for s in scope.shapes if s.has_text_frame)
    assert 'Supplementary sources' in visible and 'Unassessed information' in visible and 'p. 12' in visible
    assert 'has not been incorporated' not in visible and 'Review status' not in visible
    notes, _ = json.JSONDecoder().raw_decode(scope.notes_slide.notes_text_frame.text)
    assert notes['source_review_checks'][0]['evidence_quotes']
    assert notes['topic_coverage_checks'][0]['message'] == 'Retained unresolved coverage record'
    stream = io.BytesIO()
    deck.save(stream)
    restored = Presentation(io.BytesIO(stream.getvalue()))
    fold_sparse_text_pages(restored)
    assert len(restored.slides) == 2
    protected = [s for s in restored.slides if s.element.cSld.get('name') == 'ada:scope_limits']
    assert len(protected) == 1
    assert 'Unassessed information' in '\n'.join(s.text for s in protected[0].shapes if s.has_text_frame)
    assert result.model_dump() == before


def test_cited_source_without_pending_review_does_not_manufacture_a_scope_warning():
    result, constraint = sample()
    result.executive_brief = ExecutiveBrief.model_validate(payload(constraint, 12, 'Constraint'))
    assert editorial_closing_items(result) == []
