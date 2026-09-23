"""Source-PDF image selection remains cited, conservative and optional."""

import io
from hashlib import sha256

import pymupdf
import pytest
from PIL import Image
from pptx import Presentation

from adaptive_document_agent.models import DocumentPage
from adaptive_document_agent.services.export import _build_cache_key
from adaptive_document_agent.services.pptx_export import build_presentation
from adaptive_document_agent.services.presentation_source_visual import select_company_source_visual
from tests.test_p1_theme_planning import themed_result


def _document_with_image(*, page_text: str, image_rect=(40, 150, 550, 570)):
    image = io.BytesIO()
    Image.new("RGB", (1200, 800), (54, 96, 145)).save(image, format="JPEG")
    pdf = pymupdf.open()
    for _ in range(3):
        pdf.new_page(width=595, height=842)
    pdf[0].insert_text((40, 80), page_text, fontsize=11)
    pdf[0].insert_image(pymupdf.Rect(image_rect), stream=image.getvalue())
    data = pdf.tobytes()
    pdf.close()
    result = themed_result()
    result.document.sha256 = sha256(data).hexdigest()
    result.document.page_count = 3
    result.document.pages = [DocumentPage(page_number=i, text=page_text if i == 1 else "Other content")
                             for i in range(1, 4)]
    result.presentation_plan.company.name = "ACME DEVICES LIMITED"
    result.presentation_plan.company.identity_state = "RESOLVED"
    result.presentation_plan.company.field_source_pages = {"name": [1], "products": [1]}
    return data, result


def test_source_pdf_image_appears_once_on_company_slide_with_page_citation():
    pdf, result = _document_with_image(page_text=(
        "ACME DEVICES LIMITED supplies industrial devices. "
        "Our products include imaging equipment."
    ))
    selected = select_company_source_visual(pdf, result)
    assert selected is not None and selected.page == 1
    assert _build_cache_key(result, "template") != _build_cache_key(result, "template", source_pdf=pdf)

    deck = Presentation(io.BytesIO(build_presentation(result, source_pdf=pdf)))
    sourced = [(slide, shape) for slide in deck.slides for shape in slide.shapes
               if shape.name.startswith("source_document_image:")]
    assert len(sourced) == 1
    slide, shape = sourced[0]
    assert slide.name == "picture_profile" and shape.name == "source_document_image:p1"
    assert not any(item.name == "picture_profile_continued" for item in deck.slides)
    assert "Image from source document, p. 1" in " ".join(
        item.text for item in slide.shapes if item.has_text_frame
    )
    assert "page 1" in slide.notes_slide.notes_text_frame.text


def test_source_image_skips_scanned_or_tiny_graphics_and_rejects_wrong_pdf():
    pdf, result = _document_with_image(page_text="", image_rect=(40, 150, 80, 190))
    assert select_company_source_visual(pdf, result) is None
    result.document.sha256 = "0" * 64
    with pytest.raises(ValueError, match="does not match"):
        select_company_source_visual(pdf, result)
