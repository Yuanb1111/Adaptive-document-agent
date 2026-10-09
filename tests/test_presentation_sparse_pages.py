"""Short status pages disappear without losing their warning or source record."""
import io
import json

from pptx import Presentation
from pptx.util import Inches

from adaptive_document_agent.services.presentation_sparse_pages import fold_sparse_text_pages
from adaptive_document_agent.services.presentation_summary import render_complete_summary
from adaptive_document_agent.services.presentation_brief import BriefItem
from adaptive_document_agent.services.pptx_export import _render_contents_entries
from tests.test_presentation_brief import blank_deck


def test_short_status_moves_to_existing_page_and_agenda_is_renumbered():
    deck = blank_deck()
    _render_contents_entries(deck, ['Findings', 'Review status', 'Evidence'], '')
    render_complete_summary(deck, 'Findings', [
        BriefItem('Finding', 'The complete source finding explains the comparison and its limitations. ' * 8, [2])])
    render_complete_summary(deck, 'Review status', [
        BriefItem('Coverage review incomplete', 'Material omissions may remain.', [9])], notes='Original review record')
    assert fold_sparse_text_pages(deck) == ['Review status']
    stream = io.BytesIO()
    deck.save(stream)
    restored = Presentation(io.BytesIO(stream.getvalue()))
    assert len(restored.slides) == 2
    entries = [s.text for s in restored.slides[0].shapes if s.name.startswith('contents:entry:')]
    assert entries == ['Findings', 'Evidence']
    visible = '\n'.join(s.text for s in restored.slides[1].shapes if s.has_text_frame)
    assert 'Material omissions may remain.' in visible
    assert 'p. 9' in visible
    assert 'Original review record' in restored.slides[1].notes_slide.notes_text_frame.text
    assert fold_sparse_text_pages(restored) == []


def test_short_copy_with_native_evidence_is_not_removed():
    deck = blank_deck()
    page = render_complete_summary(deck, 'Evidence', [BriefItem('', 'Reported values.', [3])])[0]
    page.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(1))
    assert fold_sparse_text_pages(deck) == []
    assert len(deck.slides) == 1


def test_short_page_keeps_complete_notes_when_no_visible_margin_fits():
    deck = blank_deck()
    page = render_complete_summary(deck, 'Evidence', [BriefItem('', 'Reported values.', [3])])[0]
    page.shapes.add_table(8, 2, Inches(1), Inches(2), Inches(6), Inches(4.6))
    render_complete_summary(deck, 'Limit', [BriefItem('', 'The source omits the denominator.', [7])], notes='Exact source')
    assert fold_sparse_text_pages(deck) == ['Limit']
    assert len(deck.slides) == 1
    assert 'The source omits the denominator.' in page.notes_slide.notes_text_frame.text
    assert 'Exact source' in page.notes_slide.notes_text_frame.text


def test_merge_and_rebase_preserve_structured_notes_after_save():
    from types import SimpleNamespace
    deck = blank_deck()
    render_complete_summary(deck, 'Limit', [BriefItem('', 'The denominator is unspecified.', [7])],
                            notes='Exact source record')
    page = render_complete_summary(deck, 'Evidence', [BriefItem('', 'Reported values.', [3])])[0]
    page.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(1))
    record = dict(topic_id='operations', slide_numbers=[2], observation_ids=['source-1'])
    page.notes_slide.notes_text_frame.text = json.dumps({
        'observations': [{'id': 'source-1', 'raw_value': '12.30'}],
        'ADA_EXPORT_TRACE_V1': {**record, 'physical_slide': 2}})
    result = SimpleNamespace(presentation_export_trace=[record],
                             profile=SimpleNamespace(source_coverage=None))
    assert fold_sparse_text_pages(deck, result) == ['Limit']
    stream = io.BytesIO()
    deck.save(stream)
    restored = Presentation(io.BytesIO(stream.getvalue()))
    notes = json.loads(restored.slides[0].notes_slide.notes_text_frame.text)
    assert notes['observations'] == [{'id': 'source-1', 'raw_value': '12.30'}]
    assert 'Exact source record' in notes['merged_short_pages'][0]['original_notes']
    assert notes['ADA_EXPORT_TRACE_V1']['physical_slide'] == 1
    assert notes['ADA_EXPORT_TRACE_V1']['slide_numbers'] == [1]
    assert notes['ADA_PRIOR_EXPORT_TRACES_V1'][0]['physical_slide'] == 2
