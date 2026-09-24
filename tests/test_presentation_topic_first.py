"""A full series inventory selects questions before the chart quota is applied."""

import json
import io

import pytest
from pptx import Presentation

from adaptive_document_agent.agent.chart_planner import ChartPlanner
from adaptive_document_agent.agent.presentation_plan_recovery import PresentationPlanRecovery
from adaptive_document_agent.agent.presentation_planner import PresentationPlanner
from adaptive_document_agent.agent.presentation_topic_selector import PresentationTopicSelector, series_directory
from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import PresentationPlan, PresentationSlide, PresentationTopic, PresentationTopicSelection
from adaptive_document_agent.models.page import DocumentPage
from adaptive_document_agent.services.pptx_export import build_presentation
from adaptive_document_agent.validation.presentation_plan_validator import PresentationPlanValidator
from tests.test_pptx_export import _result


def test_topic_selector_receives_complete_series_directory_and_exact_ids() -> None:
    result = _result()
    directory, _ = series_directory(result)
    selected_id = directory[0]["id"]

    class Gateway:
        def generate_structured(self, messages, response_model, **kwargs):
            self.messages = messages
            return response_model.model_validate({
                "topics": [{"id": "growth", "title": "Revenue movement",
                            "question": "How did revenue move?",
                            "rationale": "The source reports comparable annual revenue.",
                            "series_ids": [selected_id]}],
            })

    gateway = Gateway()
    selection = PresentationTopicSelector(gateway).select(result)
    payload = json.loads(gateway.messages[-1]["content"].split("\n", 1)[1].rsplit("\n", 1)[0])
    assert {row[0] for row in payload["all_extracted_series"]} == {item["id"] for item in directory}
    assert selection.topics[0].series_ids == [selected_id]


def test_topic_selector_keeps_background_excerpts_outside_primary_analysis(monkeypatch) -> None:
    from adaptive_document_agent.services.company_discovery import CompanyProfileDiscovery

    result = _result()
    result.document.pages = [
        DocumentPage(page_number=1, text="Summary: products and applications"),
        DocumentPage(page_number=234, text="Financial Information: reported revenue"),
    ]
    monkeypatch.setattr(
        CompanyProfileDiscovery, "rank_profile_pages",
        staticmethod(lambda *args, **kwargs: [1, 234]),
    )
    selected_id = series_directory(result)[0][0]["id"]

    class Gateway:
        def generate_structured(self, messages, response_model, **kwargs):
            self.messages = messages
            return response_model(topics=[PresentationTopic(
                id="growth", title="Revenue movement", question="How did revenue move?",
                rationale="Comparable annual revenue is available.", series_ids=[selected_id],
            )])

    gateway = Gateway()
    PresentationTopicSelector(gateway).select(result)
    payload = json.loads(gateway.messages[-1]["content"].split("\n", 1)[1].rsplit("\n", 1)[0])
    assert payload["background_page_excerpts"] == [
        {"page": 1, "text": "Summary: products and applications"}
    ]


def test_topic_selector_rejects_unknown_evidence_id() -> None:
    with pytest.raises(ValueError, match="unknown"):
        PresentationTopicSelector._validate(
            PresentationTopicSelection(topics=[PresentationTopic(
                id="bad", title="Unknown", question="What changed?",
                rationale="Candidate topic", series_ids=["made_up"])]),
            {},
        )


def test_topic_takeaway_cannot_add_unsupported_numbers() -> None:
    result = _result()
    directory, lookup = series_directory(result)
    selection = PresentationTopicSelection(topics=[PresentationTopic(
        id="growth", title="Revenue movement", question="How did revenue move?",
        rationale="Reported annual values exist.", takeaway="Revenue grew by 9999%",
        series_ids=[directory[0]["id"]],
    )])
    with pytest.raises(ValueError, match="unsupported numeric claims"):
        PresentationTopicSelector._validate(selection, lookup)


def test_topic_selector_rejects_background_metric_as_primary_evidence() -> None:
    result = _result()
    outside = [item.model_copy(deep=True) for item in result.observations]
    for item in outside:
        for evidence in item.evidence:
            evidence.page = 1
    selection = PresentationTopicSelection(topics=[PresentationTopic(
        id="background", title="Business background", question="What changed?",
        rationale="This is background evidence.", series_ids=["background-series"],
    )])
    with pytest.raises(ValueError, match="outside the primary analysis scope"):
        PresentationTopicSelector._validate(
            selection, {"background-series": outside}, primary_pages={234},
        )


def test_topic_selector_repairs_unsupported_numeric_claim_once() -> None:
    result = _result()
    directory, _ = series_directory(result)
    series_id = directory[0]["id"]

    class Gateway:
        calls = 0

        def generate_structured(self, messages, response_model, **kwargs):
            self.calls += 1
            return response_model(topics=[PresentationTopic(
                id="growth", title="Revenue reached 999", question="How did revenue move?",
                rationale="Comparable annual revenue is available.", series_ids=[series_id],
            )]) if self.calls == 1 else response_model(topics=[PresentationTopic(
                id="growth", title="Revenue movement", question="How did revenue move?",
                rationale="Comparable annual revenue is available.", series_ids=[series_id],
            )])

    gateway = Gateway()
    assert PresentationTopicSelector(gateway).select(result).topics[0].title == "Revenue movement"
    assert gateway.calls == 2


def test_requested_topic_series_are_not_lost_to_generic_chart_limit() -> None:
    result = _result()
    revenue = result.observations
    orders = [item.model_copy(deep=True) for item in revenue]
    for item in orders:
        item.id = item.id.replace("revenue", "orders")
        item.metric_original = "Orders"
        item.metric_canonical = "orders"
        item.unit = "count"
        item.raw_unit = "orders"
        item.currency = None
    charts = ChartPlanner().plan(
        [], [], DocumentIndex([*revenue, *orders]), maximum=1,
        requested_series=[revenue, orders],
    )
    assert len(charts) >= 2
    assert {identifier for chart in charts[:2] for identifier in chart.observation_ids} == {
        item.id for item in [*revenue, *orders]
    }
    only_selected = ChartPlanner().plan(
        [], [], DocumentIndex([*revenue, *orders]), maximum=10,
        requested_series=[orders], only_requested=True,
    )
    assert {identifier for chart in only_selected for identifier in chart.observation_ids} == {
        item.id for item in orders
    }


def test_failed_slide_writing_can_retain_selected_question() -> None:
    result = _result()
    directory, lookup = series_directory(result)
    selected_id = directory[0]["id"]
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(
        id="growth", title="Revenue movement", question="How did revenue move?",
        rationale="The source reports a comparable annual revenue series.",
        takeaway="Revenue increased across the reported years",
        series_ids=[selected_id],
    )])
    result.charts = ChartPlanner().plan(
        [], [], DocumentIndex(result.observations), maximum=1,
        requested_series=[lookup[selected_id]],
    )
    plan = PresentationPlanRecovery().from_selected_topics(result)
    assert plan.planning_origin == "topic_recovery"
    assert plan.themes[0].id == "growth"
    assert plan.slides[3].chart_ids
    assert plan.slides[3].analytical_question == "How did revenue move?"
    assert plan.slides[3].title == "Revenue increased across the reported years"
    PresentationPlanValidator().validate(plan, result)


def test_topic_recovery_uses_neutral_title_when_endpoint_is_not_on_slide() -> None:
    result = _result()
    prototype = result.observations[0]
    result.observations = []
    for year in range(2021, 2036):
        item = prototype.model_copy(deep=True)
        item.id = f"revenue-{year}"
        item.period = f"FY{year}"
        item.value = float(year * 1000)
        item.raw_value = str(year * 1000)
        result.observations.append(item)
    series_id = series_directory(result)[0][0]["id"]
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(
        id="growth", title="Revenue movement", question="How did revenue move?",
        rationale="The reported series spans several years.",
        takeaway="Revenue peaked in 2033", series_ids=[series_id],
    )])
    chart = result.charts[0]
    chart.observation_ids = ["revenue-2034", "revenue-2035"]

    plan = PresentationPlanRecovery().from_selected_topics(result)
    slide = next(item for item in plan.slides if item.theme_id == "growth")

    assert slide.title == "Revenue movement"
    assert "revenue-2033" not in slide.observation_ids
    PresentationPlanValidator().validate(plan, result)


def test_fallback_distinguishes_same_metric_categories_in_visual_blocks() -> None:
    result = _result()
    originals = result.observations
    result.observations = []
    result.charts = []
    for category, suffix in (("Inventory", "stock"), ("Receivables", "debtors")):
        group = [item.model_copy(deep=True) for item in originals]
        for item in group:
            item.id = f"{suffix}-{item.period}"
            item.metric_original = "Current assets"
            item.metric_canonical = "Current assets"
            item.dimensions["category"] = category
        result.observations.extend(group)
        chart = _result().charts[0].model_copy(deep=True)
        chart.id = f"chart-{suffix}"
        chart.title = "Current assets — Reported Values"
        chart.observation_ids = [item.id for item in group]
        result.charts.append(chart)

    plan = PresentationPlanRecovery().fallback(result)
    titles = [block.title for slide in plan.slides if slide.slide_type == "analysis" for block in slide.visual_blocks]

    assert "Current assets: Inventory" in titles
    assert "Current assets: Receivables" in titles


def test_recovered_question_shows_third_series_and_keeps_detail_in_notes() -> None:
    result = _result()
    original = result.observations
    for metric in ("Orders", "Returns"):
        copies = [item.model_copy(deep=True) for item in original]
        for item in copies:
            item.id = item.id.replace("revenue", metric.casefold())
            item.metric_original = metric
            item.metric_canonical = metric.casefold()
        result.observations.extend(copies)
    directory, lookup = series_directory(result)
    selected = [entry["id"] for entry in directory]
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(
        id="movement", title="Reported movements", question="How did the three measures move?",
        rationale="The three comparable series describe related reported movements.",
        series_ids=selected,
    )])
    result.charts = ChartPlanner().plan([], [], DocumentIndex(result.observations),
        requested_series=[lookup[sid] for sid in selected], only_requested=True)
    result.presentation_plan = PresentationPlanRecovery().from_selected_topics(result)
    slide = next(item for item in result.presentation_plan.slides if item.theme_id == "movement")
    assert len(slide.chart_ids) == 2
    support = [block for block in slide.visual_blocks if block.role == "table"]
    assert len(support) == 1
    assert set(support[0].observation_ids) == set(slide.observation_ids)
    uncharted = set(slide.observation_ids)
    assert set(support[0].observation_ids) <= uncharted

    deck = Presentation(io.BytesIO(build_presentation(result)))
    analysis = next(page for page in deck.slides if page.name.startswith("composed_"))
    assert any(shape.has_table for shape in analysis.shapes)
    assert all(identifier in analysis.notes_slide.notes_text_frame.text for identifier in uncharted)
    all_text = [" ".join(shape.text for shape in page.shapes if shape.has_text_frame) for page in deck.slides]
    assert not any("Source Data Appendix" in content for content in all_text)
    assert "Appendix" not in all_text[1]


def test_sourced_company_identity_cannot_be_called_unnamed() -> None:
    result = _result()
    plan = PresentationPlanRecovery().fallback(result)
    plan.company.name = "DOBOT"
    plan.company.source_pages = [1]
    plan.company.one_line_description = "Unnamed issuer with robotics products"
    with pytest.raises(ValueError, match="contradicts the sourced company identity"):
        PresentationPlanValidator().validate(plan, result)


def test_recovery_removes_source_footnote_marker_from_audience_label() -> None:
    result = _result()
    for item in result.observations:
        item.metric_original = "Average distributor value(1) (RMB in thousands)"
        item.metric_canonical = item.metric_original
    result.charts = ChartPlanner().plan(
        [], [], DocumentIndex(result.observations),
        requested_series=[result.observations], only_requested=True,
    )
    plan = PresentationPlanRecovery().fallback(result)
    analysis = next(slide for slide in plan.slides if slide.slide_type == "analysis")
    assert "(1)" not in analysis.section_title
    assert all("(1)" not in block.title for block in analysis.visual_blocks)


def test_selected_topics_avoid_a_second_full_slide_plan_request() -> None:
    result = _result()
    selected_id = series_directory(result)[0][0]["id"]
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(
        id="growth", title="Revenue movement", question="How did revenue move?",
        rationale="Comparable annual evidence", series_ids=[selected_id],
    )])

    class Gateway:
        calls = 0

        def generate_structured(self, messages, response_model, **kwargs):
            self.calls += 1
            return response_model(title="Invalid empty slide draft")

    gateway = Gateway()
    try:
        PresentationPlanner(gateway).plan(result)
    except ValueError:
        pass  # The orchestrator retains the model-selected questions.
    assert gateway.calls == 1


def test_thematic_appendix_keeps_selected_evidence_without_unrelated_conflicts() -> None:
    result = _result()
    directory, lookup = series_directory(result)
    selected_id = directory[0]["id"]
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(
        id="growth", title="Revenue movement", question="How did revenue move?",
        rationale="Comparable annual revenue is available.", series_ids=[selected_id],
    )])
    result.charts = ChartPlanner().plan(
        [], [], DocumentIndex(result.observations), requested_series=[lookup[selected_id]],
        only_requested=True,
    )
    result.presentation_plan = PresentationPlanRecovery().from_selected_topics(result)
    conflicting = [item.model_copy(deep=True) for item in result.observations[:2]]
    for ordinal, item in enumerate(conflicting):
        item.id = f"unrelated-{ordinal}"
        item.metric_original = "Others"
        item.metric_canonical = "Others"
        item.period = "FY2023"
        item.value = float(ordinal + 1)
        item.raw_value = str(ordinal + 1)
    result.observations.extend(conflicting)
    deck = Presentation(io.BytesIO(build_presentation(result)))
    tables = [shape.table for slide in deck.slides for shape in slide.shapes if shape.has_table]
    assert tables
    table_text = " ".join(cell.text for table in tables for row in table.rows for cell in row.cells)
    assert "Others" not in table_text
    assert all(item.period.replace("FY", "") in table_text for item in lookup[selected_id])


def test_conclusion_precedes_data_index_and_final_thank_you() -> None:
    result = _result()
    result.presentation_plan = PresentationPlan(title="Revenue review", slides=[
        PresentationSlide(id="cover", slide_type="cover", title="Revenue review"),
        PresentationSlide(id="company", slide_type="company_overview", title="Document at a Glance"),
        PresentationSlide(id="summary", slide_type="executive_summary", title="Summary",
                          insight_ids=["insight-1"], source_pages=[234]),
        PresentationSlide(id="analysis", slide_type="analysis", title="Revenue movement",
                          section_title="Performance", message="How did revenue move?",
                          chart_ids=["chart-1"], source_pages=[234]),
        PresentationSlide(id="close", slide_type="risks", title="Conclusions and Watch Items",
                          section_title="Conclusions", message="Indicators to monitor.",
                          bullets=["Monitor whether the revenue trend continues."],
                          insight_ids=["insight-1"], source_pages=[234]),
        PresentationSlide(id="quality", slide_type="data_quality", title="Data quality"),
        PresentationSlide(id="appendix", slide_type="appendix", title="Source data"),
    ])
    deck = Presentation(io.BytesIO(build_presentation(result)))
    last_text = " ".join(shape.text for shape in deck.slides[-1].shapes if shape.has_text_frame)
    texts = [" ".join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in deck.slides]
    closing = next(i for i, text in enumerate(texts) if "Monitor whether the revenue trend continues" in text)
    appendix = [i for i, slide in enumerate(deck.slides) if any(shape.has_table for shape in slide.shapes)]
    assert "Conclusions and Watch Items" in texts[closing]
    assert "Watch item 1" not in texts[closing]
    assert appendix and all(closing < i < len(texts) - 1 for i in appendix)
    assert "Data Index" in texts[appendix[0]]
    assert "THANK YOU" in last_text


def test_closing_reserves_space_for_lower_ranked_watch_items():
    result = _result()
    template = result.insights[0]
    result.insights = [template.model_copy(update={
        "id": f"finding-{i}", "title": label, "importance": 1 - i / 10,
        "implication": label + " remains relevant to the reported performance.",
        "watch_item": "Monitor the pace of reported changes." if i == 3 else None,
    }) for i, label in enumerate(("Demand", "Costs", "Capacity", "Liquidity"))]
    closing = PresentationPlanRecovery._risks_slide(
        result, PresentationSlide(id="summary", slide_type="executive_summary", title="Summary"),
    )
    assert "Monitor the pace of reported changes." in closing.bullets
    from adaptive_document_agent.services.presentation_closing import render_closing
    deck = Presentation()
    render_closing(deck, result, closing)
    text = " ".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert "Watch items" in text
    assert all(bullet in text for bullet in closing.bullets)


def test_topic_recovery_selects_composition_across_category_series():
    from tests.test_p0_composition import matrix, result_for
    observations, _ = matrix()
    result = result_for(observations, [])
    directory, lookup = series_directory(result)
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(
        id="mix", title="Category mix", question="How did the category mix change?",
        rationale="The source reports a complete category composition.",
        series_ids=[entry["id"] for entry in directory],
    )])
    result.charts = ChartPlanner().plan([], [], DocumentIndex(observations),
        requested_series=list(lookup.values()), only_requested=True)
    plan = PresentationPlanRecovery().from_selected_topics(result)
    selected = {cid for slide in plan.slides if slide.slide_type == "analysis" for cid in slide.chart_ids}
    assert any(chart.id in selected and chart.chart_type == "stacked_percent" for chart in result.charts)
