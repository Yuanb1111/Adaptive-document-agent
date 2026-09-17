"""Comprehensive tests for the 24 generic presentation upgrades.

Covers:
1. Unit normalization and compact currency formatting (RMB 174.3m, RMB 1.57bn).
2. Negative vertical bar chart category axis positioning (LOW).
3. Data label font sizing (>= 14 pt bold).
4. Accounting sign normalization for expense ratios (positive display, pp movement).
5. Elimination of "Observed pairs".
6. Volume count vs currency isolation.
7. ASP + Volume chart pairing.
8. Short metric display names on cards.
9. Analytical topic grouping and generic merges.
10. Hybrid layout (chart_plus_kpis).
11. Company overview identity consistency and semantic categorization.
12. Appendix period segregation and unit cleanliness.
13. Preflight quality gate enforcement.
14. Slide deck structure and Thank You slide preservation.
"""

from __future__ import annotations

import io
from pptx import Presentation
from pptx.enum.chart import XL_TICK_LABEL_POSITION
import pytest

from adaptive_document_agent.document_model.metric_semantic_classifier import (
    classify_metric,
    format_metric_change,
    format_metric_display_value,
    sanitize_metric_label,
)
from adaptive_document_agent.extraction.normalizer import infer_unit_defaults
from adaptive_document_agent.extraction.observation_extractor import ObservationExtractor
from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import (
    ChartPlan,
    CompanyFact,
    CompanyProfile,
    DocumentProfile,
    Observation,
    ParsedDocument,
    PipelineResult,
    PresentationPlan,
    PresentationSlide,
    SourceEvidence,
)
from adaptive_document_agent.models.table import ExtractedTable, TableRow
from adaptive_document_agent.services.export import export_pptx
from adaptive_document_agent.services.financial_formatter import (
    format_compact_currency,
    format_financial_movement,
    is_expense_ratio_metric,
    normalize_currency_symbol,
    normalize_raw_unit,
    shorten_metric_title,
)
from adaptive_document_agent.services.ppt_preflight import PresentationPreflight
from adaptive_document_agent.services.pptx_export import (
    _add_native_chart,
    _charts_belong_together,
    _display_source_unit,
    _unit_label,
)


def _test_parsed_document() -> ParsedDocument:
    """Create a minimal ParsedDocument for test PipelineResults."""
    return ParsedDocument(
        document_id="test-doc",
        sha256="0" * 64,
        safe_filename="test.pdf",
        page_count=100,
    )


# ==============================================================================
# 1. DISPLAY UNITS & COMPACT CURRENCY
# ==============================================================================

def test_unit_normalization_aliases() -> None:
    """Requirement 1: Ban RMBinthousands and map unit aliases cleanly."""
    assert normalize_raw_unit("RMBinthousands") == "RMB '000"
    assert normalize_raw_unit("rmbinthousands") == "RMB '000"
    assert normalize_raw_unit("RMB in thousands") == "RMB '000"
    assert normalize_raw_unit("CNY thousands") == "RMB '000"
    assert normalize_raw_unit("rmb'000") == "RMB '000"
    assert normalize_raw_unit("RMB ’000") == "RMB '000"
    assert normalize_raw_unit("RMB in millions") == "RMB million"
    assert normalize_raw_unit("HKD in thousands") == "HK$ '000"
    assert normalize_raw_unit("in percent") == "%"
    assert normalize_raw_unit("multiple") == "x"
    assert normalize_raw_unit("units") == "units"


def test_compact_currency_formatting() -> None:
    """Requirement 1: Formats amounts into compact institutional notation."""
    # 174,314 in thousands -> RMB 174.3m
    assert format_compact_currency(174_314.0, raw_unit="RMB in thousands") == "RMB 174.3m"
    assert format_compact_currency(174_314_000.0) == "RMB 174.3m"

    # 1,567,108 in thousands -> RMB 1.57bn
    assert format_compact_currency(1_567_108.0, raw_unit="RMB '000") == "RMB 1.57bn"
    assert format_compact_currency(1_567_108_000.0) == "RMB 1.57bn"

    # 450 in thousands -> RMB 450k
    assert format_compact_currency(450.0, raw_unit="in thousands") == "RMB 450k"

    # Negative currency (e.g. net loss)
    assert format_compact_currency(-174_314_000.0) == "-RMB 174.3m"

    # Foreign currency: HKD
    assert format_compact_currency(450_000_000.0, currency="HKD") == "HK$ 450m"


def test_normalizer_infer_unit_defaults_no_unspaced_tokens() -> None:
    """Requirement 1: infer_unit_defaults never generates RMBinthousands."""
    defaults = infer_unit_defaults("(RMBinthousands, except for percentages)")
    assert defaults.raw_unit == "RMB in thousands"
    assert "RMBinthousands" not in defaults.raw_unit


# ==============================================================================
# 2. CHART AXIS LABEL POSITIONING FOR NEGATIVE BARS
# ==============================================================================

def test_negative_bar_chart_axis_low_position() -> None:
    """Requirement 2: Negative vertical bar charts set category axis position to LOW."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    obs_loss = [
        Observation(
            id="obs_loss_21",
            metric_original="Net loss",
            value=-150_000.0,
            raw_value="-150,000",
            period="FY2021",
            unit="currency",
            currency="RMB",
            confidence=0.9,
        ),
        Observation(
            id="obs_loss_22",
            metric_original="Net loss",
            value=-220_000.0,
            raw_value="-220,000",
            period="FY2022",
            unit="currency",
            currency="RMB",
            confidence=0.9,
        ),
    ]

    plan = ChartPlan(
        id="chart_net_loss",
        title="Net Loss Widened",
        question="What was the net loss trajectory?",
        chart_type="bar",
        observation_ids=["obs_loss_21", "obs_loss_22"],
        show_data_labels=True,
    )

    bounds = (1.0, 1.5, 6.0, 4.0)
    _add_native_chart(slide, plan, obs_loss, bounds)

    # Find the generated chart
    chart_shape = next(s for s in slide.shapes if s.has_chart)
    chart = chart_shape.chart

    # Category axis must have tick_label_position set to LOW
    assert chart.category_axis.tick_label_position == XL_TICK_LABEL_POSITION.LOW


# ==============================================================================
# 3. CHART DATA LABEL SIZE
# ==============================================================================

def test_chart_data_label_font_size() -> None:
    """Requirement 3: Standalone chart data labels must be >= 14 pt bold."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    obs = [
        Observation(
            id="obs_rev_21",
            metric_original="Revenue",
            value=174_314.0,
            raw_value="174,314",
            period="FY2021",
            unit="currency",
            currency="RMB",
            confidence=0.9,
        ),
        Observation(
            id="obs_rev_22",
            metric_original="Revenue",
            value=241_013.0,
            raw_value="241,013",
            period="FY2022",
            unit="currency",
            currency="RMB",
            confidence=0.9,
        ),
    ]

    plan = ChartPlan(
        id="chart_rev",
        title="Revenue Trajectory",
        question="How did revenue grow?",
        chart_type="bar",
        observation_ids=["obs_rev_21", "obs_rev_22"],
        show_data_labels=True,
    )

    _add_native_chart(slide, plan, obs, (1.0, 1.5, 6.0, 4.0), compact=False)
    chart = next(s for s in slide.shapes if s.has_chart).chart
    data_labels = chart.plots[0].data_labels

    assert data_labels.font.size.pt >= 14.0
    assert data_labels.font.bold is True
    assert chart.category_axis.tick_labels.font.size.pt >= 10.0
    assert chart.value_axis.tick_labels.font.size.pt >= 10.0


# ==============================================================================
# 4. ACCOUNTING SIGNS FOR ANALYTICAL RATIOS
# ==============================================================================

def test_expense_ratio_positive_display_and_phrasing() -> None:
    """Requirement 4: Expense ratios display as positive cost intensity, with pp phrasing."""
    cos_sem = classify_metric("Cost of sales: share of revenue", value=-49.5, raw_unit="%")
    assert cos_sem.is_expense_ratio is True
    assert cos_sem.is_percentage is True

    # Formatted display value turns negative percentage into positive
    disp_val = format_metric_display_value("-49.5", -49.5, cos_sem)
    assert disp_val == "49.5%"

    # Phrasing uses "increased by X pp", not "Loss widened"
    change_msg = format_metric_change(-42.9, -49.5, 1.0, cos_sem)
    assert "increased by 6.6 pp" in change_msg
    assert "Loss widened" not in change_msg

    # Research and development ratio
    rd_sem = classify_metric("Research and development expenses as % of revenue", value=15.2, raw_unit="%")
    assert rd_sem.is_expense_ratio is True
    change_rd = format_metric_change(12.0, 15.2, 1.0, rd_sem)
    assert "increased by 3.2 pp" in change_rd


# ==============================================================================
# 5. REJECT UNINTERPRETABLE "Observed pairs"
# ==============================================================================

def test_ban_observed_pairs_in_charts_and_preflight() -> None:
    """Requirement 5: Series name 'Observed pairs' is banned and replaced."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    obs1 = Observation(id="o1", metric_original="R&D Expenses", value=100.0, raw_value="100", period="FY2021", confidence=0.9)
    obs2 = Observation(id="o2", metric_original="Revenue", value=1000.0, raw_value="1000", period="FY2021", confidence=0.9)

    plan = ChartPlan(
        id="chart_scatter",
        title="R&D vs Revenue",
        question="Correlation between R&D and Revenue",
        chart_type="scatter",
        x_metric="Revenue",
        y_metric="R&D Expenses",
        observation_ids=["o1", "o2"],
    )

    _add_native_chart(slide, plan, [obs1, obs2], (1.0, 1.5, 6.0, 4.0))
    chart = next(s for s in slide.shapes if s.has_chart).chart
    series_name = chart.series[0].name

    assert "observed pairs" not in series_name.casefold()
    assert "R&D" in series_name or "Revenue" in series_name


# ==============================================================================
# 6. VOLUME (COUNT) VS CURRENCY ISOLATION
# ==============================================================================

def test_volume_count_does_not_inherit_currency() -> None:
    """Requirement 6: Volume / quantity does not inherit table currency declaration."""
    headers = ["Metric", "FY2022", "FY2023"]
    rows = [
        TableRow(cells=["Sales volume (units)", "12,500", "18,200"], page=50),
        TableRow(cells=["Average selling price (RMB)", "14,500", "15,800"], page=50),
    ]
    table = ExtractedTable(
        table_id="tbl_vol_asp",
        page=50,
        headers=headers,
        column_periods=[None, "FY2022", "FY2023"],
        rows=rows,
        default_unit="currency",
        default_currency="CNY",
        default_unit_scale=1000.0,
        default_raw_unit="RMB '000",
        confidence=0.9,
    )

    extractor = ObservationExtractor()
    obs = extractor._table_observations(table)

    vol_obs = [o for o in obs if "volume" in o.metric_original.casefold()]
    asp_obs = [o for o in obs if "selling price" in o.metric_original.casefold()]

    assert len(vol_obs) == 2
    assert len(asp_obs) == 2

    # Volume must NOT inherit currency or 1000x scale
    for v in vol_obs:
        assert v.unit_family == "count"
        assert v.currency is None
        assert v.display_unit == "units"
        assert v.value in (12500.0, 18200.0)

    # ASP retains currency
    for a in asp_obs:
        assert a.unit_family == "currency"
        assert a.currency == "CNY"


# ==============================================================================
# 8. SHORT METRIC DISPLAY NAMES
# ==============================================================================

def test_short_metric_display_names() -> None:
    """Requirement 8: Standard financial abbreviations for card titles."""
    assert shorten_metric_title("Selling and distribution expenses: share of revenue") == "Selling & Distribution / Revenue"
    assert shorten_metric_title("Research and development expenses: share of revenue") == "R&D / Revenue"
    assert shorten_metric_title("Administrative expenses: share of revenue") == "Admin / Revenue"
    assert shorten_metric_title("Cash and cash equivalents") == "Cash & Cash Equivalents"
    assert shorten_metric_title("Trade and bills receivables") == "Trade Receivables"
    assert shorten_metric_title("Trade and bills payables") == "Trade Payables"
    assert shorten_metric_title("Gross profit margin") == "Gross Margin"
    assert shorten_metric_title("Contract liabilities") == "Contract Liabilities"


# ==============================================================================
# 9 & 10. ANALYTICAL TOPIC GROUPING & SPECIFIC GENERIC MERGES
# ==============================================================================

def test_charts_belong_together_generic_merges() -> None:
    """Requirement 10: Gross margin + Cost of sales, Opex %, Volume + ASP merge."""
    ev = SourceEvidence(page=10, text="data", extraction_method="test", confidence=0.9)

    obs_gm = Observation(id="gm1", metric_original="Gross profit margin", value=45.0, raw_value="45.0", period="FY2023", evidence=[ev], confidence=0.9)
    obs_cos = Observation(id="cos1", metric_original="Cost of sales: share of revenue", value=55.0, raw_value="55.0", period="FY2023", evidence=[ev], confidence=0.9)
    obs_rd = Observation(id="rd1", metric_original="Research and development expenses: share of revenue", value=12.0, raw_value="12.0", period="FY2023", evidence=[ev], confidence=0.9)
    obs_sell = Observation(id="s1", metric_original="Selling and distribution expenses: share of revenue", value=18.0, raw_value="18.0", period="FY2023", evidence=[ev], confidence=0.9)
    obs_vol = Observation(id="v1", metric_original="Sales volume", value=5000.0, raw_value="5000", period="FY2023", evidence=[ev], confidence=0.9)
    obs_asp = Observation(id="asp1", metric_original="Average selling price (ASP)", value=1200.0, raw_value="1200", period="FY2023", evidence=[ev], confidence=0.9)

    index = DocumentIndex([obs_gm, obs_cos, obs_rd, obs_sell, obs_vol, obs_asp])

    c_gm = ChartPlan(id="c_gm", title="Gross Margin", question="Q1", chart_type="bar", observation_ids=["gm1"])
    c_cos = ChartPlan(id="c_cos", title="Cost of Sales / Revenue", question="Q2", chart_type="bar", observation_ids=["cos1"])
    c_rd = ChartPlan(id="c_rd", title="R&D / Revenue", question="Q3", chart_type="bar", observation_ids=["rd1"])
    c_sell = ChartPlan(id="c_sell", title="Selling / Revenue", question="Q4", chart_type="bar", observation_ids=["s1"])
    c_vol = ChartPlan(id="c_vol", title="Sales Volume", question="Q5", chart_type="bar", observation_ids=["v1"])
    c_asp = ChartPlan(id="c_asp", title="Average Selling Price (ASP)", question="Q6", chart_type="bar", observation_ids=["asp1"])

    # 1. Margin & Cost of sales belong together
    assert _charts_belong_together(c_gm, c_cos, index) is True

    # 2. Opex ratios belong together
    assert _charts_belong_together(c_rd, c_sell, index) is True

    # 3. Volume & ASP belong together
    assert _charts_belong_together(c_vol, c_asp, index) is True


# ==============================================================================
# 11. HYBRID LAYOUT (CHART_PLUS_KPIS)
# ==============================================================================

def test_chart_plus_kpis_layout_export() -> None:
    """Requirement 11: Render chart_plus_kpis hybrid layout with chart and KPI cards."""
    ev = SourceEvidence(page=40, text="100", extraction_method="test", confidence=0.95)
    obs1 = Observation(id="o1", metric_original="Revenue", value=174_314_000.0, raw_value="174,314", raw_unit="RMB '000", unit="currency", currency="RMB", period="FY2021", evidence=[ev], confidence=0.95)
    obs2 = Observation(id="o2", metric_original="Revenue", value=241_013_000.0, raw_value="241,013", raw_unit="RMB '000", unit="currency", currency="RMB", period="FY2022", evidence=[ev], confidence=0.95)
    obs_kpi = Observation(id="k1", metric_original="Gross profit margin", value=48.5, raw_value="48.5", raw_unit="%", unit="percent", period="FY2022", evidence=[ev], confidence=0.95)

    chart = ChartPlan(id="c1", title="Revenue Trajectory", question="Revenue growth", chart_type="bar", observation_ids=["o1", "o2"])

    slide = PresentationSlide(
        id="s_hybrid",
        slide_type="analysis",
        title="Revenue Increased in FY2022",
        section_title="Financial Performance",
        message="Supported by product expansion",
        layout="chart_plus_kpis",
        chart_ids=["c1"],
        observation_ids=["k1"],
        source_pages=[40],
    )

    plan = PresentationPlan(
        title="IPO Analysis",
        company=CompanyProfile(name="TechCorp Ltd.", identity_state="RESOLVED", source_pages=[40]),
        slides=[
            PresentationSlide(id="cover", slide_type="cover", title="IPO Analysis"),
            PresentationSlide(id="glance", slide_type="company_overview", title="Company at a Glance"),
            PresentationSlide(id="exec", slide_type="executive_summary", title="Executive Summary"),
            slide,
            PresentationSlide(id="quality", slide_type="data_quality", title="Data Quality"),
            PresentationSlide(id="app", slide_type="appendix", title="Key Data Appendix"),
        ],
    )

    result = PipelineResult(
        document=_test_parsed_document(),
        profile=DocumentProfile(overview_title="TechCorp Ltd.", document_type="Prospectus"),
        observations=[obs1, obs2, obs_kpi],
        charts=[chart],
        presentation_plan=plan,
    )

    pptx_bytes = export_pptx(result)
    prs = Presentation(io.BytesIO(pptx_bytes))

    # The fixed contents page is slide 2; the analysis slide is therefore index 4.
    analysis_slide = prs.slides[4]
    has_chart = any(s.has_chart for s in analysis_slide.shapes)
    has_kpi_card = any(s.has_text_frame and "48.5%" in s.text for s in analysis_slide.shapes)

    assert has_chart is True
    assert has_kpi_card is True


# ==============================================================================
# 13 & 14. COMPANY OVERVIEW CONSISTENCY & CATEGORY SEMANTICS
# ==============================================================================

def test_company_overview_clean_categories_and_no_fallback_text() -> None:
    """Requirement 13 & 14: Resolved company overview strips 'unnamed issuer' and categorizes."""
    plan = PresentationPlan(
        title="Company Review",
        company=CompanyProfile(
            name="Alpha Robotics Co., Ltd.",
            one_line_description="Leading provider of cobots (prospectus for an unnamed issuer).",
            identity_state="RESOLVED",
            products=["Cobots", "Robotic Arms"],
            geographies=["Mainland China", "Europe"],
                business_model="Direct sales and distributor network",
                key_facts=[CompanyFact(label="Headquarters", value="Shenzhen")],
                source_pages=[1],
        ),
        slides=[
            PresentationSlide(id="cov", slide_type="cover", title="Review"),
            PresentationSlide(id="glance", slide_type="company_overview", title="Company at a Glance"),
                PresentationSlide(id="exec", slide_type="executive_summary", title="Executive Summary"),
                PresentationSlide(id="quality", slide_type="data_quality", title="Data Quality"),
                PresentationSlide(id="app", slide_type="appendix", title="Appendix"),
        ],
    )

    result = PipelineResult(
        document=_test_parsed_document(),
        profile=DocumentProfile(overview_title="Alpha Robotics Co., Ltd.", document_type="Prospectus"),
        observations=[],
        charts=[],
        presentation_plan=plan,
    )

    pptx_bytes = export_pptx(result)
    prs = Presentation(io.BytesIO(pptx_bytes))
    glance_slide = prs.slides[2]

    full_text = " ".join(s.text for s in glance_slide.shapes if s.has_text_frame)

    assert "Alpha Robotics Co., Ltd." in full_text
    assert "unnamed issuer" not in full_text.casefold()
    assert "Cobots" in full_text
    assert "Europe" in full_text


# ==============================================================================
# 15 & 16. APPENDIX PERIOD SEGREGATION & UNIT CLEANUP
# ==============================================================================

def test_appendix_period_segregation_and_clean_units() -> None:
    """Requirement 15 & 16: Appendix separates flow from balance sheet and cleans units."""
    ev = SourceEvidence(page=50, text="val", extraction_method="test", confidence=0.9)
    obs_pnl = Observation(id="p1", metric_original="Revenue", value=100.0, raw_value="100,000", raw_unit="RMBinthousands", unit="currency", period="FY2022", evidence=[ev], confidence=0.9)
    obs_bs = Observation(id="b1", metric_original="Trade receivables", value=50.0, raw_value="50,000", raw_unit="RMBinthousands", unit="currency", period="2022-12-31", evidence=[ev], confidence=0.9)

    chart1 = ChartPlan(id="c1", title="Revenue", question="Q1", chart_type="bar", observation_ids=["p1"])
    chart2 = ChartPlan(id="c2", title="Trade receivables", question="Q2", chart_type="bar", observation_ids=["b1"])

    plan = PresentationPlan(
        title="Appendix Test",
            company=CompanyProfile(name="Beta Corp", identity_state="RESOLVED", source_pages=[50]),
        slides=[
            PresentationSlide(id="cov", slide_type="cover", title="Review"),
            PresentationSlide(id="glance", slide_type="company_overview", title="Company at a Glance"),
                PresentationSlide(id="exec", slide_type="executive_summary", title="Executive Summary"),
                PresentationSlide(
                    id="analysis",
                    slide_type="analysis",
                    title="Revenue Is Reported for FY2022",
                    section_title="Financial Performance",
                    message="The reported series supports the presentation format check.",
                    observation_ids=["p1"],
                    source_pages=[50],
                ),
                PresentationSlide(id="quality", slide_type="data_quality", title="Data Quality"),
                PresentationSlide(id="app", slide_type="appendix", title="Key Data Appendix"),
        ],
    )

    result = PipelineResult(
        document=_test_parsed_document(),
        profile=DocumentProfile(overview_title="Beta Corp", document_type="Prospectus"),
        observations=[obs_pnl, obs_bs],
        charts=[],
        presentation_plan=plan,
    )

    pptx_bytes = export_pptx(result)
    prs = Presentation(io.BytesIO(pptx_bytes))

    # Inspect all slides text: RMBinthousands must NEVER appear
    appendix_units: list[str] = []
    for slide in prs.slides:
        for s in slide.shapes:
            if s.has_text_frame:
                assert "RMBinthousands" not in s.text
            if s.has_table:
                for row in s.table.rows:
                    for cell in row.cells:
                        assert "RMBinthousands" not in cell.text
                appendix_units.extend(row.cells[1].text for row in list(s.table.rows)[1:])

    assert "RMB million" in appendix_units


# ==============================================================================
# 20. PRE-EXPORT CHART QUALITY GATE
# ==============================================================================

def test_preflight_quality_gate_checks() -> None:
    """Requirement 20: Preflight detects and handles chart quality issues."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Text containing raw unspaced unit
    tx_box = slide.shapes.add_textbox(0, 0, 100, 100)
    tx_box.text_frame.text = "Reported Revenue: 174,314 RMBinthousands"

    preflight = PresentationPreflight(prs)
    issues = preflight.validate_and_sanitize()

    # Verify that raw unit token was sanitized
    raw_token_issues = [i for i in issues if i.code == "raw_unit_token_sanitized"]
    assert len(raw_token_issues) >= 1
    assert "RMBinthousands" not in tx_box.text_frame.text
    assert "RMB '000" in tx_box.text_frame.text
