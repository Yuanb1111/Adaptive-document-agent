"""Selected topics can choose advanced native layouts without losing safe plans."""

from io import BytesIO

from pptx import Presentation

from adaptive_document_agent.agent.presentation_planner import PresentationPlanner
from adaptive_document_agent.agent.presentation_topic_selector import PresentationTopicSelector, series_directory
from adaptive_document_agent.agent.presentation_visual_enrichment import VisualSelection, VisualSuggestion
from adaptive_document_agent.models import (
    DocumentPage, Observation, PresentationTopic, PresentationTopicSelection, SourceEvidence,
)
from adaptive_document_agent.models.presentation import CompanySummaryItem
from adaptive_document_agent.services.pptx_export import build_presentation
from adaptive_document_agent.validation.presentation_plan_validator import PresentationPlanValidator
from tests.test_p0_composition import result_for
from tests.test_presentation_matrix import _records
from tests.test_presentation_waterfall import _bridge
from tests.test_presentation_horizon import _example as _horizon_example


def _result():
    observations, block = _records()
    result = result_for(observations, [])
    directory, _ = series_directory(result)
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(
        id="service", title="Service comparison", question="How do services compare?",
        rationale="Compare all reported category and measure cells.",
        series_ids=[row["id"] for row in directory],
    )])
    return result, block


def test_matrix_choice_reaches_compiled_plan_and_editable_ppt():
    result, block = _result()

    class Gateway:
        calls = 0

        def generate_structured(self, messages, response_model, **kwargs):
            self.calls += 1
            return VisualSelection(suggestions=[VisualSuggestion(
                topic_id="service", kind="matrix", observation_ids=block.observation_ids,
                matrix_dimension="service",
            )])

    gateway = Gateway()
    result.presentation_plan = PresentationPlanner(gateway).plan(result)
    PresentationPlanValidator().validate(result.presentation_plan, result)
    assert gateway.calls == 1
    slide = next(item for item in result.presentation_plan.slides if item.slide_type == "analysis")
    assert [visual.role for visual in slide.visual_blocks] == ["matrix"]
    deck = Presentation(BytesIO(build_presentation(result)))
    assert any(shape.name.startswith("table:comparison_matrix:")
               for page in deck.slides for shape in page.shapes)


def test_invalid_advanced_choice_keeps_existing_chart_or_data_plan():
    result, block = _result()

    class Gateway:
        def generate_structured(self, messages, response_model, **kwargs):
            return VisualSelection(suggestions=[VisualSuggestion(
                topic_id="service", kind="matrix", observation_ids=block.observation_ids[:-1],
                matrix_dimension="service",
            )])

    plan = PresentationPlanner(Gateway()).plan(result)
    slide = next(item for item in plan.slides if item.slide_type == "analysis")
    assert len(slide.visual_blocks) == 1
    assert slide.visual_blocks[0].role == "matrix"
    assert set(slide.visual_blocks[0].observation_ids) == set(block.observation_ids)
    PresentationPlanValidator().validate(plan, result)


def test_exact_reconciliation_choice_reaches_native_waterfall():
    observations, block = _bridge()
    result = result_for(observations, [])
    directory, _ = series_directory(result)
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(
        id="bridge", title="Reported reconciliation", question="How do components reconcile?",
        rationale="One source table supplies the reported components.",
        series_ids=[row["id"] for row in directory],
    )])

    class Gateway:
        def generate_structured(self, messages, response_model, **kwargs):
            return VisualSelection(suggestions=[VisualSuggestion(
                topic_id="bridge", kind="waterfall", observation_ids=block.observation_ids,
            )])

    result.presentation_plan = PresentationPlanner(Gateway()).plan(result)
    slide = next(item for item in result.presentation_plan.slides if item.slide_type == "analysis")
    assert [visual.role for visual in slide.visual_blocks] == ["waterfall"]
    deck = Presentation(BytesIO(build_presentation(result)))
    assert any(shape.name == "chart:waterfall" for page in deck.slides for shape in page.shapes)


def test_stock_flow_future_choice_reaches_separate_horizon_columns():
    document, index, block = _horizon_example()
    future = "The supplier agreement requires future payments of 30 million through 2028."
    document.pages[0].text = document.pages[0].text.replace(future, "")
    document.pages.append(DocumentPage(page_number=2, text=future))
    document.page_count = 2
    block.horizon_items[2].source_pages = [2]
    result = result_for(list(index.values()), [])
    result.document = document
    result.profile.document_summary_pages = [2]
    directory, _ = series_directory(result)
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(
        id="resources", title="Reported resources and commitments",
        question="What resources and commitments are reported?",
        rationale="Compare a dated balance, period flow and future agreement.",
        series_ids=[row["id"] for row in directory],
    )])

    class Gateway:
        def generate_structured(self, messages, response_model, **kwargs):
            return VisualSelection(suggestions=[VisualSuggestion(
                topic_id="resources", kind="horizon", observation_ids=block.observation_ids,
                horizon_items=block.horizon_items,
            )])

    result.presentation_plan = PresentationPlanner(Gateway()).plan(result)
    slide = next(item for item in result.presentation_plan.slides if item.slide_type == "analysis")
    assert [visual.role for visual in slide.visual_blocks] == ["horizon"]
    deck = Presentation(BytesIO(build_presentation(result)))
    assert sum(shape.name == "horizon:kind" for page in deck.slides for shape in page.shapes) == 3


def test_source_quoted_operating_flow_is_added_to_full_deck():
    result, _ = _result()
    passages = ["We provide shared tools to subscribers.",
                "Organisations purchase access for their teams.",
                "Subscribers pay annual service fees."]
    result.document.pages = [DocumentPage(page_number=1, text="\n".join(passages))]

    class Gateway:
        def generate_structured(self, messages, response_model, **kwargs):
            return VisualSelection()

    result.presentation_plan = PresentationPlanner(Gateway()).plan(result)
    result.presentation_plan.company.value_chain = [CompanySummaryItem(
        label=label, text=passage, source_quote=passage, source_pages=[1])
        for label, passage in zip(("Offer", "Customers", "Revenue"), passages)]
    result.presentation_plan.company.source_pages = [1]
    deck = Presentation(BytesIO(build_presentation(result)))
    assert any(page.name == "company_value_chain" for page in deck.slides)


def test_four_by_four_business_matrix_can_use_complete_selected_series():
    observations = []
    for business in ("A", "B", "C", "D"):
        for metric in ("Contribution", "Growth", "Margin", "Customers"):
            observations.append(Observation(
                id=f"{business}-{metric}", metric_original=metric, value=10,
                raw_value="10", unit="count", raw_unit="units", period="FY2025",
                period_type="fiscal_year", period_basis="FY",
                category_dimensions={"business": business}, validation_status="valid",
                confidence=.99, evidence=[SourceEvidence(page=3, text="10", row_label=business,
                                                         extraction_method="digital_table", confidence=.99)],
            ))
    result = result_for(observations, [])
    directory, lookup = series_directory(result)
    assert len(directory) == 16
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(
        id="businesses", title="Business comparison", question="How do businesses compare?",
        rationale="The reported cells form a complete comparison.",
        series_ids=[row["id"] for row in directory],
    )])
    PresentationTopicSelector._validate(result.presentation_topics, lookup)

    class Gateway:
        def generate_structured(self, messages, response_model, **kwargs):
            return VisualSelection(suggestions=[VisualSuggestion(
                topic_id="businesses", kind="matrix",
                observation_ids=[item.id for item in observations], matrix_dimension="business",
            )])

    result.presentation_plan = PresentationPlanner(Gateway()).plan(result)
    slide = next(item for item in result.presentation_plan.slides if item.slide_type == "analysis")
    assert [visual.role for visual in slide.visual_blocks] == ["matrix"]
    deck = Presentation(BytesIO(build_presentation(result)))
    table = next(shape.table for page in deck.slides for shape in page.shapes
                 if shape.name.startswith("table:comparison_matrix:"))
    assert len(table.rows) == 5 and len(table.columns) == 5
