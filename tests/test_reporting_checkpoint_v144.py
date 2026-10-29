"""Whole-report checkpoints stop at missing new coverage responses.

Private artifacts stay outside the repository. Recorded semantics, canonical
artifact transport, new author metadata and synthetic complete-flow verdicts are
kept distinct. A historical rejection never becomes fresh semantic acceptance.
"""
from contextlib import contextmanager
from copy import deepcopy
from io import BytesIO
import hashlib
import json
import os
from pathlib import Path
from threading import Event

import pytest
from pydantic import ValidationError

from adaptive_document_agent.agent.orchestrator import DocumentOrchestrator
from adaptive_document_agent.models import AnalysisResult, AnalysisTask, PipelineResult
from adaptive_document_agent.services.llm import LLMGateway, LLMSettings, ProviderName
from adaptive_document_agent.services.llm.base import LLMResponse
from adaptive_document_agent.services.llm.exceptions import LLMStructuredOutputError
from adaptive_document_agent.services.llm.usage import LLMUsage
from adaptive_document_agent.services.pipeline_completion import is_complete
from adaptive_document_agent.services.source_quotes import normalize_quote
from adaptive_document_agent.utils.pipeline_version import PIPELINE_VERSION, EXTRACTION_VERSION, ANALYSIS_VERSION
from tests.test_reporting_checkpoint_v142 import (
    RecordedResponseUnavailable, SyntheticReportingClient, _gateway, _payloads,
    _synthetic_checkpoint, _warnings,
)
from tests.test_reporting_checkpoint_v143 import EarlyCheckpointClient


SUPPORTED_RECORDED_VERSION = '2026-10-15-introduction-review-convergence-v143'


def _meaning_keys(decisions):
    return [(decision['part_id'], index) for decision in decisions
            for index, _ in enumerate(decision['required_meanings'])]


def _adapt_recorded_author_refs(draft, source, verdict):
    """Add only literal, owned metadata from retained covered=true verdicts.

    The prose and facts remain byte-for-byte equivalent JSON values. Missing
    meanings stay unmapped; no semantic inference supplies a missing quotation.
    Even an added reference is an author metadata adaptation, not new approval.
    """
    data = deepcopy(draft)
    owners = {fact['id']: part['id'] for part in source['parts'] for fact in part['facts']}
    items = [item for page in data['pages'] for item in page['items']]
    for item in items:
        item['meaning_refs'] = []
    assigned = set()
    for part in verdict['parts']:
        for coverage in part['coverage']:
            if not coverage['covered'] or not coverage['summary_quote'].strip():
                continue
            quote = normalize_quote(coverage['summary_quote'])
            candidates = [item for item in items if part['part_id'] in {owners[fid] for fid in item['fact_ids']}
                          and quote in normalize_quote(item['text'])]
            if not candidates:
                continue
            item = candidates[0]
            literal = coverage['summary_quote'] if coverage['summary_quote'] in item['text'] else item['text']
            item['meaning_refs'].append(dict(part_id=part['part_id'],
                meaning_index=coverage['meaning_index'], summary_quote=literal))
            assigned.add((part['part_id'], coverage['meaning_index']))
    stripped = deepcopy(data)
    for page in stripped['pages']:
        for item in page['items']:
            item.pop('meaning_refs')
    assert stripped == draft
    return data, assigned


class RecordedCoverageCheckpointClient(EarlyCheckpointClient):
    """Retained v143 responses followed by an explicitly unavailable v144 edit."""

    def __init__(self, checkpoint):
        super().__init__(checkpoint, failure_kind='binding')
        self.bound_reviews = [record['introduction_audience_review'] for record in self.intro['plan_audits']
                              if 'introduction_audience_review' in record]
        assert len(self.bound_reviews) == 2 and all(row['status'] == 'bound' for row in self.bound_reviews)
        self.audience_attempts = [dict(decisions=row['decisions'], review=row['review']) for row in self.bound_reviews]
        self.compact = [record for record in self.intro['plan_audits'] if 'compact_attempt' in record]
        self.coverage = _warnings(checkpoint, 'presentation_topic_coverage_review')[0]
        self.coverage_positions, self.takeaway_positions = {}, {}
        self.first_author_source, self.adapted_author, self.assigned_refs = None, None, set()
        self.missing_responses = []

    def generate_structured(self, messages, response_model, **kwargs):
        schema, name = response_model, response_model.__name__
        if name in {'IntroductionPages', 'SummaryReadBatch', 'IntroductionMeaningRevision'}:
            return super().generate_structured(messages, response_model, **kwargs)
        self.calls.append(name)
        values = _payloads(messages)
        if name == 'IntroductionAudienceSelection':
            decisions = deepcopy(self.bound_reviews[0]['decisions'])
            for decision in decisions:
                for meaning in decision['required_meanings']:
                    # Restore the authoring contract, then let Python bind the
                    # same retained owning passage. Semantics and owners stay exact.
                    assert not meaning.pop('proposed_source_quote', '')
                    meaning['source_quote'] = ''
            return self._response(schema, dict(decisions=decisions), 'retained_selection_author_transport_adaptation')
        if name == 'IntroductionAudienceReview':
            payload = values[-1]
            self.review_inputs.append(deepcopy(payload))
            if self.review_position >= len(self.bound_reviews):
                raise RecordedResponseUnavailable('No additional current-source review retained')
            record = self.bound_reviews[self.review_position]
            assert payload['decisions'] == record['decisions']
            assert payload['current_selection_digest'] == record['selection_digest']
            self.review_position += 1
            return self._response(schema, record['review'], 'recorded_bound_independent_verdict')
        if name == 'SummarySynthesisDraft':
            assert self.adapted_author is None
            self.first_author_source = deepcopy(values[0])
            self.adapted_author, self.assigned_refs = _adapt_recorded_author_refs(
                self.compact[0]['draft'], values[0], self.compact[0]['compact_semantic_review'])
            return self._response(schema, self.adapted_author, 'historical_draft_with_literal_covered_quote_metadata_only')
        if name in {'SummarySynthesisRevision', 'SummaryCoverageReview', 'SummaryLateOmissionReview'}:
            # A legacy response has no complete author mapping for this new
            # request. New mandatory metadata/semantics cannot be reconstructed.
            self.missing_responses.append(name)
            self.provenance.append(dict(operation=name, provenance='unavailable_new_coverage_protocol_response'))
            raise RecordedResponseUnavailable('No new meaning-complete correction response was recorded')
        if name == 'InsightList':
            ids = {record['task_id'] for record in values[0]}
            insights = [item.model_dump(mode='json') for item in self.checkpoint.insights
                        if set(item.result_ids) <= ids]
            return self._response(schema, dict(insights=insights), 'retained_validated_insights_artifact')
        if name == 'ReportPlan':
            return self._response(schema, self.checkpoint.report_plan.model_dump(mode='json'), 'retained_report_plan_artifact')
        if name == 'PresentationTopicSelection':
            return self._response(schema, self.coverage['original_selection'], 'retained_validated_topic_selection_artifact')
        if name == 'TopicCoverageReview':
            requested = tuple(values[0]['series_requiring_coverage_decision'])
            batch = next(row for row in self.coverage['batches'] if tuple(row['series_ids']) == requested)
            position = self.coverage_positions.get(requested, 0)
            self.coverage_positions[requested] = position + 1
            if position >= len(batch['attempts']):
                raise RecordedResponseUnavailable('No further retained coverage response for these exact series')
            attempt = batch['attempts'][position]
            if 'review' in attempt:
                return self._response(schema, attempt['review'], 'recorded_topic_coverage_response')
            try:
                schema.model_validate_json(attempt['invalid_response'])
            except ValidationError as cause:
                self.provenance.append(dict(operation=name, provenance='recorded_invalid_topic_coverage_response'))
                raise LLMStructuredOutputError('Retained coverage schema failure', response=LLMResponse(
                    text=attempt['invalid_response'], usage=LLMUsage(provider='mock', model='offline-checkpoint'))) from cause
            raise AssertionError('A retained invalid coverage response unexpectedly passed its original schema')
        if name == 'TakeawayCorrection':
            topic = values[0]['topic']
            records = [row for row in self.checkpoint.validation_warnings
                       if row.code == 'presentation_topic_takeaway_repair' and topic['id'] in row.related_ids]
            position = self.takeaway_positions.get(topic['id'], 0)
            self.takeaway_positions[topic['id']] = position + 1
            if position >= len(records):
                raise RecordedResponseUnavailable('No retained correction for this source topic')
            record = json.loads(records[position].message)
            assert (record['original_title'], record['original_takeaway']) == (topic['title'], topic['takeaway'])
            return self._response(schema, dict(title=record['proposed_title'], takeaway=record['proposed_takeaway']),
                                  'recorded_takeaway_correction_artifact')
        if name == 'VisualSelection':
            return self._response(schema, dict(suggestions=[]), 'synthetic_no_optional_visual_change')
        if name in {'BriefSourcePages', 'ExecutiveBrief', 'ExecutiveBriefPatch', 'BriefMeaningReview'}:
            raise AssertionError('Known failed introduction must stop new briefing requests: ' + name)
        raise RecordedResponseUnavailable('No retained response for new operation ' + name)


@pytest.mark.skipif(not os.environ.get('ADA_REPORT_COVERAGE_CHECKPOINT'), reason='External v143 reporting checkpoint not configured')
def test_external_v143_whole_reporting_retry_stops_at_new_author_mapping_boundary():
    path = Path(os.environ['ADA_REPORT_COVERAGE_CHECKPOINT'])
    pdf_path = Path(os.environ['ADA_REPORT_SOURCE_PDF'])
    original_bytes, raw = path.read_bytes(), pdf_path.read_bytes()
    checkpoint_hash = hashlib.sha256(original_bytes).hexdigest()
    source = PipelineResult.model_validate_json(original_bytes)
    assert source.pipeline_version == SUPPORTED_RECORDED_VERSION
    assert (EXTRACTION_VERSION, ANALYSIS_VERSION) == ('extraction-v20', 'analysis-v76')
    assert source.pipeline_state == 'partial' and source.retry_from == 'reporting'
    assert hashlib.sha256(raw).hexdigest() == source.document.sha256
    assert len(source.observations) == 1623 and len(source.analysis_results) == 5
    assert len(source.insights) == 5 and len(source.charts) == 25 and len(source.executive_brief.items) == 5
    original = source.model_copy(deep=True)
    retry = source.model_copy(deep=True)
    # Fixture-only migration of the unchanged evidence/analysis contracts. The
    # actual version, raw-PDF, privacy, source and retry guards remain active.
    retry.pipeline_version = PIPELINE_VERSION
    client = RecordedCoverageCheckpointClient(source)
    gateway = _gateway(client)
    assert gateway.discovery_workers == 1  # Real deferred-attachment branch.
    result = DocumentOrchestrator(gateway).analyse_pdf(raw, resume=retry)
    assert source == original and hashlib.sha256(path.read_bytes()).hexdigest() == checkpoint_hash
    assert hashlib.sha256(pdf_path.read_bytes()).hexdigest() == source.document.sha256
    assert result.document == original.document
    assert result.observations == original.observations and result.analysis_results == original.analysis_results
    assert not is_complete(result) and result.pipeline_state == 'partial' and result.retry_from == 'reporting'
    assert result.executive_brief is None
    assert any(issue.code == 'editorial_stages_skipped' for issue in result.validation_warnings)
    assert client.missing_responses and client.missing_responses[0] == 'SummarySynthesisRevision'
    assert not any(name in client.calls for name in ('BriefSourcePages', 'ExecutiveBrief', 'BriefMeaningReview'))
    assert client.review_position == 2
    retained = client.intro['audience_selection']['decisions']
    required = set(_meaning_keys(retained))
    assert len(required) == 55
    assert len(client.first_author_source['meaning_inventory']) == 55
    historical_missing = {(part['part_id'], check['meaning_index'])
        for part in client.compact[0]['compact_semantic_review']['parts'] for check in part['coverage'] if not check['covered']}
    assert historical_missing and not historical_missing & client.assigned_refs
    unmapped = required - client.assigned_refs
    assert len(client.assigned_refs) == 43 and len(unmapped) == 12
    assert historical_missing == unmapped
    summary = _warnings(result, 'company_introduction_summary_audit')[0]
    assert summary['parts'] == client.intro['parts']
    assert set(_meaning_keys(summary['audience_selection']['decisions'])) == required
    # An accepted selection cannot bypass missing visible-copy coverage.
    assert summary['status'] == 'incomplete'
    from adaptive_document_agent.services.pptx_export import build_presentation
    with pytest.raises(ValueError, match='source-verified opening introduction'):
        build_presentation(result)


def _coverage_fixture():
    source, raw = _synthetic_checkpoint()
    source.document.pages[1].text = (
        'SUMMARY\nWe provide scheduling software for operating teams.\n'
        'Customers pay annual subscription fees.\n'
        'Staff training is required before service delivery.\n'
        'The scheduling service is available only in the home market.\n')
    inputs = source.observations[:3]
    source.analysis_plan = [AnalysisTask(id='synthetic-sum', title='Subject A observed total',
        description='Sum the three supplied Subject A observations.', analysis_type='sum', tool_name='sum_values',
        observation_query={'observation_ids': [item.id for item in inputs]},
        reason='Explicit synthetic calculation over supplied compatible counts.', expected_output='Total count.')]
    source.analysis_results = [AnalysisResult(task_id='synthetic-sum', title='Subject A observed total',
        result=sum(item.value for item in inputs), input_observation_ids=[item.id for item in inputs],
        evidence=[quote.model_copy(deep=True) for item in inputs for quote in item.evidence], confidence=.99)]
    assert source.analysis_results[0].result == 33
    return source, raw


class SyntheticCoverageFlowClient(SyntheticReportingClient):
    """Self-reported mappings remain distinct from fresh semantic judgments."""
    supports_concurrent_requests = True
    intro_operations = {'IntroductionPages', 'SummaryReadBatch', 'IntroductionAudienceSelection',
        'IntroductionAudienceInventoryCorrection', 'IntroductionAudienceReview', 'IntroductionMeaningRevision',
        'SummarySynthesisDraft', 'SummarySynthesisRevision', 'SummaryCoverageReview', 'SummaryLateOmissionReview'}

    def __init__(self, source, *, failure_mode=None):
        super().__init__(source, accept_introduction=True)
        from adaptive_document_agent.agent.summary_late_omission import SummaryLateOmissionReview
        self.additional_schemas = {'SummaryLateOmissionReview': SummaryLateOmissionReview}
        self.failure_mode, self.parts, self.candidate = failure_mode, None, 0
        self.review_inputs, self.revision_inputs, self.late_inputs = [], [], []
        self.known_failure = Event()
        self.after_failure_calls = []

    def _typed(self, schema, value):
        from tests.summary_author_helpers import writer_test_value
        typed = schema.model_validate(writer_test_value(schema, value))
        return typed, LLMResponse(text=typed.model_dump_json(), usage=LLMUsage(provider='mock', model='synthetic-v144'))

    def _prose(self, candidate):
        lines = [part['required_meanings'][0]['meaning'] for part in self.parts]
        if candidate == 0:
            lines[:3] = ['We provide scheduling software.', 'Customers pay fees.', 'Staff training is required.']
        elif candidate == 1:
            lines[2] = 'Staff training is required.'
        else:
            lines[3] = 'The scheduling service is available.'
        return ' '.join(lines)

    def _item(self, text, *, only=None):
        parts = self.parts if only is None else [self.parts[only]]
        refs = [dict(part_id=part['id'], meaning_index=0, summary_quote=text) for part in parts]
        if self.failure_mode == 'missing_mapping':
            refs = refs[:-1]
        return dict(label='Software, delivery and scope', text=text,
            fact_ids=[fact['id'] for part in parts for fact in part['facts']], meaning_refs=refs)

    def generate_structured(self, messages, response_model, **kwargs):
        schema, name = response_model, response_model.__name__
        if self.known_failure.is_set() and name not in self.intro_operations:
            self.after_failure_calls.append(name)
            raise AssertionError('A new dependent editorial request started after known intro failure: ' + name)
        if name not in {'InsightList', 'ReportPlan', 'SummarySynthesisDraft', 'SummarySynthesisRevision',
                        'SummaryCoverageReview', 'SummaryLateOmissionReview'}:
            return super().generate_structured(messages, response_model, **kwargs)
        self.calls.append(name)
        if name == 'ReportPlan':
            # The ordinary planner retains its deterministic evidence-bound
            # outline when the fixture proposes no extra organization.
            return self._typed(schema, dict(title='Synthetic source analysis', sections=[]))
        payload = _payloads(messages)[-1]
        if name == 'InsightList':
            # The model may choose no additional interpretation; the actual
            # Python sum and all source evidence remain retained in the report.
            return self._typed(schema, dict(insights=[]))
        if name == 'SummarySynthesisDraft':
            payload = next(value for value in _payloads(messages) if isinstance(value, dict) and 'parts' in value)
            self.parts = deepcopy(payload['parts'])
            assert len(self.parts) == 4 and len(payload['meaning_inventory']) == 4
            return self._typed(schema, dict(pages=[dict(title='Software and service scope',
                items=[self._item(self._prose(0))])]))
        if name == 'SummarySynthesisRevision':
            self.revision_inputs.append(deepcopy(payload))
            if self.failure_mode in {'missing_mapping', 'false_mapping'}:
                self.known_failure.set()
                raise RecordedResponseUnavailable('No synthetic approval or edit is available for this negative fixture')
            self.candidate += 1
            if self.candidate < 3:
                item_id = payload['editable_groups'][0]['items'][0]['item_id']
                return self._typed(schema, dict(replacements=[dict(item_id=item_id,
                    replacement=self._item(self._prose(self.candidate)))], additions=[], titles=[]))
            assert self.candidate == 3
            assert schema.model_json_schema()['properties']['replacements']['maxItems'] == 0
            group_id = payload['editable_groups'][0]['group_id']
            text = ('The scheduling service is available.' if self.failure_mode == 'late_still_missing'
                    else self.parts[3]['required_meanings'][0]['meaning'])
            return self._typed(schema, dict(replacements=[], titles=[], additions=[dict(group_id=group_id,
                item=self._item(text, only=3))]))
        if name == 'SummaryLateOmissionReview':
            self.late_inputs.append(deepcopy(payload))
            by_id = {part['id']: part for part in self.parts}
            previous = [dict(**row, resolved=True,
                summary_quote=by_id[row['part_id']]['required_meanings'][row['meaning_index']]['meaning'],
                reason='The exact previously missing assertion is now visible in its owned paragraph.')
                for row in payload['previous_deficits']]
            current = [dict(**row, missing_only=self.failure_mode != 'late_not_authorized',
                reason='A source geographic limitation is absent from otherwise supported current copy.')
                for row in payload['current_deficits']]
            if self.failure_mode == 'late_not_authorized':
                self.known_failure.set()
            return self._typed(schema, dict(previous_deficits=previous, current_deficits=current))
        self.review_inputs.append(deepcopy(payload))
        # Every round must re-review the entire selected source. A previous
        # accepted part cannot conceal qualifiers lost in an edited paragraph.
        assert {part['id'] for part in payload['source']['parts']} == {part['id'] for part in self.parts}
        missing = ({0, 1, 2} if self.candidate == 0 else {2} if self.candidate == 1 else
                   {3} if self.candidate == 2 or self.failure_mode == 'late_still_missing' else set())
        index_by_id = {part['id']: index for index, part in enumerate(self.parts)}
        verdicts = []
        for part in payload['source']['parts']:
            covered = index_by_id[part['id']] not in missing
            verdicts.append(dict(part_id=part['id'], accepted=covered,
                reason='The source assertion is complete.' if covered else 'A necessary source qualification is missing.',
                coverage=[dict(meaning_index=0, covered=covered,
                    summary_quote=part['required_meanings'][0]['meaning'] if covered else '',
                    reason='The complete assertion is literal in its owned copy.' if covered else 'The qualification is absent.')]))
        if self.candidate == 3 and self.failure_mode == 'late_still_missing':
            self.known_failure.set()
        return self._typed(schema, dict(parts=verdicts, identity_accepted=None,
            opening_accepted=True, opening_reason='The generated opening explains scheduling software.',
            opening_slide_id=payload['opening_target']['slide_id'],
            opening_item_texts=payload['opening_target']['item_texts']))


def _run_coverage_fixture(monkeypatch, source, raw, client, *, background=False, wait_for_failure=False):
    from adaptive_document_agent.agent import company_introduction
    prepared_instances = []
    if background:
        original_prepare = company_introduction.prepare_company_introduction

        @contextmanager
        def capture(*args, **kwargs):
            with original_prepare(*args, **kwargs) as prepared:
                prepared_instances.append(prepared)
                yield prepared

        monkeypatch.setattr(company_introduction, 'prepare_company_introduction', capture)
        # This is still an in-memory MockLLMClient. Non-local settings exercise
        # the actual concurrent gateway path without credentials, SDK or network.
        gateway = LLMGateway(client, LLMSettings(provider=ProviderName.DEEPSEEK,
            model='offline-memory-adapter', discovery_workers=4, presentation_strategy='full_catalog'))
    else:
        gateway = _gateway(client)
    assert (gateway.discovery_workers > 1) == background
    runner = DocumentOrchestrator(gateway)
    monkeypatch.setattr(runner, '_load_document', lambda data, digest:
        source.document if data == raw and digest == source.document.sha256 else pytest.fail('Wrong source boundary'))
    if wait_for_failure:
        validate = runner._validate_analysis

        def validate_then_wait(profile, observations, results):
            issues = validate(profile, observations, results)
            assert len(prepared_instances) == 1
            with pytest.raises((ValueError, RecordedResponseUnavailable)):
                prepared_instances[0]._future.result(timeout=60)
            assert client.known_failure.is_set()
            return issues

        monkeypatch.setattr(runner, '_validate_analysis', validate_then_wait)
    result = runner.analyse_pdf(raw, resume=source)
    assert all(prepared._pool is None for prepared in prepared_instances)
    return result


@pytest.mark.parametrize('background', [False, True])
def test_synthetic_whole_report_late_qualification_repair_is_visible_in_native_ppt(monkeypatch, background):
    source, raw = _coverage_fixture()
    original = source.model_copy(deep=True)
    client = SyntheticCoverageFlowClient(source)
    result = _run_coverage_fixture(monkeypatch, source, raw, client, background=background)
    assert source == original and result.observations == original.observations
    assert result.analysis_results == original.analysis_results and result.analysis_results[0].result == 33
    assert is_complete(result), [issue.message for issue in result.validation_warnings
                                 if issue.code == 'company_introduction_unavailable']
    assert len(client.review_inputs) == 4 and len(client.late_inputs) == 1
    assert len(client.revision_inputs) == 3
    assert all(len(payload['source']['parts']) == 4 for payload in client.review_inputs)
    assert result.executive_brief is not None and not client.after_failure_calls
    summary = _warnings(result, 'company_introduction_summary_audit')[0]
    late = [row['late_omission_review'] for row in summary['plan_audits'] if 'late_omission_review' in row]
    assert len(late) == 1 and late[0]['authorized'] is True
    from adaptive_document_agent.services.pptx_export import build_presentation
    from pptx import Presentation
    exported = build_presentation(result)
    reopened = Presentation(BytesIO(exported))
    assert len(reopened.slides) > 0
    intro = [' '.join(shape.text for shape in slide.shapes if shape.has_text_frame)
             for slide in reopened.slides if any(shape.has_text_frame and 'Software and service scope' in shape.text
                                                 for shape in slide.shapes)]
    assert intro
    visible = ' '.join(intro)
    assert 'only in the home market' in visible
    assert 'before service delivery' in visible and 'annual subscription fees' in visible


@pytest.mark.parametrize(('failure_mode', 'background'), [
    ('missing_mapping', False), ('false_mapping', False), ('false_mapping', True),
    ('late_not_authorized', False), ('late_still_missing', False),
])
def test_whole_report_known_intro_failure_preserves_analysis_and_stops_brief(monkeypatch, failure_mode, background):
    source, raw = _coverage_fixture()
    original = source.model_copy(deep=True)
    client = SyntheticCoverageFlowClient(source, failure_mode=failure_mode)
    result = _run_coverage_fixture(monkeypatch, source, raw, client,
        background=background, wait_for_failure=background)
    assert source == original and result.observations == original.observations
    assert result.analysis_results == original.analysis_results and result.analysis_results[0].result == 33
    assert not is_complete(result) and result.pipeline_state == 'partial' and result.retry_from == 'reporting'
    assert client.known_failure.is_set() and not client.after_failure_calls
    assert result.executive_brief is None
    assert not any(name in client.calls for name in ('BriefSourcePages', 'ExecutiveBrief', 'BriefMeaningReview'))
    assert any(issue.code == 'editorial_stages_skipped' for issue in result.validation_warnings)
    if failure_mode == 'missing_mapping':
        assert not client.review_inputs
        summary = _warnings(result, 'company_introduction_summary_audit')[0]
        authors = [row['authored_meaning_coverage'] for row in summary['plan_audits'] if 'authored_meaning_coverage' in row]
        assert authors[0]['missing'] and authors[0]['semantic_approval'] is False
    elif failure_mode == 'false_mapping':
        assert 1 <= len(client.review_inputs) <= 2
    else:
        assert len(client.late_inputs) == 1
        assert len(client.review_inputs) == (4 if failure_mode == 'late_still_missing' else 3)
    if background:
        assert 'InsightList' not in client.calls and 'PresentationTopicSelection' not in client.calls
    from adaptive_document_agent.services.pptx_export import build_presentation
    with pytest.raises(ValueError, match='source-verified opening introduction'):
        build_presentation(result)
