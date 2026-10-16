"""Layout recovery preserves meaning/evidence and never invents absent verdicts."""
from contextlib import contextmanager
from copy import deepcopy
from io import BytesIO
import hashlib
import os
from pathlib import Path
from threading import Event

import pytest

from adaptive_document_agent.agent.orchestrator import DocumentOrchestrator
from adaptive_document_agent.models import PipelineResult
from adaptive_document_agent.services.llm import LLMGateway, LLMSettings, ProviderName
from adaptive_document_agent.services.pipeline_completion import is_complete
from adaptive_document_agent.utils.pipeline_version import PIPELINE_VERSION, EXTRACTION_VERSION, ANALYSIS_VERSION
from tests.test_reporting_checkpoint_v142 import RecordedResponseUnavailable, _payloads, _warnings
from tests.test_reporting_checkpoint_v143 import EarlyCheckpointClient
from tests.test_reporting_checkpoint_v144 import (
    RecordedCoverageCheckpointClient, SyntheticCoverageFlowClient, _coverage_fixture, _meaning_keys, _run_coverage_fixture,
)


SUPPORTED_RECORDED_VERSION = '2026-10-15-introduction-coverage-failfast-v144'


class RecordedLayoutCheckpointClient(RecordedCoverageCheckpointClient):
    """Retained introduction protocol plus two historical overflowing drafts.

    A live harness may route only retained_intro_operations here, and route new
    authoring/review/reporting through a separate real gateway. This offline
    adapter itself has no network-capable transport or credentials.
    """
    retained_intro_operations = {'IntroductionPages', 'SummaryReadBatch', 'IntroductionAudienceSelection',
        'IntroductionAudienceReview', 'IntroductionMeaningRevision'}

    def __init__(self, checkpoint):
        # The v144 checkpoint has no accepted topic selection or coverage audit.
        # Reuse only its exact retained introduction protocol, never invent one.
        EarlyCheckpointClient.__init__(self, checkpoint, failure_kind='binding')
        self.bound_reviews = [record['introduction_audience_review'] for record in self.intro['plan_audits']
                              if 'introduction_audience_review' in record]
        assert len(self.bound_reviews) == 2 and all(row['status'] == 'bound' for row in self.bound_reviews)
        self.audience_attempts = [dict(decisions=row['decisions'], review=row['review']) for row in self.bound_reviews]
        self.compact = [record for record in self.intro['plan_audits'] if 'compact_attempt' in record]
        assert len(self.compact) == 1
        self.layout_drafts = [deepcopy(row['draft']) for row in self.compact[0]['layout_attempts']]
        assert len(self.layout_drafts) == 2
        self.author_inputs, self.missing_responses = [], []
        self.known_failure = Event()

    def generate_structured(self, messages, response_model, **kwargs):
        name = response_model.__name__
        if name in self.retained_intro_operations:
            return super().generate_structured(messages, response_model, **kwargs)
        self.calls.append(name)
        if name == 'SummarySynthesisDraft':
            self.author_inputs.append(deepcopy(_payloads(messages)))
            position = len(self.author_inputs) - 1
            if position < len(self.layout_drafts):
                provenance = ('recorded_initial_author_artifact' if position == 0 else
                              'historical_layout_response_under_new_feedback_transport')
                return self._response(response_model, self.layout_drafts[position], provenance)
        self.missing_responses.append(name)
        self.provenance.append(dict(operation=name, provenance='unavailable_new_compression_or_independent_response'))
        self.known_failure.set()
        raise RecordedResponseUnavailable('No new compressed draft or independent semantic response was recorded')


def _wait_at_real_validation_boundary(monkeypatch, runner, client):
    """Control fixture scheduling only; leave values and all production gates real."""
    from adaptive_document_agent.agent import company_introduction
    original_prepare = company_introduction.prepare_company_introduction
    prepared_instances, validations = [], []

    @contextmanager
    def capture(*args, **kwargs):
        with original_prepare(*args, **kwargs) as prepared:
            prepared_instances.append(prepared)
            yield prepared

    monkeypatch.setattr(company_introduction, 'prepare_company_introduction', capture)
    validate = runner._validate_analysis

    def validate_then_finish(profile, observations, results):
        issues = validate(profile, observations, results)
        validations.append(issues)
        assert len(prepared_instances) == 1 and prepared_instances[0]._future is not None
        with pytest.raises((ValueError, RecordedResponseUnavailable)):
            prepared_instances[0]._future.result(timeout=60)
        assert client.known_failure.is_set()
        return issues

    monkeypatch.setattr(runner, '_validate_analysis', validate_then_finish)
    return prepared_instances, validations


@pytest.mark.skipif(not os.environ.get('ADA_REPORT_LAYOUT_CHECKPOINT'), reason='External v144 layout checkpoint not configured')
def test_external_v144_whole_retry_replays_two_drafts_then_stops_at_new_compression_reply(monkeypatch):
    path, pdf_path = Path(os.environ['ADA_REPORT_LAYOUT_CHECKPOINT']), Path(os.environ['ADA_REPORT_SOURCE_PDF'])
    original_bytes, raw = path.read_bytes(), pdf_path.read_bytes()
    original_hash = hashlib.sha256(original_bytes).hexdigest()
    source = PipelineResult.model_validate_json(original_bytes)
    assert source.pipeline_version == SUPPORTED_RECORDED_VERSION
    # Source row/unit and direction contracts changed in v146. Historical
    # reporting replay below still retains its exact input evidence.
    assert (EXTRACTION_VERSION, ANALYSIS_VERSION) == ('extraction-v20', 'analysis-v75')
    assert source.pipeline_state == 'partial' and source.retry_from == 'reporting'
    assert hashlib.sha256(raw).hexdigest() == source.document.sha256
    assert len(source.observations) == 1623 and len(source.analysis_results) == 12
    original, retry = source.model_copy(deep=True), source.model_copy(deep=True)
    # Only this supported unchanged extraction/analysis fixture adopts the new
    # reporting version. The real PDF SHA/version/retry guards remain enabled.
    retry.pipeline_version = PIPELINE_VERSION
    client = RecordedLayoutCheckpointClient(source)
    gateway = LLMGateway(client, LLMSettings(provider=ProviderName.DEEPSEEK,
        model='offline-memory-layout-replay', discovery_workers=4))
    runner = DocumentOrchestrator(gateway)
    prepared, validations = _wait_at_real_validation_boundary(monkeypatch, runner, client)
    result = runner.analyse_pdf(raw, resume=retry)
    assert validations and all(issue in result.validation_warnings for issue in validations[0])
    assert all(item._pool is None for item in prepared)
    assert source == original and hashlib.sha256(path.read_bytes()).hexdigest() == original_hash
    assert hashlib.sha256(pdf_path.read_bytes()).hexdigest() == original.document.sha256
    assert result.document == original.document
    assert result.observations == original.observations and result.analysis_results == original.analysis_results
    assert client.review_position == 2
    assert len(client.author_inputs) >= 3 and client.missing_responses[0] == 'SummarySynthesisDraft'
    assert not any(name in client.calls for name in ('InsightList', 'ReportPlan', 'PresentationTopicSelection',
        'SummaryCoverageReview', 'BriefSourcePages', 'ExecutiveBrief', 'BriefMeaningReview'))
    summary = _warnings(result, 'company_introduction_summary_audit')[0]
    required = _meaning_keys(client.intro['audience_selection']['decisions'])
    assert len(required) == 41 and _meaning_keys(summary['audience_selection']['decisions']) == required
    assert summary['parts'] == client.intro['parts']
    attempts = [item for row in summary['plan_audits'] for item in row.get('layout_attempts', [])]
    assert len(attempts) >= 2 and all(item['status'] == 'overflow' for item in attempts[:2])
    assert all(sum(group['placement'] == 'body' for group in item['physical_groups']) == 7 for item in attempts[:2])
    assert attempts[0]['layout_feedback']['version'] == 'physical-meaning-layout-v1'
    assert attempts[1]['compression_round'] == 2
    assert result.pipeline_state == 'partial' and result.retry_from == 'reporting' and not is_complete(result)
    assert result.executive_brief is None and any(i.code == 'editorial_stages_skipped' for i in result.validation_warnings)
    from adaptive_document_agent.services.pptx_export import build_presentation
    with pytest.raises(ValueError, match='source-verified opening introduction'):
        build_presentation(result)


def _layout_claims():
    return ['We provide scheduling software for operating teams.'] + [
        f'The {chr(65 + index // 26)}{chr(65 + index % 26)} team requires staff training before service delivery.'
        for index in range(40)]


def _layout_fixture():
    source, raw = _coverage_fixture()
    claims = _layout_claims()
    paragraphs = [' '.join(claims[start:start + 6]) for start in range(0, len(claims), 6)]
    source.document.pages[1].text = 'SUMMARY\n' + '\n'.join(paragraphs) + '\n'
    return source, raw


class SyntheticLayoutFlowClient(SyntheticCoverageFlowClient):
    """Forty-one explicit synthetic assertions survive bounded compression."""

    def __init__(self, source, *, failure_mode=None):
        super().__init__(source, failure_mode=failure_mode)
        self.author_inputs, self.author_responses, self.entries = [], [], []

    def _item_for(self, entries):
        owners = {row['part_id'] for row in entries}
        text = ' '.join(row['meaning'] for row in entries)
        refs = [dict(part_id=row['part_id'], meaning_index=row['meaning_index'], summary_quote=row['meaning'])
                for row in entries]
        return dict(label='Scheduling and training prerequisites', text=text,
            fact_ids=[fact['id'] for part in self.parts if part['id'] in owners for fact in part['facts']],
            meaning_refs=refs)

    def _draft(self, *, compressed):
        groups = []
        if not compressed:
            for index, part in enumerate(self.parts):
                entries = [row for row in self.entries if row['part_id'] == part['id']]
                groups.append(dict(title='Scheduling prerequisites ' + chr(65 + index), items=[self._item_for(entries)]))
        else:
            for index, offset in enumerate(range(0, len(self.entries), 9)):
                entries = self.entries[offset:offset + 9]
                item = self._item_for(entries)
                if offset + len(entries) == len(self.entries):
                    if self.failure_mode == 'drops_meaning':
                        # A syntactically valid self-reference is deliberately
                        # dishonest: only independent review can judge coverage.
                        item['text'] = ' '.join(row['meaning'] for row in entries[:-1])
                        item['meaning_refs'][-1]['summary_quote'] = entries[-2]['meaning']
                    elif self.failure_mode == 'drops_evidence':
                        item['fact_ids'] = [self.parts[0]['facts'][0]['id']]
                groups.append(dict(title='Scheduling prerequisites ' + chr(65 + index), items=[item]))
        return dict(pages=groups)

    def generate_structured(self, messages, response_model, **kwargs):
        schema, name = response_model, response_model.__name__
        if name not in {'IntroductionAudienceSelection', 'SummarySynthesisDraft', 'SummaryCoverageReview',
                        'SummarySynthesisRevision'}:
            return super().generate_structured(messages, response_model, **kwargs)
        self.calls.append(name)
        payloads = _payloads(messages)
        if name == 'IntroductionAudienceSelection':
            parts = payloads[0]['parts']
            claims = _layout_claims()
            assert len(parts) == 7
            decisions = []
            for index, part in enumerate(parts):
                meanings = claims[index * 6:index * 6 + 6]
                assert part['source_text'].strip() == ' '.join(meanings)
                decisions.append(dict(part_id=part['id'], include=True,
                    reason='The synthetic service and explicit training prerequisites define the audience scope.',
                    required_meanings=[dict(meaning=meaning) for meaning in meanings]))
            return self._typed(schema, dict(decisions=decisions))
        if name == 'SummarySynthesisDraft':
            source = next(value for value in payloads if isinstance(value, dict) and 'parts' in value)
            self.parts = deepcopy(source['parts'])
            self.entries = [dict(part_id=part['id'], meaning_index=index, meaning=meaning['meaning'])
                for part in self.parts for index, meaning in enumerate(part['required_meanings'])]
            assert len(self.entries) == 41 and len(source['meaning_inventory']) == 41
            self.author_inputs.append(deepcopy(payloads))
            compressed = len(self.author_inputs) > 1 and self.failure_mode != 'still_seven_pages'
            value = self._draft(compressed=compressed)
            self.author_responses.append(deepcopy(value))
            if self.failure_mode == 'still_seven_pages' and len(self.author_inputs) == 3:
                self.known_failure.set()
            if self.failure_mode == 'drops_evidence' and compressed:
                self.known_failure.set()
            return self._typed(schema, value)
        if name == 'SummarySynthesisRevision':
            self.known_failure.set()
            raise RecordedResponseUnavailable('Negative compressed meaning has no approved corrective response')
        payload = payloads[-1]
        self.review_inputs.append(deepcopy(payload))
        assert {part['id'] for part in payload['source']['parts']} == {part['id'] for part in self.parts}
        final = self.entries[-1]
        verdicts = []
        for part in payload['source']['parts']:
            checks = []
            for index, meaning in enumerate(part['required_meanings']):
                covered = not (self.failure_mode == 'drops_meaning' and
                               (part['id'], index) == (final['part_id'], final['meaning_index']))
                checks.append(dict(meaning_index=index, covered=covered,
                    summary_quote=meaning['meaning'] if covered else '',
                    reason='The exact supported requirement is visible.' if covered else
                           'The final selected team prerequisite was removed during compression.'))
            verdicts.append(dict(part_id=part['id'], accepted=all(check['covered'] for check in checks),
                reason='Every requirement is visible.' if all(check['covered'] for check in checks)
                       else 'Compression omitted a selected source requirement.', coverage=checks))
        return self._typed(schema, dict(parts=verdicts, identity_accepted=None,
            opening_accepted=True, opening_reason='The opening explains scheduling software and its intended users.',
            opening_slide_id=payload['opening_target']['slide_id'],
            opening_item_texts=payload['opening_target']['item_texts']))


@pytest.mark.parametrize('background', [False, True])
def test_whole_report_compression_retains_all_41_meanings_in_five_native_intro_pages(monkeypatch, background):
    source, raw = _layout_fixture()
    original = source.model_copy(deep=True)
    client = SyntheticLayoutFlowClient(source)
    result = _run_coverage_fixture(monkeypatch, source, raw, client, background=background)
    assert source == original and result.observations == original.observations
    assert result.analysis_results == original.analysis_results and result.analysis_results[0].result == 33
    assert is_complete(result), [i.message for i in result.validation_warnings if i.code == 'company_introduction_unavailable']
    assert len(client.author_inputs) == 2 and len(client.review_inputs) == 1
    review = result.presentation_plan.company.summary_review
    assert len(_meaning_keys(review.audience_selection['decisions'])) == 41
    assert len(result.presentation_plan.company.summary_pages) == 5
    attempts = [item for row in review.plan_audits for item in row.get('layout_attempts', [])]
    assert [item['status'] for item in attempts] == ['overflow', 'success']
    assert len(attempts[0]['physical_groups']) == 7
    assert attempts[0]['layout_feedback']['version'] == 'physical-meaning-layout-v1'
    from adaptive_document_agent.services.pptx_export import build_presentation
    from pptx import Presentation
    reopened = Presentation(BytesIO(build_presentation(result)))
    titles = {page.title for page in result.presentation_plan.company.summary_pages}
    intro = [' '.join(shape.text for shape in slide.shapes if shape.has_text_frame)
        for slide in reopened.slides if any(shape.has_text_frame and shape.text.strip() in titles
                                           for shape in slide.shapes)
        # Contents lists each group in separate textboxes too. Count only the
        # actual introduction slides and inspect their native paragraph text.
        and not any(shape.has_text_frame and shape.text.strip() == 'Contents' for shape in slide.shapes)]
    assert len(intro) == 5
    visible = ' '.join(intro)
    assert all(claim in visible for claim in _layout_claims())


@pytest.mark.parametrize('failure_mode', ['drops_meaning', 'drops_evidence', 'still_seven_pages'])
def test_whole_report_bad_compression_cannot_bypass_meaning_evidence_or_page_caps(monkeypatch, failure_mode):
    source, raw = _layout_fixture()
    original = source.model_copy(deep=True)
    client = SyntheticLayoutFlowClient(source, failure_mode=failure_mode)
    result = _run_coverage_fixture(monkeypatch, source, raw, client)
    assert source == original and result.observations == original.observations
    assert result.analysis_results == original.analysis_results and result.analysis_results[0].result == 33
    assert not is_complete(result) and result.pipeline_state == 'partial' and result.retry_from == 'reporting'
    assert result.executive_brief is None
    assert not any(name in client.calls for name in ('BriefSourcePages', 'ExecutiveBrief', 'BriefMeaningReview'))
    if failure_mode == 'drops_meaning':
        assert client.review_inputs
    else:
        assert not client.review_inputs
    if failure_mode == 'still_seven_pages':
        assert len(client.author_inputs) == 3
    from adaptive_document_agent.services.pptx_export import build_presentation
    with pytest.raises(ValueError, match='source-verified opening introduction'):
        build_presentation(result)
