"""Reuse reduces paid work without promoting mock verdicts to real evidence."""
import json
from copy import deepcopy

import pytest

from adaptive_document_agent.agent import retained_analysis as reuse
from adaptive_document_agent.agent import retained_analysis_review as review
from adaptive_document_agent.models import ValidationIssue, DocumentPage
from adaptive_document_agent.services.llm import LLMGateway, LLMSettings, MockLLMClient, ProviderName, PrivacyMode
from adaptive_document_agent.services.pipeline_completion import require_topic_coverage
from tests.test_scoped_presentation_recovery import _linked_result as linked_fixture


def _linked_result():
    result = linked_fixture()
    result.document.pages = [DocumentPage(page_number=234,
        text='Revenue RMB thousands FY2023 267,025 FY2024 325,257 FY2025 521,747')]
    return result


def selection(result):
    ids = list(reuse.findings(result))
    return dict(decisions=[dict(insight_id=i, decision='include', reason='Synthetic complete input coverage')
                           for i in ids], ordered_insight_ids=ids)


def verdict(result, *, accepted=True):
    return dict(verdicts=[dict(slide_id=s.id, accepted=accepted, reason='Synthetic source-review protocol fixture')
                         for s in result.presentation_plan.slides if s.slide_type == 'analysis'],
                omissions_accepted=True, omission_reason='Synthetic omission protocol fixture')


def gateway(responses):
    client = MockLLMClient(responses)
    return LLMGateway(client, LLMSettings(provider=ProviderName.MOCK, model='offline',
                                         privacy_mode=PrivacyMode.LOCAL_ONLY))


def prepared():
    result = _linked_result()
    gw = gateway([selection(result)])
    reuse.select_retained_analysis(gw, result)
    result.presentation_plan = reuse.compile_retained_analysis(result)
    return result, gw


def test_two_bounded_calls_preserve_inputs_and_require_independent_review():
    result = _linked_result()
    original = result.model_dump(include={'observations', 'insights', 'analysis_results'})
    gw = gateway([selection(result)])
    assert reuse.can_reuse(gw, result)
    reuse.select_retained_analysis(gw, result)
    result.presentation_plan = reuse.compile_retained_analysis(result)
    with pytest.raises(ValueError, match='Retained-analysis evidence'):
        require_topic_coverage(result)
    gw.client.responses.append(json.dumps(verdict(result)))
    review.review_retained_analysis(gw, result)
    require_topic_coverage(result)
    review.review_retained_analysis(gw, result)
    assert len(gw.client.calls) == 2
    assert [u['operation'] for u in gw.usage] == ['RetainedAnalysisSelection', 'RetainedAnalysisReview']
    assert result.model_dump(include=set(original)) == original
    assert result.presentation_plan.themes
    assert not result.presentation_topics


@pytest.mark.parametrize('change', ['value', 'copy', 'page', 'calculation', 'selection', 'requirements', 'audit', 'chart'])
def test_any_changed_evidence_or_claim_invalidates_approval(change):
    result, _ = prepared()
    review.review_retained_analysis(gateway([verdict(result)]), result)
    if change == 'value':
        result.observations[0].raw_value = '987654321'
    elif change == 'copy':
        next(s for s in result.presentation_plan.slides if s.slide_type == 'analysis').message = 'An unrelated claim'
    elif change == 'page':
        result.document.pages[0].text += ' New source qualification'
    elif change == 'calculation':
        result.analysis_results[0].result = -123
    elif change == 'selection':
        issue = next(w for w in result.validation_warnings if w.code == reuse.SELECTION_CODE)
        audit = json.loads(issue.message)
        audit['selection']['decisions'][0]['reason'] = 'Different scope'
        issue.message = json.dumps(audit)
    elif change == 'requirements':
        result.profile.analysis_focus = 'New material priority'
    elif change == 'audit':
        result.validation_warnings.append(ValidationIssue(code='report_audit_scope_review', stage='presentation', message='changed'))
    elif change == 'chart':
        result.charts[0].title = 'Changed chart claim'
    assert not review.review_complete(result)
    with pytest.raises(ValueError):
        require_topic_coverage(result)


@pytest.mark.parametrize('failure', ['reject', 'missing', 'duplicate', 'omit'])
def test_rejected_or_incomplete_verdict_cannot_publish(failure):
    result, _ = prepared()
    response = verdict(result)
    if failure == 'reject':
        response['verdicts'][0]['accepted'] = False
    elif failure == 'missing':
        response['verdicts'] = []
    elif failure == 'duplicate':
        response['verdicts'] *= 2
    else:
        response['omissions_accepted'] = False
    gw = gateway([response])
    with pytest.raises(ValueError):
        review.review_retained_analysis(gw, result)
    assert len(gw.client.calls) == 1
    assert not review.review_complete(result)


def test_old_failure_is_never_erased_by_changing_routes():
    result, gw = prepared()
    result.validation_warnings.append(ValidationIssue(code='presentation_topic_selection_failed', stage='presentation', message='failure'))
    assert not reuse.can_reuse(gw, result)
    review.review_retained_analysis(gateway([verdict(result)]), result)
    with pytest.raises(ValueError, match='topic review incomplete'):
        require_topic_coverage(result)


def test_selection_rejects_unknown_and_duplicate_ids_without_retry():
    result = _linked_result()
    invalid = selection(result)
    invalid['decisions'] *= 2
    gw = gateway([invalid])
    with pytest.raises(ValueError):
        reuse.select_retained_analysis(gw, result)
    assert len(gw.client.calls) == 1 and reuse.selected_record(result) is None


def test_selection_invalidates_after_input_change():
    result, _ = prepared()
    result.observations[0].raw_value = '1234'
    with pytest.raises(ValueError, match='stale'):
        reuse.compile_retained_analysis(result)


def test_legacy_and_full_catalog_mode_remain_available():
    result = _linked_result()
    gw = gateway([])
    gw.settings.presentation_strategy = 'full_catalog'
    assert not reuse.can_reuse(gw, result)
    assert review.review_complete(result)


def test_native_composition_preserves_reviewed_analysis_and_emits_pptx():
    from io import BytesIO
    from pptx import Presentation
    from adaptive_document_agent.agent.presentation_budget_refinement import measure
    from adaptive_document_agent.services.pptx_export import build_presentation
    result, _ = prepared()
    # Native preparation runs before review in production too.
    measure(result)
    review.review_retained_analysis(gateway([verdict(result)]), result)
    pptx = build_presentation(result)
    assert len(Presentation(BytesIO(pptx)).slides) > 0
    assert review.review_complete(result)


def test_display_translation_metadata_does_not_repeat_source_review():
    from adaptive_document_agent.models.customization import ReportRequirements
    result, _ = prepared()
    review.review_retained_analysis(gateway([verdict(result)]), result)
    result.profile.report_requirements = ReportRequirements(translation_language='en',
        copy_translations={'Revenue': 'Revenue'})
    assert review.review_complete(result)


def test_missing_original_page_cannot_even_request_approval():
    result, _ = prepared()
    result.document.pages = []
    gw = gateway([])
    with pytest.raises(ValueError, match='missing original source pages'):
        review.review_retained_analysis(gw, result)
    assert not gw.client.calls


@pytest.mark.parametrize('value', [
    [dict(raw='123.40', bbox=[1, 2, 3, 4], optional=None) for _ in range(10)],
    {'$record_table': {'columns': ['not encoding'], 'rows': [['literal']]}},
    [{'$literal_object': [['raw', 'unchanged']]}] * 4,
    [{'one': None}, {'two': 0}],
    {'source': [dict(header='FY2025', cells=[dict(raw='-1', column=2)] * 8)] * 8},
])
def test_column_encoding_preserves_every_literal_field_and_geometry(value):
    from adaptive_document_agent.utils.record_encoding import encode_records, decode_records
    assert decode_records(encode_records(value)) == value


def test_record_encoding_compacts_real_review_without_dropping_fields():
    from adaptive_document_agent.utils.record_encoding import encode_records, decode_records
    result, _ = prepared()
    context = review.review_context(result)
    encoded = encode_records(context)
    assert decode_records(encoded) == context
    assert len(json.dumps(encoded)) < len(json.dumps(context))


def test_rejected_semantic_response_is_not_reused_from_disk(tmp_path):
    from adaptive_document_agent.utils.caching import DiskCache
    result, _ = prepared()
    gw = gateway([verdict(result, accepted=False), verdict(result)])
    gw.cache = DiskCache(tmp_path)
    with pytest.raises(ValueError):
        review.review_retained_analysis(gw, result)
    assert not list(tmp_path.glob('llm-*.json'))
    review.review_retained_analysis(gw, result)
    assert len(gw.client.calls) == 2 and review.review_complete(result)


def test_source_review_keeps_full_page_and_header_geometry():
    result, _ = prepared()
    page = result.document.pages[0]
    page.text += '\n' + 'Source qualifier ' * 2000
    page.text_blocks = [dict(lines=[dict(text='FY2025', bbox=[20.125, 40.5, 80.375, 51.5]),
                                   dict(text='Unaudited', bbox=[120.5, 40.5, 180.5, 51.5])])]
    context = review.review_context(result)
    source = next(p for p in context['source_pages'] if p['page'] == page.page_number)
    assert source['text'] == page.text
    assert {line[0] for line in source['layout_lines']} == {'FY2025', 'Unaudited'}
    review.review_retained_analysis(gateway([verdict(result)]), result)
    page.text_blocks[0]['lines'][0]['bbox'][0] += .00001
    assert not review.review_complete(result)


def test_new_source_review_preserves_reasoning_while_selection_reuses_existing_analysis():
    from adaptive_document_agent.services.llm.reasoning_policy import request_reasoning_policy
    settings = LLMSettings(provider=ProviderName.DEEPSEEK, model='deepseek-flash')
    selector = request_reasoning_policy(settings, stage='presentation',
        operation='RetainedAnalysisSelection', model=settings.model)
    reviewer = request_reasoning_policy(settings, stage='presentation',
        operation='RetainedAnalysisReview', model=settings.model)
    assert selector['intent'] == 'reduce'
    assert reviewer['intent'] == 'preserve' and reviewer['options'] == {}


@pytest.mark.parametrize('accepted', [True, False, 'selection_failure'])
def test_default_route_runs_complete_reporting_with_real_gateway_and_export_gates(monkeypatch, accepted):
    from adaptive_document_agent.agent.orchestrator import DocumentOrchestrator
    from adaptive_document_agent.services.pipeline_completion import is_complete
    from adaptive_document_agent.services.pptx_export import build_presentation
    from adaptive_document_agent.utils.record_encoding import decode_records
    from tests.test_reporting_checkpoint_v144 import _coverage_fixture, SyntheticCoverageFlowClient
    from tests.test_reporting_checkpoint_v142 import _payloads
    source, raw = _coverage_fixture()

    class Client(SyntheticCoverageFlowClient):
        def generate_structured(self, messages, schema, **kwargs):
            name = schema.__name__
            if name not in {'RetainedAnalysisSelection', 'RetainedAnalysisReview'}:
                return super().generate_structured(messages, schema, **kwargs)
            self.calls.append(name)
            payload = decode_records(_payloads(messages)[-1])
            if name == 'RetainedAnalysisSelection':
                if accepted == 'selection_failure':
                    from adaptive_document_agent.services.llm.exceptions import LLMResponseError
                    raise LLMResponseError('Synthetic bounded selection failure')
                ids = [f['insight']['id'] for f in payload['findings'] if f['selectable']]
                value = dict(decisions=[dict(insight_id=f['insight']['id'],
                    decision='include' if f['insight']['id'] in ids else 'omit',
                    reason='Synthetic selection of complete calculation inputs') for f in payload['findings']],
                    ordered_insight_ids=ids)
            else:
                value = dict(verdicts=[dict(slide_id=s['id'], accepted=accepted,
                    reason='Synthetic independent evidence verdict') for s in payload['slides']],
                    omissions_accepted=accepted, omission_reason='Synthetic independent omission verdict')
            return self._typed(schema, value)

    client = Client(source)
    client.additional_schemas.update({s.__name__: s for s in
        (reuse.RetainedAnalysisSelection, review.RetainedAnalysisReview)})
    gw = LLMGateway(client, LLMSettings(provider=ProviderName.MOCK, model='offline-v216'))
    runner = DocumentOrchestrator(gw)
    monkeypatch.setattr(runner, '_load_document', lambda *args: source.document)
    result = runner.analyse_pdf(raw, resume=source)
    assert client.calls.count('RetainedAnalysisSelection') == 1
    if accepted == 'selection_failure':
        assert result.pipeline_state == 'partial' and result.failed_stage == 'presentation_selection'
        assert result.retry_from == 'reporting' and result.llm_usage
        assert result.observations == source.observations
        assert result.analysis_results == source.analysis_results
        assert 'ExecutiveBrief' not in client.calls and 'RetainedAnalysisReview' not in client.calls
        assert 'PresentationTopicSelection' not in client.calls
        return
    assert client.calls.count('RetainedAnalysisReview') == 1, result.failure_details
    assert 'PresentationTopicSelection' not in client.calls and 'TopicCoverageDelta' not in client.calls
    assert result.observations == source.observations
    assert result.analysis_results == source.analysis_results
    assert is_complete(result) is accepted, result.failure_details
    if accepted:
        from io import BytesIO
        from pptx import Presentation
        deck = Presentation(BytesIO(build_presentation(result)))
        assert len(deck.slides) > 0
    else:
        assert result.pipeline_state == 'partial' and result.failed_stage == 'presentation_coverage'
        with pytest.raises(ValueError):
            build_presentation(result)
