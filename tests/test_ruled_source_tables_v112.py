"""Decoded ruled-table geometry, exact original headers and independent inventory."""
from io import BytesIO

import pymupdf
import pypdfium2
import pytest
from pptx import Presentation

from adaptive_document_agent.extraction.ruled_source_tables import extract_ruled_source_tables
from adaptive_document_agent.extraction.source_table_audit import page_inventory_errors
from adaptive_document_agent.models import DocumentPage
from adaptive_document_agent.services.requested_tables import render_requested_tables, verify_requested_tables
from tests.test_report_customization import table_result,presentation


def ruled_pdf(*,rotation=0,mixed=False):
    doc=pymupdf.open()
    page=doc.new_page(width=600,height=500)
    for y in (65,120,240):
        page.draw_line((50,y),(550,y))
    page.insert_text((340,82),'As of December 31,',fontsize=9)
    for x,year in ((270,'2021'),(360,'2022'),(470,'6M2023')):
        page.insert_text((x,105),year,fontsize=9)
        page.insert_text((x,135),'USD000',fontsize=9)
    for y,label,values in ((160,'1st quarter',['12','15','18']),
                           (185,'3rd quarter',['10','20','N/A']),
                           (210,'Total',['22','35','18'])):
        page.insert_text((50,y),label,fontsize=9)
        for x,value in zip((270,360,470),values):
            page.insert_text((x,y),value,fontsize=9)
    page.insert_text((50,260),'Note: Totals retain the original reporting units.',fontsize=7)
    page.insert_text((50,275),'This second source annotation line must also survive.',fontsize=7)
    page.insert_text((50,315),'Subsequent body prose is not a table footnote.',fontsize=10)
    if mixed:
        for y,label,values in ((340,'Remote sites',['60','70','80']),
                               (365,'Local sites',['90','95','98'])):
            page.insert_text((50,y),label,fontsize=9)
            for x,value in zip((270,360,470),values):
                page.insert_text((x,y),value,fontsize=9)
    if rotation:
        output=pymupdf.open()
        target=output.new_page(width=500,height=600)
        target.show_pdf_page(target.rect,doc,0,rotate=rotation)
        raw=output.tobytes();output.close()
    else:
        raw=doc.tobytes()
    doc.close()
    return raw


@pytest.mark.parametrize('rotation',[0,90])
def test_geometry_recovers_first_row_total_na_and_literal_multi_tier_headers(rotation):
    raw=ruled_pdf(rotation=rotation)
    with pymupdf.open(stream=raw,filetype='pdf') as native,pypdfium2.PdfDocument(raw) as decoded:
        page=decoded[0]
        tp=page.get_textpage()
        tables,audit=extract_ruled_source_tables(tp,native[0],1)
        tp.close();page.close()
    assert len(tables)==1
    table=tables[0]
    assert any(row==['1st quarter','12','15','18'] for row in table.raw_cells)
    assert any(row==['Total','22','35','18'] for row in table.raw_cells)
    assert any('6M2023' in row for row in table.raw_header_cells)
    assert 'annotation line' in '\n'.join(table.raw_footnotes)
    assert 'Subsequent body prose' not in '\n'.join(table.raw_footnotes)
    source=DocumentPage(page_number=1,raw_tables=tables,source_table_inventory=audit)
    assert not page_inventory_errors(source)
    table.raw_cells[-1][-1]='19'
    assert any('changed' in e for e in page_inventory_errors(source))


def test_merged_original_headers_survive_native_save_reload_and_period_tampering():
    result=table_result()
    page=next(p for p in result.document.pages if p.raw_tables)
    table=page.raw_tables[0]
    table.raw_header_cells=[['','As of December 31,',''],['','2021','2022'],['','USD000','USD000']]
    table.raw_header_spans=[(0,1,2)]
    deck=presentation()
    render_requested_tables(deck,result)
    stream=BytesIO();deck.save(stream)
    restored=Presentation(stream)
    assert verify_requested_tables(restored,result)
    shape=next(s for sl in restored.slides for s in sl.shapes if s.has_table)
    assert shape.table.cell(0,1).is_merge_origin
    assert shape.table.cell(0,1).text=='As of December 31,'
    shape.table.cell(1,2).text='FY2023'
    with pytest.raises(ValueError,match='omitted or changed'):
        verify_requested_tables(restored,result)


def test_close_adjacent_column_headings_do_not_form_one_wrong_header():
    from adaptive_document_agent.extraction.source_geometry import Glyph,GeometryLine
    from adaptive_document_agent.extraction.ruled_source_tables import _header
    line=GeometryLine([Glyph('Number',117,10,143,17,7),Glyph('of',146,10,152,17,7),
                       Glyph('revenue',158,10,182,17,7),Glyph('per',185,10,195,17,7)],10,17,7)
    row,spans=_header(line,[(88,108,'right'),(148,152,'right'),(179,195,'right')],[],50,210)
    assert row==['','','Number of','revenue per']
    assert not spans


def test_wide_header_above_short_dash_values_stays_in_the_data_column():
    from adaptive_document_agent.extraction.source_geometry import Glyph,GeometryLine
    from adaptive_document_agent.extraction.ruled_source_tables import _header
    line=GeometryLine([Glyph('2020',225,10,240,17,7),Glyph('2021',295,10,310,17,7)],10,17,7)
    row,_=_header(line,[(245,250,'right'),(315,320,'right')],[],85,350)
    assert row==['','2020','2021']


def test_ruled_fast_path_preserves_a_separate_unruled_grid_on_the_same_page():
    from adaptive_document_agent.extraction.table_extractor import TableExtractor
    extractor=TableExtractor()
    extractor.extract(ruled_pdf(mixed=True))
    retained=extractor.raw_tables_by_page[1]
    assert any('1st quarter' in row[0] for t in retained for row in t.raw_cells)
    assert any('Remote sites' in (row[0] or '') for t in retained for row in t.raw_cells)
    assert extractor.source_inventory_by_page[1]['status']=='unchecked'


def test_analytical_mutations_cannot_change_the_original_source_grid():
    from adaptive_document_agent.extraction.table_extractor import TableExtractor
    extractor=TableExtractor()
    analytical=extractor.extract(ruled_pdf())
    analytical[1][0].raw_cells[0][1]='999'
    assert extractor.raw_tables_by_page[1][0].raw_cells[0][1]=='12'


@pytest.mark.parametrize('target',['table_footer','notes_footer','provenance'])
def test_native_page_references_are_verified_not_just_the_cell_values(target):
    import json
    from adaptive_document_agent.services.pptx_export import _source_footer
    result=table_result()
    table=next(t for p in result.document.pages for t in p.raw_tables)
    table.raw_footnotes=['Note: Original source annotation.']
    deck=presentation();render_requested_tables(deck,result)
    assert verify_requested_tables(deck,result)
    source_slide=next(s for s in deck.slides if any(sh.has_table for sh in s.shapes))
    if target=='provenance':
        metadata=json.loads(source_slide.notes_slide.notes_text_frame.text)
        metadata['CUSTOM_SOURCE_TABLES_V1'][0]['page']=99
        source_slide.notes_slide.notes_text_frame.text=json.dumps(metadata)
    else:
        slide=source_slide if target=='table_footer' else next(s for s in deck.slides
            if any(sh.name.startswith('customization:source_footnote:') for sh in s.shapes))
        next(sh for sh in slide.shapes if sh.name=='customization:source_footer').text=_source_footer([99])
    with pytest.raises(ValueError,match='page'):
        verify_requested_tables(deck,result)
