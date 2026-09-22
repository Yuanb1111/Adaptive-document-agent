"""Generic claim binding and evidence/display separation regressions."""

import io
import json

import pytest
from pptx import Presentation

from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import Insight, PresentationPlan, PresentationSlide, PresentationVisualBlock
from adaptive_document_agent.services.pptx_export import build_presentation, _base_slide
from adaptive_document_agent.services.slide_compositor import Rect, _put_commentary, render_composed_slide
from adaptive_document_agent.services.presentation_style import readable_axis_bounds
from adaptive_document_agent.validation.claim_validator import ClaimValidator, repair_presentation_plan
from adaptive_document_agent.validation.cross_slide_validator import CrossSlideValidator
from adaptive_document_agent.validation.scoped_narrative_values import scoped_value_errors
from tests.test_p0_composition import observation, paired_result


@pytest.mark.parametrize("left,right", [("Units shipped", "Service hours"), ("Survey responses", "Completion rate"), ("Revenue", "Gross profit")])
@pytest.mark.parametrize("reverse", [False, True])
def test_repairs_target_only_the_wrong_metric_token(left, right, reverse):
    obs = [observation(f"{prefix}{i}", label, v, f"FY{2022+i}", unit="count")
           for prefix, label, values in [("a", left, (40, 80)), ("b", right, (70, 35))]
           for i, v in enumerate(values)]
    if reverse:
        obs.reverse()
    slide = PresentationSlide(id="summary", slide_type="executive_summary", title="Summary",
        message=f"{left} increased and {right} increased.", observation_ids=[o.id for o in obs])
    detail = slide.model_copy(update={"id": "detail", "slide_type": "analysis", "message": ""})
    plan = PresentationPlan(title="Review", slides=[slide, detail])
    original = [o.model_dump() for o in obs]
    plan, repairs = repair_presentation_plan(plan, obs)
    assert repairs
    assert slide.message == f"{left} increased and {right} decreased."
    CrossSlideValidator(plan, obs).validate_and_repair()
    assert slide.message == f"{left} increased and {right} decreased."
    assert not repair_presentation_plan(plan, obs)[1]
    assert not ClaimValidator().validate_plan(plan, obs)
    assert original == [o.model_dump() for o in obs]


def test_adjusted_and_reported_loss_do_not_share_repair_aliases():
    obs = [observation(f"{prefix}{i}", label, v, f"FY{2022+i}")
           for prefix, label, values in [("a", "Loss for the year/period", (-30, -90)),
                                         ("b", "Adjusted net loss", (-80, -20))]
           for i, v in enumerate(values)]
    slide = PresentationSlide(id="summary", slide_type="executive_summary", title="Summary",
        message="Loss for the year widened and adjusted net loss narrowed.", observation_ids=[o.id for o in obs])
    plan = PresentationPlan(title="Review", slides=[slide])
    assert not repair_presentation_plan(plan, obs)[1]
    CrossSlideValidator(plan, obs).validate_and_repair()
    assert slide.title == "Summary"
    assert slide.message == "Loss for the year widened and adjusted net loss narrowed."


def test_chart_only_references_enter_direction_validation():
    result = paired_result()
    slide = result.presentation_plan.slides[3]
    slide.message = "Units shipped decreased."
    assert not slide.observation_ids
    assert ClaimValidator().validate_plan(result.presentation_plan, result.observations, result.charts)
    repair_presentation_plan(result.presentation_plan, result.observations, result.charts)
    assert slide.message == "Units shipped increased."


@pytest.mark.parametrize("metric", ["Cash and cash equivalents", "Reservoir volume", "Survey responses"])
def test_continuous_decline_claim_is_blocked_when_last_period_rebounds(metric):
    obs = [observation(str(i), metric, v, f"FY{2021+i}", unit="count") for i, v in enumerate((120, 70, 90))]
    slide = PresentationSlide(id="s", slide_type="analysis", title="Review",
        message=f"{metric} declined at each subsequent reported date.", observation_ids=[o.id for o in obs])
    plan = PresentationPlan(title="Review", slides=[slide])
    before = slide.message
    repair_presentation_plan(plan, obs)
    assert slide.message == before
    assert "non_monotonic_claim" in {i.code for i in ClaimValidator().validate_plan(plan, obs)}
    slide.message = f"{metric} declined overall, with a partial rebound in the final year."
    assert not ClaimValidator().validate_plan(plan, obs)


@pytest.mark.parametrize("label,other", [("Maintenance expense", "Maintenance expenditure"),
    ("Completed responses", "Invited respondents"), ("Net output", "Gross output")])
def test_endpoint_must_belong_to_its_named_metric(label, other):
    obs = [observation(f"{prefix}{i}", name, v, f"FY{2022+i}", unit="count")
           for prefix, name, values in [("a", label, (110, 230)), ("b", other, (115, 240))]
           for i, v in enumerate(values)]
    slide = PresentationSlide(id="s", slide_type="analysis", title="Review",
        message=f"{label} increased from 115 to 240.")
    assert len(scoped_value_errors(slide, obs)) == 2
    slide.message = f"{label} increased from 110 to 230; {other} increased from 115 to 240."
    assert not scoped_value_errors(slide, obs)
    # A qualified longer label must not be interpreted as the nested short one.
    obs[2].metric_original = obs[3].metric_original = "Adjusted " + label
    slide.message = f"Adjusted {label} increased from 115 to 240."
    assert not scoped_value_errors(slide, obs)


def test_endpoint_scale_and_accounting_magnitude_are_preserved():
    own = observation("own", "Maintenance expense", -230000)
    own.raw_value = "(230)"
    own.raw_unit = "USD '000"
    other = observation("other", "Maintenance expenditure", 240000)
    other.raw_value = "240"
    slide = PresentationSlide(id="s", slide_type="analysis", title="Review", message="Maintenance expense rose to USD0.23 million.")
    assert not scoped_value_errors(slide, [own, other])
    slide.message = "Maintenance expense rose to USD240 thousand."
    assert scoped_value_errors(slide, [own, other])


def test_unreferenced_metric_name_cannot_borrow_selected_metric_values():
    expense = observation("expense", "Maintenance expense", 230)
    expenditure = observation("expenditure", "Maintenance expenditure", 240)
    slide = PresentationSlide(id="s", slide_type="analysis", title="Review", message="Maintenance expense rose to 240.")
    assert scoped_value_errors(slide, [expenditure], [expense, expenditure])


def test_alternative_source_label_and_clause_boundary_do_not_borrow_prior_metric():
    obs = [observation("profit", "Gross profit", 150), observation("loss", "Loss for the year/period", -90)]
    slide = PresentationSlide(id="s", slide_type="analysis", title="Review",
        message="Gross profit rose to 150, but the loss for the year widened to 90.")
    assert not scoped_value_errors(slide, obs)


def test_references_do_not_create_raw_data_or_duplicate_narrative_pages():
    result = paired_result()
    extra = [observation(f"ref{i}", f"Context measure {i}", i + 3, unit="count") for i in range(20)]
    result.observations.extend(extra)
    insight = Insight(id="context", title="Detailed context", narrative="Source context retained in full. " * 100,
        kind="reported_fact", evidence=extra[0].evidence)
    result.insights = [insight]
    slide = result.presentation_plan.slides[3]
    slide.observation_ids = [o.id for o in extra]
    slide.insight_ids = [insight.id]
    original = result.model_dump()
    deck = Presentation(io.BytesIO(build_presentation(result)))
    composed = [s for s in deck.slides if s.name.startswith("composed_")]
    assert len(composed) == 1
    assert not any(s.name.startswith("kpi:ref") for s in composed[0].shapes)
    body = " ".join(s.text for s in composed[0].shapes if s.has_text_frame)
    assert "Source context retained in full" not in body
    notes = json.loads(composed[0].notes_slide.notes_text_frame.text)
    assert {o["id"] for o in notes["observations"]} >= {o.id for o in extra}
    assert notes["insights"][0]["narrative"] == insight.narrative
    assert notes["observations"][0]["evidence"]
    assert result.model_dump()["observations"] == original["observations"]
    assert result.model_dump()["insights"] == original["insights"]
    assert slide.observation_ids == [o.id for o in extra]


def test_explicit_commentary_role_remains_visible_even_with_bullets():
    result = paired_result()
    result.insights = [Insight(id="c", title="Context", narrative="A distinct material caveat.", kind="reported_fact")]
    result.presentation_plan.slides[3].visual_blocks.append(PresentationVisualBlock(role="commentary", insight_ids=["c"]))
    deck = Presentation(io.BytesIO(build_presentation(result)))
    assert "A distinct material caveat." in " ".join(s.text for sl in deck.slides for s in sl.shapes if s.has_text_frame)


def test_commentary_prefers_complete_sentence_at_break():
    deck = Presentation()
    slide = _base_slide(deck, "Review")
    text = "A short complete sentence. " + "Supporting details " * 12 + "remain in the following sentence."
    remaining = _put_commentary(slide, text, Rect(.5, 1, 4, 1))
    displayed = next(s.text for s in slide.shapes if s.name.startswith("composed:"))
    assert displayed == "A short complete sentence. "
    assert displayed + remaining == text


def test_data_only_fallback_keeps_visible_observations():
    result = paired_result()
    slide = PresentationSlide(id="data", slide_type="analysis", title="Reported values", observation_ids=["a0", "a1"])
    deck = Presentation()
    pages = render_composed_slide(deck, slide, [], result, DocumentIndex(result.observations))
    assert {s.name for page in pages for s in page.shapes} >= {"kpi:a0", "kpi:a1"}


@pytest.mark.parametrize("low,high", [(-217.38, 0), (-.0043, .00023), (-35.98, 71.3), (-8600, 0)])
def test_negative_axis_has_regular_ticks_and_preserves_headroom(low, high):
    floor, ceiling, step = readable_axis_bounds(low, high)
    assert floor <= low and ceiling >= high
    assert floor / step == pytest.approx(round(floor / step))
    assert ceiling / step == pytest.approx(round(ceiling / step))


def test_continuous_qualifier_belongs_to_only_its_own_metric():
    obs = [observation(f"{prefix}{i}", name, v, f"FY{2021+i}", unit="count")
           for prefix, name, values in [("a", "Units shipped", (120, 70, 90)), ("b", "Service hours", (20, 40, 60))]
           for i, v in enumerate(values)]
    slide = PresentationSlide(id="s", slide_type="analysis", title="Review",
        message="Units shipped declined overall and Service hours steadily increased.")
    assert not ClaimValidator().validate_slide(slide, obs)


def test_mid_series_rebound_is_not_claimed_to_be_in_final_year():
    obs = [observation(str(i), "Reservoir volume", v, f"FY{2020+i}", unit="count") for i, v in enumerate((100, 60, 80, 50))]
    slide = PresentationSlide(id="s", slide_type="analysis", title="Review", message="Reservoir volume increased.", observation_ids=[o.id for o in obs])
    plan = PresentationPlan(title="Review", slides=[slide])
    repair_presentation_plan(plan, obs)
    assert slide.message == "Reservoir volume decreased."


def test_wrong_metric_endpoint_enters_existing_bounded_model_repair():
    from adaptive_document_agent.agent.presentation_planner import PresentationPlanner
    from tests.test_p1_theme_planning import gateway_for, themed_result
    result = themed_result()
    wrong = result.presentation_plan.model_copy(deep=True)
    wrong.slides[3].message = "Units shipped increased from 60 to 80."
    gateway, client = gateway_for(wrong.model_dump(mode="json"), result.presentation_plan.model_dump(mode="json"))
    fixed = PresentationPlanner(gateway).plan(result)
    assert len(client.calls) == 2
    assert "not its own retained values" in client.calls[1][2]["content"]
    assert fixed.planning_origin == "repaired"
    assert fixed.slides[3].message == result.presentation_plan.slides[3].message
