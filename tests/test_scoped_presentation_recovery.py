"""A rejected topic cannot restart expensive planning or borrow unrelated facts."""

import json
import io

import pytest

from adaptive_document_agent.agent.presentation_insight_recovery import recover_insight_narrative
from adaptive_document_agent.agent.presentation_plan_recovery import PresentationPlanRecovery
from adaptive_document_agent.agent.presentation_planner import PresentationPlanner
from adaptive_document_agent.agent.presentation_topic_selector import PresentationTopicSelector, series_directory
from adaptive_document_agent.models import AnalysisResult, AnalysisTask, PresentationTopic, PresentationTopicSelection, ValidationIssue
from adaptive_document_agent.services.llm import LLMGateway, LLMSettings, MockLLMClient, PrivacyMode, ProviderName
from adaptive_document_agent.services.presentation_evidence import calculation_catalog
from adaptive_document_agent.validation.presentation_plan_validator import PresentationPlanValidator
from tests.test_pptx_export import _result


def _topic(identifier, series_id, **changes):
    return PresentationTopic(id=identifier, title="Revenue movement", question="How did revenue move?",
        rationale="Comparable annual revenue supports this question.", series_ids=[series_id]).model_copy(update=changes)


def _linked_result():
    result = _result()
    result.analysis_results = [AnalysisResult(task_id="growth", title="Revenue growth", result=95.39256623911618,
        input_observation_ids=[o.id for o in result.observations], evidence=result.insights[0].evidence, confidence=.9)]
    result.analysis_plan = [AnalysisTask(id="growth", title="Revenue growth", description="Annual movement",
        analysis_type="percentage_change", tool_name="percentage_change", reason="Reported annual series",
        expected_output="Growth")]
    insight = result.insights[0]
    insight.result_ids = ["growth"]
    insight.narrative = "Revenue increased from FY2023 to FY2025."
    result.presentation_plan = PresentationPlanRecovery().fallback(result)
    return result


def test_partial_topic_repair_is_scoped_and_preserves_valid_sibling_and_raw_draft():
    result = _result()
    sid = series_directory(result)[0][0]["id"]
    good, bad = _topic("valid", sid), _topic("invalid", sid, title="Revenue reached 999999")

    class Gateway:
        calls = []

        def generate_structured(self, messages, response_model, **kwargs):
            self.calls.append(messages)
            if len(self.calls) == 1:
                return response_model(topics=[good, bad])
            return response_model(topics=[_topic("invalid", sid)])

    gateway = Gateway()
    selection = PresentationTopicSelector(gateway).select(result)
    assert len(gateway.calls) == 2
    assert selection.topics[0] == good
    assert [t.id for t in selection.topics] == ["valid", "invalid"]
    repair = gateway.calls[1][-1]["content"]
    assert "all_extracted_series" not in repair and '"id":"valid"' not in repair
    assert sid in repair and "validation_error" in repair
    audit = json.loads(next(w.message for w in result.validation_warnings if w.code == "presentation_topic_validation"))
    assert audit["topic"] == bad.model_dump(mode="json")
    assert "999999" in audit["error"]


def test_repeated_bad_takeaway_retains_original_question_with_review_status():
    result = _result()
    sid = series_directory(result)[0][0]["id"]
    topic = _topic("growth", sid, takeaway="Revenue increased 999999%")

    class Gateway:
        calls = 0

        def generate_structured(self, messages, response_model, **kwargs):
            self.calls += 1
            return response_model(topics=[topic])

    gateway = Gateway()
    result.presentation_topics = PresentationTopicSelector(gateway).select(result)
    assert gateway.calls == 2
    assert result.presentation_topics.topics[0].question == topic.question
    assert result.presentation_topics.topics[0].takeaway == ""
    assert topic.takeaway == "Revenue increased 999999%"  # original draft is unchanged
    plan = PresentationPlanner(gateway).plan(result)
    assert gateway.calls == 2
    assert plan.editorial_status == "needs_review"
    assert any("unsupported optional claims" in note for note in plan.editorial_notes)


def test_invalid_unique_topic_is_audited_and_does_not_erase_valid_sibling():
    result = _result()
    sid = series_directory(result)[0][0]["id"]
    good, bad = _topic("valid", sid), _topic("bad", sid, title="Revenue reached 999999")
    client = MockLLMClient([PresentationTopicSelection(topics=[good, bad]).model_dump(),
                            PresentationTopicSelection(topics=[bad]).model_dump()])
    gateway = LLMGateway(client, LLMSettings(provider=ProviderName.MOCK, model="mock"))
    result.presentation_topics = PresentationTopicSelector(gateway).select(result)
    assert [t.id for t in result.presentation_topics.topics] == ["valid"]
    plan = PresentationPlanner(gateway).plan(result)
    assert len(client.calls) == 2
    assert plan.editorial_status == "needs_review"
    assert any("bad" in note and "unsupported" in note for note in plan.editorial_notes)
    drafts = [json.loads(w.message)["topic"] for w in result.validation_warnings if w.code == "presentation_topic_validation"]
    assert len(drafts) == 2 and all(d == bad.model_dump(mode="json") for d in drafts)


def test_local_only_correction_stays_on_same_gateway_and_preserves_valid_topics():
    result = _result()
    sid = series_directory(result)[0][0]["id"]
    good, bad = _topic("valid", sid), _topic("bad", sid, title="Revenue reached 999999")
    client = MockLLMClient([PresentationTopicSelection(topics=[good, bad]).model_dump(),
                            PresentationTopicSelection(topics=[bad]).model_dump()])
    gateway = LLMGateway(client, LLMSettings(provider=ProviderName.OLLAMA, model="local",
        base_url="http://127.0.0.1:11434", privacy_mode=PrivacyMode.LOCAL_ONLY))
    selection = PresentationTopicSelector(gateway).select(result)
    assert len(client.calls) == 2
    assert [t.id for t in selection.topics] == ["valid"]
    assert all(row["provider"] == "mock" for row in gateway.usage)
    assert gateway.settings.privacy_mode == PrivacyMode.LOCAL_ONLY


def test_failed_topic_selection_reuses_findings_without_full_plan_model_calls():
    result = _linked_result()
    result.validation_warnings.append(ValidationIssue(code="presentation_topic_selection_failed",
        message="Invalid numerical takeaway", stage="presentation"))

    class Gateway:
        def generate_structured(self, *args, **kwargs):
            pytest.fail("A failed topic pass must not start another full-plan model pipeline")

    plan = PresentationPlanner(Gateway()).plan(result)
    page = next(s for s in plan.slides if s.slide_type == "analysis")
    assert page.message == result.insights[0].narrative
    assert page.insight_ids == [result.insights[0].id]
    assert plan.themes
    PresentationPlanValidator().validate(plan, result)


def test_offline_recovery_preserves_raw_records_and_is_idempotent():
    result = _linked_result()
    raw = result.model_dump(include={"observations", "insights", "charts", "llm_usage"})
    changed = recover_insight_narrative(result)
    assert changed
    assert recover_insight_narrative(result) == []
    assert result.model_dump(include=set(raw)) == raw
    audit = next(w for w in result.validation_warnings if w.code == "presentation_insight_recovery")
    assert "Evidence-backed comparison" in audit.message


def test_recovery_cannot_borrow_same_metric_or_same_page_with_different_inputs():
    result = _linked_result()
    copies = [o.model_copy(update={"id": "other-" + o.id}, deep=True) for o in result.observations]
    result.observations.extend(copies)
    result.analysis_results[0].input_observation_ids = [o.id for o in copies]
    before = result.presentation_plan.model_dump()
    assert recover_insight_narrative(result) == []
    assert result.presentation_plan.model_dump() == before


def test_recovery_rejects_unsupported_numbers_in_retained_insight():
    result = _linked_result()
    result.insights[0].narrative = "Revenue increased 999999% from FY2023 to FY2025."
    assert recover_insight_narrative(result) == []


@pytest.mark.parametrize("currency,scale", [("USD", "million"), ("CNY", "billion")])
def test_recovery_rejects_currency_or_scale_borrowed_from_matching_bare_numbers(currency, scale):
    result = _linked_result()
    result.insights[0].narrative = (f"Revenue increased from {currency} 267.025 {scale} in FY2023 "
                                  f"to {currency} 521.747 {scale} in FY2025.")
    assert recover_insight_narrative(result) == []


def test_invalid_currency_copy_can_fall_back_to_valid_original_movement():
    result = _linked_result()
    result.insights[0].narrative = "Revenue increased from USD 267.025 million to USD 521.747 million."
    result.insights[0].movement = "Increased from FY2023 to FY2025"
    assert recover_insight_narrative(result)
    slide = next(s for s in result.presentation_plan.slides if s.slide_type == "analysis")
    assert "USD" not in slide.message and all("USD" not in b for b in slide.bullets)


@pytest.mark.parametrize("copy", [
    "Revenue increased by 267025 percent from FY2023 to FY2025.",
    "Revenue increased by 95.4 percentage points from FY2023 to FY2025.",
    "Revenue margin was 95.4% in FY2025.",
    "Revenue was 0 in FY2025.",
    "Revenue reached 95.4 million in FY2025.",
])
def test_recovery_does_not_lend_values_across_quantity_types(copy):
    result = _linked_result()
    result.insights[0].narrative = copy
    assert recover_insight_narrative(result) == []


def test_recovery_requires_every_original_task_input():
    result = _linked_result()
    result.analysis_results[0].input_observation_ids.append("missing-input")
    assert recover_insight_narrative(result) == []


def test_validated_movement_is_visible_but_rejected_original_prose_stays_in_notes():
    from pptx import Presentation
    from adaptive_document_agent.services.pptx_export import build_presentation

    result = _linked_result()
    original = "Revenue increased 999999% from FY2023 to FY2025."
    result.insights[0].narrative = original
    result.insights[0].movement = "Increased from FY2023 to FY2025"
    changed = recover_insight_narrative(result)
    assert changed
    recovered = next(s for s in result.presentation_plan.slides if s.id in changed)
    assert recovered.bullets == [recovered.message]
    deck = Presentation(io.BytesIO(build_presentation(result)))
    page = next(s for s in deck.slides if s.name.startswith("composed_")
                and recovered.id in s.notes_slide.notes_text_frame.text)
    visible = "\n".join(shape.text for shape in page.shapes if shape.has_text_frame)
    assert "999999" not in visible
    assert "Increased from FY2023 to FY2025" in visible
    assert original in page.notes_slide.notes_text_frame.text
    assert result.insights[0].narrative == original


def test_recovered_plan_skips_repeated_insight_validation(monkeypatch):
    result = _linked_result()
    assert recover_insight_narrative(result)

    def unexpected(*args, **kwargs):
        pytest.fail("An already recovered plan must not repeat insight calculations on every export")

    monkeypatch.setattr("adaptive_document_agent.agent.presentation_insight_recovery._linked_findings", unexpected)
    assert recover_insight_narrative(result) == []


def test_bound_negative_delta_accepts_positive_magnitude_but_rejects_wrong_direction():
    result = _result()
    for o, value in zip(result.observations, (30, 26, 25)):
        o.metric_original = "Service margin"
        o.value = value
        o.raw_value = f"{value}%"
        o.unit = "percent"
        o.currency = None
        o.raw_unit = "%"
        o.unit_scale = 1
    result.insights = []
    plan = PresentationPlanRecovery().fallback(result)
    slide = next(s for s in plan.slides if s.slide_type == "analysis")
    calc = next(c for c in calculation_catalog(result.observations).values() if c["kind"] == "absolute_change")
    slide.title = "Service margin movement"
    slide.message = "Service margin declined by 5 percentage points from FY2023 to FY2025."
    slide.calculation_ids = [calc["id"]]
    PresentationPlanValidator().validate(plan, result)
    slide.message = "Service margin increased by 5 percentage points from FY2023 to FY2025."
    with pytest.raises(ValueError, match="direction|increased|decreased"):
        PresentationPlanValidator().validate(plan, result)
