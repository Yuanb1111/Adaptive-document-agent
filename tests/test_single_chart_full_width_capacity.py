"""A dense single composition must use the available page before failing."""

from io import BytesIO
import json

import pytest
from pptx import Presentation

from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import (
    ChartPlan, PresentationSlide, PresentationTheme, PresentationVisualBlock,
)
from adaptive_document_agent.services.chart_space import (
    chart_space_requirements, composition_plot_height,
)
from adaptive_document_agent.services.slide_compositor import (
    _ChartCapacityError, _render_composed_slide, render_composed_slide,
    validate_composed_geometry,
)
from tests.test_chart_capacity_pagination import _deck
from tests.test_p0_composition import observation, result_for


CATEGORIES = (
    "Long term regional equipment maintenance services",
    "Specialized industrial component repair services",
    "Remote equipment inspection and monitoring services",
    "Regional installation and commissioning services",
    "Technical training and operational support services",
    "Replacement component supply and logistics services",
    "Scheduled workshop inspection and testing services",
    "Other separately disclosed equipment support services",
)
PERIODS = ("FY2021", "FY2022", "FY2023")
SHARES = (
    (21, 19, 16, 14, 11, 9, 6, 4),
    (20, 18, 17, 15, 11, 9, 6, 4),
    (19, 17, 18, 16, 11, 9, 6, 4),
)
COMMENTARY = (
    "The source presents each service category separately throughout the disclosed "
    "reporting periods. The chart retains the original category names so that the "
    "reported composition remains traceable to the underlying source records.",
    "The comparison describes the reported service portfolio. Categories retain "
    "their individual definitions, and the accompanying explanation remains visible "
    "when the chart requires a separate physical page for readable labels.",
)
CAVEATS = (
    "Source scope: the reported service categories cover the disclosed regional "
    "operations and exclude separately reported overseas operations.",
)


def _case(*, with_support=False):
    records = [
        observation(f"share-{year}-{number}", "Service portfolio share", value,
                    period, unit="percent", category=category)
        for year, (period, shares) in enumerate(zip(PERIODS, SHARES))
        for number, (category, value) in enumerate(zip(CATEGORIES, shares))
    ]
    chart = ChartPlan(
        id="service-composition", title="Service portfolio share",
        question="How does the reported service composition compare?",
        chart_type="stacked_percent", series_dimension="segment",
        observation_ids=[record.id for record in records], source_pages=[3],
    )
    support = [
        observation(f"support-{year}", "Service locations", value, f"FY{year}", unit="count")
        for year, value in ((2021, 32), (2022, 35), (2023, 37), (2024, 41))
    ] if with_support else []
    records.extend(support)
    for record in records:
        record.validation_status = "valid"
    blocks = [PresentationVisualBlock(role="hero", chart_ids=[chart.id])]
    if support:
        blocks.append(PresentationVisualBlock(
            role="table", observation_ids=[record.id for record in support]))
    plan = PresentationSlide(
        id="portfolio", slide_type="analysis", theme_id="portfolio",
        title="Service portfolio composition", section_title="Service portfolio composition",
        message="Reported category shares across the disclosed reporting periods.",
        bullets=list(COMMENTARY), layout="single", chart_ids=[chart.id],
        source_pages=[3], visual_blocks=blocks,
    )
    result = result_for(records, [chart], plan)
    result.presentation_plan.themes = [PresentationTheme(
        id="portfolio", title=plan.title, question=chart.question,
        rationale="Compare complete source category shares across periods",
        chart_ids=[chart.id], caveats=list(CAVEATS),
    )]
    return result, plan, chart, support


def _visible_text(slides):
    return " ".join(
        shape.text for slide in slides for shape in slide.shapes if shape.has_text_frame
    ).split()


@pytest.mark.parametrize("with_support", [False, True])
def test_dense_single_chart_uses_full_width_and_preserves_complete_evidence(with_support):
    result, plan, chart, support = _case(with_support=with_support)
    before = result.model_dump()
    index = DocumentIndex(result.observations)
    # This is a real legend/plot failure after the existing compact-header retry,
    # rather than a mocked exception that could hide the missing geometry path.
    with pytest.raises(_ChartCapacityError, match="plot area is too small"):
        _render_composed_slide(_deck(), plan, [chart], result, index, compact_header=True)

    deck = _deck()
    sentinel = deck.slides.add_slide(deck.slide_layouts[0])
    pages = render_composed_slide(deck, plan, [chart], result, index)
    assert deck.slides[0] is sentinel
    assert len(deck.slides) == len(pages) + 1
    native = [shape for page in pages for shape in page.shapes if shape.has_chart]
    assert len(native) == 1
    shape = native[0]
    assert shape.name == f"chart:{chart.id}"
    assert shape.width.inches > deck.slide_width.inches * .8
    values = [index.get(oid) for oid in chart.observation_ids]
    _, minimum_height, categories = chart_space_requirements(chart, values, [])
    assert composition_plot_height(
        (0, 0, shape.width.inches, shape.height.inches), categories, compact=False
    ) >= minimum_height - 1e-6
    expected_values = {
        category: [shares[number] / 100 for shares in SHARES]
        for number, category in enumerate(CATEGORIES)
    }
    assert len(shape.chart.series) == len(expected_values)
    assert {series.name for series in shape.chart.series} == set(expected_values)
    assert [category.label for category in shape.chart.plots[0].categories] == list(PERIODS)
    for series in shape.chart.series:
        assert list(series.values) == pytest.approx(expected_values[series.name])

    visible = " ".join(_visible_text(pages))
    for paragraph in COMMENTARY:
        assert " ".join(paragraph.split()) in visible
    chart_page = next(page for page in pages if any(item.has_chart for item in page.shapes))
    unit = next(item for item in chart_page.shapes
                if item.has_text_frame and item.text == "Share (%)")
    annotation = next(item for item in chart_page.shapes
                      if item.name == f"annotation:{chart.id}")
    assert unit.top == annotation.top
    assert unit.left + unit.width <= annotation.left
    assert annotation.left + annotation.width <= shape.left + shape.width
    assert max(unit.top + unit.height, annotation.top + annotation.height) <= shape.top
    for caveat in CAVEATS:
        assert " ".join(caveat.split()) in " ".join(_visible_text([chart_page]))
    assert all("Source:" in " ".join(_visible_text([page])) for page in pages)
    table_ids = [
        oid for page in pages for item in page.shapes if item.has_table
        for oid in item.name.removeprefix("table:").split(",")
    ]
    assert sorted(table_ids) == sorted(record.id for record in support)
    if support:
        assert len(pages) > 1
        cells = " ".join(
            cell.text for page in pages for item in page.shapes if item.has_table
            for row in item.table.rows for cell in row.cells
        )
        assert all(str(int(record.value)) in cells and record.period in cells for record in support)
    notes = json.loads(chart_page.notes_slide.notes_text_frame.text)
    assert {item["id"] for item in notes["observations"]} == {item.id for item in result.observations}
    assert result.model_dump() == before
    validate_composed_geometry(deck)
    # Saving and reopening also exercises relationship cleanup after failed
    # partial-page attempts; only the accepted chart may survive in the deck.
    output = BytesIO()
    deck.save(output)
    reopened = Presentation(BytesIO(output.getvalue()))
    assert len(reopened.slides) == len(pages) + 1
    assert sum(item.has_chart for page in reopened.slides for item in page.shapes) == 1


def test_truly_unfit_single_chart_rolls_back_every_retry_without_losing_prior_pages():
    result, plan, chart, _ = _case()
    # Even a full-width page cannot fit this amount of mandatory visible source
    # context. The renderer must fail honestly instead of dropping qualifications.
    result.presentation_plan.themes[0].caveats = [
        f"Source qualification {number}: the reported categories retain separate "
        "regional definitions and source coverage that must remain visible beside "
        "the underlying comparison for this particular reporting scope."
        for number in range(15)
    ]
    before = result.model_dump()
    deck = _deck()
    sentinel = deck.slides.add_slide(deck.slide_layouts[0])
    original_shape_count = len(sentinel.shapes)
    with pytest.raises(_ChartCapacityError):
        render_composed_slide(deck, plan, [chart], result, DocumentIndex(result.observations))
    assert len(deck.slides) == 1 and deck.slides[0] is sentinel
    assert len(sentinel.shapes) == original_shape_count
    assert not any(shape.has_chart for shape in sentinel.shapes)
    assert result.model_dump() == before
