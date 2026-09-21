"""Single-metric slides show supported calculations and editable evidence."""

import io

import pytest

pytestmark = pytest.mark.usefixtures("local_render_stub")
from pptx import Presentation

from adaptive_document_agent.models import (
    ChartPlan, DocumentProfile, Observation, ParsedDocument, PipelineResult,
    PresentationPlan, PresentationSlide, PresentationVisualBlock, SourceEvidence,
)
from adaptive_document_agent.services.export import export_pptx
from adaptive_document_agent.services.qa_reporter import CriticalQAError, run_comprehensive_qa
from adaptive_document_agent.services.single_metric_analysis import single_metric_analysis
from adaptive_document_agent.services.single_metric_slides import enrich_single_metric_slides


def series(values=(100, 120, 144), periods=("FY2022", "FY2023", "FY2024"), metric="Revenue", unit="currency"):
    return [Observation(
        id=f"v{i}", metric_original=metric, value=v, raw_value=str(v), period=p,
        unit=unit, unit_family="percentage" if unit == "percent" else unit,
        currency="USD" if unit == "currency" else None, confidence=0.95,
        evidence=[SourceEvidence(page=3, text=str(v), row_label=metric, extraction_method="digital_table", confidence=0.95)],
    ) for i, (v, p) in enumerate(zip(values, periods))]


def result_for(observations=None):
    obs = observations if observations is not None else series()
    return PipelineResult(
        document=ParsedDocument(document_id="generic", sha256="a" * 64, safe_filename="generic.pdf", page_count=5),
        profile=DocumentProfile(document_purpose="Review the reported metric"), observations=obs,
        presentation_plan=PresentationPlan(title="Metric review", slides=[
            PresentationSlide(id="cover", slide_type="cover", title="Metric review"),
            PresentationSlide(id="company", slide_type="company_overview", title="Document overview"),
            PresentationSlide(id="summary", slide_type="executive_summary", title="Executive summary"),
            PresentationSlide(id="analysis", slide_type="analysis", title=f"{obs[0].metric_original} performance",
                message="The chart shows the reported movement across comparable periods.", section_title="Performance",
                layout="data_overview", observation_ids=[o.id for o in obs], source_pages=[3]),
            PresentationSlide(id="quality", slide_type="data_quality", title="Data quality"),
            PresentationSlide(id="appendix", slide_type="appendix", title="Source data"),
        ]),
    )


def test_annual_statistics_use_elapsed_years_and_keep_provenance():
    analysis = single_metric_analysis(series())
    assert analysis.absolute_change == 44
    assert analysis.percentage_change == pytest.approx(44)
    assert analysis.cagr == pytest.approx(20)
    assert [c.percentage_change for c in analysis.changes] == pytest.approx([20, 20])
    assert all(c.is_yoy for c in analysis.changes)
    assert analysis.peak.id == "v2" and analysis.trough.id == "v0"
    assert all(o.evidence[0].page == 3 for o in analysis.observations)
    skipped_year = single_metric_analysis(series((100, 144), ("FY2022", "FY2024")))
    assert skipped_year.cagr == pytest.approx(20)
    assert not skipped_year.changes[0].is_yoy


@pytest.mark.parametrize("values", [(0, 120, 144), (-100, -80, -50), (100, -10, 120)])
def test_nonpositive_values_never_produce_cagr(values):
    analysis = single_metric_analysis(series(values))
    assert analysis.cagr is None
    if values[0] <= 0:
        assert analysis.percentage_change is None


def test_interim_yoy_is_not_annual_cagr_and_quarters_are_not_yoy():
    interim = single_metric_analysis(series(periods=("6M2022", "6M2023", "6M2024")))
    assert interim.cagr is None
    assert all(c.is_yoy for c in interim.changes)
    quarters = single_metric_analysis(series(periods=("Q12024", "Q22024", "Q32024")))
    assert quarters is not None and quarters.cagr is None
    assert not any(c.is_yoy for c in quarters.changes)


@pytest.mark.parametrize("mutation", ["period", "currency", "unit", "missing", "evidence", "conflict"])
def test_incompatible_or_incomplete_inputs_do_not_generate_statistics(mutation):
    obs = series()
    if mutation == "period":
        obs[-1].period = "6M2024"
    elif mutation == "currency":
        obs[-1].currency = "EUR"
    elif mutation == "unit":
        obs[-1].unit = "count"
    elif mutation == "missing":
        obs[-1].value = None
    elif mutation == "evidence":
        obs[-1].evidence = []
    else:
        obs.append(obs[-1].model_copy(update={"id": "conflict", "value": 999}))
    assert single_metric_analysis(obs) is None


def test_margin_changes_use_percentage_points_and_turning_point():
    analysis = single_metric_analysis(series((20, 30, 25), metric="Gross profit margin", unit="percent"))
    assert analysis.is_percentage
    assert analysis.absolute_change == 5
    assert analysis.cagr is None and analysis.percentage_change is None
    assert analysis.turning_periods == ["FY2023"]


@pytest.mark.parametrize("periods", [("Prior", "Current", "Forecast"), ("FY2022", "Q12023", "FY2024")])
def test_uncertain_or_mixed_period_durations_do_not_generate_statistics(periods):
    assert single_metric_analysis(series(periods=periods)) is None


def test_overflowing_changes_are_not_exported_as_derived_values():
    assert single_metric_analysis(series((-1e308, 0, 1e308))) is None


def test_enrichment_is_idempotent_and_preserves_evidence_and_company():
    result = result_for()
    original = [o.model_dump() for o in result.observations]
    company = result.presentation_plan.company.model_dump()
    assert enrich_single_metric_slides(result) == ["analysis"]
    assert enrich_single_metric_slides(result) == []
    assert len(result.charts) == 1
    assert result.presentation_plan.slides[3].layout == "single_metric_hero"
    assert result.charts[0].observation_ids == [o.id for o in result.observations]
    assert [o.model_dump() for o in result.observations] == original
    assert result.presentation_plan.company.model_dump() == company


def test_single_metric_table_exports_as_large_editable_hero_chart():
    result = result_for()
    deck = Presentation(io.BytesIO(export_pptx(result)))
    hero = next(s for s in deck.slides if s.name == "single_metric_hero")
    text = "\n".join(s.text for s in hero.shapes if s.has_text_frame)
    for expected in ["Start value", "End value", "Absolute change", "CAGR", "+20.0%", "FY2023 YoY", "Peak:", "Source:"]:
        assert expected in text
    chart = next(s for s in hero.shapes if s.has_chart)
    assert chart.width.inches > 10
    assert chart.height.inches > 2
    assert list(chart.chart.series[0].values) == [100, 120, 144]
    assert not any(s.has_table for s in hero.shapes)
    assert "731876%" not in text
    # Explicit layout bands must not overlap the native chart.
    for shape in hero.shapes:
        if shape.has_text_frame and shape.text.strip() in {"Start value", "End value", "CAGR"}:
            assert shape.top > chart.top + chart.height


def test_existing_valid_chart_keeps_identity_when_enriched():
    result = result_for()
    chart = ChartPlan(id="existing", title="Revenue", question="Revenue over time", chart_type="bar",
        observation_ids=[o.id for o in result.observations], source_pages=[3])
    result.charts = [chart]
    result.presentation_plan.slides[3].chart_ids = [chart.id]
    assert enrich_single_metric_slides(result) == ["analysis"]
    assert result.charts == [chart]


def test_non_financial_count_hero_preserves_display_scale():
    result = result_for(series((10_000, 12_000, 14_400), metric="Units sold", unit="count"))
    deck = Presentation(io.BytesIO(export_pptx(result)))
    hero = next(s for s in deck.slides if s.name == "single_metric_hero")
    text = "\n".join(s.text for s in hero.shapes if s.has_text_frame)
    assert "thousands of units" in text
    assert "RMB" not in text


@pytest.mark.parametrize("block", [False, True])
def test_chart_title_or_block_data_mismatch_blocks_export(block):
    result = result_for()
    slide = result.presentation_plan.slides[3]
    chart = ChartPlan(id="bad", title="Cash and cash equivalents", question="Cash", chart_type="bar",
        observation_ids=[o.id for o in result.observations], source_pages=[3])
    result.charts = [chart]
    if block:
        slide.visual_blocks = [PresentationVisualBlock(role="hero", chart_ids=[chart.id])]
    else:
        slide.chart_ids = [chart.id]
    qa = run_comprehensive_qa(result, auto_repair=False)
    assert any(i.code == "chart_title_data_mismatch" for i in qa.critical_errors)
    with pytest.raises(CriticalQAError):
        export_pptx(result)


def test_duplicate_slides_are_reported_and_hero_not_fabricated_for_wrong_title():
    result = result_for()
    slide = result.presentation_plan.slides[3]
    result.presentation_plan.slides.insert(4, slide.model_copy(update={"id": "duplicate"}, deep=True))
    qa = run_comprehensive_qa(result, auto_repair=False)
    assert any(i.code == "redundant_analysis_slide" for i in qa.warnings)
    slide.title = "Cash and cash equivalents"
    assert "analysis" not in enrich_single_metric_slides(result)


@pytest.mark.parametrize("metric,value", [("Revenue %", 10_350_986), ("Income tax expense %", 731_876)])
def test_implausible_monetary_percentages_block_export(metric, value):
    result = result_for(series((value, value + 100, value + 200), metric=metric, unit="percent"))
    for obs in result.observations:
        obs.evidence[0].row_label = metric.removesuffix(" %")
    qa = run_comprehensive_qa(result, auto_repair=False)
    assert any(i.code == "monetary_percentage_mismatch" for i in qa.critical_errors)
    assert any(i.code == "implausible_percentage" for i in qa.critical_errors)
    with pytest.raises(CriticalQAError):
        export_pptx(result)


def test_genuine_expense_ratio_passes_export_qa():
    result = result_for(series((12, 11, 10), metric="Selling expense ratio", unit="percent"))
    assert not run_comprehensive_qa(result).critical_errors
