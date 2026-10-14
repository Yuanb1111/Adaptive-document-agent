from io import BytesIO

import pytest
from pptx import Presentation

from adaptive_document_agent.models.customization import ReportRequirement, ReportRequirements
from adaptive_document_agent.services.slide_budgets import (
    introduction_limit, resolved_limits, require_exported_budgets)
from adaptive_document_agent.services.customization_checks import check_requirements
from adaptive_document_agent.services.brief_table_evidence import _quoted_complete_row
from adaptive_document_agent.services.source_quotes import normalize_quote
from tests.test_compact_summary_v111 import fixture


def requirements(intro=5):
    return ReportRequirements(items=[
        ReportRequirement(id='body', kind='slide_limit', request_quote='Body 20',
                          description='Body maximum', max_slides=20, slide_scope='body'),
        ReportRequirement(id='intro', kind='slide_limit', request_quote=f'Introduction {intro}',
                          description='Introduction maximum', max_slides=intro,
                          slide_section='introduction', slide_scope='body')])


@pytest.mark.parametrize('limit', [1, 2, 5])
def test_explicit_section_budget_does_not_shrink_the_report_budget(limit):
    original = requirements(limit)
    restored = ReportRequirements.model_validate_json(original.model_dump_json())
    assert resolved_limits(restored) == [20]
    assert introduction_limit(restored) == limit


def test_physical_section_and_report_limits_survive_roundtrip():
    import json
    from adaptive_document_agent.services.summary_export import _PREFIX
    result, _, _ = fixture(1)
    result.profile.report_requirements = requirements()
    deck = Presentation()
    for n in range(25):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        if n < 5 or n >= 20:
            shape = slide.shapes.add_textbox(0, 0, 100, 100)
            shape.element.xpath('.//p:cNvPr')[0].set('descr', _PREFIX + json.dumps({'page_id': str(n)}))
        if n >= 20:
            slide.element.cSld.set('name', 'ada:appendix')
    stream = BytesIO(); deck.save(stream)
    deck = Presentation(stream)
    require_exported_budgets(deck, result)
    check_requirements(result, deck)
    checks = {c.requirement_id: c for c in result.customization_report}
    assert checks['body'].status == checks['intro'].status == 'satisfied'
    assert len(checks['body'].slide_numbers) == 20
    assert len(checks['intro'].slide_numbers) == 5
    result.profile.report_requirements.items[1].max_slides = 4
    with pytest.raises(ValueError, match='introduction body has 5'):
        require_exported_budgets(deck, result)
    from adaptive_document_agent.services.customization_checks import mark_customization_export_verified
    with pytest.raises(ValueError, match='introduction body has 5'):
        mark_customization_export_verified(result, stream.getvalue())
    assert next(c for c in result.customization_report if c.requirement_id == 'intro').status == 'not_met'
    result.profile.report_requirements.items[1].max_slides = 5
    deck.slides[20].element.cSld.set('name', '')
    with pytest.raises(ValueError, match='report body has 21'):
        require_exported_budgets(deck, result)


def test_single_table_leader_preserves_all_cells_and_accounting_signs():
    cells = ['Provision (charge)/credit', '(12)', '(9.4)', '3', '0.1']
    assert _quoted_complete_row(cells, normalize_quote('Provision (charge)/credit . (12) (9.4) 3 0.1'))
    assert not _quoted_complete_row(cells, normalize_quote('Provision (charge)/credit . 12 (9.4) 3 0.1'))
    assert not _quoted_complete_row(cells, normalize_quote('Provision (charge)/credit . omitted words (12) (9.4) 3 0.1'))
    assert not _quoted_complete_row(cells, normalize_quote('Provision (charge)/credit . (12) (9.4) 3'))


@pytest.mark.parametrize('repair', [False, True])
def test_coverage_budget_stays_trusted_and_full_source_survives_targeted_repair(repair):
    from adaptive_document_agent.agent.topic_coverage_review import _review_batch
    from adaptive_document_agent.agent.report_requirements import instruction_messages
    from adaptive_document_agent.agent.presentation_topic_selector import PresentationTopicSelector
    from adaptive_document_agent.agent.coverage_review_context import compact_context, batch_context
    from adaptive_document_agent.models import PresentationTopicSelection
    from tests.test_bounded_topic_coverage import context_fixture
    from tests.test_topic_coverage_review import _topic, _review
    result, lookup, annual, interim, adverse, payload = context_fixture()
    result.profile.report_requirements = requirements()
    current = PresentationTopicSelection(topics=[_topic(annual)])
    instructions = instruction_messages(result.profile, purpose='coverage review')
    class Client:
        calls = []
        def generate_structured(self, messages, model, **kwargs):
            self.calls.append(messages)
            assert messages[1] == instructions[0]
            assert messages[1]['role'] == 'user'
            assert '20' in messages[1]['content'] and 'introduction' in messages[1]['content']
            assert '<UNTRUSTED_DOCUMENT_CONTENT>' in messages[2]['content']
            assert 'all_extracted_series' in messages[2]['content']
            answer = _review(annual, interim, adverse)
            if repair and len(self.calls) == 1:
                answer.topics = current.topics
            return answer
    client = Client()
    _review_batch(batch_context(compact_context(payload), current, [interim, adverse]),
                  current, [interim, adverse], lookup, set(), client,
                  PresentationTopicSelector._validate, {}, instructions=instructions)
    assert len(client.calls) == (2 if repair else 1)
