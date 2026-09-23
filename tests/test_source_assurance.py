from adaptive_document_agent.extraction.borderless_layout import SourceLine, geometric_audit_statuses
from adaptive_document_agent.extraction.observation_extractor import ObservationExtractor
from adaptive_document_agent.document_model.period_semantic_validator import format_period_label
from adaptive_document_agent.models.table import ExtractedTable, TableRow

import io
import pytest
from pptx import Presentation
from adaptive_document_agent.models import CompanyFact, CompanyProfile, PresentationSlide, PresentationVisualBlock
from adaptive_document_agent.services.presentation_evidence import build_evidence_catalog
from adaptive_document_agent.services.presentation_editorial import review_presentation
from adaptive_document_agent.validation.claim_validator import ClaimValidator
from tests.test_p0_composition import observation, paired_result
from tests.test_presentation_brief import blank_deck, visible


def test_assurance_is_column_evidence_not_interim_date():
    years = SourceLine("2020 2021 2022", [(0, 4, 90, 110), (5, 9, 190, 210), (10, 14, 290, 310)])
    label = SourceLine("(unaudited)", [(0, 11, 180, 220)])
    assert geometric_audit_statuses([label], years, 6) == ["unknown", "unknown", "unaudited", "unaudited", "unknown", "unknown"]
    assert geometric_audit_statuses([SourceLine("(unaudited)")], years, 3) == ["unknown"] * 3
    assert format_period_label("2022-06-30", is_balance_sheet=True) == "30 Jun 2022"
    assert format_period_label("2022-06-30", is_balance_sheet=True, is_unaudited=True) == "30 Jun 2022*"
    assert format_period_label("6M2022", is_unaudited=True) == "6M2022*"


def test_operators_are_not_metric_parents_and_unknown_is_not_audited():
    table = ExtractedTable(table_id="sample", page=2, headers=["Metric", "2020", "2021"],
        column_periods=[None, "6M2020", "6M2021"],
        column_audit_statuses=["unknown", "unaudited", "unknown"],
        rows=[TableRow(cells=["Add:", None, None], page=2),
              TableRow(cells=["Adjusted net loss", "(8)", "(6)"], page=2)])
    values = ObservationExtractor()._table_observations(table)
    assert [o.audited_status for o in values] == ["unaudited", "unknown"]
    assert all(o.parent_section is None and "section" not in o.dimensions for o in values)
    assert [o.raw_value for o in values] == ["(8)", "(6)"]


def test_conflicting_or_unscoped_assurance_is_not_assigned():
    years = SourceLine("2020 2021", [(0, 4, 90, 110), (5, 9, 190, 210)])
    labels = [SourceLine("audited", [(0, 7, 90, 110)]), SourceLine("unaudited", [(0, 9, 90, 110)])]
    assert geometric_audit_statuses(labels, years, 2) == ["unknown", "unknown"]
    assert geometric_audit_statuses([SourceLine("unaudited", [(0, 9, 140, 160)])], years, 2) == ["unknown", "unknown"]


def test_continuation_tables_do_not_borrow_assurance():
    from adaptive_document_agent.extraction.table_reconstructor import TableReconstructor
    left = ExtractedTable(table_id="left", page=2, headers=["Metric", "2020", "2021"],
        column_periods=[None, "FY2020", "FY2021"], column_audit_statuses=["unknown", "unaudited", "unknown"])
    right = left.model_copy(deep=True, update={"table_id": "right", "page": 3})
    assert TableReconstructor._compatible(left, right)
    right.column_audit_statuses = []
    assert not TableReconstructor._compatible(left, right)


@pytest.mark.parametrize("metric", ["Cash balance", "Reservoir volume", "Units shipped"])
def test_through_latest_date_cannot_hide_rebound(metric):
    values = [observation(str(i), metric, v, f"FY{2020+i}", unit="count")
              for i, v in enumerate((140, 280, 110, 70, 80))]
    slide = PresentationSlide(id="s", slide_type="analysis", title="Review",
        message=f"{metric} peaked in the second year and then declined through the latest reported date.")
    assert "non_monotonic_claim" in {i.code for i in ClaimValidator().validate_slide(slide, values)}
    slide.message = f"{metric} declined overall, with a partial rebound at the latest date."
    assert not ClaimValidator().validate_slide(slide, values)


def test_latest_evidence_review_is_scope_aware_and_can_explain_exclusion():
    result = paired_result()
    later = result.observations[0].model_copy(deep=True, update={"id": "later", "period": "6M2030"})
    result.observations.append(later)
    code = "presentation_recent_evidence_omitted"
    assert code in {f.code for f in review_presentation(result.presentation_plan, result)}
    later.entity = "Unrelated subsidiary"
    assert code not in {f.code for f in review_presentation(result.presentation_plan, result)}
    later.entity = result.observations[0].entity
    result.presentation_plan.coverage_notes = ["Latest interim scope excluded pending comparable-period evidence."]
    assert code not in {f.code for f in review_presentation(result.presentation_plan, result)}


def test_catalog_exposes_assurance_and_different_definition_counterparts():
    result = paired_result()
    result.charts = []
    result.observations = [observation("reported", "Reported output", 12), observation("adjusted", "Adjusted output", 16)]
    result.observations[0].ifrs_status = "IFRS"
    result.observations[1].ifrs_status = "ADJUSTED"
    result.observations[1].audited_status = "unknown"
    catalog = build_evidence_catalog(result)
    assert all(s["related_evidence"] and not s["related_evidence"][0]["same_definition_basis"] for s in catalog["series"])
    assert next(o for o in catalog["observations"] if o["id"] == "adjusted")["audited_status"] == "unknown"
    assert all(len(c["observation_ids"]) == 1 for c in catalog["series"])


def test_recent_chart_subject_basis_is_not_crowded_out_by_unrelated_metrics():
    result = paired_result()
    # Match the selected source series, while retaining the distinct period.
    selected = result.observations[0]
    latest = selected.model_copy(deep=True, update={"id": "recent", "period": "6M2028", "period_basis": "6M"})
    result.observations.extend([observation(f"other-{i}", f"Other measure {i}", i+1, unit="count") for i in range(80)])
    result.observations.append(latest)
    catalog = build_evidence_catalog(result, max_series=6, max_observations=30)
    assert "recent" in {o["id"] for o in catalog["observations"]}
    recent_series = next(s for s in catalog["series"] if "recent" in s["observation_ids"])
    assert recent_series["observation_ids"] == ["recent"]
    assert len(catalog["series"]) <= 6 and len(catalog["observations"]) <= 30


def test_profile_paginates_all_supported_facts_at_readable_size():
    from adaptive_document_agent.services.presentation_brief import BriefItem, render_profile
    deck = blank_deck()
    items = [BriefItem(f"Fact {i}", f"Source-supported detail number {i}. " + "Complete scope qualification. " * 4, [i+1]) for i in range(8)]
    render_profile(deck, "Company overview", items)
    assert len(deck.slides) > 1
    all_text = "\n".join(visible(s) for s in deck.slides)
    assert all(i.text in all_text for i in items)
    assert all(p.font.size.pt >= 18 for s in deck.slides for shape in s.shapes if shape.name == "brief:body" for p in shape.text_frame.paragraphs)


def test_company_overview_uses_richer_fields_without_duplicate_fact_continuation():
    from adaptive_document_agent.models import DocumentPage
    from adaptive_document_agent.services.pptx_export import _add_company_at_a_glance
    from tests.test_company_overview_pipeline import _make_dummy_pipeline_result

    company = CompanyProfile(
        name="Example Robotics Limited",
        one_line_description="Develops collaborative robots for industrial and education applications.",
        industry="Industrial automation",
        products=["Six-axis cobots", "Four-axis cobots"],
        segments=["Industrial", "Education"],
        business_model="Direct sales and distributors",
        geographies=["Mainland China", "Europe"],
        listing_market="HKEX",
        listing_facts=["H-share global offering"],
        identity_state="RESOLVED",
        source_pages=[1],
        field_source_pages={key: [1] for key in (
            "name", "one_line_description", "industry", "products", "segments",
            "business_model", "geographies", "listing_market", "listing_facts"
        )},
        key_facts=[
            CompanyFact(label="Main products/services", value="Six-axis and four-axis cobots", source_pages=[1]),
            CompanyFact(label="Listing market", value="HKEX", source_pages=[1]),
        ],
    )
    result = _make_dummy_pipeline_result([
        DocumentPage(page_number=1, text=(
            "EXAMPLE ROBOTICS LIMITED\nIndustrial automation\n"
            "Develops collaborative robots for industrial and education applications.\n"
            "Six-axis cobots and four-axis cobots. Direct sales and distributors.\n"
            "Mainland China and Europe. HKEX H-share global offering."
        ))
    ], company)
    deck = blank_deck()

    _add_company_at_a_glance(deck, result, result.presentation_plan.slides[1])

    assert len(deck.slides) == 1
    text = visible(deck.slides[0])
    assert "Develops collaborative robots" in text
    assert "Segments: Industrial, Education" in text
    assert "Listing details: H-share global offering" in text
    assert text.count("Listing market") == 0


def test_four_watch_items_render_on_one_page_without_continuation():
    from adaptive_document_agent.services.pptx_export import _add_planned_text_slide

    result = paired_result()
    bullets = [
        "Operating cash outflows could absorb available liquidity.",
        "Higher borrowings could reduce financial flexibility.",
        "Inventory growth increases the importance of stock conversion.",
        "The latest adjusted loss improvement may not persist.",
    ]
    slide = PresentationSlide(
        id="risks",
        slide_type="risks",
        title="Risks and Watch Items",
        message="The reported movements identify four areas to monitor.",
        bullets=bullets,
        source_pages=[1, 2],
    )
    deck = blank_deck()

    _add_planned_text_slide(deck, result, slide)

    assert len(deck.slides) == 1
    text = visible(deck.slides[0])
    assert all(bullet in text for bullet in bullets)
    assert "continued" not in text.casefold()


def test_no_chart_kpis_keep_commentary_and_readable_font():
    from adaptive_document_agent.services.pptx_export import build_presentation
    result = paired_result()
    result.charts = []
    plan = result.presentation_plan.slides[3]
    plan.chart_ids = []
    plan.visual_blocks = [PresentationVisualBlock(role="kpi", observation_ids=["a0", "a1"])]
    plan.bullets = ["A source-supported qualification remains visible."]
    deck = Presentation(io.BytesIO(build_presentation(result)))
    assert plan.bullets[0] in "\n".join(visible(s) for s in deck.slides)
    shape = next(sh for s in deck.slides for sh in s.shapes if sh.has_text_frame and plan.bullets[0] in sh.text)
    assert all(p.font.size.pt >= 18 for p in shape.text_frame.paragraphs)


def test_three_period_kpis_fit_same_page_as_chart():
    from adaptive_document_agent.document_model import DocumentIndex
    from adaptive_document_agent.services.slide_compositor import render_composed_slide, validate_composed_geometry
    result = paired_result()
    plan = result.presentation_plan.slides[3]
    plan.chart_ids = ["b"]
    plan.visual_blocks = [PresentationVisualBlock(role="hero", chart_ids=["b"]),
                          PresentationVisualBlock(role="kpi", observation_ids=["a0", "a1", "a2"])]
    plan.bullets = ["Supporting context remains visible."]
    deck = blank_deck()
    pages = render_composed_slide(deck, plan, [result.charts[1]], result, DocumentIndex(result.observations))
    assert len(pages) == 1
    table = next(sh for sh in pages[0].shapes if sh.name == "table:a0,a1,a2")
    assert len(table.table.rows) == 3
    validate_composed_geometry(deck)


@pytest.mark.parametrize("text", ["A sentence.\n\nA material caveat.\r\n", "汉字与英文 figures 12.4% mixed。\n", "x" * 400, "    Multiple   spaces\tremain."])
def test_font_capacity_preserves_every_character(text):
    from adaptive_document_agent.services.text_capacity import wrap_copy
    assert "".join(wrap_copy(text, 2, 18)) == text


def test_customer_sector_and_incomplete_address_are_not_company_facts():
    from adaptive_document_agent.models import CompanyProfile, DocumentPage
    from adaptive_document_agent.services.company_extractor import extract_structured_company_fields, validate_headquarters
    from tests.test_company_overview_pipeline import _make_dummy_pipeline_result
    result = _make_dummy_pipeline_result([
        DocumentPage(page_number=1, text="GLOBAL OFFERING\nEXAMPLE SYSTEMS LIMITED\nStock code: 5678"),
        DocumentPage(page_number=2, text="BUSINESS OVERVIEW\nOur Company supplies equipment to consumer electronics customers.")], CompanyProfile())
    assert extract_structured_company_fields(CompanyProfile(), result).industry == ""
    assert validate_headquarters("The PRC No") == ""
    assert validate_headquarters("The country Building") == ""


def test_quality_notes_do_not_rewrite_actual_validation_warnings():
    from adaptive_document_agent.models import ValidationIssue
    from adaptive_document_agent.services.pptx_export import _add_quality_slide
    result = paired_result()
    message = "Conflicting audited column evidence requires review."
    result.validation_warnings = [ValidationIssue(code="test", severity="warning", stage="test", message=message)]
    deck = blank_deck()
    _add_quality_slide(deck, result)
    assert message in visible(deck.slides[0])
