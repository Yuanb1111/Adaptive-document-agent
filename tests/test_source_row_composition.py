"""Complete source-row mixes and compact multi-chart topics remain evidence-bound."""
import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from adaptive_document_agent.agent.chart_planner import ChartPlanner
from adaptive_document_agent.agent.presentation_topic_selector import series_directory
from adaptive_document_agent.agent.topic_plan_compiler import compile_topic_plan
from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import PresentationTopic, PresentationTopicSelection
from adaptive_document_agent.services.composition_candidates import presentation_compositions
from adaptive_document_agent.services.composition_data import composition_data
from adaptive_document_agent.services.pptx_export import _add_native_chart
from adaptive_document_agent.services.slide_compositor import render_composed_slide
from adaptive_document_agent.ui.charts import render_chart
from test_p0_composition import observation, result_for


def source_matrix():
    values = []
    for period in ('FY2023', 'FY2024'):
        for name, value in [('Hardware', 40), ('Services', 30), ('Support', 20), ('Other', 10)]:
            o = observation(f'{period}-{name}', f'{name}: % of Total', value, period, unit='percent')
            o.evidence[0].row_label = name
            o.evidence[0].column_label = '% of Total'
            o.evidence[0].table_id = 'source-mix'
            o.dimensions = {'table_context': 'Reported sales mix'}
            values.append(o)
    return values


def test_source_row_matrix_and_snapshot_keep_raw_facts_and_editable_data():
    obs = source_matrix()
    before = [o.model_dump() for o in obs]
    charts = presentation_compositions(DocumentIndex(obs))
    stack = next(c for c in charts if c.chart_type == 'stacked_percent')
    pie = next(c for c in charts if c.chart_type == 'pie')
    assert stack.composition_table_id == 'source-mix'
    assert composition_data(stack, obs).categories == ['Hardware', 'Other', 'Services', 'Support']
    snapshot = [o for o in obs if o.id in pie.observation_ids]
    assert {o.period for o in snapshot} == {'FY2024'}
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    _add_native_chart(slide, pie, snapshot, (.5, .5, 7, 4))
    native = next(s.chart for s in slide.shapes if s.has_chart)
    ui = render_chart(pie, DocumentIndex(obs))
    assert native.chart_type == XL_CHART_TYPE.PIE
    assert list(native.series[0].values) == list(ui.data[0].values) == [40, 10, 30, 20]
    assert ui.data[0].hole == 0
    assert [o.model_dump() for o in obs] == before


@pytest.mark.parametrize('mutation', ['missing', 'table', 'labels', 'negative', 'dimension', 'column'])
def test_source_row_composition_rejects_unproved_matrices(mutation):
    obs = source_matrix()
    if mutation == 'missing':
        obs.pop()
    elif mutation == 'table':
        obs[-1].evidence[0].table_id = 'unrelated'
    elif mutation == 'labels':
        obs[-1].evidence.append(obs[-1].evidence[0].model_copy(update={'row_label': 'Conflicting row'}))
    elif mutation == 'negative':
        obs[-1].value = -10
    elif mutation == 'dimension':
        obs[-1].dimensions['region'] = 'East'
    else:
        for o in obs:
            o.metric_original = o.metric_original.replace('% of Total', 'Gross margin')
            o.evidence[0].column_label = 'Gross margin'
    assert presentation_compositions(DocumentIndex(obs)) == []


def test_source_row_matrices_separate_annual_and_interim_periods():
    obs = source_matrix()
    interim = [o.model_copy(deep=True) for o in obs]
    for o in interim:
        o.id = 'interim-' + o.id
        o.period = o.period.replace('FY', '6M')
        o.period_basis = '6M'
    charts = presentation_compositions(DocumentIndex(obs + interim))
    assert len(charts) == 4
    for c in charts:
        periods = {o.period[:2] for o in obs + interim if o.id in c.observation_ids}
        assert len(periods) == 1


def test_matrix_uses_one_selection_slot_and_three_native_charts_fit_one_page():
    obs = source_matrix()
    for metric in ('Units shipped', 'Service hours'):
        obs.extend(observation(f'{metric}-{i}', metric, v, f'FY{2022+i}', unit='count')
                   for i, v in enumerate((100, 120, 140)))
    result = result_for(obs, [])
    directory, lookup = series_directory(result)
    ids = [next(d['id'] for d in directory if d['visual_kind'] == 'stacked_percent')]
    ids += [next(d['id'] for d in directory if d['metric'] == m) for m in ('Units shipped', 'Service hours')]
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(
        id='activity', title='Sales mix and operating activity', question='How do mix and activity compare?',
        rationale='Compare complementary reported measures.', series_ids=ids)])
    result.charts = ChartPlanner().plan([], [], DocumentIndex(obs), requested_series=[lookup[s] for s in ids], only_requested=True)
    result.presentation_plan = compile_topic_plan(result)
    slide_plan = next(s for s in result.presentation_plan.slides if s.slide_type == 'analysis')
    assert slide_plan.layout == 'three_up'
    assert len(slide_plan.chart_ids) == 3
    assert not slide_plan.observation_ids
    ppt = Presentation()
    ppt.slide_width, ppt.slide_height = Inches(13.333), Inches(7.5)
    selected = [c for c in result.charts if c.id in slide_plan.chart_ids]
    render_composed_slide(ppt, slide_plan, selected, result, DocumentIndex(obs))
    assert len(ppt.slides) == 1
    native = [s.chart for s in ppt.slides[0].shapes if s.has_chart]
    assert len(native) == 3
    assert XL_CHART_TYPE.COLUMN_STACKED_100 in {c.chart_type for c in native}

def test_amount_mix_directory_retains_denominator_through_topic_compilation():
    from test_p0_composition import matrix
    obs, _ = matrix(metric='Revenue')
    for o in obs:
        o.unit = o.unit_family = 'currency'
        o.currency = 'USD'
    # Totals have the same reported metric plus an explicit aggregate category.
    totals = [observation(f'total-{p}', 'Revenue', 100, p, category='Total')
              for p in ('FY2023', 'FY2024')]
    result = result_for(obs + totals, [])
    directory, lookup = series_directory(result)
    chosen = next(d for d in directory if d['visual_kind'] == 'stacked_percent')
    assert {o.id for o in totals} <= {o.id for o in lookup[chosen['id']]}
    result.presentation_topics = PresentationTopicSelection(topics=[PresentationTopic(
        id='mix', title='Revenue mix', question='How does the reported revenue mix compare?',
        rationale='Compare the complete reported composition.', series_ids=[chosen['id']])])
    result.charts = ChartPlanner().plan([], [], DocumentIndex(result.observations),
        requested_series=[lookup[chosen['id']]], only_requested=True)
    plan = compile_topic_plan(result)
    slide = next(s for s in plan.slides if s.slide_type == 'analysis')
    assert len(slide.chart_ids) == 1
    assert not slide.observation_ids
    chart = next(c for c in result.charts if c.id == slide.chart_ids[0])
    assert chart.chart_type == 'stacked_percent'
    assert set(chart.total_observation_ids) == {o.id for o in totals}

def test_small_composition_parts_do_not_hide_every_label():
    obs = source_matrix()
    for o in obs:
        o.value = 2 if o.id.endswith('Other') else 48 if o.id.endswith('Hardware') else o.value
        o.raw_value = str(o.value)
    charts = presentation_compositions(DocumentIndex(obs))
    ppt = Presentation()
    for plan in charts:
        selected = [o for o in obs if o.id in plan.observation_ids]
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        _add_native_chart(slide, plan, selected, (.5, .5, 7, 4))
        native = next(s.chart for s in slide.shapes if s.has_chart)
        assert native.plots[0].has_data_labels
        if plan.chart_type == 'pie':
            from pptx.enum.chart import XL_DATA_LABEL_POSITION
            assert native.series[0].points[1].data_label.position == XL_DATA_LABEL_POSITION.OUTSIDE_END
        else:
            label = native.series[1].points[0].data_label._dLbl
            assert not label.xpath('c:delete')
            assert label.xpath('c:showVal')[0].get('val') == '0'
            assert label.xpath('c:showPercent')[0].get('val') == '0'
