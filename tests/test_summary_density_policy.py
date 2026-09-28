"""Sparse duplicate summaries can disappear without losing source records."""

import io

import pytest
from pptx import Presentation

from adaptive_document_agent.models import PresentationPlan, PresentationSlide, PresentationVisualBlock
from adaptive_document_agent.services.presentation_brief import omit_redundant_summary
from adaptive_document_agent.services.pptx_export import (
    BUNDLED_TEMPLATE_PATH,
    _add_thank_you_slide,
    _build_planned_presentation,
    _number_slides,
)
from tests.test_deck_standardization import _build_test_pipeline_result


def _plan():
    summary = PresentationSlide(
        id="summary", slide_type="executive_summary", title="Executive Summary",
        bullets=["Response time shortened.", "Service coverage"],
        observation_ids=["retained_record"], source_pages=[12],
    )
    analysis = PresentationSlide(
        id="analysis", slide_type="analysis", title="Response time shortened.",
        section_title="Service coverage",
    )
    return PresentationPlan(title="Service review", planning_origin="topic_recovery",
                            slides=[summary, analysis]), summary


@pytest.mark.parametrize("origin", ["topic_recovery", "topic_compilation"])
def test_short_repeated_final_headings_are_omitted_without_mutating_plan(origin):
    plan, summary = _plan()
    plan.planning_origin = origin
    summary.bullets = ["  RESPONSE  time\nshortened.  ", "Service coverage", "  "]
    before = plan.model_dump()

    assert omit_redundant_summary(plan, summary)
    assert plan.model_dump() == before


@pytest.mark.parametrize("origin", ["model", "repaired", "fallback", "legacy"])
def test_independently_planned_summaries_are_retained(origin):
    plan, summary = _plan()
    plan.planning_origin = origin

    assert not omit_redundant_summary(plan, summary)


@pytest.mark.parametrize("bullets", [
    [],
    ["  "],
    ["Response time shortened.", "Coverage excludes overseas operations."],
    ["Response time shortened"],
    ["Response time shortened. Coverage excludes overseas operations."],
    ["Response time"],
])
def test_independent_or_partially_matching_copy_is_not_discarded(bullets):
    plan, summary = _plan()
    summary.bullets = bullets

    assert not omit_redundant_summary(plan, summary)


@pytest.mark.parametrize("field,value", [
    ("insight_ids", ["independent_insight"]),
    ("chart_ids", ["summary_chart"]),
    ("visual_blocks", [PresentationVisualBlock(role="kpi", observation_ids=["retained_record"])]),
    ("visual_blocks", [PresentationVisualBlock(role="table", observation_ids=["retained_record"])]),
    ("visual_blocks", [PresentationVisualBlock(role="commentary", title="Independent context")]),
])
def test_explicit_summary_content_is_retained(field, value):
    plan, summary = _plan()
    setattr(summary, field, value)

    assert not omit_redundant_summary(plan, summary)


@pytest.mark.parametrize("bullets", [
    ["Response time shortened.", "Service coverage", "Customer adoption", "Delivery capacity"],
    ["Complete service coverage and response-time qualification " * 6],
    ["Response time fell from 8 days to 5 days."],
])
def test_rich_or_value_bearing_summaries_stay_even_when_headings_repeat(bullets):
    plan, summary = _plan()
    summary.bullets = bullets
    plan.slides.extend(PresentationSlide(id=f"detail_{i}", slide_type="analysis", title=text)
                       for i, text in enumerate(bullets))

    assert not omit_redundant_summary(plan, summary)


def _text(slide):
    return "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)


def test_export_omits_duplicate_summary_and_contents_but_retains_full_provenance():
    from adaptive_document_agent.services.company_extractor import extract_structured_company_fields

    result = _build_test_pipeline_result(slide_title="Current ratio weakened while revenue grew.")
    plan = result.presentation_plan
    plan.planning_origin = "topic_compilation"
    # Existing introduction rendering normalizes legacy company fields first.
    plan.company = extract_structured_company_fields(plan.company, result)
    summary = next(slide for slide in plan.slides if slide.slide_type == "executive_summary")
    analysis = next(slide for slide in plan.slides if slide.slide_type == "analysis")
    summary.bullets = [analysis.title, analysis.section_title]
    summary.bullet_observation_ids = [["rev_2022", "rev_2023"], ["cr_2022", "cr_2023"]]
    original_summary = summary.model_dump_json(indent=2)
    original_plan = plan.model_dump()
    original_observations = [item.model_dump() for item in result.observations]

    deck = Presentation(str(BUNDLED_TEMPLATE_PATH))
    for slide_id in list(deck.slides._sldIdLst):
        deck.part.drop_rel(slide_id.rId)
        deck.slides._sldIdLst.remove(slide_id)
    _build_planned_presentation(deck, result)
    _add_thank_you_slide(deck)
    _number_slides(deck)
    stream = io.BytesIO()
    deck.save(stream)
    deck = Presentation(io.BytesIO(stream.getvalue()))

    assert all("Executive Summary" not in _text(slide) for slide in deck.slides)
    assert any(analysis.title in _text(slide) for slide in deck.slides)
    entries = [shape for slide in deck.slides for shape in slide.shapes
               if shape.name.startswith("contents:entry:")]
    assert entries and analysis.section_title in [shape.text for shape in entries]
    assert [shape.name for shape in entries] == [f"contents:entry:{i}" for i in range(1, len(entries) + 1)]
    assert any(original_summary in slide.notes_slide.notes_text_frame.text for slide in deck.slides)
    assert plan.model_dump() == original_plan
    assert [item.model_dump() for item in result.observations] == original_observations
    for number, slide in enumerate(deck.slides, start=1):
        page_numbers = [shape.text for shape in slide.shapes if shape.has_text_frame
                        and abs(shape.top.inches - 6.53) < .01
                        and abs(shape.left.inches - 11.60) < .01]
        assert page_numbers == ([] if number in {1, len(deck.slides)} else [str(number)])
