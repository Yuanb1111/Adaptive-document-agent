"""Whitespace ownership must not erase content or suppress real reading gaps."""
from copy import deepcopy
import json
from threading import Event

import pytest

from adaptive_document_agent.agent.summary_reader import SummaryReadBatch, _read_batch
from adaptive_document_agent.models.summary import SummaryReview
from adaptive_document_agent.services.summary_source import source_blocks
from adaptive_document_agent.services.summary_validation import reading_errors
from tests.test_complete_summary_reading import source, evidence_payload


def fixture():
    result = source(1)
    text = 'SUMMARY\n\nCedar builds scheduling software for manufacturing teams.\n\nDelivery requires customer training.\n\n-- i --'
    result.document.pages[0].text = text
    blocks = source_blocks(result, [1])
    bid = blocks[0].id
    def part(start, end, role, text='', quote_line=None):
        return dict(block_id=bid, start_line=start, end_line=end, role=role,
            title='Operations' if role == 'content' else 'Source layout',
            reading_note='The literal source passage describes this subject.',
            facts=[dict(label='Operations', text=text, quote_start_line=quote_line,
                        quote_end_line=quote_line)] if text else [])
    rows = [part(1,1,'heading'), part(3,3,'content', text.splitlines()[2],3),
            part(5,7,'content',text.splitlines()[4],5), part(7,7,'layout')]
    return result, blocks, rows, part


class Responses:
    def __init__(self, responses):
        self.responses = iter(responses)
        self.calls = []

    def generate_structured(self, messages, response_model, **kwargs):
        self.calls.append(deepcopy(messages))
        value = next(self.responses)
        if isinstance(value, Exception):
            raise value
        return response_model.model_validate({'parts': deepcopy(value)})


def assert_complete(result, blocks, outcome):
    assert outcome.error is None
    review = SummaryReview(document_id=result.document.document_id,
        document_sha256=result.document.sha256, source_pages=[1], source_blocks=blocks,
        parts=outcome.parts)
    assert not reading_errors(review, result)
    return review


def test_blank_gaps_and_duplicate_layout_need_no_model_retry():
    result, blocks, rows, _ = fixture()
    before = deepcopy(rows)
    client = Responses([rows])
    outcome = _read_batch(client, result, blocks, Event())
    assert_complete(result, blocks, outcome)
    assert len(client.calls) == 1
    assert outcome.audit['attempts'][0]['response']['parts'] == SummaryReadBatch(parts=before).model_dump()['parts']
    actual = [p for p in outcome.parts if p.role == 'content']
    assert [(p.start_line,p.end_line) for p in actual] == [(3,3),(5,7)]
    assert [p.facts[0].text for p in actual] == [r['facts'][0]['text'] for r in rows if r['facts']]
    assert {r['action'] for r in outcome.audit['attempts'][0]['layout_reconciliation']} == {
        'assign_blank_lines', 'remove_redundant_layout'}


@pytest.mark.parametrize('missing_content', [False, True])
def test_missing_nonblank_line_is_read_by_model_and_existing_parts_are_preserved(missing_content):
    result, blocks, rows, part = fixture()
    rows[2]['end_line'] = 5
    rows.pop()  # Footer unassigned.
    if missing_content:
        # The program must not assume a short final line is a footer.
        result.document.pages[0].text = result.document.pages[0].text.replace('-- i --', 'Service excludes hardware.')
        blocks = source_blocks(result,[1])
        for row in rows:
            row['block_id'] = blocks[0].id
    patch = part(7,7,'content' if missing_content else 'layout',
                 'Service excludes hardware.' if missing_content else '', 7)
    patch['block_id'] = blocks[0].id
    client = Responses([rows, rows, [patch]])
    outcome = _read_batch(client, result, blocks, Event())
    assert_complete(result, blocks, outcome)
    assert len(client.calls) == 3
    assert outcome.audit['missing_range_read']['missing_ranges'] == [
        dict(block_id=blocks[0].id, start_line=7,end_line=7)]
    assert evidence_payload(client.calls[-1]) == evidence_payload(client.calls[0])
    assert 'Read ONLY' in client.calls[-1][-1]['content']
    recovered = next(p for p in outcome.parts if p.start_line == 7)
    assert recovered.role == ('content' if missing_content else 'layout')
    assert bool(recovered.facts) is missing_content


@pytest.mark.parametrize('invalid_patch', ['overwrite', 'missing_fact', 'wrong_quote', 'overlap'])
def test_missing_range_read_cannot_change_accepted_parts_or_invent_support(invalid_patch):
    result, blocks, rows, part = fixture()
    rows[2]['end_line'] = 5
    rows.pop()
    patch = [part(7,7,'layout')]
    if invalid_patch == 'overwrite':
        patch = [part(5,7,'layout')]
    elif invalid_patch == 'missing_fact':
        patch = [part(7,7,'content')]
    elif invalid_patch == 'wrong_quote':
        patch = [part(7,7,'content','Cedar has 999 factories.',3)]
    else:
        patch *= 2
    outcome = _read_batch(Responses([rows, rows, patch]), result, blocks, Event())
    assert outcome.error is not None and not outcome.parts
    assert outcome.audit['missing_range_read']['error']


def test_substantive_overlap_still_fails_after_bounded_targeted_read():
    result, blocks, rows, _ = fixture()
    rows.append(deepcopy(rows[1]))
    client = Responses([rows,rows,[rows[1],rows[1]]])
    outcome = _read_batch(client,result,blocks,Event())
    assert outcome.error is not None and len(client.calls) == 3
    assert 'overlap' in str(outcome.error)
    assert 'missing_range_read' not in outcome.audit


def test_model_repartitions_overlap_without_changing_accepted_parts():
    result, blocks, rows, _ = fixture()
    rows.append(deepcopy(rows[1]))
    client = Responses([rows, rows, [rows[1]]])
    outcome = _read_batch(client, result, blocks, Event())
    assert_complete(result, blocks, outcome)
    assert len(client.calls) == 3
    assert outcome.audit['rejected_part_read']['ranges'] == [
        dict(block_id=blocks[0].id, start_line=3, end_line=3)]
    assert next(p for p in outcome.parts if p.start_line == 5).facts[0].text == rows[2]['facts'][0]['text']


def test_missing_range_read_privacy_failure_is_preserved():
    from adaptive_document_agent.services.llm.exceptions import PrivacyViolationError
    result, blocks, rows, _ = fixture()
    rows[2]['end_line'] = 5
    rows.pop()
    client = Responses([rows,rows,PrivacyViolationError('Local Only')])
    outcome = _read_batch(client,result,blocks,Event())
    assert isinstance(outcome.error, PrivacyViolationError)
    assert outcome.audit['missing_range_read']['error'] == 'Local Only'
    assert len(client.calls) == 3


def test_missing_range_output_truncation_is_audited_without_another_retry():
    from adaptive_document_agent.services.llm.base import LLMResponse
    from adaptive_document_agent.services.llm.usage import LLMUsage
    from adaptive_document_agent.services.llm.exceptions import LLMStructuredOutputError
    result, blocks, rows, _ = fixture()
    rows[2]['end_line'] = 5
    rows.pop()
    failure = LLMStructuredOutputError('Output limit', response=LLMResponse(text='{"parts":[',
        usage=LLMUsage(provider='mock',model='mock',finish_reason='length')))
    client = Responses([rows,rows,failure])
    outcome = _read_batch(client,result,blocks,Event())
    assert outcome.error is failure and len(client.calls) == 3
    assert outcome.audit['missing_range_read']['raw_response'] == '{"parts":['


def test_arbitrary_outside_line_endpoint_stays_bounded_and_invalid():
    result, blocks, rows, _ = fixture()
    rows[-1]['end_line'] = 10**12
    outcome = _read_batch(Responses([rows,rows]),result,blocks,Event())
    assert outcome.error is not None
    assert outcome.audit['attempts'][0]['line_coverage'][0]['outside_lines'] == [10**12]


def test_recovered_reading_reaches_native_ppt_with_business_opening_and_source_notes():
    from io import BytesIO
    from pptx import Presentation
    from adaptive_document_agent.agent.company_introduction import ensure_company_introduction
    from adaptive_document_agent.models import PresentationPlan, PresentationSlide
    from adaptive_document_agent.services.company_summary import require_verified_introduction
    from adaptive_document_agent.services.llm.base import LLMResponse
    from adaptive_document_agent.services.pptx_export import build_presentation
    from tests.test_complete_summary_reading import SummaryClient, gateway

    result, blocks, rows, _ = fixture()
    class Client(SummaryClient):
        def generate_structured(self, messages, response_model, **kwargs):
            if response_model is SummaryReadBatch:
                response = SummaryReadBatch(parts=rows)
                return response, LLMResponse(text=response.model_dump_json())
            return super().generate_structured(messages, response_model, **kwargs)

    plan = PresentationPlan(title='Source report', slides=[PresentationSlide(
        id=kind,slide_type=kind,title=kind.replace('_',' ').title()) for kind in
        ('cover','company_overview','executive_summary','data_quality','appendix')])
    ensure_company_introduction(gateway(Client(), local=True), result, plan)
    result.presentation_plan = plan
    require_verified_introduction(result)
    assert plan.company.summary_review.status == 'complete'
    deck = Presentation(BytesIO(build_presentation(result)))
    text = '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert 'scheduling software' in text and 'customer training' in text
    assert any('Source' in slide.notes_slide.notes_text_frame.text for slide in deck.slides)
