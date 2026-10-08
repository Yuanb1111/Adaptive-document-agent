"""Summary capacity and renderer line spacing use the same absolute units."""

import io
import os
from copy import deepcopy

import pymupdf
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from adaptive_document_agent.services import export
from adaptive_document_agent.services.pdf_render_layout import pdf_rendered_pages
from adaptive_document_agent.services.presentation_brief import BriefItem, LINE_HEIGHT_FACTOR
from adaptive_document_agent.services.presentation_rendering import configured_renderer
from adaptive_document_agent.services.presentation_summary import render_complete_summary
from adaptive_document_agent.services.presentation_visual_qa import inspect_pages
from tests.ppt_render_stub import LocalRenderStub
from tests.test_presentation_brief import blank_deck
from tests.test_pptx_export import _result


COPY = (
    "The measured outcome covers the stated population and reporting period. "
    "The source retains every qualification and excludes unreported observations. "
)


def _summary(case):
    if case == "compact":
        items = [BriefItem(f"Finding {i}", COPY * 3, [i + 1]) for i in range(3)]
    elif case == "wrapped-heading":
        items = [BriefItem("A qualified finding across the reported populations and observation periods " * 3,
                           COPY * 3, [7])]
    elif case == "continuation":
        items = [BriefItem("Evidence scope", COPY * 30, [9]),
                 BriefItem("Next finding", "The next complete finding remains visible.", [10])]
    elif case == "cjk":
        items = [BriefItem("证据范围", "仅包括已披露样本，不补充缺失数据。结论保留期间、单位和证据。" * 12, [11])]
    else:
        items = [BriefItem(f"Finding {i}", COPY * 3, [i + 1]) for i in range(2)]
    before = deepcopy(items)
    deck = blank_deck()
    slides = render_complete_summary(deck, "Selected findings", items, single_column=case == "compact")
    assert items == before
    bodies = [s for slide in slides for s in slide.shapes if s.name == "brief:body"]
    assert "".join(s.text for s in bodies) == "".join(item.text for item in items)
    assert {p.font.size.pt for s in bodies for p in s.text_frame.paragraphs} == {16}
    if case == "compact":
        assert len(slides) > 1
    if case == "continuation":
        assert len(slides) > 1
    stream = io.BytesIO()
    deck.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize("case", ["compact", "standard", "wrapped-heading", "continuation", "cjk"])
def test_summary_serializes_exact_point_leading_for_every_text_size(case):
    deck = Presentation(io.BytesIO(_summary(case)))
    sizes = set()
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.name not in {"brief:heading", "brief:body"}:
                continue
            for paragraph in shape.text_frame.paragraphs:
                size = paragraph.font.size.pt
                sizes.add(size)
                assert paragraph.line_spacing == Pt(size * LINE_HEIGHT_FACTOR)
                assert paragraph._p.xpath("./a:pPr/a:lnSpc/a:spcPts")
                assert not paragraph._p.xpath("./a:pPr/a:lnSpc/a:spcPct")
    assert sizes == {16, 18}


def test_old_proportional_spacing_build_cache_is_not_reused(monkeypatch):
    from adaptive_document_agent.services.pptx_export import _resolve_template_path
    from adaptive_document_agent.utils.pipeline_version import PIPELINE_VERSION
    import hashlib

    result, cache = _result(), {}
    first = export.export_pptx_with_report(result, renderer=LocalRenderStub(), build_cache=cache)
    current_key = next(iter(cache))
    legacy_key = (f"ppt-build-v8:{PIPELINE_VERSION}", *current_key[1:])
    assert current_key != legacy_key
    cache = {legacy_key: cache[current_key]}
    calls = []
    original = export.build_presentation

    def build(*args, **kwargs):
        calls.append(True)
        return original(*args, **kwargs)

    monkeypatch.setattr(export, "build_presentation", build)
    rebuilt = export.export_pptx_with_report(result, renderer=LocalRenderStub(), build_cache=cache)
    assert calls == [True] and not rebuilt.build_cache_hit
    assert first.report.slide_count == rebuilt.report.slide_count
    digest = hashlib.sha256(_resolve_template_path(None).read_bytes()).hexdigest()
    assert set(cache) == {export._build_cache_key(result, digest)}


@pytest.mark.skipif(not os.getenv("PPTX_QA_LIBREOFFICE_INTEGRATION"), reason="Real Linux LibreOffice integration")
@pytest.mark.parametrize("case", ["compact", "standard", "wrapped-heading", "continuation", "cjk"])
def test_real_summary_text_remains_inside_its_rendered_region(tmp_path, monkeypatch, case):
    monkeypatch.setenv("PPTX_QA_BACKEND", "libreoffice")
    payload = _summary(case)
    pages = configured_renderer().render(payload, tmp_path)
    assert not inspect_pages(payload, pages)
    assert all(element.get("renderedTextPresent") is not False
               for page in pages for element in page.layout["elements"])


@pytest.mark.parametrize("damage,code", [("missing", "MISSING_RENDER_TEXT"), ("off-canvas", "OUT_OF_BOUNDS")])
def test_summary_gate_still_rejects_missing_copy_and_off_canvas_shapes(tmp_path, damage, code):
    payload = _summary("standard")
    deck = Presentation(io.BytesIO(payload))
    # This controlled PDF omits the source text entirely. Geometry alone cannot
    # certify visibility, even though summary line spacing is now corrected.
    pdf_path = tmp_path / "missing.pdf"
    with pymupdf.open() as pdf:
        for _ in deck.slides:
            page = pdf.new_page(width=deck.slide_width.pt, height=deck.slide_height.pt)
            page.insert_text((25, 35), "Different visible content", fontsize=14)
        pdf.save(pdf_path)
    if damage == "off-canvas":
        body = next(s for s in deck.slides[0].shapes if s.name == "brief:body")
        body.left = Inches(-1)
        stream = io.BytesIO()
        deck.save(stream)
        payload = stream.getvalue()
    issues = inspect_pages(payload, pdf_rendered_pages(payload, pdf_path))
    assert any(issue.code == code and issue.severity == "critical" for issue in issues)
