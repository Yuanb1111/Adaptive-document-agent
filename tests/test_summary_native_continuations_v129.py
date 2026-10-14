from io import BytesIO
import pytest
from pptx import Presentation
from pptx.util import Inches
from adaptive_document_agent.agent.compact_summary import generate_compact_summary
from adaptive_document_agent.models import PresentationPlan
from adaptive_document_agent.models.presentation import CompanyProfile
from adaptive_document_agent.services.summary_export import bind_summary_pages, verify_summary_export
from tests.test_compact_summary_v111 import fixture, Reviewer
from tests.test_presentation_brief import blank_deck


def split_deck():
    result, review, draft = fixture(1)
    company = generate_compact_summary(Reviewer(draft), result, review, CompanyProfile(summary_review=review))
    review.status = 'complete'
    result.presentation_plan = PresentationPlan(title='Verified source', company=company)
    page = company.summary_pages[0]
    first, second = page.items[0].text.split(';', 1)
    deck = blank_deck()
    shapes = []
    for text in [first + ';', second.lstrip()]:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3))
        shape.name, shape.text = 'brief:body', text
        shapes.append(shape)
    bind_summary_pages(list(deck.slides), page)
    return deck, result, shapes


def test_native_geometry_continuations_retain_complete_visible_evidence_after_reload():
    deck, result, shapes = split_deck()
    assert verify_summary_export(deck, result) == {'part-0': [1, 2]}
    stream = BytesIO(); deck.save(stream)
    assert verify_summary_export(Presentation(stream), result) == {'part-0': [1, 2]}


@pytest.mark.parametrize('change', ['missing', 'duplicate', 'reordered', 'copy', 'unbound'])
def test_native_continuation_tampering_never_passes(change):
    deck, result, shapes = split_deck()
    if change == 'missing':
        ident = deck.slides._sldIdLst[1]; deck.part.drop_rel(ident.rId); deck.slides._sldIdLst.remove(ident)
    elif change == 'duplicate':
        shapes[1].element.xpath('.//p:cNvPr')[0].set('descr', shapes[0].element.xpath('.//p:cNvPr')[0].get('descr'))
    elif change == 'reordered':
        ids = deck.slides._sldIdLst; ids.insert(0, ids[1])
    elif change == 'copy': shapes[1].text += ' Invented delivery guarantee.'
    else: shapes[1].element.xpath('.//p:cNvPr')[0].set('descr', '')
    with pytest.raises(ValueError): verify_summary_export(deck, result)
