"""Native segmentation preserves reviewed copy without assuming a language."""
from io import BytesIO

import pytest
from pptx import Presentation
from pptx.util import Inches

from adaptive_document_agent.models import (
    DocumentProfile, ParsedDocument, PipelineResult, PresentationPlan,
)
from adaptive_document_agent.models.presentation import CompanyProfile
from adaptive_document_agent.models.summary import (
    SummaryPart, SummaryReview, SummarySlideItem, SummarySlidePage,
)
from adaptive_document_agent.models.summary_display import SummaryDisplay
from adaptive_document_agent.services.presentation_brief import BriefItem, render_profile
from adaptive_document_agent.services.summary_export import bind_summary_pages, verify_summary_export
from tests.test_presentation_brief import blank_deck


def summary_result(texts):
    """Synthetic reviewed copy; no model acceptance or document-specific fixture."""
    parts = [SummaryPart(id=f'part-{i}', title='Study scope', role='content',
        block_id='source', start_line=i + 1, end_line=i + 1, source_page=1,
        source_text=text, reading_note='Synthetic evidence for copy binding.')
        for i, text in enumerate(texts)]
    review = SummaryReview(document_id='study', document_sha256='synthetic',
        source_pages=[1], status='complete', parts=parts, max_intro_slides=5)
    page = SummarySlidePage(id='scope', title='Study scope', source_section='Scope',
        items=[SummarySlideItem(label=f'Scope {i + 1}', text=text, source_quote=text,
            source_pages=[1], part_ids=[parts[i].id]) for i, text in enumerate(texts)])
    result = PipelineResult(
        document=ParsedDocument(document_id='study', sha256='synthetic',
            safe_filename='study.pdf', page_count=1), profile=DocumentProfile(),
        presentation_plan=PresentationPlan(title='Study',
            company=CompanyProfile(summary_review=review, summary_pages=[page])))
    return result, page


def native_segments(deck, texts, *, separate_slides=False):
    slide = None
    shapes = []
    for i, text in enumerate(texts):
        if slide is None or separate_slides:
            slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1 + i), Inches(8), Inches(1))
        shape.name, shape.text = 'brief:body', text
        shapes.append(shape)
    return shapes


@pytest.mark.parametrize('layout', ['paragraphs', 'comparison', 'timeline', 'relationship', 'key_metrics'])
@pytest.mark.parametrize('copy', [
    '实验涵盖两个样本组。结果仅适用于已测量的参与者。未测量的群体不在研究范围内。',
    '調査は二つの集団を対象とする。結果は測定済みの参加者に限られる。',
    '实验使用Sensor-A。The measured range is 12.5 kg.结果仅适用于本次样本。',
    'The study covers two cohorts. Results apply only to measured participants.',
])
def test_structured_multilingual_copy_binds_and_survives_pptx_reload(layout, copy):
    result, page = summary_result(['The protocol defines the study scope.', copy])
    before = result.model_dump_json()
    deck = blank_deck()
    slides = render_profile(deck, page.title, [
        BriefItem(item.label, item.text, item.source_pages, display=SummaryDisplay(layout=layout))
        for item in page.items])
    shapes = [shape for slide in slides for shape in slide.shapes if shape.name == 'brief:body']
    assert len(shapes) > len(page.items)
    assert ''.join(shape.text for shape in shapes) == ''.join(item.text for item in page.items)
    bind_summary_pages(slides, page)
    assert set(verify_summary_export(deck, result)) == {'part-0', 'part-1'}
    stream = BytesIO()
    deck.save(stream)
    stream.seek(0)
    assert verify_summary_export(Presentation(stream), result) == verify_summary_export(deck, result)
    assert result.model_dump_json() == before


@pytest.mark.parametrize('copy,segments', [
    ('样本仅含已测量对象。结论限于本次实验。', ['样本仅含已测量对象。', '结论限于本次实验。']),
    ('対象は測定済みの参加者。結論には限界がある。', ['対象は測定済みの参加者。', '結論には限界がある。']),
    ('ขอบเขตการศึกษาครั้งนี้', ['ขอบเขต', 'การศึกษา', 'ครั้งนี้']),
    ('The protocol is unchanged.', ['The proto', 'col is unchanged.']),
    ('Results are limited; further study is required.', ['Results are limited;', ' further study is required.']),
    # Persisted older native slides may have trimmed boundary whitespace.
    ('Results are limited; further study is required.', ['Results are limited;', 'further study is required.']),
    ('样本范围。The study is limited.结果不可推广。', ['样本范围。', 'The study is limited.', '结果不可推广。']),
])
@pytest.mark.parametrize('separate_slides', [False, True])
def test_native_boundaries_do_not_invent_word_separators(copy, segments, separate_slides):
    result, page = summary_result([copy])
    deck = blank_deck()
    shapes = native_segments(deck, segments, separate_slides=separate_slides)
    bind_summary_pages(list(deck.slides), page)
    assert verify_summary_export(deck, result) == {'part-0': list(range(1, len(shapes) + 1)) if separate_slides else [1] * len(shapes)}


@pytest.mark.parametrize('fault', ['missing', 'duplicate', 'reordered', 'number', 'unit', 'word_space', 'added_space', 'extra'])
@pytest.mark.parametrize('after_binding', [False, True])
def test_missing_changed_or_reordered_copy_still_blocks_export(fault, after_binding):
    original = ['样本限于本次实验。', 'The measured range is 12.5 kg. ', 'The study has no effect.']
    result, page = summary_result([''.join(original)])
    deck = blank_deck()
    shapes = native_segments(deck, original)
    if after_binding:
        bind_summary_pages(list(deck.slides), page)
    if fault == 'missing':
        shapes[0].element.getparent().remove(shapes[0].element)
    elif fault == 'duplicate':
        shapes[1].text += shapes[1].text
    elif fault == 'reordered':
        shapes[0].text, shapes[1].text = shapes[1].text, shapes[0].text
    elif fault == 'number':
        shapes[1].text = shapes[1].text.replace('12.5', '125')
    elif fault == 'unit':
        shapes[1].text = shapes[1].text.replace('kg', 'g')
    elif fault == 'word_space':
        shapes[2].text = shapes[2].text.replace('no effect', 'noeffect')
    elif fault == 'added_space':
        shapes[0].text = shapes[0].text.replace('本次', '本 次')
    else:
        shapes[2].text += ' An unsupported guarantee.'
    with pytest.raises(ValueError):
        if after_binding:
            verify_summary_export(deck, result)
        else:
            bind_summary_pages(list(deck.slides), page)
