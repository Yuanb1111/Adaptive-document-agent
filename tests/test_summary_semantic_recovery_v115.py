"""Rejected diagram/qualitative spans must be re-read, never silently omitted."""
from copy import deepcopy
from threading import Event

import pytest
from pydantic import ValidationError

from adaptive_document_agent.agent.summary_reader import _read_batch
from adaptive_document_agent.models.summary import SummarySlideItem
from adaptive_document_agent.services.llm.exceptions import PrivacyViolationError
from adaptive_document_agent.services.summary_source import source_blocks
from tests.test_summary_line_recovery_v114 import Responses, assert_complete, fixture


def diagram_fixture():
    result, _, _, _ = fixture()
    diagram = ('Architecture components: scheduling engine, execution console, integration connectors, '
               'monitoring services, operator interface, route planning, device control, telemetry '
               'and safety systems. These components support warehouse automation.')
    lead = 'The following section introduces operating results and their principal contributing factors.'
    result.document.pages[0].text = 'SUMMARY\nCedar makes warehouse software.\n' + diagram + '\n' + lead
    blocks = source_blocks(result, [1])
    def row(line, role, text=''):
        return dict(block_id=blocks[0].id, start_line=line, end_line=line, role=role,
            title='Source topic', reading_note='This passage describes the source topic.',
            facts=[dict(label='Topic', text=text, quote_start_line=line, quote_end_line=line)] if text else [])
    rows = [row(1,'heading'), row(2,'content','Cedar makes warehouse software.'),
            row(3,'layout'), row(4,'content')]
    fixes = [row(3,'content',diagram), row(4,'content',lead)]
    return result, blocks, rows, fixes


def test_repeated_diagram_and_empty_qualitative_fact_failures_get_one_targeted_read():
    result, blocks, rows, fixes = diagram_fixture()
    original = deepcopy(rows)
    client = Responses([rows, rows, fixes[:1]])
    outcome = _read_batch(client, result, blocks, Event())
    assert_complete(result, blocks, outcome)
    assert len(client.calls) == 3
    assert rows == original
    assert outcome.parts[1].facts[0].text == rows[1]['facts'][0]['text']
    assert [p.role for p in outcome.parts] == ['heading','content','content','content']
    audit = outcome.audit['rejected_part_read']
    assert [r['start_line'] for r in audit['ranges']] == [3]
    assert outcome.parts[-1].facts[0].evidence_only
    assert outcome.parts[-1].facts[0].text == result.document.pages[0].text.splitlines()[-1]
    assert audit['errors'] == []
    assert all(p.facts[0].source_pages == [1] for p in outcome.parts if p.role == 'content')


@pytest.mark.parametrize('bad', ['unchanged','overwrite','invention','outside_quote','omit','overlap'])
def test_targeted_read_still_rejects_missing_substance_and_invalid_evidence(bad):
    result, blocks, rows, fixes = diagram_fixture()
    fixes = fixes[:1]
    if bad == 'unchanged':
        fixes = rows[2:3]
    elif bad == 'overwrite':
        fixes.insert(0, rows[1])
    elif bad == 'invention':
        fixes[0]['facts'][0]['text'] = 'There are 999 devices.'
    elif bad == 'outside_quote':
        fixes[0]['facts'][0]['quote_start_line'] = 2
    elif bad == 'omit':
        fixes.pop()
    else:
        fixes.append(deepcopy(fixes[0]))
    client = Responses([rows, rows, fixes])
    outcome = _read_batch(client, result, blocks, Event())
    assert outcome.error is not None and not outcome.parts
    assert len(client.calls) == 3 and outcome.audit['rejected_part_read']['error']


def test_invalid_quote_can_be_reread_without_borrowing_another_parts_source():
    result, blocks, rows, fixes = diagram_fixture()
    rows[2:] = deepcopy(fixes)
    rows[2]['facts'][0]['quote_end_line'] = 4
    outcome = _read_batch(Responses([rows,rows,fixes]), result, blocks, Event())
    assert_complete(result, blocks, outcome)
    assert outcome.parts[2].facts[0].source_quote == result.document.pages[0].text.splitlines()[2]


def test_internal_metadata_is_retained_without_relaxing_visible_label_limit():
    result, blocks, rows, fixes = diagram_fixture()
    rows[2:] = fixes
    rows[1]['facts'][0]['label'] = 'A detailed description of the source subject ' * 4
    rows[1]['reading_note'] = 'The passage identifies what the company provides. ' * 25
    outcome = _read_batch(Responses([rows]), result, blocks, Event())
    assert_complete(result, blocks, outcome)
    assert outcome.parts[1].reading_note == rows[1]['reading_note']
    assert outcome.parts[1].facts[0].label == rows[1]['facts'][0]['label']
    assert len(outcome.audit['attempts']) == 1
    with pytest.raises(ValidationError):
        SummarySlideItem(**outcome.parts[1].facts[0].model_dump(), part_ids=[outcome.parts[1].id])


def test_targeted_read_privacy_error_is_not_salvaged_or_retried():
    result, blocks, rows, _ = diagram_fixture()
    client = Responses([rows, rows, PrivacyViolationError('Local Only')])
    outcome = _read_batch(client, result, blocks, Event())
    assert isinstance(outcome.error, PrivacyViolationError)
    assert len(client.calls) == 3 and outcome.audit['rejected_part_read']['error'] == 'Local Only'


def test_targeted_recovery_reaches_native_ppt_with_source_bound_business_opening():
    from io import BytesIO
    from pptx import Presentation
    from adaptive_document_agent.agent.company_introduction import ensure_company_introduction
    from adaptive_document_agent.agent.summary_reader import SummaryReadBatch
    from adaptive_document_agent.services.llm.base import LLMResponse
    from adaptive_document_agent.services.pptx_export import build_presentation
    from tests.test_complete_summary_reading import SummaryClient, gateway, complete_deck
    from adaptive_document_agent.models import PresentationPlan
    result, _, rows, fixes = diagram_fixture()
    replies = iter([rows, rows, fixes[:1]])
    class Client(SummaryClient):
        def generate_structured(self, messages, response_model, **kwargs):
            if response_model is SummaryReadBatch:
                response = SummaryReadBatch(parts=next(replies))
                return response, LLMResponse(text=response.model_dump_json())
            return super().generate_structured(messages, response_model, **kwargs)
    plan = PresentationPlan(title='Source report')
    ensure_company_introduction(gateway(Client(),local=True), result, plan)
    result.presentation_plan = plan
    complete_deck(result)
    assert plan.company.summary_review.status == 'complete'
    deck = Presentation(BytesIO(build_presentation(result)))
    text = '\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
    assert 'Cedar makes warehouse software.' in text
    assert 'integration connectors' in text
    assert 'principal contributing factors' in text


@pytest.mark.parametrize('include_headers', [True, False])
def test_table_repartition_binds_period_and_units_inside_repaired_part(include_headers):
    result, _, _, make = fixture()
    result.document.pages[0].text = 'SUMMARY\nYear ended 2024 (RMB million)\nRevenue 30\n-- i --'
    blocks = source_blocks(result,[1])
    rows = [make(1,1,'heading'),make(2,2,'layout'),
            make(3,3,'content','Revenue was RMB 30 million in 2024.',3), make(4,4,'layout')]
    for row in rows:
        row['block_id'] = blocks[0].id
    rows[2]['facts'][0]['quote_start_line'] = 2  # Real header, wrong original partition.
    patch = deepcopy(rows[2])
    patch['start_line'] = 2
    patch['facts'][0]['quote_start_line'] = 2 if include_headers else 3
    if include_headers:
        patch['facts'][0]['text'] = 'Year ended 2024 (RMB million)\nRevenue 30'
    outcome = _read_batch(Responses([rows,rows,[patch]]),result,blocks,Event())
    audit = outcome.audit['rejected_part_read']
    assert audit['ranges'] == [dict(block_id=blocks[0].id,start_line=2,end_line=3)]
    if include_headers:
        assert_complete(result,blocks,outcome)
        assert outcome.parts[1].facts[0].source_quote == 'Year ended 2024 (RMB million)\nRevenue 30'
    else:
        assert outcome.error is not None


def test_long_internal_label_keeps_literal_fact_through_reading_layout():
    from adaptive_document_agent.agent.summary_reading_layout import reading_presentation
    from adaptive_document_agent.agent.company_introduction import IntroductionDraft
    result, blocks, rows, fixes = diagram_fixture()
    rows[2:] = fixes
    rows[1]['facts'][0]['label'] = 'Detailed internal source subject ' * 4
    outcome = _read_batch(Responses([rows]),result,blocks,Event())
    review = assert_complete(result,blocks,outcome)
    company = reading_presentation(review,result,IntroductionDraft)
    assert company.summary_pages[0].items[0].text == 'Cedar makes warehouse software.'
    assert len(company.summary_pages[0].items[0].label) <= 60
    assert review.parts[1].facts[0].label == rows[1]['facts'][0]['label']
