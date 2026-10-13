"""Adaptive theme coverage, independent meaning proofs, and native export."""
from io import BytesIO
from threading import Event

import pytest
from pptx import Presentation

from adaptive_document_agent.agent.summary_themes import generate_thematic_summary, plan_themes, SummaryThemePlan
from adaptive_document_agent.agent.compact_summary import CompactSummaryDraft, CompactSummaryReview
from adaptive_document_agent.models import PresentationPlan, PresentationSlide
from adaptive_document_agent.models.presentation import CompanyProfile
from adaptive_document_agent.services.summary_validation import presentation_errors
from adaptive_document_agent.services.summary_export import verify_summary_export
from adaptive_document_agent.services.pptx_export import BUNDLED_TEMPLATE_PATH, _add_company_at_a_glance
from tests.test_compact_summary_v111 import fixture
from tests.test_complete_summary_reading import evidence_payload


class Themes:
    discovery_workers = 2

    def __init__(self, *, bad_plan=False, reject=None, opening=True):
        self.bad_plan, self.reject, self.opening = bad_plan, reject, opening
        self.calls = []

    def generate_structured(self, messages, schema, **kwargs):
        self.calls.append(schema.__name__)
        payload = evidence_payload(messages)
        if schema is SummaryThemePlan:
            ids = [p['id'] for p in payload['parts']]
            if self.bad_plan:
                ids[-1] = ids[0]
            return schema(themes=[dict(title='Service prerequisites', part_ids=ids[i:i+10])
                                  for i in range(0, len(ids), 10)])
        if schema is CompactSummaryDraft:
            facts = [f for p in payload['parts'] for f in p['facts']]
            quotes = {q['id']: q['text'] for q in payload['source_quotes']}
            names = [(f.get('text') or quotes[f['quote_ref']]).split()[1] for f in facts]
            return schema(pages=[dict(title='Service prerequisites', items=[dict(
                label='Scheduling', text='Teams '+', '.join(names)+' support scheduling; staff training '
                'is required before service delivery.', fact_ids=[f['id'] for f in facts])])])
        assert schema is CompactSummaryReview
        return schema(opening_slide_id=(payload.get('opening_target') or {}).get('slide_id'), opening_item_texts=(payload.get('opening_target') or {}).get('item_texts'), opening_accepted=self.opening, opening_reason='Scheduling and training are explained.', parts=[dict(part_id=p['id'],
            accepted=p['id'] != self.reject, reason='Training dependency is retained.')
            for p in payload['source']['parts']])


def test_complete_introduction_can_exceed_five_pages_with_verified_native_bindings():
    result, review, _ = fixture()
    review.max_intro_slides = None
    company = generate_thematic_summary(Themes(), result, review, CompanyProfile(summary_review=review))
    review.status = 'complete'
    assert len(company.summary_pages) == 7
    assert not presentation_errors(review, company.summary_pages, result)
    result.presentation_plan = PresentationPlan(title='Source report', company=company)
    deck = Presentation(str(BUNDLED_TEMPLATE_PATH))
    for ident in list(deck.slides._sldIdLst):
        deck.part.drop_rel(ident.rId)
        deck.slides._sldIdLst.remove(ident)
    _add_company_at_a_glance(deck, result, PresentationSlide(id='intro', slide_type='company_overview', title='Introduction'))
    stream = BytesIO()
    deck.save(stream)
    assert len(verify_summary_export(Presentation(stream), result)) == 70
    company.summary_pages[1].items[0].text += ' Unreviewed claim.'
    assert any('independent' in e for e in presentation_errors(review, company.summary_pages, result))


def test_explicit_page_budget_still_fails_instead_of_omitting_parts():
    result, review, _ = fixture()
    with pytest.raises(ValueError, match='five-page limit'):
        generate_thematic_summary(Themes(), result, review, CompanyProfile(summary_review=review))


def test_outline_missing_or_duplicate_parts_has_only_one_repair():
    result, review, _ = fixture()
    client = Themes(bad_plan=True)
    with pytest.raises(ValueError, match='exactly once'):
        plan_themes(client, result, review)
    assert client.calls == ['SummaryThemePlan'] * 2


@pytest.mark.parametrize('options', [{'reject': 'part-15'}, {'opening': False}])
def test_meaning_and_business_opening_remain_mandatory(options):
    result, review, _ = fixture(20)
    review.max_intro_slides = None
    with pytest.raises(ValueError):
        generate_thematic_summary(Themes(**options), result, review, CompanyProfile(summary_review=review))
    assert not any('theme_review_proofs' in a for a in review.plan_audits)


def test_cancelled_outline_makes_no_request():
    from concurrent.futures import CancelledError
    result, review, _ = fixture()
    signal = Event()
    signal.set()
    client = Themes()
    with pytest.raises(CancelledError):
        plan_themes(client, result, review, cancelled=signal)
    assert not client.calls


def test_outline_privacy_failure_is_not_retried():
    from adaptive_document_agent.services.llm.exceptions import PrivacyViolationError
    result, review, _ = fixture()
    class Private:
        calls = 0
        def generate_structured(self, *args, **kwargs):
            self.calls += 1
            raise PrivacyViolationError('Local Only')
    client = Private()
    with pytest.raises(PrivacyViolationError):
        plan_themes(client, result, review)
    assert client.calls == 1


def test_whole_item_pagination_preserves_copy_and_evidence():
    from adaptive_document_agent.agent.compact_summary import expand_compact, compact_input
    from adaptive_document_agent.agent.summary_theme_layout import pack_theme_pages
    from adaptive_document_agent.services.compact_summary import compact_layout_errors
    result, review, _ = fixture(4)
    _, catalog = compact_input(review)
    draft = CompactSummaryDraft(pages=[dict(title='Service prerequisites', items=[dict(
        label='Scheduling', text=('排班服务须经客户集成和员工培训后交付。' * 9), fact_ids=[f'f{i}'])
        for i in range(4)])])
    pages = expand_compact(draft, catalog, review)
    packed = pack_theme_pages(pages)
    assert len(packed) > 1 and not compact_layout_errors(packed)
    assert [i.model_dump() for p in packed for i in p.items] == [i.model_dump() for p in pages for i in p.items]
