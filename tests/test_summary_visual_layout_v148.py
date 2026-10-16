"""Visual introductions retain complete authored claims and evidence ownership."""
from copy import deepcopy
from dataclasses import replace
import io

import pytest
from pptx import Presentation
from pptx.util import Pt

from adaptive_document_agent.models.summary_display import SummaryDisplay
from adaptive_document_agent.services.presentation_brief import BriefItem, render_profile
from adaptive_document_agent.services.summary_visual_layout import display_segments
from tests.test_presentation_brief import blank_deck


COPY = (
    'The company develops equipment for shared workspaces and supplies industrial customers. '
    'Its product portfolio comprises four reported series, each serving the stated applications. '
    'The source reports a 13.0% global shipment share in 2023, according to the cited study. '
    'The revenue-based share was 3.6% for that same year, with a different measurement basis. '
    'These measures describe different denominators and cannot be interchanged. '
    'Commercialization remains at an early stage and future performance is uncertain. '
    'All statements retain the stated reporting boundary, attribution and source qualifications.'
)


def bodies(slides):
    return [shape for slide in slides for shape in slide.shapes if shape.name == 'brief:body']


@pytest.mark.parametrize('layout', ['paragraphs', 'comparison', 'timeline', 'relationship', 'key_metrics'])
def test_complete_qualified_copy_stays_editable_in_one_native_slide(layout):
    deck = blank_deck()
    item = BriefItem('Reported operations and scope', COPY, [10, 11],
                     display=SummaryDisplay(layout=layout))
    before = deepcopy(item)
    slides = render_profile(deck, 'Operations and scope', [item])
    assert len(slides) == 1
    assert len(bodies(slides)) > 1
    assert ''.join(shape.text for shape in bodies(slides)) == item.text
    assert item == before
    content = [shape for slide in slides for shape in slide.shapes
               if shape.name in {'brief:body', 'brief:heading'}]
    for shape in content:
        expected = 16 if shape.name == 'brief:body' else 18
        assert all(p.font.size.pt == expected for p in shape.text_frame.paragraphs)
        assert all(p.line_spacing == Pt(expected * 1.25) for p in shape.text_frame.paragraphs)
        assert shape.top.inches + shape.height.inches <= deck.slide_height.inches - 1.28 + .001
    for i, first in enumerate(content):
        for second in content[i + 1:]:
            dx = min(first.left + first.width, second.left + second.width) - max(first.left, second.left)
            dy = min(first.top + first.height, second.top + second.height) - max(first.top, second.top)
            assert dx <= 1 or dy <= 1
    stream = io.BytesIO()
    deck.save(stream)
    reopened = Presentation(io.BytesIO(stream.getvalue()))
    assert ''.join(shape.text for shape in bodies(reopened.slides)) == item.text
    assert not any(shape.shape_type == 13 for slide in reopened.slides for shape in slide.shapes)


@pytest.mark.parametrize('anchors', [[], ['does not exist.'], ['scope.'], ['share'], ['year,'],
    ['source qualifications.', 'industrial customers.']])
def test_invalid_break_hints_fall_back_losslessly_without_editing_copy(anchors):
    display = SummaryDisplay(layout='comparison', break_after=anchors)
    parts = display_segments(COPY, display)
    assert ''.join(parts) == COPY
    assert len(parts) > 1


def test_author_breaks_preserve_every_character_and_parenthetical_qualifications():
    text = ('A reported measure is 1.25 units (the stated boundary excludes missing inputs. '
            'The original limitation still applies). A separate comparison follows.')
    assert display_segments(text) == [text[:text.index(' A separate') + 1], text[text.index('A separate'):]]
    parts = display_segments(COPY, SummaryDisplay(break_after=['industrial customers.', 'early stage and future performance is uncertain.']))
    assert len(parts) == 3 and ''.join(parts) == COPY


def test_overflow_uses_existing_complete_copy_pagination_without_empty_draft_pages():
    item = BriefItem('Reported scope', COPY * 8, [8], display=SummaryDisplay())
    deck = blank_deck()
    slides = render_profile(deck, 'Scope', [item])
    baseline = blank_deck()
    old = render_profile(baseline, 'Scope', [replace(item, display=None)])
    assert len(slides) == len(old)
    assert ''.join(shape.text for shape in bodies(slides)) == item.text
    assert len(deck.slides) == len(slides)


def test_structure_changes_require_a_fresh_independent_part_review():
    from adaptive_document_agent.agent.summary_review_reuse import part_review_signatures
    from adaptive_document_agent.agent.compact_summary import expand_compact
    from tests.test_summary_compression_v145 import compression_fixture
    _, review, catalog, _, draft = compression_fixture()
    before = part_review_signatures(review, expand_compact(draft, catalog, review))
    draft.pages[0].items[0].display = SummaryDisplay(layout='comparison')
    after = part_review_signatures(review, expand_compact(draft, catalog, review))
    assert before != after


def test_display_expansion_retains_every_meaning_address_and_exact_source_binding():
    from adaptive_document_agent.agent.compact_summary import expand_compact
    from tests.test_summary_compression_v145 import compression_fixture
    _, review, catalog, _, draft = compression_fixture()
    before = deepcopy(draft)
    plain = expand_compact(draft, catalog, review)
    draft.pages[0].items[0].display = SummaryDisplay(layout='key_metrics', break_after=['Not a supported anchor.'])
    structured = expand_compact(draft, catalog, review)
    for previous, current in zip(before.pages, draft.pages):
        for first, second in zip(previous.items, current.items):
            assert first.meaning_refs == second.meaning_refs
            assert first.fact_ids == second.fact_ids
            assert first.text == second.text
    for previous, current in zip(plain, structured):
        for first, second in zip(previous.items, current.items):
            assert first.model_dump(exclude={'display'}) == second.model_dump(exclude={'display'})


def test_absent_display_preserves_historical_digest_but_real_structure_changes_it():
    from hashlib import sha256
    import json
    from adaptive_document_agent.agent.compact_summary import expand_compact
    from adaptive_document_agent.services.compact_summary import compact_digest
    from tests.test_summary_compression_v145 import compression_fixture
    _, review, catalog, _, draft = compression_fixture()
    pages = expand_compact(draft, catalog, review)
    historical = [page.model_dump(mode='json') for page in pages]
    for page in historical:
        for item in page['items']:
            item.pop('display')
    expected = sha256(json.dumps(historical, sort_keys=True, ensure_ascii=False).encode('utf-8')).hexdigest()
    assert compact_digest(pages) == expected
    pages[0].items[0].display = SummaryDisplay(layout='comparison')
    assert compact_digest(pages) != expected


def test_budgeted_author_wire_requires_an_explicit_supported_display_choice():
    from pydantic import ValidationError
    from adaptive_document_agent.agent.compact_summary import SummarySynthesisDraft
    from adaptive_document_agent.agent.summary_writer_copy import writer_schema
    schema = writer_schema(SummarySynthesisDraft)
    item_schema = schema.model_json_schema()['$defs']['SummarySynthesisBodyItem']
    assert 'display' in item_schema['required']
    assert 'anyOf' not in item_schema['properties']['display']
    assert 'layout' in schema.model_json_schema()['$defs']['SummaryDisplay']['required']
    value = dict(pages=[dict(title='Source scope', items=[dict(label='Scope', text='The source defines this scope.', fact_ids=['f0'])])])
    with pytest.raises(ValidationError):
        schema.model_validate(value)
    item = value['pages'][0]['items'][0]
    item['display'] = None
    with pytest.raises(ValidationError):
        schema.model_validate(value)
    item['display'] = {'layout': 'paragraphs'}
    assert schema.model_validate(value).pages[0].items[0].display.layout == 'paragraphs'
    # Persisted historical copy remains valid without manufacturing a choice.
    item['display'] = None
    assert SummarySynthesisDraft.model_validate(value).pages[0].items[0].display is None


@pytest.mark.parametrize('layout', ['paragraphs', 'comparison', 'timeline', 'relationship', 'key_metrics'])
def test_multi_item_page_retains_each_complete_copy_and_local_evidence_segments(layout):
    from adaptive_document_agent.models.summary import SummarySlideItem, SummarySlidePage
    from adaptive_document_agent.services.summary_export import bind_summary_pages
    from adaptive_document_agent.services.presentation_brand_qa import validate_generated_brand
    texts = ['The study covers the first cohort. Its scope excludes unmeasured participants.',
             'The second disclosed cohort has a separate measurement boundary. No missing values are inferred.']
    page = SummarySlidePage(id='multiple', title='Study scope', source_section='Scope', items=[
        SummarySlideItem(label=f'Cohort {index}', text=text, source_quote=text,
                         source_pages=[index + 2], part_ids=[f'part-{index}'])
        for index, text in enumerate(texts)])
    items = [BriefItem(item.label, item.text, item.source_pages, display=SummaryDisplay(layout=layout))
             for item in page.items]
    deck = blank_deck()
    slides = render_profile(deck, page.title, items)
    validate_generated_brand(deck)
    assert len(slides) == 1 and len(bodies(slides)) == 4
    assert len([shape for shape in slides[0].shapes if shape.name == 'summary:panel']) == 2
    assert all(shape.width.inches < deck.slide_width.inches / 2 for shape in bodies(slides))
    assert ''.join(shape.text for shape in bodies(slides)) == ''.join(texts)
    bind_summary_pages(slides, page)
    for shape in bodies(slides):
        assert shape.top.inches + shape.height.inches <= deck.slide_height.inches - 1.28 + .001
        assert all(paragraph.font.size.pt == 16 for paragraph in shape.text_frame.paragraphs)
    text = '\n'.join(shape.text for shape in slides[0].shapes if shape.has_text_frame)
    assert 'p. 2-3' in text


def test_panel_capacity_fallback_keeps_all_copy_and_creates_no_partial_groups():
    items = [BriefItem(f'Source scope {index}', COPY * 2, [index + 2], display=SummaryDisplay())
             for index in range(2)]
    deck = blank_deck()
    slides = render_profile(deck, 'Complete source scope', items)
    assert len(deck.slides) == len(slides)
    assert ''.join(shape.text for shape in bodies(slides)) == ''.join(item.text for item in items)
    assert not any(shape.name == 'summary:panel' for slide in slides for shape in slide.shapes)


def test_native_binding_proves_complete_segment_union_and_detects_missing_qualifier():
    from adaptive_document_agent.models.summary import SummarySlideItem, SummarySlidePage
    from adaptive_document_agent.services.summary_export import bind_summary_pages
    page = SummarySlidePage(id='visual', title='Scope', source_section='Scope', items=[SummarySlideItem(
        label='Reported scope', text=COPY, source_quote=COPY, source_pages=[10], part_ids=['source-part'])])
    deck = blank_deck()
    slides = render_profile(deck, page.title, [BriefItem(page.items[0].label, COPY, [10], display=SummaryDisplay())])
    bind_summary_pages(slides, page)
    for index, shape in enumerate(bodies(slides)):
        assert 'source-part' in shape.element.xpath('.//p:cNvPr')[0].get('descr')
        assert f'"segment_index": {index}' in shape.element.xpath('.//p:cNvPr')[0].get('descr')
    bodies(slides)[-1].text = 'A qualifier was removed.'
    with pytest.raises(ValueError, match='complete ordered visible native copy'):
        bind_summary_pages(slides, page)
