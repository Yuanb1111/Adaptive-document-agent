"""Synthetic topic scope collisions must recover without weakening claim gates."""

from __future__ import annotations

import io
import json

import pytest
from pptx import Presentation

from adaptive_document_agent.agent.presentation_plan_recovery import PresentationPlanRecovery
from adaptive_document_agent.agent.presentation_topic_selector import series_directory
from adaptive_document_agent.agent.topic_plan_compiler import compile_topic_plan
from adaptive_document_agent.models import (
    ChartPlan,
    DocumentProfile,
    ParsedDocument,
    PipelineResult,
    PresentationPlan,
    PresentationTopic,
    PresentationTopicSelection,
    ReportPlan,
)
from adaptive_document_agent.services.export import export_pptx
from adaptive_document_agent.services.llm import LLMGateway
from adaptive_document_agent.services.presentation_claim_evidence import prepare_presentation_claims
from adaptive_document_agent.services.qa_reporter import CriticalQAError, run_comprehensive_qa
from adaptive_document_agent.validation.claim_validator import ClaimValidator, repair_presentation_plan
from adaptive_document_agent.validation.presentation_plan_validator import PresentationPlanValidator
from tests.test_segmented_direction_claims import _series, _slide


_SOURCE_FIELDS = {"observations", "presentation_topics", "charts", "insights", "llm_usage"}


@pytest.fixture(autouse=True)
def forbid_model_requests(monkeypatch):
    """Recovery uses retained selections, never another model request."""
    def unexpected_request(*args, **kwargs):
        pytest.fail("Question-first recovery must not make a model request")

    monkeypatch.setattr(LLMGateway, "generate_text", unexpected_request)
    monkeypatch.setattr(LLMGateway, "generate_structured", unexpected_request)


def _synthetic_result(collision: str) -> PipelineResult:
    if collision == "metric_alias":
        groups = [
            _series("Loss for the year", [-140, -180, -230]),
            _series("Loss for the year %", [-8, -10, -12]),
        ]
        question = "How did loss amounts and percentages compare across FY2023 to FY2025?"
        takeaway = "Loss for the year widened from FY2023 to FY2025."
        title = "Reported loss measures"
        for item in groups[0]:
            item.currency, item.raw_unit, item.unit_family = "USD", "USD", "currency"
        percentages = groups[1]
    else:
        groups = []
        for category, values in (("North", [10, 12, 14]), ("South", [20, 23, 25]),
                                 ("West", [30, 28, 25])):
            group = _series("Gross profit margin", values)
            for item in group:
                item.id = f"{category.casefold()}_{item.id}"
                item.category_dimensions = {"region": category}
            groups.append(group)
        question = "How did gross margins vary by region across FY2023 to FY2025?"
        takeaway = "North and South gross margins rose from FY2023 to FY2025 while West gross margin fell."
        title = "Regional gross margins"
        percentages = [item for group in groups for item in group]
    for item in percentages:
        item.unit, item.raw_unit, item.currency = "percent", "%", None
        item.unit_family = "percentage"
        item.raw_value = f"{item.value:g}%"
        item.evidence[0].column_label = "%"

    volume = _series("Units", [14, 18, 22])
    for item in volume:
        item.unit, item.raw_unit, item.currency, item.unit_family = "count", "units", None, "count"
    groups.append(volume)
    result = PipelineResult(
        document=ParsedDocument(document_id="synthetic-scope", sha256="synthetic-scope",
                                safe_filename="synthetic_scope.pdf", page_count=10),
        profile=DocumentProfile(document_type="Document", document_purpose="Review reported synthetic measures.",
                                analysis_page_ranges=[(8, 8)]),
        report_plan=ReportPlan(title="Synthetic evidence review"),
        observations=[item for group in groups for item in group],
        charts=[ChartPlan(
            id=f"synthetic-chart-{index}", title=group[0].metric_original,
            chart_type="line", available_chart_types=["line", "bar", "table"],
            question="How did the reported values vary?",
            observation_ids=[item.id for item in group], source_pages=[8],
            x_axis_title="Period", y_axis_title=group[0].raw_unit,
        ) for index, group in enumerate(groups)],
    )
    _, series = series_directory(result)
    affected_ids = {item.id for group in groups[:-1] for item in group}
    affected_series = [sid for sid, members in series.items() if {item.id for item in members} <= affected_ids]
    volume_series = [sid for sid, members in series.items() if {item.id for item in members} == {item.id for item in volume}]
    assert len(affected_series) == len(groups) - 1
    assert len(volume_series) == 1
    result.presentation_topics = PresentationTopicSelection(topics=[
        PresentationTopic(id="scope", title=title, question=question, takeaway=takeaway,
                          rationale="Retain the separately reported measures and their source scopes.",
                          series_ids=affected_series),
        PresentationTopic(id="volume", title="Reported volume",
                          question="How did units vary over the reported periods?",
                          takeaway="Units rose from FY2023 to FY2025.",
                          rationale="The reported units support a comparable-period review.",
                          series_ids=volume_series),
    ])
    return result


def _analysis(plan: PresentationPlan, theme_id: str = "scope"):
    return next(slide for slide in plan.slides if slide.slide_type == "analysis" and slide.theme_id == theme_id)


def _assert_question_first(result: PipelineResult, plan: PresentationPlan) -> None:
    topic, unaffected = result.presentation_topics.topics
    analysis = _analysis(plan)
    summary = next(slide for slide in plan.slides if slide.slide_type == "executive_summary")
    assert analysis.title == topic.question
    assert analysis.message == topic.question
    assert analysis.analytical_question == topic.question
    assert analysis.selection_reason == topic.rationale
    assert summary.bullets == [topic.question, unaffected.takeaway]
    assert _analysis(plan, "volume").title == unaffected.takeaway
    assert {theme.id for theme in plan.themes} == {topic.id, unaffected.id}
    assert plan.editorial_status == ("degraded" if plan.planning_origin == "topic_recovery" else "needs_review")
    assert any(f"Authored takeaway: {topic.takeaway}" in note for note in plan.editorial_notes)
    assert not ClaimValidator().validate_plan(plan, result.observations, result.charts)
    PresentationPlanValidator().validate(plan, result)
    selected_ids = {item.id for item in result.observations if not item.id.startswith("units_")}
    theme = next(theme for theme in plan.themes if theme.id == topic.id)
    assert set(theme.observation_ids) == selected_ids
    assert set(analysis.chart_ids) == {chart.id for chart in result.charts if set(chart.observation_ids) <= selected_ids}
    assert set(summary.bullet_observation_ids[0]) == selected_ids


def _stale_complete_cache(result: PipelineResult) -> None:
    """Reconstruct an old full-coverage fallback without calling unsafe code."""
    seed = result.model_copy(deep=True)
    seed.presentation_topics.topics[0].takeaway = seed.presentation_topics.topics[0].question
    plan = PresentationPlanRecovery().from_selected_topics(seed)
    _analysis(plan).title = result.presentation_topics.topics[0].takeaway
    summary = next(slide for slide in plan.slides if slide.slide_type == "executive_summary")
    summary.bullets[0] = result.presentation_topics.topics[0].takeaway
    plan.planning_origin = "topic_recovery"
    plan.editorial_status = "needs_review"
    result.presentation_plan = plan
    assert {theme.id for theme in plan.themes} == {topic.id for topic in result.presentation_topics.topics}
    assert any(issue.code == "direction_scope_ambiguous"
               for issue in ClaimValidator().validate_plan(plan, result.observations, result.charts))


@pytest.mark.parametrize("collision", ["metric_alias", "shared_category"])
def test_collision_is_still_a_strict_direction_scope_error(collision):
    result = _synthetic_result(collision)
    topic = result.presentation_topics.topics[0]
    observations = [item for item in result.observations if not item.id.startswith("units_")]
    issues = ClaimValidator().validate_slide(_slide(topic.takeaway, observations, "title"), observations)
    assert any(issue.code == "direction_scope_ambiguous" and issue.severity == "error" for issue in issues)
    assert all("explicit period" not in issue.message for issue in issues)


@pytest.mark.parametrize("collision", ["metric_alias", "shared_category"])
@pytest.mark.parametrize("entrypoint", ["compile", "fallback"])
def test_selected_scope_collision_retains_question_and_all_original_evidence(collision, entrypoint):
    result = _synthetic_result(collision)
    original = result.model_dump(include=_SOURCE_FIELDS)
    plan = (compile_topic_plan(result) if entrypoint == "compile"
            else PresentationPlanRecovery().from_selected_topics(result))
    _assert_question_first(result, plan)
    assert result.model_dump(include=_SOURCE_FIELDS) == original
    result.presentation_plan = plan
    for _ in range(2):
        prepare_presentation_claims(result)
        qa = run_comprehensive_qa(result, auto_repair=True)
        assert not qa.has_critical_errors, [item.message for item in qa.critical_errors]
        _assert_question_first(result, result.presentation_plan)
        assert result.model_dump(include=_SOURCE_FIELDS) == original


@pytest.mark.parametrize("collision", ["metric_alias", "shared_category"])
def test_full_coverage_degraded_cache_is_repaired_once_without_a_model(collision):
    result = _synthetic_result(collision)
    _stale_complete_cache(result)
    original = result.model_dump(include=_SOURCE_FIELDS)
    company = result.presentation_plan.company.model_dump()
    recovery = PresentationPlanRecovery()
    assert recovery.recover_missing_plan(result)
    _assert_question_first(result, result.presentation_plan)
    assert result.presentation_plan.company.model_dump() == company
    assert result.model_dump(include=_SOURCE_FIELDS) == original
    first = result.model_dump()
    assert not recovery.recover_missing_plan(result)
    assert result.model_dump() == first


@pytest.mark.parametrize("collision", ["metric_alias", "shared_category"])
@pytest.mark.parametrize("cached", [False, True])
def test_public_export_keeps_recovered_question_and_editable_evidence(collision, cached, local_render_stub):
    result = _synthetic_result(collision)
    if cached:
        _stale_complete_cache(result)
    original = result.model_dump(include=_SOURCE_FIELDS)
    deck = Presentation(io.BytesIO(export_pptx(result)))
    _assert_question_first(result, result.presentation_plan)
    visible = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert result.presentation_topics.topics[0].question in visible
    assert result.presentation_topics.topics[0].takeaway not in visible
    assert any(shape.has_chart for slide in deck.slides for shape in slide.shapes)
    assert result.model_dump(include=_SOURCE_FIELDS) == original
    snapshot = result.model_dump()
    export_pptx(result)
    assert result.model_dump() == snapshot


@pytest.mark.parametrize("question", [
    "Did loss for the year widen from FY2023 to FY2025?",
    "Did loss for the year reach 999999 from FY2023 to FY2025?",
])
@pytest.mark.parametrize("entrypoint", ["compile", "fallback"])
def test_unsafe_question_cannot_be_used_as_a_validation_bypass(question, entrypoint):
    result = _synthetic_result("metric_alias")
    result.presentation_topics.topics[0].question = question
    original = result.model_dump(include=_SOURCE_FIELDS)
    with pytest.raises(ValueError):
        if entrypoint == "compile":
            compile_topic_plan(result)
        else:
            PresentationPlanRecovery().from_selected_topics(result)
    assert result.model_dump(include=_SOURCE_FIELDS) == original


def test_failed_cached_recovery_stays_transactional_and_export_stays_blocked(local_render_stub):
    result = _synthetic_result("metric_alias")
    _stale_complete_cache(result)
    result.presentation_topics.topics[0].question = "Did loss for the year widen from FY2023 to FY2025?"
    _analysis(result.presentation_plan).message = result.presentation_topics.topics[0].question
    _analysis(result.presentation_plan).analytical_question = result.presentation_topics.topics[0].question
    original = result.model_dump(include=_SOURCE_FIELDS)
    plan = result.presentation_plan.model_dump()
    assert not PresentationPlanRecovery().recover_missing_plan(result)
    assert result.presentation_plan.model_dump() == plan
    assert result.model_dump(include=_SOURCE_FIELDS) == original
    with pytest.raises(CriticalQAError):
        export_pptx(result)
    assert result.model_dump(include=_SOURCE_FIELDS) == original
    assert any(issue.code == "presentation_cached_recovery_failed" for issue in result.validation_warnings)


def test_real_direction_contradiction_remains_detected_and_repaired():
    observations = _series("Net loss", [-90, -110, -160])
    wrong = "Net loss narrowed from FY2023 to FY2025."
    slide = _slide(wrong, observations, "title")
    issues = ClaimValidator().validate_slide(slide, observations)
    assert any(issue.code == "directional_contradiction" for issue in issues)
    plan, repairs = repair_presentation_plan(PresentationPlan(title="Synthetic review", slides=[slide]), observations)
    assert repairs
    assert plan.slides[0].title == "Net loss widened from FY2023 to FY2025."
    assert not ClaimValidator().validate_plan(plan, observations)


def test_long_takeaway_does_not_revive_when_shortened_title_already_equals_question(local_render_stub):
    result = _synthetic_result("metric_alias")
    topic = result.presentation_topics.topics[0]
    topic.title = topic.question
    topic.takeaway = (
        "Loss for the year widened from FY2023 to FY2025 across the retained "
        "synthetic reporting periods shown in the cited source records."
    )
    assert len(topic.takeaway.split()) > 18
    original = result.model_dump(include=_SOURCE_FIELDS)
    result.presentation_plan = compile_topic_plan(result)
    _assert_question_first(result, result.presentation_plan)
    markers = [note for note in result.presentation_plan.editorial_notes
               if note.startswith("[claim_scope_withheld] ")]
    assert len(markers) == 1
    assert json.loads(markers[0].split(" ", 1)[1])["takeaway"] == topic.takeaway
    for _ in range(2):
        prepare_presentation_claims(result)
        assert not run_comprehensive_qa(result).has_critical_errors
        _assert_question_first(result, result.presentation_plan)
    export_pptx(result)
    _assert_question_first(result, result.presentation_plan)
    assert result.model_dump(include=_SOURCE_FIELDS) == original


def test_cached_retry_for_another_topic_preserves_prior_withheld_audit():
    result = _synthetic_result("metric_alias")
    topic, other = result.presentation_topics.topics
    topic.title = topic.question
    topic.takeaway += " The retained synthetic source records show the reported measures separately for every cited reporting period."
    result.presentation_plan = compile_topic_plan(result)
    marker = next(note for note in result.presentation_plan.editorial_notes
                  if note.startswith("[claim_scope_withheld] "))
    result.presentation_plan.planning_origin = "topic_recovery"
    other.takeaway = "Units rose through FY2025."
    _analysis(result.presentation_plan, "volume").title = other.takeaway
    next(slide for slide in result.presentation_plan.slides
         if slide.slide_type == "executive_summary").bullets[1] = other.takeaway
    original = result.model_dump(include=_SOURCE_FIELDS)
    assert PresentationPlanRecovery().recover_missing_plan(result)
    assert marker in result.presentation_plan.editorial_notes
    prepare_presentation_claims(result)
    assert not run_comprehensive_qa(result).has_critical_errors
    summary = next(slide for slide in result.presentation_plan.slides if slide.slide_type == "executive_summary")
    assert summary.bullets == [topic.question, other.question]
    assert _analysis(result.presentation_plan).title == topic.question
    assert _analysis(result.presentation_plan, "volume").title == other.question
    assert result.model_dump(include=_SOURCE_FIELDS) == original
    stable = result.model_dump()
    assert not PresentationPlanRecovery().recover_missing_plan(result)
    assert result.model_dump() == stable


@pytest.mark.parametrize("overview_state", ["missing", "deep_dive"])
def test_summary_rebuild_cannot_revive_withheld_claim_without_an_overview(overview_state):
    from adaptive_document_agent.agent.presentation_summary_selection import rebuild_selected_topic_summary

    result = _synthetic_result("metric_alias")
    result.presentation_plan = compile_topic_plan(result)
    plan = result.presentation_plan
    topic = result.presentation_topics.topics[0]
    analysis = _analysis(plan)
    assert any(note.startswith("[claim_scope_withheld] ") for note in plan.editorial_notes)
    if overview_state == "missing":
        plan.slides = [slide for slide in plan.slides if slide.id != analysis.id]
    else:
        analysis.slide_role = "deep_dive"
    summary = next(slide for slide in plan.slides if slide.slide_type == "executive_summary")
    summary.bullets[0] = topic.takeaway
    original = result.model_dump(include=_SOURCE_FIELDS)
    assert rebuild_selected_topic_summary(result, plan)
    assert summary.bullets[0] == topic.question
    assert topic.takeaway not in summary.bullets
    assert result.model_dump(include=_SOURCE_FIELDS) == original
    stable = plan.model_dump()
    assert rebuild_selected_topic_summary(result, plan)
    assert plan.model_dump() == stable


def test_summary_only_ambiguity_is_preserved_in_withheld_audit_errors():
    result = _synthetic_result("metric_alias")
    topic = result.presentation_topics.topics[0]
    topic.title = topic.question
    topic.takeaway = (
        "Loss for the year widened from FY2023 to FY2025 across the retained "
        "synthetic reporting periods shown in the cited source records."
    )
    assert len(topic.takeaway.split()) > 18
    plan = compile_topic_plan(result)
    _assert_question_first(result, plan)
    entry = json.loads(next(note for note in plan.editorial_notes
                            if note.startswith("[claim_scope_withheld] ")).split(" ", 1)[1])
    assert entry["topic_id"] == topic.id
    assert entry["takeaway"] == topic.takeaway
    assert entry["errors"]
    assert any("slide_executive_summary" in message and "no unique source metric" in message
               for message in entry["errors"])


def test_nonnumeric_research_question_with_direction_words_cannot_drop_topics():
    result = _synthetic_result('metric_alias')
    topic = result.presentation_topics.topics[1]
    topic.takeaway = 'Units rose through FY2025.'
    topic.question = 'How did Units increase over the reported periods?'
    original = result.model_dump(include=_SOURCE_FIELDS)
    plan = compile_topic_plan(result)
    assert {t.id for t in plan.themes} == {'scope', 'volume'}
    analysis = _analysis(plan, 'volume')
    assert analysis.analytical_question == topic.question
    assert analysis.title in {topic.title, topic.question}
    assert not ClaimValidator().validate_plan(plan, result.observations, result.charts)
    assert any(topic.takeaway in note for note in plan.editorial_notes)
    assert result.model_dump(include=_SOURCE_FIELDS) == original


def test_local_withdrawal_keeps_unrelated_approved_summary_copy():
    from adaptive_document_agent.agent.presentation_topic_scope_recovery import recover_unscoped_topic_claims
    result = _synthetic_result('metric_alias')
    plan = compile_topic_plan(result)
    topic, other = result.presentation_topics.topics
    other.takeaway = 'Units rose.'
    summary = next(s for s in plan.slides if s.slide_type == 'executive_summary')
    summary.bullets = [topic.takeaway, other.title]
    _analysis(plan).title = topic.takeaway
    issues = ClaimValidator().validate_plan(plan, result.observations, result.charts)
    assert issues
    original = result.model_dump(include=_SOURCE_FIELDS)
    assert recover_unscoped_topic_claims(plan, result, issues)
    assert next(s for s in plan.slides if s.id == summary.id).bullets == [topic.question, other.title]
    assert not ClaimValidator().validate_plan(plan, result.observations, result.charts)
    assert result.model_dump(include=_SOURCE_FIELDS) == original
