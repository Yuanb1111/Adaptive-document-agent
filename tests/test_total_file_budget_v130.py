"""Complete reading is retained while a reviewed audience deck remains bounded."""
from tests.introduction_review_helpers import review_response
from io import BytesIO
from copy import deepcopy
from types import SimpleNamespace
import json

import pytest
from pptx import Presentation

from adaptive_document_agent.agent.introduction_audience import select_introduction_audience
from adaptive_document_agent.agent.compact_summary import generate_compact_summary
from adaptive_document_agent.document_model import DocumentIndex
from adaptive_document_agent.models import PresentationPlan, PresentationSlide
from adaptive_document_agent.models.presentation import CompanyProfile
from adaptive_document_agent.models.customization import PresentationBudget
from adaptive_document_agent.services.introduction_audience import audience_part_ids
from adaptive_document_agent.services.presentation_requirements import apply_presentation_budget
from adaptive_document_agent.services.slide_budgets import require_exported_budgets
from adaptive_document_agent.services.summary_validation import presentation_errors
from adaptive_document_agent.services.summary_export import verify_summary_export
from adaptive_document_agent.services.pptx_export import _add_company_at_a_glance, BUNDLED_TEMPLATE_PATH
from tests.test_compact_summary_v111 import fixture, Reviewer


def choose(result, review, selected=('part-0',), *, reject_first=False):
    class Client:
        calls = []
        def generate_structured(self, messages, schema, **kwargs):
            self.calls.append(schema.__name__)
            if schema.__name__ == 'IntroductionAudienceSelection':
                return schema(decisions=[dict(part_id=p.id, include=p.id in selected,
                    reason='Shared service and training condition are represented by the selected source.'
                    if p.id not in selected else 'Core service and required training belong in the introduction.',
                    required_meanings=[dict(meaning=p.reading_note, source_quote=p.source_text.strip())]
                    if p.id in selected else [])
                    for p in review.parts if p.role == 'content'])
            return review_response(schema, messages, accepted=not reject_first or self.calls.count(schema.__name__) > 1,
                          reason='The selected union retains service scope and training dependency.')
    client = Client()
    select_introduction_audience(client, result, review)
    return client


def test_selection_keeps_entire_reading_and_requires_independent_approval():
    result, review, _ = fixture(3)
    original = review.model_dump()['parts']
    client = choose(result, review, reject_first=True)
    assert client.calls == ['IntroductionAudienceSelection', 'IntroductionAudienceReview'] * 2
    assert review.model_dump()['parts'] == original
    assert audience_part_ids(review) == {'part-0'}
    assert len(review.audience_selection['decisions']) == 3


def test_reviewer_rejection_of_missing_material_condition_cannot_be_bypassed():
    result, review, _ = fixture(3)
    class Client:
        def generate_structured(self, messages, schema, **kwargs):
            if schema.__name__ == 'IntroductionAudienceSelection':
                return schema(decisions=[dict(part_id=p.id, include=p.id == 'part-0', reason='Proposed relevance decision.',
                    required_meanings=[dict(meaning=p.reading_note, source_quote=p.source_text.strip())]
                    if p.id == 'part-0' else [])
                    for p in review.parts])
            return review_response(schema, messages, accepted=False, reason='A material delivery condition is missing from the selected union.')
    with pytest.raises(ValueError, match='material delivery condition'):
        select_introduction_audience(Client(), result, review)
    assert review.audience_part_ids is None and not review.audience_selection


def test_truncated_selection_regenerates_once_from_full_source_with_bounded_capacity():
    from adaptive_document_agent.agent.introduction_audience import complete_response, IntroductionAudienceReview
    from adaptive_document_agent.services.llm.exceptions import LLMStructuredOutputError
    from adaptive_document_agent.services.llm.base import LLMResponse
    from adaptive_document_agent.services.llm.usage import LLMUsage
    messages = [{'role':'system','content':'Review the entire supplied source.'},
                {'role':'user','content':'Full source; a truncated response must not replace it.'}]
    class Client:
        calls = []
        def generate_structured(self, request, schema, **kwargs):
            self.calls.append((deepcopy(request), kwargs))
            if len(self.calls) == 1:
                raise LLMStructuredOutputError('Truncated output', response=LLMResponse(text='{"accepted":',
                    usage=LLMUsage(provider='mock', model='mock', finish_reason='length')))
            return schema(accepted=True, reason='Complete source conditions are retained.')
    client = Client()
    assert complete_response(client, messages, IntroductionAudienceReview).accepted
    assert [kwargs['max_tokens'] for _, kwargs in client.calls] == [32000, 40000]
    assert all(request == messages and kwargs['allow_repair'] is False for request, kwargs in client.calls)


@pytest.mark.parametrize('mutation', ['source', 'missing', 'duplicate', 'review', 'unknown', 'reason'])
def test_unreviewed_stale_or_incomplete_selection_is_rejected(mutation):
    result, review, _ = fixture(3)
    choose(result, review)
    if mutation == 'source': review.parts[1].source_text += ' Essential condition.'
    elif mutation == 'missing': review.audience_selection['decisions'].pop()
    elif mutation == 'duplicate': review.audience_part_ids.append('part-0')
    elif mutation == 'review': review.audience_selection['independent_review']['accepted'] = False
    elif mutation == 'unknown': review.audience_part_ids.append('other-source')
    elif mutation == 'reason': review.audience_selection['decisions'][1]['reason'] = ' '
    with pytest.raises(ValueError, match='complete current source'):
        audience_part_ids(review)


def test_native_export_binds_every_selected_part_without_reclassifying_omitted_content():
    result, review, _ = fixture(3)
    choose(result, review)
    draft = dict(pages=[dict(title='Scheduling service', items=[dict(label='Service scope',
        text=review.parts[0].source_text.strip(), fact_ids=['f0'])])])
    company = generate_compact_summary(Reviewer(draft), result, review, CompanyProfile(summary_review=review))
    review.status = 'complete'
    assert not presentation_errors(review, company.summary_pages, result)
    assert all(p.role == 'content' for p in review.parts)
    assert 'training' in review.decisions[1].reason
    result.presentation_plan = PresentationPlan(title='Source deck', company=company)
    deck = Presentation(str(BUNDLED_TEMPLATE_PATH))
    for ident in list(deck.slides._sldIdLst):
        deck.part.drop_rel(ident.rId); deck.slides._sldIdLst.remove(ident)
    _add_company_at_a_glance(deck, result, PresentationSlide(id='company', slide_type='company_overview', title='Introduction'))
    stream = BytesIO(); deck.save(stream)
    reopened = Presentation(stream)
    assert set(verify_summary_export(reopened, result)) == {'part-0'}
    assert len(reopened.slides) == 1 and len(review.parts) == 3
    # Claim/evidence checks still apply to the smaller selected union.
    company.summary_pages[0].items[0].text += ' Revenue was RMB999 million.'
    assert any('unsupported numeric' in error for error in presentation_errors(review, company.summary_pages, result))


def test_added_source_part_cannot_bypass_the_audience_and_part_meaning_review():
    from adaptive_document_agent.models.summary import SummaryFactRef
    result, review, _ = fixture(3)
    choose(result, review)
    draft = dict(pages=[dict(title='Scheduling service', items=[dict(label='Service scope',
        text=review.parts[0].source_text.strip(), fact_ids=['f0'])])])
    company = generate_compact_summary(Reviewer(draft), result, review, CompanyProfile(summary_review=review))
    item = company.summary_pages[0].items[0]
    item.part_ids.append('part-1')
    item.source_refs.append(SummaryFactRef(part_id='part-1', fact_index=0))
    item.source_quote += '\n\n' + review.parts[1].facts[0].source_quote
    review.decisions[1].decision = 'include'
    errors = presentation_errors(review, company.summary_pages, result, verify_compact_review=False)
    assert any('outside its independently reviewed audience selection' in e for e in errors)


def test_default_budget_counts_appendices_and_analysis_continuations_after_roundtrip():
    result, _, _ = fixture(1)
    result.profile.report_requirements = apply_presentation_budget(None, '', PresentationBudget(body_slides=20, introduction_slides=5))
    result.presentation_plan = PresentationPlan(title='Deck', slides=[PresentationSlide(id='charts', slide_type='analysis', title='Charts')])
    deck = Presentation()
    for n in range(21):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        if n >= 19: slide.element.cSld.set('name', 'ada:data-index')
    stream = BytesIO(); deck.save(stream)
    with pytest.raises(ValueError, match='report all has 21'):
        require_exported_budgets(Presentation(stream), result)
    result.profile.report_requirements = apply_presentation_budget(None, '', PresentationBudget(body_slides=22, introduction_slides=5))
    result.presentation_export_trace = [dict(planned_slide_id='charts', slide_numbers=list(range(3,14)))]
    with pytest.raises(ValueError, match='analysis all has 11'):
        require_exported_budgets(Presentation(stream), result)
    result.profile.report_requirements = apply_presentation_budget(None, '', PresentationBudget(body_slides=20, introduction_slides=5, slide_scope='body'))
    result.presentation_export_trace = []
    require_exported_budgets(Presentation(stream), result)


def test_data_index_retains_matrix_sources_and_exact_raw_values_without_repeating_tables():
    from adaptive_document_agent.services.presentation_data_index import index_entries, render_data_index
    from tests.test_p0_composition import observation
    result, _, _ = fixture(1)
    obs = observation('raw', 'Output', 100, 'FY2023', unit='count')
    obs.raw_value = '100.00'
    obs.raw_unit = 'units'
    result.observations = [obs]
    result.presentation_plan = PresentationPlan(title='Deck', slides=[PresentationSlide(id='matrix', slide_type='analysis', title='Output mix')])
    result.presentation_export_trace = [dict(planned_slide_id='matrix', slide_numbers=[4], observation_ids=['raw'], chart_ids=[], source_pages=[1])]
    original = deepcopy(result.observations)
    deck = Presentation(str(BUNDLED_TEMPLATE_PATH))
    for ident in list(deck.slides._sldIdLst):
        deck.part.drop_rel(ident.rId); deck.slides._sldIdLst.remove(ident)
    entries = index_entries(result, DocumentIndex(result.observations))
    pages = render_data_index(deck, result, DocumentIndex(result.observations))
    assert len(pages) == 1 and entries[0]['observation_ids'] == ['raw']
    notes = json.loads(pages[0].notes_slide.notes_text_frame.text)
    assert notes['data_index'][0]['observations'][0]['raw_value'] == '100.00'
    visible = ' '.join(s.text for s in pages[0].shapes if s.has_text_frame)
    assert all(t in visible for t in ['Output mix', 'units', 'FY2023', 'Data Index'])
    assert result.observations == original


def test_budgeted_native_structure_has_one_final_conclusion_and_no_duplicate_opening_summary():
    from tests.test_executive_brief import result_for, gateway, payload
    from tests.test_p0_composition import observation
    from adaptive_document_agent.agent.executive_brief import ExecutiveBriefWriter
    from adaptive_document_agent.services.pptx_export import _build_planned_presentation
    from adaptive_document_agent.services.presentation_sparse_pages import fold_sparse_text_pages
    from adaptive_document_agent.models import ChartPlan
    source = 'Cedar provides scheduling software. Revenue was 100 units in 2023 and 120 units in 2024.'
    result = result_for([source])
    result.profile.document_summary = 'Cedar provides scheduling software.'
    result.profile.document_summary_pages = [1]
    result.observations = [observation('a', 'Revenue', 100, '2023', unit='count'),
                           observation('b', 'Revenue', 120, '2024', unit='count')]
    result.charts = [ChartPlan(id='chart', title='Revenue', chart_type='line',
        question='Reported revenue', observation_ids=['a', 'b'], source_pages=[1])]
    result.presentation_plan = PresentationPlan(title='Review', slides=[
        PresentationSlide(id='cover', slide_type='cover', title='Review'),
        PresentationSlide(id='intro', slide_type='company_overview', title='Company'),
        PresentationSlide(id='brief', slide_type='executive_summary', title='Old summary'),
        PresentationSlide(id='chartpage', slide_type='analysis', title='Revenue', chart_ids=['chart']),
        PresentationSlide(id='closing', slide_type='risks', title='Closing findings'),
        PresentationSlide(id='index', slide_type='appendix', title='Data Index')])
    g, _ = gateway([payload('Cedar provides scheduling software.', label='Business')])
    result.executive_brief = ExecutiveBriefWriter(g).generate(result)
    result.profile.report_requirements = apply_presentation_budget(None, '', PresentationBudget(body_slides=20, introduction_slides=5))
    deck = Presentation(str(BUNDLED_TEMPLATE_PATH))
    for ident in list(deck.slides._sldIdLst):
        deck.part.drop_rel(ident.rId); deck.slides._sldIdLst.remove(ident)
    _build_planned_presentation(deck, result)
    fold_sparse_text_pages(deck, result)
    titles = [next((s.text for s in slide.placeholders if s.has_text_frame and s.placeholder_format.idx in (14,15)), '')
              for slide in deck.slides]
    assert titles[-1] == 'Conclusion' and any(t == 'Data Index' for t in titles)
    assert titles.index('Data Index') < titles.index('Conclusion')
    assert 'Old summary' not in titles and 'Key takeaways' not in titles and 'Key figures' not in titles
    assert len([s for s in deck.slides if s.element.cSld.get('name') == 'ada:conclusion']) == 1
    assert result.executive_brief.items[0].text in deck.slides[-1].notes_slide.notes_text_frame.text
    require_exported_budgets(deck, result)


def test_data_index_coalesces_only_same_authored_topic_and_preserves_all_raw_records():
    from adaptive_document_agent.services.presentation_data_index import index_entries
    from tests.test_p0_composition import observation
    result,_,_=fixture(1)
    a=observation('annual','Output',100,'FY2023',unit='count')
    b=observation('recent','Output',55,'6M2024',unit='count')
    result.observations=[a,b]
    result.presentation_plan=PresentationPlan(title='Deck',slides=[
        PresentationSlide(id='first',slide_type='analysis',title='Annual view',section_title='Output',theme_id='output'),
        PresentationSlide(id='second',slide_type='analysis',title='Recent view',section_title='Output',theme_id='output')])
    result.presentation_export_trace=[dict(planned_slide_id=s.id,slide_numbers=[n],observation_ids=[o.id],source_pages=[1])
        for s,n,o in zip(result.presentation_plan.slides,[5,6],[a,b])]
    entries=index_entries(result,DocumentIndex(result.observations))
    assert len(entries)==1 and entries[0]['slides']==[5,6]
    assert entries[0]['periods']==['FY2023','6M2024']
    assert [o['raw_value'] for o in entries[0]['observations']]==[a.raw_value,b.raw_value]
    result.presentation_plan.slides[1].theme_id='independent'
    assert len(index_entries(result,DocumentIndex(result.observations)))==2
