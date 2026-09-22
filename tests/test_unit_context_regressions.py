"""Source-generic regressions for unit, hierarchy and cache failures in exports."""

import pytest
from pptx import Presentation

from adaptive_document_agent.document_model import DocumentIndex, group_comparable_series
from adaptive_document_agent.extraction.observation_extractor import ObservationExtractor
from adaptive_document_agent.models import ChartPlan, DocumentProfile, ParsedDocument, PipelineResult
from adaptive_document_agent.models.table import ExtractedTable, TableRow
from adaptive_document_agent.services.composition_candidates import reported_composition_charts
from adaptive_document_agent.services.export import _build_cache_key
from adaptive_document_agent.services.movement_formatter import FinancialMovementFormatter as Formatter
from adaptive_document_agent.services.pptx_export import _add_evidence_table_slides
from adaptive_document_agent.validation.claim_validator import are_observations_compatible
from adaptive_document_agent.validation.presentation_data_qa import validate_presentation_data
from tests.test_presentation_accuracy_regressions import obs


def result_with(values, charts=()):
    return PipelineResult(document=ParsedDocument(document_id="test", sha256="a"*64, safe_filename="source.pdf", page_count=3),
        profile=DocumentProfile(), observations=values, charts=list(charts))


@pytest.mark.parametrize("start,end", [(21000, 2390000), (10000, 20000), (-21000, -2390000), (0, 21000)])
def test_normalized_observations_never_rescaled_by_chart_unit(start, end):
    a, b = obs("a", start, "2023-12-31", "Bills receivables"), obs("b", end, "2024-06-30", "Bills receivables")
    actual = Formatter.format_movement_narrative(a, b, scale=1e6, unit="USD millions")
    expected = Formatter.format_movement_narrative(a, b)
    assert actual == expected
    assert "bn" not in actual


def test_input_scale_is_explicit_not_guessed_from_magnitude():
    assert "bn" not in Formatter.format_movement("Loss", -1.57, -.83)
    assert "0.74bn" in Formatter.format_movement("Loss", -1.57, -.83, scale=1e9)
    assert Formatter._format_currency_value(21000, currency="USD", scale=1, unit="currency") == "US$ 21k"
    assert Formatter._format_currency_value(210000, currency="USD", scale=1000, unit="currency") == "US$ 210m"


def test_observation_narrative_does_not_invent_missing_currency():
    a, b = obs("a", 100, "FY2023", "Output"), obs("b", 120, "FY2024", "Output")
    for item in (a, b):
        item.currency = None
        item.unit = "count"
    text = Formatter.format_movement_narrative(a, b)
    assert "RMB" not in text and "US$" not in text
    assert "100.0" in text and "120.0" in text


@pytest.mark.parametrize("axis", [None, "period"])
def test_same_label_in_different_source_contexts_does_not_form_one_series(axis):
    values = [obs(str(i), i+10, f"FY{2020+i}", "Raw materials") for i in range(4)]
    for i, value in enumerate(values):
        value.dimensions["table_context"] = "Production costs" if i < 2 else "Inventory balances"
    assert len(group_comparable_series(values)) == 2
    assert not are_observations_compatible(values[0], values[-1])[0]
    chart = ChartPlan(id="mixed", title="Raw materials", chart_type="line", question="Trend", x_dimension=axis, observation_ids=[o.id for o in values])
    assert any(issue.code == "mixed_metric_context" for issue in validate_presentation_data(result_with(values, [chart])))


def test_group_total_closes_hierarchy_without_losing_values():
    rows = [
        ["Non-current liabilities", None, None],
        ["Deferred income", "80", "90"],
        ["Lease liabilities", "20", "30"],
        ["Total non-current liabilities", "100", "120"],
        ["Net assets", "300", "400"],
        ["Equity", None, None],
        ["Share capital", "300", "400"],
    ]
    table = ExtractedTable(table_id="table", page=3, headers=["", "2023", "2024"],
        column_periods=[None, "2023-12-31", "2024-06-30"], column_types=["label", "amount", "amount"],
        default_unit="currency", default_currency="USD", default_unit_scale=1000,
        rows=[TableRow(cells=row, page=3) for row in rows])
    values = ObservationExtractor()._table_observations(table)
    assert len(values) == 10
    net = [o for o in values if o.metric_original == "Net assets"]
    assert [o.value for o in net] == [300000, 400000]
    assert all(not o.category_dimensions for o in net)
    charts = reported_composition_charts(DocumentIndex(values))
    liability_chart = next(c for c in charts if c.title == "Non-current liabilities")
    assert len(liability_chart.observation_ids) == 4
    assert len(liability_chart.total_observation_ids) == 2


def test_composition_is_checked_even_without_presentation_plan():
    values = [obs(f"{p}-{c}", v, p, category=c) for p in ("FY2023", "FY2024")
              for c, v in (("East", 40), ("West", 60), ("Total resources", 100))]
    chart = ChartPlan(id="bad", title="Resources", question="Mix", chart_type="stacked_bar",
        series_dimension="component", observation_ids=[o.id for o in values])
    assert any(i.code == "invalid_composition_chart" for i in validate_presentation_data(result_with(values, [chart])))


def test_appendix_preserves_categories_and_rejects_overwriting_conflicts():
    values = [obs("a", 168e6, category="Total liabilities"), obs("b", 315e6, category="Net assets")]
    deck = Presentation()
    _add_evidence_table_slides(deck, result_with(values), [])
    labels = [row.cells[0].text for slide in deck.slides for shape in slide.shapes if shape.has_table for row in shape.table.rows]
    assert any("Total liabilities" in label for label in labels)
    assert any("Net assets" in label for label in labels)
    values[1].category_dimensions = values[0].category_dimensions.copy()
    with pytest.raises(ValueError, match="Conflicting appendix"):
        _add_evidence_table_slides(Presentation(), result_with(values), [])


def test_build_cache_includes_loaded_pipeline_contract(monkeypatch):
    from adaptive_document_agent.utils import pipeline_version
    result = result_with([])
    before = _build_cache_key(result, "template")
    monkeypatch.setattr(pipeline_version, "PIPELINE_VERSION", "new-contract")
    assert _build_cache_key(result, "template") != before
